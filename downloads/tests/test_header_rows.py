"""Header rows are located, not assumed -- and failures say which run they belong to.

Synthetic workbooks only, so these run on a clean checkout. The shapes mirror what
production showed: ML inserts a banner above the Ventas table when it flags the export
as incomplete, and every row below moves down by one.
"""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from mx_sales import __main__ as cli
from mx_sales import pipeline
from mx_sales.config import Settings
from mx_sales.readers.excel import find_anchor_row, locate_header_row
from mx_sales.reports import facturacion, ventas_mx

BANNER = "No pudimos mostrar algunos datos. Intenta generando el Excel de ventas nuevamente."
VENTAS_NAME = "20260901_Ventas_MX_Mercado_Libre_y_Mercado_Shops_2026-09-01_10-00hs_123456789.xlsx"

#: (group label or None, header) per column. A group label sits only above its first column.
VENTAS_COLUMNS = [
    ("Ventas", "# de venta"), (None, "Fecha de venta"), (None, "Estado"), (None, "Unidades"),
    (None, "Ingresos por productos (MXN)"), (None, "Cargo por venta e impuestos (MXN)"),
    (None, "Ingresos por envío (MXN)"), (None, "Costos de envío (MXN)"),
    (None, "Descuentos y bonificaciones"), (None, "Anulaciones y reembolsos (MXN)"),
    (None, "Total (MXN)"), ("Publicaciones", "SKU"), (None, "Variante"),
    ("Facturación al comprador", "Factura adjunta"), (None, "Tipo y número de documento"),
    (None, "CFDI"),
    ("Envíos", "Fecha en camino"), (None, "Fecha entregado"),
    ("Devoluciones", "Unidades"), (None, "Fecha en camino"), (None, "Fecha entregado"),
    (None, "Fecha de revisión"),
]
VENTAS_ROWS = [
    ["2000000000000001", "3 de agosto de 2026 09:24 hs.", "Entregado", 1,
     100, -24, 0, -10, " ", 0, 66, "SKU-A", " ",
     "Factura no adjunta", "RFC: XAXX010101000", "S01 Sin efectos fiscales.",
     "4 de agosto", "6 de agosto", " ", " ", " ", " "],
    ["2000000000000002", "5 de agosto de 2026 10:00 hs.", "Entregado", 2,
     200, -48, 0, -20, " ", 0, 132, "SKU-B", " ",
     "Factura no adjunta", "RFC: XAXX010101000", "S01 Sin efectos fiscales.",
     "6 de agosto", "8 de agosto", " ", " ", " ", " "],
]


def write_ventas(path: Path, *, banner_rows: int = 0, header: bool = True) -> Path:
    """A Ventas workbook with the table at its usual place, pushed down by ``banner_rows``."""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = ventas_mx.SHEET_NAME
    ws["A1"] = "Ventas"
    ws["A2"] = "Estado de tus ventas al 1 de septiembre de 2026"
    for extra in range(banner_rows):
        ws.cell(4 + extra, 1, BANNER)
    group = ventas_mx.GROUP_ROW + 1 + banner_rows  # openpyxl rows are 1-indexed
    if header:
        for column, (label, name) in enumerate(VENTAS_COLUMNS, 1):
            if label:
                ws.cell(group, column, label)
            ws.cell(group + 1, column, name)
    for offset, row in enumerate(VENTAS_ROWS, group + 2):
        for column, value in enumerate(row, 1):
            ws.cell(offset, column, value)
    wb.save(path)
    return path


FACTURACION_HEAD = ["N° de factura fiscal", "Fecha del cargo", "Número del cargo", "Detalle",
                    "Valor del cargo", "Número de venta"]


def write_facturacion(path: Path, *, banner_rows: int = 0) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = facturacion.SHEET_NAME
    ws["A1"] = "Reporte de facturación"
    for extra in range(banner_rows):
        ws.cell(3 + extra, 1, "Aviso: este reporte puede tardar en actualizarse.")
    head = facturacion.HEADER_ROW + 1 + banner_rows
    for column, name in enumerate(FACTURACION_HEAD, 1):
        ws.cell(head, column, name)
    ws.append(["F1", "2026-08-03", "9001", "Cargo por venta", 15, "2000000000000001"])
    ws.append(["F1", "2026-08-03", "9002", "Cargo por envíos de Mercado Libre", 10, "2000000000000001"])
    wb.save(path)
    return path


class TestVentasHeader:
    def test_normal_layout(self, tmp_path):
        raw = ventas_mx.load(write_ventas(tmp_path / VENTAS_NAME))
        clean = ventas_mx.transform(raw)

        assert list(clean["order_id"]) == ["2000000000000001", "2000000000000002"]
        assert list(clean["sale_month"]) == ["2026-08", "2026-08"]
        assert clean["total_mxn"].sum() == 198
        assert list(clean["net_margin_mxn"]) == [66, 132]

    @pytest.mark.parametrize("banner_rows", [1, 2])
    def test_banner_pushes_the_table_down(self, tmp_path, banner_rows):
        # The production failure: with the rows fixed, the group row was read as the
        # header and transform raised "expected columns missing ... ventas__de_venta".
        plain = ventas_mx.load(write_ventas(tmp_path / "plain" / VENTAS_NAME))
        shifted = ventas_mx.load(
            write_ventas(tmp_path / "shifted" / VENTAS_NAME, banner_rows=banner_rows)
        )

        assert list(shifted.columns) == list(plain.columns)
        pd.testing.assert_frame_equal(ventas_mx.transform(shifted), ventas_mx.transform(plain))

    def test_group_row_moves_with_the_header(self, tmp_path):
        raw = ventas_mx.load(write_ventas(tmp_path / VENTAS_NAME, banner_rows=1))

        # "Unidades" under Ventas and under Devoluciones stay apart only if the group
        # row was taken from the row above the header that was actually found.
        assert {"ventas__unidades", "devoluciones__unidades"} <= set(raw.columns)
        assert "publicaciones__sku" in raw.columns

    def test_no_recognisable_header_fails_clearly(self, tmp_path):
        path = write_ventas(tmp_path / VENTAS_NAME, header=False)

        with pytest.raises(ValueError) as caught:
            ventas_mx.load(path)

        message = str(caught.value)
        assert message.startswith("ventas_mx: no usable header")
        assert "# de venta" in message and "Fecha de venta" in message

    def test_empty_sheet_fails_clearly(self, tmp_path):
        path = tmp_path / VENTAS_NAME
        wb = Workbook()
        wb.active.title = ventas_mx.SHEET_NAME
        wb.active["A1"] = "Ventas"
        wb.save(path)

        with pytest.raises(ValueError, match="no usable header"):
            ventas_mx.load(path)

    def test_header_in_another_language_fails_clearly(self, tmp_path):
        path = write_ventas(tmp_path / VENTAS_NAME)
        wb = Workbook()
        ws = wb.active
        ws.title = ventas_mx.SHEET_NAME
        ws.cell(ventas_mx.GROUP_ROW + 1, 1, "Sales")
        for column, name in enumerate(["Order #", "Sale date", "Status"], 1):
            ws.cell(ventas_mx.HEADER_ROW + 1, column, name)
        ws.append(["1", "August 3, 2026", "Delivered"])
        wb.save(path)

        with pytest.raises(ValueError) as caught:
            ventas_mx.load(path)
        # What was read instead is named, so the cause is visible without opening the file.
        assert "sales__order" in str(caught.value)


class TestAnchorRow:
    def test_matches_regardless_of_accents_case_and_padding(self):
        frame = pd.DataFrame([["Título", None], [" detalle ", "VALOR DEL CARGO"], ["x", 1]])
        assert find_anchor_row(frame, ("Detalle", "Valor del cargo")) == 1

    def test_needs_every_anchor_on_the_same_row(self):
        frame = pd.DataFrame([["Detalle", None], [None, "Valor del cargo"]])
        assert find_anchor_row(frame, ("Detalle", "Valor del cargo")) is None

    def test_stops_at_the_search_window(self):
        frame = pd.DataFrame([[None, None]] * 5 + [["a", "b"]])
        assert find_anchor_row(frame, ("a", "b"), search_rows=5) is None
        assert find_anchor_row(frame, ("a", "b")) == 5


class TestFacturacionHeader:
    def test_normal_layout(self, tmp_path):
        clean = facturacion.transform(facturacion.load(write_facturacion(tmp_path / "f.xlsx")))
        assert list(clean["charge_amount_mxn"]) == [15, 10]

    def test_inserted_row_is_followed(self, tmp_path):
        plain = facturacion.load(write_facturacion(tmp_path / "plain.xlsx"))
        shifted = facturacion.load(write_facturacion(tmp_path / "shifted.xlsx", banner_rows=1))
        pd.testing.assert_frame_equal(shifted, plain)

    def test_display_loader_follows_it_too(self, tmp_path):
        path = write_facturacion(tmp_path / "shifted.xlsx", banner_rows=1)

        assert locate_header_row(
            path, facturacion.SHEET_NAME, facturacion.HEADER_ANCHORS, default=facturacion.HEADER_ROW
        ) == facturacion.HEADER_ROW + 1
        shown = facturacion.load_display(path)
        assert list(shown.columns) == FACTURACION_HEAD
        assert len(shown) == 2

    def test_falls_back_to_the_usual_row_without_anchors(self, tmp_path):
        path = tmp_path / "odd.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = facturacion.SHEET_NAME
        ws.cell(facturacion.HEADER_ROW + 1, 1, "Otra cosa")
        wb.save(path)

        assert locate_header_row(
            path, facturacion.SHEET_NAME, facturacion.HEADER_ANCHORS, default=facturacion.HEADER_ROW
        ) == facturacion.HEADER_ROW


@pytest.fixture
def tree(tmp_path):
    """Two run folders for one store: the older holds a Ventas file that cannot be read."""
    raw = tmp_path / "downloads"
    write_ventas(raw / "STORE_A" / "20260815_090000" / VENTAS_NAME, header=False)
    write_ventas(raw / "STORE_A" / "20261001_090000_m202609" / VENTAS_NAME, banner_rows=1)
    return Settings(raw_dir=raw, processed_dir=tmp_path / "processed",
                    warehouse_path=tmp_path / "w.duckdb")


class TestFailuresByRun:
    def test_run_records_the_relative_path_of_each_failure(self, tree):
        (result,) = pipeline.run("ventas_mx", tree, write=False)

        assert result.files == 1 and result.rows == 2
        assert result.failed_files == [f"STORE_A/20260815_090000/{VENTAS_NAME}"]
        assert len(result.errors) == 1 and result.errors[0].startswith(VENTAS_NAME)

    def test_from_runs(self):
        runs = frozenset({"20261001_090000_m202609"})
        assert pipeline.from_runs("A/20261001_090000_m202609/v.xlsx", runs)
        assert not pipeline.from_runs("A/20260815_090000/v.xlsx", runs)
        # The file's own name never counts as a run folder.
        assert not pipeline.from_runs("A/x/20261001_090000_m202609", runs)

    def build(self, tree, monkeypatch, capsys, current_runs):
        results = pipeline.run("ventas_mx", tree, write=False)
        monkeypatch.setattr(cli.pipeline, "run", lambda report, only=None: results)
        code = cli._build(None, None, warehouse_refresh=False, current_runs=current_runs)
        out = capsys.readouterr().out
        (marker,) = [l for l in out.splitlines() if l.startswith(cli.BUILD_RESULT_MARKER)]
        return code, out, json.loads(marker[len(cli.BUILD_RESULT_MARKER):])

    def test_without_the_flag_every_failure_is_an_error(self, tree, monkeypatch, capsys):
        code, out, detail = self.build(tree, monkeypatch, capsys, None)

        assert code == 1
        assert "  ERROR " + VENTAS_NAME in out
        assert "note:" not in out
        assert [f["current"] for f in detail["failed"]] == [True]
        assert detail["written"] == 1

    def test_older_failure_becomes_a_note(self, tree, monkeypatch, capsys):
        code, out, detail = self.build(tree, monkeypatch, capsys, ["20261001_090000_m202609"])

        # Same exit code as ever: the split travels in the JSON line.
        assert code == 1
        assert "ERROR" not in out
        assert "note: 1 file(s) from earlier runs" in out
        assert f"STORE_A/20260815_090000/{VENTAS_NAME}" in out
        assert "(0 from this run, 1 older)" in out
        assert [(f["file"], f["current"]) for f in detail["failed"]] == [
            (f"STORE_A/20260815_090000/{VENTAS_NAME}", False)
        ]

    def test_failure_in_the_named_run_stays_an_error(self, tree, monkeypatch, capsys):
        code, out, detail = self.build(
            tree, monkeypatch, capsys, ["C:\\dl\\STORE_A\\20260815_090000\\"]
        )

        assert code == 1
        assert "  ERROR " + VENTAS_NAME in out and "note:" not in out
        assert detail["failed"][0]["current"] is True
        assert detail["current_runs"] == ["20260815_090000"]
