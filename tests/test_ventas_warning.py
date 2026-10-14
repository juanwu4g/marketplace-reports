"""A platform warning must not discard otherwise readable Ventas rows."""

import pytest
from openpyxl import Workbook

import pandas as pd

import report_build
import run_reports
from test_reports_build import NOTICE, make_download_folder


@pytest.mark.parametrize("notice,header_row", [
    ("", 6),
    ("No pudimos mostrar algunos datos. Intenta generando el Excel de ventas nuevamente.", 7),
])
def test_load_ventas_keeps_rows_and_warning(tmp_path, notice, header_row):
    path = tmp_path / "Ventas_MX_example.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = report_build.VENTAS_SHEET
    if notice:
        ws["A2"] = notice
    for col, value in enumerate((report_build.V_ORDER_ID, report_build.V_DATE,
                                 report_build.V_STATUS, report_build.V_TOTAL), 1):
        ws.cell(header_row, col, value)
    for col, value in enumerate(("123", "1 de agosto de 2026 10:00 hs.",
                                 "Entregado", 100), 1):
        ws.cell(header_row + 1, col, value)
    wb.save(path)

    selected, ventas = report_build.load_ventas(tmp_path)

    assert selected == str(path)
    assert len(ventas) == 1
    assert ventas.iloc[0]["ym"] == "2026-08"
    assert ventas.attrs["incomplete_notice"] == notice


def _validated(folder):
    """真走一遍 derive -> validate：不是拿一个手写的 check 字典冒充。"""
    p0, p1 = report_build.month_bounds("2026-08")
    S = report_build.derive("TEST", str(folder), p0, p1)
    return S, report_build.validate([S])


def test_validate_flags_banner_export_at_warn_level(tmp_path):
    folder = make_download_folder(tmp_path / "TEST" / "20260901_100000", banner=True)

    S, checks = _validated(folder)

    # 横幅把表头推下去一行，行还是都读到了 —— 数据没有因为警告被丢掉。
    assert S["ventas_warning"] == NOTICE
    assert len(S["ventas"]) == 2
    assert S["m"]["ing"] == 300
    hit = [c for c in checks if c["key"] == "ventas_incomplete"]
    assert len(hit) == 1
    assert hit[0]["ok"] is False and hit[0]["severity"] == "warn"
    assert NOTICE in hit[0]["detail"]
    assert "Ventas_MX" in hit[0]["detail"]          # 说明里带源文件名
    # warn 级：不 strict 时它自己不构成"校验未过"。
    assert not [c for c in checks if not c["ok"] and c["severity"] == "error"]


def test_validate_is_silent_without_the_banner(tmp_path):
    folder = make_download_folder(tmp_path / "TEST" / "20260901_100000")

    S, checks = _validated(folder)

    assert S["ventas_warning"] == ""
    assert not [c for c in checks if c["key"] == "ventas_incomplete"]


def test_banner_does_not_change_any_number(tmp_path):
    plain, _ = _validated(make_download_folder(tmp_path / "a" / "20260901_100000"))
    flagged, _ = _validated(make_download_folder(tmp_path / "b" / "20260901_100000",
                                                 banner=True))

    numbers = {k: v for k, v in plain["m"].items() if isinstance(v, (int, float))}
    assert numbers
    for key, value in numbers.items():
        other = flagged["m"][key]
        assert other == value or (pd.isna(other) and pd.isna(value)), key


def test_strict_refuses_the_banner_export_but_plain_run_writes_it(tmp_path, capsys):
    folder = make_download_folder(tmp_path / "TEST" / "20260901_100000", banner=True)
    argv = ["--month", "2026-08", "--store", "TEST=%s" % folder, "--no-recalc"]

    loose = tmp_path / "loose.xlsx"
    assert report_build.main(argv + ["--out", str(loose)]) == report_build.EXIT_OK
    assert loose.is_file()

    strict = tmp_path / "strict.xlsx"
    code = report_build.main(argv + ["--out", str(strict), "--strict"])
    assert code == report_build.EXIT_STRICT_ABORT
    assert not strict.exists()
    assert "--strict 模式下不生成报表：ventas_incomplete" in capsys.readouterr().out


def test_run_summary_marks_generated_report_as_incomplete(monkeypatch, capsys, tmp_path):
    check = {"key": "ventas_incomplete", "store": "TEST", "name": "Ventas 导出数据不完整",
             "ok": False, "severity": "warn", "detail": "平台提示数据不完整"}
    monkeypatch.setattr(run_reports.store_config, "stores_for_region",
                        lambda region="MX": ["TEST"])
    monkeypatch.setattr(run_reports, "resolve", lambda names, month: ([("TEST", str(tmp_path))], []))
    monkeypatch.setattr(run_reports.financial_report, "has_libreoffice", lambda: False)
    monkeypatch.setattr(run_reports.financial_report, "build", lambda *args, **kwargs: {
        "ok": True, "wrote": True, "out": str(tmp_path / "report.xlsx"), "seconds": 1,
        "data_incomplete": True, "checks": [check], "returncode": 0})
    monkeypatch.setattr("sys.argv", ["run_reports.py", "--month", "2026-08",
                                    "--no-combined", "--out-dir", str(tmp_path)])

    assert run_reports.main() == 1
    output = capsys.readouterr().out
    assert "数据不全" in output
    assert "Ventas 源数据不完整" in output
