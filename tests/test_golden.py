# -*- coding: utf-8 -*-
"""Golden files: real store-months rebuilt with the current code must equal the
accepted workbook cell by cell (values, formulas as text, number formats).

The data are real exports with customer details, so they live OUTSIDE the repo:

    $REPORT_GOLDEN_ROOT (default ~/report-golden)/
        mx/BOCINA_TA02/                 the store's download folder for 2026-09
        mx/BOCINA_TA02.baseline.json    the accepted workbook, dumped
        cl/DAB_CL/ … cl/DAB_CL.baseline.json
        …

Missing data -> the case is SKIPPED with a message (other machines, CI). Run
`pytest tests/test_golden.py -rs` to see which cases ran.

A diff here means a number or a cell changed. If finance accepted the change,
regenerate that baseline deliberately:

    python tests/test_golden.py --rebaseline cl/DAB_CL

The build date is the only cell that changes with the day; it is pinned.

MercadoLibre's bill basis (platforms/mercadolibre/bill_report) has its own
baselines next to the order-based ones, built from the same folders:

        mx/BOCINA_TA02.bill.baseline.json        store workbook, bill basis
        cl.bill-combined.baseline.json           a country's 全店汇总, bill basis

    python tests/test_golden.py --rebaseline-bill mx/BOCINA_TA02
    python tests/test_golden.py --rebaseline-bill-combined cl
"""
import datetime
import json
import os
import sys

import pytest
from openpyxl import load_workbook

if __name__ == "__main__":                              # run as a script from tests/
    sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from platforms.mercadolibre import report_build

ROOT = os.environ.get("REPORT_GOLDEN_ROOT", os.path.expanduser("~/report-golden"))
BUILD_DATE = "2026-10-14"

# case id -> (store name as in config.json, site, month)          MercadoLibre
# case id -> (store name, "<platform code>", month)                 other platforms
# (other platforms build with platforms.<pkg>.report_build.build(pairs, month, out,
#  build_date=...); the case folder is the store's download folder for the month)
CASES = {
    "cl-fa/DAB_CL-Fala": ("DAB_CL-Fala (Coco)", "FA", "2026-09"),
    "mx-az/SM-Amazon": ("SM-Amazon", "AZ", "2026-09"),
    "mx-lv/Liverpool Report": ("Liverpool Report", "LV", "2026-09"),
    "mx/BOCINA_TA02": ("BOCINA_TA02", "mx", "2026-09"),
    "cl/DAB_CL": ("DAB_CL (Coco) 子账号", "cl", "2026-09"),
    "cl/Tecnohub_CL": ("Tecnohub_CL (个人店)", "cl", "2026-09"),
    "pe/IDG_PE": ("IDG_PE", "pe", "2026-09"),
}


# bill basis: the MercadoLibre cases above, and a country's combined workbook
BILL_CASES = ("mx/BOCINA_TA02", "cl/DAB_CL", "cl/Tecnohub_CL", "pe/IDG_PE")
BILL_COMBINED = {"cl": ("cl/DAB_CL", "cl/Tecnohub_CL")}
BILL_DATE = BUILD_DATE + " 00:00"


def dump(path):
    """{sheet: {cell: value}, sheet#fmt: {cell: number format}, _sheets: [...]}"""
    wb = load_workbook(path)
    out = {"_sheets": wb.sheetnames}
    for ws in wb.worksheets:
        cells, fm = {}, {}
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is not None:
                    if isinstance(v, (datetime.date, datetime.datetime)):
                        v = v.isoformat()
                    elif not isinstance(v, (int, float, str, bool)):
                        v = str(v)
                    cells[c.coordinate] = v
                if c.number_format != "General":
                    fm[c.coordinate] = c.number_format
        out[ws.title] = cells
        out[ws.title + "#fmt"] = fm
    return json.loads(json.dumps(out))


def build(case, out):
    store, site, month = CASES[case]
    if site.isupper():                                   # another platform's report
        import importlib
        from core import markets
        pkg = markets.platform(site).package
        rb = importlib.import_module(pkg + ".report_build")
        r = rb.build([(store, os.path.join(ROOT, case))], month, out, build_date=BUILD_DATE)
        return report_build.EXIT_OK if r["outcome"] == "ok" else report_build.EXIT_CHECKS
    # start from another site on purpose: the build must switch by the folder
    report_build.set_site("cl" if site != "cl" else "mx")
    rc = report_build.main(["--month", month, "--store", "%s=%s" % (store, os.path.join(ROOT, case)),
                            "--out", out, "--no-recalc", "--build-date", BUILD_DATE])
    report_build.set_site("mx")
    return rc


def build_bill(case, out_dir):
    from platforms.mercadolibre.bill_report import build as BB
    store, site, month = CASES[case]
    return BB.run(store, os.path.join(ROOT, case), month, site, out_dir=out_dir, recalc=False,
                  build_date=BILL_DATE)["out"]


def build_bill_combined(region, out_dir):
    from platforms.mercadolibre.bill_report import build as BB
    prepared = []
    for case in BILL_COMBINED[region]:
        store, site, month = CASES[case]
        prepared.append(BB.prepare(store, os.path.join(ROOT, case), month, site))
    return BB.write_combined(prepared, month, region.upper(), out_dir=out_dir, recalc=False,
                             build_date=BILL_DATE)["out"]


def diff(want, got):
    out = []
    for key in sorted(set(want) | set(got)):
        a, b = want.get(key), got.get(key)
        if isinstance(a, dict) and isinstance(b, dict):
            for c in sorted(set(a) | set(b)):
                if a.get(c) != b.get(c):
                    out.append("%s!%s: %r -> %r" % (key, c, a.get(c), b.get(c)))
        elif a != b:
            out.append("%s: %r -> %r" % (key, a, b))
    return out


@pytest.mark.parametrize("case", sorted(CASES))
def test_golden_workbook_identical(case, tmp_path):
    folder, baseline = os.path.join(ROOT, case), os.path.join(ROOT, case + ".baseline.json")
    if not (os.path.isdir(folder) and os.path.isfile(baseline)):
        pytest.skip("no golden data for %s under %s" % (case, ROOT))
    out = str(tmp_path / "golden.xlsx")
    rc = build(case, out)
    assert rc in (report_build.EXIT_OK, report_build.EXIT_CHECKS)
    with open(baseline, encoding="utf-8") as fh:
        want = json.load(fh)
    d = diff(want, dump(out))
    assert not d, "%s changed in %d cells, first ones:\n%s" % (case, len(d), "\n".join(d[:20]))


def check_against(baseline, out, what):
    with open(baseline, encoding="utf-8") as fh:
        want = json.load(fh)
    d = diff(want, dump(out))
    assert not d, "%s changed in %d cells, first ones:\n%s" % (what, len(d), "\n".join(d[:20]))


@pytest.mark.parametrize("case", BILL_CASES)
def test_golden_bill_workbook_identical(case, tmp_path):
    folder, baseline = os.path.join(ROOT, case), os.path.join(ROOT, case + ".bill.baseline.json")
    if not (os.path.isdir(folder) and os.path.isfile(baseline)):
        pytest.skip("no golden bill-basis data for %s under %s" % (case, ROOT))
    check_against(baseline, build_bill(case, str(tmp_path)), case + " (bill basis)")


@pytest.mark.parametrize("region", sorted(BILL_COMBINED))
def test_golden_bill_combined_identical(region, tmp_path):
    baseline = os.path.join(ROOT, region + ".bill-combined.baseline.json")
    if not (all(os.path.isdir(os.path.join(ROOT, c)) for c in BILL_COMBINED[region])
            and os.path.isfile(baseline)):
        pytest.skip("no golden bill-basis combined data for %s under %s" % (region, ROOT))
    check_against(baseline, build_bill_combined(region, str(tmp_path)), region + " combined (bill basis)")


def _write_baseline(path, out):
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(dump(out), fh, ensure_ascii=False)
    print("baseline written:", path)


if __name__ == "__main__":
    import tempfile
    if len(sys.argv) == 3 and sys.argv[1] == "--rebaseline-bill" and sys.argv[2] in BILL_CASES:
        _write_baseline(os.path.join(ROOT, sys.argv[2] + ".bill.baseline.json"),
                        build_bill(sys.argv[2], tempfile.mkdtemp()))
    elif len(sys.argv) == 3 and sys.argv[1] == "--rebaseline-bill-combined" and sys.argv[2] in BILL_COMBINED:
        _write_baseline(os.path.join(ROOT, sys.argv[2] + ".bill-combined.baseline.json"),
                        build_bill_combined(sys.argv[2], tempfile.mkdtemp()))
    elif len(sys.argv) == 3 and sys.argv[1] == "--rebaseline" and sys.argv[2] in CASES:
        case = sys.argv[2]
        out = os.path.join(tempfile.mkdtemp(), "golden.xlsx")
        print("build ->", build(case, out))
        with open(os.path.join(ROOT, case + ".baseline.json"), "w", encoding="utf-8") as fh:
            json.dump(dump(out), fh, ensure_ascii=False)
        print("baseline written for", case)
    else:
        print("usage: python tests/test_golden.py --rebaseline <%s>\n"
              "       python tests/test_golden.py --rebaseline-bill <%s>\n"
              "       python tests/test_golden.py --rebaseline-bill-combined <%s>"
              % ("|".join(sorted(CASES)), "|".join(BILL_CASES), "|".join(sorted(BILL_COMBINED))))
