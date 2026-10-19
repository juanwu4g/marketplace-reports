# -*- coding: utf-8 -*-
"""The shared SKU cost table (core/sku_costs.py), its effect on the Falabella and
Amazon workbooks, and tools/sku_cost_template. Synthetic data only."""
import datetime
import io
import json
import os
import shutil

import pytest
from openpyxl import Workbook, load_workbook

import test_az_report as AZT
import test_fa_report as FAT
from core import report_kit as K
from core import sku_costs as SC
from core import store_config
from platforms.amazon import report_build as AZ
from platforms.falabella import report_build as FA

MONTH = "2026-09"
END = datetime.date(2026, 9, 30)
HEAD = ["平台", "国家", "SKU", "单位成本", "币种", "生效日期", "备注"]


def xlsx(path, rows, head=HEAD, title_rows=()):
    wb = Workbook()
    ws = wb.active
    for t in title_rows:
        ws.append(t)
    ws.append(head)
    for r in rows:
        ws.append(list(r))
    wb.save(str(path))
    return str(path)


def configure(monkeypatch, tmp_path, cost_path):
    cfg = tmp_path / "cfg_with_costs.json"
    cfg.write_text(json.dumps({"stores": {"default": "X", "allowed": []},
                               "reports": {"sku_costs": cost_path}}), encoding="utf-8")
    monkeypatch.setattr(store_config, "CONFIG_PATH", str(cfg))


def by_key(checks):
    out = {}
    for c in checks:
        out.setdefault(c["key"], c)
    return out


def items(*rows):
    return [SC.Item(*r) for r in rows]


# ------------------------------------------------------------------ reading the file

def test_xlsx_with_english_aliases_numbers_and_extra_columns(tmp_path):
    p = xlsx(tmp_path / "c.xlsx", [
        ["fa", "cl", "  sku1 ", 1000, "clp", None, "x", "参考"],
        [None, None, 12345.0, "1,234.5", "CLP", "2026-01", None, None],
        ["FA", "CL", "SKU-EMPTY", None, "CLP", None, "还没填", None],       # template row: skipped
        [None, None, None, None, None, None, None, None],                    # blank row
    ], head=["Platform", "Country", "Seller SKU", "Unit cost (CLP)", "Currency", "Effective date", "Notes",
             "商品名称（参考）"], title_rows=[["成本表"]])
    T = SC.load(p)
    assert T.usable and not T.problems and not T.duplicates
    assert [(r.platform, r.country, r.key, r.cost, r.currency) for r in T.rows] == [
        ("FA", "CL", "SKU1", 1000.0, "CLP"), ("", "", "12345", 1234.5, "CLP")]
    assert T.rows[1].effective == datetime.date(2026, 1, 1)
    assert T.rows[0].line == 3                          # the sheet row, for the operator


def test_csv_utf8_bom_semicolons(tmp_path):
    p = tmp_path / "c.csv"
    p.write_bytes(u"﻿平台;国家;SKU;单位成本;币种;生效日期;备注\nAZ;MX;A-1;12.5;MXN;2026-09-01;\n"
                  u";;b-2;3;MXN;;\n".encode("utf-8"))
    T = SC.load(str(p))
    assert T.usable
    assert [(r.platform, r.key, r.cost, r.effective) for r in T.rows] == [
        ("AZ", "A-1", 12.5, datetime.date(2026, 9, 1)), ("", "B-2", 3.0, None)]
    assert T.rows[0].line == 2


def test_missing_unreadable_and_headerless_files(tmp_path):
    assert "找不到" in SC.load(str(tmp_path / "nope.xlsx")).error
    bad = tmp_path / "c.xlsx"
    bad.write_bytes(b"not a zip")
    assert "读不出来" in SC.load(str(bad)).error
    ods = tmp_path / "x.ods"
    ods.write_text("x")
    assert "只认" in SC.load(str(ods)).error
    p = xlsx(tmp_path / "h.xlsx", [["A", 1]], head=["物品", "价格"])
    assert "表头" in SC.load(p).error


def test_effective_dates_latest_on_or_before_month_end(tmp_path):
    p = xlsx(tmp_path / "c.xlsx", [
        ["FA", "CL", "S", 100, "CLP", None, "最早"],
        ["FA", "CL", "S", 200, "CLP", datetime.date(2026, 3, 1), ""],
        ["FA", "CL", "S", 300, "CLP", "2026-09-30", "月底当天生效"],
        ["FA", "CL", "S", 400, "CLP", "2026-10-01", "下月"],
    ])
    T = SC.load(p)
    assert SC.lookup(T, "FA", "CL", "s", END)[1].cost == 300
    assert SC.lookup(T, "FA", "CL", "s", datetime.date(2026, 8, 31))[1].cost == 200
    assert SC.lookup(T, "FA", "CL", "s", datetime.date(2025, 1, 31))[1].cost == 100
    assert SC.lookup(T, "FA", "CL", "s", datetime.date(2026, 12, 31))[1].cost == 400


def test_scoping_platform_country_then_broader(tmp_path):
    p = xlsx(tmp_path / "c.xlsx", [
        ["", "", "X", 1, "CLP"], ["", "CL", "X", 2, "CLP"], ["FA", "", "X", 3, "CLP"],
        ["FA", "CL", "X", 4, "CLP"], ["AZ", "MX", "Y", 5, "MXN"], ["", "", "Y", 6, "CLP"],
    ])
    T = SC.load(p)
    cost = lambda pl, c, s: (SC.lookup(T, pl, c, s, END)[1] or SC.Row(0, "", "", "", "", None, "", None)).cost
    assert cost("FA", "CL", "x") == 4
    assert cost("FA", "PE", "x") == 3
    assert cost("ML", "CL", "x") == 2
    assert cost("ML", "MX", "x") == 1
    assert cost("AZ", "MX", "y") == 5
    assert cost("FA", "CL", "y") == 6
    assert SC.lookup(T, "FA", "CL", "nothing", END)[0] == SC.NO_COST
    # a future row at the specific level does not hide the broader one
    T2 = SC.load(xlsx(tmp_path / "d.xlsx", [["FA", "CL", "Z", 9, "CLP", "2027-01-01"], ["", "", "Z", 7, "CLP"]]))
    assert SC.lookup(T2, "FA", "CL", "Z", END)[1].cost == 7


def test_configured_path(tmp_path, monkeypatch):
    assert SC.configured_path() is None                 # the neutral test config has no key
    assert SC.load_configured() is None
    configure(monkeypatch, tmp_path, "data/costs.xlsx")
    from core.paths import ROOT
    assert SC.configured_path() == os.path.join(ROOT, "data", "costs.xlsx")
    configure(monkeypatch, tmp_path, str(tmp_path / "abs.xlsx"))
    assert SC.configured_path() == str(tmp_path / "abs.xlsx")
    assert SC.load_configured().error                   # configured, not there
    monkeypatch.setattr(store_config, "CONFIG_PATH", str(tmp_path / "no_config.json"))
    assert SC.configured_path() is None


# ------------------------------------------------------------------ validation

def test_validation_checks(tmp_path):
    p = xlsx(tmp_path / "c.xlsx", [
        ["FA", "CL", "A", 100, "CLP"],
        ["FA", "CL", "B", -5, "CLP"],                   # negative
        ["FA", "CL", "C", "abc", "CLP"],                # not a number
        ["FA", "", "D", 10, "CLP", "2026-09-01"],
        ["FA", "", "d", 11, "CLP", "2026-09-01"],       # duplicate key and date
        ["FA", "CL", "E", 10, "USD"],                   # currency of another report
        ["ZZ", "", "F", 10, "CLP"],                     # unknown platform: shown everywhere
        ["AZ", "MX", "G", -1, "MXN"],                   # another platform's bad row: not shown here
        ["FA", "CL", "H", 1, "CLP", "someday"],         # bad date
    ])
    T = SC.load(p)
    res = SC.apply(T, "FA", "CL", "CLP", MONTH, items(
        ("A", "a", 3, 1, 3000.0), ("B", "b", 1, 0, 100.0), ("D", "d", 2, 0, 900.0),
        ("E", "e", 1, 0, 500.0), ("Z", "z", 0, 0, 0.0)))
    C = by_key(SC.checks(T, [("店", res)], "FA", "CL"))
    assert C["sku_costs_file"]["ok"]
    rows = C["sku_costs_rows"]
    assert not rows["ok"] and rows["severity"] == "error"
    for frag in ("第 3 行", "负数", "第 4 行", "不是数字", "ZZ", "someday"):
        assert frag in rows["detail"], frag
    assert "第 9 行" not in rows["detail"]              # AZ's row is AZ's business
    dup = C["sku_costs_dup"]
    assert not dup["ok"] and dup["severity"] == "error" and "第 5、6 行" in dup["detail"]
    assert res.line("D").status == SC.CONFLICT and res.line("D").cost == 0
    cur = C["sku_costs_currency"]
    assert not cur["ok"] and cur["severity"] == "error" and "E（USD" in cur["detail"]
    assert res.line("E").status == SC.CURRENCY and res.line("E").unit_cost is None
    cov = C["sku_costs_coverage"]
    assert not cov["ok"] and cov["severity"] == "warn"
    # uncovered by sales, the SKU with nothing sold is not counted
    assert cov["detail"].index("D（2 件") < cov["detail"].index("E（1 件") < cov["detail"].index("B（1 件")
    assert "Z（" not in cov["detail"]
    assert res.cogs == 200.0                            # A: 100 × (3 − 1)
    assert "SKU 1/4" in cov["detail"]
    # missing file: one error, nothing applied
    M = SC.load(str(tmp_path / "gone.xlsx"))
    res = SC.apply(M, "FA", "CL", "CLP", MONTH, items(("A", "a", 3, 0, 1.0)))
    C = SC.checks(M, [("店", res)], "FA", "CL")
    assert [(c["key"], c["ok"], c["severity"]) for c in C] == [("sku_costs_file", False, "error")]
    assert res.cogs == 0
    # not configured: no checks at all
    assert SC.checks(None, [("店", None)], "FA", "CL") == []
    assert SC.apply(None, "FA", "CL", "CLP", MONTH, []) is None


def test_full_coverage_is_an_ok_line(tmp_path):
    T = SC.load(xlsx(tmp_path / "c.xlsx", [["", "", "A", 2.5, "MXN"]]))
    res = SC.apply(T, "AZ", "MX", "MXN", MONTH, items(("a", "", 4, 1, 100.0)))
    C = by_key(SC.checks(T, [("店", res)], "AZ", "MX"))
    assert all(c["ok"] for c in C.values())
    assert "SKU 1/1、件数 4/4、销售额 100.0% 已有成本" in C["sku_costs_coverage"]["detail"]
    assert res.cogs == 7.5


# ------------------------------------------------------------------ Falabella workbook

@pytest.fixture
def fa_folder(tmp_path):
    d = tmp_path / "downloads" / FAT.STORE / "20261016_120000_m202609"
    d.mkdir(parents=True)
    FAT.write_tx(str(d / "NewReportTransaction_2026-10-16T05_29_08.xlsx"), FAT.TX)
    FAT.write_order_pieces(str(d))
    FAT.write_invoice(str(d / "InvoiceReport_FACL-1-2026-08-21-2026-09-20_2026-09-24T08_18_05.xlsx"), FAT.TX)
    FAT.write_liq(str(d / "S_FSC_LIQ_BOLETA_NCR_REPORT_2026-09-01_2026-09-30.xlsx"))
    FAT.write_stock(str(d / "Product_details_1.xlsx"))
    (d / "run_meta.json").write_text(json.dumps({"platform": "FA", "month": MONTH, "missing": []}),
                                     encoding="utf-8")
    return str(d)


def build_fa(folder, tmp_path):
    out = str(tmp_path / "out" / "fa.xlsx")
    return FA.build([(FAT.STORE, folder)], MONTH, out), out


def pl_rows(ws):
    return dict((ws.cell(row=i, column=1).value, i) for i in range(1, ws.max_row + 1))


def sku_rows(ws):
    return dict((ws.cell(row=i, column=2).value, i) for i in range(5, ws.max_row + 1)
                if ws.cell(row=i, column=2).value)


def test_fa_without_a_table_is_as_before(fa_folder, tmp_path):
    r, out = build_fa(fa_folder, tmp_path)
    assert not any(c["key"].startswith("sku_costs") for c in r["checks"])
    wb = load_workbook(out)
    pl = wb[K.SHEET_NAMES[1]]
    at = pl_rows(pl)["减：商品采购成本"]
    assert pl.cell(row=at, column=3).value == 0
    assert pl.cell(row=at, column=4).value.startswith("⚠ 平台数据不含此项")
    sk = wb[K.SHEET_NAMES[6]]
    assert sk.max_column == 17 and sk.cell(row=4, column=15).value == "单位成本（录入）"
    assert sk.cell(row=sku_rows(sk)["SKU1"], column=16).value == "=-O%d*D%d" % ((sku_rows(sk)["SKU1"],) * 2)


def test_fa_full_coverage(fa_folder, tmp_path, monkeypatch):
    configure(monkeypatch, tmp_path, xlsx(tmp_path / "c.xlsx", [
        ["FA", "CL", "sku1", 4000, "CLP"], ["FA", "", "SKU2", 9000, "CLP"]]))
    r, out = build_fa(fa_folder, tmp_path)
    C = by_key(r["checks"])
    assert C["sku_costs_coverage"]["ok"] and C["sku_costs_currency"]["ok"]
    # SKU1: 2 delivered, 0 returned -> 8,000; SKU2: 1 delivered, 1 returned -> 0
    assert r["metrics"][FAT.STORE]["cogs"] == 8000
    assert r["metrics"][FAT.STORE]["profit"] == r["metrics"][FAT.STORE]["contrib"] - 8000
    wb = load_workbook(out)
    pl = wb[K.SHEET_NAMES[1]]
    rows = pl_rows(pl)
    c = pl.cell(row=rows["减：商品采购成本"], column=3)
    assert c.value == 8000 and c.fill.fgColor.rgb.endswith(K.C_INPUT)       # still an input cell
    assert "c.xlsx" in pl.cell(row=rows["减：商品采购成本"], column=4).value
    assert "100.0% 销售额已有成本" in pl.cell(row=rows["减：商品采购成本"], column=4).value
    assert pl.cell(row=rows["税前净利润"], column=3).value == "=C%d-C%d" % (
        rows["经营贡献毛利（未扣商品成本）"], rows["减：商品采购成本"])
    sk = wb[K.SHEET_NAMES[6]]
    s = sku_rows(sk)
    assert sk.cell(row=4, column=18).value == "成本来源"
    assert sk.cell(row=s["SKU1"], column=15).value == 4000
    assert sk.cell(row=s["SKU1"], column=16).value == "=-O{0}*(D{0}-E{0})".format(s["SKU1"])
    assert sk.cell(row=s["SKU1"], column=18).value.startswith("成本表第 2 行（FA/CL")
    assert sk.cell(row=s[FA.NO_SKU], column=15).value is None
    assert "不计商品成本" in sk.cell(row=s[FA.NO_SKU], column=18).value
    rd = wb[K.SHEET_NAMES[0]]
    labels = dict((rd.cell(row=i, column=2).value, rd.cell(row=i, column=3).value)
                  for i in range(1, rd.max_row + 1))
    assert labels["商品成本表"].startswith("c.xlsx")
    assert "退货件数" in labels["商品成本"] and "到岸成本" in labels["商品成本"]


def test_fa_partial_coverage_and_wrong_currency(fa_folder, tmp_path, monkeypatch):
    configure(monkeypatch, tmp_path, xlsx(tmp_path / "c.xlsx", [["FA", "CL", "SKU2", 9000, "CLP"],
                                                               ["", "", "SKU1", 3, "USD"]]))
    r, out = build_fa(fa_folder, tmp_path)
    C = by_key(r["checks"])
    assert not C["sku_costs_currency"]["ok"] and "SKU1（USD" in C["sku_costs_currency"]["detail"]
    assert r["outcome"] == "checks_failed"
    assert not C["sku_costs_coverage"]["ok"] and "SKU1（2 件" in C["sku_costs_coverage"]["detail"]
    sk = load_workbook(out)[K.SHEET_NAMES[6]]
    s = sku_rows(sk)["SKU1"]
    assert sk.cell(row=s, column=15).value is None
    assert sk.cell(row=s, column=15).fill.fgColor.rgb.endswith(K.C_INPUT)
    assert sk.cell(row=s, column=18).value.startswith("未录入成本（成本表币种 USD")


def test_fa_no_coverage_and_missing_file(fa_folder, tmp_path, monkeypatch):
    configure(monkeypatch, tmp_path, xlsx(tmp_path / "c.xlsx", [["FA", "CL", "OTHER", 1, "CLP"]]))
    r, out = build_fa(fa_folder, tmp_path)
    assert r["metrics"][FAT.STORE]["cogs"] == 0
    assert "SKU 0/2" in by_key(r["checks"])["sku_costs_coverage"]["detail"]
    sk = load_workbook(out)[K.SHEET_NAMES[6]]
    assert sk.cell(row=sku_rows(sk)["SKU1"], column=18).value == "未录入成本"
    configure(monkeypatch, tmp_path, str(tmp_path / "gone.xlsx"))
    r, out = build_fa(fa_folder, tmp_path)
    assert r["outcome"] == "checks_failed"
    assert not by_key(r["checks"])["sku_costs_file"]["ok"]
    pl = load_workbook(out)[K.SHEET_NAMES[1]]
    at = pl_rows(pl)["减：商品采购成本"]
    assert pl.cell(row=at, column=3).value == 0 and "读不出来" in pl.cell(row=at, column=4).value


# ------------------------------------------------------------------ Amazon workbook

@pytest.fixture
def az_folder(tmp_path, monkeypatch):
    def fake(path):
        with io.open(path, encoding="utf-8") as fh:
            return fh.read(), "test"
    monkeypatch.setattr(AZ, "pdf_text", fake)
    return AZT.make_folder(tmp_path)


def build_az(folder, tmp_path, stores=None):
    out = str(tmp_path / "out" / "az.xlsx")
    return AZ.build(stores or [(AZT.STORE, folder)], MONTH, out), out


def test_az_full_partial_and_none(az_folder, tmp_path, monkeypatch):
    r, out = build_az(az_folder, tmp_path)
    assert not any(c["key"].startswith("sku_costs") for c in r["checks"])
    assert load_workbook(out)[K.SHEET_NAMES[6]].max_column == 17
    # SKU-A: 1 sold, 1 refunded -> 0; SKU-B: 2 sold -> 2 × 120.55
    configure(monkeypatch, tmp_path, xlsx(tmp_path / "c.xlsx", [["AZ", "MX", "SKU-A", 99, "MXN"],
                                                               ["", "MX", "sku-b", 120.55, "MXN"]]))
    r, out = build_az(az_folder, tmp_path)
    assert r["outcome"] == "ok"
    assert r["metrics"][AZT.STORE]["cogs"] == 241.10
    wb = load_workbook(out)
    pl = wb[K.SHEET_NAMES[1]]
    rows = pl_rows(pl)
    assert pl.cell(row=rows["减：商品采购成本"], column=3).value == 241.10
    sk = wb[K.SHEET_NAMES[6]]
    s = sku_rows(sk)
    assert sk.cell(row=s["SKU-A"], column=16).value == "=-O{0}*(D{0}-E{0})".format(s["SKU-A"])
    assert (sk.cell(row=s["SKU-A"], column=4).value, sk.cell(row=s["SKU-A"], column=5).value) == (1, 1)
    assert sk.cell(row=s["SKU-B"], column=18).value.startswith("成本表第 3 行（MX）")
    # partial
    configure(monkeypatch, tmp_path, xlsx(tmp_path / "c2.xlsx", [["AZ", "MX", "SKU-A", 99, "MXN"]]))
    r, out = build_az(az_folder, tmp_path)
    cov = by_key(r["checks"])["sku_costs_coverage"]
    assert not cov["ok"] and cov["severity"] == "warn" and "SKU-B（2 件" in cov["detail"]
    assert r["outcome"] == "ok"                         # a warning, not an error
    assert r["metrics"][AZT.STORE]["cogs"] == 0
    # wrong currency
    configure(monkeypatch, tmp_path, xlsx(tmp_path / "c3.xlsx", [["AZ", "", "SKU-B", 5, "CLP"]]))
    r, out = build_az(az_folder, tmp_path)
    assert not by_key(r["checks"])["sku_costs_currency"]["ok"] and r["outcome"] == "checks_failed"


def test_az_two_stores_and_formulas_evaluate(tmp_path, monkeypatch):
    def fake(path):
        with io.open(path, encoding="utf-8") as fh:
            return fh.read(), "test"
    monkeypatch.setattr(AZ, "pdf_text", fake)
    f1 = AZT.make_folder(tmp_path, AZT.STORE)
    f2 = AZT.make_folder(tmp_path, "TEST_AZ_2")
    configure(monkeypatch, tmp_path, xlsx(tmp_path / "c.xlsx", [["AZ", "MX", "SKU-B", 100, "MXN"]]))
    r, out = build_az(None, tmp_path, [(AZT.STORE, f1), ("TEST_AZ_2", f2)])
    wb = load_workbook(out)
    pl = wb[K.SHEET_NAMES[1]]
    at = pl_rows(pl)["减：商品采购成本"]
    assert [pl.cell(row=at, column=j).value for j in (3, 4)] == [200, 200]
    assert pl.cell(row=at, column=5).value == "=SUM(C%d:D%d)" % (at, at)
    if not shutil.which("soffice"):
        pytest.skip("no LibreOffice to evaluate the formulas")
    ok, _msg = K.recalc(out)
    assert ok
    n, n_err, _none, bad = K.scan_formula_errors(out)
    assert n and not n_err, bad
    v = load_workbook(out, data_only=True)[K.SHEET_NAMES[1]]
    rows = pl_rows(v)
    for col in (3, 4, 5):
        contrib = v.cell(row=rows["经营贡献毛利（未扣商品成本）"], column=col).value
        assert v.cell(row=rows["税前净利润"], column=col).value == pytest.approx(
            contrib - v.cell(row=rows["减：商品采购成本"], column=col).value)


# ------------------------------------------------------------------ the template tool

def test_template_merges_and_never_deletes(az_folder, tmp_path, monkeypatch):
    from tools import sku_cost_template as T
    cfg = tmp_path / "cfg_tpl.json"
    cfg.write_text(json.dumps({"stores": {"default": AZT.STORE, "allowed": [AZT.STORE],
                                          "catalog": {AZT.STORE: {"platform": "AZ", "country": "MX"}}}}),
                   encoding="utf-8")
    monkeypatch.setattr(store_config, "CONFIG_PATH", str(cfg))
    base = str(tmp_path / "downloads")
    out = str(tmp_path / "costs.xlsx")
    # existing: a cost for SKU-A with a note and an extra column, an old SKU not sold
    # this month, and an ML row that has nothing to do with Amazon
    xlsx(out, [["AZ", "MX", "sku-a", 50, "MXN", datetime.date(2026, 1, 1), "财务", "张三"],
               ["AZ", "MX", "OLD-1", 7, "MXN", None, None, None],
               ["ML", "MX", "SKU-B", 9, "MXN", None, None, None]], head=HEAD + ["负责人"])
    assert T.main(["--out", out, "--month", MONTH, "--downloads", base]) == 0
    ws = load_workbook(out).worksheets[0]
    head = [c.value for c in ws[1]]
    assert head[:7] == HEAD and head[-1] == "负责人" and "本月件数（参考）" in head
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    col = dict((h, i) for i, h in enumerate(head))
    assert [r[col["SKU"]] for r in rows] == ["sku-a", "OLD-1", "SKU-B", "SKU-B"]
    a = rows[0]
    assert (a[col["单位成本"]], a[col["备注"]], a[col["负责人"]]) == (50, "财务", "张三")
    assert a[col["生效日期"]].date() == datetime.date(2026, 1, 1)
    assert (a[col["本月件数（参考）"]], a[col["店铺（参考）"]]) == (1, AZT.STORE)
    new = rows[3]
    assert (new[col["平台"]], new[col["国家"]], new[col["单位成本"]], new[col["币种"]],
            new[col["本月件数（参考）"]]) == ("AZ", "MX", None, "MXN", 2)
    assert ws.cell(row=5, column=4).fill.fgColor.rgb.endswith(K.C_INPUT)
    # the written file reads back as a cost table: blank costs are "not entered", not errors
    Tb = SC.load(out)
    assert Tb.usable and not Tb.problems and len(Tb.rows) == 3
    # run again: nothing added, nothing lost
    assert T.main(["--out", out, "--month", MONTH, "--downloads", base]) == 0
    assert load_workbook(out).worksheets[0].max_row == 5
    # a generic row covering a SKU stops it from being added
    out2 = str(tmp_path / "costs2.xlsx")
    xlsx(out2, [["", "", "SKU-B", 1, "MXN"]])
    assert T.main(["--out", out2, "--month", MONTH, "--downloads", base, "--market", "AZ-MX"]) == 0
    assert [c.value for c in load_workbook(out2).worksheets[0]["C"]][1:] == ["SKU-B", "SKU-A"]


def test_template_skips_platforms_without_sku_list(tmp_path, monkeypatch, capsys):
    from tools import sku_cost_template as T
    monkeypatch.setattr(T, "platform_modules", lambda code: None)
    assert T.main(["--out", str(tmp_path / "x.xlsx"), "--month", MONTH, "--market", "ML-MX"]) == 2
    assert "跳过" in capsys.readouterr().out
    assert not os.path.exists(str(tmp_path / "x.xlsx"))
