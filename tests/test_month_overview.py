# -*- coding: utf-8 -*-
"""本月总览_<月>.xlsx and the single ！需要处理_<年月>.txt (core/month_overview).

Synthetic results via tests/period_data.py; both layouts; nothing opens a store.
"""
import io
import json
import os
import types

import pytest
from openpyxl import load_workbook

import period_data as P
from core import month_overview as MO
from core import results as R
from core import runlog

MONTH = "2026-09"
YM = "202609"


def _alert(folder):
    return os.path.join(folder, MO.alert_file_name(MONTH))


def _read(folder):
    with io.open(_alert(folder), encoding="utf-8") as f:
        return f.read()


def _sheet_rows(path, name, first_col=1):
    wb = load_workbook(path)
    ws = wb[name]
    rows = []
    for row in ws.iter_rows(min_row=MO.HEADER_ROW + 1, values_only=True):
        if any(v is not None for v in row):
            rows.append(row)
    return wb.sheetnames, rows


def _folder(tmp_path):
    from core import layout
    d = layout.alert_dir(MONTH)
    os.makedirs(d, exist_ok=True)
    return d


def _put_doc(store, built, **kw):
    """P.put but allowing a different currency (to prove subtotals never mix)."""
    cur = kw.pop("currency", None)
    d = P.doc(store, MONTH, built, **kw)
    if cur:
        d["currency"] = cur
    R.write(d)
    return d


BILL = [("revenue", 120.0), ("fees_total", -20.0), ("contrib", 100.0, "f"), ("cogs", 0.0, "in"),
        ("profit", 100.0, "f")]
ORDER = [("net_sales", 110.0), ("order_fees_total", -8.0), ("bill_fees_total", -2.0),
         ("contrib", 100.0, "f"), ("profit", 100.0, "f")]
MAIN = [("net_sales", 1000.0), ("fee_total", -100.0), ("contrib", 900.0, "f"), ("profit", 900.0, "f")]


def _write_run(log_dir, run_id, batch_rows, month=MONTH, region="MX", stores="all"):
    os.makedirs(runlog.runs_dir(log_dir), exist_ok=True)
    batch = os.path.join(log_dir, "batch_%s.json" % run_id)
    with io.open(batch, "w", encoding="utf-8") as f:
        json.dump({"stores": batch_rows}, f, ensure_ascii=False)
    runlog.write_json(runlog.run_paths(run_id, log_dir)["json"],
                      {"run_id": run_id, "month": month, "region": region, "stores": stores,
                       "batch_summary": batch, "status": "done", "exit": 1, "steps": []})
    open(runlog.run_paths(run_id, log_dir)["log"], "w").close()
    return batch


# ------------------------------------------------------------ 总览

@pytest.mark.parametrize("v2", [True, False])
def test_overview_has_a_row_for_every_configured_store_in_both_layouts(tmp_path, monkeypatch, v2):
    P.configure(tmp_path, monkeypatch, v2=v2)
    folder = _folder(tmp_path)
    _put_doc("S1", "20261003_000000", pl=BILL)
    _put_doc("S1", "20261003_000000", basis="order", pl=ORDER)
    _put_doc("F1", "20261003_000000", pl=MAIN)
    got = MO.build(MONTH, folder, say=lambda t: None, log_dir=str(tmp_path / "logs"))
    assert got["error"] is None and os.path.isfile(got["path"])
    assert os.path.basename(got["path"]) == "本月总览_%s.xlsx" % MONTH
    names, rows = _sheet_rows(got["path"], MO.SHEET_OVERVIEW)
    assert names == [MO.SHEET_OVERVIEW, MO.SHEET_ACTIONS, MO.SHEET_HISTORY, MO.SHEET_NOTES]
    stores = [r[2] for r in rows if not str(r[2]).startswith("小计")]
    # every configured store (the ML ones twice: bill then order), platform -> country -> store
    assert stores == ["BOCINA_SM", "BOCINA_SM", "S1", "S1", "S2", "S2", "S3", "S3", "F1", "F2", "A1"]
    by = dict(((r[2], r[3]), r) for r in rows)
    s1 = by[("S1", "账单基准 正式")]
    assert s1[4] == "完整" and s1[5] == "是" and s1[9] == "MXN"
    assert s1[10:14] == (120.0, -20.0, 100.0, 100.0)
    assert by[("S1", "订单基准 参考")][10:14] == (110.0, -10.0, 100.0, 100.0)
    s2 = by[("S2", "账单基准 正式")]
    assert s2[4] == "没有报表" and s2[10:14] == (None, None, None, None)
    assert "本月没有下载目录或没出过报表" in s2[14] and s2[15] == "运营" and "下载并出报表" in s2[16]
    f1 = by[("F1", "—")]
    assert f1[4] == "完整" and f1[9] == "CLP" and f1[10:14] == (1000.0, -100.0, 900.0, 900.0)
    # bill before order within a store; the official basis is marked
    assert rows[2][3] == "账单基准 正式" and rows[3][3] == "订单基准 参考"


def test_subtotals_are_per_platform_country_basis_and_currency(tmp_path, monkeypatch):
    P.configure(tmp_path, monkeypatch, v2=True)
    folder = _folder(tmp_path)
    _put_doc("S1", "20261003_000000", pl=BILL)
    _put_doc("S2", "20261003_000000", pl=BILL)
    _put_doc("S1", "20261003_000000", basis="order", pl=ORDER)
    _put_doc("F1", "20261003_000000", pl=MAIN)
    _put_doc("F2", "20261003_000000", pl=MAIN, currency="USD")        # a second currency in one market
    got = MO.build(MONTH, folder, say=lambda t: None, log_dir=str(tmp_path / "logs"))
    _, rows = _sheet_rows(got["path"], MO.SHEET_OVERVIEW)
    subs = [r for r in rows if str(r[2]).startswith("小计")]
    by = dict(((r[0], r[1], r[2], r[3]), r) for r in subs)
    # ML Mexico: bill rows add up, order rows add up — never bill + order
    assert by[("美客多", "墨西哥", "小计（MXN）", "账单基准 正式")][10:14] == (240.0, -40.0, 200.0, 200.0)
    assert by[("美客多", "墨西哥", "小计（MXN）", "订单基准 参考")][10:14] == (110.0, -10.0, 100.0, 100.0)
    # Falabella Chile: one line per currency, nothing added across them
    assert by[("Falabella", "智利", "小计（CLP）", "—")][10:14] == (1000.0, -100.0, 900.0, 900.0)
    assert by[("Falabella", "智利", "小计（USD）", "—")][10:14] == (1000.0, -100.0, 900.0, 900.0)
    assert not [r for r in subs if "CLP" in str(r[2]) and r[10] == 2000.0]
    # Amazon Mexico has no version at all: no number, no subtotal with numbers
    az = [r for r in subs if r[0] == "亚马逊"]
    assert not az or az[0][10:14] == (None, None, None, None)


def test_headline_keys_exist_in_every_platform_pl_spec():
    """HEADLINE must only name rows the platform's results file really carries."""
    from platforms.mercadolibre.report import sheet_pl
    order_keys = set()
    for key, label, _es, _kind, _note in sheet_pl.build_pl_rows():
        order_keys.add(key if key is not None else sheet_pl.PL_FORMULA_KEYS[label])
    from platforms.mercadolibre.bill_report import workbook as BW
    S = {"prof": types.SimpleNamespace(vat_name="IVA", withholding=True, country_cn="墨西哥")}
    bill_keys = set(r.key for r in BW.pl_rows(S))
    specs = {("ML", "order"): order_keys, ("ML", "bill"): bill_keys}
    import importlib
    for pf, pkg in (("FA", "falabella"), ("AZ", "amazon"), ("LV", "liverpool")):
        mod = importlib.import_module("platforms.%s.report_build" % pkg)
        specs[(pf, "main")] = set(r.key for r in mod.PL_ROWS)
    assert set(HEADLINE_KEYS()) == set(specs), "every (platform, basis) with a results file has a HEADLINE"
    for key, spec in MO.HEADLINE.items():
        assert set(spec) == set(MO.FIGURES)
        for figure, keys in spec.items():
            missing = [k for k in keys if k not in specs[key]]
            assert not missing, "%s %s: %s not in the P&L spec" % (key, figure, missing)


def HEADLINE_KEYS():
    return list(MO.HEADLINE)


def test_figures_are_copied_and_blank_when_a_key_is_missing():
    doc = {"platform": "ML", "basis": "order",
           "pl": [{"key": "net_sales", "value": 10.0}, {"key": "order_fees_total", "value": -1.0},
                  {"key": "contrib", "value": 9.0}]}      # no bill_fees_total, no profit
    got = MO.figures(doc)
    assert got == {"revenue": 10.0, "fees": None, "contrib": 9.0, "profit": None}
    assert MO.figures({"platform": "XX", "basis": "main", "pl": []}) == dict((k, None) for k in MO.FIGURES)


# ------------------------------------------------------------ 需要你处理

def test_actions_list_only_what_a_person_must_do(tmp_path, monkeypatch):
    P.configure(tmp_path, monkeypatch, v2=True)
    folder = _folder(tmp_path)
    logs = str(tmp_path / "logs")
    # S1: a person must log in (pending human action from the last download)
    _write_run(logs, "20261003_020000_scheduled", [
        {"store": "S1", "status": "failed", "missing": [{"form": "ventas", "cause": "PAGE_BLOCKED"}]},
        {"store": "S2", "status": "failed", "errors": ["找不到店铺 S2"]},
        {"store": "S3", "status": "ok", "missing": []},
        {"store": "BOCINA_SM", "status": "ok", "missing": []}])
    # S3: a failed version; BOCINA_SM: gaps whose only reasons are non-blocking; F1 provisional; F2 complete
    _put_doc("S3", "20261003_000000", level="failed", reasons=["报表没有生成：崩了"], pl=[])
    _put_doc("S3", "20261003_000000", basis="order", level="failed", reasons=["报表没有生成：崩了"], pl=[])
    _put_doc("BOCINA_SM", "20261003_000000", level="gaps", pl=BILL,
             reasons=["缺下载（不影响金额）：stock/general（x）", "提醒（不影响金额）：y"])
    _put_doc("BOCINA_SM", "20261003_000000", basis="order", level="gaps", pl=ORDER,
             reasons=["缺下载（不影响金额）：stock/general（x）"])
    _put_doc("F1", "20261003_000000", level="provisional", reasons=["账单尚未关账：x（关账后重出）"], pl=MAIN)
    _put_doc("F2", "20261003_000000", pl=MAIN)
    _put_doc("A1", "20261003_000000", pl=MAIN)
    got = MO.build(MONTH, folder, run_notes=["[MX] 下载失败（退出码 2：紫鸟），报表未生成。"],
                   say=lambda t: None, log_dir=logs, entrance="manual")
    acts = got["actions"]
    stores = [(a["store"], a["problem"]) for a in acts]
    # (i) the human action, with who / how from download_diag
    human = [a for a in acts if a["store"] == "S1" and a["source"].startswith("下载记录")]
    assert human and human[0]["problem"] == "要登录 / 验证" and "打开勾选的店铺" in human[0]["next"]
    assert human[0]["who"] == "账号" and human[0]["unit"] == "墨西哥"
    # (iv) the runner's line, tagged with its country
    note = [a for a in acts if a["source"] == "本轮运行"]
    assert note and note[0]["unit"] == "墨西哥" and note[0]["problem"].startswith("下载失败")
    # (ii) stores without a version: S1, S2 (reason from the batch summary)
    s2 = [a for a in acts if a["store"] == "S2"]
    assert s2 and "紫鸟里找不到这家店" in s2[0]["problem"] and "config.json" in s2[0]["next"]
    assert [a for a in acts if a["store"] == "S1" and a["source"] == "月度结果"]
    # (iii) the failed version
    assert [a for a in acts if a["store"] == "S3" and a["problem"].startswith("报表没生成")]
    # excluded: non-blocking gaps, provisional (platform delay), complete stores
    assert not [s for s, _ in stores if s in ("BOCINA_SM", "F1", "F2", "A1")]
    # the sheet says the same
    _, rows = _sheet_rows(got["path"], MO.SHEET_ACTIONS)
    assert [r[1] for r in rows if r[1] == "S3"]
    assert not [r for r in rows if r[1] == "F2"]
    # and the alert file carries each item with who / next / source
    text = _read(folder)
    assert text.startswith("%s 需要处理的事（%d 项）" % (MONTH, len(acts)))
    assert "这一轮是：手动" in text and "谁来处理：账号" in text and "来源：期间报表" not in text
    assert "完整清单见 本月总览_%s.xlsx 的「需要你处理」页" % MONTH in text


def test_period_results_that_are_incomplete_or_failed_are_actions(tmp_path, monkeypatch):
    P.configure(tmp_path, monkeypatch, v2=True, catalog={"S1": P.CATALOG["S1"]})
    folder = _folder(tmp_path)
    for m in ("2026-07", "2026-09"):
        P.put("S1", m, "20261003_000000", pl=BILL)
    P.put("S1", MONTH, "20261003_000000", basis="order", pl=ORDER)
    from core import period_report as PR
    res = PR.generate(PR.quarter("2026-Q3"), "ML-MX", out_dir=str(tmp_path / "reports"))
    assert res["status"] == PR.INCOMPLETE
    got = MO.build(MONTH, folder, period_results=[res], say=lambda t: None, log_dir=str(tmp_path / "logs"))
    per = [a for a in got["actions"] if a["source"] == "期间报表"]
    assert len(per) == 1 and "S1 2026-08 无完整版本" in per[0]["problem"] and per[0]["who"] == "财务"
    # a complete one is not an action
    P.put("S1", "2026-08", "20261003_000000", pl=BILL)
    res = PR.generate(PR.quarter("2026-Q3"), "ML-MX", out_dir=str(tmp_path / "reports"))
    got = MO.build(MONTH, folder, period_results=[res], say=lambda t: None, log_dir=str(tmp_path / "logs"))
    assert not [a for a in got["actions"] if a["source"] == "期间报表"]


# ------------------------------------------------------------ 告警文件

def test_alert_is_written_then_removed_and_legacy_files_go(tmp_path, monkeypatch):
    P.configure(tmp_path, monkeypatch, v2=True, catalog={"S1": P.CATALOG["S1"]})
    folder = _folder(tmp_path)
    for name in ("！报表异常_%s.txt" % YM, "！报表异常_CL_FA_%s.txt" % YM, "！期间报表异常_MX_ML_2026-Q3_账单基准.txt",
                 "！期间报表异常_%s.txt" % YM, "！报表异常_202608.txt", "！期间报表异常_CL_FA_2026-Q2.txt"):
        open(os.path.join(folder, name), "w").close()
    said = []
    got = MO.build(MONTH, folder, say=said.append, log_dir=str(tmp_path / "logs"))
    assert os.path.isfile(_alert(folder)) and got["alert_path"] == _alert(folder)
    assert "S1" in _read(folder) and "本月没有报表" in _read(folder)
    gone = sorted(os.path.basename(p) for p in got["removed"])
    assert gone == sorted(["！报表异常_%s.txt" % YM, "！报表异常_CL_FA_%s.txt" % YM,
                           "！期间报表异常_MX_ML_2026-Q3_账单基准.txt", "！期间报表异常_%s.txt" % YM])
    assert os.path.exists(os.path.join(folder, "！报表异常_202608.txt")), "another month's file is not ours"
    assert os.path.exists(os.path.join(folder, "！期间报表异常_CL_FA_2026-Q2.txt"))
    assert any(t.startswith("已写出待处理清单") for t in said)
    # the store gets both bases: nothing left to do -> the alert goes
    _put_doc("S1", "20261004_000000", pl=BILL)
    _put_doc("S1", "20261004_000000", basis="order", pl=ORDER)
    said[:] = []
    got = MO.build(MONTH, folder, say=said.append, log_dir=str(tmp_path / "logs"))
    assert got["actions"] == [] and got["alert_path"] is None
    assert not os.path.exists(_alert(folder))
    assert any(t.startswith("没有需要处理的事，删掉了") for t in said)
    _, rows = _sheet_rows(got["path"], MO.SHEET_ACTIONS)
    assert rows and rows[0][0] == "这个月没有需要处理的事。"


def test_a_store_no_longer_in_config_is_listed_but_not_an_action(tmp_path, monkeypatch):
    P.configure(tmp_path, monkeypatch, v2=True)
    _put_doc("S2", "20261003_000000", pl=BILL)
    P.configure(tmp_path, monkeypatch, v2=True, catalog={"S1": P.CATALOG["S1"]})
    folder = _folder(tmp_path)
    got = MO.build(MONTH, folder, say=lambda t: None, log_dir=str(tmp_path / "logs"))
    _, rows = _sheet_rows(got["path"], MO.SHEET_OVERVIEW)
    assert [r for r in rows if r[2] == "S2（已不在配置里）" and r[4] == "完整"]
    assert not [a for a in got["actions"] if "S2" in a["store"]]


# ------------------------------------------------------------ 版本历史

def test_history_shows_deltas_and_flags_a_changed_final(tmp_path, monkeypatch):
    P.configure(tmp_path, monkeypatch, v2=True, catalog={"S1": P.CATALOG["S1"]})
    folder = _folder(tmp_path)
    _put_doc("S1", "20261001_000000", pl=BILL)                                          # first final
    _put_doc("S1", "20261002_000000", level="gaps", reasons=["缺下载：ventas（x）"],
             pl=[("revenue", 130.0), ("fees_total", -20.0), ("contrib", 110.0, "f"), ("profit", 110.0, "f")])
    _put_doc("S1", "20261003_000000",
             pl=[("revenue", 125.0), ("fees_total", -20.0), ("contrib", 105.0, "f"), ("profit", 105.0, "f")])
    got = MO.build(MONTH, folder, say=lambda t: None, log_dir=str(tmp_path / "logs"))
    _, rows = _sheet_rows(got["path"], MO.SHEET_HISTORY)
    rows = [r for r in rows if r[1] == "账单基准 正式"]
    assert [r[4][:10] for r in rows] == ["2026-10-03", "2026-10-02", "2026-10-01"], "newest first"
    assert [r[6] for r in rows] == ["是", "否", "否"]
    assert rows[0][7:11] == (125.0, -20.0, 105.0, 105.0)
    assert rows[0][11:15] == (-5.0, 0.0, -5.0, -5.0)          # vs the version below it
    assert rows[1][11:15] == (10.0, 0.0, 10.0, 10.0)
    assert rows[2][11:15] == (None, None, None, None)         # the oldest has nothing to compare with
    assert rows[0][15] == MO.CHANGED_AFTER_FINAL and rows[1][15] is None and rows[2][15] is None
    wb = load_workbook(got["path"])
    ws = wb[MO.SHEET_HISTORY]
    cell = [c for c in ws[MO.HEADER_ROW + 1] if c.value == MO.CHANGED_AFTER_FINAL][0]
    assert cell.font.color.rgb.endswith("C00000") and cell.font.bold


def test_a_final_with_the_same_figures_as_the_old_final_is_not_flagged(tmp_path, monkeypatch):
    P.configure(tmp_path, monkeypatch, v2=True, catalog={"S1": P.CATALOG["S1"]})
    folder = _folder(tmp_path)
    _put_doc("S1", "20261001_000000", pl=BILL)
    _put_doc("S1", "20261003_000000", pl=BILL)
    got = MO.build(MONTH, folder, say=lambda t: None, log_dir=str(tmp_path / "logs"))
    _, rows = _sheet_rows(got["path"], MO.SHEET_HISTORY)
    assert [r[15] for r in rows] == [None, None]


# ------------------------------------------------------------ 稳健

def test_locked_workbook_gets_the_fallback_name(tmp_path, monkeypatch):
    P.configure(tmp_path, monkeypatch, v2=True, catalog={"S1": P.CATALOG["S1"]})
    folder = _folder(tmp_path)
    target = os.path.join(folder, MO.overview_file_name(MONTH))
    real = os.replace

    def locked(src, dst):
        if dst == target:
            raise PermissionError(13, "正被 Excel 打开着", dst)
        return real(src, dst)
    monkeypatch.setattr(MO.os, "replace", locked)
    said = []
    got = MO.build(MONTH, folder, say=said.append, log_dir=str(tmp_path / "logs"))
    assert got["path"] == os.path.join(folder, "本月总览_%s(新).xlsx" % MONTH) and os.path.isfile(got["path"])
    assert any("写不进" in t and "(新)" in t for t in said)
    assert not [n for n in os.listdir(folder) if n.endswith(".part")], "no temp file left behind"


def test_build_never_raises(tmp_path, monkeypatch):
    P.configure(tmp_path, monkeypatch, v2=True)
    folder = str(tmp_path / "elsewhere")
    # results root missing entirely: a plain overview with every store 没有报表
    got = MO.build(MONTH, folder, say=lambda t: None, log_dir=str(tmp_path / "logs"))
    assert got["error"] is None and os.path.isfile(got["path"]) and got["rows"] == 11
    # the gathering itself blowing up: reported, alert still written with the runner's lines
    monkeypatch.setattr(MO, "gather", lambda *a, **k: 1 / 0)
    said = []
    got = MO.build(MONTH, folder, run_notes=["运行器内部出错，这一轮没跑完。"], say=said.append,
                   log_dir=str(tmp_path / "logs"))
    assert got["path"] is None and "ZeroDivisionError" in got["error"]
    text = _read(folder)
    assert "本月总览没生成" in text and "运行器内部出错" in text
    assert any(t.startswith("[警告] 本月总览没生成") for t in said)
    # config.json unreadable: a note, not a crash
    monkeypatch.undo()
    from core import store_config
    P.configure(tmp_path, monkeypatch, v2=True)

    def broken(*a, **k):
        raise store_config.StoreConfigError("坏了")
    monkeypatch.setattr(store_config, "stores_for_market", broken)
    got = MO.build(MONTH, folder, say=lambda t: None, log_dir=str(tmp_path / "logs"))
    assert got["error"] is None and os.path.isfile(got["path"])
    assert any("config.json 有问题" in a["problem"] and "坏了" in a["problem"] for a in got["actions"])


def test_config_problem_stores_are_shown_and_listed(tmp_path, monkeypatch):
    cat = dict(P.CATALOG)
    cat["WEIRD"] = {"platform": "ZZ", "country": "MX"}
    P.configure(tmp_path, monkeypatch, v2=True, catalog=cat)
    folder = _folder(tmp_path)
    got = MO.build(MONTH, folder, say=lambda t: None, log_dir=str(tmp_path / "logs"))
    _, rows = _sheet_rows(got["path"], MO.SHEET_OVERVIEW)
    weird = [r for r in rows if r[2] == "WEIRD"]
    assert weird and weird[0][4] == MO.CONFIG_BAD and "未知平台" in weird[0][14]
    assert [a for a in got["actions"] if a["store"] == "WEIRD" and a["source"] == "config.json"]


# ------------------------------------------------------------ 和运行器接在一起

from test_region_all import env, cfg, plan, child  # noqa: E402,F401  the --region ALL fixtures


def test_a_whole_run_leaves_one_overview_and_one_alert(env):  # noqa: F811
    code = env.run(["--region", "ALL", "--entrance", "scheduled"], {
        "MX": plan(("download", child(0)), ("report", child(0))),
        "CL": plan(("download", child(2, "紫鸟客户端错误：起不来")), ("report", child(0))),
    })
    assert code == 2
    names = sorted(os.listdir(env.out_dir))
    assert names == sorted(["本月总览_%s.xlsx" % MONTH, "！需要处理_%s.txt" % YM])
    assert not [n for n in names if n.startswith("！报表异常") or n.startswith("！期间报表异常")]
    text = _read(env.out_dir)
    assert "智利 —— 下载失败（退出码 2" in text
    # every configured store without a report is listed, each once
    for s in ("BOCINA_SM", "EWTTO_SM", "Tecnohub_CL", "Andes_CL"):
        assert text.count(" %s —— 本月没有报表" % s) == 1
    _, doc, recs = env.last()
    msgs = [r["message"] for r in recs]
    assert any(m.startswith("[MX] 已更新本月总览：") for m in msgs), "each country's section logs its refresh"
    assert any(m.startswith("[CL] 已写出待处理清单：") for m in msgs)
    assert any(m.startswith("已更新本月总览：") for m in msgs), "and once more at the very end"
    # the window's button finds the same file
    from core import run_monthly
    assert os.path.join(run_monthly.alert_dir(env.out_dir, MONTH), MO.overview_file_name(MONTH)) == \
        os.path.join(env.out_dir, "本月总览_%s.xlsx" % MONTH)
