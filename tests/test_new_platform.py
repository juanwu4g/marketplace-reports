# -*- coding: utf-8 -*-
"""tools/new_platform: the generated package imports, keeps the platform
contract with --dry-run, and stops real work with a clear message.

It is generated into a temporary folder, made importable as platforms.<pkg>
and registered in a copy of core.markets.PLATFORMS for the test only."""
import collections
import io
import os
import re
import sys

import pytest
from openpyxl import Workbook, load_workbook

import platforms
from core import markets
from core import report_kit as K
from tools import new_platform as NP

import test_platform_contract as contract

PKG = "zz_scaffold"
CODE = "ZZ"


@pytest.fixture
def scaffold(tmp_path, monkeypatch):
    root = tmp_path / "platforms"
    folder = NP.generate(CODE, PKG, "测试平台", "MX", root=str(root), name_en="Zeta Zone")
    monkeypatch.setattr(platforms, "__path__", list(platforms.__path__) + [str(root)])
    reg = collections.OrderedDict(markets.PLATFORMS)
    line = NP.registry_line(CODE, PKG, "测试平台", "MX", "Zeta Zone")
    ns = {"Platform": markets.Platform, "re": re}
    exec("entries = [\n%s\n]" % line, ns)          # the printed line is valid Python
    code, p = ns["entries"][0]
    reg[code] = p
    monkeypatch.setattr(markets, "PLATFORMS", reg)
    yield p, folder
    for name in [m for m in sys.modules if m.startswith("platforms." + PKG)]:
        del sys.modules[name]


def test_generates_every_file(scaffold):
    _p, folder = scaffold
    assert sorted(os.listdir(folder)) == sorted(NP.FILES)
    readme = io.open(os.path.join(folder, "README.md"), encoding="utf-8").read()
    assert "ZZ-MX" in readme and "MX_ZZ_全店汇总_YYYY-MM.xlsx" in readme
    for name in NP.FILES:
        text = io.open(os.path.join(folder, name), encoding="utf-8").read()
        assert not re.search(r"__[A-Z_]+__", text.replace("__init__", "").replace("__name__", "")
                             .replace("__main__", "").replace("__file__", "")), name


def test_registry_line_matches_ziniao_names(scaffold):
    p, _ = scaffold
    assert p.package == "platforms." + PKG and p.countries == ("MX",)
    assert markets.ziniao_matches(CODE, "Zeta Zone-墨西哥")
    assert markets.ziniao_matches(CODE, "zetazone")
    assert not markets.ziniao_matches(CODE, "falabella-智利")


def test_scaffold_keeps_the_platform_contract(scaffold, tmp_path, monkeypatch, capsys):
    p, _ = scaffold
    # the same checks test_platform_contract runs for registered platforms
    from core import paths, pending_store, store_config
    cfg = contract.config_for_all()
    monkeypatch.setattr(store_config, "_stores", lambda: cfg)
    monkeypatch.setattr(paths, "ROOT", str(tmp_path))
    monkeypatch.setattr(pending_store, "prune", lambda *a, **k: [])
    import core.ziniao_client as zc

    def trap(*a, **k):
        raise AssertionError("must not start Ziniao")
    monkeypatch.setattr(zc, "ZiniaoClient", trap)

    contract.test_package_has_runner_modules(p)
    contract.test_registry_entry_is_consistent(p)
    contract.test_periods_declare_the_expected_forms(p)
    for stores in (None, [contract.own_store(p)]):
        for step in contract.plan(p, str(tmp_path / "out"), stores):
            module, args = contract.runner_args(step["cmd"])
            code = contract.run_cli(module, args + ["--dry-run"], monkeypatch)
            assert code in ((0,) if module.endswith("run_batch") else (0, 2)), (step, code)
    mk = ["--market", "ZZ-MX", "--month", "2026-09", "--dry-run"]
    batch, reports = "platforms.%s.run_batch" % PKG, "platforms.%s.run_reports" % PKG
    assert contract.run_cli(batch, mk + ["--stores", "NOPE"], monkeypatch) == 3
    assert contract.run_cli(batch, mk + ["--stores", contract.own_store(
        markets.PLATFORMS["FA"])], monkeypatch) == 4
    assert contract.run_cli(batch, mk[:-1] + ["--market", "FA-CL", "--stores",
                                              contract.own_store(p)], monkeypatch) == 4
    got = contract.run_cli(reports, mk + ["--stores", "NOPE", "--out-dir", str(tmp_path)],
                           monkeypatch)
    assert isinstance(got, store_config.StoreNotAllowed) or got == 3

    # a real download stops before Ziniao, saying it is not built yet
    capsys.readouterr()
    assert contract.run_cli(batch, mk[:-1] + ["--stores", contract.own_store(p)],
                            monkeypatch) == 4
    assert "还没实现" in capsys.readouterr().out


def test_scaffold_report_lays_out_ten_sheets(scaffold, tmp_path):
    import importlib
    rb = importlib.import_module("platforms.%s.report_build" % PKG)
    with pytest.raises(NotImplementedError) as e:
        rb.build([("S1", str(tmp_path))], "2026-09", str(tmp_path / "x.xlsx"))
    assert "还没实现" in str(e.value)
    stores = [{"store": s, "meta": {}, "m": {"pl": {k: 10.0 for k in rb.DATA_KEYS}}}
              for s in ("S1", "S2")]
    ctx = rb.build_context(stores, "2026-09", "2026-10-01")
    ctx["checks"] = [K.check("x", "S1", "示例", False, "说明")]
    out = rb.build_workbook(stores, ctx, str(tmp_path / "w.xlsx"))
    wb = load_workbook(out)
    assert wb.sheetnames == K.SHEET_NAMES
    pl = wb[K.SHEET_NAMES[1]]
    assert pl["E5"].value == "合计"
    assert wb[K.SHEET_NAMES[0]]["B1"].value.startswith("测试平台 月度财务报表")


def test_refuses_clashes(tmp_path):
    with pytest.raises(ValueError) as e:
        NP.generate("FA", "falabella2", "x", "CL", root=str(tmp_path))
    assert "已经被" in str(e.value)
    with pytest.raises(ValueError):
        NP.generate("QQ", "falabella", "x", "CL", root=str(tmp_path))
    with pytest.raises(ValueError):
        NP.generate("QQ", "Bad-Name", "x", "CL", root=str(tmp_path))
    with pytest.raises(ValueError):
        NP.generate("QQ", "qq", "x", "XX", root=str(tmp_path))
    (tmp_path / "qq").mkdir()
    with pytest.raises(ValueError):
        NP.generate("QQ", "qq", "x", "CL", root=str(tmp_path))
    assert NP.main(["QQ", "qq", "x", "CL", "--root", str(tmp_path)]) == 2


def test_cli_prints_the_registry_line(tmp_path, capsys):
    assert NP.main(["qq", "qq_shop", "QQ 商城", "CL", "--root", str(tmp_path)]) == 0
    out = capsys.readouterr().out
    assert '("QQ", Platform("QQ", "platforms.qq_shop", "QQ 商城", "Qq Shop",' in out
    assert r're.compile(r"qq\s*shop", re.I), ("CL",))),' in out
    assert (tmp_path / "qq_shop" / "run_batch.py").exists()
    # markets.py itself is not touched
    assert "QQ" not in markets.PLATFORMS
