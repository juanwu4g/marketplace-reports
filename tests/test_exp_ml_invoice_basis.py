# -*- coding: utf-8 -*-
"""experiments/ml_invoice_basis: the invoice-basis experiment on synthetic folders.

Synthetic data only (real downloads carry buyers' names and addresses). The
Mexican folder, month 2026-09:

  Ventas: O1 delivered (1,000; buyer paid 50 shipping), O2 a two-item pack
          (parent row 600 = 400 + 200, two item rows without money), O3 cancelled
          (500), O4 ordered 09-30 and billed on 10-01, O5 ordered 08-30.
  Sept invoice: commission + shipping of O1; the two pack items; O3 charged and
          reversed in the invoice; a sale Ventas does not list (300), charged and
          reversed; O5 charged 09-01; Product Ads 200; an unregistered fee 10.
  Sept credit notes: 10 % of O1's commission returned (a partial refund); a full
          reversal of a sale whose charge is in an invoice not downloaded (found
          through the listing's category rate, 13 % of 1,000 = 130 exactly).
  Oct invoice (still open): O4's commission.
"""
import contextlib
import io
import json
import os

import pytest
from openpyxl import Workbook, load_workbook

from core import report_kit as K
from platforms.mercadolibre.report import site
from platforms.mercadolibre.report.catalog import BILLING_HEADER_ROW, VENTAS_HEADER_ROW

from experiments.ml_invoice_basis import __main__ as cli
from experiments.ml_invoice_basis import common, compare, invoice, report, stability, workbook

MONTH = "2026-09"
STORE = "EXP_TEST"

B_HEAD = ["N° de factura fiscal", "Fecha del cargo", "Número del cargo", "Detalle",
          "Descontado de la operación", "Estado del cargo", "Cargo que bonifica",
          "Valor del cargo", "Porcentaje por categoría", "Costo por unidad vendida",
          "Número de venta", "Fecha de venta", "Cantidad vendida", "Precio unitario",
          "Total de la venta", "Número de paquete", "Envío a cargo del cliente",
          "Número de publicación", "Título de publicación"]


def bill(no, date, det, amount, ch="Si", state=None, rev=None, sale=None, sdate=None, qty=None,
         price=None, total=None, pack=None, buyer_ship=None, pub=None, pct=None):
    return ["A1", date, no, det, ch, state, rev, amount, pct, 0.0 if pct is not None else None,
            sale, sdate, qty, price, total, pack, buyer_ship, pub,
            "Producto %s" % (pub or "") if sale else None]


CV, SHIP, ANUL = "Cargo por venta", "Cargo por envíos de Mercado Libre", "Anulación del cargo por venta"
O1, O2, O3, O4, O5, GHOST = ("2000000000000001", "2000000000000002", "2000000000000003",
                             "2000000000000004", "2000000000000005", "2000000000000009")
ITEM_A, ITEM_B = "3000000000000021", "3000000000000022"

SEPT = [
    bill("101", "2026-09-05", CV, 130, sale=O1, sdate="2026-09-05", qty=1, price=1000, total=1000,
         pub="MLM1", pct=0.13),
    bill("102", "2026-09-05", SHIP, 100, sale=O1, sdate="2026-09-05", qty=1, price=1000, total=1000,
         buyer_ship=50, pub="MLM1"),
    bill("103", "2026-09-10", CV, 52, sale=ITEM_A, sdate="2026-09-10", qty=1, price=400, total=400,
         pack=O2, pub="MLM2", pct=0.13),
    bill("104", "2026-09-10", CV, 26, sale=ITEM_B, sdate="2026-09-10", qty=1, price=200, total=200,
         pack=O2, pub="MLM3", pct=0.13),
    bill("105", "2026-09-12", CV, 65, ch="No aplica", state="Anulado en factura", sale=O3,
         sdate="2026-09-12", qty=1, price=500, total=500, pub="MLM1", pct=0.13),
    bill("106", "2026-09-12", ANUL, -65, ch="No aplica", rev="105", sale=O3, sdate="2026-09-12",
         qty=1, price=500, total=500, pub="MLM1"),
    bill("107", "2026-09-15", CV, 39, ch="No aplica", state="Anulado en factura", sale=GHOST,
         sdate="2026-09-15", qty=1, price=300, total=300, pub="MLM2", pct=0.13),
    bill("108", "2026-09-15", ANUL, -39, ch="No aplica", rev="107", sale=GHOST, sdate="2026-09-15",
         qty=1, price=300, total=300, pub="MLM2"),
    bill("109", "2026-09-01", CV, 104, sale=O5, sdate="2026-08-30", qty=1, price=800, total=800,
         pub="MLM1", pct=0.13),
    bill("110", "2026-09-03", "Cargo por campaña de publicidad de Product Ads", 200, ch="No"),
    bill("111", "2026-09-20", "Cargo por algo nuevo", 10, ch="No"),
]
SEPT_NC = [
    bill("201", "2026-09-25", ANUL, -13, ch="No aplica", rev="101", sale=O1, sdate="2026-09-05",
         qty=1, price=1000, total=1000, pub="MLM1"),
    bill("202", "2026-09-26", ANUL, -130, ch="No aplica", rev="999", sale="2000000000000007",
         sdate="2026-08-20", qty=1, price=1000, total=1000, pub="MLM1"),
]
OCT = [bill("301", "2026-10-01", CV, 91, sale=O4, sdate="2026-09-30", qty=1, price=700,
            total=700, pub="MLM1", pct=0.13)]

V_HEAD = ["# de venta", "Fecha de venta", "Estado", "Paquete de varios productos", "Unidades",
          "Ingresos por productos ({c})", "Cargo por venta e impuestos ({c})",
          "Ingresos por envío ({c})", "Costos de envío ({c})", "Anulaciones y reembolsos ({c})",
          "Total ({c})", "SKU", "Título de la publicación"]


def ventas_rows(o1_status="Entregado", o1_refund=0.0):
    # (id, day, month, status, pack flag, units, revenue, commission+tax, ship income,
    #  ship cost, refund, sku)
    return [
        (O1, 5, "septiembre", o1_status, "No", 1, 1000, -150, 50, -100, o1_refund, "SKU-1"),
        (O2, 10, "septiembre", "Paquete de 2 productos", "", None, 600, -90, None, None, None, ""),
        (ITEM_A, 10, "septiembre", "Entregado", "Sí", 1, None, None, None, None, None, "SKU-A"),
        (ITEM_B, 10, "septiembre", "Entregado", "Sí", 1, None, None, None, None, None, "SKU-B"),
        (O3, 12, "septiembre", "Cancelada por el comprador", "No", 1, 500, -75, None, None, -425,
         "SKU-1"),
        (O4, 30, "septiembre", "Entregado", "No", 1, 700, -105, None, -70, None, "SKU-1"),
        (O5, 30, "agosto", "Entregado", "No", 1, 800, -120, None, -80, None, "SKU-1"),
    ]


def _billing_file(path, rows, banner=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "REPORT"
    ws["A1"] = "Fecha de actualización"
    if banner:
        ws["A3"] = "Este reporte podría tener cambios hasta el cierre del período de facturación"
    for i, h in enumerate(B_HEAD, 1):
        ws.cell(BILLING_HEADER_ROW + 1, i, h)
    for r, row in enumerate(rows, BILLING_HEADER_ROW + 2):
        for i, v in enumerate(row, 1):
            ws.cell(r, i, v)
    wb.save(path)


def make_folder(folder, site_code="mx", ventas=None, bills=None, oct_open=True):
    folder = str(folder)
    os.makedirs(folder, exist_ok=True)
    cur = {"mx": "MXN", "cl": "CLP", "pe": "PEN"}[site_code]
    with io.open(os.path.join(folder, "run_meta.json"), "w", encoding="utf-8") as fh:
        json.dump({"site": site_code, "month": MONTH, "downloaded_at": "2026-10-14T10:00:00"}, fh)
    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas %s" % site_code.upper()
    ws.cell(VENTAS_HEADER_ROW, 1, "Ventas")
    for i, h in enumerate(V_HEAD, 1):
        ws.cell(VENTAS_HEADER_ROW + 1, i, h.format(c=cur))
    for r, (oid, d, mes, st, pk, u, rev, com, si, sc, ref, sku) in enumerate(
            ventas or ventas_rows(), VENTAS_HEADER_ROW + 2):
        total = None if rev is None else rev + (com or 0) + (si or 0) + (sc or 0) + (ref or 0)
        row = [oid, "%d de %s de 2026 10:00 hs." % (d, mes), st, pk, u, rev, com, si, sc, ref,
               total, sku, "Producto " + sku if sku else None]
        for i, v in enumerate(row, 1):
            ws.cell(r, i, v)
    wb.save(os.path.join(folder, "20261014_Ventas_%s_Mercado_Libre_y_Mercado_Shops_"
                                 "2026-10-14_10-00hs_123.xlsx" % site_code.upper()))
    sept, sept_nc, octo = bills or (SEPT, SEPT_NC, OCT)
    _billing_file(os.path.join(folder, "Reporte_Facturacion_MercadoLibre_Sept2026.xlsx"), sept)
    if sept_nc:
        _billing_file(os.path.join(folder, "Reporte_Notas_Credito_MercadoLibre_Sept2026.xlsx"),
                      sept_nc)
    if octo:
        _billing_file(os.path.join(folder, "Reporte_Facturacion_MercadoLibre_Oct2026.xlsx"), octo,
                      banner=oct_open)
    return folder


@pytest.fixture(autouse=True)
def _back_to_mx():
    yield
    site.set_site("mx")


def _compute(folder, store=STORE):
    with contextlib.redirect_stdout(io.StringIO()):
        return report.compute(store, folder, MONTH)


# ════════════════════════════════════════════════════════════════════════
# the invoice basis
# ════════════════════════════════════════════════════════════════════════

def test_invoice_basis_sales_reversals_fees(tmp_path):
    D = _compute(make_folder(tmp_path / STORE / "20261014_100000_m202609"))
    assert D["site"] == "mx" and D["cycle_day"] == 1
    P = D["invs"]["period"]
    n = P["n"]
    assert (n["n_sales"], n["units"], n["gross_sales"]) == (6, 6.0, 3200.0)
    assert P["view"]["start"].day == 1 and P["view"]["end"].day == 30
    # 106 / 108 full (ID), 201 = 13/130 of O1, 202 = original charge not downloaded:
    # MLM1's 13 % of 1,000 is exactly 130 -> a full reversal
    R = P["reversals"].set_index("charge_no")
    assert round(R.loc["201", "ratio"], 6) == 0.1 and round(R.loc["201", "est_sales"], 2) == 100
    assert R.loc["202", "method"] == invoice.LINK_RATE and R.loc["202", "est_sales"] == 1000
    assert n["rev_methods"][invoice.LINK_ID] == 3 and n["rev_methods"][invoice.LINK_RATE] == 1
    assert round(n["rev_sales"], 2) == 500 + 300 + 100 + 1000
    assert n["ship_income"] == 50 and n["ship_income_rev"] == 0
    # every invoice row lands in a catalogue line or in the unknown list, never lost
    rows_total = float(P["view"]["rows"]["amount"].sum())
    assert round(n["fee_total"], 2) == round(rows_total, 2)
    assert round(sum(f["net"] for f in P["fees"]) + sum(u["amount"] for u in P["unknown"]), 2) \
        == round(rows_total, 2)
    assert [u["es"] for u in P["unknown"]] == ["Cargo por algo nuevo"]
    assert n["unk_bill"] == 10 and n["unk_order"] == 0
    assert round(n["contribution"], 2) == round(3200 - 1900 + 50 - rows_total, 2)
    # Mexico bills by calendar month: both readings are the same rows
    assert D["invs"]["calendar"]["n"]["gross_sales"] == n["gross_sales"]
    assert P["view"]["closed"] and not any("关账" in i for i in P["view"]["issues"])


def test_reconciliation_splits_the_gmv_difference_exactly(tmp_path):
    D = _compute(make_folder(tmp_path / STORE / "20261014_100000_m202609"))
    rec = D["recs"]["period"]
    E = rec["extra"].set_index("ksale")
    assert E.loc[GHOST, "code"] == "absent_rev"         # cancelled at once, not in Ventas
    assert E.loc[O5, "code"] == "prev"                  # August order charged 09-01
    M = rec["missing"].set_index("okey")
    assert list(M.index) == [O4] and M.loc[O4, "code"] == "elsewhere"
    assert "Octubre 2026" in M.loc[O4, "why"]
    T = rec["matched"].set_index("okey")
    assert sorted(T.index) == sorted([O1, O2, O3])      # the pack matched through Número de paquete
    assert T.loc[O2, "inv_total"] == 600 and T.loc[O2, "units"] == 2
    s = compare.rec_summary(rec)
    gmv = D["S"]["m"]["ing"]
    assert rec["month_gmv"] == gmv == 2800
    assert round(D["invs"]["period"]["n"]["gross_sales"] - gmv, 2) == round(
        s["extra_total"] - s["missing_gmv"] + s["matched_diff"], 2)


def test_comparison_lines_add_up_to_both_contributions(tmp_path):
    D = _compute(make_folder(tmp_path / STORE / "20261014_100000_m202609"))
    lines = dict((ln["key"], ln) for ln in D["lines"])
    money = [ln for ln in D["lines"] if ln["kind"] == "m"]
    cur = sum(ln["cur"] for ln in money)
    assert round(cur, 2) == round(D["S"]["contrib"], 2)
    for v in ("period", "calendar"):
        inv = sum(ln["inv"][v] or 0.0 for ln in money)
        assert round(inv, 2) == round(D["invs"][v]["n"]["contribution"], 2)
    assert lines["coupon"]["inv"]["period"] is None           # not in the invoice
    assert lines["gmv"]["cur"] == 2800 and lines["gmv"]["inv"]["period"] == 3200
    labels = [u[0] for u in D["unavailable"]]
    assert any("代扣" in x for x in labels) and any("MercadoPago" in x for x in labels)


def test_order_basis_is_the_monthly_reports_own(tmp_path):
    """derive() through the experiment gives exactly what the monthly report gives."""
    from platforms.mercadolibre import report_build as rb
    folder = make_folder(tmp_path / STORE / "20261014_100000_m202609")
    D = _compute(folder)
    rb.set_site("mx")
    p0, p1 = rb.month_bounds(MONTH)
    with contextlib.redirect_stdout(io.StringIO()):
        S = rb.derive(STORE, folder, p0, p1)
        ctx = rb.build_context([S], MONTH)
    for k in ("ing", "refund", "com_gross", "env_cost", "tax_ret", "bill_total"):
        assert D["S"]["m"][k] == S["m"][k], k
    assert D["S"]["contrib"] == ctx["contrib"][STORE]


def test_chile_invoice_period_runs_from_the_cut_off_day(tmp_path):
    sept = [bill("401", "2026-08-08", "Cargo por campaña de publicidad de Product Ads", 1000, ch="No"),
            bill("402", "2026-09-03", CV, 1300, sale=O1, sdate="2026-09-03", qty=1, price=10000,
                 total=10000, pub="MLC1", pct=0.13)]
    octo = [bill("501", "2026-09-08", "Cargo por campaña de publicidad de Product Ads", 500, ch="No"),
            bill("502", "2026-09-20", CV, 650, sale=O3, sdate="2026-09-20", qty=1, price=5000,
                 total=5000, pub="MLC1", pct=0.13)]
    ventas = [(O1, 3, "septiembre", "Entregado", "No", 1, 10000, -1300, None, None, None, "S1"),
              (O3, 20, "septiembre", "Entregado", "No", 1, 5000, -650, None, None, None, "S1")]
    folder = make_folder(tmp_path / "CL_TEST" / "20261014_100000_m202609", "cl", ventas,
                         (sept, None, octo), oct_open=False)
    D = _compute(folder, "CL_TEST")
    assert D["site"] == "cl" and D["cycle_day"] == 8 and D["dec"] == 0
    P, C = D["invs"]["period"], D["invs"]["calendar"]
    assert (P["view"]["start"].strftime("%m-%d"), P["view"]["end"].strftime("%m-%d")) == ("08-08",
                                                                                         "09-07")
    assert P["n"]["gross_sales"] == 10000 and C["n"]["gross_sales"] == 15000
    assert C["view"]["bills"] == [(2026, 9), (2026, 10)] and C["view"]["closed"]
    # Product Ads: 08-08 in the period only, 09-08 in the calendar cut only
    ads = "Cargo por campaña de publicidad de Product Ads"
    assert [f["net"] for f in P["fees"] if f["es"] == ads] == [1000]
    assert [f["net"] for f in C["fees"] if f["es"] == ads] == [500]
    assert site.SITE == "mx"                              # restored after the Chilean folder
    # without the October invoice the calendar cut says what is missing
    os.remove(os.path.join(folder, "Reporte_Facturacion_MercadoLibre_Oct2026.xlsx"))
    D = _compute(folder, "CL_TEST")
    iss = D["invs"]["calendar"]["view"]["issues"]
    assert not D["invs"]["calendar"]["view"]["complete"]
    assert any("Octubre 2026" in i and "09/08" in i for i in iss)


def test_open_invoice_is_reported(tmp_path):
    sept, nc, octo = SEPT, SEPT_NC, OCT
    folder = make_folder(tmp_path / STORE / "x_m202609", bills=(sept, nc, octo))
    _billing_file(os.path.join(folder, "Reporte_Facturacion_MercadoLibre_Sept2026.xlsx"), sept,
                  banner=True)
    D = _compute(folder)
    v = D["invs"]["period"]["view"]
    assert not v["closed"] and any("还没关账" in i for i in v["issues"])


def test_listing_terms_prefer_the_exact_rate():
    """A later rate of the same listing must not change an earlier reversal."""
    import pandas as pd
    terms = {(0.145, 0.0, pd.Timestamp("2026-08-28")), (0.11, 0.0, pd.Timestamp("2026-10-02"))}
    assert invoice._from_terms(terms, 406.89, 3699.0, 1.0, pd.Timestamp("2026-08-29"), 0.05) \
        == 406.89
    assert invoice._from_terms(terms, 100.0, 3699.0, 1.0, pd.Timestamp("2026-08-29"), 0.05) \
        == round(0.145 * 3699, 2)


# ════════════════════════════════════════════════════════════════════════
# stability
# ════════════════════════════════════════════════════════════════════════

def test_stability_order_basis_moves_invoice_basis_does_not(tmp_path):
    a = make_folder(tmp_path / STORE / "20261013_100000_m202609")
    b = make_folder(tmp_path / STORE / "20261020_100000_m202609",
                    ventas=ventas_rows("Devolución finalizada. Pusimos el producto de nuevo a la venta",
                                       -850))
    with contextlib.redirect_stdout(io.StringIO()):
        A = stability.analyse(STORE, MONTH, stability.default_folders(STORE, MONTH, str(tmp_path)))
    assert [os.path.basename(D["folder"])[:8] for D in A["snaps"]] == ["20261013", "20261020"]
    k0, k1 = A["kns"]
    assert k0["order"]["refund"] == -425 and k1["order"]["refund"] == -1275
    for v in ("period", "calendar"):
        assert k0[v] == k1[v]
    p = A["pairs"][0]
    assert [x["okey"] for x in p["ventas"]] == [O1] and "状态" in p["ventas"][0]["kind"]
    assert p["invoice"]["period"][0] == []
    assert "逐行相同" in p["why"]["period"] and any("状态变化 1 单" in s for s in p["why"]["order"])
    assert a and b


def test_stability_sees_a_changed_invoice_row(tmp_path):
    make_folder(tmp_path / STORE / "20261013_100000_m202609")
    sept = [list(r) for r in SEPT]
    sept[0][7] = 140                                    # 101's amount changed
    make_folder(tmp_path / STORE / "20261020_100000_m202609", bills=(sept, SEPT_NC, OCT))
    with contextlib.redirect_stdout(io.StringIO()):
        A = stability.analyse(STORE, MONTH, stability.default_folders(STORE, MONTH, str(tmp_path)))
    diff, na, nb = A["pairs"][0]["invoice"]["period"]
    assert (na, nb) == (13, 13) and len(diff) == 1 and diff[0]["charge"] == "101"
    assert "金额" in diff[0]["kind"] and "1 行不同" in A["pairs"][0]["why"]["period"]


# ════════════════════════════════════════════════════════════════════════
# command lines and workbooks
# ════════════════════════════════════════════════════════════════════════

def _run(main, argv):
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        rc = main(argv)
    return rc, buf.getvalue()


def test_cli_writes_the_experiment_workbook(tmp_path):
    base = tmp_path / "downloads"
    make_folder(base / STORE / "20261013_100000_m202609")
    make_folder(base / STORE / "20261020_100000_m202609")
    out_dir = tmp_path / "out"
    rc, text = _run(cli.main, ["--store", STORE, "--month", MONTH, "--base", str(base),
                               "--out-dir", str(out_dir), "--no-recalc"])
    assert rc == 0, text
    out = out_dir / ("%s_%s_账单口径实验.xlsx" % (STORE, MONTH))
    assert out.exists()
    assert out.name != K.store_workbook_name(STORE, MONTH)
    wb = load_workbook(str(out))
    assert wb.sheetnames == workbook.SHEETS
    assert "20261020" in text                                 # the newest folder was used
    pl = wb[workbook.SHEETS[1]]
    labels = [pl.cell(row=i, column=1).value for i in range(1, pl.max_row + 1)]
    assert "开票销售额" in labels and "残差（应为 0）" in labels and "减：⚠ 未登记科目" in labels
    st = wb[workbook.SHEETS[5]]
    assert any("A. 现行订单口径" in str(c.value or "") for c in st["A"])


def test_cli_stability_workbook(tmp_path):
    a = make_folder(tmp_path / STORE / "20261013_100000_m202609")
    b = make_folder(tmp_path / STORE / "20261020_100000_m202609")
    rc, text = _run(stability.main, ["--store", STORE, "--month", MONTH, "--folders", a, b,
                                     "--out-dir", str(tmp_path / "o"), "--no-recalc"])
    assert rc == 0, text
    wb = load_workbook(str(tmp_path / "o" / ("%s_%s_口径稳定性.xlsx" % (STORE, MONTH))))
    assert wb.sheetnames == workbook.STAB_SHEETS


def test_cli_errors_are_chinese_and_exit_2(tmp_path):
    rc, text = _run(cli.main, ["--store", "NOPE", "--month", MONTH, "--base", str(tmp_path),
                               "--no-recalc"])
    assert rc == 2 and "[错误]" in text and "没找到" in text
    rc, text = _run(cli.main, ["--store", STORE, "--month", "2026-9", "--no-recalc"])
    assert rc == 2 and "YYYY-MM" in text
    one = make_folder(tmp_path / STORE / "20261013_100000_m202609")
    rc, text = _run(stability.main, ["--store", STORE, "--month", MONTH, "--folders", one,
                                     "--no-recalc"])
    assert rc == 2 and "至少两次下载" in text


def test_default_output_folder_is_git_ignored():
    assert common.DEFAULT_OUT_DIR.replace("\\", "/").endswith("downloads/data/reports/experiments")
    p = common.out_path(None, "A/B", MONTH, common.REPORT_SUFFIX)
    assert os.path.basename(p) == "A_B_2026-09_账单口径实验.xlsx"


@pytest.mark.skipif(not __import__("shutil").which("soffice"), reason="LibreOffice not installed")
def test_workbook_formulas_recalculate_without_errors(tmp_path):
    folder = make_folder(tmp_path / STORE / "20261014_100000_m202609")
    rc, text = _run(cli.main, ["--store", STORE, "--month", MONTH, "--folder", folder,
                               "--out-dir", str(tmp_path / "o")])
    assert rc == 0, text
    assert "错误 0 条" in text
    wb = load_workbook(str(tmp_path / "o" / ("%s_%s_账单口径实验.xlsx" % (STORE, MONTH))),
                       data_only=True)
    for name in (workbook.SHEETS[1], workbook.SHEETS[2]):
        ws = wb[name]
        for row in ws.iter_rows():
            if row[0].value == "残差（应为 0）":
                assert all(abs(c.value or 0) < 0.005 for c in row[2:5]
                           if isinstance(c.value, (int, float))), name
