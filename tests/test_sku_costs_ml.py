# -*- coding: utf-8 -*-
"""The shared SKU cost table (core/sku_costs.py) in the MercadoLibre workbook
(Mexico and Chile) and in tools/sku_cost_template on MercadoLibre folders.
Synthetic data only.

The folder has five orders in September: SKU-A delivered (2 units) and cancelled
before shipping (1 unit), SKU-B delivered (1) and returned back on sale (1), SKU-C
returned but not reviewed yet (3). Cost units are ⑦ 表五's 销售件数: A 2, B 1, C 3.
"""
import contextlib
import io
import json
import os
import shutil

import pytest
from openpyxl import Workbook, load_workbook

from core import report_kit as K
from core import sku_costs as SC
from core import store_config
from platforms.mercadolibre import report_build as rb

MONTH = "2026-09"
HEAD = ["平台", "国家", "SKU", "单位成本", "币种", "生效日期", "备注"]
SHEETS = rb.SHEET_NAMES
# (order, status, units, revenue, commission, shipping cost, refund, SKU)
ORDERS = [
    ("2000000000000001", "Entregado", 2, 200, -24, -20, 0, "SKU-A"),
    ("2000000000000002", "Entregado", 1, 300, -36, -30, 0, "SKU-B"),
    ("2000000000000003", "Cancelada por el comprador", 1, 100, 0, 0, -100, "SKU-A"),
    ("2000000000000004", "Devolución finalizada. Pusimos el producto de nuevo a la venta",
     1, 300, 0, 0, -300, "SKU-B"),
    ("2000000000000005", "Devolución en camino", 3, 90, 0, 0, -90, "SKU-C"),
]
UNITS = {"SKU-A": 2, "SKU-B": 1, "SKU-C": 3}
STORE = "ML_TEST"


@pytest.fixture(autouse=True)
def _back_to_mx():
    yield
    rb.set_site("mx")


def make_folder(folder, site):
    """A download folder of one site with the five orders and one bill."""
    folder = str(folder)
    os.makedirs(folder, exist_ok=True)
    before = rb.SITE
    rb.set_site(site)
    try:
        with io.open(os.path.join(folder, "run_meta.json"), "w", encoding="utf-8") as fh:
            json.dump({"site": site, "month": MONTH}, fh)
        wb = Workbook()
        ws = wb.active
        ws.title = rb.VENTAS_SHEET
        head = [rb.V_ORDER_ID, rb.V_DATE, rb.V_STATUS, rb.V_UNITS, rb.V_REVENUE, rb.V_COMMISSION,
                rb.V_SHIP_INCOME, rb.V_SHIP_COST, rb.V_REFUND, rb.V_TOTAL, rb.V_SKU, rb.V_TITLE]
        ws.cell(rb.VENTAS_HEADER_ROW, 1, "Ventas")
        for i, h in enumerate(head, 1):
            ws.cell(rb.VENTAS_HEADER_ROW + 1, i, h)
        for r, (oid, st, u, rev, com, ship, ref, sku) in enumerate(ORDERS, rb.VENTAS_HEADER_ROW + 2):
            row = [oid, "%d de septiembre de 2026 10:00 hs." % (r % 20 + 1), st, u, rev, com, 0,
                   ship, ref, rev + com + ship + ref, sku, "Producto " + sku]
            for i, v in enumerate(row, 1):
                ws.cell(r, i, v)
        wb.save(os.path.join(folder, "20261001_Ventas_%s_Mercado_Libre_y_Mercado_Shops_"
                                     "2026-10-01_10-00hs_123456789.xlsx" % site.upper()))
        wb = Workbook()
        ws = wb.active
        ws.title = "REPORT"
        bhead = [rb.F_INVOICE, rb.F_CHARGE_DATE, rb.F_CHARGE_NO, rb.F_DETALLE, rb.F_CHANNEL,
                 rb.F_STATE, rb.F_REVERSES, rb.F_AMOUNT, rb.F_SALE_NO, rb.F_PACK_NO, rb.F_SALE_DATE]
        for i, h in enumerate(bhead, 1):
            ws.cell(rb.BILLING_HEADER_ROW + 1, i, h)
        bills = [("F1", "2026-09-05", "9001", "Cargo por venta", "Si", "Facturado", None, 24,
                  ORDERS[0][0], None, "2026-09-05"),
                 ("F1", "2026-09-05", "9002", "Cargo por venta", "Si", "Facturado", None, 36,
                  ORDERS[1][0], None, "2026-09-05")]
        for r, row in enumerate(bills, rb.BILLING_HEADER_ROW + 2):
            for i, v in enumerate(row, 1):
                ws.cell(r, i, v)
        wb.save(os.path.join(folder, "Reporte_Facturacion_MercadoLibre_Sep2026.xlsx"))
    finally:
        rb.set_site(before)
    return folder


def costs_file(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(HEAD)
    for r in rows:
        ws.append(list(r))
    wb.save(str(path))
    return str(path)


def configure(monkeypatch, tmp_path, cost_path):
    cfg = tmp_path / "cfg_with_costs.json"
    cfg.write_text(json.dumps({"stores": {"default": "X", "allowed": []},
                               "reports": {"sku_costs": cost_path}}), encoding="utf-8")
    monkeypatch.setattr(store_config, "CONFIG_PATH", str(cfg))


def no_config(monkeypatch, tmp_path):
    monkeypatch.setattr(store_config, "CONFIG_PATH", str(tmp_path / "no_config.json"))


def build(folder, tmp_path, site, name="ml.xlsx", stores=None, recalc=False):
    out = str(tmp_path / "out" / name)
    metrics = str(tmp_path / "out" / (name + ".json"))
    argv = ["--month", MONTH, "--out", out, "--site", site, "--metrics-out", metrics]
    for s, f in stores or [(STORE, folder)]:
        argv += ["--store", "%s=%s" % (s, f)]
    if not recalc:
        argv.append("--no-recalc")
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        rc = rb.run(argv)
    text = buf.getvalue()
    line = [x for x in text.splitlines() if x.startswith("[REPORT_CHECKS_JSON] ")]
    checks = json.loads(line[0].split(" ", 1)[1]) if line else []
    with io.open(metrics, encoding="utf-8") as fh:
        recs = json.load(fh)
    return rc, checks, dict((r["store"], r) for r in recs), out


def by_key(checks):
    out = {}
    for c in checks:
        out.setdefault(c["key"], c)
    return out


def pl_rows(ws):
    return dict((ws.cell(row=i, column=1).value, i) for i in range(1, ws.max_row + 1))


def table5(ws):
    """{SKU: row} of ⑦ 表五, plus the header row and the store-total row."""
    t5 = [r for r in range(1, ws.max_row + 1)
          if str(ws.cell(row=r, column=1).value or "").startswith("表五")][0]
    head = [r for r in range(t5, ws.max_row + 1) if ws.cell(row=r, column=1).value == "店铺"][0]
    tot = [r for r in range(head, ws.max_row + 1)
           if str(ws.cell(row=r, column=2).value or "").endswith("合计（＝表一同名行）")][0]
    rows = dict((ws.cell(row=r, column=2).value, r) for r in range(head + 1, tot))
    return head, rows, tot


def readme_labels(ws):
    return dict((ws.cell(row=i, column=2).value, ws.cell(row=i, column=3).value)
                for i in range(1, ws.max_row + 1))


@pytest.fixture
def mx_folder(tmp_path):
    return make_folder(tmp_path / "downloads" / STORE / "20261001_100000", "mx")


@pytest.fixture
def cl_folder(tmp_path):
    return make_folder(tmp_path / "downloads" / "ML_CL_TEST" / "20261001_100000", "cl")


# ------------------------------------------------------------------ the units basis

def test_sku_list_is_table_five_and_keeps_the_site(mx_folder, cl_folder):
    rb.set_site("pe")
    with contextlib.redirect_stdout(io.StringIO()):
        p, c, cur, items = rb.sku_list(STORE, mx_folder, MONTH)
    assert (p, c, cur) == ("ML", "MX", "MXN")
    assert dict((it.sku, (it.units, it.returned, it.sales)) for it in items) == {
        "SKU-A": (2, 0, 200.0), "SKU-B": (1, 0, 300.0), "SKU-C": (3, 0, 90.0)}
    assert rb.SITE == "pe"                       # put back as it was
    with contextlib.redirect_stdout(io.StringIO()):
        p, c, cur, items = rb.sku_list("ML_CL_TEST", cl_folder, MONTH)
    assert (p, c, cur) == ("ML", "CL", "CLP") and len(items) == 3


# ------------------------------------------------------------------ Mexico

def test_mx_without_a_table_is_as_before(mx_folder, tmp_path, monkeypatch):
    no_config(monkeypatch, tmp_path)
    rc, checks, m, out = build(mx_folder, tmp_path, "mx")
    assert not any(c["key"].startswith("sku_costs") for c in checks)
    assert "cogs" not in m[STORE]
    wb = load_workbook(out)
    pl = wb[SHEETS[1]]
    at = pl_rows(pl)["减：商品采购成本"]
    assert pl.cell(row=at, column=3).value == 0
    assert pl.cell(row=at, column=6).value.startswith("⚠ 平台数据不含此项")
    sk = wb[SHEETS[6]]
    assert sk.max_column == 30
    head, rows, _tot = table5(sk)
    assert sk.cell(row=rows["SKU-A"], column=13).value == 0
    assert sk.cell(row=rows["SKU-A"], column=13).fill.fgColor.rgb.endswith(K.C_INPUT)
    assert "补齐商品采购成本" in [c.value for c in wb[SHEETS[9]]["B"]]


def test_mx_full_coverage(mx_folder, tmp_path, monkeypatch):
    configure(monkeypatch, tmp_path, costs_file(tmp_path / "c.xlsx", [
        ["ML", "MX", "sku-a", 50, "MXN"], ["", "MX", "SKU-B", 80, "MXN"], ["ML", "", "SKU-C", 10, "MXN"]]))
    rc, checks, m, out = build(mx_folder, tmp_path, "mx")
    C = by_key(checks)
    assert C["sku_costs_file"]["ok"] and C["sku_costs_currency"]["ok"]
    assert C["sku_costs_coverage"]["ok"] and "SKU 3/3、件数 6/6" in C["sku_costs_coverage"]["detail"]
    # A 2 × 50 (the cancelled unit has no cost), B 1 × 80 (the resellable return has none),
    # C 3 × 10 (not reviewed yet: cost kept)
    assert m[STORE]["cogs"] == 210
    assert m[STORE]["profit"] == pytest.approx(m[STORE]["contribution"] - 210)
    wb = load_workbook(out)
    pl = wb[SHEETS[1]]
    rows = pl_rows(pl)
    c = pl.cell(row=rows["减：商品采购成本"], column=3)
    assert c.value == 210 and c.fill.fgColor.rgb.endswith(K.C_INPUT)        # still an input cell
    note = pl.cell(row=rows["减：商品采购成本"], column=6).value
    assert "c.xlsx" in note and "100.0% 销售额已有成本" in note
    assert pl.cell(row=rows["税前净利润"], column=3).value == "=C%d-C%d" % (
        rows["经营贡献毛利（未扣商品成本）"], rows["减：商品采购成本"])
    sk = wb[SHEETS[6]]
    head, t5, tot = table5(sk)
    assert sk.cell(row=head, column=31).value == "成本来源"
    for sku, unit in (("SKU-A", 50), ("SKU-B", 80), ("SKU-C", 10)):
        assert sk.cell(row=t5[sku], column=5).value == UNITS[sku]
        assert sk.cell(row=t5[sku], column=13).value == unit
        assert sk.cell(row=t5[sku], column=14).value == "=E{0}*M{0}".format(t5[sku])
        assert not sk.cell(row=t5[sku], column=13).fill.fgColor.rgb.endswith(K.C_INPUT)
    assert sk.cell(row=t5["SKU-A"], column=31).value.startswith("成本表第 2 行（ML/MX")
    assert sk.cell(row=t5["SKU-B"], column=31).value.startswith("成本表第 3 行（MX）")
    assert sk.cell(row=t5["取消订单残留费用"], column=31).value == "不计商品成本"
    assert sk.cell(row=t5["退货可再售残留费用"], column=31).value == "不计商品成本"
    rd = readme_labels(wb[SHEETS[0]])
    assert rd["商品成本表"].startswith("c.xlsx")
    assert "表五" in rd["商品成本"] and "到岸成本" in rd["商品成本"] and "IVA" in rd["商品成本"]
    assert "商品采购成本 210" in rd["商品成本与利润"]
    assert "七、商品成本（会计处理）" in rd and "八、重要提示" in rd
    gaps = [c.value for c in wb[SHEETS[9]]["B"]]
    assert "补齐商品采购成本" not in gaps and "补录商品成本表" not in gaps


def test_mx_partial_coverage_and_wrong_currency(mx_folder, tmp_path, monkeypatch):
    configure(monkeypatch, tmp_path, costs_file(tmp_path / "c.xlsx", [
        ["ML", "MX", "SKU-A", 50, "MXN"], ["", "", "SKU-B", 4, "USD"]]))
    rc, checks, m, out = build(mx_folder, tmp_path, "mx")
    C = by_key(checks)
    assert not C["sku_costs_currency"]["ok"] and "SKU-B（USD" in C["sku_costs_currency"]["detail"]
    cov = C["sku_costs_coverage"]
    assert not cov["ok"] and cov["severity"] == "warn"
    assert "SKU-B（1 件，销售额 300）" in cov["detail"] and "SKU-C（3 件" in cov["detail"]
    assert m[STORE]["cogs"] == 100
    wb = load_workbook(out)
    sk = wb[SHEETS[6]]
    _head, t5, _tot = table5(sk)
    b = t5["SKU-B"]
    assert sk.cell(row=b, column=13).value == 0
    assert sk.cell(row=b, column=13).fill.fgColor.rgb.endswith(K.C_INPUT)
    assert sk.cell(row=b, column=31).value.startswith("未录入成本（成本表币种 USD")
    assert sk.cell(row=t5["SKU-C"], column=31).value == "未录入成本"
    pl = wb[SHEETS[1]]
    rows = pl_rows(pl)
    assert pl.cell(row=rows["减：商品采购成本"], column=3).value == 100
    assert "SKU 1/3" in pl.cell(row=rows["减：商品采购成本"], column=6).value
    gaps = wb[SHEETS[9]]
    at = [r for r in range(1, gaps.max_row + 1) if gaps.cell(row=r, column=2).value == "补录商品成本表"]
    assert at and "ML_TEST SKU-B" in gaps.cell(row=at[0], column=3).value


def test_mx_no_coverage_and_missing_file(mx_folder, tmp_path, monkeypatch):
    configure(monkeypatch, tmp_path, costs_file(tmp_path / "c.xlsx", [["ML", "MX", "OTHER", 1, "MXN"]]))
    rc, checks, m, out = build(mx_folder, tmp_path, "mx")
    assert m[STORE]["cogs"] == 0
    assert "SKU 0/3" in by_key(checks)["sku_costs_coverage"]["detail"]
    configure(monkeypatch, tmp_path, str(tmp_path / "gone.xlsx"))
    rc, checks, m, out = build(mx_folder, tmp_path, "mx", name="gone.xlsx")
    assert rc == rb.EXIT_CHECKS and not by_key(checks)["sku_costs_file"]["ok"]
    pl = load_workbook(out)[SHEETS[1]]
    at = pl_rows(pl)["减：商品采购成本"]
    assert pl.cell(row=at, column=3).value == 0 and "读不出来" in pl.cell(row=at, column=6).value


def test_mx_two_stores_and_formulas_evaluate(tmp_path, monkeypatch):
    f1 = make_folder(tmp_path / "d" / "S1" / "20261001_100000", "mx")
    f2 = make_folder(tmp_path / "d" / "S2" / "20261001_100000", "mx")
    configure(monkeypatch, tmp_path, costs_file(tmp_path / "c.xlsx", [
        ["ML", "MX", "SKU-A", 50, "MXN"], ["ML", "MX", "SKU-C", 10.5, "MXN"]]))
    soffice = bool(shutil.which("soffice"))
    rc, checks, m, out = build(None, tmp_path, "mx", stores=[("S1", f1), ("S2", f2)], recalc=soffice)
    assert m["S1"]["cogs"] == m["S2"]["cogs"] == 131.5
    wb = load_workbook(out)
    pl = wb[SHEETS[1]]
    at = pl_rows(pl)["减：商品采购成本"]
    assert [pl.cell(row=at, column=j).value for j in (3, 4)] == [131.5, 131.5]
    assert pl.cell(row=at, column=5).value == "=SUM(C%d:D%d)" % (at, at)
    note = pl.cell(row=at, column=7).value
    assert note.startswith("S1 来自成本表") and "；S2 来自成本表" in note
    if not soffice:
        pytest.skip("no LibreOffice to evaluate the formulas")
    n, n_err, _none, bad = K.scan_formula_errors(out)
    assert n and not n_err, bad
    v = load_workbook(out, data_only=True)
    pl = v[SHEETS[1]]
    rows = pl_rows(pl)
    for col in (3, 4, 5):
        contrib = pl.cell(row=rows["经营贡献毛利（未扣商品成本）"], column=col).value
        assert pl.cell(row=rows["税前净利润"], column=col).value == pytest.approx(
            contrib - pl.cell(row=rows["减：商品采购成本"], column=col).value)
    sk = v[SHEETS[6]]
    t5 = [r for r in range(1, sk.max_row + 1)
          if str(sk.cell(row=r, column=1).value or "").startswith("表五")][0]
    totals = [sk.cell(row=r, column=14).value for r in range(t5, sk.max_row + 1)
              if str(sk.cell(row=r, column=2).value or "").endswith("合计（＝表一同名行）")]
    assert totals == [pytest.approx(131.5)] * 2          # Σ ⑦ cost = ② cost, per store
    t1 = [sk.cell(row=r, column=14).value for r in range(7, t5)
          if str(sk.cell(row=r, column=2).value or "") in ("S1 合计", "S2 合计")]
    assert t1 == [pytest.approx(131.5)] * 2              # 表一 picks the same costs up


# ------------------------------------------------------------------ Chile

def test_cl_full_partial_and_wrong_currency(cl_folder, tmp_path, monkeypatch):
    store = "ML_CL_TEST"
    no_config(monkeypatch, tmp_path)
    rc, checks, m, out = build(cl_folder, tmp_path, "cl", stores=[(store, cl_folder)])
    assert not any(c["key"].startswith("sku_costs") for c in checks)
    assert load_workbook(out)[SHEETS[6]].max_column == 30
    # full
    configure(monkeypatch, tmp_path, costs_file(tmp_path / "c.xlsx", [
        ["ML", "CL", "SKU-A", 5000, "CLP"], ["ML", "CL", "SKU-B", 8000, "CLP"],
        ["", "CL", "SKU-C", 1000, "CLP"], ["ML", "MX", "SKU-A", 50, "MXN"]]))
    rc, checks, m, out = build(cl_folder, tmp_path, "cl", name="full.xlsx", stores=[(store, cl_folder)])
    C = by_key(checks)
    assert C["sku_costs_coverage"]["ok"] and C["sku_costs_currency"]["ok"]
    assert C["sku_costs_currency"]["name"].endswith("（CLP）")
    assert m[store]["cogs"] == 2 * 5000 + 8000 + 3 * 1000
    wb = load_workbook(out)
    pl = wb[SHEETS[1]]
    assert pl.cell(row=pl_rows(pl)["减：商品采购成本"], column=3).value == 21000
    _head, t5, _tot = table5(wb[SHEETS[6]])
    assert wb[SHEETS[6]].cell(row=t5["SKU-A"], column=31).value.startswith("成本表第 2 行（ML/CL")
    assert "IVA" in readme_labels(wb[SHEETS[0]])["商品成本"]
    # partial, and a Mexican-peso row is never used for a Chilean report
    configure(monkeypatch, tmp_path, costs_file(tmp_path / "c2.xlsx", [
        ["ML", "CL", "SKU-B", 8000, "CLP"], ["ML", "", "SKU-A", 50, "MXN"]]))
    rc, checks, m, out = build(cl_folder, tmp_path, "cl", name="part.xlsx", stores=[(store, cl_folder)])
    C = by_key(checks)
    assert m[store]["cogs"] == 8000
    assert not C["sku_costs_currency"]["ok"] and "SKU-A（MXN" in C["sku_costs_currency"]["detail"]
    assert not C["sku_costs_coverage"]["ok"] and "SKU 1/3" in C["sku_costs_coverage"]["detail"]
    assert rc == rb.EXIT_CHECKS


# ------------------------------------------------------------------ the template tool

def test_template_lists_mercadolibre_skus(tmp_path, monkeypatch, capsys):
    from tools import sku_cost_template as T
    base = tmp_path / "downloads"
    make_folder(base / "MX1" / "20261001_100000", "mx")
    make_folder(base / "CL1" / "20261001_100000", "cl")
    cfg = tmp_path / "cfg_tpl.json"
    cfg.write_text(json.dumps({"stores": {"default": "MX1", "allowed": ["MX1", "CL1"],
                                          "regions": {"MX1": "MX", "CL1": "CL"}}}), encoding="utf-8")
    monkeypatch.setattr(store_config, "CONFIG_PATH", str(cfg))
    out = str(tmp_path / "costs.xlsx")
    costs_file(out, [["ML", "MX", "SKU-A", 50, "MXN"]])
    assert T.main(["--out", out, "--month", MONTH, "--downloads", str(base)]) == 0
    assert "跳过" not in capsys.readouterr().out
    ws = load_workbook(out).worksheets[0]
    head = [c.value for c in ws[1]]
    col = dict((h, i) for i, h in enumerate(head))
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    got = [(r[col["平台"]], r[col["国家"]], r[col["SKU"]], r[col["币种"]], r[col["本月件数（参考）"]])
           for r in rows]
    assert got == [("ML", "MX", "SKU-A", "MXN", 2),           # the existing row, refreshed
                   ("ML", "CL", "SKU-B", "CLP", 1), ("ML", "CL", "SKU-A", "CLP", 2),
                   ("ML", "CL", "SKU-C", "CLP", 3),
                   ("ML", "MX", "SKU-B", "MXN", 1), ("ML", "MX", "SKU-C", "MXN", 3)]
    assert rows[0][col["单位成本"]] == 50
    assert rb.SITE == "mx"
    # the written file is a cost table for the MercadoLibre report
    Tb = SC.load(out)
    assert Tb.usable and not Tb.problems and len(Tb.rows) == 1
    assert T.platform_modules("ML") is not None
