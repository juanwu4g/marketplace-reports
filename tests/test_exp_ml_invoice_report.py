# -*- coding: utf-8 -*-
"""experiments/ml_invoice_report — the billing-statement baseline report.

Synthetic download folders only (no real data): invoices with sales, reversals,
store-level charges and an out-of-period row; a sales export with a multi-product
pack; a MercadoPago collection file; a storage detail; invoice payments.
"""
import datetime
import os
import re
import shutil

import pytest
from openpyxl import Workbook, load_workbook

from core import report_kit as K
from experiments.ml_invoice_report import build, catalog, mapping, model, readers, status

D = datetime.datetime

INV_HDR = ["N° de factura fiscal", "Fecha del cargo", "Número del cargo", "Detalle", "Descontado de la operación",
           "Estado del cargo", "Cargo que bonifica", "Valor del cargo", "Valor del descuento", "Número de venta",
           "Pago", "Fecha de venta", "Cantidad vendida", "Precio unitario", "Total de la venta", "Número de envío",
           "Número de paquete", "Envío a cargo del cliente", "Número de publicación", "Título de publicación"]


def inv_row(day, no, det, amount, sale=None, ref="", deducted="Si", month=9):
    """sale = (sale_no, pago, qty, price, buyer_ship, listing, title, pack, sale_day)"""
    r = {"N° de factura fiscal": "A1", "Fecha del cargo": D(2026, month, day), "Número del cargo": no,
         "Detalle": det, "Descontado de la operación": deducted, "Valor del cargo": amount,
         "Cargo que bonifica": ref}
    if sale:
        sn, pago, qty, price, ship, listing, title, pack, sday = sale
        r.update({"Número de venta": sn, "Pago": pago, "Fecha de venta": D(2026, 9 if sday > 0 else 8, abs(sday)),
                  "Cantidad vendida": qty, "Precio unitario": price, "Total de la venta": qty * price,
                  "Número de paquete": pack, "Envío a cargo del cliente": ship, "Número de publicación": listing,
                  "Título de publicación": title})
    return r


def write_invoice(path, rows, header_row=8):
    wb = Workbook()
    ws = wb.active
    ws.title = "REPORT"
    ws.cell(row=1, column=1, value="Fecha de actualización")
    ws.cell(row=1, column=2, value=D(2026, 10, 3, 2, 59))
    ws.cell(row=5, column=1, value="Informacion sobre Mercado Envíos Full")
    for j, h in enumerate(INV_HDR, 1):
        ws.cell(row=header_row, column=j, value=h)
    for i, r in enumerate(rows, header_row + 1):
        for j, h in enumerate(INV_HDR, 1):
            ws.cell(row=i, column=j, value=r.get(h))
    wb.save(path)


SA = ("1001", "9001", 2, 100.0, 10.0, "MLM111", "Bocina A", "", 2)
SB = ("1002", "9002", 1, 300.0, 0.0, "MLM222", "Bocina B", "", 5)
SC_ = ("1003", "9003", 1, 100.0, 0.0, "MLM111", "Bocina A", "5003", 10)     # item of a 2-product pack
SD = ("1004", "9004", 1, 50.0, 0.0, "MLM333", "Cable D", "", 28)          # not in the sales export
SE = ("1005", "9005", 1, 80.0, 0.0, "MLM222", "Bocina B", "", 30)         # reversed after month end
SF = ("1006", "9006", 1, 70.0, 0.0, "MLM222", "Bocina B", "", -30)        # billed in August


def sept_rows():
    return [
        inv_row(31, "c00", "Cargo por venta", 10.5, SF, month=8),            # out of period
        inv_row(2, "c01", "Cargo por venta", 30.0, SA),
        inv_row(2, "c02", "Cargo por envíos de Mercado Libre", 50.0, SA),
        inv_row(5, "c03", "Cargo por venta", 45.0, SB),
        inv_row(5, "c04", "Cargo por envíos de Mercado Libre", 60.0, SB),
        inv_row(6, "c05", "Anulación del cargo por venta", -45.0, SB, ref="c03", deducted="No aplica"),
        inv_row(6, "c06", "Anulación del cargo por envíos de Mercado Libre", -60.0, SB, ref="c04",
                deducted="No aplica"),
        inv_row(10, "c07", "Cargo por venta", 15.0, SC_),
        inv_row(20, "c09", "Cargo por devolución", 20.0, SC_),
        inv_row(28, "c10", "Cargo por venta", 7.5, SD),
        inv_row(29, "c11", "Anulación del cargo por venta", -3.75, SD, ref="c10", deducted="No aplica"),
        inv_row(30, "c12", "Cargo por venta", 12.0, SE),
        inv_row(3, "c20", "Cargo por campaña de publicidad de Product Ads", 100.0, deducted="No"),
        inv_row(1, "c21", "Cargo por mantenimiento de Mi página", 40.0, deducted="No"),
        inv_row(2, "c22", "Cargo por servicio de almacenamiento Full", 9.0, deducted="No"),
        inv_row(12, "c23", "Anulación del cargo por venta", -10.5, SF, ref="c00", deducted="No aplica"),
    ]


def write_nc(path, rows):
    write_invoice(path, rows)


def write_ventas(path):
    hdr = ["# de venta", "Fecha de venta", "Estado", "Descripción del estado", "Paquete de varios productos",
           "Pertenece a un kit", "Unidades", "Ingresos por productos (MXN)", "SKU", "# de publicación",
           "Título de la publicación", "Variante", "Precio unitario de venta de la publicación (MXN)",
           "Forma de entrega", "Fecha en camino", "Fecha entregado", "Revisado por Mercado Libre",
           "Fecha de revisión", "Dinero a favor", "Resultado", "Destino", "Motivo del resultado"]
    rows = [
        ["1001", "2 de septiembre de 2026 10:00 hs.", "Entregado", "", "No", "No", 2, 200, "SKU-A", "MLM111",
         "Bocina A", "", 100, "Mercado Envíos Full", "3 de septiembre", "4 de septiembre", "", "", "", "", "", ""],
        ["1002", "5 de septiembre de 2026 10:00 hs.", "Cancelada por el comprador", "", "No", "No", 1, 300,
         "SKU-B", "MLM222", "Bocina B", "", 300, "Mercado Envíos Full", "", "", "", "", "", "", "", ""],
        ["5003", "10 de septiembre de 2026 10:00 hs.", "Paquete de 2 productos", "", "", "", None, 250, "", "",
         "", "", None, "", "", "", "", "", "", "", "", ""],
        ["1003", "10 de septiembre de 2026 10:00 hs.",
         "Devolución finalizada. Pusimos el producto de nuevo a la venta", "", "Sí", "No", 1, None, "SKU-A",
         "MLM111", "Bocina A", "", 100, "Mercado Envíos Full", "11 de septiembre", "12 de septiembre", "Si",
         "18 de septiembre", "Del comprador", "Apto para la venta", "Vendedor", "está en buenas condiciones"],
        ["1013", "10 de septiembre de 2026 10:00 hs.", "Entregado", "", "Sí", "No", 1, None, "SKU-B", "MLM222",
         "Bocina B", "", 150, "Mercado Envíos Full", "11 de septiembre", "12 de septiembre", "", "", "", "", "",
         ""],
        ["1005", "30 de septiembre de 2026 10:00 hs.", "Entregado", "", "No", "No", 1, 80, "SKU-B", "MLM222",
         "Bocina B", "", 80, "Mercado Envíos Full", "1 de octubre", "2 de octubre", "", "", "", "", "", ""],
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas MX"
    ws.cell(row=1, column=1, value="Ventas")
    ws.cell(row=5, column=1, value="Ventas")         # the group row above the header
    for j, h in enumerate(hdr, 1):
        ws.cell(row=6, column=j, value=h)
    for i, r in enumerate(rows, 7):
        for j, v in enumerate(r, 1):
            ws.cell(row=i, column=j, value=v)
    wb.save(path)


def write_collection(path):
    cols = ["Fecha de compra (date_created)", "Número de operación de Mercado Pago (operation_id)",
            "Número de venta en Mercado Libre (order_id)", "SKU Producto (seller_custom_field)",
            "Identificador de producto (item_id)", "Estado de la operación (status)",
            "Detalle del estado de la operación (status_detail)", "Estado del envío (shipment_status)",
            "Monto devuelto (amount_refunded)", "Valor del producto (transaction_amount)"]
    rows = [["28/09/2026 10:00:00", "9004", "1004", "SKU-D", "MLM333", "approved", "accredited", "delivered",
             "25", "50"]]
    with open(path, "w", encoding="utf-8-sig") as fh:
        fh.write(";".join(cols) + "\n")
        for r in rows:
            fh.write(";".join(r) + "\n")


def write_storage(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws2 = wb.create_sheet("Detalle")
    hdr = ["Código ML", "Código universal", "SKU", "Número de publicación", "Producto", "Tamaño de la unidad",
           "Costos acumulados hasta 30 sep", "01/09/2026", "02/09/2026"]
    for j, h in enumerate(hdr, 1):
        ws2.cell(row=5, column=j, value=h)
    for i, r in enumerate([["ABCD11111", "", "SKU-A", "111", "Bocina A", "Grande", 6, 6.0, 6.0],
                           ["ABCD22222", "", "SKU-B", "222", "Bocina B", "Grande", 3, 3.0, 3.0]], 6):
        for j, v in enumerate(r, 1):
            ws2.cell(row=i, column=j, value=v)
    wb.save(path)


def write_pagos(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Pagos y notas de crédito"
    hdr = ["Número de pago", "Número de nota de crédito", "Tipo de pago", "Medio de pago",
           "Fecha de pago / Emisión de NC", "Estado", "Fecha de anulación", "Importe total",
           "Importe aplicado a este mes", "Importe aplicado a otro mes"]
    for j, h in enumerate(hdr, 1):
        ws.cell(row=10, column=j, value=h)
    for j, v in enumerate(["p1", None, "Pago voluntario", "Dinero", D(2026, 10, 8), "Pago aprobado", None,
                           149.0, 149.0, 0.0], 1):
        ws.cell(row=11, column=j, value=v)
    for j, v in enumerate([None, "B100", "Nota de crédito", None, D(2026, 9, 21), "Facturado", None, 15.0, 15.0,
                           0.0], 1):
        ws.cell(row=12, column=j, value=v)
    ws2 = wb.create_sheet("Detalle de Pagos del mes")
    hdr2 = ["Número del pago", "Fecha del pago", "Parte del pago aplicado a cargos", "Número del cargo",
            "Descripción del cargo", "Fecha del cargo"]
    for j, h in enumerate(hdr2, 1):
        ws2.cell(row=10, column=j, value=h)
    for i, (amt, no, det) in enumerate([(100.0, "c20", "Product Ads"), (40.0, "c21", "Mi página"),
                                        (9.0, "c22", "almacenamiento")], 11):
        for j, v in enumerate(["p1", D(2026, 10, 8), amt, no, det, D(2026, 9, 2)], 1):
            ws2.cell(row=i, column=j, value=v)
    wb.save(path)


@pytest.fixture
def mx_folder(tmp_path):
    f = tmp_path / "STORE_MX" / "20261013_m202609"
    f.mkdir(parents=True)
    write_invoice(str(f / "Reporte_Facturacion_MercadoLibre_Sept2026.xlsx"), sept_rows())
    write_nc(str(f / "Reporte_Notas_Credito_MercadoLibre_Sept2026.xlsx"), [
        dict(inv_row(20, "c08", "Anulación del cargo por venta", -15.0, SC_, ref="c07", deducted="No aplica"),
             **{"N° de factura fiscal": "B100"})])
    write_invoice(str(f / "Reporte_Facturacion_MercadoLibre_Oct2026.xlsx"), [
        inv_row(3, "d01", "Anulación del cargo por venta", -12.0, SE, ref="c12", deducted="No aplica", month=10),
        inv_row(3, "d02", "Cargo por campaña de publicidad de Product Ads", 30.0, deducted="No", month=10)])
    write_nc(str(f / "Reporte_Notas_Credito_MercadoLibre_Oct2026.xlsx"), [])
    write_ventas(str(f / "20261013_Ventas_MX_Mercado_Libre_y_Mercado_Shops_x.xlsx"))
    write_collection(str(f / "collection-20261013-abcd.csv"))
    write_storage(str(f / "01-09-26_30-09-26_Costos_por_servicio_almacenamiento.xlsx"))
    write_pagos(str(f / "Reporte_Pagos_Facturas_Sept2026.xlsx"))
    return str(f)


def derive(folder, site="mx"):
    S = model.load("STORE_MX", folder, "2026-09", site)
    model.derive(S)
    return S


# ------------------------------------------------------------------ readers

def test_readers_small_helpers():
    assert readers.spanish_date("13 de octubre de 2026 13:35 hs.") == datetime.date(2026, 10, 13)
    assert readers.spanish_date("2 de setiembre de 2026") == datetime.date(2026, 9, 2)
    assert readers.file_month("Reporte_Facturacion_MercadoLibre_Sept2026.xlsx") == "2026-09"
    assert readers.file_month("Reporte_Notas_Credito_MercadoLibre_Oct2026.xlsx") == "2026-10"
    assert readers.idstr(2000018220767084.0) == "2000018220767084"
    assert readers.listing_digits("MLM5929009984") == "5929009984"


def test_catalog_unknown_is_flagged():
    e, known = catalog.entry("Cargo por venta")
    assert known and e.line == "fee_commission" and e.kind == "commission"
    e, known = catalog.entry("Cargo misterioso")
    assert not known and e.line == "fee_other"
    assert set(e.line for e in catalog.CATALOG.values()) <= set(catalog.PL_FEE_KEYS)


# ------------------------------------------------------------------ the numbers

def test_pl_from_invoice_rows(mx_folder):
    S = derive(mx_folder)
    pl = S["m"]["pl"]
    assert pl["gmv_billed"] == pytest.approx(200 + 300 + 100 + 50 + 80)
    assert pl["rev_cancel"] == pytest.approx(-300)          # sale B: export says cancelled, nothing shipped
    assert pl["rev_return"] == pytest.approx(-100)          # sale C: returned, reviewed, back on sale
    # sale F: billed in August, reversed on Sep 12 (counts in September); no evidence and 13 days -> unproven
    assert pl["rev_unproven"] == pytest.approx(-70)
    assert pl["rev_partial"] == pytest.approx(-25)          # sale D: half the commission reversed
    assert pl["net_gmv"] == pytest.approx(730 - 495)
    assert pl["ship_income"] == pytest.approx(10)
    ledger = 30 + 50 + 45 + 60 - 45 - 60 + 15 - 15 + 20 + 7.5 - 3.75 + 12 + 100 + 40 + 9 - 10.5
    assert S["m"]["ledger"] == pytest.approx(ledger)
    assert pl["fees_total"] == pytest.approx(-ledger)
    assert pl["fee_commission"] == pytest.approx(-(30 + 45 - 45 + 15 - 15 + 7.5 - 3.75 + 12 - 10.5))
    # post-period: sale E reversed on Oct 3 stays in September GMV, listed as a memo
    assert S["m"]["post_n"] == 1 and S["m"]["post_amount"] == pytest.approx(-80)
    for k, v in model.residuals(S).items():
        assert abs(v) < 1e-6, k


def test_status_and_sku_mapping(mx_folder):
    S = derive(mx_folder)
    st = dict((sn, S["status"][sn]["cls"]) for sn in S["sales"])
    assert st["1001"] == status.COMPLETED
    assert st["1002"] == status.CANCEL
    assert st["1003"] == status.RETURN
    assert st["1004"] == status.PARTIAL
    tiers = dict((sn, S["map"][sn].tier) for sn in S["sales"])
    assert tiers["1001"] == mapping.T1
    assert tiers["1003"] == mapping.T1                      # item of a pack, by its own sale number
    assert tiers["1004"] == mapping.T1B                     # only MercadoPago knows it
    assert S["sales"]["1004"]["sku"] == "SKU-D"
    assert S["m"]["mapped_share"] == pytest.approx(1.0)
    assert S["status"]["1003"]["back_to_stock"]


def test_store_level_costs_direct_and_allocated(mx_folder):
    S = derive(mx_folder)
    sk = S["sku"]
    # storage: charge of Sep 2 = storage of Sep 1, split 6:3 between SKU-A and SKU-B
    assert sk["SKU-A"]["d:fee_full_storage"] == pytest.approx(-6.0)
    assert sk["SKU-B"]["d:fee_full_storage"] == pytest.approx(-3.0)
    # ads + Mi página allocated by net GMV share (negative net GMV counts as 0)
    alloc = sum(x["alloc"] for x in sk.values())
    assert alloc == pytest.approx(-140.0)
    base = dict((k, max(0.0, x["net_gmv"])) for k, x in sk.items())
    tot = sum(base.values())
    assert sk["SKU-A"]["a:fee_ads"] == pytest.approx(-100.0 * base["SKU-A"] / tot)
    assert sum(x["contrib"] for x in sk.values()) == pytest.approx(S["m"]["pl"]["contrib"])


def test_unknown_detalle_is_an_error_but_ties(mx_folder):
    p = os.path.join(mx_folder, "Reporte_Facturacion_MercadoLibre_Sept2026.xlsx")
    write_invoice(p, sept_rows() + [inv_row(15, "c99", "Cargo misterioso", 5.0, deducted="No")])
    S = derive(mx_folder)
    checks = model.validate(S)
    bad = [c for c in checks if c["key"] == "unknown_detalle"][0]
    assert not bad["ok"] and bad["severity"] == "error"
    assert S["m"]["pl"]["fee_other"] == pytest.approx(-5.0)
    assert all(abs(v) < 1e-6 for v in model.residuals(S).values())


def test_payments_bridge(mx_folder):
    S = derive(mx_folder)
    doc = [d for d in S["docs"] if d["name"].endswith("Sept2026.xlsx")][0]
    assert doc["invoiced"] == pytest.approx(149.0) and doc["paid"] == pytest.approx(149.0)
    assert S["nc_match"]["matched"] and not S["nc_match"]["nc_only"]


# ------------------------------------------------------------------ Chile: two invoices make one month

def test_chile_month_from_two_invoices(tmp_path):
    f = tmp_path / "CL"
    f.mkdir()
    s1 = ("2001", "8001", 1, 10000.0, 0.0, "MLC1", "Producto", "", 3)
    s2 = ("2002", "8002", 1, 20000.0, 0.0, "MLC1", "Producto", "", 20)
    s0 = ("2000", "8000", 1, 5000.0, 0.0, "MLC1", "Producto", "", -20)
    write_invoice(str(f / "Reporte_Facturacion_MercadoLibre_Sept2026.xlsx"), [
        inv_row(20, "x0", "Cargo por venta", 500.0, s0, month=8),
        inv_row(3, "x1", "Cargo por venta", 1000.0, s1)])
    write_invoice(str(f / "Reporte_Facturacion_MercadoLibre_Oct2026.xlsx"), [
        inv_row(20, "x2", "Cargo por venta", 2000.0, s2),
        inv_row(5, "x3", "Cargo por campaña de publicidad de Product Ads", 300.0, deducted="No", month=10)])
    S = model.load("DAB", str(f), "2026-09", "cl")
    model.derive(S)
    assert S["m"]["pl"]["gmv_billed"] == pytest.approx(30000.0)
    assert S["m"]["ledger"] == pytest.approx(3000.0)
    assert S["m"]["pl"]["fees_total"] == pytest.approx(-3000.0)   # Aug 20 and Oct 5 rows are not September
    os.remove(str(f / "Reporte_Facturacion_MercadoLibre_Oct2026.xlsx"))
    S = model.load("DAB", str(f), "2026-09", "cl")
    model.derive(S)
    ck = dict((c["key"], c) for c in model.validate(S))
    assert not ck["invoice_cover"]["ok"] and ck["invoice_cover"]["severity"] == "error"


def test_chile_cover_passes_with_cut_day(tmp_path):
    f = tmp_path / "CL2"
    f.mkdir()
    rows_sep = [inv_row(d, "s%d" % d, "Cargo por servicio de almacenamiento Full", 10.0, deducted="No",
                        month=(8 if d >= 8 else 9)) for d in (8, 20, 31, 1, 7)]
    rows_oct = [inv_row(d, "o%d" % d, "Cargo por servicio de almacenamiento Full", 10.0, deducted="No",
                        month=(9 if d >= 8 else 10)) for d in (8, 30, 1, 7)]
    write_invoice(str(f / "Reporte_Facturacion_MercadoLibre_Sept2026.xlsx"), rows_sep)
    write_invoice(str(f / "Reporte_Facturacion_MercadoLibre_Oct2026.xlsx"), rows_oct)
    S = model.load("DAB", str(f), "2026-09", "cl")
    model.derive(S)
    ck = dict((c["key"], c) for c in model.validate(S))
    assert ck["invoice_cover"]["ok"], ck["invoice_cover"]["detail"]
    assert S["m"]["ledger"] == pytest.approx(40.0)          # Sep 1, 7 (Sept file) + Sep 8, 30 (Oct file)


# ------------------------------------------------------------------ workbook

def test_workbook_is_a_sister(mx_folder, tmp_path):
    res = build.run("STORE_MX", mx_folder, "2026-09", "mx", out_dir=str(tmp_path / "out"), recalc=False)
    assert os.path.basename(res["out"]) == "STORE_MX_2026-09_账单基准报表.xlsx"
    wb = load_workbook(res["out"])
    assert wb.sheetnames == K.SHEET_NAMES
    labels = [c.value for c in wb[K.SHEET_NAMES[1]]["A"] if c.value]
    for want in ("商品销售额（本期计费的销售）", "减：发货前取消", "减：退货退款", "净商品销售额", "平台费用合计",
                 "经营贡献（未扣商品成本）"):
        assert want in labels
    n_resid = sum(1 for ws in wb.worksheets for row in ws.iter_rows()
                  for c in row[:3] if isinstance(c.value, str) and c.value.strip().startswith("残差"))
    assert n_resid >= 20


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_workbook_recalculates_with_zero_residuals(mx_folder, tmp_path):
    res = build.run("STORE_MX", mx_folder, "2026-09", "mx", out_dir=str(tmp_path / "out"), recalc=True)
    n_f, n_err, n_none, bad = res["scan"]
    assert n_err == 0, bad
    wb = load_workbook(res["out"], data_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            if any(isinstance(c.value, str) and c.value.strip().startswith("残差") for c in row[:3]):
                for c in row:
                    if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                        assert abs(c.value) < 1e-6, (ws.title, c.coordinate, c.value)


def test_sku_cost_table(mx_folder, tmp_path):
    p = tmp_path / "costs.csv"
    p.write_text("平台,国家,SKU,单位成本,币种\nML,MX,SKU-A,40,MXN\nML,MX,SKU-B,90,MXN\n", encoding="utf-8")
    res = build.run("STORE_MX", mx_folder, "2026-09", "mx", out_dir=str(tmp_path / "out"), sku_costs=str(p),
                    recalc=False)
    cost = res["S"]["cost"]
    # SKU-A: 3 units billed, 1 returned and back on sale -> 2 cost units
    assert cost.line("SKU-A").cost_units == 2 and cost.line("SKU-A").cost == pytest.approx(80.0)
    # SKU-B: 300 (cancelled, nothing shipped) + 80 -> 1 cost unit
    assert cost.line("SKU-B").cost_units == 1
    assert cost.cogs == pytest.approx(80.0 + 90.0)


def test_cli_needs_a_site_for_old_folders(mx_folder, capsys):
    from experiments.ml_invoice_report.__main__ import main
    assert main(["--store", "S", "--month", "2026-09", "--folder", mx_folder]) == 2
    assert "--site" in capsys.readouterr().out


def test_independent_of_the_official_report():
    here = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "experiments",
                        "ml_invoice_report")
    for name in os.listdir(here):
        if name.endswith(".py"):
            src = open(os.path.join(here, name), encoding="utf-8").read()
            assert not re.search(r"^\s*(from|import)\s+platforms", src, re.M), name
            assert "ml_invoice_basis" not in src, name


def test_status_rules():
    ev = status.Evidence(None, "refunded", "bpp_refunded", "cancelled", 100.0, None, 0.0)
    assert status.classify_reversal(1.0, ev, 1)[0] == status.CANCEL       # refunded before dispatch
    ev = status.Evidence(None, "refunded", "bpp_refunded", "delivered", 100.0, None, 0.0)
    assert status.classify_reversal(1.0, ev, 9)[0] == status.RETURN
    ev = status.Evidence()
    assert status.classify_reversal(1.0, ev, 2)[2] == "推定"
    assert status.classify_reversal(1.0, ev, 12)[0] == status.UNPROVEN
    assert status.classify_reversal(0.4, ev, 2)[0] == status.PARTIAL
