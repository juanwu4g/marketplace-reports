# -*- coding: utf-8 -*-
"""Amazon Mexico report on synthetic files shaped exactly like the real exports.

A tiny September 2026 (all names, ids, RFCs and amounts made up):

  order A   1 Sep, SKU-A x1: price 1 000 + shipping 50, promotion −100, fees
  order B  15 Sep, SKU-B x2: price 2 000, fees (inside the fee period 14–28 Sep)
  refund A 20 Sep: everything of A back, fees refunded net of the admin fee
  store level: ads (Amazon's own net/IVA split), subscription, storage,
            long-term storage, inbound transport
  FBA reimbursement +300 and a clawback −50 (SKU-B)
  Ajuste without description: ISR of August −150, IVA adjustment +1.23
  payout −1 000

The summary PDF is replaced by its pdftotext text (report_build.pdf_text is
patched); fee CFDIs and the withholding certificate are small XML files with the
attributes the report reads.
"""
import json
import os

import pytest
from openpyxl import load_workbook

from core import report_kit
from platforms.amazon import report_build as R
from platforms.amazon import run_reports
from platforms.amazon import tx_catalog as TC

MONTH = "2026-09"
STORE = "TEST_AZ_1"
STORE2 = "TEST_AZ_2"

PREAMBLE = [
    '"Incluye  transacciones de Amazon Marketplace, Fulfillment by Amazon (FBA) y Amazon Webstore"',
    '"Todos los importes en MXN, a menos que se especifique"',
    '"Definiciones:"',
    '"fecha/hora: fecha y hora publicadas de la transacción"',
    '"Impuesto de venta cobrado: Incluye impuesto de venta cobrado a los compradores."',
    '"Tarifas de venta: Incluye tarifas de cierre variable y de remisión."',
    '"Tarifas de otra transacción: Incluye reintegros de envío."',
    '"Otro: Incluye otros importes de transacción no relacionados con el pedido."',
    '"Nota: los reportes se indexan por fecha de publicación."',
]
HEAD = ["fecha/hora", "Id. de liquidación", "tipo", "Id. del pedido", "sku", "descripción", "cantidad",
        "marketplace", "cumplimiento", "ciudad del pedido", "estado del pedido", "código postal del pedido",
        "modelo de recaudación de impuestos"] + TC.MONEY_COLS + ["total", "Estado de la transacción",
                                                                "Fecha de liberación de la transacción"]
COL = {"sales": "ventas de productos", "tax": "impuesto de ventas de productos",
       "ship": "créditos de envío", "ship_tax": "impuesto de abono de envío",
       "promo": "descuentos promocionales", "promo_tax": "impuesto de reembolsos promocionales",
       "wh": "impuesto de retenciones en la plataforma", "sell": "tarifas de venta",
       "fba": "tarifas fba", "other_tx": "tarifas de otra transacción", "otro": "otro"}


def row(when, tipo, order="", sku="", desc="", qty="", state="Lanzado", total=None, **amounts):
    r = dict.fromkeys(HEAD, "")
    r.update({"fecha/hora": when, "Id. de liquidación": "10000000001", "tipo": tipo,
              "Id. del pedido": order, "sku": sku, "descripción": desc, "cantidad": str(qty),
              "Estado de la transacción": state})
    if tipo in ("Pedido", "Reembolso"):
        r.update({"marketplace": "amazon.com.mx", "cumplimiento": "Amazon",
                  "ciudad del pedido": "CIUDAD SINTETICA", "estado del pedido": "ESTADO SINTETICO",
                  "código postal del pedido": "CP-SINTETICO", "modelo de recaudación de impuestos":
                  "MarketplaceFacilitator"})
    s = 0.0
    for c in TC.MONEY_COLS:
        r[c] = "0"
    for k, v in amounts.items():
        r[COL[k]] = "{:,.2f}".format(v)
        s += v
    r["total"] = "{:,.2f}".format(s if total is None else total)
    return r


RFD = "Reembolso de inventario de Logística de Amazon (%s)"
TX = [
    row("1 sep 2026 9:00:00 a.m. GMT-7", "Pedido", "900-0000001-0000001", "SKU-A", "Producto A", 1,
        sales=1000.00, tax=160.00, ship=50.00, ship_tax=8.00, promo=-100.00, promo_tax=-16.00, wh=-80.00,
        sell=-150.80, fba=-58.00),
    row("3 sep 2026 2:00:00 p.m. GMT-7", "Tarifa de inventario FBA", desc="Cargos por transporte de envíos a Amazon",
        otro=-23.20),
    row("5 sep 2026 8:00:00 a.m. GMT-7", "Tarifas de transacciones de Logística de Amazon",
        desc="tarifa por almacenamiento prolongado de FBA", fba=-58.00),
    row("11 sep 2026 1:17:24 p.m. GMT-7", "Ajuste", otro=-150.00),              # ISR of August
    row("13 sep 2026 12:56:58 p.m. GMT-7", "Tarifa de servicio", desc="Suscripción", otro=-600.00),
    row("15 sep 2026 10:00:00 a.m. GMT-7", "Pedido", "900-0000002-0000002", "SKU-B", "Producto B", 2,
        sales=2000.00, tax=320.00, wh=-160.00, sell=-278.40, fba=-116.00),
    row("16 sep 2026 12:23:28 p.m. GMT-7", "Ajuste", otro=1.23),                 # IVA adjustment
    row("16 sep 2026 1:05:41 p.m. GMT-7", "Tarifas de transacciones de Logística de Amazon",
        desc="Tarifa por almacenamiento de inventario de FBA", fba=-116.00),
    row("17 sep 2026 12:02:08 a.m. GMT-7", "Ajuste", sku="SKU-B", desc=RFD % "Extraviado en el almacén",
        otro=300.00),
    row("20 sep 2026 11:00:00 p.m. GMT-7", "Reembolso", "900-0000001-0000001", "SKU-A", "Producto A", 1,
        sales=-1000.00, tax=-160.00, ship=-50.00, ship_tax=-8.00, promo=100.00, promo_tax=16.00,
        sell=116.00),
    row("26 sep 2026 9:57:49 p.m. GMT-7", "Tarifa de servicio", desc="Costo de la publicidad",
        other_tx=-500.00, otro=-80.00),
    row("28 sep 2026 8:53:16 a.m. GMT-7", "Ajuste", sku="SKU-B", desc=RFD % "ajuste general", otro=-50.00),
    row("29 sep 2026 12:22:51 p.m. GMT-7", "Trasferir", desc="Para la cuenta que termina en 123",
        otro=-1000.00),
]

SUMMARY = """                                                       Página 1 de 1
Actividad de la cuenta desde Sep 1, 2026 00:00 CST hasta Sep 30, 2026 23:59 CST     Todos los importes en MXN
  Resúmenes                       Puede incluir transacciones de Amazon Marketplace
                                                                                         Totales
 Ingresos             Ventas netas, crédito y reembolsos                              {ing}
 Gastos               Tarifas netas, incluidas tarifas de servicio de Amazon          {gas}
 Transferencias       Depósitos netos y retiros                                        -1,000.00
 Impuesto             Impuestos netos cobrados sobre los servicios y las ventas          -68.77

  Ingresos                                  2,300.00          Gastos                       -1,914.40
                                 Débitos      Créditos                              Débitos     Créditos
 Ventas de producto FBA                       3,000.00        Tarifas de venta FBA       -429.20
 Reembolsos de venta de producto FBA   -1,000.00              Reembolsos de tarifa de venta          150.80
 Crédito de inventario del FBA                  300.00        Tarifas de administración de reembolso   -34.80
                                                              Costo de la publicidad      -580.00
  Transferencias                           -1,000.00          Impuesto                         -68.77
 Transferencias a cuenta bancaria      -1,000.00              Impuestos de producto, envío y envoltura para regalo retenidos     472.00
                                                              Impuestos de producto, envío y envoltura para regalo reembolsados   -152.00
                                                              Impuesto al valor agregado (IVA) obligatorio de Amazon retenido    -240.00
                                                              Impuesto sobre la Renta retenido     -150.00
                                                              Ajuste del IVA del producto                    1.23
"""

CFDI = """<?xml version="1.0" encoding="UTF-8"?><cfdi:Comprobante xmlns:cfdi="http://www.sat.gob.mx/cfd/4"
 Fecha="2026-09-29T07:16:18" Serie="Seller" SubTotal="{sub}" TipoDeComprobante="{kind}" Total="{tot}" Version="4.0">
 <cfdi:Emisor Nombre="PLATAFORMA SINTETICA" Rfc="XAXX010101000"/>
 <cfdi:Receptor Nombre="VENDEDOR SINTETICO" Rfc="XEXX010101000"/>
 <cfdi:Impuestos TotalImpuestosTrasladados="{iva}"/>
 <cfdi:Complemento>{pago}<tfd:TimbreFiscalDigital xmlns:tfd="http://www.sat.gob.mx/TimbreFiscalDigital" UUID="{uuid}"/></cfdi:Complemento>
</cfdi:Comprobante>"""
PAGO = ('<pago20:Pagos xmlns:pago20="http://www.sat.gob.mx/Pagos20"><pago20:Pago FechaPago="2026-09-20T00:00:00" '
        'Monto="{m}"/></pago20:Pagos>')

CERT = """<?xml version="1.0" encoding="UTF-8"?><retenciones:Retenciones
 xmlns:retenciones="http://www.sat.gob.mx/esquemas/retencionpago/2"
 xmlns:plataformasTecnologicas="http://www.sat.gob.mx/esquemas/retencionpago/1/PlataformasTecnologicas10"
 CveRetenc="26" FechaExp="2026-10-06T02:21:41" Version="2.0">
 <retenciones:Emisor NomDenRazSocE="PLATAFORMA SINTETICA" RfcE="XAXX010101000"/>
 <retenciones:Periodo Ejercicio="2026" MesFin="{mes}" MesIni="{mes}"/>
 <retenciones:Totales MontoTotOperacion="2950.00" MontoTotRet="{tot}">
  <retenciones:ImpRetenidos BaseRet="472.00" ImpuestoRet="002" MontoRet="{iva}" TipoPagoRet="01"/>
  <retenciones:ImpRetenidos BaseRet="2950.00" ImpuestoRet="001" MontoRet="73.75" TipoPagoRet="04"/>
 </retenciones:Totales>
 <retenciones:Complemento>
  <plataformasTecnologicas:ServiciosPlataformasTecnologicas NumServ="2" TotalIVARetenido="{iva}"
   TotalISRRetenido="73.75" TotalIVATrasladado="472.00" MonTotalporUsoPlataforma="429.20" Version="1.0"/>
  <tfd:TimbreFiscalDigital xmlns:tfd="http://www.sat.gob.mx/TimbreFiscalDigital" UUID="00000000-0000-0000-0000-00000000CE27"/>
 </retenciones:Complemento>
</retenciones:Retenciones>"""

STOCK_HEAD = ["snapshot-date", "sku", "fnsku", "asin", "product-name", "condition", "available",
              "inv-age-0-to-90-days", "inv-age-91-to-180-days", "inv-age-181-to-270-days",
              "inv-age-271-to-365-days", "inv-age-365-plus-days", "currency", "estimated-ltsf-next-charge",
              "units-shipped-t7", "units-shipped-t30", "units-shipped-t60", "units-shipped-t90", "your-price",
              "days-of-supply", "estimated-storage-cost-next-month", "inbound-quantity",
              "Reserved FC Processing", "Reserved Customer Order"]


def write_tx(path, rows=TX, head=HEAD):
    lines = list(PREAMBLE) + [",".join('"%s"' % h for h in head)]
    for r in rows:
        lines.append(",".join('"%s"' % str(r.get(h, "")).replace('"', '""') for h in head))
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        fh.write("\r\n".join(lines) + "\r\n")


def write_cfdi(path, kind, sub, iva, uuid, pago=None):
    tot = round(sub + iva, 2) if kind != "P" else 0
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(CFDI.format(kind=kind, sub="%.2f" % sub, iva="%.2f" % iva, tot="%.2f" % tot, uuid=uuid,
                             pago=PAGO.format(m="%.2f" % pago) if pago else ""))


def write_invoices(folder):
    """Fee period 14–28 Sep (UTC 19:07:36): invoice I = fees of order B + storage,
    credit note E = the fee refund of refund A. A payment complement of 20 Sep."""
    i_tot, e_tot = 278.40 + 116.00 + 116.00, 116.00
    docs = [("I", i_tot / 1.16, "00000000-0000-0000-0000-0000000000A1"),
            ("E", e_tot / 1.16, "00000000-0000-0000-0000-0000000000E1")]
    index = []
    for kind, sub, uuid in docs:
        sub = round(sub, 2)
        write_cfdi(os.path.join(folder, "fee_%s.xml" % uuid), kind, sub,
                   round((i_tot if kind == "I" else e_tot) - sub, 2), uuid)
        index.append({"type": "Factura de la tarifa de vendedor para México", "format": "XML", "number": uuid,
                      "start": "2026-09-14T19:07:36", "end": "2026-09-28T19:07:36",
                      "file": "fee_%s.xml" % uuid})
    p = "00000000-0000-0000-0000-0000000000P1"
    write_cfdi(os.path.join(folder, "fee_%s.xml" % p), "P", 0, 0, p, pago=394.40)
    index.append({"type": "Complemento de pago de la tarifa de vendedor para México", "format": "XML",
                  "number": p, "start": "2026-09-20T19:07:36", "end": "2026-09-20T19:07:36",
                  "file": "fee_%s.xml" % p})
    with open(os.path.join(folder, "fee_invoices.json"), "w", encoding="utf-8") as fh:
        json.dump(index, fh, ensure_ascii=False)


def write_cert(path, iva=236.00, mes="09"):
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(CERT.format(iva="%.2f" % iva, tot="%.2f" % (iva + 73.75), mes=mes))


def write_stock(path):
    rows = [["2026-10-19", "SKU-A", "X0", "B0SYNTH0001", "Producto A", "New", "10", "4", "0", "6", "0", "0",
             "MXN", "0.0", "1", "3", "5", "9", "1160.00", "40", "12.50", "5", "1", "0"],
            ["2026-10-19", "SKU-B", "X1", "B0SYNTH0002", "Producto B", "New", "1", "1", "0", "0", "0", "0",
             "MXN", "0.0", "1", "6", "8", "9", "1160.00", "5", "1.00", "0", "0", "1"],
            ["2026-10-19", "SKU-C", "X2", "B0SYNTH0003", "Producto C", "New", "7", "0", "7", "0", "0", "0",
             "MXN", "0.0", "0", "0", "0", "0", "99.00", "", "3.00", "0", "0", "0"]]
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        fh.write(",".join('"%s"' % h for h in STOCK_HEAD) + "\r\n")
        for r in rows:
            fh.write(",".join('"%s"' % v for v in r) + "\r\n")


def make_folder(root, store=STORE, stamp="20261019_120000", month=MONTH):
    y, m = month.split("-")
    d = root / "downloads" / store / ("%s_m%s%s" % (stamp, y, m))
    d.mkdir(parents=True)
    stem = R.periods.file_stem(month)
    write_tx(str(d / ("%sMonthlyTransaction.csv" % stem)))
    (d / ("%sMonthlySummary.pdf" % stem)).write_text(SUMMARY.format(ing="2,300.00", gas="-1,914.40"),
                                                     encoding="utf-8")
    write_invoices(str(d))
    write_cert(str(d / "retenciones_00000000-0000-0000-0000-00000000CE27.xml"))
    write_stock(str(d / "stock_20261019.csv"))
    (d / "run_meta.json").write_text(json.dumps({
        "platform": "AZ", "market": "AZ-MX", "country": "MX", "month": month, "store": store,
        "downloaded_at": "2026-10-19T12:00:00", "missing": [], "flags": {"language_switched": True}}),
        encoding="utf-8")
    return str(d)


@pytest.fixture(autouse=True)
def text_pdf(monkeypatch):
    """The 'PDF' of these tests is the pdftotext text itself."""
    def fake(path):
        with open(path, encoding="utf-8") as fh:
            return fh.read(), "test"
    monkeypatch.setattr(R, "pdf_text", fake)


@pytest.fixture
def folder(tmp_path):
    return make_folder(tmp_path)


def build(folder, tmp_path, store=STORE, **kw):
    out = str(tmp_path / "out" / report_kit.store_workbook_name(store, MONTH))
    r = R.build([(store, folder)], MONTH, out, **kw)
    return r, out


def checks(r):
    out = {}
    for c in r["checks"]:
        out.setdefault(c["key"], c)
    return out


def failed(r, severity="error"):
    return [c for c in r["checks"] if not c["ok"] and not c["skipped"] and c["severity"] == severity]


# ------------------------------------------------------------------ the numbers

def test_pl_lines_net_of_iva(folder, tmp_path):
    r, _ = build(folder, tmp_path)
    m = r["metrics"][STORE]
    pl = m["pl"]
    assert pl["sales"] == 3000.00 and pl["refund"] == -1000.00
    assert pl["shipping"] == 0.0 and pl["promo"] == 0.0           # charged and refunded
    assert m["net_sales"] == 2000.00
    # fee columns are IVA-inclusive: net = gross / 1.16
    assert pl["commission"] == -270.00                         # (−150.80 − 278.40 + 116.00) / 1.16
    assert pl["fba_fee"] == -150.00 and pl["storage"] == -100.00 and pl["storage_lt"] == -50.00
    assert pl["logistics_other"] == -20.00 and pl["subscription"] == -517.24
    assert pl["ads"] == -500.00                                # Amazon's own split, not divided again
    assert pl["other_income"] == 250.00                        # reimbursement 300, clawback −50
    assert pl["unregistered"] == 0.0
    assert m["fees"] == -1607.24 and m["contrib"] == 642.76
    # memo lines: never in the P&L
    assert m["iva_out"] == 320.00 and m["iva_withheld"] == -240.00
    assert m["fee_iva"] == -257.16                             # −43.20 −24 −16 −8 −3.20 −80 −82.76
    assert m["isr"] == -150.00 and m["iva_adj"] == 1.23 and m["payout"] == -1000.00
    assert m["cert_isr"] == -73.75
    assert r["outcome"] == "ok", failed(r) + failed(r, "warn")


def test_bridges_tie(folder, tmp_path):
    r, _ = build(folder, tmp_path)
    m = r["metrics"][STORE]
    assert m["bridge_a"]["resid"] == 0 and m["bridge_a"]["o"] == 2950.00 and m["bridge_a"]["r"] == -950.00
    assert m["tx_total"] == -683.17 and m["balance_calc"] == -683.17 and m["bridge_c_resid"] == 0
    assert m["groups_csv"] == {"ingresos": 2300.00, "gastos": -1914.40, "transferencias": -1000.00,
                               "impuesto": -68.77}
    (B,) = m["bridge_b"]
    assert B["full"] and B["inv"] == 394.40 and B["tx_fees"] == 394.40 and B["diff"] == 0
    assert B["tx_ads"] == 580.00                               # shown next to it, not compared
    assert len(m["invoice_payments"]) == 1
    c = checks(r)
    for k in ("bridge_a", "bridge_c", "bridge_d", "isr_lump", "fee_invoice", "taxcert", "iva_rate",
              "withholding_rate", "row_total", "unregistered", "file_stock", "deferred", "tx_in_month"):
        assert c[k]["ok"] and not c[k]["skipped"], (k, c[k])


def test_isr_lump_classification():
    assert TC.classify("Ajuste", "", "", -11814.82).key == "isr"
    assert TC.classify("Ajuste", "", "", 2.69).key == "iva_adj"
    assert TC.classify("Ajuste", None, None, -100.0).key == "iva_adj"      # threshold is strict
    assert TC.classify("Ajuste", "", "SKU-X", -500.0) is None              # a product without reason
    e = TC.classify("Ajuste", RFD % "Dañado en el almacén", "SKU-X", 10)
    assert e.key == "fba_reimbursement" and e.line == TC.OTHER_INCOME
    assert TC.reimburse_reason(RFD % "Extraviado: envío entrante")[1] == "入库途中丢失"
    assert TC.classify("tarifa de SERVICIO", "costo de la publicidad", "", 0).key == "ads"   # case, accents
    assert TC.column_target(TC.BY_KEY["ads"], "otro") == (TC.ADS, TC.IVA)
    assert TC.column_target(TC.BY_KEY["order"], "otro") == (TC.UNREGISTERED, TC.NET)


def test_summary_mismatch_and_wrong_isr_split_are_errors(folder, tmp_path):
    stem = R.periods.file_stem(MONTH)
    with open(os.path.join(folder, "%sMonthlySummary.pdf" % stem), "w", encoding="utf-8") as fh:
        fh.write(SUMMARY.format(ing="2,310.00", gas="-1,914.40").replace("-150.00", "-160.00"))
    r, out = build(folder, tmp_path)
    c = checks(r)
    assert not c["bridge_d"]["ok"] and "ingresos" in c["bridge_d"]["detail"]
    assert not c["isr_lump"]["ok"]
    assert r["outcome"] == "checks_failed" and os.path.exists(out)


# ------------------------------------------------------------------ errors and gaps

def test_unknown_type_and_unexpected_column_are_errors(folder, tmp_path):
    stem = R.periods.file_stem(MONTH)
    rows = TX + [row("12 sep 2026 9:00:00 a.m. GMT-7", "Cargo nuevo", desc="Algo que no existía", otro=-11.60),
                 row("12 sep 2026 9:30:00 a.m. GMT-7", "Pedido", "900-0000003-0000003", "SKU-A", "Producto A", 1,
                     sales=100.00, tax=16.00, otro=-5.00)]
    write_tx(os.path.join(folder, "%sMonthlyTransaction.csv" % stem), rows)
    r, out = build(folder, tmp_path)
    c = checks(r)
    m = r["metrics"][STORE]
    assert not c["unregistered"]["ok"] and c["unregistered"]["severity"] == "error"
    assert "Cargo nuevo" in c["unregistered"]["detail"] and "otro" in c["unregistered"]["detail"]
    assert m["pl"]["unregistered"] == -16.60                    # kept, not split
    assert m["bridge_c_resid"] == 0                             # still ties
    assert r["outcome"] == "checks_failed"
    assert build(folder, tmp_path, strict=True)[0]["outcome"] == "strict_aborted"


def test_row_total_mismatch_is_an_error(folder, tmp_path):
    stem = R.periods.file_stem(MONTH)
    rows = TX + [row("12 sep 2026 9:00:00 a.m. GMT-7", "Tarifa de servicio", desc="Suscripción", otro=-1.00,
                     total=-2.00)]
    write_tx(os.path.join(folder, "%sMonthlyTransaction.csv" % stem), rows)
    r, _ = build(folder, tmp_path)
    assert not checks(r)["row_total"]["ok"] and not checks(r)["bridge_c"]["ok"]


def test_english_file_is_an_error(folder, tmp_path):
    stem = R.periods.file_stem(MONTH)
    head = ["date/time"] + HEAD[1:]
    write_tx(os.path.join(folder, "%sMonthlyTransaction.csv" % stem), [], head=head)
    r, out = build(folder, tmp_path)
    c = checks(r)
    assert not c["language"]["ok"] and c["language"]["severity"] == "error"
    assert "date/time" in c["language"]["detail"] and "Español" in c["language"]["detail"]
    assert not c["file_order"]["ok"]
    assert os.path.exists(out)


def test_empty_month_no_data(tmp_path):
    d = tmp_path / "downloads" / STORE / "20261019_120000_m202609"
    d.mkdir(parents=True)
    (d / "run_meta.json").write_text(json.dumps({
        "platform": "AZ", "month": MONTH, "missing": [
            {"form": "order", "cause": "NO_DATA"},
            {"form": "taxcert", "cause": "NOT_AVAILABLE", "explain": "凭证还没出", "action": "下次下载",
             "in_report": True}]}), encoding="utf-8")
    r, out = build(str(d), tmp_path)
    c = checks(r)
    m = r["metrics"][STORE]
    assert c["file_order"]["ok"] and "没有交易" in c["file_order"]["detail"]
    assert m["contrib"] == 0 and m["tx_total"] == 0 and m["bridge_c_resid"] == 0
    assert c["bridge_d"]["skipped"] and c["taxcert"]["skipped"]
    assert not failed(r)
    wb = load_workbook(out)
    assert wb.sheetnames == report_kit.SHEET_NAMES


def test_missing_order_file_is_an_error(folder, tmp_path):
    os.remove(os.path.join(folder, "%sMonthlyTransaction.csv" % R.periods.file_stem(MONTH)))
    r, out = build(folder, tmp_path)
    assert not checks(r)["file_order"]["ok"] and r["outcome"] == "checks_failed"
    assert os.path.exists(out)


def test_missing_summary_and_taxcert_are_skipped(folder, tmp_path):
    stem = R.periods.file_stem(MONTH)
    os.remove(os.path.join(folder, "%sMonthlySummary.pdf" % stem))
    os.remove(os.path.join(folder, "retenciones_00000000-0000-0000-0000-00000000CE27.xml"))
    r, out = build(folder, tmp_path)
    c = checks(r)
    assert c["bridge_d"]["skipped"] and c["taxcert"]["skipped"] and "待取得" in c["taxcert"]["detail"]
    assert r["metrics"][STORE]["cert_isr"] == 0.0
    assert r["outcome"] == "ok"
    wb = load_workbook(out)
    pl = wb[report_kit.SHEET_NAMES[1]]
    notes = " ".join(str(c.value) for row_ in pl.iter_rows() for c in row_ if c.value)
    assert "待取得" in notes


def test_unreadable_pdf_is_skipped_not_failed(folder, tmp_path, monkeypatch):
    monkeypatch.setattr(R, "pdf_text", lambda path: (None, "本机没有能读 PDF 的工具"))
    r, _ = build(folder, tmp_path)
    c = checks(r)
    assert c["bridge_d"]["skipped"] and "pdftotext" in c["bridge_d"]["detail"]
    assert r["outcome"] == "ok"


def test_taxcert_difference_and_wrong_month_warn(folder, tmp_path):
    cert = os.path.join(folder, "retenciones_00000000-0000-0000-0000-00000000CE27.xml")
    write_cert(cert, iva=400.00, mes="08")
    r, _ = build(folder, tmp_path)
    c = checks(r)
    assert not c["taxcert"]["ok"] and c["taxcert"]["severity"] == "warn"
    assert not c["taxcert_period"]["ok"]
    assert r["outcome"] == "ok"                                 # warnings only


def test_fee_period_straddling_the_month_is_partial(folder, tmp_path):
    idx = os.path.join(folder, "fee_invoices.json")
    data = json.load(open(idx, encoding="utf-8"))
    for x in data:
        if x["type"].startswith("Factura"):
            x["start"], x["end"] = "2026-09-28T19:07:36", "2026-10-12T19:07:36"
    json.dump(data, open(idx, "w", encoding="utf-8"))
    r, _ = build(folder, tmp_path)
    (B,) = r["metrics"][STORE]["bridge_b"]
    assert not B["full"]
    assert checks(r)["fee_invoice"]["skipped"]


def test_deferred_lines_warn(folder, tmp_path):
    stem = R.periods.file_stem(MONTH)
    rows = [dict(x) for x in TX]
    rows[5]["Estado de la transacción"] = "Diferido"
    write_tx(os.path.join(folder, "%sMonthlyTransaction.csv" % stem), rows)
    r, _ = build(folder, tmp_path)
    c = checks(r)
    assert not c["deferred"]["ok"] and c["deferred"]["severity"] == "warn" and "Diferido" in c["deferred"]["detail"]
    assert r["metrics"][STORE]["pl"]["sales"] == 3000.00        # still in: posting basis


# ------------------------------------------------------------------ the workbook

def test_workbook_sheets_rows_and_no_buyer_location(folder, tmp_path):
    r, out = build(folder, tmp_path)
    wb = load_workbook(out)
    assert wb.sheetnames == report_kit.SHEET_NAMES
    text = " ".join(str(c.value) for ws in wb for row_ in ws.iter_rows() for c in row_ if c.value)
    assert "CIUDAD SINTETICA" not in text and "CP-SINTETICO" not in text and "ESTADO SINTETICO" not in text
    assert "XEXX010101000" not in text
    pl = wb[report_kit.SHEET_NAMES[1]]
    labels = [pl.cell(row=i, column=1).value for i in range(6, 6 + len(R.PL_ROWS))]
    assert labels == [x.label for x in R.PL_ROWS]               # fixed row positions
    assert "FBA 库存赔付" in " ".join(str(c.value) for row_ in wb[report_kit.SHEET_NAMES[5]].iter_rows()
                                     for c in row_ if c.value)
    sku = wb[report_kit.SHEET_NAMES[6]]
    assert sku["P5"].value == "=-O5*D5" and sku["Q5"].value == "=N5+P5"


def test_two_stores_combined_workbook(tmp_path):
    f1 = make_folder(tmp_path, STORE)
    f2 = make_folder(tmp_path, STORE2)
    out = str(tmp_path / "out" / report_kit.combined_workbook_name("MX", "AZ", MONTH))
    assert os.path.basename(out) == "MX_AZ_全店汇总_2026-09.xlsx"
    r = R.build([(STORE, f1), (STORE2, f2)], MONTH, out)
    assert r["outcome"] == "ok"
    wb = load_workbook(out)
    pl = wb[report_kit.SHEET_NAMES[1]]
    assert [pl.cell(row=5, column=j).value for j in (3, 4, 5)] == [STORE, STORE2, "合计"]
    assert pl.cell(row=6, column=5).value == 6000.00            # data rows: sum of stores
    br = wb[report_kit.SHEET_NAMES[2]]
    assert br.cell(row=5, column=4).value == "合计"


def test_history_from_an_earlier_month(folder, tmp_path):
    aug = make_folder(tmp_path, STORE, stamp="20260910_120000", month="2026-08")
    # an August file is named 2026Aug…; the September rows inside do not matter here
    assert os.path.exists(os.path.join(aug, "2026AugMonthlyTransaction.csv"))
    r, out = build(folder, tmp_path)
    wb = load_workbook(out)
    tr = wb[report_kit.SHEET_NAMES[7]]
    assert "2026-08" in [c.value for row_ in tr.iter_rows() for c in row_]


def test_files_from_a_later_run_of_the_same_month(folder, tmp_path):
    later = os.path.join(os.path.dirname(folder), "20261020_090000_m202609")
    os.makedirs(later)
    os.replace(os.path.join(folder, "stock_20261019.csv"), os.path.join(later, "stock_20261020.csv"))
    r, _ = build(folder, tmp_path)
    c = checks(r)
    assert c["file_stock"]["ok"] and "20261020_090000_m202609" in c["file_stock"]["detail"]


def test_parse_when():
    import datetime
    assert R.parse_when("31 ago 2026 11:28:19 p.m. GMT-7") == datetime.datetime(2026, 9, 1, 6, 28, 19)
    assert R.parse_when("1 sept 2026 12:05:00 a.m. GMT-6") == datetime.datetime(2026, 9, 1, 6, 5, 0)
    assert R.parse_when("15 dic 2026 12:00:00 p.m. GMT-8") == datetime.datetime(2026, 12, 15, 20, 0, 0)
    assert R.parse_when("garbage") is None


def test_parse_summary_reads_the_layout():
    v = R.parse_summary(SUMMARY.format(ing="2,300.00", gas="-1,914.40"))
    assert v["ingresos"] == 2300.00 and v["gastos"] == -1914.40
    assert v["transferencias"] == -1000.00 and v["impuesto"] == -68.77
    assert v["isr"] == -150.00 and v["iva_adj"] == 1.23 and v["iva_withheld"] == -240.00
    assert v["refund_admin"] == -34.80 and v["_from"] == "sep 1 2026"
    assert R.parse_summary("Summary Income Expenses Transfers")["_lang"] == "en"


# ------------------------------------------------------------------ run_reports

def test_run_reports_folder_and_exit_codes(folder, tmp_path):
    base = os.path.dirname(os.path.dirname(folder))
    assert run_reports.latest_folder(STORE, MONTH, base) == folder
    later = os.path.join(os.path.dirname(folder), "20261020_090000_m202609")   # stock-only rerun
    os.makedirs(later)
    open(os.path.join(later, "run_meta.json"), "w").write("{}")
    assert run_reports.latest_folder(STORE, MONTH, base) == folder        # the one with the report
    assert run_reports.latest_folder(STORE, "2026-08", base) is None
    usable, unusable = run_reports.resolve([STORE, "OTHER_AZ"], MONTH, base)
    assert usable == [(STORE, folder)] and unusable[0][0] == "OTHER_AZ"
    ok = {"outcome": "ok"}
    assert run_reports.exit_code([ok], usable, [], per_store=True) == 0
    assert run_reports.exit_code([ok], usable, unusable, per_store=True) == 2
    assert run_reports.exit_code([ok], usable, unusable, per_store=False) == 1
    assert run_reports.exit_code([{"outcome": "checks_failed"}], usable, []) == 1
    assert run_reports.exit_code([{"outcome": "crashed"}], usable, []) == 2


def test_a_month_folder_without_the_report_still_resolves(tmp_path):
    d = tmp_path / "downloads" / STORE / "20261019_120000_m202609"
    d.mkdir(parents=True)
    (d / "run_meta.json").write_text("{}", encoding="utf-8")
    assert run_reports.latest_folder(STORE, MONTH, str(tmp_path / "downloads")) == str(d)


def test_names():
    assert report_kit.combined_workbook_name("mx", "az", MONTH) == "MX_AZ_全店汇总_2026-09.xlsx"
    assert report_kit.alert_file_name("MX", "AZ", MONTH) == "！报表异常_MX_AZ_202609.txt"
