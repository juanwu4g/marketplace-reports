# -*- coding: utf-8 -*-
"""core/report_family: the shared sheet writers, on tiny synthetic stores."""
import pytest
from openpyxl import Workbook

from core import report_family as RF
from core import report_kit as K

A = {"store": "A店", "m": {"sales": 1000.0, "fees": -300.0, "orders": 10, "seen": 990.0},
     "meta": {"missing": [{"form": "billing/x", "cause": "UI_NOT_FOUND", "explain": "按钮没找到",
                           "action": "重跑", "in_report": True},
                          {"form": "invoice", "cause": "NO_DATA", "in_report": False},
                          {"cause": "noise without a form"}]}}
B = {"store": "B店", "m": {"sales": 500.0, "fees": -100.0, "orders": 4, "seen": 400.0},
     "meta": {}}

SPEC = [RF.PLRow("sales", "商品销售收入", "Sales", "d", "说明 1"),
        RF.PLRow("fees", "减：平台费用", "Fees", "d", ""),
        RF.PLRow("fee_total", "费用小计", "", "sum", "", RF.Span("fees", "fees")),
        RF.PLRow("contrib", "经营贡献毛利", "", "f", "", ("sales", "fee_total")),
        RF.PLRow("cogs", "减：商品成本", "", "in", "录入"),
        RF.PLRow("profit", "税前净利润", "", "f", "", ("contrib", "-cogs"))]


def pl(stores):
    wb = Workbook()
    wb.remove(wb.active)
    fam = RF.Family(stores, money=K.MNY0)
    sh, r, at, note_col = fam.pl_sheet(wb, K.SHEET_NAMES[1], SPEC,
                                       value=lambda S, k: S["m"][k], title="损益", subtitle="口径")
    return fam, sh, r, at, note_col


def test_pl_one_store_has_no_total_column():
    fam, sh, r, at, note_col = pl([A])
    ws = sh.ws
    assert [ws.cell(row=5, column=j).value for j in range(1, 5)] == ["项目", "平台科目（原文）", "A店", "说明"]
    assert note_col == 4
    assert at == {"sales": 6, "fees": 7, "fee_total": 8, "contrib": 9, "cogs": 10, "profit": 11}
    assert ws["C6"].value == 1000.0 and ws["C6"].number_format == K.MNY0
    assert ws["C8"].value == "=SUM(C7:C7)"
    assert ws["C9"].value == "=C6+C8"
    assert ws["C11"].value == "=C9-C10"
    assert ws["C10"].value == 0 and ws["C10"].fill.fgColor.rgb.endswith(K.C_INPUT)
    assert ws["C8"].fill.fgColor.rgb.endswith(K.C_SUB)      # subtotal
    assert ws["C9"].fill.fgColor.rgb.endswith(K.C_CALC)     # computed
    assert ws["D6"].value == "说明 1"
    assert ws["B6"].value == "Sales"
    assert r == 12


def test_pl_several_stores_add_a_total_column():
    fam, sh, r, at, note_col = pl([A, B])
    ws = sh.ws
    assert [ws.cell(row=5, column=j).value for j in range(3, 7)] == ["A店", "B店", "合计", "说明"]
    assert ws["E6"].value == 1500.0 and ws["E6"].font.b          # data total = sum
    assert ws["E9"].value == "=E6+E8"                            # formulas run in the total column too
    assert ws["E10"].value == "=SUM(C10:D10)"                    # input rows total by formula
    assert ws["E10"].fill.fgColor.rgb in (None, "00000000")      # … and are not yellow
    assert note_col == 6


def test_pl_spec_errors_are_loud():
    wb = Workbook()
    fam = RF.Family([A])
    with pytest.raises(ValueError):
        fam.pl_sheet(wb, "x", [RF.PLRow("a", "A", "", "f", "")], value=lambda S, k: 0,
                     title="t", subtitle="s")
    with pytest.raises(ValueError):
        fam.pl_sheet(wb, "y", [RF.PLRow("a", "A", "", "d", ""), RF.PLRow("a", "A", "", "d", "")],
                     value=lambda S, k: 0, title="t", subtitle="s")
    with pytest.raises(KeyError):          # a formula naming a row that is not above it
        fam.pl_sheet(wb, "z", [RF.PLRow("a", "A", "", "f", "", ("nope",))],
                     value=lambda S, k: 0, title="t", subtitle="s")


def test_ratio_and_memo_blocks():
    fam, sh, r, at, note_col = pl([A, B])
    ws = sh.ws
    r2 = fam.ratio_block(sh, r + 1, "费用率", [
        ("费用", lambda S: -S["m"]["fees"], lambda S: S["m"]["sales"]),
        ("零基数", lambda S: 1.0, lambda S: 0.0)], note_col)
    row = r + 2
    assert ws.cell(row=r + 1, column=1).value == "费用率"
    assert ws.cell(row=r + 1, column=1).fill.fgColor.rgb.endswith(K.C_BAND)
    assert ws.cell(row=row, column=3).value == pytest.approx(0.3)
    assert ws.cell(row=row, column=5).value == pytest.approx(400 / 1500.0)   # Σ/Σ, not mean
    assert ws.cell(row=row, column=3).number_format == K.PCT
    assert ws.cell(row=row + 1, column=3).value == 0.0
    r3 = fam.memo_block(sh, r2 + 1, "备查项", [("订单数", "", lambda S: S["m"]["orders"], "n")], note_col)
    assert ws.cell(row=r2 + 2, column=5).value == 14
    assert ws.cell(row=r2 + 2, column=note_col).value == "n"
    assert r3 == r2 + 3


def test_bridge_lines_styles_and_residual():
    wb = Workbook()
    fam = RF.Family([A, B], money=K.MNY0)
    L = RF.Line
    ws = fam.bridge_sheet(wb, K.SHEET_NAMES[2], "平账", "残差应为 0", [
        ("桥 A", [L("订单", lambda S: S["m"]["sales"], "来源"),
                 L("＝ 推算", lambda S: S["m"]["sales"], style="subtotal"),
                 L("结算", lambda S: S["m"]["seen"]),
                 L("残差（应为 0）", lambda S: S["m"]["sales"] - S["m"]["seen"], style="residual")]),
        ("桥 B", [L("待开票", lambda S: 1.0, "下月", style="pending")])])
    assert [ws.cell(row=5, column=j).value for j in range(1, 6)] == ["项目", "A店", "B店", "合计", "说明"]
    assert ws["A6"].value == "桥 A"
    res = 10
    assert ws.cell(row=res, column=1).value == "残差（应为 0）"
    assert ws.cell(row=res, column=4).value == 110.0
    for col in (2, 3, 4):
        c = ws.cell(row=res, column=col)
        assert c.font.b and c.fill.fgColor.rgb.endswith(K.C_TOTAL)
    assert ws.cell(row=8, column=2).fill.fgColor.rgb.endswith(K.C_SUB)
    assert ws.cell(row=7, column=5).value == "来源"
    assert ws.cell(row=11, column=1).value is None                  # empty row between bridges
    assert ws["A12"].value == "桥 B"
    assert ws.cell(row=13, column=2).fill.fgColor.rgb.endswith(K.C_WARN)
    assert not ws.cell(row=13, column=2).font.b
    with pytest.raises(ValueError):
        fam.bridge_sheet(wb, "x", "t", "s", [("b", [L("x", lambda S: 0, style="fancy")])])


def test_long_table_columns():
    wb = Workbook()
    fam = RF.Family([A], money=K.MNY0)
    cols = [RF.Col("店铺", 20), RF.Col("金额", 14, RF.MONEY), RF.Col("率", 9, K.PCT, border=False),
            RF.Col("成本（录入）", 12, RF.MONEY, fill=K.C_INPUT),
            RF.Col("毛利", 12, RF.MONEY, border=False, fill=K.C_CALC, formula="=B{r}-D{r}")]
    ws = fam.table_sheet(wb, "④", "表", "副标题", cols)
    assert ws["A1"].value == "表" and ws["A2"].value == "副标题"
    assert [c.value for c in ws[4]] == ["店铺", "金额", "率", "成本（录入）", "毛利"]
    assert ws.column_dimensions["E"].width == 12
    cells = fam.write_row(ws, 5, cols, ["A店", 10.0, 0.5, 0, None])
    assert ws["B5"].number_format == K.MNY0 and ws["B5"].border.left.style == "thin"
    assert ws["C5"].number_format == K.PCT and ws["C5"].border.left.style is None
    assert ws["D5"].fill.fgColor.rgb.endswith(K.C_INPUT)
    assert ws["E5"].value == "=B5-D5"
    assert len(cells) == 5


def test_readme_sections_and_warning():
    wb = Workbook()
    fam = RF.Family([A])
    ws = fam.readme_sheet(wb, K.SHEET_NAMES[0], "标题", "生成于 x", [
        ("一、", [("a", "文字"), ("b", "长" * 120)]), ("二、", [("c", "")])], warning="有 1 项 error")
    assert ws["B1"].value == "标题"
    assert ws["B4"].value == "一、" and "B4:C4" in [str(m) for m in ws.merged_cells.ranges]
    assert ws["C5"].value == "文字"
    assert ws.row_dimensions[6].height > 15                      # long text gets a taller row
    assert ws["B7"].value == "" and ws["B8"].value == "二、"     # empty row between sections
    assert ws["B11"].value == "⚠ 注意" and ws["C11"].value == "有 1 项 error"
    ws2 = fam.readme_sheet(wb, "x", "t", "s", [("一、", [("a", "b")])])
    assert ws2["B7"].value is None


def test_gaps_sheet_orders_checks_and_lists_missing_files():
    wb = Workbook()
    fam = RF.Family([A, B])
    checks = [K.check("ok", "A店", "通过的", True, "d"),
              K.check("w", "A店", "警告", False, "d", severity="warn"),
              K.check("e", "B店", "错误", False, "d"),
              K.check("s", "B店", "跳过的", True, "d", skipped=True)]
    ws = fam.gaps_sheet(wb, K.SHEET_NAMES[9], "缺口", checks,
                        notes=[("报税提示", [("提示 1", "文字")])])
    assert [ws.cell(row=r, column=2).value for r in range(5, 9)] == ["错误", "警告", "通过的", "跳过的"]
    assert ws["D5"].value == "未通过" and ws["A5"].fill.fgColor.rgb.endswith(K.C_BAD)
    assert ws["A6"].fill.fgColor.rgb.endswith(K.C_WARN)
    assert ws["D8"].value == "跳过" and ws["A8"].fill.fgColor.rgb.endswith(K.C_OK)
    assert ws["A10"].value.startswith("没有下载到的文件")
    assert (ws["A11"].value, ws["B11"].value, ws["D11"].value) == ("A店", "billing/x", "影响报表")
    assert ws["E11"].value == "按钮没找到　下一步：重跑"
    assert ws["D12"].value == "报表不读"
    assert ws["A14"].value == "报税提示" and ws["B15"].value == "提示 1" and ws["E15"].value == "文字"


def test_gaps_sheet_says_so_when_nothing_is_missing():
    wb = Workbook()
    ws = RF.Family([B]).gaps_sheet(wb, "⑩", "缺口", [])
    assert ws["B7"].value == "没有。应下的文件全部到齐。"


def test_per_store_and_not_applicable():
    assert RF.Family([A]).per_store(lambda S: "x") == "x"
    assert RF.Family([A, B]).per_store(lambda S: S["m"]["orders"]) == "A店 10；B店 4"
    wb = Workbook()
    ws = RF.not_applicable_sheet(wb, K.SHEET_NAMES[5], "退货质检", "不适用：没有这种报表")
    assert ws["A3"].value.startswith("不适用")


def test_core_report_family_names_no_platform():
    src = open(RF.__file__, encoding="utf-8").read().lower()
    for word in ("falabella", "mercadolibre", "美客多", "clp", "iva"):
        assert word not in src, word
