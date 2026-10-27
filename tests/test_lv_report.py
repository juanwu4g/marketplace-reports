# -*- coding: utf-8 -*-
"""Liverpool's report on a synthetic month (no real export: real ones carry
customer data). The folder has every file the downloads write, with the real
column names, and a story whose numbers are worked out by hand:

  O1  SKU-A 1,160 (with IVA), ordered and posted 2 Sept, commission 197.20,
      card surcharge 41.18, withheld IVA 80 / ISR 25 (8 % / 2.5 % of 1,000)
  O2  SKU-B 2,320, Liverpool promotion 232: the ledger has the sale at 2,088 and a
      separate promotion credit of 232 with the same reference; commission on 2,320
  O3  SKU-A 1,160 posted 10 Sept and cancelled 12 Sept: a return document gives
      back the price, the commission and the withholding
  O4  500, cancelled before it was paid: never posted
  fees: shipping 300 (orders O1 100, O2 200, document 1800000001) + 50 not in the
      shipping list; extra shipping 60 (O2); subscription 600; a logistics refund 40
  payments: one attempt sent back by the bank, one payment of 2,000 net of 900 of fees
"""
import datetime
import os
import shutil

import pandas as pd
import pytest
from openpyxl import load_workbook

from core import report_kit as K
from platforms.liverpool import catalog as C
from platforms.liverpool import readers as R
from platforms.liverpool import report_build as RB

STORE = "LV_TEST"
MONTH = "2026-09"
LEDGER_COLS = ["ID", "Soc.", "Lib.mayor", "Acreedor", "Nº doc.", "Año", "Asignación", "BP", "CT",
               "Clase doc.", "VP", "Texto", "II", "Doc.comp.", "Referencia", "Fe.pago", "Fe.contab.",
               "Fecha doc.", "Importe Debe/Haber", "ML", "Impte.mon.extranj.", "Mon.", "Compens."]
P, Rc = C.PAYABLE, C.RECEIVABLE


def D(day):
    return datetime.datetime(2026, 9, day)


def L(account, doc, cls, text, tax, ref, day, amount, clearing=""):
    return {"Lib.mayor": account, "Nº doc.": doc, "Clase doc.": cls, "Texto": text, "II": tax,
            "Referencia": ref, "Fe.contab.": D(day), "Fecha doc.": D(day), "Importe Debe/Haber": amount,
            "Doc.comp.": clearing, "Mon.": "MXP", "Soc.": "110"}


LEDGER = [
    # O1
    L(P, "9500000001", "YV", "F/ 90000001  21 625                0021", "D2", "0590000001", 2, -1160),
    L(P, "9500000001", "YV", "COM 90000001 21 625   18           0021", "Q6", "0590000001", 2, 80),
    L(P, "9500000001", "YV", "COM 90000001 21 625   18           0021", "Q8", "0590000001", 2, 25),
    L(Rc, "9500000001", "YV", "0021/Comisión % DILISA            0021", "D2", "0590000001", 2, 197.2),
    L(Rc, "9500000001", "YV", "0021/Sobrecomis % T.Exter         0021", "D2", "0590000001", 2, 41.18),
    # O2 with a Liverpool promotion
    L(P, "9500000002", "YV", "F/ 90000002  21 625                0021", "D2", "0590000002", 5, -2088),
    L(P, "9500000002", "YV", "COM 90000002 21 625   18           0021", "Q6", "0590000002", 5, 160),
    L(P, "9500000002", "YV", "COM 90000002 21 625   18           0021", "Q8", "0590000002", 5, 50),
    L(Rc, "9500000002", "YV", "0021/Comisión % DILISA            0021", "D2", "0590000002", 5, 394.4),
    L(P, "9500000009", "YV", "F/", "C2", "0590000002", 5, -232),
    # O3 posted and cancelled
    L(P, "9500000003", "YV", "F/ 90000003  46 625                0046", "D2", "0590000003", 10, -1160),
    L(P, "9500000003", "YV", "COM 90000003 46 625   18           0046", "Q6", "0590000003", 10, 80),
    L(P, "9500000003", "YV", "COM 90000003 46 625   18           0046", "Q8", "0590000003", 10, 25),
    L(Rc, "9500000003", "YV", "0046/Comisión % DILISA            0046", "D2", "0590000003", 10, 197.2),
    L(P, "9600000003", "YD", "DEV 90000003 46 625   14           0046", "C5", "0190000003", 12, 1160),
    L(P, "9600000003", "YD", "F/ 90000003  46 625                0046", "Q6", "0190000003", 12, -80),
    L(P, "9600000003", "YD", "F/ 90000003  46 625                0046", "Q8", "0190000003", 12, -25),
    L(P, "9600000003", "YD", "F/ 90000003  46 625                0046", "C5", "0190000003", 12, -197.2),
    # fees
    L(Rc, "1800000001", "DY", "Rec Logistica 15/09/2026", "C2", "MKPC20260915", 15, 300),
    L(Rc, "1800000002", "DY", "Rec Logística 16/09/2026", "D2", "", 16, 50),
    L(Rc, "1800000003", "DY", "Balance logistico Septiembre 2026", "C2", "", 20, 60),
    L(Rc, "1800000004", "DY", "Suscripcion MKP Septiembre 2026", "C2", "", 8, 600),
    L(P, "3100000001", "RW", "Ajuste log Agosto 2026", "C2", "", 14, -40),
    # payments: an attempt sent back by the bank (nets to 0), then a payment
    L(P, "2000000001", "ZP", "", "", "", 7, 3000, "2000000001"),
    L(P, "2000000002", "ZP", "", "", "", 7, -3000, "2000000001"),
    L(Rc, "2000000001", "ZP", "RECHAZO BANCO", "", "", 7, -900, "2000000001"),
    L(Rc, "2000000002", "ZP", "PAGO RECHAZADO", "", "", 7, 900, "2000000001"),
    L(P, "2000000010", "ZP", "", "", "", 28, 2900, "2000000010"),
    L(Rc, "2000000010", "ZP", "", "", "", 28, -900, "2000000010"),
]
SALES_COLS = ["Almacén", "Fecha", "Proveedor", "Artículo", "Nº documento", "Referencia", "Venta con impuestos",
              "Descuento", "Venta con descuento", "Comisión Dilisa", "Pago neto"]


def S_(doc, day, article, order, gross, disc, comm):
    return {"Almacén": "0021", "Fecha": D(day), "Proveedor": "0000163073", "Artículo": article,
            "Nº documento": doc, "Referencia": "0000000000000000000" + order, "Venta con impuestos": gross,
            "Descuento": disc, "Venta con descuento": gross + disc, "Comisión Dilisa": comm,
            "Pago neto": gross + disc - comm}


SALES = [S_("9500000001", 2, "1000000111", "1110000001", 1160, 0, 197.2),
         S_("9500000002", 5, "1000000222", "2220000002", 2320, -232, 394.4),
         S_("9500000003", 10, "1000000111", "3330000003", 1160, 0, 197.2),
         S_("9600000003", 12, "1000000111", "3330000003", -1160, 0, -197.2)]


def O_(order, day, article, sku, subtotal, status, promo=0, delivered=None):
    return {"Fecha de creación": "%02d/09/2026 - 10:00 AM" % day, "Id del pedido": order, "Estado": status,
            "Canal de venta": "Website", "Método de envío": "Home delivery", "No de SKU Liverpool": article,
            "Nombre del producto": "Producto %s" % sku, "SKU de la oferta": sku, "Cantidad": 1,
            "Precio por unidad": subtotal, "Subtotal": subtotal, "Promoción Liverpool": promo,
            "Fecha de entrega": delivered, "Nombre completo del cliente que recibe": "Fulano de Tal",
            "Calle 1": "Insurgentes 1", "Teléfono principal": "5555555555"}


ORDERS = [O_("1110000001", 2, "1000000111", "SKU-A", 1160, "Recibido", delivered="08/09/2026 - 01:00 PM"),
          O_("2220000002", 5, "1000000222", "SKU-B", 2320, "Recibido", 232, "02/10/2026 - 01:00 PM"),
          O_("3330000003", 10, "1000000111", "SKU-A", 1160, "Cancelado"),
          O_("4440000004", 30, "1000000222", "SKU-B", 500, "Cancelado")]


def make_folder(tmp_path, drop=()):
    folder = os.path.join(str(tmp_path), STORE, "20261005_100000_m202609")
    os.makedirs(folder)
    led = pd.DataFrame(LEDGER, columns=LEDGER_COLS)
    led = pd.concat([led, pd.DataFrame([{"Importe Debe/Haber": led["Importe Debe/Haber"].sum()}])],
                    ignore_index=True)            # the export's total row
    files = {
        "ledger": (R.LEDGER % "202609", led),
        "sales": (R.SALES % "202609", pd.concat([pd.DataFrame(SALES, columns=SALES_COLS),
                                                 pd.DataFrame([{"Venta con impuestos": 3480}])],
                                                ignore_index=True)),
        "orders": ("LV_pedidos_20261005.xlsx", pd.DataFrame(ORDERS)),
        "logistics": (R.LOGISTICS % "202609", pd.DataFrame([
            {"Nombre paquetera": "FEDEX_MKP", "Pedido": "1110000001", "Tipo": "Envío", "Fecha Cobro": "15/09/2026 - 07:00 PM",
             "Costos Logisticos": 100.0, "Costo de seguro": 0.0, "Estado de cobro": "COBRADO", "Doc contable": "1800000001"},
            {"Nombre paquetera": "FEDEX_MKP", "Pedido": "2220000002", "Tipo": "Envío", "Fecha Cobro": "15/09/2026 - 07:00 PM",
             "Costos Logisticos": 200.0, "Costo de seguro": 0.0, "Estado de cobro": "COBRADO", "Doc contable": "1800000001"}])),
        "balance": (R.BALANCE % "202609", pd.DataFrame([
            {"Id ajuste": "330001", "Fecha de registro": "20/09/2026 - 06:00 AM", "Estatus": "COBRADO",
             "Remisión": "2220000002", "Monto": 60.0, "Motivo": "COMPLEMENTARY", "Doc contable": "1800000003"}])),
        "payments": (R.PAYMENTS % "202609", pd.DataFrame([
            {"Fe.contab.": D(28), "Fecha pago": D(28), "Doc.comp.": "2000000010", "Importe": 2900.0,
             "Tipo de pago": "Transferencia"}])),
    }
    for key, (name, df) in files.items():
        if key not in drop:
            df.to_excel(os.path.join(folder, name), index=False)
    inv = os.path.join(folder, R.INVOICE_DIR % "202609")
    os.makedirs(inv)
    with open(os.path.join(inv, "DLI931201MI9_2026_MCOM_000000000001.xml"), "w", encoding="utf-8") as fh:
        fh.write('<?xml version="1.0" encoding="utf-8"?><cfdi:Comprobante xmlns:cfdi="http://www.sat.gob.mx/cfd/4" '
                 'Serie="MCOM" Folio="000000000001" Fecha="2026-09-08T10:00:00" SubTotal="510.00" Total="591.60" '
                 'TipoDeComprobante="I"><cfdi:Conceptos><cfdi:Concepto Descripcion="Importe de Comisiones"/>'
                 '</cfdi:Conceptos><cfdi:Impuestos TotalImpuestosTrasladados="81.60"/></cfdi:Comprobante>')
    return folder


def build(folder, tmp_path, recalc=False):
    out = os.path.join(str(tmp_path), "out", "%s_%s.xlsx" % (STORE, MONTH))
    r = RB.build([(STORE, folder)], MONTH, out)
    if recalc:
        K.recalc(out)
    return r, out


# ------------------------------------------------------------------ catalog and readers

def test_catalog_registers_every_line_of_the_story():
    for row in LEDGER:
        assert C.lookup(row["Lib.mayor"], row["Clase doc."], row["Texto"], row["II"]), row
    assert C.lookup(P, "YV", "F/ 1 21 625", "D2")[0] == C.SALES
    assert C.lookup(P, "YV", "F/", "C2")[0] == C.PROMO
    assert C.lookup(Rc, "YV", "0458/Comisión % DILISA      0458", "D2")[0] == C.COMMISSION
    assert C.lookup(P, "YD", "F/ 1 46 625", "C5")[0] == C.COMMISSION
    assert C.lookup(Rc, "DY", "Cargo misterioso", "C2") is None


def test_order_list_drops_customer_data(tmp_path):
    folder = make_folder(tmp_path)
    o = R.read_orders(os.path.join(folder, "LV_pedidos_20261005.xlsx"))
    assert not {"Nombre completo del cliente que recibe", "Calle 1", "Teléfono principal"} & set(o.columns)
    assert "Fulano" not in o.to_string()
    assert list(o["sku"][:2]) == ["SKU-A", "SKU-B"]


# ------------------------------------------------------------------ the numbers

def test_pl_from_the_ledger(tmp_path):
    r, _ = build(make_folder(tmp_path), tmp_path)
    m = r["metrics"][STORE]
    pl = m["pl"]
    assert pl[C.SALES] == pytest.approx(1160 + 2088 + 1160)
    assert pl[C.PROMO] == pytest.approx(232)
    assert m["revenue"] == pytest.approx(4640)
    assert pl[C.REFUND] == pytest.approx(-1160)
    assert m["net_sales"] == pytest.approx(3480)
    assert pl[C.COMMISSION] == pytest.approx(-(197.2 + 394.4 + 197.2 - 197.2))   # returned on O3
    assert pl[C.CARD_FEE] == pytest.approx(-41.18)
    assert pl[C.LOGISTICS] == pytest.approx(-350)
    assert pl[C.LOGISTICS_EXTRA] == pytest.approx(-60)
    assert pl[C.LOGISTICS_ADJ] == pytest.approx(40)
    assert pl[C.SUBSCRIPTION] == pytest.approx(-600)
    assert m["fees"] == pytest.approx(-591.6 - 41.18 - 350 - 60 + 40 - 600)
    assert pl[C.WH_IVA] == pytest.approx(-240) and pl[C.WH_ISR] == pytest.approx(-75)
    assert pl["withheld"] == pytest.approx(-315)
    assert m["contrib"] == pytest.approx(3480 + m["fees"])
    # payments: the attempt sent back nets to 0; the real one is 2,900 − 900 of fees
    assert m["cash"] == pytest.approx(2000)
    assert m["rejected"] == [("2000000001", 3000.0)]
    assert not RB.RF.failed_errors(r["checks"])


def test_bridges_tie(tmp_path):
    r, _ = build(make_folder(tmp_path), tmp_path)
    m = r["metrics"][STORE]
    A = m["bridge_a"]
    assert A["orders"] == pytest.approx(1160 + 2320 + 1160 + 500)
    assert A["not_posted"] == pytest.approx(500) and A["not_posted_states"] == {"Cancelado": 1}
    assert A["resid"] == pytest.approx(0)
    assert m["bridge_b"]["vm"] == pytest.approx(m["bridge_b"]["ledger"])
    D_ = m["bridge_d"]
    assert D_["matched_list"] == pytest.approx(D_["matched_ledger"]) == pytest.approx(-300)
    assert D_["ledger_only"] == pytest.approx(-50)
    checks = dict((c["key"], c) for c in r["checks"])
    for k in ("pl_tie", "bridge_a", "bridge_b", "bridge_d", "unregistered", "sale_link"):
        assert checks[k]["ok"], (k, checks[k]["detail"])
    assert not checks["payments_rejected"]["ok"]                 # a warning, with the attempt
    # posted in September, delivered in October (O2): the timing memo
    assert m["undelivered"] == pytest.approx(2320) and m["undelivered_n"] == 1


def test_sku_table_adds_up_to_the_pl(tmp_path):
    r, _ = build(make_folder(tmp_path), tmp_path)
    m = r["metrics"][STORE]
    g = RB.derive(RB.load_store(STORE, make_folder(tmp_path / "again"), MONTH))["sku"]
    assert g["contrib"].sum() == pytest.approx(m["contrib"])
    assert g["withheld"].sum() == pytest.approx(m["pl"]["withheld"])
    a, b = g.loc["SKU-A"], g.loc["SKU-B"]
    assert a["units"] == 2 and a["returns"] == 1                # O1 and O3, O3 cancelled
    assert b[C.SALES] + b[C.PROMO] == pytest.approx(2320)
    # shipping by the shipping list (O2 200) and extra shipping (O2 60) on SKU-B, plus allocated shares
    assert b[C.LOGISTICS] < -200 and b[C.LOGISTICS_EXTRA] == pytest.approx(-60)


def test_unregistered_line_is_an_error_and_still_ties(tmp_path):
    global LEDGER
    saved = list(LEDGER)
    LEDGER = LEDGER + [L(Rc, "1800000099", "DY", "Cargo misterioso", "C2", "", 21, 77)]
    try:
        r, _ = build(make_folder(tmp_path), tmp_path)
    finally:
        LEDGER = saved
    m = r["metrics"][STORE]
    assert m["pl"][C.UNREGISTERED] == pytest.approx(-77)
    bad = dict((c["key"], c) for c in RB.RF.failed_errors(r["checks"]))
    assert "unregistered" in bad and "pl_tie" not in bad


def test_marketplace_files_missing_keeps_the_money(tmp_path):
    full, _ = build(make_folder(tmp_path / "a"), tmp_path / "a")
    part, _ = build(make_folder(tmp_path / "b", drop=("orders", "logistics", "balance")), tmp_path / "b")
    assert part["metrics"][STORE]["pl"] == full["metrics"][STORE]["pl"]
    assert not RB.RF.failed_errors(part["checks"])
    warn = dict((c["key"], c) for c in part["checks"])
    assert not warn["file_orders"]["ok"] and warn["file_orders"]["severity"] == "warn"
    from core import grading
    for form in ("orders", "logistics", "balance"):
        assert grading.non_blocking_form("LV", "main", form)
    assert grading.non_blocking_form("LV", "main", "ledger") is None


def test_no_ledger_is_an_error(tmp_path):
    r, _ = build(make_folder(tmp_path, drop=("ledger",)), tmp_path)
    assert "file_ledger" in [c["key"] for c in RB.RF.failed_errors(r["checks"])]


# ------------------------------------------------------------------ the workbook

def test_workbook_is_a_sister(tmp_path):
    r, out = build(make_folder(tmp_path), tmp_path)
    wb = load_workbook(out)
    assert wb.sheetnames == K.SHEET_NAMES
    labels = [c.value for c in wb[K.SHEET_NAMES[1]]["A"] if c.value]
    for want in ("商品销售收入（买家实付）", "加：Liverpool 出资的促销补偿", "减：取消与退货退款", "净销售收入",
                 "平台费用小计", "经营贡献毛利（未扣商品成本）", "减：商品采购成本",
                 "平台代扣代缴税金（预缴税款，不计入贡献）"):
        assert want in labels
    text = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value)
    assert "Fulano" not in text and "Insurgentes" not in text


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_workbook_recalculates_with_zero_residuals(tmp_path):
    r, out = build(make_folder(tmp_path), tmp_path, recalc=True)
    n_f, n_err, n_none, bad = K.scan_formula_errors(out)
    assert n_err == 0, bad
    wb = load_workbook(out, data_only=True)
    ws = wb[K.SHEET_NAMES[2]]
    for row in ws.iter_rows(values_only=True):
        if row and isinstance(row[0], str) and row[0].startswith("残差"):
            assert all(abs(v) < 0.01 for v in row[1:] if isinstance(v, (int, float))), row
    pl = dict((row[0], row[2]) for row in wb[K.SHEET_NAMES[1]].iter_rows(values_only=True) if row and row[0])
    assert pl["经营贡献毛利（未扣商品成本）"] == pytest.approx(r["metrics"][STORE]["contrib"])
