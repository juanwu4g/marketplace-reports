# -*- coding: utf-8 -*-
"""Period reports (core/period_report, run_period.py): built only from the final
monthly results; totals to the cent; missing store-months; changed versions;
names and folders in both layouts; the command line's exit codes."""
import decimal
import io
import os

import pytest
from openpyxl import load_workbook

import period_data as P
from core import results as R
from core import report_kit as K
from core.period_report import choice as PC
from core.period_report import cli
from core.period_report import output as G
from core.period_report import model as MD
from core.period_report import workbook as WB

Q3 = MD.quarter("2026-Q3")


@pytest.fixture
def v2(tmp_path, monkeypatch):
    return P.configure(tmp_path, monkeypatch, v2=True)


@pytest.fixture
def legacy(tmp_path, monkeypatch):
    return P.configure(tmp_path, monkeypatch, v2=False)


def sheet_rows(path, name):
    wb = load_workbook(path)
    return [list(r) for r in wb[name].iter_rows(values_only=True)]


def cents(values, places=2):
    """What 'to the cent' means: the decimal figures of the results files added
    exactly, rounded half-up."""
    total = sum((decimal.Decimal(repr(float(v))) for v in values), decimal.Decimal(0))
    q = decimal.Decimal(1).scaleb(-places)
    return float(total.quantize(q, rounding=decimal.ROUND_HALF_UP))


def block(rows, title):
    """The rows of one ② block: {label: [month values…, total]}."""
    out, inside = {}, False
    for r in rows:
        if r[0] == title or (isinstance(r[0], str) and r[0].startswith(title + "（")):
            inside = True
            continue
        if inside and r[0] == "项目":
            continue
        if inside and (r[0] is None or r[0] == "备查项（不进损益）"):
            break
        if inside:
            out[r[0]] = r[2:]
    return out


def noisy_quarter():
    """Three stores, values that do not add up exactly in binary floating point."""
    vals = {("S1", "2026-07"): 0.1, ("S1", "2026-08"): 0.2, ("S1", "2026-09"): 0.7,
            ("S2", "2026-07"): 1234.565, ("S2", "2026-08"): 0.105, ("S2", "2026-09"): 9999.995,
            ("S3", "2026-07"): -0.015, ("S3", "2026-08"): 33.33, ("S3", "2026-09"): 66.67}
    for (s, m), v in vals.items():
        P.put(s, m, "2026%s01_120000" % ("%02d" % (int(m[5:]) + 1)),
              pl=[("gmv", v), ("fee", -v / 3.0), ("net", v - v / 3.0, "f")])
    return vals


# ------------------------------------------------------------------ totals

def test_period_totals_equal_the_sum_of_the_final_monthly_results(v2):
    vals = noisy_quarter()
    res = G.generate(Q3, "MX")
    assert res["status"] == G.COMPLETE and res["path"].endswith("MX_ML_2026-Q3_账单基准.xlsx")
    rows = sheet_rows(res["path"], K.SHEET_NAMES[1])
    for key, f in (("GMV", lambda v: v), ("FEE", lambda v: -v / 3.0), ("NET", lambda v: v - v / 3.0)):
        for s in ("S1", "S2", "S3"):
            want = cents(f(vals[(s, m)]) for m in Q3.months)
            assert block(rows, s)[key][3] == want
            assert [round(x, 2) for x in block(rows, s)[key][:3]] == [round(f(vals[(s, m)]), 2) for m in Q3.months]
        market = block(rows, WB.MARKET_TOTAL)[key]
        assert market[3] == cents(f(v) for v in vals.values())
        for i, m in enumerate(Q3.months):
            assert market[i] == cents(f(vals[(s, m)]) for s in ("S1", "S2", "S3"))
    # the half-cent case: 1234.565 + 0.105 + 9999.995 = 11234.665 -> 11234.67 (half-up)
    assert block(rows, "S2")["GMV"][3] == 11234.67
    # ③: every total is within half a cent of an independent binary sum
    checks = sheet_rows(res["path"], K.SHEET_NAMES[2])
    assert [r[2] for r in checks[5:] if r[2]] == ["通过"] * 5
    # ⑧: the same market totals by month
    trend = sheet_rows(res["path"], K.SHEET_NAMES[7])
    assert trend[4][3] == cents(vals[(s, "2026-07")] for s in ("S1", "S2", "S3"))
    assert trend[-1][0] == "期间合计" and trend[-1][3] == cents(vals.values())


def test_clp_has_no_decimals_and_one_store_has_no_market_block(v2):
    for m, v in (("2026-07", 1000.4), ("2026-08", 2000.4), ("2026-09", 3000.4)):
        P.put("F1", m, "20261001_000000", pl=[("gmv", v)])
    res = G.generate(Q3, "FA-CL")
    assert res["path"].endswith(os.path.join("FA-CL", "CL_FA_2026-Q3.xlsx"))
    rows = sheet_rows(res["path"], K.SHEET_NAMES[1])
    assert block(rows, "F1")["GMV"][3] == 6001
    assert not any(r[0] and str(r[0]).startswith(WB.MARKET_TOTAL) for r in rows)
    wb = load_workbook(res["path"])
    assert wb[K.SHEET_NAMES[1]].cell(row=7, column=6).number_format == K.MNY0


def test_non_final_versions_are_ignored(v2):
    P.put("S1", "2026-07", "20260805_000000", pl=[("gmv", 10.0)])
    P.put("S1", "2026-07", "20260812_000000", level=R.GAPS, pl=[("gmv", 999.0)], reasons=["缺 Ventas"])
    P.put("S1", "2026-08", "20260905_000000", pl=[("gmv", 20.0)], level=R.COMPLETE_WARN)
    P.put("S1", "2026-08", "20260915_000000", pl=[("gmv", 25.0)])            # newer complete wins
    P.put("S1", "2026-09", "20261005_000000", level=R.PROVISIONAL, pl=[("gmv", 777.0)],
          reasons=["9 月账单未关账"])
    P.put("S1", "2026-09", "20261006_000000", level=R.FAILED, pl=[("gmv", 666.0)])
    P.put("S1", "2026-09", "20261005_000000", basis="order", pl=[("gmv", 555.0)])   # other basis
    D = MD.gather("MX", Q3)
    assert D.final("S1", "2026-07")["pl"][0]["value"] == 10.0
    assert D.final("S1", "2026-08")["version"].endswith("20260915_000000")
    assert D.final("S1", "2026-09") is None
    assert "暂定" not in D.cell("S1", "2026-09").reason()                # the newest is the failed one
    assert "失败" in D.cell("S1", "2026-09").reason()
    assert D.store_total("S1", "gmv") == 35.0
    order = MD.gather("MX", Q3, basis="order")
    assert order.store_total("S1", "gmv") == 555.0 and order.basis == "order"


# ------------------------------------------------------------------ missing

def test_missing_store_month_is_left_out_banner_and_listed(v2):
    P.put("S1", "2026-07", "20260805_000000", pl=[("gmv", 10.0)])
    P.put("S1", "2026-08", "20260905_000000", pl=[("gmv", 20.0)])
    P.put("S1", "2026-09", "20261005_000000", level=R.GAPS, pl=[("gmv", 30.0)],
          reasons=["Ventas 没下载到"])
    P.put("S2", "2026-08", "20260905_000000", pl=[("gmv", 5.0)])
    res = G.generate(Q3, "MX")
    assert res["status"] == G.INCOMPLETE
    assert [(c.store, c.month) for c in res["missing"]] == [
        ("S1", "2026-09"), ("S2", "2026-07"), ("S2", "2026-09")]
    wb = load_workbook(res["path"])
    readme = wb[K.SHEET_NAMES[0]]
    assert readme["B3"].value.startswith("⚠ 有 3 个店铺月份没有完整版本")
    assert readme["B3"].fill.fgColor.rgb.endswith(K.C_RED)
    text = [r for r in sheet_rows(res["path"], K.SHEET_NAMES[0]) if r[1] == "S1 · 2026-09"][0][2]
    assert text.startswith("无完整版本") and "有缺口" in text and "Ventas 没下载到" in text
    pl = sheet_rows(res["path"], K.SHEET_NAMES[1])
    s1 = block(pl, "S1")["GMV"]
    assert s1 == [10.0, 20.0, None, 30.0, None]
    assert block(pl, WB.MARKET_TOTAL)["GMV"][:4] == [10.0, 25.0, None, 35.0]
    heads = [r for r in pl if r[0] == "项目"]
    assert heads[0][2:5] == ["2026-07\n1/2 家完整", "2026-08", "2026-09\n0/2 家完整"]
    assert heads[1][4] == "2026-09\n无完整版本"
    gaps = sheet_rows(res["path"], K.SHEET_NAMES[9])
    listed = [(r[0], r[1]) for r in gaps if r[2] == WB.NO_FINAL]
    assert listed == [("S1", "2026-09"), ("S2", "2026-07"), ("S2", "2026-09")]


def test_listed_store_without_any_result_is_missing_everywhere(v2):
    P.put("S1", "2026-07", "20260805_000000")
    D = MD.gather("MX", Q3, stores=["S1", "S3"])
    assert D.stores == ["S1", "S3"]
    assert "没有任何版本" in D.cell("S3", "2026-08").reason()
    assert len(D.missing()) == 5


def test_pl_rows_follow_the_latest_month_and_missing_rows_count_zero(v2):
    P.put("S1", "2026-07", "20260805_000000", pl=[("gmv", 10.0), ("old", 1.0)])
    P.put("S1", "2026-08", "20260905_000000", pl=[("gmv", 20.0), ("ads", -2.0)])
    P.put("S1", "2026-09", "20261005_000000", pl=[("ads", -3.0), ("gmv", 30.0)])
    D = MD.gather("MX", Q3)
    assert [r["key"] for r in D.pl_spec()] == ["ads", "gmv", "old"]
    assert D.store_total("S1", "ads") == -5.0 and D.pl_value("S1", "2026-07", "ads") == 0
    assert sorted((m, k) for _, m, k, _ in D.pl_gaps()) == [
        ("2026-07", "ads"), ("2026-08", "old"), ("2026-09", "old")]
    res = G.generate(Q3, "MX")
    pl = sheet_rows(res["path"], K.SHEET_NAMES[1])
    note = [r for r in pl if r[0] == "ADS"][0][-1]
    assert note == "2026-07 的结果里没有这一行，按 0 计"


def test_other_sections_add_up(v2):
    def put(m, built, units, cost, stock):
        P.put("S1", m, built, pl=[("gmv", 1.0)],
              fees=[{"name": "Cargo por venta", "label": "佣金", "pl_key": "fee", "amount": -1.25}],
              returns=[{"key": "cancel", "label": "取消", "count": 2, "amount": -3.5}],
              memo=[{"key": "iva", "label": "IVA", "value": 0.16}],
              sku=[{"sku": "SKU1", "title": "Bocina " + m, "units": units, "units_returned": 1,
                    "net_sales": 10.1, "fees": -1.1, "contribution": 9.0, "cost": cost}],
              inventory={"as_of": m + "-30", "rows": [{"sku": "SKU1", "title": "B", "stock": stock,
                                                       "sold_30d": 4}]},
              checks=[K.check("bridge_b", "S1", "桥 B", False, "差 3.20", "warn")])
    put("2026-07", "20260805_000000", 3, -5.0, 10)
    put("2026-08", "20260905_000000", 4, None, 20)
    put("2026-09", "20261005_000000", 5, -7.0, 30)
    res = G.generate(Q3, "MX")
    D = res["data"]
    sku = D.sku_rows("S1")[0]
    assert (sku["units"], sku["net_sales"], sku["cost"], sku["title"]) == (12, 30.3, -12.0, "Bocina 2026-09")
    assert sku["cost_missing"] == {"2026-08"}
    rows = sheet_rows(res["path"], K.SHEET_NAMES[6])
    assert rows[4][:9] == ["S1", "SKU1", "Bocina 2026-09", 12, 3, 30.3, -3.3, 27.0, -12.0]
    assert rows[4][9] == "2026-08 没有成本，未计入"
    fees = sheet_rows(res["path"], K.SHEET_NAMES[3])
    assert fees[4] == ["S1", "佣金", "Cargo por venta", "fee", -1.25, -1.25, -1.25, -3.75]
    ret = sheet_rows(res["path"], K.SHEET_NAMES[4])
    assert ret[4][-1] == 6 and ret[5][-1] == -10.5
    inv = sheet_rows(res["path"], K.SHEET_NAMES[8])
    assert inv[4][:7] == ["S1", "2026-09-30", "2026-09", "SKU1", "B", 30, 4]     # not added up
    gaps = sheet_rows(res["path"], K.SHEET_NAMES[9])
    assert sum(1 for r in gaps if r[2] == "桥 B") == 3
    assert res["carried"] == 3
    memo = [r for r in sheet_rows(res["path"], K.SHEET_NAMES[1]) if r[0] == "IVA"][0]
    assert memo[5] == 0.48


# ------------------------------------------------------------------ changed versions

def test_regeneration_lists_versions_changed_since_the_previous_file(v2):
    P.put("S1", "2026-07", "20260805_000000")
    P.put("S1", "2026-08", "20260905_000000")
    first = G.generate(Q3, "MX")
    assert first["first"] and first["changed"] == []
    gaps = sheet_rows(first["path"], K.SHEET_NAMES[9])
    assert any(r[2] == "这是第一次生成，没有上一份可比。" for r in gaps)
    again = G.generate(Q3, "MX")
    assert not again["first"] and again["changed"] == []
    P.put("S1", "2026-08", "20260920_000000")                                # a newer complete version
    P.put("S1", "2026-09", "20261005_000000")                                # a month gets one
    third = G.generate(Q3, "MX")
    assert [(s, m, what) for s, m, _, _, what in third["changed"]] == [
        ("S1", "2026-08", "换成了更新的完整版本"), ("S1", "2026-09", "有了完整版本")]
    old, new = third["changed"][0][2:4]
    assert old.endswith("20260905_000000") and new.endswith("20260920_000000")
    gaps = sheet_rows(third["path"], K.SHEET_NAMES[9])
    assert [r[2] for r in gaps if r[0] == "S1" and r[2] and "版本" in r[2]] == [
        "换成了更新的完整版本", "有了完整版本"]
    # the record the next generation reads is hidden
    wb = load_workbook(third["path"])
    assert wb[WB.VERSIONS_SHEET].sheet_state == "hidden"
    assert wb.sheetnames[:10] == K.SHEET_NAMES
    assert WB.versions_of(third["path"])[("S1", "2026-08")] == new


def test_unreadable_previous_file_is_said_not_guessed(v2):
    P.put("S1", "2026-07", "20260805_000000")
    path = G.target(Q3, "ML-MX", "bill")
    os.makedirs(os.path.dirname(path))
    io.open(path, "wb").write(b"not a workbook")
    res = G.generate(Q3, "MX")
    assert res["status"] == G.INCOMPLETE and not res["first"]
    rows = sheet_rows(res["path"], K.SHEET_NAMES[9])
    assert any(r[2] == "上一份同名文件读不出版本记录，没法比较。" for r in rows)


def test_file_open_in_excel_is_an_error_and_leaves_no_partial_file(v2, monkeypatch):
    P.put("S1", "2026-07", "20260805_000000")

    def locked(a, b):
        raise PermissionError("[WinError 32] being used by another process")
    monkeypatch.setattr(G.os, "replace", locked)
    res = G.generate(Q3, "MX")
    assert res["status"] == G.ERROR and res["path"] is None and "Excel" in res["message"]
    folder = os.path.dirname(G.target(Q3, "ML-MX", "bill"))
    assert os.listdir(folder) == []


# ------------------------------------------------------------------ names and folders

def test_names_and_folders_in_the_new_layout(v2):
    root = v2
    year, cust = MD.year("2026"), MD.custom("2026-07", "2026-11")
    assert G.target(Q3, "MX", "bill") == os.path.join(root, "quarterly", "2026-Q3", "ML-MX",
                                                      "MX_ML_2026-Q3_账单基准.xlsx")
    assert G.target(year, "MX", "order") == os.path.join(root, "annual", "2026", "ML-MX",
                                                         "MX_ML_2026_订单基准.xlsx")
    assert G.target(cust, "FA-CL", "main") == os.path.join(root, "custom", "2026-07_2026-11", "FA-CL",
                                                           "CL_FA_2026-07_2026-11.xlsx")
    assert G.target(Q3, "AZ-MX", "main", out_dir="/share/r") == os.path.join(
        os.path.abspath("/share/r"), "quarterly", "2026-Q3", "AZ-MX", "MX_AZ_2026-Q3.xlsx")
    assert G.alert_name("MX", Q3, "bill") == "！期间报表异常_MX_ML_2026-Q3_账单基准.txt"


def test_names_and_folders_in_the_old_layout(legacy, tmp_path):
    flat = os.path.join(str(tmp_path / "code"), "downloads", "data", "reports", "financial")
    assert legacy == flat
    assert G.target(Q3, "CL", "bill") == os.path.join(flat, "期间报表", "quarterly", "2026-Q3", "ML-CL",
                                                      "CL_ML_2026-Q3_账单基准.xlsx")
    P.put("S1", "2026-07", "20260805_000000")
    res = G.generate(Q3, "MX")
    assert res["path"] == os.path.join(flat, "期间报表", "quarterly", "2026-Q3", "ML-MX",
                                       "MX_ML_2026-Q3_账单基准.xlsx")
    assert os.path.exists(res["path"])


def test_periods():
    assert MD.quarter("2026q4") == MD.Period("quarterly", "2026-Q4", ["2026-10", "2026-11", "2026-12"])
    assert MD.year("2026").months[0] == "2026-01" and MD.year("2026").label == "2026"
    assert MD.custom("2026-11", "2027-02").label == "2026-11_2027-02"
    for bad in (lambda: MD.year("26"), lambda: MD.custom("2026-13", "2027-01"),
                lambda: MD.custom("2026-09", "2026-08"), lambda: MD.quarter("2026-Q5")):
        with pytest.raises(ValueError):
            bad()
    assert MD.period_cn(Q3) == "2026 年第 3 季度（2026-07 至 2026-09）"
    assert G.auto_periods("2026-09") == [Q3]
    assert [p.label for p in G.auto_periods("2026-12")] == ["2026-Q4", "2026"]
    assert G.auto_periods("2026-10") == [] and G.auto_periods("2026-11") == []
    assert G.markets_of("ALL")[:3] == ["ML-MX", "ML-CL", "ML-PE"]
    assert G.markets_of("MX,FA-CL,MX") == ["ML-MX", "FA-CL"]


def test_store_no_longer_in_config_is_still_read_by_its_results_folder(tmp_path, monkeypatch):
    P.configure(tmp_path, monkeypatch, v2=True)
    P.put("S1", "2026-07", "20260805_000000", pl=[("gmv", 7.0)])
    P.configure(tmp_path, monkeypatch, v2=True, catalog={"S2": P.CATALOG["S2"]})   # S1 removed
    D = MD.gather("MX", Q3)
    assert D.stores == ["S1"] and D.store_total("S1", "gmv") == 7.0


# ------------------------------------------------------------------ command line

def run_cli(argv, capsys):
    code = cli.main(argv)
    return code, capsys.readouterr().out


def test_cli_exit_codes(v2, capsys):
    code, out = run_cli(["--quarter", "2026-Q3"], capsys)
    assert code == 2 and "一份都没生成" in out
    for m in Q3.months:
        P.put("S1", m, "20261005_000000")
    code, out = run_cli(["--quarter", "2026-Q3", "--region", "MX"], capsys)
    assert code == 0 and "已写出" in out and "S1" in out and "✓ 完整" in out
    P.put("F1", "2026-07", "20260805_000000")
    code, out = run_cli(["--quarter", "2026-Q3"], capsys)                  # ALL: FA-CL is incomplete
    assert code == 1 and "✗ F1 2026-08 无完整版本" in out
    assert os.path.exists(G.target(Q3, "FA-CL", "main"))
    code, out = run_cli(["--quarter", "2026-Q3", "--stores", "S1", "NOPE"], capsys)
    assert code == 2 and "NOPE" in out


def test_cli_preview_writes_nothing(v2, capsys):
    P.put("S1", "2026-07", "20260805_000000")
    code, out = run_cli(["--from", "2026-07", "--to", "2026-08", "--region", "MX", "--preview"], capsys)
    assert code == 1 and "2026-08 ✗ 无完整版本" in out
    assert not os.path.exists(os.path.join(v2, "custom"))


def test_cli_error_is_exit_2(v2, capsys, monkeypatch):
    P.put("S1", "2026-07", "20260805_000000")
    monkeypatch.setattr(WB, "build", lambda *a, **k: 1 / 0)
    code, out = run_cli(["--year", "2026", "--region", "MX", "--basis", "order"], capsys)
    assert code == 2                                                       # order basis: nothing there
    code, out = run_cli(["--year", "2026", "--region", "MX"], capsys)
    assert code == 2 and "ZeroDivisionError" in out


def test_cli_rejects_bad_arguments(v2, capsys):
    for argv in (["--quarter", "2026-Q9"], ["--from", "2026-07"], ["--quarter", "2026-Q3", "--region", "XX"],
                 []):
        with pytest.raises(SystemExit) as e:
            cli.main(argv)
        assert e.value.code == 2
    capsys.readouterr()


def test_cli_out_dir_and_the_window_setting(v2, tmp_path, capsys):
    P.put("S1", "2026-07", "20260805_000000")
    code, _ = run_cli(["--quarter", "2026-Q3", "--region", "MX", "--out-dir", str(tmp_path / "share")], capsys)
    assert code == 1
    assert os.path.exists(os.path.join(str(tmp_path / "share"), "quarterly", "2026-Q3", "ML-MX",
                                       "MX_ML_2026-Q3_账单基准.xlsx"))
    (tmp_path / "gui_settings.json").write_text('{"out_dir": "%s"}' % str(tmp_path / "chosen").replace("\\", "/"),
                                                encoding="utf-8")
    run_cli(["--quarter", "2026-Q3", "--region", "MX"], capsys)
    assert os.path.isdir(os.path.join(str(tmp_path / "chosen"), "quarterly"))


def test_run_period_is_ascii_and_at_the_root():
    from core.paths import ROOT
    p = os.path.join(ROOT, "run_period.py")
    assert os.path.exists(p)
    io.open(p, encoding="ascii").read()


# ------------------------------------------------------------------ the tab's rules

def test_tab_choices_command_and_grid(v2):
    assert PC.quarters("2026-09", 3) == ["2026-Q3", "2026-Q2", "2026-Q1"]
    assert PC.quarters("2026-01", 2) == ["2026-Q1", "2025-Q4"]
    assert PC.years("2026-09", 2) == ["2026", "2025"]
    assert PC.period_of("custom", first="2026-07", last="2026-08").label == "2026-07_2026-08"
    with pytest.raises(ValueError):
        PC.period_of("custom", first="2026-09", last="2026-08")
    cmd = PC.command("py", "/r", Q3, "MX,FA-CL", ["S1"], "order", "/out")
    assert cmd == ["py", "-u", os.path.join("/r", "run_period.py"), "--quarter", "2026-Q3", "--region",
                   "MX,FA-CL", "--stores", "S1", "--basis", "order", "--out-dir", "/out"]
    assert "--basis" not in PC.command("py", "/r", MD.year("2026"), "ALL", basis="bill")
    P.put("S1", "2026-07", "20260805_000000")
    P.put("S1", "2026-08", "20260905_000000", level=R.GAPS, reasons=["缺账单"])
    g = PC.grid(Q3, "ALL")
    assert [b["market"] for b in g] == ["ML-MX"] and g[0]["missing"] == 2
    store, cells = g[0]["rows"][0]
    assert store == "S1" and [c[2] for c in cells] == ["✓ 完整", "✗ 无完整版本", "✗ 无完整版本"]
    assert "缺账单" in cells[1][3]
    assert PC.folder(Q3, "MX") is None                                        # nothing generated yet
    G.generate(Q3, "MX")
    assert PC.folder(Q3, "MX") == os.path.dirname(G.target(Q3, "MX", "bill"))
    assert PC.folder(Q3, "ALL") == os.path.dirname(os.path.dirname(G.target(Q3, "MX", "bill")))
    assert PC.exit_text(1)[1] == "warn" and PC.exit_text(9)[1] == "bad"
