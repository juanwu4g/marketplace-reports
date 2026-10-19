# -*- coding: utf-8 -*-
"""Falabella report on synthetic files shaped exactly like the real exports.

A tiny September: two orders settled in September (one returned), one August
order settled in September, one September order not settled yet, a cancelled
order, store-level ads, and one fee invoice period (21 Aug - 20 Sep) while the
next one is still open.
"""
import json
import os

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from core import report_kit
from platforms.falabella import fee_catalog
from platforms.falabella import report_build as R
from platforms.falabella import run_reports

MONTH = "2026-09"
STORE = "TEST_CL-Fala"

TX_HEAD = ["Fecha creación de la orden", "Nombre del producto", "Nº de orden", "SKU vendedor",
           "SKU Falabella", "Id Artículo", "Falabella-Id", "Estado de la orden", "Fecha de Entrega",
           "Tipo de transacción", "Fecha de transacción", "Divisa", "Monto (Sin IVA)", "IVA",
           "Monto con IVA", "Tipo de comisión", "% comisión", "Talla", "Opscore", "Estado de pago",
           "Nº estado de cuenta", "Referencia de pago Fpay", "Comentarios", "Tipo de envío",
           "Proveedor logístico", "Modalidad logística", "Seller ID", "Tipo de Seller",
           "Tipo de Documento Tributario", " N° Documento Tributario ", "Categoría de transacciones"]


def tx(order, item, sku, kind, date, gross, doc=None, created="2026-09-02 10:00:00", mode="FBF",
       pay="Pagado"):
    net = round(gross / 1.19)
    row = dict.fromkeys(TX_HEAD, "")
    row.update({"Fecha creación de la orden": created, "Nombre del producto": "Producto %s" % sku,
                "Nº de orden": order, "SKU vendedor": sku, "Id Artículo": item,
                "Estado de la orden": "entregada", "Tipo de transacción": kind,
                "Fecha de transacción": date, "Divisa": "CLP", "Monto (Sin IVA)": net,
                "IVA": gross - net, "Monto con IVA": gross, "Estado de pago": pay,
                "Nº estado de cuenta": "FACL-TEST-20260915-CLP", "Modalidad logística": mode,
                "Tipo de Documento Tributario": doc, " N° Documento Tributario ": "1" if doc else None})
    return row


TX = [
    # order A: sold 2 Sep, settled 5 Sep
    tx("100", "1", "SKU1", "Pago por precio del producto", "2026-09-05", 11900),
    tx("100", "1", "SKU1", "Cobro por comisión por venta", "2026-09-05", -1785, "Facturas"),
    tx("100", "1", "SKU1", "Cobro por cofinanciamiento logístico, Fulfillment by Falabella",
       "2026-09-05", -2380, "Facturas"),
    tx("100", "1", "SKU1", "Pago de envío comprador", "2026-09-05", 2380),
    tx("100", "1", "SKU1", "Reversa de pago de envío comprador", "2026-09-05", -2380, "Facturas"),
    # order B: sold 3 Sep, settled 6 Sep, returned 25 Sep (open period)
    tx("200", "2", "SKU2", "Pago por precio del producto", "2026-09-06", 23800, created="2026-09-03 09:00:00"),
    tx("200", "2", "SKU2", "Cobro por comisión por venta", "2026-09-06", -3570, "Facturas"),
    tx("200", "2", "SKU2", "Descuento por devolución de producto", "2026-09-25", -23800),
    tx("200", "2", "SKU2", "Reembolso por comisión por venta", "2026-09-25", 3570),
    # order C: sold in August, settled 2 Sep
    tx("300", "3", "SKU1", "Pago por precio del producto", "2026-09-02", 11900, created="2026-08-30 20:00:00"),
    tx("300", "3", "SKU1", "Cobro por comisión por venta", "2026-09-02", -1785, "Facturas",
       created="2026-08-30 20:00:00"),
    # store level: ads 1 Sep, storage 30 Sep (open period)
    tx(None, None, None, "Cobro por Productos patrocinados", "2026-09-01", -5950, "Facturas"),
    tx(None, None, None, "Cobro por almacenamiento de productos, Fulfillment by Falabella",
       "2026-09-30", -1190),
    # a line from October must be ignored
    tx("100", "1", "SKU1", "Cobro por comisión por venta", "2026-10-01", -119, "Facturas"),
]

ORDERS = [  # Item id, order, created, state, price paid
    ("1", "100", "Sep 02, 2026 10:00", "Entregadas", "11,900.00"),
    ("2", "200", "Sep 03, 2026 09:00", "Devoluciones", "23,800.00"),
    ("4", "400", "Sep 29, 2026 22:00", "Enviados", "35,700.00"),        # not settled yet
    ("5", "500", "Sep 10, 2026 12:00", "Canceladas", "5,950.00"),
]


def write_tx(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte Transacciones"
    ws.append(["Fecha de actualización", "2026-10-16 01:41:32"])
    ws.append([])
    ws.append(["Este reporte podría tener cambios hasta el cierre del período de facturación"])
    ws.append([])
    ws.append(["Información de la orden"])
    ws.append(TX_HEAD)
    for r in rows:
        ws.append([r[h] for h in TX_HEAD])
    g = wb.create_sheet("Glosario de Terminos Contables")
    g.append(["Transacciones", "Descripción"])
    wb.save(path)


def write_orders(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet 1"
    head = ["Item id", "SKU del vendedor", "SKU de Falabella", "Fecha de creacion", "N° orden",
            "Nombre del cliente", "Documento requerido", "Precio pagado", "Precio unitario",
            "Costo de envio", "Cupon/Descuentos", "Impuesto", "Producto", "Estado",
            "Tipo de Fulfillment", "Motivo de cancelacion"]
    ws.append(head)
    for item, order, created, state, paid in ORDERS:
        ws.append([item, "SKU%s" % item, "9%s" % item, created, order, "Cliente Sintético",
                   "BOLETA", paid, paid, "0.00", "0.00", "0.00", "Producto", state,
                   "Fulfilled by Falabella", "Arrepentimiento" if state == "Canceladas" else None])
    wb.save(path)


def write_invoice(path, rows):
    """Factura falabella.com 21 Aug - 20 Sep: the invoiced lines up to 20 Sep."""
    wb = Workbook()
    ws = wb.active
    ws.title = "settlement invoice"
    ws.append(["Numero de documento", "Tipo de documento", "Fecha de Transaccion", "Descripcion Factura",
               "Tipo de Transaccion", "Nombre del producto", "SKU vendedor", "SKU Falabella",
               "Monto (Sin IVA)", "IVA", "Monto con IVA", "Divisa", "N estado de cuenta", "N de orden"])
    for r in rows:
        if r["Tipo de Documento Tributario"] and r["Fecha de transacción"] <= "2026-09-20":
            ws.append(["1", "Factura", r["Fecha de transacción"] + " 10:00:00.0", "x",
                       r["Tipo de transacción"], "", r["SKU vendedor"], "", r["Monto (Sin IVA)"],
                       r["IVA"], r["Monto con IVA"], "CLP", "", r["Nº de orden"]])
    ws.append(["2", "Factura", "2026-08-25 10:00:00.0", "x", "Cobro por comisión por venta", "",
               "SKU9", "", -1000, -190, -1190, "CLP", "", "900"])          # August part
    wb.save(path)


def write_liq(path):
    with pd.ExcelWriter(path) as xw:
        pd.DataFrame([["S", "2026-09-01", "2026-09-30", "41", "GLOBAL", "SKU1", "x", 1, 11900, 35700]],
                     columns=["ID del vendedor", "Fecha desde", "Fecha hasta", "Folio",
                              "Tipo de documento referenciado", "SKU del vendedor", "Artículo",
                              "Cantidad", "Total artículo", "Total liquidación"]) \
            .to_excel(xw, sheet_name="Liquidaciones", index=False)
        pd.DataFrame([["S", "2026-09-02 10:00:00", "BOLETA", "7", "100", "f", "SKU1", "x", 1, 11900, 11900],
                      ["S", "2026-09-03 10:00:00", "BOLETA", "8", "200", "f", "SKU2", "x", 1, 23800, 23800]],
                     columns=["ID del vendedor", "Fecha y hora de emisión", "Tipo de documento", "Folio",
                              "N° de orden", "Falabella-Id", "SKU del vendedor", "Artículo", "Cantidad",
                              "Monto línea", "Monto total documento"]) \
            .to_excel(xw, sheet_name="Docs. liquidados", index=False)


def write_stock(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Product details"
    ws.append(["Marca", "Categoria", "Codigo Universal", "Nombre", "SKU Vendedor", "SKU GSC",
               "Barcode Etiqueta de Producto", "Modelo", "Color", "Tamano", "Precio",
               "Ventas ultimas 4 semanas", "Stock Disponible"])
    ws.append(["B", "{}", "", "Producto SKU1", "SKU1", "91", "", "", None, None, 11900, 8, 2])
    ws.append(["B", "{}", "", "Producto SKU2", "SKU2", "92", "", "", None, None, 23800, 0, 5])
    wb.save(path)


@pytest.fixture
def folder(tmp_path):
    d = tmp_path / "downloads" / STORE / "20261016_120000_m202609"
    d.mkdir(parents=True)
    write_tx(str(d / "NewReportTransaction_2026-10-16T05_29_08.xlsx"), TX)
    write_orders(str(d / "Ordenes_2026-09-01_2026-09-14.xlsx"))
    write_invoice(str(d / "InvoiceReport_FACL-1-2026-08-21-2026-09-20_2026-09-24T08_18_05.xlsx"), TX)
    write_liq(str(d / "S_FSC_LIQ_BOLETA_NCR_REPORT_2026-09-01_2026-09-30.xlsx"))
    write_stock(str(d / "Product_details_1.xlsx"))
    (d / "run_meta.json").write_text(json.dumps({
        "platform": "FA", "month": MONTH, "downloaded_at": "2026-10-16T12:00:00",
        "missing": [{"form": "billing1/2026-09-21_2026-10-20", "cause": "NOT_AVAILABLE",
                     "explain": "账期还没结束", "action": "无需处理", "in_report": True}]}),
        encoding="utf-8")
    return str(d)


def net(g):
    return round(g / 1.19)


def build(folder, tmp_path, **kw):
    out = str(tmp_path / "out" / "x.xlsx")
    r = R.build([(STORE, folder)], MONTH, out, **kw)
    return r, out


def checks(r):
    return dict((c["key"] + ("" if c["key"] != "bridge_b" else c["name"][-12:]), c) for c in r["checks"])


def test_pl_lines(folder, tmp_path):
    r, out = build(folder, tmp_path)
    m = r["metrics"][STORE]
    assert m["pl"]["sales"] == net(11900) + net(23800) + net(11900)
    assert m["pl"]["refund"] == net(-23800)
    assert m["pl"]["ship_income"] == net(2380) and m["pl"]["ship_buyer"] == net(-2380)
    assert m["pl"]["ship_net"] == 0                         # pass-through: not revenue
    assert m["revenue"] == m["pl"]["sales"]
    assert m["fees"] == sum(m["pl"][k] for k in R.FEE_ROWS)  # excludes the pass-through
    assert m["pl"]["commission"] == 2 * net(-1785) + net(-3570) + net(3570)   # A, C; B charged and refunded
    assert m["pl"]["ads"] == net(-5950) and m["pl"]["fulfilment"] == net(-1190)
    assert m["contrib"] == m["tx_net"]                    # every line lands on one row
    assert m["lines"] == len(TX) - 1                      # the October line is excluded
    assert r["outcome"] == "ok"


def test_bridges(folder, tmp_path):
    r, _ = build(folder, tmp_path)
    m = r["metrics"][STORE]
    A = m["bridge_a"]
    assert A["orders"] == 11900 + 23800 + 35700           # cancelled order excluded
    assert A["not_settled"] == 35700 and A["earlier"] == 11900
    assert A["resid"] == 0
    (B,) = m["bridge_b"]
    assert B["inv_net"] == B["tx_net"] and B["inv_lines"] == B["tx_lines"]
    (P,) = m["bridge_b_pending"]
    assert P["period"] == "2026-09-21_2026-10-20" and P["tx_lines"] == 1   # storage on the 30th
    c = {x["key"]: x for x in r["checks"]}
    assert c["bridge_b"]["ok"] and not c["bridge_b_open"]["ok"]
    assert c["bridge_b_open"]["severity"] == "warn"
    assert m["iva_out"] > 0 and m["iva_in"] > 0
    assert abs(m["contrib"] + m["iva_out"] - m["iva_in"] - m["tx_gross"]) < 1e-6


def test_workbook_sheets_and_no_customer_data(folder, tmp_path):
    r, out = build(folder, tmp_path)
    wb = load_workbook(out)
    assert wb.sheetnames == report_kit.SHEET_NAMES
    text = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value)
    assert "Cliente Sintético" not in text                 # customer columns never reach the workbook
    assert "不适用" in " ".join(str(c.value) for row in wb[report_kit.SHEET_NAMES[5]].iter_rows()
                                for c in row if c.value)
    pl = wb[report_kit.SHEET_NAMES[1]]
    labels = [pl.cell(row=i, column=1).value for i in range(6, 6 + len(R.PL_ROWS))]
    assert labels == [x[1] for x in R.PL_ROWS]             # fixed row positions


def test_unregistered_type_is_an_error(folder, tmp_path):
    rows = TX + [tx("100", "1", "SKU1", "Cobro por algo nuevo", "2026-09-10", -1190, "Facturas")]
    write_tx(os.path.join(folder, "NewReportTransaction_2026-10-16T05_29_08.xlsx"), rows)
    r, _ = build(folder, tmp_path)
    c = {x["key"]: x for x in r["checks"]}
    assert not c["unregistered"]["ok"] and c["unregistered"]["severity"] == "error"
    assert r["metrics"][STORE]["pl"]["unregistered"] == net(-1190)
    assert c["pl_tie"]["ok"]                               # still ties: the amount is kept
    assert r["outcome"] == "checks_failed"
    r2, _ = build(folder, tmp_path, strict=True)
    assert r2["outcome"] == "strict_aborted"


def test_missing_orders_is_an_error_not_a_crash(folder, tmp_path):
    os.remove(os.path.join(folder, "Ordenes_2026-09-01_2026-09-14.xlsx"))
    r, out = build(folder, tmp_path)
    c = {x["key"]: x for x in r["checks"]}
    assert not c["file_order"]["ok"] and c["bridge_a"]["skipped"]
    assert os.path.exists(out)


def test_catalogue_is_complete_for_the_glossary_shape():
    assert fee_catalog.lookup("Pago por precio del producto")[1] == fee_catalog.SALES
    assert fee_catalog.lookup("  cobro por comision por venta ")[1] == fee_catalog.COMMISSION
    assert fee_catalog.lookup("Cobro por algo nuevo") is None
    assert len(fee_catalog.CATALOG) >= 147


def test_run_reports_folder_and_exit_codes(folder, tmp_path, monkeypatch):
    base = os.path.dirname(os.path.dirname(folder))
    assert run_reports.latest_folder(STORE, MONTH, base) == folder
    assert run_reports.latest_folder(STORE, "2026-08", base) is None
    usable, unusable = run_reports.resolve([STORE, "OTHER_CL-Fala"], MONTH, base)
    assert usable == [(STORE, folder)] and unusable[0][0] == "OTHER_CL-Fala"
    ok = {"outcome": "ok"}
    assert run_reports.exit_code([ok], usable, [], per_store=True) == 0
    assert run_reports.exit_code([ok], usable, unusable, per_store=True) == 2
    assert run_reports.exit_code([ok], usable, unusable, per_store=False) == 1
    assert run_reports.exit_code([{"outcome": "checks_failed"}], usable, []) == 1
    assert run_reports.exit_code([{"outcome": "crashed"}], usable, []) == 2


def test_names():
    assert report_kit.combined_workbook_name("cl", "fa", MONTH) == "CL_FA_全店汇总_2026-09.xlsx"
    assert report_kit.alert_file_name("CL", "FA", MONTH) == "！报表异常_CL_FA_202609.txt"


def test_files_from_a_later_run_of_the_same_month(folder, tmp_path):
    later = os.path.join(os.path.dirname(folder), "20261016_180000_m202609")
    os.makedirs(later)
    os.replace(os.path.join(folder, "Product_details_1.xlsx"), os.path.join(later, "Product_details_1.xlsx"))
    os.makedirs(os.path.join(os.path.dirname(folder), "20261016_190000_m202608"))   # other month: ignored
    r, out = build(folder, tmp_path)
    c = {x["key"]: x for x in r["checks"]}
    assert c["file_stock"]["ok"]
