# -*- coding: utf-8 -*-
"""The grade banner on sheet ① (core/report_kit.stamp_grade, docs/REPORT_FAMILY.md
"Grade banner"): the block itself on a small workbook, and the three runner paths
that stamp it — MercadoLibre order basis (report_build --results-version),
MercadoLibre bill basis (run_reports.bill_reports with a stamp) and the shared
ReportRunner of the other platforms — each proving the banner carries the same
grade as the results file. Builders never write it: report_build without
--results-version leaves ① as built (the golden rule)."""
import json
import os
import sys
import types

import pytest
from openpyxl import Workbook, load_workbook

from core import paths
from core import platform_reports as PR
from core import report_kit as K
from core import results as R
from core import store_config

import test_ml_bill_report as BT
import test_reports_build as OT

STAMP = "20261022_101500"
S1 = "RUN_S1"
SCRATCH = os.path.dirname(os.path.abspath(__file__))


@pytest.fixture
def v2(tmp_path, monkeypatch):
    """New layout; the test stores registered in a catalog."""
    stores = {"STORE_MX": ("ML", "MX"), "TEST": ("ML", "MX"), S1: ("FA", "CL")}
    cfg = tmp_path / "config.json"
    cfg.write_text(json.dumps({
        "paths": {"data_root": str(tmp_path / "data")},
        "stores": {"allowed": sorted(stores), "catalog": dict(
            (s, {"platform": p, "country": c}) for s, (p, c) in stores.items())}}), encoding="utf-8")
    monkeypatch.setattr(store_config, "CONFIG_PATH", str(cfg))
    monkeypatch.setattr(paths, "ROOT", str(tmp_path / "code"))
    return tmp_path


@pytest.fixture
def no_recalc(monkeypatch):
    monkeypatch.setattr(K, "recalc", lambda out, **kw: (False, "skipped"))


@pytest.fixture
def formula_errors(monkeypatch):
    """A recalculation that 'finds' a formula error (without LibreOffice)."""
    calls = []
    monkeypatch.setattr(K, "recalc", lambda out, **kw: calls.append(out) or (True, "已重算"))
    monkeypatch.setattr(K, "scan_formula_errors", lambda out: (3, 1, 0, ["②!C9 = #REF!"]))
    return calls


def banner(path):
    """{label: value} of the banner on sheet ①, plus its column; {} when there is none."""
    ws = load_workbook(path).worksheets[0]
    col = None
    for c in ws[1]:
        if c.value == K.GRADE_LABEL:
            col = c.column
    if col is None:
        return {}
    out = {"_col": col}
    for r in range(1, K.GRADE_ROWS_MAX + 1):
        label = ws.cell(row=r, column=col).value
        if label:
            out[label] = ws.cell(row=r, column=col + 1).value
    return out


def info(level, reasons=(), basis="order", version="20260901_000000_m202608__" + STAMP):
    return R.stamp_info({"level": level, "reasons": list(reasons)}, version, "20260901_000000_m202608", basis,
                        built_at=None)


def small_workbook(path):
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(K.SHEET_NAMES[0])
    ws["A1"] = "Mercado Libre 墨西哥店铺 月度财务报表"
    ws["A2"] = "会计期间：2026-08"
    ws["A3"] = "⚠ 平台提示导出不完整"
    ws.merge_cells("A3:C3")
    ws["B5"] = "一、覆盖范围"
    ws["C6"] = "EWTTO_SM"
    ws["C40"] = "=SUM(1,2)"
    wb.create_sheet(K.SHEET_NAMES[1])["A1"] = "② title"
    wb.save(path)
    return path


def cells(path):
    ws = load_workbook(path).worksheets[0]
    return dict((c.coordinate, c.value) for row in ws.iter_rows() for c in row if c.value is not None)


# ------------------------------------------------------------------ the block

def test_run_time_text():
    assert R.run_time_text("20261021_233002_m202609") == "2026-10-21 23:30"
    assert R.run_time_text("/data/raw/ML/MX/X/20261021_233002_m202609/") == "2026-10-21 23:30"
    assert R.run_time_text("20261021_233002_m202609__20261022_101500") == "2026-10-21 23:30"
    assert R.run_time_text("20261013_m202609") == "" and R.run_time_text("") == "" and R.run_time_text(None) == ""


def test_stamp_info():
    i = R.stamp_info({"level": R.COMPLETE_WARN, "reasons": ["提醒 a"]}, "v", "20261021_233002_m202609", "bill")
    assert i["level"] == R.COMPLETE_WARN and i["reasons"] == ["提醒 a"] and i["final"] is True
    assert i["basis_cn"] == "账单基准" and i["downloaded_at"] == "2026-10-21 23:30" and len(i["built_at"]) == 16
    assert R.stamp_info({"level": R.PROVISIONAL}, "v", "x", "main")["final"] is False
    assert R.stamp_info({"level": R.GAPS}, "v", "x", "main")["basis_cn"] == ""


def test_stamp_block_after_the_content_touches_no_existing_cell(tmp_path):
    p = small_workbook(str(tmp_path / "w.xlsx"))
    before = cells(p)
    assert load_workbook(p).worksheets[0].max_column == 3
    rng = K.stamp_grade(p, info(R.COMPLETE_WARN, ["缺下载（不影响金额）：stock"], "order"))
    assert rng == "E1:F8"                                   # D is the spacer
    after = cells(p)
    assert all(after[k] == v for k, v in before.items()) and len(after) == len(before) + 15
    b = banner(p)
    assert b["_col"] == 5
    assert b["完整度"] == "完整·有提醒" and b["原因"] == "缺下载（不影响金额）：stock"
    assert b["报表版本"] == "20260901_000000_m202608__" + STAMP and b["口径"] == "订单基准"
    assert b["数据下载于"] == "2026-09-01 00:00" and len(b["报表生成于"]) == 16
    assert b["是否最终版本"].startswith("是（")
    wb = load_workbook(p)
    ws = wb.worksheets[0]
    assert ws["E8"].value == K.GRADE_NOTE and ws["F8"].value is None
    assert ws.sheet_properties.tabColor.rgb.endswith(K.C_GREEN)
    assert wb.properties.subject == "完整度：完整·有提醒" and wb.properties.keywords == b["报表版本"]
    assert wb.properties.description == "缺下载（不影响金额）：stock"
    assert ws["F1"].fill.fgColor.rgb.endswith(K.C_OK) and ws["F1"].font.b and ws["F1"].font.sz == 12
    assert ws["E1"].font.b and ws["E1"].fill.fgColor.rgb.endswith(K.C_SUB)
    assert ws.column_dimensions["F"].width >= 60
    assert wb.worksheets[1]["A1"].value == "② title" and K.GRADE_LABEL not in [c.value for c in wb.worksheets[1][1]]


def test_stamp_is_idempotent_and_shrinks(tmp_path):
    p = small_workbook(str(tmp_path / "w.xlsx"))
    assert K.stamp_grade(p, info(R.GAPS, ["缺下载：billing", "校验未过：桥A"], "bill")) == "E1:F8"
    b = banner(p)
    assert b["完整度"] == "有缺口" and b["原因"] == "缺下载：billing；校验未过：桥A" and b["口径"] == "账单基准"
    assert b["是否最终版本"].startswith("否（")
    ws = load_workbook(p).worksheets[0]
    assert ws["F1"].fill.fgColor.rgb.endswith(K.C_BAD) and ws["F1"].font.color.rgb.endswith(K.C_RED)
    assert ws.sheet_properties.tabColor.rgb.endswith(K.C_RED)
    # again, as the final grade: same block, no second one, the 口径 row gone
    assert K.stamp_grade(p, info(R.COMPLETE, [], basis="main")) == "E1:F7"
    ws = load_workbook(p).worksheets[0]
    assert [c.value for c in ws[1]].count(K.GRADE_LABEL) == 1 and ws.max_column == 6
    b = banner(p)
    assert b["完整度"] == "完整" and b["原因"] == "—" and "口径" not in b
    assert ws["E7"].value == K.GRADE_NOTE and ws["E8"].value is None and ws["F8"].value is None
    assert ws.sheet_properties.tabColor.rgb.endswith(K.C_GREEN)
    assert load_workbook(p).properties.description == "—"


def test_stamp_colours_by_grade(tmp_path):
    p = small_workbook(str(tmp_path / "w.xlsx"))
    K.stamp_grade(p, info(R.PROVISIONAL, ["账单尚未关账：X"]))
    ws = load_workbook(p).worksheets[0]
    assert ws["F1"].value == "暂定" and ws["F1"].fill.fgColor.rgb.endswith(K.C_WARN)
    assert ws.sheet_properties.tabColor.rgb.endswith(K.C_AMBER)
    K.stamp_grade(p, info(R.FAILED, ["报表没有生成"]))
    ws = load_workbook(p).worksheets[0]
    assert ws["F1"].value == "失败" and ws["F1"].fill.fgColor.rgb.endswith(K.C_BAD)
    assert ws.sheet_properties.tabColor.rgb.endswith(K.C_RED)


def test_stamp_on_an_empty_first_sheet_and_a_missing_file(tmp_path, capsys):
    p = str(tmp_path / "e.xlsx")
    Workbook().save(p)                                      # one empty sheet (openpyxl cannot save none)
    assert K.stamp_grade(p, info(R.COMPLETE, basis="main")) == "C1:D7"
    assert banner(p)["完整度"] == "完整"
    assert K.stamp_grade_safely(str(tmp_path / "nope.xlsx"), info(R.COMPLETE)) is None
    assert "[提示] 完整度标记没写上" in capsys.readouterr().out


@pytest.mark.skipif(K.soffice_path() is None, reason="LibreOffice not installed")
def test_stamp_survives_recalculation(tmp_path):
    """The promise behind 'stamp after save, before recalc': LibreOffice keeps the
    banner and fills the formula cache."""
    p = small_workbook(str(tmp_path / "w.xlsx"))
    assert K.stamp_grade(p, info(R.COMPLETE)) == "E1:F8"
    ok, _msg = K.recalc(p)
    assert ok
    b = banner(p)
    assert b["完整度"] == "完整" and b["报表版本"] == "20260901_000000_m202608__" + STAMP
    assert load_workbook(p, data_only=True).worksheets[0]["C40"].value == 3


# ------------------------------------------------------------------ MercadoLibre, order basis

def order_folder(tmp_path, banner_=False):
    folder = OT.make_download_folder(tmp_path / "TEST" / "20260901_000000_m202608", banner=banner_)
    with open(os.path.join(folder, "run_meta.json"), "w", encoding="utf-8") as fh:
        json.dump({"site": "mx", "month": "2026-08", "expected": ["sales", "billing"], "missing": []}, fh)
    return folder


def test_order_basis_banner_matches_the_results_file(v2, tmp_path, capsys):
    from platforms.mercadolibre import report_build
    folder = order_folder(tmp_path)
    out = str(tmp_path / "out" / "TEST_2026-08.xlsx")
    version = R.version_id(folder, STAMP)
    rc = report_build.run(["--month", "2026-08", "--store", "TEST=%s" % folder, "--out", out, "--no-recalc",
                           "--results-version", version])
    assert rc == report_build.EXIT_OK
    [doc] = R.versions("TEST", "2026-08", "order")
    b = banner(out)
    assert b["完整度"] == R.GRADE_CN[doc["grade"]["level"]] == "完整"
    assert b["报表版本"] == version == doc["version"] and b["口径"] == "订单基准"
    assert b["数据下载于"] == "2026-09-01 00:00" and b["是否最终版本"].startswith("是（")
    assert b["_col"] == 5                                   # ML's ① lives in A–C
    assert "[完整度] 完整" in capsys.readouterr().out


def test_order_basis_gap_banner(v2, tmp_path):
    from platforms.mercadolibre import report_build
    folder = order_folder(tmp_path, banner_=True)           # the Ventas export says it is incomplete
    out = str(tmp_path / "x.xlsx")
    report_build.run(["--month", "2026-08", "--store", "TEST=%s" % folder, "--out", out, "--no-recalc",
                      "--results-version", R.version_id(folder, STAMP)])
    [doc] = R.versions("TEST", "2026-08", "order")
    b = banner(out)
    assert doc["grade"]["level"] == R.GAPS and b["完整度"] == "有缺口"
    assert b["原因"] == "；".join(doc["grade"]["reasons"]) and b["是否最终版本"].startswith("否（")


def test_order_basis_formula_errors_restamp(v2, tmp_path, monkeypatch, capsys):
    """Formula errors are found after the first stamp: the banner is written
    again with the final grade and the workbook recalculated once more."""
    from platforms.mercadolibre import report_build
    calls = []
    monkeypatch.setattr(report_build, "recalc", lambda out: calls.append(out) or (True, "已重算"))
    monkeypatch.setattr(report_build, "scan_formula_errors", lambda out: (3, 1, 0, ["②!C9 = #REF!"]))
    folder = order_folder(tmp_path)
    out = str(tmp_path / "x.xlsx")
    rc = report_build.run(["--month", "2026-08", "--store", "TEST=%s" % folder, "--out", out,
                           "--results-version", R.version_id(folder, STAMP)])
    assert rc == report_build.EXIT_CHECKS and len(calls) == 2
    [doc] = R.versions("TEST", "2026-08", "order")
    assert doc["grade"]["level"] == R.GAPS and "公式错误" in doc["grade"]["reasons"][0]
    assert banner(out)["完整度"] == "有缺口"
    text = capsys.readouterr().out
    assert "[完整度] 完整" in text and "[完整度] 有缺口" in text


def test_report_build_without_results_version_writes_no_banner(v2, tmp_path):
    """The golden rule: a builder run (no --results-version) leaves ① as built."""
    from platforms.mercadolibre import report_build
    folder = order_folder(tmp_path)
    out = str(tmp_path / "x.xlsx")
    report_build.run(["--month", "2026-08", "--store", "TEST=%s" % folder, "--out", out, "--no-recalc"])
    ws = load_workbook(out).worksheets[0]
    assert ws.max_column == 3 and K.GRADE_LABEL not in [c.value for c in ws[1]]
    assert ws.sheet_properties.tabColor is None
    # a two-store workbook is not a store workbook: no banner even with the option
    out2 = str(tmp_path / "y.xlsx")
    report_build.run(["--month", "2026-08", "--store", "TEST=%s" % folder, "--store", "TEST2=%s" % folder,
                      "--out", out2, "--no-recalc", "--results-version", R.version_id(folder, STAMP)])
    ws = load_workbook(out2).worksheets[0]
    assert ws.max_column == 3 and K.GRADE_LABEL not in [c.value for c in ws[1]]


# ------------------------------------------------------------------ MercadoLibre, bill basis

def test_bill_basis_banner_matches_the_results_file(v2, tmp_path, no_recalc, capsys):
    from platforms.mercadolibre import run_reports
    folder = BT.make_mx_folder(tmp_path)
    res = run_reports.bill_reports([("STORE_MX", folder)], "2026-09", "MX", str(tmp_path / "out"),
                                   combined=False, stamp=STAMP)
    assert res[0]["outcome"] in ("ok", "checks_failed")
    [doc] = R.versions("STORE_MX", "2026-09", "bill")
    b = banner(res[0]["out"])
    assert b["完整度"] == R.GRADE_CN[doc["grade"]["level"]]
    assert b["报表版本"] == doc["version"] == R.version_id(folder, STAMP) and b["口径"] == "账单基准"
    assert b["原因"] == ("；".join(doc["grade"]["reasons"]) or "—")
    assert b["数据下载于"] == "—"                           # the test folder's name has no HHMMSS
    assert "[完整度] %s" % b["完整度"] in capsys.readouterr().out


def test_bill_basis_without_stamp_writes_no_banner(v2, tmp_path, no_recalc):
    from platforms.mercadolibre import run_reports
    from platforms.mercadolibre.bill_report import build as BB
    folder = BT.make_mx_folder(tmp_path)
    res = run_reports.bill_reports([("STORE_MX", folder)], "2026-09", "MX", str(tmp_path / "out"), combined=False)
    assert banner(res[0]["out"]) == {}
    r = BB.run("STORE_MX", folder, "2026-09", "mx", out_dir=str(tmp_path / "o2"), recalc=False)
    assert banner(r["out"]) == {}


def test_bill_basis_formula_errors_restamp(v2, tmp_path, formula_errors):
    from platforms.mercadolibre import run_reports
    folder = BT.make_mx_folder(tmp_path)
    res = run_reports.bill_reports([("STORE_MX", folder)], "2026-09", "MX", str(tmp_path / "out"),
                                   combined=False, stamp=STAMP)
    assert res[0]["outcome"] == "formula_errors" and len(formula_errors) == 2
    [doc] = R.versions("STORE_MX", "2026-09", "bill")
    assert doc["grade"]["level"] == R.GAPS and banner(res[0]["out"])["完整度"] == "有缺口"


# ------------------------------------------------------------------ the other platforms (ReportRunner)

@pytest.fixture
def legacy(tmp_path, monkeypatch):
    cfg = {"default": S1, "allowed": [S1], "catalog": {S1: {"platform": "FA", "country": "CL"}},
           "batches": {"FA-CL": [S1]}}
    monkeypatch.setattr(store_config, "_stores", lambda: cfg)
    monkeypatch.setattr(paths, "ROOT", str(tmp_path))
    d = os.path.join(str(tmp_path), "downloads", S1, "20261001_083000_m202609")
    os.makedirs(d)
    with open(os.path.join(d, "run_meta.json"), "w", encoding="utf-8") as fh:
        json.dump({"expected": ["main"], "missing": []}, fh)
    return tmp_path


@pytest.fixture
def fake_build(monkeypatch):
    """A report_build.build that writes a real (small) family workbook."""
    mod = types.ModuleType("fake_platform_pkg.report_build")

    def build(pairs, month, out, strict=False):
        d = os.path.dirname(out)
        if d and not os.path.isdir(d):
            os.makedirs(d)
        wb = Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet(K.SHEET_NAMES[0])
        ws["B1"] = "月度财务报表 · %s" % month
        ws["C5"] = "、".join(n for n, _ in pairs)
        wb.create_sheet(K.SHEET_NAMES[1])["A1"] = "②"
        wb.save(out)
        return {"outcome": "ok", "out": out, "checks": [],
                "results": dict((n, {"figures": {"pl": []}, "meta": {"expected": ["main"], "missing": []}})
                                for n, _ in pairs)}
    mod.build = build
    monkeypatch.setitem(sys.modules, "fake_platform_pkg", types.ModuleType("fake_platform_pkg"))
    monkeypatch.setitem(sys.modules, "fake_platform_pkg.report_build", mod)
    return mod


def runner():
    periods = types.SimpleNamespace(previous_month=lambda today: "2026-09",
                                    parse_month=lambda m: (int(m[:4]), int(m[5:7])))
    return PR.ReportRunner("FA", "fake_platform_pkg", "FA-CL", periods)


def test_runner_banner_matches_the_results_file(legacy, fake_build, no_recalc, capsys):
    out_dir = str(legacy / "flat")
    assert runner().main(["--month", "2026-09", "--out-dir", out_dir]) == 0
    store_wb = os.path.join(out_dir, "%s_2026-09.xlsx" % S1)
    combined = os.path.join(out_dir, "CL_FA_全店汇总_2026-09.xlsx")
    [doc] = R.versions(S1, "2026-09", "main")
    b = banner(store_wb)
    assert b["完整度"] == R.GRADE_CN[doc["grade"]["level"]] == "完整"
    assert b["报表版本"] == doc["version"] and b["数据下载于"] == "2026-10-01 08:30"
    assert "口径" not in b and b["_col"] == 5                # the family's ① lives in A–C
    assert banner(combined) == {}                           # the combined workbook is not graded
    assert "[完整度] 完整" in capsys.readouterr().out


def test_runner_store_grade_is_the_results_grade(legacy, fake_build):
    r = fake_build.build([(S1, "/d")], "2026-09", str(legacy / "o" / "x.xlsx"))
    rn = runner()
    assert rn.store_grade(S1, r) == {"level": R.COMPLETE, "reasons": []}
    assert rn.store_grade(S1, r, formula_errors=True)["level"] == R.GAPS
    assert rn.store_grade(S1, {"outcome": "crashed", "checks": [], "error": "boom"})["level"] == R.FAILED
    doc = rn.results_document(S1, "/d/20261001_083000_m202609", "2026-09", "FA-CL", "v__20261022_101500",
                              dict(r, outcome="formula_errors"))
    assert doc["grade"]["level"] == R.GAPS


def test_runner_formula_errors_restamp(legacy, fake_build, formula_errors):
    out_dir = str(legacy / "flat")
    assert runner().main(["--month", "2026-09", "--out-dir", out_dir, "--no-combined"]) == 1
    store_wb = os.path.join(out_dir, "%s_2026-09.xlsx" % S1)
    assert formula_errors == [store_wb, store_wb]           # recalc, re-stamp, recalc again
    [doc] = R.versions(S1, "2026-09", "main")
    assert doc["grade"]["level"] == R.GAPS and banner(store_wb)["完整度"] == "有缺口"


def test_runner_safe_build_without_version_writes_no_banner(legacy, fake_build, no_recalc):
    out = str(legacy / "o" / "x.xlsx")
    r = runner().safe_build([(S1, "/d/20261001_083000_m202609")], "2026-09", out)
    assert r["outcome"] == "ok" and banner(out) == {}
