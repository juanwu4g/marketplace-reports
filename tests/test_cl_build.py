# -*- coding: utf-8 -*-
"""智利报表在两家真实店铺的下载目录上：10 页齐全、损益表勾稽、各条校验按事实过 / 不过。

数据（含真实订单）不进仓库，不在就跳过。路径可用环境变量覆盖：
    MELI_CL_DAB_DIR   DAB_CL (Coco) 子账号，2026-09，账单切期 8 号
    MELI_CL_TEC_DIR   Tecnohub_CL (个人店)，2026-09，账单切期 24 号、Octubre 账单未付款
下面断言的数字都是 2026-10-14 对着这两份目录逐项核过的。
"""
import io
import contextlib
import os

import pytest
from openpyxl import load_workbook

import report_build as rb

_SCRATCH = ("/tmp/claude-1001/-home-hat-meli-mx-pipeline/"
            "60f85e89-371b-4dd8-a325-f85924a7174a/scratchpad")
DAB = os.environ.get("MELI_CL_DAB_DIR", os.path.join(_SCRATCH, "cl_dab"))
TEC = os.environ.get("MELI_CL_TEC_DIR", os.path.join(_SCRATCH, "cl_tec"))
DAB_NAME, TEC_NAME = "DAB_CL (Coco) 子账号", "Tecnohub_CL (个人店)"

need_dab = pytest.mark.skipif(not os.path.isdir(DAB), reason="没有 DAB_CL 下载目录")
need_both = pytest.mark.skipif(not (os.path.isdir(DAB) and os.path.isdir(TEC)),
                               reason="没有两家智利店的下载目录")


@pytest.fixture(autouse=True)
def _back_to_mx():
    yield
    rb.set_site("mx")


def _derive(name, folder):
    rb.set_site("cl")
    p0, p1 = rb.month_bounds("2026-09")
    with contextlib.redirect_stdout(io.StringIO()):
        return rb.derive(name, folder, p0, p1)


def _checks(S):
    return {c["key"]: c for c in rb.validate([S])}


@need_dab
def test_dab_numbers_and_checks():
    S = _derive(DAB_NAME, DAB)
    m = S["m"]
    assert S["cycle_day"] == 8
    assert [(y, mo) for y, mo, _a, _b in S["bill_cycles"]] == [(2026, 9), (2026, 10)]
    assert round(m["ing"]) == 47978708
    assert abs(m["resid"]) <= 1
    # 桥 B：两张账单的付款按计费日拆回 9 月，分毫不差；每张账单整张也各自平
    assert round(m["payable"]) == round(m["pay_auto"]) == 2739295
    assert [round(c["payable"] - c["paid"]) for c in m["pay_cycles"]] == [0, 0]
    # 没有代扣：佣金列 − 账单佣金 ≈ 退货处理费
    assert round(m["tax_ret"]) == 56228 and round(m["vcol_dev"]) == 58402
    assert (m["com_match_n"], m["com_match_of"]) == (2389, 2406)
    c = _checks(S)
    for key in ("bridge_a", "bridge_b", "fee_catalog", "bill_block_balance",
                "no_withholding", "cargos_full", "billing_present", "sales_window"):
        assert c[key]["ok"] and not c[key]["skipped"], (key, c[key]["detail"])
    assert "tax_rate" not in c                       # 墨西哥的代扣税率校验不适用
    assert not c["ventas_incomplete"]["ok"]          # 平台横幅：数据不全（如实报）
    # MercadoPago 逐单：250 单不一致里 216 单正好差买家另付的运费
    MO = S["mp_orders"]
    assert (MO["linked"], MO["agree"], len(MO["differ"])) == (2434, 2184, 250)
    assert MO["ship_explained"][0] == 216


@need_both
def test_tecnohub_open_bill_and_cycle_24():
    S = _derive(TEC_NAME, TEC)
    assert S["cycle_day"] == 24
    assert S["bill_open"] == ["Octubre 2026"]
    c = _checks(S)
    assert c["bridge_b"]["skipped"]                  # Octubre 账单还没付款：不能下结论
    assert "Octubre 2026" in c["bridge_b"]["detail"]
    assert c["no_withholding"]["ok"] and c["cargos_full"]["ok"]
    assert S["m"]["storage_days"] == 30              # 两期仓储报表合起来覆盖整月


@need_both
def test_combined_workbook_ties_and_speaks_clp(tmp_path):
    out = str(tmp_path / "cl.xlsx")
    with contextlib.redirect_stdout(io.StringIO()):
        rc = rb.main(["--month", "2026-09", "--store", "%s=%s" % (DAB_NAME, DAB),
                      "--store", "%s=%s" % (TEC_NAME, TEC),
                      "--out", out, "--no-recalc", "--build-date", "2026-10-14"])
    assert rc in (rb.EXIT_OK, rb.EXIT_CHECKS)
    wb = load_workbook(out)
    assert wb.sheetnames == rb.SHEET_NAMES
    pl = wb[rb.SHEET_NAMES[1]]
    rows = {pl.cell(r, 1).value: r for r in range(5, pl.max_row + 1) if pl.cell(r, 1).value}
    # 经营贡献毛利 = 收入行到账单级兜底行全部数据行之和（小计行是公式，跳过）
    for col, want in (("C", 33065458.3), ("D", -23771.1)):
        tot = sum(pl["%s%d" % (col, r)].value for r in range(
            rows["商品销售收入 (GMV)"], rows["减：账单级未登记科目 ⚠"] + 1)
            if isinstance(pl["%s%d" % (col, r)].value, (int, float)))
        assert round(tot, 1) == want
        assert pl["%s%d" % (col, rows["口径尾差（应为 0）"])].value == 0
    assert not any("代扣代缴" in str(k) for k in rows)  # 没有代扣税行
    assert pl["C%d" % rows["商品销售收入 (GMV)"]].number_format == "#,##0;(#,##0);-"
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    assert "MXN" not in cell.value and "墨西哥比索" not in cell.value, \
                        (ws.title, cell.coordinate)
                    assert not cell.value.startswith("= "), (ws.title, cell.coordinate)
