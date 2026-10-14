# -*- coding: utf-8 -*-
"""墨西哥报表在加入智利之后必须逐格不变。

拿一份真实的 MX 下载目录（BOCINA_TA02，2026-09）用当前代码出报表，和改动前
那一版代码出的报表逐格比对：每一页每个单元格的值（公式按字符串比）和数字格式。
基线 JSON 是加智利之前用旧代码生成的，存在仓库外（含真实店铺数据，不进仓库）。

两样东西任一不在（别的机器、CI）就跳过。路径可用环境变量覆盖：
    MELI_MX_GOLDEN_DIR       MX 下载目录
    MELI_MX_GOLDEN_BASELINE  基线 JSON

编制日期是唯一会随运行日期变的单元格（① 阅读说明!A2），这里用 --build-date
固定成基线生成时的 2026-10-14，所以要求是"一个格子都不差"。
"""
import datetime
import json
import os

import pytest
from openpyxl import load_workbook

import report_build

_SCRATCH = ("/tmp/claude-1001/-home-hat-meli-mx-pipeline/"
            "60f85e89-371b-4dd8-a325-f85924a7174a/scratchpad")
GOLDEN_DIR = os.environ.get(
    "MELI_MX_GOLDEN_DIR",
    os.path.join(_SCRATCH, "golden/downloads/BOCINA_TA02/20261014_025745_m202609"))
BASELINE = os.environ.get(
    "MELI_MX_GOLDEN_BASELINE", os.path.join(_SCRATCH, "clwork/mx_baseline.json"))
BUILD_DATE = "2026-10-14"


def dump(path):
    """{页名: {坐标: 值}, 页名#fmt: {坐标: 数字格式}, _sheets: [页名...]}，与基线同一格式。"""
    wb = load_workbook(path)
    out = {"_sheets": wb.sheetnames}
    for ws in wb.worksheets:
        cells, fm = {}, {}
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is not None:
                    if isinstance(v, (datetime.date, datetime.datetime)):
                        v = v.isoformat()
                    elif not isinstance(v, (int, float, str, bool)):
                        v = str(v)
                    cells[c.coordinate] = v
                if c.number_format != "General":
                    fm[c.coordinate] = c.number_format
        out[ws.title] = cells
        out[ws.title + "#fmt"] = fm
    return json.loads(json.dumps(out))


@pytest.fixture(autouse=True)
def _back_to_mx():
    yield
    report_build.set_site("mx")


@pytest.mark.skipif(not (os.path.isdir(GOLDEN_DIR) and os.path.isfile(BASELINE)),
                    reason="没有 MX 黄金目录或基线 JSON（只在有真实数据的机器上跑）")
def test_mx_workbook_identical_cell_for_cell(tmp_path):
    # 先切到智利再出墨西哥：main() 必须按目录把站点切回来，不能沿用上一次的。
    report_build.set_site("cl")
    out = str(tmp_path / "mx.xlsx")
    rc = report_build.main(["--month", "2026-09", "--store", "BOCINA_TA02=%s" % GOLDEN_DIR,
                            "--out", out, "--no-recalc", "--build-date", BUILD_DATE])
    assert rc in (report_build.EXIT_OK, report_build.EXIT_CHECKS)
    with open(BASELINE, encoding="utf-8") as fh:
        want = json.load(fh)
    got = dump(out)
    assert got["_sheets"] == want["_sheets"]
    diffs = []
    for key in sorted(set(want) | set(got)):
        a, b = want.get(key), got.get(key)
        if isinstance(a, dict) and isinstance(b, dict):
            for c in sorted(set(a) | set(b)):
                if a.get(c) != b.get(c):
                    diffs.append("%s!%s: %r -> %r" % (key, c, a.get(c), b.get(c)))
        elif a != b:
            diffs.append("%s: %r -> %r" % (key, a, b))
    assert not diffs, "MX 报表变了 %d 处，前几处：\n%s" % (len(diffs), "\n".join(diffs[:20]))
