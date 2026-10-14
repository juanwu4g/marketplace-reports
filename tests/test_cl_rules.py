# -*- coding: utf-8 -*-
"""智利规则的单元测试：会计月归属（账单周期）、CLP 格式与容差、"不代扣"的判定。

全部用合成数据，不需要真实下载目录。
"""
import os

import pandas as pd
import pytest

import financial_report
import report_build as rb


@pytest.fixture(autouse=True)
def _back_to_mx():
    yield
    rb.set_site("mx")


T = pd.Timestamp


# ---------------------------------------------------------------- 会计月归属

def test_mx_bill_is_the_calendar_month():
    rb.set_site("mx")
    assert rb.bill_cycles(2026, 9) == [(2026, 9, T(2026, 9, 1), T(2026, 9, 30))]


def test_cl_month_spans_two_bills_cut_on_the_8th():
    rb.set_site("cl")
    assert rb.bill_cycles(2026, 9, 8) == [
        (2026, 9, T(2026, 8, 8), T(2026, 9, 7)),     # "Septiembre" 账单 → 9/1–9/7
        (2026, 10, T(2026, 9, 8), T(2026, 10, 7)),   # "Octubre" 账单 → 9/8–9/30
    ]


def test_cl_cut_day_is_per_store():
    rb.set_site("cl")
    assert rb.bill_cycles(2026, 9, 24) == [
        (2026, 9, T(2026, 8, 24), T(2026, 9, 23)),
        (2026, 10, T(2026, 9, 24), T(2026, 10, 23)),
    ]
    # 跨年
    assert [(y, m) for y, m, _a, _b in rb.bill_cycles(2026, 12, 8)] == [(2026, 12), (2027, 1)]


def _bill(rows):
    """合成账单：[(文件名, 计费日), ...]"""
    return pd.DataFrame({"_file": [f for f, _ in rows], "_doc": "FAC",
                         "cdate": pd.to_datetime([d for _, d in rows])})


def test_infer_cycle_day_from_bill_files(tmp_path):
    rb.set_site("cl")
    B = _bill([("Reporte_Facturacion_MercadoLibre_Sept2026.xlsx", "2026-08-24"),
               ("Reporte_Facturacion_MercadoLibre_Sept2026.xlsx", "2026-09-23"),
               ("Reporte_Facturacion_MercadoLibre_Oct2026.xlsx", "2026-09-24")])
    assert rb.infer_cycle_day(B, str(tmp_path)) == 24


def test_infer_cycle_day_from_storage_file_names(tmp_path):
    rb.set_site("cl")
    for n in ("08-09-26_07-10-26_Costos_por_servicio_almacenamiento.xlsx",
              "08-10-26_07-11-26_Costos_por_servicio_almacenamiento.xlsx"):
        open(os.path.join(str(tmp_path), n), "w").close()
    assert rb.infer_cycle_day(_bill([]), str(tmp_path)) == 8


def test_infer_cycle_day_falls_back_to_site_default(tmp_path):
    rb.set_site("cl")
    assert rb.infer_cycle_day(_bill([]), str(tmp_path)) == rb.SITE_PROFILES["cl"]["BILL_CYCLE_DAY"]


def test_financial_report_wants_both_cl_bills(tmp_path):
    with open(os.path.join(str(tmp_path), "run_meta.json"), "w") as fh:
        fh.write('{"site": "cl"}')
    assert financial_report.bills_needed(str(tmp_path), (2026, 12)) == {(2026, 12), (2027, 1)}
    mx = tmp_path / "mx"
    mx.mkdir()
    assert financial_report.bills_needed(str(mx), (2026, 9)) == {(2026, 9)}   # 没有 site = mx


def test_financial_report_accepts_ventas_cl(tmp_path):
    for n in ("20261014_Ventas_CL_x.xlsx", "Reporte_Facturacion_MercadoLibre_Sept2026.xlsx"):
        open(os.path.join(str(tmp_path), n), "w").close()
    assert financial_report._usable(str(tmp_path)) == (True, [])


# ---------------------------------------------------------------- CLP 格式与容差

def test_set_site_switches_columns_format_and_tolerances():
    rb.set_site("cl")
    assert rb.VENTAS_SHEET == "Ventas CL"
    assert rb.V_REVENUE == "Ingresos por productos (CLP)"
    assert rb.V_TOTAL == "Total (CLP)"
    assert rb.MNY == "#,##0;(#,##0);-"
    assert rb.fmt(1234567.4) == "1,234,567"
    assert rb.fmt(1234.5, 2) == "1,234.50"         # 显式要小数位时照给
    assert (rb.TOL_CENT, rb.CN_TOL, rb.BRIDGE_A_TOL) == (1.0, 1.0, 1.0)
    rb.set_site("mx")
    assert rb.V_REVENUE == "Ingresos por productos (MXN)"
    assert rb.fmt(1234567.456) == "1,234,567.46"
    assert (rb.TOL_CENT, rb.CN_TOL, rb.BRIDGE_A_TOL) == (0.01, 0.05, 0.05)


def test_unknown_site_is_refused():
    with pytest.raises(ValueError):
        rb.set_site("ar")


def test_bridge_a_tolerance_is_one_peso():
    S = {"store": "X", "m": {"resid": 0.6}}
    rb.set_site("cl")
    assert rb._chk("bridge_a", "X", "", abs(S["m"]["resid"]) <= rb.BRIDGE_A_TOL, "")["ok"]
    rb.set_site("mx")
    assert not abs(S["m"]["resid"]) <= rb.BRIDGE_A_TOL


def test_cl_reversal_column_is_cargo_que_anula(tmp_path):
    """智利账单把冲销指向列叫 Cargo que anula（墨西哥叫 Cargo que bonifica）。"""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "REPORT"
    head = [rb.F_CHARGE_DATE, rb.F_CHARGE_NO, rb.F_DETALLE, rb.F_CHANNEL,
            "Cargo que anula", rb.F_AMOUNT]
    for i, h in enumerate(head, 1):
        ws.cell(rb.BILLING_HEADER_ROW + 1, i, h)
    rows = [("2026-09-02", "1", "Cargo por campaña de publicidad de Product Ads", "No", None, 100),
            ("2026-09-03", "2", "Anulación del cargo por campaña de publicidad de Product Ads",
             "No aplica", "1", -100)]
    for r, row in enumerate(rows, rb.BILLING_HEADER_ROW + 2):
        for i, v in enumerate(row, 1):
            ws.cell(r, i, v)
    wb.save(str(tmp_path / "Reporte_Facturacion_MercadoLibre_Sept2026.xlsx"))
    B = rb.load_billing(str(tmp_path))
    rev = B[B["is_rev"]].iloc[0]
    assert rev["reverses"] == "1" and rev["parent_channel"] == "No"


# ---------------------------------------------------------------- 不代扣的判定

def _S(tax_ret, dev, ing=1000000.0, mp_tax=0.0):
    return {"store": "X", "mp": {"tax": mp_tax} if mp_tax is not None else None,
            "m": {"tax_ret": tax_ret, "vcol_dev": dev, "vcol_rest": tax_ret - dev,
                  "ing": ing, "com_match_n": 99, "com_match_of": 100}}


def test_no_withholding_holds_when_gap_is_return_fees():
    rb.set_site("cl")
    c = rb._no_withholding_check(_S(56228, 58402, ing=47978708))
    assert c["ok"] and c["key"] == "no_withholding" and c["severity"] == "error"


def test_no_withholding_fails_if_platform_starts_withholding():
    rb.set_site("cl")
    # 墨西哥式的代扣 ~9% 会让剩余远超 0.5% 的上限
    assert not rb._no_withholding_check(_S(90000, 0))["ok"]
    # 流水里出现了税金
    assert not rb._no_withholding_check(_S(0, 0, mp_tax=-500))["ok"]
    # 没有流水不算失败（只是少一条证据）
    assert rb._no_withholding_check(_S(0, 0, mp_tax=None))["ok"]


def test_cl_pl_has_no_withholding_rows():
    rb.set_site("mx")
    mx_keys = [r[0] for r in rb.build_pl_rows()]
    rb.set_site("cl")
    assert [r[0] for r in rb.build_pl_rows()] == mx_keys       # 行结构两国相同
    notes = dict((r[0], r[4]) for r in rb.build_pl_rows())
    assert "代扣" not in notes["cn_adj"] and "不代扣" in notes["com_neg"]
    memo = rb.pl_memo_no_withholding([{"store": "X", "m": {"vcol_dev": 1.0, "vcol_rest": 0.0}}])
    labels = [x[0] for x in memo]
    assert not any("代扣代缴" in l for l in labels)
    assert any("IVA" in l for l in labels)
    assert all(x[3] != "sens" for x in memo)      # 没有"代扣税不可抵扣"的敏感性行
