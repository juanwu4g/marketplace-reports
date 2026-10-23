# -*- coding: utf-8 -*-
"""The monthly results files (core/results, docs/RESULTS_AND_PERIODS.md §2) of every
platform and basis, on the synthetic download folders of the report tests: the
file's figures are the workbook's own (data rows and 备查项 cell by cell, computed
rows equal to the report's own totals), and failed builds leave a failed version."""
import json
import os
import subprocess

import pytest
from openpyxl import load_workbook

from core import paths
from core import report_kit as K
from core import results as R
from core import store_config

import test_az_report as AZT
import test_fa_report as FAT
import test_ml_bill_report as BT
import test_reports_build as OT

STAMP = "20261022_101500"


@pytest.fixture
def v2(tmp_path, monkeypatch):
    """New layout; the test stores registered in a catalog."""
    stores = {"STORE_MX": ("ML", "MX"), "TEST": ("ML", "MX"), FAT.STORE: ("FA", "CL"),
              AZT.STORE: ("AZ", "MX")}
    cfg = tmp_path / "config.json"
    cfg.write_text(json.dumps({
        "paths": {"data_root": str(tmp_path / "data")},
        "stores": {"allowed": sorted(stores), "catalog": dict(
            (s, {"platform": p, "country": c}) for s, (p, c) in stores.items())}}), encoding="utf-8")
    monkeypatch.setattr(store_config, "CONFIG_PATH", str(cfg))
    monkeypatch.setattr(paths, "ROOT", str(tmp_path / "code"))
    return tmp_path


def pl_sheet(path, col=3):
    """② column `col` by row label (the first occurrence of each label)."""
    ws = load_workbook(path)[K.SHEET_NAMES[1]]
    out = {}
    for row in ws.iter_rows(min_row=1):
        a = row[0].value
        if isinstance(a, str) and len(row) >= col:
            out.setdefault(a, row[col - 1].value)
    return out


def assert_data_rows_and_memo(doc_like, path):
    """Every ② data / input row and every 备查项 amount equals the sheet's cell."""
    cells = pl_sheet(path)
    n = 0
    for r in doc_like["pl"]:
        if r["kind"] in ("d", "in"):
            assert cells[r["label"]] == pytest.approx(r["value"], abs=0.005), r["label"]
            n += 1
    for r in doc_like["memo"]:
        v = cells[r["label"]]
        if isinstance(v, str) and v.startswith("="):           # the sensitivity row is a formula
            continue
        assert v == pytest.approx(r["value"], abs=0.005), r["label"]
        n += 1
    assert n > 5


def pl_of(doc):
    return dict((r["key"], r["value"]) for r in doc["pl"])


# ------------------------------------------------------------------ MercadoLibre, bill basis

def test_bill_basis_results_file(v2, tmp_path):
    from platforms.mercadolibre import run_reports
    folder = BT.make_mx_folder(tmp_path)
    out_dir = str(tmp_path / "out")
    res = run_reports.bill_reports([("STORE_MX", folder)], "2026-09", "MX", out_dir, combined=False, stamp=STAMP)
    assert res[0]["outcome"] in ("ok", "checks_failed")
    [doc] = R.versions("STORE_MX", "2026-09", "bill")
    assert doc["version"] == R.version_id(folder, STAMP) and doc["run"] == os.path.basename(folder)
    assert doc["platform"] == "ML" and doc["market"] == "ML-MX" and doc["currency"] == "MXN"
    assert doc["workbook"] == os.path.abspath(res[0]["out"]) and doc["source_folders"] == [os.path.abspath(folder)]
    assert_data_rows_and_memo(doc, doc["workbook"])
    S = BT.derive(folder)
    for k, v in S["m"]["pl"].items():                       # computed rows = the model's totals
        assert pl_of(doc)[k] == pytest.approx(v), k
    p = pl_of(doc)
    assert p["profit"] == pytest.approx(p["contrib"] - p["cogs"])
    assert sum(f["amount"] for f in doc["fees"]) == pytest.approx(p["fees_total"])
    assert all(f["pl_key"] in p for f in doc["fees"])
    assert sum(x["amount"] for x in doc["returns"]) == pytest.approx(
        p["rev_cancel"] + p["rev_return"] + p["rev_partial"] + p["rev_unproven"])
    assert sum(x["contribution"] for x in doc["sku"]) == pytest.approx(p["contrib"])
    assert all(x["cost"] is None for x in doc["sku"])        # no cost table
    assert doc["inventory"] is None or doc["inventory"]["rows"]
    assert doc["grade"]["level"] in R.GRADES and doc["checks"]


def test_bill_basis_crash_leaves_a_failed_version(v2, tmp_path, monkeypatch):
    from platforms.mercadolibre import run_reports
    from platforms.mercadolibre.bill_report import workbook as WB
    folder = BT.make_mx_folder(tmp_path)

    def boom(*a, **k):
        raise RuntimeError("disk full")
    monkeypatch.setattr(WB, "build", boom)
    res = run_reports.bill_reports([("STORE_MX", folder)], "2026-09", "MX", str(tmp_path / "o"), combined=False,
                                   stamp=STAMP)
    assert res[0]["outcome"] == "crashed"
    [doc] = R.versions("STORE_MX", "2026-09", "bill")
    assert doc["grade"]["level"] == R.FAILED and "disk full" in doc["grade"]["reasons"][0]
    assert doc["workbook"] is None and doc["pl"] == []


def test_bill_basis_without_stamp_writes_no_results_file(v2, tmp_path):
    from platforms.mercadolibre import run_reports
    run_reports.bill_reports([("STORE_MX", BT.make_mx_folder(tmp_path))], "2026-09", "MX", str(tmp_path / "o"),
                             combined=False)
    assert R.versions("STORE_MX", "2026-09") == []


# ------------------------------------------------------------------ MercadoLibre, order basis

def test_order_basis_results_file_written_by_report_build(v2, tmp_path):
    from platforms.mercadolibre import report_build
    folder = OT.make_download_folder(tmp_path / "TEST" / "20260901_000000_m202608")
    out, metrics = tmp_path / "out" / "TEST_2026-08.xlsx", tmp_path / "m.json"
    version = R.version_id(folder, STAMP)
    rc = report_build.run(["--month", "2026-08", "--store", "TEST=%s" % folder, "--out", str(out), "--no-recalc",
                           "--results-version", version, "--metrics-out", str(metrics)])
    assert rc == report_build.EXIT_OK
    [doc] = R.versions("TEST", "2026-08", "order")
    assert doc["version"] == version and doc["basis"] == "order" and doc["workbook"] == str(out)
    assert doc["grade"] == {"level": R.COMPLETE, "reasons": []}
    assert_data_rows_and_memo(doc, str(out))
    p = pl_of(doc)
    rec = json.load(open(str(metrics), encoding="utf-8"))[0]
    assert p["contrib"] == pytest.approx(rec["contribution"], abs=0.01)
    assert p["revenue"] == pytest.approx(rec["ing"] + rec["env_ing"] - rec["coupon"], abs=0.01)
    assert p["profit"] == pytest.approx(p["contrib"])           # no cost table: 商品采购成本 0
    assert [r["key"] for r in doc["pl"]][:4] == ["ing", "env_ing", "coupon_neg", "revenue"]
    assert {m["key"] for m in doc["memo"]} >= {"withheld_taxes", "contrib_if_tax_not_deductible"}
    sens = dict((m["key"], m["value"]) for m in doc["memo"])
    assert sens["contrib_if_tax_not_deductible"] == pytest.approx(p["contrib"] - sens["withheld_taxes"])
    assert [f["pl_key"] for f in doc["fees"]] and all(f["pl_key"] in p for f in doc["fees"])
    assert sum(f["amount"] for f in doc["fees"]) == pytest.approx(-75.0)       # the August bill: 15+10+30+20
    assert {x["sku"] for x in doc["sku"]} == {"SKU-A", "SKU-B"}
    assert doc["returns"] == [] and doc["inventory"] is None


def test_order_basis_banner_is_a_gap(v2, tmp_path):
    from platforms.mercadolibre import report_build
    folder = OT.make_download_folder(tmp_path / "TEST" / "20260901_000000_m202608", banner=True)
    report_build.run(["--month", "2026-08", "--store", "TEST=%s" % folder, "--out", str(tmp_path / "x.xlsx"),
                      "--no-recalc", "--results-version", R.version_id(folder, STAMP)])
    [doc] = R.versions("TEST", "2026-08", "order")
    assert doc["grade"]["level"] == R.GAPS and doc["flags"]["data_incomplete"] is True


def test_no_results_file_without_the_option_or_for_a_combined_workbook(v2, tmp_path):
    from platforms.mercadolibre import report_build
    folder = OT.make_download_folder(tmp_path / "TEST" / "20260901_000000_m202608")
    report_build.run(["--month", "2026-08", "--store", "TEST=%s" % folder, "--out", str(tmp_path / "x.xlsx"),
                      "--no-recalc"])
    report_build.run(["--month", "2026-08", "--store", "TEST=%s" % folder, "--store", "TEST2=%s" % folder,
                      "--out", str(tmp_path / "y.xlsx"), "--no-recalc", "--results-version",
                      R.version_id(folder, STAMP)])
    assert R.versions("TEST", "2026-08") == []


def test_financial_report_passes_the_version_only_for_one_store(monkeypatch, tmp_path):
    from platforms.mercadolibre import financial_report as FR
    seen = []

    def fake_run(cmd, **kw):
        seen.append(cmd)
        return subprocess.CompletedProcess(cmd, 0, stdout="", stderr=None)
    monkeypatch.setattr(FR.subprocess, "run", fake_run)
    FR.build([("A", "/d/a")], "2026-09", str(tmp_path / "a.xlsx"), no_recalc=True, results_version="r__20261022_000000")
    FR.build([("A", "/d/a"), ("B", "/d/b")], "2026-09", str(tmp_path / "b.xlsx"), no_recalc=True,
             results_version="r__20261022_000000")
    assert seen[0][-2:] == ["--results-version", "r__20261022_000000"]
    assert "--results-version" not in seen[1]


def test_order_failed_writes_a_failed_version_once(v2, tmp_path):
    from platforms.mercadolibre import run_reports
    version = R.version_id("/d/20261001_000000_m202609", STAMP)
    r = {"outcome": "crashed", "checks": []}
    assert run_reports.order_failed("TEST", "/d/20261001_000000_m202609", "2026-09", "MX", version, r)
    assert run_reports.order_failed("TEST", "/d/20261001_000000_m202609", "2026-09", "MX", version, r) is None
    [doc] = R.versions("TEST", "2026-09", "order")
    assert doc["grade"]["level"] == R.FAILED and "生成出错" in doc["grade"]["reasons"][0]


# ------------------------------------------------------------------ Falabella, Amazon

def test_falabella_figures_are_the_workbooks(tmp_path):
    folder = FAT.make_folder(tmp_path)
    r, out = FAT.build(folder, tmp_path)
    rec = r["results"][FAT.STORE]
    assert rec["error"] is None and rec["folder"] == folder
    fig = rec["figures"]
    assert_data_rows_and_memo(fig, out)
    m = r["metrics"][FAT.STORE]
    p = dict((x["key"], x["value"]) for x in fig["pl"])
    assert p["contrib"] == pytest.approx(m["contrib"]) and p["revenue"] == pytest.approx(m["revenue"])
    assert p["net_sales"] == pytest.approx(m["net_sales"]) and p["profit"] == pytest.approx(m["contrib"])
    assert sum(x["contribution"] for x in fig["sku"]) == pytest.approx(m["contrib"])
    assert all(f["pl_key"] in p for f in fig["fees"])
    assert "订单数（不含已取消）" not in [x["label"] for x in fig["memo"]]     # a count, not an amount
    assert fig["inventory"]["sold_window_days"] == 28 and fig["inventory"]["rows"]
    assert fig["returns"] and all(isinstance(x["count"], int) for x in fig["returns"])


@pytest.fixture
def az_text_pdf(monkeypatch):
    def fake(path):
        with open(path, encoding="utf-8") as fh:
            return fh.read(), "test"
    monkeypatch.setattr(AZT.R, "pdf_text", fake)


def test_amazon_figures_are_the_workbooks(tmp_path, az_text_pdf):
    folder = AZT.make_folder(tmp_path)
    r, out = AZT.build(folder, tmp_path)
    fig = r["results"][AZT.STORE]["figures"]
    assert_data_rows_and_memo(fig, out)
    m = r["metrics"][AZT.STORE]
    p = dict((x["key"], x["value"]) for x in fig["pl"])
    assert p["contrib"] == pytest.approx(m["contrib"], abs=0.01)
    assert p["fee_total"] == pytest.approx(sum(f["amount"] for f in fig["fees"]), abs=0.05)
    assert fig["returns"][0]["amount"] == pytest.approx(p[AZT.TC.REFUND])
    assert fig["inventory"]["rows"] and fig["inventory"]["as_of"]
    assert {"iva_out", "contrib_if_tax_not_deductible"} <= {x["key"] for x in fig["memo"]}


def test_extractor_failure_is_reported_not_raised(tmp_path, monkeypatch):
    folder = FAT.make_folder(tmp_path)

    def boom(S):
        raise KeyError("x")
    monkeypatch.setattr(FAT.R, "results_figures", boom)
    r, out = FAT.build(folder, tmp_path)
    assert r["outcome"] in ("ok", "checks_failed") and os.path.isfile(out)
    assert r["results"][FAT.STORE]["figures"] is None and "KeyError" in r["results"][FAT.STORE]["error"]


def test_runner_writes_results_in_legacy_layout_where_reports_stay_flat(tmp_path, monkeypatch):
    """Falabella's run_reports (core/platform_reports) in the old layout: workbooks
    flat as before, one results file per store, none for the summary."""
    from platforms.falabella import run_reports
    monkeypatch.setattr(paths, "ROOT", str(tmp_path))
    cfg = {"default": FAT.STORE, "allowed": [FAT.STORE],
           "catalog": {FAT.STORE: {"platform": "FA", "country": "CL"}}, "batches": {"FA-CL": [FAT.STORE]}}
    monkeypatch.setattr(store_config, "_stores", lambda: cfg)
    monkeypatch.setattr(K, "recalc", lambda out: (False, "跳过重算"))
    d = tmp_path / "downloads" / FAT.STORE
    folder = FAT.make_folder(tmp_path)          # tmp/downloads/<store>/<run>
    assert os.path.dirname(folder) == str(d)
    out = tmp_path / "flat"
    assert run_reports.main(["--month", FAT.MONTH, "--out-dir", str(out)]) in (0, 1)
    assert sorted(os.listdir(str(out))) == ["CL_FA_全店汇总_2026-09.xlsx", "%s_2026-09.xlsx" % FAT.STORE]
    [doc] = R.versions(FAT.STORE, FAT.MONTH, "main")
    assert doc["workbook"] == str(out / ("%s_2026-09.xlsx" % FAT.STORE)) and doc["market"] == "FA-CL"
    assert doc["currency"] == "CLP" and doc["decimals"] == 0 and doc["pl"]
    assert os.path.dirname(os.path.dirname(doc["_path"])).endswith(os.path.join("results", "FA", "CL", FAT.STORE))
