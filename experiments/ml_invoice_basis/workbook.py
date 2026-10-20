# -*- coding: utf-8 -*-
"""The experiment's two workbooks, in the report family's look (core/report_kit,
core/report_family): navy titles, blue section bands, grey formula rows, Chinese
labels with the platform's own term next to them.

  <store>_<YYYY-MM>_账单口径实验.xlsx   ① 说明 ② 差异摘要（给财务） ③ 账单口径损益
                                       ④ 与现行订单口径对比 ⑤ 费用明细 ⑥ 销售与冲销明细
                                       ⑦ 退款优惠券与现金 ⑧ 三种日期与应计 ⑨ 口径稳定性（多次下载）
  <store>_<YYYY-MM>_口径稳定性.xlsx     ① 说明 ② 稳定性总表 ③ 订单变化明细 ④ 账单行变化明细

Everything is written from the plain data of report.compute() / stability.analyse();
nothing here reads the MercadoLibre site globals."""

import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from core import report_family as RF
from core import report_kit as K
from core.report_kit import BOX, F, fill, wrap
from platforms.mercadolibre.report.catalog import CHANNEL_MEANING, CHANNEL_ZH, FEE_CATALOG

from experiments.ml_invoice_basis import accounting
from experiments.ml_invoice_basis import mercadopago as MP
from experiments.ml_invoice_basis.common import snapshot_label
from experiments.ml_invoice_basis.compare import CAT, order_key, rec_summary
from experiments.ml_invoice_basis.invoice import LINK_ORDER, bill_name

SHEETS = ["① 说明", "② 差异摘要（给财务）", "③ 账单口径损益", "④ 与现行订单口径对比", "⑤ 费用明细",
          "⑥ 销售与冲销明细", "⑦ 退款优惠券与现金", "⑧ 三种日期与应计", "⑨ 口径稳定性（多次下载）"]
STAB_SHEETS = ["① 说明", "② 稳定性总表", "③ 订单变化明细", "④ 账单行变化明细"]
UNAVAILABLE = "无法得到"
BUCKET_ZH = {"commission": "佣金", "shipping": "配送", "returns": "退货", "ads": "广告",
             "storage": "仓储", "fulfillment": "Full 履约", "penalty": "罚金",
             "subscription": "订阅", "misc": "其他（待确认）"}


# ════════════════════════════════════════════════════════════════════════
# small writers
# ════════════════════════════════════════════════════════════════════════

def _d(v):
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    return str(v)


def _span(view):
    return "%s–%s" % (view["start"].strftime("%Y-%m-%d"), view["end"].strftime("%m-%d"))


def _put(ws, r, c, v, fmt=None, bold=False, color="000000", fillc=None, wrap_text=False,
         border=True, italic=False, size=10):
    if v is not None and not isinstance(v, (str, int, float)):
        v = _d(v)
    if isinstance(v, float):
        v = None if pd.isna(v) else round(v, 6)     # no 3.6e-12 float noise on the page
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = F(size, bold, color, italic)
    if fmt:
        cell.number_format = fmt
    if fillc:
        cell.fill = fill(fillc)
    if wrap_text:
        cell.alignment = wrap("top")
    if border:
        cell.border = BOX
    return cell


def _head(ws, r, cols, height=30):
    K.header_cells(ws, r, cols)
    ws.row_dimensions[r].height = height
    return r + 1


def _band(ws, r, text, ncols):
    return RF.section(ws, r, text, ncols)


def _note(ws, r, text, ncols, color=K.C_GREY, bold=False):
    c = ws.cell(row=r, column=1, value=text)
    c.font = F(10, bold, color)
    c.alignment = wrap("top")
    if ncols > 1:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
    ws.row_dimensions[r].height = K.est_height(text, per_line=max(40, ncols * 14))
    return r + 1


def _fmt_of(kind, money):
    return {"m": money, "n": K.INT, "p": K.PCT, "r": "0.0000"}.get(kind)


def _table(ws, r, cols, rows, money):
    """cols: [(header, kind)] kind t/m/n/p/r/d; rows: lists of values."""
    r = _head(ws, r, [c[0] for c in cols])
    for row in rows:
        for j, ((_h, kind), v) in enumerate(zip(cols, row), 1):
            if kind == "d":
                v = _d(v)
            _put(ws, r, j, v, _fmt_of(kind, money), wrap_text=(kind == "w"))
        r += 1
    return r


def _value(ws, r, c, v, money, kind="m", bold=False, fillc=None):
    if v is None:
        return _put(ws, r, c, UNAVAILABLE, color=K.C_RED, fillc=fillc)
    return _put(ws, r, c, v, _fmt_of(kind, money), bold=bold, fillc=fillc)


def _diff_formula(a, b, r):
    return '=IF(AND(ISNUMBER(%s%d),ISNUMBER(%s%d)),%s%d-%s%d,"")' % (b, r, a, r, b, r, a, r)


# ════════════════════════════════════════════════════════════════════════
# ① 说明
# ════════════════════════════════════════════════════════════════════════

def _periods_text(D):
    P, C = D["invs"]["period"]["view"], D["invs"]["calendar"]["view"]
    y, m = int(D["month"][:4]), int(D["month"][5:7])
    if D["cycle_day"] == 1:
        p = ("%s 账单：计费日 %s 至 %s。%s的账单按自然月出，所以账单期就是自然月，两种读法数字相同。"
             % (bill_name((y, m)), _d(P["start"]), _d(P["end"]), D["country"]))
    else:
        p = ("%s 账单：计费日 %s 至 %s（本店账单切期日 = 每月 %d 号，按本店自己的账单/仓储文件推出；"
             "%s的账单不是自然月，名为“%s”的账单覆盖上月 %d 号到本月 %d 号）。"
             % (bill_name((y, m)), _d(P["start"]), _d(P["end"]), D["cycle_day"], D["country"],
                bill_name((y, m)), D["cycle_day"], D["cycle_day"] - 1))
    c = ("计费日 %s 至 %s 的账单行，取自 %s。" % (_d(C["start"]), _d(C["end"]),
                                       " + ".join(bill_name(b) for b in C["bills"])))
    return p, c


def _money(D):
    return lambda v: "{:,.{d}f}".format(v, d=D["dec"])


def _mp_files_text(D):
    F = D["mp"]["files"]
    have = ["%s：%s" % (MP.FILE_ZH[k], os.path.basename(p)) for k, p in F.items() if p]
    if not have:
        return MP.NO_FILES + "（settlement_v2、collection、reserve-release、after_collection 都没有）。" \
                             "退款、优惠券只能沿用估算或写“无法得到”，现金核对做不了。"
    return "；".join(have) + "。实际采用：%s。" % MP.source_text(D["mp"]) + \
        ("".join(D["mp"]["notes"]))


def policies(D, pct, absv):
    """The accounting policies of this experiment, as statements (①)."""
    v = D["vat"]
    r = v["rate"]
    return [
        ("", "1 收入在哪个月", "账单口径：一笔销售在平台计佣金的那天（Cargo por venta 的计费日）进入当月；"
                          "⑧ 同时给出按下单日、按交付日分月的结果。按会计准则（IFRS 15 / NIF D-1 / 智利 IFRS）"
                          "商品控制权在交付时转移，严格说应按交付日；平台数据里交付日没有年份、未交付的没有日期，"
                          "所以本实验以计费日为主、交付日作对照。"),
        ("", "2 取消与退款", "在冲销出现的账单期确认（与收入同一套凭证）。金额优先取 MercadoPago 实际退给买家的钱"
                         "（settlement_v2 的 REFUND，加上买家赢了的争议/拒付 DISPUTE）；一笔销售在几张账单里分次冲销的，"
                         "实际退款按本期冲销的佣金占比分摊。没有流水的沿用“冲销佣金 ÷ 原佣金 × 售价”的估算；"
                         "下单后未付款就取消的销售（流水里没有收款）按全额冲销。估算与实际的差逐笔列在 ⑦。"),
        ("", "3 卖家承担的优惠券", "是销售折扣，冲减收入。金额取 MercadoPago 结算流水：每笔收款的 TRANSACTION_AMOUNT "
                             "+ FEE_AMOUNT + TAXES_AMOUNT − REAL_AMOUNT（除佣金和税以外被扣的钱）。它就是月度报表由 "
                             "Ventas 的 Total 倒推的那个数（逐单验证见 ⑦）。退款时退回的优惠券同步冲回。Cobros 报表的 "
                             "coupon_fee 不是它（是整场活动的折扣，含平台出资部分），不采用。"),
        ("", "4 争议与拒付", "买家赢了的争议（mediación / contracargo）退给买家的钱并入实际退款，单独显示金额。"
                          "争议里“只扣卖家、不退买家”的那一行就是账单的退货处理费（Cargo por devolución，"
                          "同一收款号、同一金额，⑦ 有核对），已在平台费用里，不重复计。"),
        ("", "5 增值税（%s）" % v["name"],
         "收入与平台费用都改为不含税列报（÷ %g），销项税、进项税放在 ③ 的备忘栏，不进损益；"
         "含税的列保留，用来和现行报表（含税口径）对比。%s。证据见下面“增值税”一节。"
         % (1 + r, v["credit_why"])),
        ("", "6 应计（截止）", "本期销售的订单级费用（联盟佣金、退货处理费……）如果下一张账单（已下载的，关账或未关账）"
                          "里才计费，计入本期“应计费用”。还没出账单的费用无法得知，不估计——这是账单口径的局限。"),
        ("", "7 代扣税（%s）" % ("墨西哥 IVA/ISR" if v["withholding"] else "本国不适用"),
         ("不是费用，是可以抵减应缴税额的预缴税款，所以不进经营贡献毛利。正式金额以 MercadoLibre/MercadoPago 每月出具的"
          " CFDI de retenciones（Constancia de retenciones de IVA e ISR）为准——这个文件还没有下载（以后的下载步骤）；"
          "MercadoPago 流水的 TAXES_AMOUNT 只作参考（IVA 与 ISR 混在一起，不能拆）。")
         if v["withholding"] else "本国平台不代扣税（MercadoPago 流水 TAXES_AMOUNT 为 0）。"),
        ("", "8 商品成本", "配置了 SKU 成本表才计算：成本件数 = 本期开票件数 − 全额冲销的件数（部分冲销视为退款不退货），"
                        "单位成本按成本表（不含可抵扣的增值税）。没有配置时写“未配置 SKU 成本表”，不估计。"),
        ("", "9 重要性", "② 把差额 ≥ 现行净收入的 %g%% 或 ≥ %s %s 的行标为“重要”（两个标准都可在 ② 顶部改，"
                       "也可用命令行 --threshold-pct / --threshold-abs 改）。" % (pct, "{:,.0f}".format(absv), D["cur"])),
        ("", "10 现金", "⑦ 从调整后的经营贡献毛利走到 MercadoPago 本月释放给卖家的钱：每一步都是两份数据的实测差，"
                       "逐笔拆开；拆不开的写“余数”，不强行抹平。没有 settlement_v2 的下载做不了这一段。"),
    ]


def readme_sheet(wb, D, A=None, pct=None, absv=None):
    pct = accounting.DEFAULT_PCT if pct is None else pct
    absv = accounting.DEFAULT_ABS if absv is None else absv
    ws = wb.create_sheet(SHEETS[0])
    y, mth = int(D["month"][:4]), int(D["month"][5:7])
    K.title_rows(ws, "美客多 账单口径实验 · %s · %d年%d月" % (D["store"], y, mth),
                 "实验报表，不是月度财务报表，不参与每月的运行。同一店同一月的正式报表仍是 %s_%s.xlsx。"
                 % (D["store"], D["month"]))
    issues = D["invs"]["period"]["view"]["issues"] + D["invs"]["calendar"]["view"]["issues"]
    issues = [i for i in issues if "Notas de crédito 文件" not in i]
    if issues:
        ws["A3"] = "⚠ " + "；".join(sorted(set(issues)))
        ws["A3"].font = F(10, True, K.C_RED)
    RF.set_widths(ws, [2, 26, 120])
    P, C = D["invs"]["period"], D["invs"]["calendar"]
    AP, AC = D["acc"]["period"], D["acc"]["calendar"]
    S, m = D["S"], D["S"]["m"]
    pt, ct = _periods_text(D)
    money = _money(D)
    docs = []
    for d in D["docs"]:
        n = int((S["billing_all"]["_file"] == d["name"]).sum())
        docs.append(("", d["name"], "%s · %s · %d 行" % (
            "账单 Facturación" if d["doc"] == "FAC" else "贷记单 Notas de crédito",
            "还没关账（文件顶部写着 podría tener cambios）" if d["open"] else "已关账", n)))
    status = (P["view"]["issues"] + C["view"]["issues"]) or ["两种读法用到的账单都已下载、都已关账。"]
    ship_note = "（墨西哥只有少数订单买家另付运费）" if D["site"] == "mx" else ""
    cannot = [("", lab, "%s —— %s" % (
        "参考值 %s（取自现行订单口径用的文件，不计入账单口径）" % money(v) if v is not None else "两边都没有",
        why)) for lab, _es, v, why in D["unavailable"]]
    methods = P["n"]["rev_methods"]
    vt = D["vat"]
    verdict = {True: "数据证明：含税", False: "数据没有一致证明（按含税处理，要复核）",
               None: "文件不足以核对（按含税处理，属于假设）"}[D["vat_verdict"]]
    rows = [
        ("S", "一、这份报表是什么", ""),
        ("", "目的", "实验：如果把损益的依据从销售报表（Ventas）换成当月的账单（Facturación + Notas de "
                    "crédito），能做到多好、和现行报表差多少、哪个口径更稳定。现行月度报表一个数都不改。"),
        ("", "给财务的摘要", "先看 %s：从现行报表的经营贡献毛利一步一步走到账单口径（含税 → 实际 → 不含税），"
                         "每一步的金额、占净收入的比例、是否重要、原因类别。" % SHEETS[1]),
        ("", "现行订单口径", "收入、退款、运费收入、优惠券取 Ventas，按下单日归月；佣金和退货处理费按订单号到"
                         "所有已下载的账单里去找（不论哪张账单计费）；配送费取 Ventas；广告/仓储等月度费用按"
                         "计费日归自然月。Ventas 是平台的实时导出，订单状态、退款每天都在变。"),
        ("", "账单口径（本实验）", "收入、冲销、费用全部取账单里的行：每笔销售在账单上有一行“Cargo por venta”，"
                              "带着售价、件数、单价、下单日和商品；取消/退款在账单上是“Anulación del cargo"
                              " por venta”（冲销），费用按科目逐行列出。账单关账后就是定稿的凭证。账单里没有的"
                              "（实际退款、优惠券、代扣税、现金）从 MercadoPago 补，补不上的明说。"),
        ("S", "二、期间（重要）", ""),
        ("", "账单期", pt),
        ("", "自然月切", ct + "这一列给出“按自然月看”的读法，方便和现行报表对比；② 的摘要用这一列。"),
        ("", "现行订单口径", "下单日 %s 至 %s（Ventas 导出于 %s）。"
         % (_d(D["S"]["p0"]), _d(D["S"]["p1"]), _d(S["meta"].get("downloaded_at", "")[:10]) or "?")),
        ("", "账单状态", "；".join(status)),
        ("S", "三、用到的文件", ""),
        ("", "下载目录", D["folder"]),
    ] + docs + [
        ("", "Ventas（现行口径用）", "%s%s" % (os.path.basename(S["files"].get("ventas") or ""),
                                         "；⚠ 平台提示导出可能缺数据：%s" % S["ventas_warning"]
                                         if S.get("ventas_warning") else "")),
        ("", "Pagos_Facturas", "；".join(os.path.basename(p) for p in sorted(
            set(x["file"] for x in D["cash"]["pagos"]))) or "（见 ⑦ 资金去向）"),
        ("", "MercadoPago", _mp_files_text(D)),
        ("S", "四、增值税（%s %g%%）：平台费用的 Valor del cargo 含不含税？—— %s"
         % (vt["name"], round(vt["rate"] * 100, 2), verdict), ""),
    ] + [("", "证据 %d" % (i + 1) if not t.startswith(("结论", "假设")) else t[:2], t)
         for i, t in enumerate(D["vat_lines"])] + [
        ("", "现行报表怎么处理", "现行月度报表全部按含税口径列示：GMV 含税（桥 F 另算“其中 %s 成分 = GMV × %g/%g”与"
                             "“不含税收入（参考）”，即 derive() 的 iva_in_gmv / net_revenue_exvat），平台费用直接用 Valor del "
                             "cargo、不拆税；桥 A 里的“账单口径纯佣金”指的是不含代扣税（墨西哥 IVA/ISR 预扣），不是不含 %s。"
                             "智利版桥 F 还写着“账单只给一个 Valor del cargo，没有单列 IVA”——上面的证据回答了这个问题。"
         % (vt["name"], round(vt["rate"] * 100, 2), round(100 + vt["rate"] * 100, 2), vt["name"])),
        ("", "本实验怎么列", "③ 损益表同时给出含税两列（与现行报表、④ 对比）和不含税两列（收入 ÷ %g；费用%s）；"
                           "销项税、进项税、应纳税额在 ③ 的“增值税备忘”。"
         % (1 + vt["rate"], " ÷ %g，进项税可抵扣" % (1 + vt["rate"]) if vt["credit"] else "保持含税（进项税不可抵扣）")),
        ("S", "五、账单能给什么", ""),
        ("", "开票销售", "每笔被计佣金的销售：售价（Total de la venta）、件数、单价、下单日、商品、包裹号。"),
        ("", "取消与退款", "冲销行（Anulación del cargo por venta，在账单或贷记单里）。冲掉的销售额按比例估算，见“方法”。"),
        ("", "平台费用", "全部科目：佣金、配送费、退货处理费、广告、Full 仓储/揽收/罚金、页面维护费……以及它们的冲销。"
                      "科目名称和中文与现行报表一致（FEE_CATALOG）；字典里没有的科目单独列出、标 ⚠，绝不并进别的行。"),
        ("", "买家支付的运费", "账单的“Envío a cargo del cliente”列。实测与 Ventas 的“Ingresos por envío”逐单相等"
                          "（DAB_CL 592/592、BOCINA_TA02 13/13）%s。" % ship_note),
        ("S", "六、账单不能给什么（发票口径无法得到）", ""),
    ] + cannot + [
        ("S", "七、会计政策（本实验的处理；财务可以提出不同意见）", ""),
    ] + policies(D, pct, absv) + [
        ("S", "八、方法", ""),
        ("", "期间", "账单期 = 名为本月的账单文件（Facturación + 同名的 Notas de crédito）的全部行；"
                   "自然月切 = 所有已下载账单里计费日（Fecha del cargo）在本月的行。"),
        ("", "销售", "每个 Número de venta 算一笔（以 Cargo por venta 行为准；联盟佣金行只加佣金，不重复计销售）。"
                   "多件包裹在账单上每件一笔，和 Ventas 对照时按 Número de paquete 并回包裹。"),
        ("", "冲销销售额（估算）", "冲销比例 = 冲销的佣金 ÷ 这笔销售的原佣金；估算冲销销售额 = 比例 × 售价。"
                            "依据：平台按退款比例同时退还佣金（现行报表的贷记单核对已用 MercadoPago 现金逐单证实）。"
                            "原佣金的找法依次是：%s。本期 %d 笔冲销：%s。"
         % ("；".join("（%d）%s" % (i + 1, x) for i, x in enumerate(LINK_ORDER)), P["n"]["n_rev"],
            "，".join("%s %d 笔" % (x.split("（")[0].split("：")[0], methods[x])
                     for x in LINK_ORDER if methods.get(x)) or "无")),
        ("", "MercadoPago ↔ 账单", "账单每一行的“Pago”就是 MercadoPago 的收款号（settlement_v2 的 SOURCE_ID、Cobros 的 "
                                 "operation_id），按它把收款、退款、争议、优惠券、代扣税挂到每笔销售上；几笔销售共用一个"
                                 "收款号时按售价分摊。"),
        ("", "运费收入", "买家支付的运费取 Cargo por envíos 行上的“Envío a cargo del cliente”；配送费被冲销时，"
                      "按同样比例算作退回买家。"),
        ("", "费用", "账单净额 = 计费 − 冲销（含贷记单）。全部账单行金额之和 = 费用合计，③ 页有核对行（残差应为 0）。"),
        ("", "经营贡献毛利", "账单口径（估算）= 净开票销售额 + 买家运费（净）− 平台费用合计；"
                          "调整后 = 估算 + 冲销估算回拨 − 实际退款 − 卖家优惠券 − 应计费用；不含税 = 调整后各行去掉增值税。"
                          "都没有扣代扣税（可抵税款）和商品成本（除非配置了成本表）。"),
        ("S", "九、本期要点", ""),
        ("", "商品销售额", "现行 %s；账单期 %s；自然月切 %s。"
         % (money(m["ing"]), money(P["n"]["gross_sales"]), money(C["n"]["gross_sales"]))),
        ("", "经营贡献毛利", "现行 %s（含税）；账单口径·自然月切：估算 %s → 调整后 %s（含税）→ 不含税 %s。"
                          "账单期：估算 %s → 调整后 %s → 不含税 %s。"
         % (money(S["contrib"]), money(C["n"]["contribution"]), money(AC["contrib_adj"]),
            money(AC["vat"]["contrib_net"]), money(P["n"]["contribution"]), money(AP["contrib_adj"]),
            money(AP["vat"]["contrib_net"]))),
        ("", "平台费用合计", "现行 %s；账单期 %s；自然月切 %s。"
         % (money(_cur_fees(D)), money(-P["n"]["fee_total"]), money(-C["n"]["fee_total"]))),
        ("", "取消与退款：估算 vs 实际", _rev_point(D, AC, money)),
        ("", "卖家优惠券", ("MercadoPago 实际（自然月切）%s：开票销售上 %s，退款退回 %s；现行（Ventas 倒推）%s。"
                         % (money(AC["coupon"]), money(AC["mp"]["coupon_gross"]), money(AC["mp"]["coupon_back"]),
                            money(m["coupon"])))
         if AC["mp"]["coupon_available"] else MP.source_text(D["mp"]) + "：优惠券无法得到（现行 Ventas 倒推 %s）"
         % money(m["coupon"])),
        ("", "应计费用", "自然月切 %s（%d 行，%d 笔销售）；账单期 %s。"
         % (money(AC["accrual"]["total"]), AC["accrual"]["n"], AC["accrual"]["n_sales"],
            money(AP["accrual"]["total"]))),
        ("", "增值税", "自然月切：销项税 %s，进项税 %s（%s），差 %s。"
         % (money(AC["vat"]["output"]), money(AC["vat"]["input"]), "可抵扣" if vt["credit"] else "不可抵扣",
            money(AC["vat"]["output"] - AC["vat"]["input"]))),
        ("", "现金", _cash_point(D, money)),
        ("", "差异怎么来的", "摘要见 %s；逐行原因见 %s；逐单名单见 %s；退款/优惠券/现金见 %s；日期与应计见 %s。"
         % (SHEETS[1], SHEETS[3], SHEETS[5], SHEETS[6], SHEETS[7])),
    ]
    if A and A.get("pairs"):
        for p in A["pairs"]:
            rows.append(("", "稳定性 %s → %s" % (snapshot_label(p["a"]["folder"])[:10],
                                              snapshot_label(p["b"]["folder"])[:10]),
                         "订单口径：%s 账单口径：%s %s" % (" ".join(p["why"]["order"][:2]),
                                                   p["why"]["period"], p["why"]["calendar"])))
    rows += [
        ("S", "十、各页", ""),
        ("", SHEETS[1], "给财务：现行 → 账单口径的每一步差额、占净收入、是否重要（标准可改）、原因类别。"),
        ("", SHEETS[2], "账单口径的损益表：账单期与自然月切，含税与不含税；调整为实际（MercadoPago）与应计；"
                        "商品成本；增值税备忘；账单得不到的数单列备查。"),
        ("", SHEETS[3], "与现行订单口径逐行对比（含税、冲销按估算）：差额、原因类别、原因说明；订单 ↔ 开票销售的逐单对照摘要。"),
        ("", SHEETS[4], "费用按科目：计费、冲销、净额；未登记科目单列。"),
        ("", SHEETS[5], "冲销明细（关联方法）、只在一边的订单/销售、金额不一致的订单、开票销售明细。"),
        ("", SHEETS[6], "MercadoPago：实际退款 vs 估算（逐笔）、优惠券、争议、代扣税；现金核对；资金去向。"),
        ("", SHEETS[7], "应计（下期账单已出现的本期费用）；每笔销售的下单日、计费日、交付日与三种分月结果。"),
        ("", SHEETS[8], "同一店同一月多次下载时，两种口径各变了多少、为什么（含实际退款、优惠券、不含税数）。"),
    ]
    RF.readme_rows(ws, 5, rows)
    return ws


def _rev_point(D, AC, money):
    t = AC["rev"]
    if not AC["mp"]["available"]:
        return "%s：%d 笔冲销全部沿用估算 %s。" % (MP.source_text(D["mp"]), t["n"], money(t["est"]))
    return ("自然月切 %d 笔冲销：估算 %s；其中 %d 笔有 MercadoPago 实际数，估算 %s 对实际 %s，误差（估算 − 实际）%s，"
            "%d 笔分毫不差；%d 笔沿用估算 %s。实际里有争议/拒付 %s。"
            % (t["n"], money(t["est"]), t["n_known"], money(t["est_known"]), money(t["actual_known"]),
               money(t["err"]), t["n_exact"], t["n_unknown"], money(t["est_unknown"]), money(t["dispute"])))


def _cash_point(D, money):
    C = D["cash"]
    if not C["have"]:
        return "%s：做不了现金核对。" % MP.source_text(D["mp"])
    L = dict((x["key"], x) for x in C["lines"])
    v = lambda k: _line_value(C["lines"], k)
    return ("调整后毛利 %s → 推算应收货款 %s；这些收款实际 %s，差 %s（逐项拆开见 ⑦）；"
            "本月按释放日释放给卖家 %s%s。"
            % (money(v("c1")), money(v("c5")), money(v("c6")), money(v("c7")), money(v("c14")),
               "（Liberaciones 报表 %s）" % money(L["c15"]["value"]) if "c15" in L else ""))


def _line_value(lines, k):
    """The value of a bridge line, its formula sums computed in Python (for texts)."""
    by = dict((x["key"], x) for x in lines)
    x = by[k]
    if x.get("sum_of"):
        tot = 0.0
        for s in x["sum_of"]:
            sign = -1.0 if s.startswith("-") else 1.0
            tot += sign * (_line_value(lines, s.lstrip("-")) or 0.0)
        return tot
    return x.get("value") or 0.0


def _cur_fees(D):
    m = D["S"]["m"]
    return -(m["com_gross"] + m["env_cost"] + m["dev_net"] + m["unc_order"]
             + sum(m["bill_fee_net"].values()) + m["unc_bill"])


# ════════════════════════════════════════════════════════════════════════
# ② 差异摘要（给财务）
# ════════════════════════════════════════════════════════════════════════

def summary_sheet(wb, D, pct, absv):
    ws = wb.create_sheet(SHEETS[1])
    money = D["money"]
    K.title_rows(ws, "差异摘要（给财务）· %s · %s" % (D["store"], D["month"]),
                 "从现行订单口径的经营贡献毛利出发，一步一步走到账单口径（%s，自然月切 %s）。"
                 "差额 = 账单口径 − 现行；正数 = 账单口径多。"
                 % (D["country"], _span(D["invs"]["calendar"]["view"])),
                 "“重要”= 差额的绝对值 ≥ 下面两个标准之一。黄色格子可以改，改了自动重算。")
    RF.set_widths(ws, [46, 17, 17, 17, 11, 9, 16, 90])
    base = accounting.base_net_revenue(D)
    _put(ws, 4, 1, "现行订单口径 · 净收入（GMV − 退款 + 修正 + 运费 − 优惠券）", bold=True)
    _put(ws, 4, 3, base, money, bold=True)
    _put(ws, 5, 1, "重要性标准一：占净收入的比例")
    _put(ws, 5, 3, pct / 100.0, K.PCT2, fillc=K.C_INPUT)
    _put(ws, 6, 1, "重要性标准二：金额（%s）" % D["cur"])
    _put(ws, 6, 3, absv, money, fillc=K.C_INPUT)
    r = _head(ws, 8, ["项目", "现行订单口径", "账单口径\n（自然月切）", "差额 / 金额", "占净收入", "重要？",
                      "原因类别", "说明"], height=36)
    first = r
    rows = accounting.summary_rows(D)
    at = {}
    for x in rows:
        k = x["kind"]
        fc = {"start": K.C_SUB, "sub": K.C_CALC, "end": K.C_TOTAL}.get(k)
        bold = k != "diff"
        _put(ws, r, 1, x["label"], bold=bold, fillc=fc)
        _put(ws, r, 2, x.get("cur") if k == "diff" else "", money, fillc=fc)
        _put(ws, r, 3, x.get("inv") if k == "diff" else "", money, fillc=fc)
        _put(ws, r, 4, x["value"], money, bold=bold, fillc=fc)
        if k == "diff":
            _put(ws, r, 5, '=IF($C$4=0,"",D%d/$C$4)' % r, K.PCT2)
            _put(ws, r, 6, '=IF(OR(ABS(D%d)>=$C$6,ABS(D%d)>=ABS($C$4)*$C$5),"⚠ 重要","")' % (r, r),
                 color=K.C_RED, bold=True)
        else:
            _put(ws, r, 5, "", fillc=fc)
            _put(ws, r, 6, "", fillc=fc)
        _put(ws, r, 7, x["cat"], fillc=fc)
        _put(ws, r, 8, x["why"], color=K.C_GREY, wrap_text=True, fillc=fc)
        if len(x["why"] or "") > 80:
            ws.row_dimensions[r].height = K.est_height(x["why"], 80)
        at.setdefault(x["key"], r)
        r += 1
    last = r - 1
    diffs = lambda a, b: ",".join("D%d" % i for i in range(a + 1, b))
    r += 1
    r = _band(ws, r, "核对（应为 0）", 8)
    checks = [("起点 + 逐行差额 − 账单口径（估算）", at["start"], at["mid"]),
              ("账单口径（估算）+ 调整 − 调整后", at["mid"], at["adj"]),
              ("调整后 + 增值税列报 − 不含税", at["adj"], at["end"])]
    for lab, a, b in checks:
        _put(ws, r, 1, "残差（应为 0）", bold=True, fillc=K.C_TOTAL)
        _put(ws, r, 2, lab, color=K.C_GREY, fillc=K.C_TOTAL)
        mid = diffs(a, b)
        _put(ws, r, 4, "=D%d%s-D%d" % (a, ("+SUM(%s)" % mid) if mid else "", b), money, bold=True,
             fillc=K.C_TOTAL)
        r += 1
    r += 1
    _put(ws, r, 1, "超过标准的差异", bold=True)
    _put(ws, r, 4, '=COUNTIF(F%d:F%d,"⚠ 重要")' % (first, last), K.INT, bold=True)
    _put(ws, r, 8, "项（按上面两个标准，任一满足即算）", color=K.C_GREY)
    r += 2
    r = _band(ws, r, "原因类别是什么意思", 8)
    for cat, text in (
            ("时间差", "同一笔业务，两边归到了不同的月份（下单日 ↔ 计费日；本月销售的费用下月才计费）。跨月会自己抵消。"),
            ("账单里没有", "账单本身不含这一项（卖家优惠券、实际退款、现金），要从 MercadoPago 补；补不上就是“无法得到”。"),
            ("估算与实际", "账单只给冲销的佣金，退了多少钱是按比例估算的；有 MercadoPago 流水时换成实际数。"),
            ("取数来源不同", "两边取自不同的表（Ventas ↔ 账单）或现行报表有专门的修正。"),
            (accounting.REASON_VAT, "现行报表含税；账单口径另给不含税的数（增值税不是卖家的钱）。只是列报方式，不是业务差异。")):
        _put(ws, r, 1, cat, bold=True)
        _put(ws, r, 2, text, color=K.C_GREY)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        r += 1
    r += 1
    notes = ["代扣税：两边的经营贡献毛利都不扣（是可抵税款）。%s"
             % ("正式金额要看 CFDI de retenciones（还没下载）。" if D["vat"]["withholding"] else "本国平台不代扣。"),
             "本页用自然月切（与现行口径同一个月）。%s"
             % ("智利的账单期（%s）在 ③、④ 另列。" % _span(D["invs"]["period"]["view"]) if D["cycle_day"] != 1
                else "墨西哥/秘鲁的账单期就是自然月。"),
             "MercadoPago：%s。" % MP.source_text(D["mp"])]
    if not D["mp"]["source"]:
        notes.append("所以“冲销：估算 → 实际”和“卖家优惠券”两行是 0：不是没有差异，是本次下载没有 MercadoPago 文件，算不出来。")
    for t in notes:
        r = _note(ws, r, t, 8)
    ws.freeze_panes = "B9"
    return ws


# ════════════════════════════════════════════════════════════════════════
# ③ 账单口径损益
# ════════════════════════════════════════════════════════════════════════

def pl_sheet(wb, D):
    ws = wb.create_sheet(SHEETS[2])
    money = D["money"]
    P, C = D["invs"]["period"], D["invs"]["calendar"]
    AP, AC = D["acc"]["period"], D["acc"]["calendar"]
    vt = D["vat"]
    div = "%g" % (1 + vt["rate"])
    pt, ct = _periods_text(D)
    K.title_rows(ws, "账单口径损益表（实验）· %s · %s" % (D["store"], D["month"]),
                 "账单期 = " + pt, "自然月切 = " + ct)
    RF.set_widths(ws, [40, 40, 17, 17, 17, 17, 80])
    ws["A4"] = ("含税两列与现行报表、④ 可比；不含税两列 = 收入 ÷ %s，费用%s。%s 的证据见 ①。"
                % (div, " ÷ %s（进项税可抵扣）" % div if vt["credit"] else "保持含税（进项税不可抵扣）", vt["name"]))
    ws["A4"].font = F(9, False, K.C_GREY)
    r = _head(ws, 5, ["项目", "原文（平台科目）", "账单期 · 含税\n%s" % _span(P["view"]),
                      "自然月切 · 含税\n%s" % _span(C["view"]), "账单期 · 不含税", "自然月切 · 不含税", "说明"],
              height=40)
    at = {}

    def row(k, label, es, kind, vp, vc, note, style=None, vat="rev", fx=None):
        """vat: rev (÷ 1+rate), fee (÷ when creditable), same (copy), blank (no net columns).
        fx: keys whose rows add up to this row (each column its own formula)."""
        nonlocal r
        fillc = {"sum": K.C_SUB, "f": K.C_CALC, "tot": K.C_TOTAL, "info": None}.get(style)
        bold = style in ("sum", "f", "tot")
        _put(ws, r, 1, label, bold=bold, fillc=fillc, color=K.C_GREY if style == "info" else "000000")
        _put(ws, r, 2, es, color=K.C_GREY, fillc=fillc)
        fm = _fmt_of(kind, money)
        for col, v in (("C", vp), ("D", vc)):
            c = "CD".index(col) + 3
            if fx:
                _put(ws, r, c, "=" + "+".join("%s%d" % (col, at[x]) for x in fx), fm, bold=bold, fillc=fillc)
            elif isinstance(v, str) and v.startswith("="):
                _put(ws, r, c, v, fm, bold=bold, fillc=fillc)
            else:
                _value(ws, r, c, v, money, kind, bold, fillc)
        for src, col in (("C", "E"), ("D", "F")):
            c = "EF".index(col) + 5
            if vat == "blank":
                _put(ws, r, c, "", fillc=fillc)
            elif fx:
                _put(ws, r, c, "=" + "+".join("%s%d" % (col, at[x]) for x in fx), fm, bold=bold, fillc=fillc)
            elif isinstance(vp if src == "C" else vc, str) and str(vp if src == "C" else vc).startswith("=SUM"):
                f = (vp if src == "C" else vc).replace(src, col)
                _put(ws, r, c, f, fm, bold=bold, fillc=fillc)
            elif vat == "same" or (vat == "fee" and not vt["credit"]):
                _put(ws, r, c, "=%s%d" % (src, r), fm, bold=bold, fillc=fillc)
            else:
                _put(ws, r, c, "=%s%d/%s" % (src, r, div), fm, bold=bold, fillc=fillc)
        _put(ws, r, 7, note, color=K.C_GREY, wrap_text=True, fillc=fillc)
        if note and len(note) > 70:
            ws.row_dimensions[r].height = K.est_height(note, 62)
        at[k] = r
        r += 1

    n = lambda inv, k: inv["n"][k]
    r = _band(ws, r, "收入（账单的 Cargo por venta / Cargo por envíos 行）", 7)
    row("n_sales", "开票销售笔数", "Número de venta", "n", n(P, "n_sales"), n(C, "n_sales"),
        "一件商品一笔；按计费日落在期间内（下单月份分布见下）", vat="same")
    row("units", "开票件数", "Cantidad vendida", "n", n(P, "units"), n(C, "units"), "", vat="same")
    row("gross", "开票销售额", "Total de la venta（Cargo por venta 行）", "m", n(P, "gross_sales"),
        n(C, "gross_sales"), "买家实付商品金额（含税售价）；每笔销售算一次")
    row("rev", "减：冲销销售额（估算）", "Anulación del cargo por venta", "m", -n(P, "rev_sales"),
        -n(C, "rev_sales"),
        "冲销佣金 ÷ 原佣金 × 售价。账单期 %d 笔（全额 %d 笔），自然月切 %d 笔；逐笔见 ⑥，实际数见下面“调整为实际”"
        % (n(P, "n_rev"), n(P, "rev_full"), n(C, "n_rev")))
    row("net_sales", "净开票销售额", "", "m", None, None, "＝ 开票销售额 − 冲销销售额", "f",
        fx=["gross", "rev"])
    row("ship", "加：买家支付的运费", "Envío a cargo del cliente（Cargo por envíos 行）", "m",
        n(P, "ship_income"), n(C, "ship_income"),
        "实测与 Ventas 的 Ingresos por envío 逐单相等；账单期 %d 笔" % n(P, "ship_n"))
    row("ship_rev", "减：运费退回买家（随配送费冲销）", "Anulación del cargo por envíos", "m",
        -n(P, "ship_income_rev"), -n(C, "ship_income_rev"), "配送费被冲销时，按同样比例算作退回买家的运费")
    row("income", "净收入合计", "", "m", None, None, "", "sum", fx=["net_sales", "ship", "ship_rev"])
    r = _band(ws, r, "平台费用（账单净额 = 计费 − 冲销，含贷记单；负数 = 费用）", 7)
    first = None
    fp = dict((x["es"], x) for x in P["fees"])
    fc = dict((x["es"], x) for x in C["fees"])
    for es, zh, bucket, order_level in FEE_CATALOG:
        a, b = fp[es], fc[es]
        row("fee:" + es, "减：" + zh, es, "m", -a["net"], -b["net"],
            "%s；账单期计费 %d 笔 %s，冲销 %d 笔 %s" % (
                "订单级" if order_level else "账单级", a["n_gross"],
                "{:,.{d}f}".format(a["gross"], d=D["dec"]), a["n_rev"],
                "{:,.{d}f}".format(a["rev"], d=D["dec"])), vat="fee")
        first = first or at["fee:" + es]
    unk = {}
    for v, inv in (("p", P), ("c", C)):
        for u in inv["unknown"]:
            unk.setdefault(u["es"], {"p": 0.0, "c": 0.0, "n": 0})[v] += u["amount"]
    if unk:
        for es, u in sorted(unk.items()):
            row("unk:" + es, "减：⚠ 未登记科目", es, "m", -u["p"], -u["c"],
                "FEE_CATALOG 里没有这个 Detalle：金额照常计入，科目要人工确认", vat="fee")
            for c in range(1, 8):
                ws.cell(row=at["unk:" + es], column=c).fill = fill(K.C_BAD)
    else:
        row("unk", "减：未登记科目 ⚠", "(未在 FEE_CATALOG 登记)", "m", 0.0, 0.0, "本期没有", vat="fee")
    last = r - 1
    row("fees", "平台费用合计", "", "m", "=SUM(C%d:C%d)" % (first, last),
        "=SUM(D%d:D%d)" % (first, last), "", "sum", vat="fee")
    row("contrib", "经营贡献毛利（账单口径，未扣商品成本）", "冲销按估算", "m", None, None,
        "＝ 净收入合计 + 平台费用合计；未扣代扣税、卖家优惠券、商品成本（见下）。与现行报表、④ 对比用这一行", "tot",
        fx=["income", "fees"])
    r += 1
    r = _band(ws, r, "核对", 7)
    row("chk", "账单行金额合计（取负）", "Σ Valor del cargo", "m", -n(P, "fee_total"), -n(C, "fee_total"),
        "期间内所有账单行（Facturación + Notas de crédito）的金额之和", vat="blank")
    row("resid", "残差（应为 0）", "", "m", "=C%d-C%d" % (at["fees"], at["chk"]),
        "=D%d-D%d" % (at["fees"], at["chk"]), "上面逐科目加总 − 全部账单行；不为 0 说明有行没归到任何科目", "tot",
        vat="blank")
    r += 1
    src = MP.source_text(D["mp"])
    have = AC["mp"]["available"]
    r = _band(ws, r, "调整为实际（MercadoPago：%s）与应计" % src, 7)
    row("adj_est", "加回：冲销销售额（估算）", "上面的估算行", "m", n(P, "rev_sales"), n(C, "rev_sales"),
        "先把估算加回来，再减实际")
    row("adj_act", "减：取消与退款（实际）", "settlement_v2 REFUND + DISPUTE", "m", -AP["rev"]["value"],
        -AC["rev"]["value"],
        ("账单期 %d 笔有实际退款记录、%d 笔沿用估算；实际里争议/拒付 %s。逐笔见 ⑦"
         % (AP["rev"]["n_known"], AP["rev"]["n_unknown"], "{:,.{d}f}".format(AP["rev"]["dispute"], d=D["dec"])))
        if have else src + "：全部沿用估算（与上面的估算行相同）")
    row("adj_coupon", "减：卖家承担优惠券（实际，净）", "settlement_v2：交易额 + 费用 + 税 − 到账",
        "m", -AP["coupon"], -AC["coupon"],
        ("开票销售上的优惠券 − 退款退回的优惠券；账单期 %d 笔销售有优惠券，%d 笔销售在流水里找不到收款"
         % (AP["mp"]["n_coupon"], AP["mp"]["n_missing"])) if AP["mp"]["coupon_available"] else
        src + "：优惠券无法得到（要 settlement_v2），按 0 计——不是没有优惠券，是算不出来")
    row("adj_accrual", "减：应计费用（下期账单已出现）", "下期账单里本期销售的订单级费用", "m",
        -AP["accrual"]["total"], -AC["accrual"]["total"],
        "账单期 %d 行、自然月切 %d 行；逐行见 ⑧。还没出账单的费用无法得知" % (AP["accrual"]["n"], AC["accrual"]["n"]),
        vat="fee")
    row("contrib_adj", "经营贡献毛利（调整后，未扣商品成本）", "", "m", None, None,
        "＝ 上面的经营贡献毛利 + 本节各行", "tot",
        fx=["contrib", "adj_est", "adj_act", "adj_coupon", "adj_accrual"])
    row("err", "备查：冲销估算的误差（估算 − 实际）", "仅有实际数的那些笔", "m",
        AP["rev"]["err"] if have else None, AC["rev"]["err"] if have else None,
        ("正数 = 估算多冲了。账单期 %d 笔里 %d 笔分毫不差，误差绝对值合计 %s"
         % (AP["rev"]["n_known"], AP["rev"]["n_exact"], "{:,.{d}f}".format(AP["rev"]["abs_err"], d=D["dec"])))
        if have else src, "info", vat="blank")
    row("dch", "备查：MercadoPago 争议里只扣卖家的费用", "DISPUTE（交易额 0）", "m",
        -AP["mp"]["dispute_charge"] if have else None, -AC["mp"]["dispute_charge"] if have else None,
        ("就是账单的退货处理费（Cargo por devolución）：自然月切 %d 笔里 %d 笔同一收款号、同一金额。"
         "已在上面的平台费用里，不重复计" % (AC["mp"]["n_dispute_charge"], AC["mp"]["n_dispute_is_dev"]))
        if have else src, "info", vat="blank")
    r += 1
    r = _band(ws, r, "商品成本", 7)
    cp, cc = AP["cogs"], AC["cogs"]
    if cp is None and cc is None:
        _put(ws, r, 1, "未配置 SKU 成本表", bold=True)
        for c in range(2, 7):
            _put(ws, r, c, "")
        _put(ws, r, 7, "config.json 里没有 reports.sku_costs（见 core/sku_costs.py）；商品成本不估计。%s"
             % (D.get("cost_note") or ""), color=K.C_GREY, wrap_text=True)
        r += 1
    else:
        row("cogs", "减：商品采购成本", "SKU 成本表", "m", -cp["cogs"], -cc["cogs"],
            "成本件数 = 开票件数 − 全额冲销件数；%s" % cc["coverage"], vat="same")
        row("gm", "毛利（扣商品成本后）", "", "m", None, None, "", "tot", fx=["contrib_adj", "cogs"])
    r += 1
    r = _band(ws, r, "增值税备忘（%s %g%%，不进损益）" % (vt["name"], round(vt["rate"] * 100, 2)), 7)
    revrows = ["gross", "rev", "ship", "ship_rev", "adj_est", "adj_act", "adj_coupon"]
    feerows = ["fees", "adj_accrual"]
    for k, label, note in (
            ("out", "销项税（收入中所含 %s）" % vt["name"], "＝ 含税收入 − 不含税收入（开票、冲销、运费、实际退款、优惠券）"),
            ("inp", "进项税（平台费用中所含 %s，%s）" % (vt["name"], "可抵扣" if vt["credit"] else "不可抵扣"),
             vt["credit_why"])):
        rows = revrows if k == "out" else feerows
        sgn = "" if k == "out" else "-"
        _put(ws, r, 1, label, bold=True)
        _put(ws, r, 2, "", color=K.C_GREY)
        for col, net in (("C", "E"), ("D", "F")):
            f = "=%s(%s)" % (sgn, "+".join("(%s%d-%s%d)" % (col, at[x], net, at[x]) for x in rows))
            _put(ws, r, "CD".index(col) + 3, f, money, bold=True)
        _put(ws, r, 5, "")
        _put(ws, r, 6, "")
        _put(ws, r, 7, note, color=K.C_GREY, wrap_text=True)
        at[k] = r
        r += 1
    _put(ws, r, 1, "平台业务的应纳增值税（销项 − 进项）", bold=True, fillc=K.C_SUB)
    _put(ws, r, 2, "", fillc=K.C_SUB)
    for col in "CD":
        _put(ws, r, "CD".index(col) + 3, "=%s%d-%s%d" % (col, at["out"], col, at["inp"]), money, bold=True,
             fillc=K.C_SUB)
    _put(ws, r, 5, "", fillc=K.C_SUB)
    _put(ws, r, 6, "", fillc=K.C_SUB)
    _put(ws, r, 7, "只含平台业务本身；采购等其他进项要财务另加。%s" % (
        "墨西哥平台已代扣一部分 IVA（见下面代扣税），申报时抵减" if vt["withholding"] else "本国平台不代扣，全额自行申报"),
        color=K.C_GREY, wrap_text=True, fillc=K.C_SUB)
    r += 2
    r = _band(ws, r, "发票口径无法得到（现行订单口径的值，仅供参考，不计入上面任何合计）", 7)
    r = _head(ws, r, ["项目", "原文", "现行订单口径（参考）", "", "", "", "为什么账单里没有"])
    for lab, es, v, why in D["unavailable"]:
        _put(ws, r, 1, lab, color=K.C_GREY)
        _put(ws, r, 2, es, color=K.C_GREY)
        if v is None:
            _put(ws, r, 3, "—", color=K.C_GREY)
        else:
            _put(ws, r, 3, v, money, color=K.C_GREY, italic=True)
        for c in (4, 5, 6):
            _put(ws, r, c, "")
        _put(ws, r, 7, why, color=K.C_GREY, wrap_text=True)
        ws.row_dimensions[r].height = K.est_height(why, 62)
        r += 1
    r += 1
    r = _band(ws, r, "检查", 7)
    for lab, val, note, color in _withholding_checks(D):
        _put(ws, r, 1, lab, bold=True)
        _put(ws, r, 2, "")
        if isinstance(val, (int, float)):
            _put(ws, r, 3, val, money)
        else:
            _put(ws, r, 3, val, color=color, bold=True)
        for c in (4, 5, 6):
            _put(ws, r, c, "")
        _put(ws, r, 7, note, color=K.C_GREY, wrap_text=True)
        ws.row_dimensions[r].height = K.est_height(note, 62)
        r += 1
    r += 1
    r = _band(ws, r, "结算通道（Descontado de la operación）", 7)
    r = _head(ws, r, ["通道", "含义", "账单期", "自然月切", "", "", ""])
    for ch in ("Si", "No", "No aplica"):
        _put(ws, r, 1, "%s（%s）" % (CHANNEL_ZH[ch], ch))
        _put(ws, r, 2, CHANNEL_MEANING[ch], color=K.C_GREY)
        _put(ws, r, 3, -P["n"]["channel"][ch], money)
        _put(ws, r, 4, -C["n"]["channel"][ch], money)
        for c in (5, 6, 7):
            _put(ws, r, c, "")
        r += 1
    r += 1
    r = _band(ws, r, "开票销售的下单月份（看出有多少是别的月份的订单）", 7)
    r = _head(ws, r, ["下单月份", "", "账单期 笔数", "自然月切 笔数", "", "", ""])
    months = sorted(set(P["n"]["sale_months"]) | set(C["n"]["sale_months"]))
    for mo in months:
        _put(ws, r, 1, mo)
        _put(ws, r, 2, "")
        _put(ws, r, 3, P["n"]["sale_months"].get(mo, 0), K.INT)
        _put(ws, r, 4, C["n"]["sale_months"].get(mo, 0), K.INT)
        for c in (5, 6, 7):
            _put(ws, r, c, "")
        r += 1
    ws.freeze_panes = "C6"
    return ws


def _withholding_checks(D):
    """(label, value or status text, note, colour) rows of ③'s check block."""
    AC = D["acc"]["calendar"]
    if not D["vat"]["withholding"]:
        return [("代扣税凭证", "不适用", "本国平台不代扣税（MercadoPago 流水 TAXES_AMOUNT 为 0），不需要代扣凭证。",
                 K.C_GREEN)]
    out = [("代扣税凭证：CFDI de retenciones", "尚未下载",
            "代扣的 IVA/ISR 账单里没有（“发票口径无法得到”）。正式来源是 MercadoLibre/MercadoPago 每月出具的 CFDI de "
            "retenciones（Constancia de retenciones de IVA e ISR，在卖家中心的税务文件里）。本次没有下载，"
            "是以后的一个下载步骤；拿到之前不编任何数字。", K.C_RED)]
    if AC["mp"]["available"] and D["mp"]["source"] == "settlement":
        g = D["invs"]["calendar"]["n"]["gross_sales"]
        out.append(("参考：MercadoPago 流水里这些销售被扣的税", AC["mp"]["tax"],
                    "自然月切开票销售的收款上 TAXES_AMOUNT 之和（扣除退款退回的），占开票销售额 %s。IVA 与 ISR 混在一起，"
                    "不能代替代扣凭证；现行报表“Ventas 佣金及税金 − 账单佣金”是 %s。"
                    % ("{:.2%}".format(AC["mp"]["tax"] / g) if g else "-",
                       "{:,.{d}f}".format(D["S"]["m"]["tax_ret"], d=D["dec"])), "000000"))
    else:
        out.append(("参考：MercadoPago 流水里的代扣税", "无法得到", MP.source_text(D["mp"]) + "。", K.C_GREY))
    return out


# ════════════════════════════════════════════════════════════════════════
# ④ 与现行订单口径对比
# ════════════════════════════════════════════════════════════════════════

def compare_sheet(wb, D):
    ws = wb.create_sheet(SHEETS[3])
    money = D["money"]
    P, C = D["invs"]["period"], D["invs"]["calendar"]
    K.title_rows(ws, "与现行订单口径逐行对比 · %s · %s" % (D["store"], D["month"]),
                 "现行 = 月度报表 ② 合并损益表的数（同一个下载目录，用它自己的 derive() 算出，未做任何改动）。"
                 "差额 = 账单口径 − 现行。",
                 "负数 = 费用/减项。“无法得到”= 账单里没有这一项。")
    RF.set_widths(ws, [32, 34, 17, 17, 17, 16, 16, 26, 62, 62])
    r = _head(ws, 5, ["项目", "原文（Ventas ↔ 账单）", "现行订单口径\n下单日 %s" % D["month"],
                      "账单口径·账单期\n%s" % _span(P["view"]), "账单口径·自然月切\n%s" % _span(C["view"]),
                      "差额\n账单期 − 现行", "差额\n自然月切 − 现行", "原因类别", "说明（账单期）",
                      "说明（自然月切）"], height=44)
    at, block = {}, []
    for ln in D["lines"]:
        k, kind = ln["key"], ln["kind"]
        style = {"sub": K.C_SUB, "tot": K.C_TOTAL}.get(kind)
        bold = style is not None
        _put(ws, r, 1, ln["label"], bold=bold, fillc=style)
        _put(ws, r, 2, ln["es"], color=K.C_GREY, fillc=style)
        if kind == "sub":
            rng = [x for x in block]
            for col in "CDE":
                _put(ws, r, " CDE".index(col) + 2, "=SUM(%s)" % ",".join("%s%d" % (col, x) for x in rng)
                     if rng else 0, money, bold=True, fillc=style)
            block = []
        elif kind == "tot":
            subs = [at["_net_sales"], at["_order_fees"], at["_bill_fees"]]
            for col in "CDE":
                _put(ws, r, " CDE".index(col) + 2, "=SUM(%s)" % ",".join("%s%d" % (col, x) for x in subs),
                     money, bold=True, fillc=style)
        else:
            _value(ws, r, 3, ln["cur"], money, "n" if kind == "n" else "m")
            _value(ws, r, 4, ln["inv"].get("period"), money, "n" if kind == "n" else "m")
            _value(ws, r, 5, ln["inv"].get("calendar"), money, "n" if kind == "n" else "m")
            if kind == "m":
                block.append(r)
        fm = K.INT if kind == "n" else money
        _put(ws, r, 6, _diff_formula("C", "D", r), fm, bold=bold, fillc=style)
        _put(ws, r, 7, _diff_formula("C", "E", r), fm, bold=bold, fillc=style)
        cats = ln["cats"]
        if isinstance(cats, dict):
            cats = list(dict.fromkeys(cats.get("period", []) + cats.get("calendar", [])))
        _put(ws, r, 8, "、".join(CAT[c] for c in cats), wrap_text=True, fillc=style)
        npd, ncl = ln["notes"].get("period", ""), ln["notes"].get("calendar", "")
        _put(ws, r, 9, npd, color=K.C_GREY, wrap_text=True, fillc=style)
        _put(ws, r, 10, "同左" if ncl == npd and ncl else ncl, color=K.C_GREY, wrap_text=True,
             fillc=style)
        longest = max(len(npd or ""), len(ncl or ""))
        if longest > 60:
            ws.row_dimensions[r].height = K.est_height("x" * longest, 52)
        at[k] = r
        r += 1
    r += 1
    r = _band(ws, r, "核对", 10)
    _put(ws, r, 1, "经营贡献毛利（独立计算）")
    _put(ws, r, 2, "现行：report_build.build_context；账单：本实验", color=K.C_GREY)
    _put(ws, r, 3, D["S"]["contrib"], money)
    _put(ws, r, 4, P["n"]["contribution"], money)
    _put(ws, r, 5, C["n"]["contribution"], money)
    chk = r
    r += 1
    _put(ws, r, 1, "残差（应为 0）", bold=True, fillc=K.C_TOTAL)
    _put(ws, r, 2, "上表合计 − 独立计算", color=K.C_GREY, fillc=K.C_TOTAL)
    for c in "CDE":
        _put(ws, r, " CDE".index(c) + 2, "=%s%d-%s%d" % (c, at["_contrib"], c, chk), money,
             bold=True, fillc=K.C_TOTAL)
    r += 2

    r = _band(ws, r, "逐单对照：本月订单（Ventas，下单日）↔ 开票销售（账单）", 10)
    r = _note(ws, r, "Ventas 的多件包裹（“Paquete de N productos”母行 + N 行子行）在这里算一单；账单每件一笔，"
                     "按 Número de paquete 并回包裹。名单见 ⑥。", 10)
    r = _head(ws, r, ["项目", "", "", "账单期", "自然月切", "", "", "", "说明", ""])
    rs = dict((v, rec_summary(D["recs"][v])) for v in ("period", "calendar"))
    R = D["recs"]
    items = [
        ("本月订单（Ventas）", R["period"]["n_month_orders"], R["calendar"]["n_month_orders"], "n",
         "多件包裹算一单"),
        ("开票销售笔数", P["n"]["n_sales"], C["n"]["n_sales"], "n", "账单上一件商品一笔"),
        ("两边都有的订单", rs["period"]["matched_n"], rs["calendar"]["matched_n"], "n", ""),
        ("只在账单：上月订单本期计费", rs["period"]["extra_prev"], rs["calendar"]["extra_prev"], "n", ""),
        ("只在账单：下月订单本期计费", rs["period"]["extra_next"], rs["calendar"]["extra_next"], "n",
         "多为时区：账单的下单日按平台时间，Ventas 按本地时间"),
        ("只在账单：Ventas 里查无此单", rs["period"]["extra_absent"], rs["calendar"]["extra_absent"], "n",
         "多为下单后随即取消、账单上计费又冲销的订单；Ventas 不列这类订单"),
        ("只在 Ventas：在别的账单计费", rs["period"]["missing_elsewhere"],
         rs["calendar"]["missing_elsewhere"], "n", "智利：账单期以外（切期日之后）的订单在下一张账单"),
        ("只在 Ventas：取消，平台未计费", rs["period"]["missing_cancel"], rs["calendar"]["missing_cancel"],
         "n", ""),
        ("只在 Ventas：其他（未计费 / 金额为 0）", rs["period"]["missing_other"],
         rs["calendar"]["missing_other"], "n", "已下载的账单里都没有这单的佣金"),
        ("两边金额不一致的订单", rs["period"]["differ_n"], rs["calendar"]["differ_n"], "n", "见 ⑥"),
    ]
    for lab, a, b, kind, note in items:
        _put(ws, r, 1, lab)
        for c in (2, 3, 6, 7, 8, 10):
            _put(ws, r, c, "")
        _put(ws, r, 4, a, _fmt_of(kind, money))
        _put(ws, r, 5, b, _fmt_of(kind, money))
        _put(ws, r, 9, note, color=K.C_GREY, wrap_text=True)
        r += 1
    r += 1
    r = _head(ws, r, ["GMV 差额拆解", "", "", "账单期", "自然月切", "", "", "", "说明", ""])
    g0 = r
    parts = [
        ("开票销售额", P["n"]["gross_sales"], C["n"]["gross_sales"], "账单"),
        ("减：本月订单 GMV（Ventas，按订单汇总）", -R["period"]["month_gmv"], -R["calendar"]["month_gmv"],
         "与现行 ② 的商品销售额同一批 Ventas 行"),
        ("＝ 差额", "=D%d+D%d" % (g0, g0 + 1), "=E%d+E%d" % (g0, g0 + 1), ""),
        ("其中：只在账单的销售", rs["period"]["extra_total"], rs["calendar"]["extra_total"], "+"),
        ("其中：只在 Ventas 的订单", -rs["period"]["missing_gmv"], -rs["calendar"]["missing_gmv"], "−"),
        ("其中：两边都有、金额不同", rs["period"]["matched_diff"], rs["calendar"]["matched_diff"], ""),
        ("残差（应为 0）", "=D%d-SUM(D%d:D%d)" % (g0 + 2, g0 + 3, g0 + 5),
         "=E%d-SUM(E%d:E%d)" % (g0 + 2, g0 + 3, g0 + 5), "拆解完整时为 0"),
    ]
    for lab, a, b, note in parts:
        last = lab.startswith("残差") or lab.startswith("＝")
        fc = K.C_TOTAL if lab.startswith("残差") else (K.C_CALC if lab.startswith("＝") else None)
        _put(ws, r, 1, lab, bold=last, fillc=fc)
        for c in (2, 3, 6, 7, 8, 10):
            _put(ws, r, c, "", fillc=fc)
        _put(ws, r, 4, a, money, bold=last, fillc=fc)
        _put(ws, r, 5, b, money, bold=last, fillc=fc)
        _put(ws, r, 9, note, color=K.C_GREY, fillc=fc)
        r += 1
    ws.freeze_panes = "C6"
    return ws


# ════════════════════════════════════════════════════════════════════════
# ⑤ 费用明细
# ════════════════════════════════════════════════════════════════════════

def fees_sheet(wb, D):
    ws = wb.create_sheet(SHEETS[4])
    money = D["money"]
    P, C = D["invs"]["period"], D["invs"]["calendar"]
    S = D["S"]
    K.title_rows(ws, "平台费用明细（账单口径）· %s · %s" % (D["store"], D["month"]),
                 "金额为正 = 平台向卖家收费（账单原样）；冲销为负。账单期 %s，自然月切 %s。"
                 % (_span(P["view"]), _span(C["view"])),
                 "现行（参考）：订单级 = 本月订单关联到的账单行净额（现行 ② 的佣金行只取毛额、配送费取 Ventas，"
                 "见 ④）；账单级 = 现行 ② 的同名行。")
    RF.set_widths(ws, [30, 46, 12, 10, 11, 16, 11, 16, 16, 16, 16])
    r = _head(ws, 5, ["科目", "Detalle 原文", "类别", "级别", "账单期\n计费笔数", "账单期\n计费金额",
                      "账单期\n冲销笔数", "账单期\n冲销金额", "账单期\n净额", "自然月切\n净额",
                      "现行订单口径\n（参考）"], height=40)
    L = S["billing_linked"]
    fc = dict((x["es"], x) for x in C["fees"])
    first = r
    for f in P["fees"]:
        es = f["es"]
        if f["order_level"]:
            cur = float(L[L["parent_detalle"].astype(str) == es]["amount"].sum()) if len(L) else 0.0
        else:
            cur = S["m"]["bill_fee_net"].get(es, 0.0)
        vals = [f["zh"], es, BUCKET_ZH.get(f["bucket"], f["bucket"]),
                "订单级" if f["order_level"] else "账单级", f["n_gross"], f["gross"], f["n_rev"],
                f["rev"], "=F%d+H%d" % (r, r), fc[es]["net"], cur]
        kinds = "tttnmnmmmmm"
        quiet = not (f["n_gross"] or f["n_rev"] or abs(fc[es]["net"]) > 0.004 or abs(cur) > 0.004)
        for j, (v, kd) in enumerate(zip(vals, kinds), 1):
            _put(ws, r, j, v, _fmt_of(kd, money), color=K.C_GREY if quiet else "000000",
                 fillc=K.C_CALC if j == 9 else None)
        r += 1
    unk = {}
    for v, inv in (("p", P), ("c", C)):
        for u in inv["unknown"]:
            x = unk.setdefault(u["es"], {"p": 0.0, "c": 0.0, "n": 0, "lvl": u["order_level"],
                                         "ch": u["channels"]})
            x[v] += u["amount"]
            if v == "p":
                x["n"] += u["n"]
    for es, u in sorted(unk.items()):
        vals = ["⚠ 未登记科目", es, "未登记", "订单级" if u["lvl"] else "账单级", u["n"], u["p"], "", "",
                "=F%d" % r, u["c"], ""]
        for j, (v, kd) in enumerate(zip(vals, "tttnmnmmmmm"), 1):
            _put(ws, r, j, v, _fmt_of(kd, money), fillc=K.C_BAD)
        r += 1
    last = r - 1
    _put(ws, r, 1, "合计", bold=True, fillc=K.C_SUB)
    for j in range(2, 12):
        col = get_column_letter(j)
        v = "=SUM(%s%d:%s%d)" % (col, first, col, last) if j >= 5 else ""
        _put(ws, r, j, v, K.INT if j in (5, 7) else money, bold=True, fillc=K.C_SUB)
    tot = r
    r += 1
    _put(ws, r, 1, "核对：账单行金额合计", bold=False)
    for j in range(2, 12):
        _put(ws, r, j, "")
    _put(ws, r, 9, P["n"]["fee_total"], money)
    _put(ws, r, 10, C["n"]["fee_total"], money)
    chk = r
    r += 1
    _put(ws, r, 1, "残差（应为 0）", bold=True, fillc=K.C_TOTAL)
    for j in range(2, 12):
        _put(ws, r, j, "", fillc=K.C_TOTAL)
    _put(ws, r, 9, "=I%d-I%d" % (tot, chk), money, bold=True, fillc=K.C_TOTAL)
    _put(ws, r, 10, "=J%d-J%d" % (tot, chk), money, bold=True, fillc=K.C_TOTAL)
    r += 2
    if not unk:
        r = _note(ws, r, "本期没有未登记科目：账单里每一个 Detalle 都在 FEE_CATALOG 里。", 11, K.C_GREEN)
    else:
        r = _note(ws, r, "⚠ 有 %d 个 Detalle 不在 FEE_CATALOG 里（红底行）：金额已计入，科目归属请人工确认，"
                         "确认后在 platforms/mercadolibre/report/catalog.py 登记。" % len(unk), 11, K.C_RED, True)
    ws.freeze_panes = "C6"
    return ws


# ════════════════════════════════════════════════════════════════════════
# ⑥ 销售与冲销明细
# ════════════════════════════════════════════════════════════════════════

def _views_to_list(D):
    """MX/PE: the two readings are the same rows -> list them once."""
    if D["cycle_day"] == 1:
        return [("period", "账单期（= 自然月）")]
    return [("period", "账单期 %s" % _span(D["invs"]["period"]["view"])),
            ("calendar", "自然月切 %s" % _span(D["invs"]["calendar"]["view"]))]


def detail_sheet(wb, D):
    ws = wb.create_sheet(SHEETS[5])
    money = D["money"]
    K.title_rows(ws, "销售与冲销明细 · %s · %s" % (D["store"], D["month"]),
                 "每一笔都能回到账单原始行（销售号 = Número de venta，计费号 = Número del cargo）。",
                 "墨西哥/秘鲁的账单按自然月，账单期和自然月切是同一批行，只列一次。"
                 if D["cycle_day"] == 1 else "智利：账单期和自然月切分别列出。")
    RF.set_widths(ws, [20, 20, 14, 14, 14, 13, 13, 11, 14, 15, 11, 44, 14, 30, 30, 15, 40])
    orders, parent_of = D["orders"], D["parent_of"]
    r = 5
    r = _band(ws, r, "A. 冲销的关联方法（找到这笔销售的原佣金，才能算冲掉了多少）", 17)
    r = _head(ws, r, ["方法", "", "", "账单期", "自然月切"] + [""] * 12)
    for mth in LINK_ORDER:
        _put(ws, r, 1, mth, wrap_text=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        _put(ws, r, 4, D["invs"]["period"]["n"]["rev_methods"].get(mth, 0), K.INT)
        _put(ws, r, 5, D["invs"]["calendar"]["n"]["rev_methods"].get(mth, 0), K.INT)
        ws.row_dimensions[r].height = 30
        r += 1
    r += 1
    rev_cols = [("冲销计费日", "d"), ("凭证", "t"), ("销售号", "t"), ("包裹号", "t"), ("下单日", "d"),
                ("冲销佣金", "m"), ("原佣金", "m"), ("冲销比例", "p"), ("售价", "m"),
                ("估算冲销销售额", "m"), ("估算件数", "r"), ("关联方法", "t"), ("原佣金计费日", "d"),
                ("Ventas 状态（参考）", "t"), ("商品", "t"), ("Ventas 退款净额（参考）", "m"),
                ("文件", "t")]
    for v, name in _views_to_list(D):
        R = D["invs"][v]["reversals"]
        r = _band(ws, r, "B. 冲销明细 · %s（%d 笔）" % (name, len(R)), 17)
        rows = []
        for _, x in R.sort_values("cdate").iterrows():
            o = orders.get(order_key(orders, parent_of, x["kpack"], x["ksale"])) or {}
            rows.append([x["cdate"], x["doc"], x["ksale"], x["kpack"], x["sdate"], x["rev_commission"],
                         x["orig_commission"], x["ratio"], x["sale_total"], x["est_sales"],
                         x["est_units"], x["method"].split("（")[0], x["parent_date"],
                         o.get("status", "Ventas 查无"), x["title"], o.get("refund"), x["file"]])
        r = _table(ws, r, rev_cols, rows, money) + 1
    rec_cols = [("销售号", "t"), ("Ventas 订单号", "t"), ("下单日（账单）", "d"), ("计费日", "d"),
                ("Ventas 下单日", "d"), ("件数", "n"), ("售价", "m"), ("原因", "t"), ("商品", "t")]
    miss_cols = [("Ventas 订单号", "t"), ("下单日", "d"), ("状态", "t"), ("件数", "n"), ("GMV", "m"),
                 ("退款净额", "m"), ("原因", "t"), ("SKU", "t"), ("商品", "t")]
    diff_cols = [("Ventas 订单号", "t"), ("下单日", "d"), ("状态", "t"), ("账单售价合计", "m"),
                 ("Ventas GMV", "m"), ("差额", "m"), ("账单件数", "n"), ("Ventas 件数", "n")]
    for v, name in _views_to_list(D):
        rec = D["recs"][v]
        E, M, T = rec["extra"], rec["missing"], rec["differ"]
        r = _band(ws, r, "C. 本期开票、但不是本月订单 · %s（%d 笔）" % (name, len(E)), 17)
        r = _table(ws, r, rec_cols, [[x["ksale"], x["okey"], x["sdate"], x["cdate"], x["odate"],
                                      x["units"], x["total"], x["why"], x["title"]]
                                     for _, x in E.sort_values(["code", "sdate"]).iterrows()],
                   money) + 1
        r = _band(ws, r, "D. 本月订单、但不在本期开票 · %s（%d 单）" % (name, len(M)), 17)
        r = _table(ws, r, miss_cols, [[x["okey"], x["odate"], x["status"], x["units"], x["gmv"],
                                       x["refund"], x["why"], x["sku"], x["title"]]
                                      for _, x in M.sort_values(["code", "odate"]).iterrows()],
                   money) + 1
        r = _band(ws, r, "E. 两边都有、金额不一致的订单 · %s（%d 单）" % (name, len(T)), 17)
        r = _table(ws, r, diff_cols, [[x["okey"], x["odate"], x["status"], x["inv_total"], x["gmv"],
                                       x["diff"], x["inv_units"], x["units"]]
                                      for _, x in T.iterrows()], money) + 1
    sale_cols = [("销售号", "t"), ("包裹号", "t"), ("Ventas 订单号", "t"), ("下单日", "d"),
                 ("计费日", "d"), ("件数", "n"), ("单价", "m"), ("售价", "m"), ("佣金", "m"),
                 ("Estado del cargo", "t"), ("通道", "t"), ("凭证", "t"), ("商品", "t")]
    for v, name in _views_to_list(D):
        sales = D["recs"][v]["sales"]
        r = _band(ws, r, "F. 开票销售明细 · %s（%d 笔）" % (name, len(sales)), 17)
        r = _table(ws, r, sale_cols, [[x["ksale"], x["kpack"], x["okey"], x["sdate"], x["cdate"],
                                       x["units"], x["price"], x["total"], x["commission"],
                                       x["state"], x["channel"], x["doc"], x["title"]]
                                      for _, x in sales.sort_values(["cdate", "ksale"]).iterrows()],
                   money) + 1
    ws.freeze_panes = "A5"
    return ws


# ════════════════════════════════════════════════════════════════════════
# ⑦ 退款优惠券与现金（MercadoPago）
# ════════════════════════════════════════════════════════════════════════

def mp_sheet(wb, D):
    ws = wb.create_sheet(SHEETS[6])
    money = D["money"]
    mp = D["mp"]
    src = MP.source_text(mp)
    K.title_rows(ws, "退款、优惠券与现金（MercadoPago）· %s · %s" % (D["store"], D["month"]),
                 "账单里没有的钱，从 MercadoPago 补：实际退给买家多少、卖家承担多少优惠券、平台代扣多少税、"
                 "最后放给卖家多少钱。",
                 "数据来源：%s。按账单每一行的 Pago（收款号）挂到每笔销售上。" % src)
    RF.set_widths(ws, [44, 18, 18, 18, 18, 18, 18, 16, 60, 22, 40])
    nc = 11
    r = 5
    r = _band(ws, r, "A. 本次下载的 MercadoPago 文件", nc)
    for k, _pat in MP.GLOBS:
        p = mp["files"].get(k)
        _put(ws, r, 1, MP.FILE_ZH[k])
        _put(ws, r, 2, os.path.basename(p) if p else "没有", color="000000" if p else K.C_GREY)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=nc)
        r += 1
    for t in mp["notes"] or ["实际采用：%s。" % src]:
        r = _note(ws, r, t, nc, K.C_RED if not mp["source"] else K.C_NAVY)
    r += 1
    views = _views_to_list(D)
    r = _band(ws, r, "B. 取消与退款：估算（账单冲销比例）vs 实际（MercadoPago）", nc)
    r = _head(ws, r, ["项目"] + [name for _v, name in views] + [""] * (nc - 1 - len(views)))
    items = [("冲销的销售（笔）", "n", "n"), ("估算冲销销售额", "est", "m"),
             ("有实际数的笔数", "n_known", "n"), ("　其估算", "est_known", "m"),
             ("　其实际（退款 + 争议/拒付）", "actual_known", "m"), ("　　其中争议/拒付", "dispute", "m"),
             ("　误差：估算 − 实际", "err", "m"), ("　误差绝对值合计", "abs_err", "m"),
             ("　分毫不差的笔数", "n_exact", "n"), ("沿用估算的笔数", "n_unknown", "n"),
             ("　其估算金额", "est_unknown", "m"), ("③ 用的“实际”（实际 + 沿用的估算）", "value", "m")]
    for lab, k, kind in items:
        _put(ws, r, 1, lab, bold=k in ("err", "value"))
        for j, (v, _n) in enumerate(views):
            _put(ws, r, 2 + j, D["acc"][v]["rev"][k], _fmt_of(kind, money), bold=k in ("err", "value"))
        for c in range(2 + len(views), nc + 1):
            _put(ws, r, c, "")
        r += 1
    _put(ws, r, 1, "按状态（笔）", bold=True)
    r += 1
    for s in accounting.STATUS_ORDER:
        _put(ws, r, 1, s, wrap_text=True)
        for j, (v, _n) in enumerate(views):
            _put(ws, r, 2 + j, D["acc"][v]["rev"]["status"].get(s, 0), K.INT)
        ws.row_dimensions[r].height = 28
        r += 1
    r += 1
    cols = [("销售号", "t"), ("包裹号", "t"), ("下单日（账单）", "d"), ("冲销计费日", "d"), ("售价", "m"),
            ("冲销比例", "p"), ("估算冲销销售额", "m"), ("实际（退款+争议）", "m"), ("其中争议/拒付", "m"),
            ("误差：估算 − 实际", "m"), ("状态", "t"), ("收款号", "t"), ("商品", "t")]
    for v, name in views:
        A = D["acc"][v]["A"]
        r = _band(ws, r, "B%s. 逐笔 · %s（%d 笔，误差大的在前）" % ("12"[views.index((v, name))], name, len(A)), nc)
        rows = []
        if len(A):
            X = A.assign(_e=A["err"].abs().fillna(-1)).sort_values(["_e", "cdate"], ascending=[False, True])
            for _, x in X.iterrows():
                rows.append([x["ksale"], x["kpack"], x["sdate"], x["cdate"], x["sale_total"], x["ratio"],
                             x["est"], x["actual"], x["dispute"], x["err"], x["status"], x["pays"], x["title"]])
        r = _table(ws, r, cols, rows, money) + 1
    r = _band(ws, r, "C. 卖家承担的优惠券", nc)
    r = _head(ws, r, ["项目"] + [name for _v, name in views] + [""] * (nc - 1 - len(views)))
    for lab, k, kind in (("开票销售的笔数", "n_sales", "n"), ("有优惠券的销售（笔）", "n_coupon", "n"),
                         ("开票销售上的优惠券", "coupon_gross", "m"), ("减：退款退回的优惠券", "coupon_back", "m"),
                         ("＝ 优惠券（净），③ 用这个数", "coupon_net", "m"),
                         ("流水里找不到收款的销售（笔）", "n_missing", "n"),
                         ("优惠券算不出来的销售（笔）", "n_unknown", "n")):
        _put(ws, r, 1, lab, bold=k == "coupon_net")
        for j, (v, _n) in enumerate(views):
            m = D["acc"][v]["mp"]
            val = (m[k] if k != "coupon_back" else -m[k]) if m["coupon_available"] else None
            _value(ws, r, 2 + j, val, money, kind, bold=k == "coupon_net")
        for c in range(2 + len(views), nc + 1):
            _put(ws, r, c, "")
        r += 1
    cc = D.get("coupon_check")
    if cc:
        r = _note(ws, r, "核对优惠券取数方法：本月订单（Ventas）里 %d 单能在结算流水里找到收款，其中 %d 单 Ventas 倒推的优惠券"
                         "（月度报表的做法）= 流水的“交易额 + 费用 + 税 − 到账”，分毫不差；两边合计 %s 对 %s。"
                  % (cc["n"], cc["eq"], "{:,.{d}f}".format(cc["ventas"], d=D["dec"]),
                     "{:,.{d}f}".format(cc["mp"], d=D["dec"])), nc, K.C_NAVY)
    else:
        r = _note(ws, r, "%s：无法逐单核对优惠券的取数方法（需要 settlement_v2）。现行 Ventas 倒推的优惠券是 %s。"
                  % (src, "{:,.{d}f}".format(D["S"]["m"]["coupon"], d=D["dec"])), nc)
    r += 1
    r = _band(ws, r, "D. 争议、拒付与售后", nc)
    for v, name in views:
        m = D["acc"][v]["mp"]
        t = D["acc"][v]["rev"]
        if not m["available"]:
            r = _note(ws, r, "%s：%s。" % (name, src), nc)
            continue
        r = _note(ws, r, "%s：买家赢了的争议/拒付退回 %s（已并入实际退款）；争议里只扣卖家的费用 %d 笔 %s，其中 %d 笔就是账单的"
                         "退货处理费（同一收款号、同一金额），已在平台费用里，不重复计；Poscobro 里本期销售还在进行中的索赔 %d 笔。"
                  % (name, "{:,.{d}f}".format(t["dispute"], d=D["dec"]), m["n_dispute_charge"],
                     "{:,.{d}f}".format(m["dispute_charge"], d=D["dec"]), m["n_dispute_is_dev"], m["open_claims"]),
                  nc)
    r += 1
    r = _band(ws, r, "E. 平台代扣税", nc)
    for lab, val, note, color in _withholding_checks(D):
        _put(ws, r, 1, lab, bold=True)
        if isinstance(val, (int, float)):
            _put(ws, r, 2, val, money)
        else:
            _put(ws, r, 2, val, color=color, bold=True)
        _put(ws, r, 3, note, color=K.C_GREY, wrap_text=True)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=nc)
        ws.row_dimensions[r].height = K.est_height(note, 120)
        r += 1
    r += 1
    r = cash_block(ws, r, D, nc)
    ws.freeze_panes = "B5"
    return ws


def cash_block(ws, r, D, nc):
    money = D["money"]
    C = D["cash"]
    r = _band(ws, r, "F. 现金核对：调整后的经营贡献毛利 → MercadoPago 本月释放给卖家的钱（自然月 %s）"
              % _span(D["invs"]["calendar"]["view"]), nc)
    r = _head(ws, r, ["项目", "来源", "金额", "", "", "", "", "", "说明", "", ""])
    at = {}
    for x in C["lines"]:
        if x["kind"] == "note":
            r = _note(ws, r, x["label"], nc, K.C_RED)
            continue
        kind = x["kind"]
        fc = {"sub": K.C_CALC, "tot": K.C_SUB, "chk": K.C_TOTAL}.get(kind)
        grey = kind in ("info", "part")
        _put(ws, r, 1, ("　" if kind == "part" else "") + x["label"], bold=kind in ("sub", "tot", "chk"),
             fillc=fc, color=K.C_GREY if kind == "info" else "000000", wrap_text=True)
        _put(ws, r, 2, x["src"], color=K.C_GREY, fillc=fc)
        if x.get("sum_of"):
            if all(s.lstrip("-") in at for s in x["sum_of"]):
                f = "=" + "".join(("-" if s.startswith("-") else "+") + "C%d" % at[s.lstrip("-")]
                                  for s in x["sum_of"]).lstrip("+")
                _put(ws, r, 3, f, money, bold=True, fillc=fc)
            else:
                _put(ws, r, 3, UNAVAILABLE, color=K.C_RED, fillc=fc)
        elif x["value"] is None:
            _put(ws, r, 3, UNAVAILABLE, color=K.C_RED, fillc=fc)
        else:
            _put(ws, r, 3, x["value"], money, fillc=fc, italic=grey and kind == "info")
        for c in range(4, 9):
            _put(ws, r, c, "", fillc=fc)
        _put(ws, r, 9, x.get("note", ""), color=K.C_GREY, wrap_text=True, fillc=fc)
        ws.merge_cells(start_row=r, start_column=9, end_row=r, end_column=nc)
        if len(x.get("note") or "") > 60 or len(x["label"]) > 40:
            ws.row_dimensions[r].height = K.est_height(x.get("note") or x["label"], 90)
        if x.get("value") is not None or x.get("sum_of"):
            at[x["key"]] = r
        r += 1
    r += 1
    r = _band(ws, r, "G. 资金去向（本月；只列出来，不强行勾稽）", nc)
    pg = C["pagos"]
    if pg:
        for x in pg:
            _put(ws, r, 1, "支付平台账单（%s，%s）" % (x["kind"], x["file"]))
            _put(ws, r, 2, _d(x["day"]), color=K.C_GREY)
            _put(ws, r, 3, -x["amount"], money)
            r += 1
    else:
        r = _note(ws, r, "本次下载的 Pagos_Facturas 里，本月没有账单付款（上月账单的付款报表没有下载；本月账单在下月付款）。", nc)
    w = C.get("withdraw")
    if w:
        _put(ws, r, 1, "提现到银行（%d 笔，已批准）" % w["n"])
        _put(ws, r, 2, "Retiros", color=K.C_GREY)
        _put(ws, r, 3, -w["amount"], money)
        r += 1
    elif D["mp"]["source"]:
        r = _note(ws, r, "本次下载没有提现报表（withdraw），或本月没有提现。", nc)
    if C.get("outgoing") is not None:
        _put(ws, r, 1, "从 MercadoPago 余额对外付款（结算流水里金额为负的 SETTLEMENT，含支付平台账单、买东西）")
        _put(ws, r, 2, "settlement_v2", color=K.C_GREY)
        _put(ws, r, 3, C["outgoing"], money)
        r += 1
    acct = D["mp"].get("account")
    if acct:
        r = _note(ws, r, "余额与流水报表（account_statement，期间见文件）：期初 %s，收入 %s，支出 %s，期末 %s。"
                  % tuple("{:,.2f}".format(acct.get(k) or 0) for k in ("INITIAL_BALANCE", "CREDITS", "DEBITS",
                                                                     "FINAL_BALANCE")), nc)
    return r + 1


# ════════════════════════════════════════════════════════════════════════
# ⑧ 三种日期与应计
# ════════════════════════════════════════════════════════════════════════

def dates_sheet(wb, D):
    ws = wb.create_sheet(SHEETS[7])
    money = D["money"]
    T = D["dates"]
    K.title_rows(ws, "三种日期与应计 · %s · %s" % (D["store"], D["month"]),
                 "每笔开票销售的下单日（Ventas）、计费日（账单）、交付日（Ventas“Fecha entregado”），"
                 "以及按三种规则分月的收入；下期账单里已经出现的本期费用（应计）。",
                 "金额是开票售价（含税，不减冲销）。交付日在 Ventas 里没有年份，取离下单日最近的那一年。")
    RF.set_widths(ws, [20, 20, 13, 13, 13, 15, 8, 30, 12, 12, 12, 44])
    nc = 12
    r = 5
    r = _band(ws, r, "A. 应计：本期销售的费用，在下一张（已下载的）账单里才计费", nc)
    for v, name in _views_to_list(D):
        acc = D["acc"][v]["accrual"]
        r = _note(ws, r, "%s：%d 行，%d 笔销售，合计 %s（③ 的“应计费用”行）。"
                  % (name, acc["n"], acc["n_sales"], "{:,.{d}f}".format(acc["total"], d=D["dec"])), nc, K.C_NAVY)
        rows = [[x["zh"], x["es"], x["file"], "未关账" if x["open"] else "已关账", x["n"], x["amount"],
                 x["from"], x["to"]] for x in acc["lines"]]
        if rows:
            r = _table(ws, r, [("科目", "t"), ("Detalle", "t"), ("账单文件", "t"), ("账单状态", "t"), ("行数", "n"),
                               ("金额", "m"), ("计费日从", "d"), ("到", "d")], rows, money)
    r = _note(ws, r, "还没出账单的费用（下一张账单还没下载、或平台还没计费）无法得知，本实验不估计。"
                     "下一张账单关账前，这一行还会变（见 ⑨ 稳定性）。", nc)
    r += 1
    r = _band(ws, r, "B. 收入按三种日期分月（全部已下载账单里的开票销售，%d 笔）" % T["n_all"], nc)
    r = _head(ws, r, ["月份", "", "按下单日\n金额", "笔数", "按计费日（账单）\n金额", "笔数", "", "按交付日\n金额",
                      "笔数", "", "", "说明"], height=36)
    for s in T["split"]:
        hi = s["month"] == D["month"]
        fc = K.C_WARN if hi else None
        _put(ws, r, 1, s["month"] + ("（本月）" if hi else ""), bold=hi, fillc=fc)
        _put(ws, r, 2, "", fillc=fc)
        _put(ws, r, 3, s["order"], money, bold=hi, fillc=fc)
        _put(ws, r, 4, s["n_order"], K.INT, fillc=fc)
        _put(ws, r, 5, s["billing"], money, bold=hi, fillc=fc)
        _put(ws, r, 6, s["n_billing"], K.INT, fillc=fc)
        _put(ws, r, 7, "", fillc=fc)
        _put(ws, r, 8, s["delivery"], money, bold=hi, fillc=fc)
        _put(ws, r, 9, s["n_delivery"], K.INT, fillc=fc)
        for c in (10, 11):
            _put(ws, r, c, "", fillc=fc)
        _put(ws, r, 12, "交付日：Ventas 里还没交付（在途、取消、退货）的没有日期" if s["month"] == accounting.NO_DATE
             else "", color=K.C_GREY, fillc=fc)
        r += 1
    im = T["in_month"]
    r = _note(ws, r, "本月收入：按下单日 %s，按计费日 %s，按交付日 %s。只有已下载的账单里的销售才在这张表里："
                     "下单日较早的月份、计费日较晚的月份都不完整（例如上月的销售在上月账单里，没有下载）。"
              % tuple("{:,.{d}f}".format(im[k], d=D["dec"]) for k in ("order", "billing", "delivery")), nc, K.C_NAVY)
    if not T["has_delivery"]:
        r = _note(ws, r, "⚠ 这份 Ventas 没有“Fecha entregado”列，交付日无法得到。", nc, K.C_RED)
    r += 1
    X = T["table"]
    r = _band(ws, r, "C. 逐笔：三个日期中任一落在本月的开票销售（%d 笔）" % len(X), nc)
    cols = [("销售号", "t"), ("Ventas 订单号", "t"), ("下单日", "d"), ("计费日", "d"), ("交付日", "d"), ("售价", "m"),
            ("件数", "n"), ("Ventas 状态", "t"), ("按下单日", "t"), ("按计费日", "t"), ("按交付日", "t"), ("商品", "t")]
    rows = [[x["ksale"], x["okey"], x["order"], x["billing"], x["delivery"], x["total"], x["units"], x["status"],
             x["m_order"], x["m_billing"], x["m_delivery"], x["title"]] for _, x in X.iterrows()]
    r = _table(ws, r, cols, rows, money)
    ws.freeze_panes = "A5"
    return ws


# ════════════════════════════════════════════════════════════════════════
# 稳定性（⑨ of the experiment report, ② of the stability workbook)
# ════════════════════════════════════════════════════════════════════════

def stability_table(ws, r, A, money):
    """One block per basis: metrics × snapshots, change last − first, %, why."""
    snaps, kns = A["snaps"], A["kns"]
    n = len(snaps)
    labels = [snapshot_label(D["folder"]) for D in snaps]
    ncols = n + 4
    RF.set_widths(ws, [44] + [17] * n + [16, 10, 80])
    for b, title in (("order", "A. 现行订单口径（Ventas，按下单日）"),
                     ("period", "B. 账单口径 · 账单期"),
                     ("calendar", "C. 账单口径 · 自然月切")):
        r = _band(ws, r, title, ncols)
        r = _head(ws, r, ["指标"] + labels + ["变化\n（末次 − 首次）", "变化 %", "说明"], height=36)
        c0, c1 = get_column_letter(2), get_column_letter(n + 1)
        for k, lab, kind in A["rows"]:
            vals = [kn[b].get(k) for kn in kns]
            if all(v is None for v in vals):
                _put(ws, r, 1, lab, color=K.C_GREY)
                for j in range(n):
                    _put(ws, r, 2 + j, UNAVAILABLE if b != "order" else "—", color=K.C_GREY)
                for j in range(n + 2, ncols + 1):
                    _put(ws, r, j, "")
                r += 1
                continue
            moved = any(v is not None and vals[0] is not None and abs(v - vals[0]) > 0.004
                        for v in vals)
            fc = K.C_WARN if moved else None
            _put(ws, r, 1, lab, fillc=fc)
            for j, v in enumerate(vals):
                _value(ws, r, 2 + j, v, money, kind, fillc=fc)
            fmtv = K.INT if kind == "n" else money
            _put(ws, r, n + 2, _diff_formula(c0, c1, r), fmtv, fillc=fc)
            _put(ws, r, n + 3, '=IF(AND(ISNUMBER(%s%d),ISNUMBER(%s%d)),IF(%s%d=0,"",(%s%d-%s%d)/ABS(%s%d)),"")'
                 % (c0, r, c1, r, c0, r, c1, r, c0, r, c0, r), K.PCT, fillc=fc)
            _put(ws, r, n + 4, "有变化" if moved else "", color=K.C_AMBER, fillc=fc)
            r += 1
        for p in A["pairs"]:
            head = "%s → %s：" % (snapshot_label(p["a"]["folder"]), snapshot_label(p["b"]["folder"]))
            texts = p["why"]["order"] if b == "order" else [p["why"][b]]
            r = _note(ws, r, head + " ".join(texts), ncols, K.C_NAVY)
        r += 1
    return r


def stability_sheet(wb, name, A, money, store, month, reason=None):
    ws = wb.create_sheet(name)
    K.title_rows(ws, "口径稳定性（同一月多次下载）· %s · %s" % (store, month),
                 "每次下载用同一套代码分别算两种口径；黄底 = 这一行在几次下载之间变了。",
                 "订单口径变 = Ventas（实时导出）里的订单、状态、退款变了，或下月账单里新计了本月订单的费用；"
                 "账单口径变 = 账单行本身变了（关账的账单不应该变）。")
    if not A:
        ws.column_dimensions["A"].width = 120
        ws["A5"] = reason or "只有一次下载，无法比较。"
        ws["A5"].font = F(11, False, K.C_GREY)
        ws["A6"] = ("要比较请把同一月的几次下载放在同一店铺目录下（文件夹名以 _m<YYYYMM> 结尾），"
                    "或运行：python -m experiments.ml_invoice_basis.stability --store <店> --month <YYYY-MM> "
                    "--folders <目录1> <目录2> …")
        ws["A6"].font = F(10, False, K.C_GREY)
        return ws
    r = 5
    for D in A["snaps"]:
        r = _note(ws, r, "%s：%s　文件：%s" % (snapshot_label(D["folder"]), D["folder"], "、".join(
            "%s%s" % (d["name"], "（未关账）" if d["open"] else "") for d in D["docs"])),
            len(A["snaps"]) + 4)
    r += 1
    stability_table(ws, r, A, money)
    ws.freeze_panes = "B5"
    return ws


# ════════════════════════════════════════════════════════════════════════
# workbooks
# ════════════════════════════════════════════════════════════════════════

def write_report_workbook(D, A, out, reason=None, pct=None, absv=None):
    """pct / absv: the materiality thresholds of ② (% of net revenue, amount)."""
    pct = accounting.DEFAULT_PCT if pct is None else pct
    absv = accounting.DEFAULT_ABS if absv is None else absv
    wb = Workbook()
    wb.remove(wb.active)
    readme_sheet(wb, D, A, pct, absv)
    summary_sheet(wb, D, pct, absv)
    pl_sheet(wb, D)
    compare_sheet(wb, D)
    fees_sheet(wb, D)
    detail_sheet(wb, D)
    mp_sheet(wb, D)
    dates_sheet(wb, D)
    stability_sheet(wb, SHEETS[8], A, D["money"], D["store"], D["month"], reason)
    return _save(wb, out)


def write_stability_workbook(A, out):
    D0 = A["snaps"][-1]
    money = D0["money"]
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(STAB_SHEETS[0])
    K.title_rows(ws, "美客多 口径稳定性实验 · %s · %s" % (A["store"], A["month"]),
                 "实验报表，不是月度财务报表。问题：Ventas（订单表）一直在变，那换成账单口径会不会稳定？")
    RF.set_widths(ws, [2, 26, 120])
    rows = [("S", "一、比较了哪几次下载", "")]
    for D in A["snaps"]:
        rows.append(("", snapshot_label(D["folder"]), "%s；账单：%s；Ventas：%s%s" % (
            D["folder"], "、".join("%s%s" % (d["name"], "（未关账）" if d["open"] else "")
                                   for d in D["docs"]),
            os.path.basename(D["S"]["files"].get("ventas") or ""),
            "（⚠ 平台提示导出可能缺数据）" if D["S"].get("ventas_warning") else "")))
    rows += [("S", "二、怎么比", ""),
             ("", "两种口径", "每次下载都用同一套代码算：现行订单口径 = 月度报表自己的 derive()（一行未改）；"
                         "账单口径 = 本实验（账单期、自然月切两种读法）。"),
             ("", "逐单 / 逐行", "订单口径按订单号比对两次 Ventas（新增、消失、状态/金额变化），见 %s；"
                            "账单口径按（凭证, Número del cargo）比对两次账单行（新增、消失、金额/Estado/通道/发票号变化），"
                            "见 %s。" % (STAB_SHEETS[2], STAB_SHEETS[3])),
             ("S", "三、结论", "")]
    for p in A["pairs"]:
        tag = "%s → %s" % (snapshot_label(p["a"]["folder"])[:10], snapshot_label(p["b"]["folder"])[:10])
        rows.append(("", "订单口径 " + tag, " ".join(p["why"]["order"])))
        rows.append(("", "账单期 " + tag, p["why"]["period"]))
        rows.append(("", "自然月切 " + tag, p["why"]["calendar"]))
    RF.readme_rows(ws, 4, rows)

    stability_sheet(wb, STAB_SHEETS[1], A, money, A["store"], A["month"])

    ws = wb.create_sheet(STAB_SHEETS[2])
    K.title_rows(ws, "订单变化明细（Ventas，本月下单的订单）", "每对相邻的两次下载之间：新增、消失、状态或金额变化的订单。")
    RF.set_widths(ws, [20, 12, 26, 34, 34, 14, 14, 14, 14, 14, 14, 9, 9, 14, 14])
    r = 4
    cols = [("订单号", "t"), ("下单日", "d"), ("变化", "t"), ("状态（前）", "t"), ("状态（后）", "t"),
            ("GMV 前", "m"), ("GMV 后", "m"), ("退款 前", "m"), ("退款 后", "m"), ("Total 前", "m"),
            ("Total 后", "m"), ("件数 前", "n"), ("件数 后", "n"), ("佣金及税金 前", "m"),
            ("佣金及税金 后", "m")]
    for p in A["pairs"]:
        r = _band(ws, r, "%s → %s（%d 单）" % (snapshot_label(p["a"]["folder"]),
                                              snapshot_label(p["b"]["folder"]), len(p["ventas"])), 15)
        rows = []
        for x in p["ventas"]:
            b, a = x["before"] or {}, x["after"] or {}
            rows.append([x["okey"], x["odate"], x["kind"], b.get("status"), a.get("status"),
                         b.get("gmv"), a.get("gmv"), b.get("refund"), a.get("refund"), b.get("total"),
                         a.get("total"), b.get("units"), a.get("units"), b.get("com_col"),
                         a.get("com_col")])
        if not rows:
            r = _note(ws, r, "两次 Ventas 里本月的订单完全相同（订单、状态、金额都没变）。", 15, K.C_GREEN)
        else:
            r = _table(ws, r, cols, rows, money)
        r += 1
    ws.freeze_panes = "A4"

    ws = wb.create_sheet(STAB_SHEETS[3])
    K.title_rows(ws, "账单行变化明细", "按（凭证, Número del cargo）比对相邻两次下载的账单行。关账的账单应当一行不变。")
    RF.set_widths(ws, [12, 8, 16, 12, 40, 16, 14, 14, 26, 26, 12, 12, 16, 16, 40])
    r = 4
    cols = [("读法", "t"), ("凭证", "t"), ("计费号", "t"), ("计费日", "d"), ("Detalle", "t"),
            ("变化", "t"), ("金额 前", "m"), ("金额 后", "m"), ("Estado 前", "t"), ("Estado 后", "t"),
            ("通道 前", "t"), ("通道 后", "t"), ("发票号 前", "t"), ("发票号 后", "t"), ("文件", "t")]
    for p in A["pairs"]:
        tag = "%s → %s" % (snapshot_label(p["a"]["folder"]), snapshot_label(p["b"]["folder"]))
        for v, name in (("period", "账单期"), ("calendar", "自然月切")):
            diff, na, nb = p["invoice"][v]
            r = _band(ws, r, "%s · %s：%d → %d 行，%d 行不同" % (tag, name, na, nb, len(diff)), 15)
            if not diff:
                r = _note(ws, r, p["why"][v], 15, K.C_GREEN) + 1
                continue
            rows = []
            for x in diff:
                b, a = x["before"] or {}, x["after"] or {}
                z = a or b
                rows.append([name, x["doc"], x["charge"], z.get("cdate"), z.get("det"), x["kind"],
                             b.get("amount"), a.get("amount"), b.get("state"), a.get("state"),
                             b.get("channel"), a.get("channel"), b.get("invoice"), a.get("invoice"),
                             z.get("file")])
            r = _table(ws, r, cols, rows, money) + 1
        lr = p["linked"]
        r = _note(ws, r, "另：现行订单口径关联到本月订单的账单行（跨全部已下载账单）新增 %d 行、消失 %d 行。"
                         % (lr["new_n"], lr["gone_n"]), 15) + 1
    ws.freeze_panes = "A4"
    return _save(wb, out)


def _save(wb, out):
    d = os.path.dirname(os.path.abspath(out))
    if d and not os.path.isdir(d):
        os.makedirs(d)
    wb.save(out)
    return out


def finish(out, no_recalc=False):
    """Recalculate with LibreOffice when available and scan for formula errors.
    Returns the exit code: 0, or 1 when a formula error is in the file."""
    if no_recalc:
        return 0
    ok, msg = K.recalc(out)
    print("[重算] %s" % msg)
    if ok:
        nf, ne, nn, bad = K.scan_formula_errors(out)
        print("[公式] 共 %d 条，错误 %d 条，未求值 %d 条" % (nf, ne, nn))
        for b in bad:
            print("       %s" % b)
        if ne:
            print("[警告] 文件里有公式错误。")
            return 1
    return 0
