# -*- coding: utf-8 -*-
"""List every SKU of a month so finance can fill in unit costs (core/sku_costs).

    python -m tools.sku_cost_template --out sku_costs.xlsx
    python -m tools.sku_cost_template --out sku_costs.xlsx --month 2026-09 --market FA-CL --market AZ-MX

Reads the newest download folder of the month of every configured store of each
market (the same folder the monthly report uses) through the platform's own
report_build — files only, it never opens a browser or a store. Writes one row per
(platform, country, SKU) with the product name, the stores, units and sales of the
month, and empty 单位成本 / 生效日期 cells to fill.

If --out already exists it is merged, never replaced: every existing row stays as
it is (costs, dates, notes, extra columns); the reference columns of rows that
match are refreshed; only SKUs not in the file yet are added at the bottom. A SKU
already covered by a broader row (blank 平台 or 国家) is not added again.

A platform takes part when its report_build has `sku_list(store, folder, month)`
and its run_reports has `latest_folder(store, month, base)`; markets of other
platforms are listed as skipped.

Exit codes: 0 written, 2 nothing to list / bad arguments, 4 configuration error.
"""
import argparse
import collections
import datetime
import importlib
import os
import sys

from core import console  # noqa: F401  中文输出编码保护
from core import markets
from core import sku_costs as SC
from core import store_config
from core.store_config import StoreConfigError, StoreNotAllowed

# reference columns: written for finance, ignored when the table is read
REF_COLS = collections.OrderedDict([
    ("name", "商品名称（参考）"), ("stores", "店铺（参考）"), ("month", "统计月份（参考）"),
    ("units", "本月件数（参考）"), ("sales", "本月销售额（参考，不含 IVA）")])


def last_month(today=None):
    d = (today or datetime.date.today()).replace(day=1) - datetime.timedelta(days=1)
    return "%04d-%02d" % (d.year, d.month)


def platform_modules(platform_code):
    """(report_build, run_reports) of a platform, or None when it cannot list SKUs."""
    pkg = markets.platform(platform_code).package
    try:
        rb = importlib.import_module(pkg + ".report_build")
        rr = importlib.import_module(pkg + ".run_reports")
    except ImportError:
        return None
    if not hasattr(rb, "sku_list") or not hasattr(rr, "latest_folder"):
        return None
    return rb, rr


def collect(market_ids, month, base=None, log=print):
    """{(platform, country, sku key): {"sku", "name", "stores", "units", "sales", "currency"}}"""
    found = collections.OrderedDict()
    for mk in market_ids:
        p, c = markets.split(mk)
        mods = platform_modules(p)
        if mods is None:
            log("  %s：这个平台还不能列出 SKU（报表未接入成本表），跳过" % mk)
            continue
        rb, rr = mods
        names = store_config.stores_for_market(mk)
        if not names:
            log("  %s：config.json 里没有这个市场的店铺" % mk)
        for store in names:
            folder = rr.latest_folder(store, month, base) if base else rr.latest_folder(store, month)
            if not folder:
                log("  %-26s 没有 %s 的下载目录，跳过" % (store, month))
                continue
            pp, cc, cur, items = rb.sku_list(store, folder, month)
            log("  %-26s %s  %d 个 SKU" % (store, folder, len(items)))
            for it in items:
                k = (pp, cc, SC.sku_key(it.sku))
                x = found.setdefault(k, {"sku": str(it.sku).strip(), "name": it.name or "", "stores": [],
                                         "units": 0, "sales": 0.0, "currency": cur})
                if store not in x["stores"]:
                    x["stores"].append(store)
                x["units"] += int(it.units or 0)
                x["sales"] += float(it.sales or 0.0)
                if not x["name"] and it.name:
                    x["name"] = it.name
    return found


def read_existing(path):
    """(header labels, rows as lists, column index of each canonical field)."""
    rows = SC._read_cells(path)
    for i, (_line, cells) in enumerate(rows[:20]):
        idx = {}
        for j, h in enumerate(cells):
            k = SC._ALIASES.get(SC._norm_header(h))
            if k and k not in idx:
                idx[k] = j
            ref = [r for r, label in REF_COLS.items() if SC._norm_header(label) == SC._norm_header(h)]
            if ref and ref[0] not in idx:
                idx[ref[0]] = j
        if "sku" in idx:
            header = [SC._text(h) for h in cells]
            body = [list(c) for _l, c in rows[i + 1:] if any(SC._text(v) for v in c)]
            return header, body, idx
    raise ValueError("%s 里找不到表头行（要有 SKU 列）" % path)


def merge(found, existing, month):
    """The rows to write: existing rows first (unchanged but for the reference
    columns), then the new SKUs. Returns (header, rows, added, refreshed)."""
    cols = list(SC.COLUMN_LABELS) + list(REF_COLS)
    header = [SC.COLUMN_LABELS[k] for k in SC.COLUMN_LABELS] + list(REF_COLS.values())
    old_header, old_rows, idx = existing if existing else ([], [], {})
    known = set(idx.values())
    extra = [j for j in range(len(old_header)) if j not in known and old_header[j]]
    header += [old_header[j] for j in extra]
    out, have, refreshed = [], set(), 0
    for cells in old_rows:
        get = lambda k: cells[idx[k]] if k in idx and idx[k] < len(cells) else None
        row = dict((k, get(k)) for k in cols)
        p, c, key = SC._text(row["platform"]).upper(), SC._text(row["country"]).upper(), SC.sku_key(row["sku"])
        have.add((p, c, key))
        x = found.get((p, c, key))
        if x:
            row.update({"name": row.get("name") or x["name"], "stores": "、".join(x["stores"]),
                        "month": month, "units": x["units"], "sales": round(x["sales"], 2)})
            refreshed += 1
        out.append([row[k] for k in cols] + [cells[j] if j < len(cells) else None for j in extra])
    new = []
    for (p, c, key), x in found.items():
        if any((pp, cc, key) in have for pp in (p, "") for cc in (c, "")):
            continue
        new.append((p, c, -x["sales"], [p, c, x["sku"], None, x["currency"], None, None, x["name"],
                                        "、".join(x["stores"]), month, x["units"], round(x["sales"], 2)]
                    + [None] * len(extra)))
    new.sort(key=lambda t: t[:3])
    out += [t[3] for t in new]
    return header, out, len(new), refreshed


def write(path, header, rows):
    from openpyxl import Workbook
    from core import report_kit as K
    from openpyxl.styles import Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "商品成本"
    ws.append(header)
    for j, _h in enumerate(header, 1):
        c = ws.cell(row=1, column=j)
        c.font = K.F(10, True, "FFFFFF")
        c.fill = K.fill(K.C_NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = [7, 7, 22, 12, 7, 12, 24, 44, 26, 10, 10, 16]
    from openpyxl.utils import get_column_letter
    for j in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(j)].width = widths[j - 1] if j <= len(widths) else 16
    ws.row_dimensions[1].height = 32
    for i, r in enumerate(rows, 2):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = K.F(10, False, K.C_GREY) if j > len(SC.COLUMN_LABELS) else K.F(10)
        if r[3] is None or SC._text(r[3]) == "":
            ws.cell(row=i, column=4).fill = K.fill(K.C_INPUT)
        ws.cell(row=i, column=4).number_format = "#,##0.00"
        ws.cell(row=i, column=6).number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=12).number_format = "#,##0.00"
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(header)), max(1, len(rows) + 1))
    info = wb.create_sheet("说明")
    info.column_dimensions["A"].width = 14
    info.column_dimensions["B"].width = 100
    for line in [
            ("这张表", "每个 SKU 的单位成本，所有平台的月报共用。报表只读第一个工作表；灰色的“参考”列不读。"),
            ("平台", "平台代码（%s）；空着 = 所有平台" % "、".join(
                "%s %s" % (k, v.name_zh) for k, v in markets.PLATFORMS.items())),
            ("国家", "国家代码（MX、CL …）；空着 = 所有国家"),
            ("SKU", "卖家 SKU，和平台文件里的一样（不分大小写，前后空格不算）"),
            ("单位成本", "到岸成本：采购价＋运费＋关税，每件，不含 IVA；必须 ≥ 0"),
            ("币种", "必须是这个平台报表的币种（MXN、CLP …）；报表不做汇率换算"),
            ("生效日期", "可空。成本变了就新加一行写新的生效日期，旧行保留：每个月用月底那天有效的那一行"),
            ("备注", "随便写，报表不读"),
            ("匹配顺序", "先找 平台＋国家 都写了的行，再找只写平台的，再找只写国家的，最后找都空着的"),
            ("更新", "再次运行 python -m tools.sku_cost_template --out <这个文件> 只会在末尾追加新 SKU，"
                     "已有的行和成本一行不删")]:
        info.append(list(line))
        info.cell(row=info.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    tmp = path + ".tmp.xlsx"
    wb.save(tmp)
    os.replace(tmp, path)


def main(argv=None):
    ap = argparse.ArgumentParser(description="列出本月所有 SKU，生成/更新商品成本表（只读下载目录，不开浏览器）")
    ap.add_argument("--out", required=True, help="成本表 .xlsx；已存在就合并（不删任何一行）")
    ap.add_argument("--month", default=None, help="YYYY-MM，默认上个月")
    ap.add_argument("--market", action="append", default=None,
                    help="市场代码，如 FA-CL、AZ-MX，可重复；默认 config.json 里所有能列出 SKU 的市场")
    ap.add_argument("--downloads", default=None, help="下载目录的根（默认 downloads/）")
    a = ap.parse_args(argv)
    month = a.month or last_month()
    if not a.out.lower().endswith(".xlsx"):
        print("--out 要是 .xlsx 文件")
        return 2
    try:
        mks = [markets.normalize(m) for m in a.market] if a.market else \
            [m for m in store_config.markets_configured() if platform_modules(markets.split(m)[0])]
        print("月份 %s；市场：%s" % (month, "、".join(mks) or "（没有）"))
        found = collect(mks, month, a.downloads)
    except (StoreConfigError, StoreNotAllowed, markets.UnknownMarket) as e:
        print("配置错误：%s" % e)
        return 4
    existing = None
    if os.path.exists(a.out):
        try:
            existing = read_existing(a.out)
        except Exception as e:
            print("读不出已有的成本表 %s（%s）：没有改动它" % (a.out, e))
            return 2
    if not found and not existing:
        print("没有找到任何 SKU，没有写文件")
        return 2
    header, rows, added, refreshed = merge(found, existing, month)
    write(a.out, header, rows)
    print("已写入 %s：共 %d 行，新增 %d 个 SKU，更新参考列 %d 行。请在黄色格填单位成本。"
          % (a.out, len(rows), added, refreshed))
    return 0


if __name__ == "__main__":
    sys.exit(main())
