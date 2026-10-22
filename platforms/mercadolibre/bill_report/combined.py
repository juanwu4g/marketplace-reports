# -*- coding: utf-8 -*-
"""The bill-basis combined workbook of one country (全店汇总_账单基准).

The same ten sheets as a store workbook (docs/REPORT_FAMILY.md: "Store workbooks
and the combined one have the same sheets; the combined one has one column per
store plus a total"):

  ① ② ③  store columns + 合计, built by core/report_family from each store's S
  ④      one row per Detalle, store columns + 合计
  ⑤      the reversal and status summaries with store columns; the per-sale
         lists of every store under a 店铺 column
  ⑥ ⑦ ⑨  every store's rows under a 店铺 column (the row builders of workbook.py)
  ⑧      the weeks of the month summed over the stores
  ⑩      every store's checks, missing files and open items

Each S is built once by build.prepare() and shared with the store workbooks,
so a figure here is the sum of the same figures in the store workbooks.
"""
import re

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from core import report_family as RF
from core import report_kit as K
from core.report_kit import F, fill
from platforms.mercadolibre.bill_report import catalog as C
from platforms.mercadolibre.bill_report import mapping as M
from platforms.mercadolibre.bill_report import status as ST
from platforms.mercadolibre.bill_report import workbook as WB

SHEETS = K.SHEET_NAMES
BOLD = WB.BOLD
GREY9 = WB.GREY9
STORE_COL = RF.Col("店铺", 20)
NUMERALS = "一二三四五六七八九十"
REV_KEY = {ST.CANCEL: "rev_cancel", ST.RETURN: "rev_return", ST.PARTIAL: "rev_partial",
           ST.UNPROVEN: "rev_unproven"}


def build(Ss, ctx, out_path):
    """Ss: the stores' S (build.prepare), one country. ctx: month, period_cn,
    checks (every store's), sku_costs, built, order_combined (the order-based
    summary's file name, for ①)."""
    wb = Workbook()
    wb.remove(wb.active)
    fam = RF.Family(Ss, money=K.money_format(Ss[0]["prof"].decimals))
    ctx["money"] = fam.money
    sheet_readme(wb, fam, Ss, ctx)
    ctx["pl_at"] = WB.sheet_pl(wb, fam, Ss[0], ctx)
    sheet_bridge(wb, fam, Ss, ctx)
    sheet_fees(wb, fam, Ss, ctx)
    sheet_returns(wb, fam, Ss, ctx)
    sheet_qc(wb, fam, Ss, ctx)
    sheet_sku(wb, fam, Ss, ctx)
    sheet_trend(wb, fam, Ss, ctx)
    sheet_inventory(wb, fam, Ss, ctx)
    sheet_gaps(wb, fam, Ss, ctx)
    wb.save(out_path)
    return out_path


# ------------------------------------------------------------------ helpers

def pl_col(fam, i=None):
    """② column letter of store i, or of 合计 (i None; the store's own column
    when there is only one store)."""
    if i is None:
        i = fam.n if fam.has_total else 0
    return get_column_letter(3 + i)


def plref(fam, ctx, key, i=None):
    return "%s!%s%d" % (WB.q(SHEETS[1]), pl_col(fam, i), ctx["pl_at"][key])


def with_store(cols):
    """The 店铺 column in front of a store table's columns; formula templates
    ("=D{r}-E{r}") move one column to the right with it."""
    out = [STORE_COL]
    for c in cols:
        if c.formula:
            c = c._replace(formula=re.sub(r"\b([A-Z]{1,2})(?=\{r\})",
                                          lambda m: get_column_letter(WB._col_idx(m.group(1)) + 1),
                                          c.formula))
        out.append(c)
    return out


def total_row(ws, r, first, cols, label="合计", int_cols=(), skip=()):
    """SUM row under rows first..r-1 for the given column numbers."""
    WB.write_cells(ws, r, 1, [label], font=BOLD, fill_color=K.C_TOTAL)
    for j in cols:
        if j in skip:
            continue
        L_ = get_column_letter(j)
        WB.write_cells(ws, r, j, ["=SUM(%s%d:%s%d)" % (L_, first, L_, r - 1) if r > first else 0],
                       fmt=K.INT if j in int_cols else None, font=BOLD, fill_color=K.C_TOTAL)
    return r


def money_fmt(ws, r, cols, fam, int_cols=()):
    for j in cols:
        if j not in int_cols:
            ws.cell(row=r, column=j).number_format = fam.money


def renumber(sections, start):
    """Give sections taken from a store's ① the combined workbook's numbers."""
    out = []
    for i, (head, lines) in enumerate(sections):
        out.append(("%s、%s" % (NUMERALS[start + i], head.split("、", 1)[-1]), lines))
    return out


# ------------------------------------------------------------------ ①

def sheet_readme(wb, fam, Ss, ctx):
    p = Ss[0]["prof"]
    dec = p.decimals
    period, points = [], []
    for S in Ss:
        secs = WB.readme_sections(S, ctx)
        period.append((S["store"], secs[1][1][0][1]))
        pl, m = S["m"]["pl"], S["m"]
        bad = [c for c in RF.failed_errors(ctx["checks"]) if c.get("store") == S["store"]]
        points.append((S["store"], "净商品销售额 %s；平台费用 %s；经营贡献 %s（%s）；SKU 映射 %s；%s。文件夹：%s" % (
            WB.fmtm(pl["net_gmv"], dec), WB.fmtm(-pl["fees_total"], dec), WB.fmtm(pl["contrib"], dec),
            WB.pct(pl["contrib"] / pl["net_gmv"]) if pl["net_gmv"] else "—", WB.pct(m["mapped_share"], 2),
            ("%d 项 error 级校验未通过（见⑩）" % len(bad)) if bad else "error 级校验全部通过", S["folder"])))
    policy = WB.readme_sections(Ss[0], ctx)
    sections = [
        ("一、这份报表是什么", [
            ("基准", "以 MercadoLibre 的账单（Reporte_Facturacion ＋ Reporte_Notas_Credito）为唯一金额基准的财务报表，"
                     "%s全部 %d 家店铺的汇总：每家店一列，最后一列合计。每家店的数字与它自己的单店报表"
                     "（<店铺>_%s_账单基准.xlsx）完全相同。" % (p.country_cn, len(Ss), ctx["month"])),
            ("两份报表的关系", "每月与订单基准汇总表（%s）一起生成，两者各自独立计算、互不引用。会计入账以本报表"
                              "（账单基准）为准；订单基准报表作经营分析参考。" % ctx.get("order_combined", "")),
            ("国家 / 币种", "%s ／ %s（%s）" % (p.country_cn, p.currency, p.currency_cn)),
            ("店铺", "、".join(S["store"] for S in Ss)),
        ]),
        ("二、期间：%s，按计费日（各店账单周期不同，逐店说明）" % ctx["period_cn"], period),
        ("三、各店要点", points),
    ] + renumber([policy[i] for i in WB.POLICY_SECTIONS], 3) + [
        ("七、各页说明", [
            (SHEETS[1], "损益：每家店一列 ＋ 合计；行位置与单店报表相同。下方费用率与备查项。"),
            (SHEETS[2], "勾稽：每家店的账本、销售、费用归属、SKU 映射、状态、账单付款、贷记单；逐个账单文件的明细见单店报表③。"),
            (SHEETS[3], "账单里出现的每个科目一行，每家店一列 ＋ 合计；下方按②行与②逐店勾稽。"),
            (SHEETS[4], "冲销与状态汇总（每家店一列），然后各店的逐单清单（店铺列）。"),
            (SHEETS[5], "各店退货逐单（店铺列）。"),
            (SHEETS[6], "表一各店的 SKU 一行（同一 SKU 在不同店铺各占一行）；表二各店每笔销售 → SKU 的映射。"),
            (SHEETS[7], "全部店铺按周合计；各店的前 10 SKU 按周见单店报表⑧。"),
            (SHEETS[8], "各店 Full 库存快照与本月净销量（店铺列）。"),
            (SHEETS[9], "各店全部校验、缺的文件、未匹配的销售与待决事项。"),
        ]),
    ]
    bad = RF.failed_errors(ctx["checks"])
    parts = [x for x in (("本期有 %d 项 error 级校验未通过（见⑩）：相关店铺的数字在修好之前不能信" % len(bad)) if bad else "",
                         WB.provisional_text(Ss)) if x]
    warning = "；".join(parts) + "。" if parts else None
    fam.readme_sheet(wb, SHEETS[0], "账单基准财务报表 · %s · %s全店汇总" % (ctx["period_cn"], p.country_cn),
                     "账单基准（会计入账口径）　生成于 %s　%d 家店铺　基准：MercadoLibre 账单 ＋ 贷记单（计费日）" % (
                         ctx["built"], len(Ss)), sections, warning=warning)


# ------------------------------------------------------------------ ③

def sheet_bridge(wb, fam, Ss, ctx):
    p = Ss[0]["prof"]
    Ln = RF.Line
    per = fam.per_store
    revs_in = lambda X: sum(e["amount"] for e in X["revs"] if e["in_period"])
    direct = lambda X: sum(v for (l, h), v in X["attributed"].items() if h == "direct")
    alloc = lambda X: sum(v for (l, h), v in X["attributed"].items() if h == "alloc")
    A = [Ln("账本合计（本月计费日的全部账单行）", lambda X: X["m"]["ledger"], "逐个账单文件见单店报表③桥 A", "subtotal"),
         Ln("② 平台费用合计（取反）", lambda X: -X["m"]["pl"]["fees_total"], ""),
         Ln("残差（应为 0）", lambda X: round(X["m"]["ledger"] + X["m"]["pl"]["fees_total"], 6), "", "residual")]
    B = [Ln("计费销售额（本月计费日的 Cargo por venta 行）", lambda X: X["m"]["pl"]["gmv_billed"],
            per(lambda X: "%d 笔" % X["m"]["n_cohort"]))]
    B += [Ln("减：%s" % ST.CLASS_CN[c], (lambda X, c=c: X["m"]["pl"]["rev_" + c]), "") for c in ST.REVERSAL_CLASSES]
    B += [Ln("净商品销售额（②）", lambda X: X["m"]["pl"]["net_gmv"], "", "subtotal"),
          Ln("⑦ 表一 各 SKU 净商品销售额合计", lambda X: X["m"]["sku_sum"]["net_gmv"], ""),
          Ln("残差（应为 0）：② − ⑦", lambda X: round(X["m"]["pl"]["net_gmv"] - X["m"]["sku_sum"]["net_gmv"], 6),
             "", "residual"),
          Ln("残差（应为 0）：② 冲销 − ⑤ 逐单冲销合计", lambda X: round(
              sum(X["m"]["pl"]["rev_" + c] for c in ST.REVERSAL_CLASSES) - revs_in(X), 6), "", "residual")]
    Cc = [Ln("直接归 SKU 的费用合计", direct, "", "subtotal"),
          Ln("按销售额分摊的费用合计", alloc, "", "subtotal"),
          Ln("② 平台费用合计", lambda X: X["m"]["pl"]["fees_total"], ""),
          Ln("残差（应为 0）：直接 ＋ 分摊 − ②", lambda X: round(direct(X) + alloc(X) - X["m"]["pl"]["fees_total"], 6),
             "", "residual"),
          Ln("残差（应为 0）：⑦ 直接列 ＋ 分摊列 − ②", lambda X: round(
              X["m"]["sku_sum"]["direct"] + X["m"]["sku_sum"]["alloc"] - X["m"]["pl"]["fees_total"], 6), "",
             "residual")]
    D = [Ln(M.TIER_CN[t], (lambda X, t=t: X["m"]["coverage"][t][1]),
            "可信度：%s；笔数 %s" % (M.CONFIDENCE[t], per(lambda X, t=t: str(X["m"]["coverage"][t][0]))))
         for t in M.TIER_ORDER]
    D += [Ln("合计", lambda X: sum(v[1] for v in X["m"]["coverage"].values()), "", "subtotal"),
          Ln("② 计费销售额", lambda X: X["m"]["pl"]["gmv_billed"],
             "覆盖率 %s（目标 ≥ 99%%）" % per(lambda X: WB.pct(X["m"]["mapped_share"], 2))),
          Ln("残差（应为 0）", lambda X: round(sum(v[1] for v in X["m"]["coverage"].values())
                                          - X["m"]["pl"]["gmv_billed"], 6), "", "residual")]
    E = [Ln(ST.CLASS_CN[c], (lambda X, c=c: X["m"]["cohort_status"][c][1]),
            "笔数 %s" % per(lambda X, c=c: str(X["m"]["cohort_status"][c][0]))) for c in ST.CLASS_CN]
    E += [Ln("合计", lambda X: sum(v[1] for v in X["m"]["cohort_status"].values()), "", "subtotal"),
          Ln("残差（应为 0）：− ② 计费销售额", lambda X: round(sum(v[1] for v in X["m"]["cohort_status"].values())
                                                       - X["m"]["pl"]["gmv_billed"], 6), "", "residual")]
    docs = lambda key: (lambda X: sum(d[key] for d in X["docs"]))
    Fb = [Ln("各张账单：月结费用（Descontado = No）", docs("invoiced"), "整张账单，非仅本月计费日；逐张见单店报表③桥 F"),
          Ln("　减：其后被冲销的部分", docs("annulled"), ""),
          Ln("　＝ 应付", docs("expected"), "", "subtotal"),
          Ln("　已付（付款报表逐费用号对上）", docs("paid"), ""),
          Ln("　待付余额（账单未到期或未付清，不是差错）",
             lambda X: round(sum(d["expected"] - d["paid"] for d in X["docs"]), 6), "", "pending")]
    G = [Ln("贷记单文件合计（全部，取正）", lambda X: X["m"]["nc_file_total"], ""),
         Ln("其中与付款报表对上的", lambda X: sum(a for _, a, _b, _h in X["nc_match"]["matched"]), ""),
         Ln("　付款报表上的对应金额", lambda X: sum(b for _, _a, b, _h in X["nc_match"]["matched"]), ""),
         Ln("　残差（应为 0）", lambda X: round(sum(a - b for _, a, b, _h in X["nc_match"]["matched"]), 6), "",
            "residual"),
         Ln("只在贷记单文件里的", lambda X: sum(a for _, a in X["nc_match"]["nc_only"]), "", "pending"),
         Ln("只在付款报表里的", lambda X: sum(a for _, a in X["nc_match"]["pagos_only"]), "", "pending")]
    st = lambda key: (lambda X: X["storage_stats"][key])
    H = [Ln("本期 Full 仓储日租（账单）",
            lambda X: X["storage_stats"]["exact"] + X["storage_stats"]["nearest"] + X["storage_stats"]["none"], ""),
         Ln("　按当天的 SKU 仓储明细归属", st("exact"), "账单 d 日 ＝ 明细 d−1 日"),
         Ln("　按最近一天的 SKU 结构归属", st("nearest"), "明细未覆盖的日子（多为月初）"),
         Ln("　无明细，按销售额分摊", st("none"), "")]
    sections = [
        ("桥 A　账本 → ② 平台费用", A),
        ("桥 B　计费销售 → 净商品销售额（②、⑤、⑦ 三处一致）", B),
        ("桥 C　平台费用的归属：直接归 SKU ＋ 按销售额分摊 ＝ ②（按科目的拆分见单店报表③）", Cc),
        ("桥 D　SKU 映射覆盖（按计费销售额）", D),
        ("桥 E　本月计费销售的状态（⑤ 二）", E),
        ("桥 F　账单 ↔ 账单付款报表（Reporte_Pagos_Facturas）", Fb),
        ("桥 G　贷记单文件 ↔ 付款报表中的贷记单", G),
    ]
    if p.has_full:
        sections.append(("桥 H　Full 仓储日租按 SKU 明细归属", H))
    fam.bridge_sheet(wb, SHEETS[2], "平账勾稽表（%s · 账单基准 · 全店汇总）" % ctx["period_cn"],
                     "每座桥最后是“残差（应为 0）”；琥珀色 = 等待中的单据（账单未付、贷记单未开），不是错误。单位：%s"
                     % p.currency, sections)


# ------------------------------------------------------------------ ④

def sheet_fees(wb, fam, Ss, ctx):
    p = Ss[0]["prof"]
    ws = wb.create_sheet(SHEETS[3])
    K.title_rows(ws, "平台费用明细（%s · 账单基准 · 全店汇总）" % ctx["period_cn"],
                 "账单里出现的每个科目一行（含贷记单），每家店一列；金额为账单原符号（费用为正、冲销为负），只计本期计费日。"
                 "②的费用行 ＝ 本表按“归入②行”相加后取反。单位：%s" % p.currency)
    first_cols = [("科目（账单原文 Detalle）", 52), ("中文", 28), ("归入②行", 22), ("归 SKU 方式", 30)]
    c0 = len(first_cols) + 1
    nval = fam.value_columns()
    RF.header_row(ws, 4, [h for h, _ in first_cols] + fam.names + (["合计"] if fam.has_total else []),
                  [w for _, w in first_cols] + [16] * nval)
    seen = []
    for S in Ss:
        seen += [d for d in dict.fromkeys(list(S["P"]["detalle"])) if d not in seen]
    order = [d for d in C.CATALOG if d in seen] + [d for d in seen if d not in C.CATALOG]
    r = 5
    first = r
    for d in order:
        e, known = C.entry(d)
        WB.write_cells(ws, r, 1, [d, e.cn if known else "⚠ 未登记科目", C.FEE_CN[e.line], C.ATTRIB_CN[e.attrib]],
                       font=WB.TEXT9, fill_color=None if known else K.C_BAD)
        fam.values_row(ws, r, c0, [float(S["P"].loc[S["P"]["detalle"] == d, "amount"].sum()) for S in Ss])
        r += 1
    cols = range(c0, c0 + nval)
    tot = total_row(ws, r, first, cols)
    money_fmt(ws, tot, cols, fam)
    r += 1
    WB.write_cells(ws, r, 1, ["残差（应为 0）：本期金额 ＋ ②平台费用合计"], font=BOLD, fill_color=K.C_TOTAL)
    for i in range(nval):
        L_ = get_column_letter(c0 + i)
        WB.write_cells(ws, r, c0 + i, ["=ROUND(%s%d+%s,6)" % (L_, tot, plref(fam, ctx, "fees_total", i))],
                       fmt=fam.money, font=BOLD, fill_color=K.C_TOTAL)
    r += 2
    r = WB.section_head(ws, r, "按②行汇总（与②逐店勾稽：每格 ＝ 本期账单金额 ＋ ② 对应值，应为 0）", c0 + nval - 1)
    K.header_cells(ws, r, ["②行", "", "", ""] + fam.names + (["合计"] if fam.has_total else []))
    r += 1
    for key, lab, _o, _n in C.PL_FEE_LINES:
        ws.cell(row=r, column=1, value=lab).border = K.BOX
        for i, S in enumerate(Ss):
            amt = float(S["P"].loc[S["P"]["line"] == key, "amount"].sum())
            fam.money_cell(ws.cell(row=r, column=c0 + i), "=ROUND(%.6f+%s,6)" % (amt, plref(fam, ctx, key, i)))
        if fam.has_total:
            fam.money_cell(ws.cell(row=r, column=c0 + fam.n), "=SUM(%s%d:%s%d)" % (
                get_column_letter(c0), r, get_column_letter(c0 + fam.n - 1), r))
        r += 1
    ws.freeze_panes = ws.cell(row=5, column=c0)


# ------------------------------------------------------------------ ⑤

def sheet_returns(wb, fam, Ss, ctx):
    p = Ss[0]["prof"]
    cols = with_store(WB.REV_COLS)
    ws = wb.create_sheet(SHEETS[4])
    K.title_rows(ws, "退货退款分析（%s · 账单基准 · 全店汇总）" % ctx["period_cn"],
                 "金额来自账单：冲销销售额 ＝ 佣金冲销比例 × 订单金额（估算，②符号）。分类规则见①。单位：%s" % p.currency)
    RF.set_widths(ws, [c.width for c in cols])
    nval = fam.value_columns()
    r = 4
    r = WB.section_head(ws, r, "一、本期冲销汇总（计费日在本月，计入②；冲销销售额，每家店一列）", len(cols))
    K.header_cells(ws, r, ["类别", ""] + fam.names + (["合计"] if fam.has_total else []))
    r += 1
    first = r
    for c in ST.REVERSAL_CLASSES:
        WB.write_cells(ws, r, 1, [ST.CLASS_CN[c]], font=F(10))
        fam.values_row(ws, r, 3, [sum(e["amount"] for e in S["revs"] if e["in_period"] and e["cls"] == c)
                                  for S in Ss])
        r += 1
    tot = total_row(ws, r, first, range(3, 3 + nval))
    money_fmt(ws, tot, range(3, 3 + nval), fam)
    r += 1
    WB.write_cells(ws, r, 1, ["残差（应为 0）：合计 − ②各冲销行"], font=BOLD, fill_color=K.C_TOTAL)
    for i in range(nval):
        L_ = get_column_letter(3 + i)
        WB.write_cells(ws, r, 3 + i, ["=ROUND(%s%d-(%s),6)" % (L_, tot, "+".join(
            plref(fam, ctx, REV_KEY[c], i) for c in ST.REVERSAL_CLASSES))], fmt=fam.money, font=BOLD,
            fill_color=K.C_TOTAL)
    r += 2
    r = WB.section_head(ws, r, "二、本月计费销售的状态（计费销售额，每家店一列；用到下载时的全部证据，含期后冲销）", len(cols))
    K.header_cells(ws, r, ["状态", ""] + fam.names + (["合计"] if fam.has_total else []))
    r += 1
    first = r
    for c in ST.CLASS_CN:
        WB.write_cells(ws, r, 1, [ST.CLASS_CN[c]])
        fam.values_row(ws, r, 3, [S["m"]["cohort_status"][c][1] for S in Ss])
        r += 1
    tot = total_row(ws, r, first, range(3, 3 + nval))
    money_fmt(ws, tot, range(3, 3 + nval), fam)
    r += 1
    WB.write_cells(ws, r, 1, ["残差（应为 0）：合计 − ②计费销售额"], font=BOLD, fill_color=K.C_TOTAL)
    for i in range(nval):
        L_ = get_column_letter(3 + i)
        WB.write_cells(ws, r, 3 + i, ["=ROUND(%s%d-%s,6)" % (L_, tot, plref(fam, ctx, "gmv_billed", i))],
                       fmt=fam.money, font=BOLD, fill_color=K.C_TOTAL)
    r += 2
    lists = [("A　发货前取消（本期冲销，计入②）", ST.CANCEL, "in"),
             ("B　退货退款（本期冲销，计入②）", ST.RETURN, "in"),
             ("C　部分退款（本期冲销，计入②）", ST.PARTIAL, "in"),
             ("D　已冲销·状态未证实（本期冲销，计入②）", ST.UNPROVEN, "in"),
             ("E　期后冲销：本月计费、月后才冲销（不计入本月，仅列示）", None, "post")]
    amount_col = 11            # 冲销销售额 (column 10 of REV_COLS, shifted by the 店铺 column)
    for title, cls, when in lists:
        rows = []
        for S in Ss:
            if when == "in":
                ev = [e for e in S["revs"] if e["in_period"] and e["cls"] == cls]
            else:
                ev = [e for e in S["revs"] if e["post"] and S["sales"].get(e["sale_no"], {}).get("billed")]
            rows += [[S["store"]] + WB._rev_values(S, e)
                     for e in sorted(ev, key=lambda e: (e["charge_date"] or S["p0"], e["sale_no"]))]
        r = WB.section_head(ws, r, "三、逐单 %s —— %d 笔" % (title, len(rows)), len(cols))
        K.header_cells(ws, r, [c.header for c in cols])
        ws.row_dimensions[r].height = 30
        r += 1
        first = r
        for vals in rows:
            fam.write_row(ws, r, cols, vals)
            r += 1
        sub = total_row(ws, r, first, (8, 9, 11, 12, 13, 14, 15), label="小计", int_cols=(8,))
        money_fmt(ws, sub, (9, 11, 12, 13, 14, 15), fam)
        r += 1
        if when == "in":
            r = WB.resid_row(fam, ws, r, len(cols), "残差（应为 0）：小计 − ②",
                             {amount_col: "=ROUND(%s%d-%s,6)" % (get_column_letter(amount_col), sub,
                                                                 plref(fam, ctx, REV_KEY[cls]))})
        r += 1
    ws.freeze_panes = "C4"


# ------------------------------------------------------------------ ⑥

def sheet_qc(wb, fam, Ss, ctx):
    p = Ss[0]["prof"]
    cols = with_store(WB.QC_COLS)
    ws = fam.table_sheet(wb, SHEETS[5], "退货质检明细（%s · 账单基准 · 全店汇总）" % ctx["period_cn"],
                         WB.QC_SUBTITLE % p.currency, cols)
    r = 5
    first = r
    for S in Ss:
        for vals in WB.qc_rows(S):
            fam.write_row(ws, r, cols, [S["store"]] + vals)
            r += 1
    tot = total_row(ws, r, first, (16, 17, 18))
    money_fmt(ws, tot, (16, 17, 18), fam)
    r += 1
    WB.resid_row(fam, ws, r, len(cols), "残差（应为 0）：退货处理费合计 − ②退货处理费",
                 {16: "=ROUND(P%d-%s,6)" % (tot, plref(fam, ctx, "fee_return"))})
    if not p.has_full:
        ws.cell(row=r + 2, column=1, value="%s没有 Full 仓：没有仓库质检文件，审核信息只来自销售报表。" %
                p.country_cn).font = GREY9
    ws.freeze_panes = "C5"


# ------------------------------------------------------------------ ⑦

def sheet_sku(wb, fam, Ss, ctx):
    p = Ss[0]["prof"]
    with_cost = all(S.get("cost") is not None for S in Ss)
    cols = with_store(WB.sku_cols(with_cost))
    ws = wb.create_sheet(SHEETS[6])
    K.title_rows(ws, "SKU 盈利分析（%s · 账单基准 · 全店汇总）" % ctx["period_cn"],
                 "表一：各店的 SKU 一行（同一 SKU 在不同店铺各占一行）。直接费用按订单/明细归属；“分摊”列是店铺级费用按"
                 "本店净商品销售额占比分摊（政策见①），两者分开列。单位：%s" % p.currency,
                 "表二（下方）：各店每笔销售 → SKU 的映射，含级别与可信度。")
    RF.set_widths(ws, [c.width for c in cols])
    r = 4
    r = WB.section_head(ws, r, "表一　SKU 盈利", len(cols))
    K.header_cells(ws, r, [c.header for c in cols])
    ws.row_dimensions[r].height = 42
    r += 1
    first = r
    for S in Ss:
        cost = S.get("cost") if with_cost else None
        for k, vals in WB.sku_rows(S, cost):
            cells = fam.write_row(ws, r, cols, [S["store"]] + vals)
            if cost is not None and (cost.line(k) is None or cost.line(k).unit_cost is None):
                cells[29].fill = fill(K.C_INPUT)
            if k.startswith(M.UNMAPPED_PREFIX) or k.startswith("（"):
                cells[1].font = F(10, False, K.C_RED)
            r += 1
    # money and count columns move one to the right (店铺 is column A)
    skip = {29} | ({30, 34} if with_cost else set())
    int_cols = {5, 6, 7, 31}
    tot = total_row(ws, r, first, range(5, len(cols) + 1), int_cols=int_cols, skip=skip)
    money_fmt(ws, tot, [j for j in range(5, len(cols) + 1) if j not in skip], fam, int_cols=int_cols)
    r += 1
    ws.cell(row=r, column=1, value="② 对应值（合计）").font = BOLD
    shift = lambda col: get_column_letter(WB._col_idx(col) + 1)
    for col, keys in WB.SKU_TIES.items():
        c = ws.cell(row=r, column=WB._col_idx(col) + 1, value="=" + "+".join(plref(fam, ctx, k) for k in keys))
        c.number_format = fam.money
        c.border = K.BOX
    r += 1
    cells = dict((WB._col_idx(col) + 1, "=ROUND(%s%d-%s%d,6)" % (shift(col), tot, shift(col), r - 1))
                 for col in WB.SKU_TIES)
    r = WB.resid_row(fam, ws, r, len(cols), "残差（应为 0）：合计 − ②", cells)
    r += 1
    map_cols = with_store(WB.MAP_COLS)
    r = WB.section_head(ws, r, "表二　账单销售 → SKU 映射（各店本月计费的销售，及本期冲销所涉的其他销售）", len(map_cols))
    K.header_cells(ws, r, [c.header for c in map_cols])
    ws.row_dimensions[r].height = 30
    r += 1
    for S in Ss:
        for vals, unmapped in WB.map_rows(S):
            fam.write_row(ws, r, map_cols, [S["store"]] + vals, fill_color=K.C_BAD if unmapped else None)
            r += 1
    ws.freeze_panes = "C6"


# ------------------------------------------------------------------ ⑧

def sheet_trend(wb, fam, Ss, ctx):
    p = Ss[0]["prof"]
    cols = ([RF.Col("周（计费日）", 14), RF.Col("计费笔数", 8, K.INT), RF.Col("件数", 8, K.INT),
             RF.Col("计费销售额", 14, RF.MONEY), RF.Col("冲销", 13, RF.MONEY),
             RF.Col("净商品销售额", 14, RF.MONEY, formula="=D{r}+E{r}", fill=K.C_CALC),
             RF.Col("买家运费（净）", 12, RF.MONEY)] +
            [RF.Col(C.FEE_CN[k], 12, RF.MONEY) for k in WB.TREND_KEYS] +
            [RF.Col("平台费用合计", 14, RF.MONEY, formula="=SUM(H{r}:P{r})", fill=K.C_CALC),
             RF.Col("经营贡献", 14, RF.MONEY, formula="=F{r}+G{r}+Q{r}", fill=K.C_CALC),
             RF.Col("费用率", 8, K.PCT, formula="=IF(F{r}=0,0,-Q{r}/F{r})", fill=K.C_CALC)])
    ws = fam.table_sheet(wb, SHEETS[7], "月度趋势（%s 按周 · 账单基准 · 全店合计）" % ctx["period_cn"],
                         "全部店铺按周（周一起，按计费日）相加。各店的周数据与前 10 SKU 见单店报表⑧。单位：%s" % p.currency,
                         cols)
    keys = ["n", "units", "gmv", "rev", "ship"] + WB.TREND_KEYS
    r = 5
    first = r
    for i, (a, b, label) in enumerate(Ss[0]["weeks"]):
        x = dict((k, sum(S["weekly"][i][k] for S in Ss)) for k in keys)
        fam.write_row(ws, r, cols, [label, x["n"], x["units"], x["gmv"], x["rev"], None, x["ship"]] +
                      [x[k] for k in WB.TREND_KEYS] + [None, None, None])
        r += 1
    tot = total_row(ws, r, first, range(2, 19), int_cols=(2, 3))
    money_fmt(ws, tot, range(4, 19), fam)
    r += 1
    ties = {"D": ("gmv_billed",), "E": ("rev_cancel", "rev_return", "rev_partial", "rev_unproven"),
            "F": ("net_gmv",), "G": ("ship_income", "ship_income_rev"), "Q": ("fees_total",), "R": ("contrib",)}
    for k, col in zip(WB.TREND_KEYS, "HIJKLMNOP"):
        ties[col] = (k,)
    ws.cell(row=r, column=1, value="② 对应值（合计）").font = BOLD
    for col, ks in ties.items():
        fam.money_cell(ws.cell(row=r, column=WB._col_idx(col)), "=" + "+".join(plref(fam, ctx, k) for k in ks))
    r += 1
    WB.resid_row(fam, ws, r, len(cols), "残差（应为 0）",
                 dict((WB._col_idx(c), "=ROUND(%s%d-%s%d,6)" % (c, tot, c, r - 1)) for c in ties))
    ws.freeze_panes = "B5"


# ------------------------------------------------------------------ ⑨

def sheet_inventory(wb, fam, Ss, ctx):
    p = Ss[0]["prof"]
    if not p.has_full:
        RF.not_applicable_sheet(wb, SHEETS[8], "库存与动销分析",
                                "不适用：%s没有 Full 仓，平台不提供库存报表（库存页 404）。本月净销量见⑦。" % p.country_cn)
        return
    cols = with_store(WB.inventory_cols(Ss[0]))
    ws = fam.table_sheet(wb, SHEETS[8], "库存与动销分析（Full 快照 vs %s 账单净销量 · 全店汇总）" % ctx["period_cn"],
                         WB.INV_SUBTITLE % ("各店下载时", "各店的对账文件"), cols)
    r = 5
    first = r
    for S in Ss:
        for vals, flag in WB.inventory_rows(S):
            fam.write_row(ws, r, cols, [S["store"]] + vals, fill_color=K.C_WARN if flag else None)
            r += 1
    tot = total_row(ws, r, first, range(4, 11), int_cols=range(4, 11))
    r += 1
    mapped = sum(S["m"]["sku_sum"]["units"] - S["m"]["sku_sum"]["units_rev"] - WB.unmapped_net_units(S) for S in Ss)
    ws.cell(row=r, column=1, value="⑦ 各店净件数合计（不含未匹配的销售）").font = BOLD
    c = ws.cell(row=r, column=10, value=mapped)
    c.number_format = K.INT
    r += 1
    WB.resid_row(fam, ws, r, len(cols), "残差（应为 0）", {10: "=ROUND(J%d-J%d,6)" % (tot, r - 1)})
    ws.freeze_panes = "C5"


# ------------------------------------------------------------------ ⑩

def sheet_gaps(wb, fam, Ss, ctx):
    p = Ss[0]["prof"]
    unm, low, flagged = [], [], []
    for S in Ss:
        dec = S["prof"].decimals
        unm += [("%s %s" % (S["store"], s["sale_no"]), "刊登 %s｜%s｜%s｜%s" % (
            s["listing"], s["title"], WB.fmtm(s["total"], dec), S["map"][s["sale_no"]].note))
            for s in S["cohort"] if S["map"][s["sale_no"]].tier == M.NONE]
        low += [("%s %s" % (S["store"], s["sale_no"]), "%s → %s（%s；%s）" % (
            s["listing"], s["sku"], M.TIER_CN[S["map"][s["sale_no"]].tier], S["map"][s["sale_no"]].note))
            for s in S["cohort"] if S["map"][s["sale_no"]].tier in (M.T2M, M.T3)]
        flagged += [("%s %s" % (S["store"], sn), S["status"][sn]["basis"][:300])
                    for sn in S["issues"]["refund_not_reversed"] + S["issues"]["cancel_not_reversed"]]
    notes = []
    if unm:
        notes.append(("未匹配 SKU 的销售（%d 笔）" % len(unm), unm[:300]))
    if low:
        notes.append(("低可信度映射（%d 笔）——请抽查" % len(low), low[:200]))
    if flagged:
        notes.append(("本月计费、有退款/取消证据但账单尚未冲销的销售（%d 笔）" % len(flagged), flagged[:200]))
    notes.append(("待你决定 / 会计判断", [
        ("冲销金额的口径", "按“佣金冲销比例 × 订单金额”估算退款（账单口径）；MercadoPago 实际退款列为旁证。"),
        ("期后冲销", "本月计费、下月才冲销的销售按计费日留在本月 GMV。"),
        ("%s" % p.vat_name, "金额含 %s；申报口径交%s会计师确认。" % (p.vat_name, p.country_cn)),
        ("店铺级费用分摊", "广告等按各店自己的净商品销售额分摊到该店的 SKU，不跨店分摊。"),
    ]))
    fam.gaps_sheet(wb, SHEETS[9], "数据缺口与待办（%s · 账单基准 · 全店汇总）" % ctx["period_cn"], ctx["checks"],
                   notes=notes)
