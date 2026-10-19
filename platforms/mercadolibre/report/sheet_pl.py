# -*- coding: utf-8 -*-
"""MercadoLibre report — ② 合并损益表: the fixed rows (PL_LINES / PL_TAIL and one
billed-fee row per FEE_CATALOG entry), their values, ratios and the memo block."""

from openpyxl.utils import get_column_letter

from core.report_family import Span, pl_formula
from core.report_kit import (BOX, C_BAD, C_CALC, C_GREY, C_INPUT, C_NAVY, C_RED, C_SUB, C_TOTAL,
    C_WARN, F, PCT, est_height, fill, wrap)

from platforms.mercadolibre.report import site
from platforms.mercadolibre.report.catalog import (FEE_BY_ES, FEE_CATALOG, SHEET_NAMES,
    TAX_RATE_NOMINAL_NET, fmt, money, pct)
from platforms.mercadolibre.report.loading import cycle_day
from platforms.mercadolibre.report.sheet_common import Sheet, agg_tax_rate, per_store, store_cols


# ════════════════════════════════════════════════════════════════════════
# ② 合并损益表的行定义 —— 写死，保证跨月行号对齐
# ════════════════════════════════════════════════════════════════════════
# (key, 中文行名, 西班牙语科目, kind, 说明)
#   kind: 'd' 取数 / 'sum' 小计(加粗浅蓝) / 'f' 计算行(灰底) / 'in' 需录入(黄底)
# key 为 None 的行由公式生成。费用行的 key 对应 metrics 里的字段名。
PL_LINES = [
    ("ing",          "商品销售收入 (GMV)",        "Ingresos por productos",           "d",
     "买家实付商品金额（含税售价），按下单日归集"),
    ("env_ing",      "运费收入",                  "Ingresos por envío",               "d",
     "买家另行支付的运费"),
    ("coupon_neg",   "减：卖家承担优惠券/折扣",    "Descuentos y bonificaciones",      "d", None),
    (None,           "营业收入合计",              "",                                 "sum",
     "＝ 商品收入 ＋ 运费收入 − 卖家优惠券"),
    ("refund_neg",   "减：退款与取消（净）",       "Anulaciones y reembolsos",         "d",
     "已退给买家的净金额（平台已返还的佣金/运费/税已抵减）"),
    ("cn_adj",       "加：部分退款重复扣佣金调整",  "(本报表修正，平台无此科目)",         "d",
     "平台对部分退款订单按退款比例退还佣金与代扣税，但销售报表的“退款与取消”列会把原佣金"
     "再扣一次。本行把多扣的加回来，逐单明细见 ⑩ 页。已用 MercadoPago 现金流水逐单核对。"),
    (None,           "净销售收入",                "",                                 "f",
     "＝ 营业收入合计 − 退款与取消 ＋ 重复扣佣金调整"),
    ("com_neg",      "减：平台销售佣金",           "Cargo por venta",                  "d",
     "账单口径纯佣金（含联盟推广佣金），按品类 13%–18%，不含代扣税"),
    ("shipcost_neg", "减：平台配送费",             "Cargo por envíos de Mercado Libre", "d",
     "Mercado Envíos 配送费（已扣平台补贴后的净额）"),
    ("dev_neg",      "减：退货处理费",             "Cargo por devolución",             "d",
     "退货产生的逆向物流/处理费"),
    ("unc_order_neg", "减：订单级未登记科目 ⚠",     "(未在 FEE_CATALOG 登记)",           "d",
     "字典里没有的订单级费用。金额照常计入以保证勾稽，但需人工确认科目归属，见 ⑩ 页"),
    (None,           "平台交易费用小计",           "",                                 "sum",
     "订单打款时即已扣除（账单 Descontado = Si）"),
    (None,           "订单毛贡献",                "",                                 "f",
     "＝ 净销售收入 − 平台交易费用"),
]
# 账单级费用行：按 FEE_CATALOG 顺序自动生成，保证新增科目时行序稳定
PL_BILL_NOTES = {
    "Cargo por campaña de publicidad de Product Ads": "已扣除当月冲销",
    "Cargo por campaña de publicidad de Display Ads": "已扣除当月冲销",
    "Cargo por servicio de almacenamiento Full": "按件按天计费；另有独立仓储费报表可交叉核对",
    "Cargo por servicio de colecta Full": "按体积 m³ 计费",
    "Cargo por comprar más espacio en Full": "主动购买的额外仓位",
    "Cargo por sobrepasar espacio Full": "占用超出配额，按超出件数罚",
    "Cargo por retiro de stock Full": "从 Full 仓取回库存的费用",
    "Cargo por stock antiguo en Full": "库龄 4 个月以上加收",
    "Cargo por incumplimiento en Envíos Full": "库存差异等履约违规罚金",
    "Cargo por mantenimiento de Mi página": "Mi página 月费",
    "Cargo por campaña de Publicidad de Seguidores": "粉丝推广广告（Seguidores）；已扣除当月冲销",
    "Cargo debitado del dinero en garantía del Programa Beneficio de Reputación":
        "信誉保障计划（Programa Beneficio de Reputación）从保证金里扣的钱；科目归属待会计确认",
}
PL_TAIL = [
    ("unc_bill_neg", "减：账单级未登记科目 ⚠",     "(未在 FEE_CATALOG 登记)",           "d",
     "字典里没有的账单级费用。金额照常计入以保证勾稽，但需人工确认科目归属，见 ⑩ 页"),
    (None,           "月度账单费用小计",           "",                                 "sum",
     "月底从 MercadoPago 余额自动扣款（账单 Descontado = No）"),
    (None,           "经营贡献毛利（未扣商品成本）", "",                                "f",
     "＝ 订单毛贡献 − 月度账单费用"),
    ("cogs",         "减：商品采购成本",           "",                                 "in",
     "⚠ 平台数据不含此项，请在【⑦ SKU盈利分析】录入单位成本后回填，或直接在此填总额"),
    (None,           "税前净利润",                "",                                 "f",
     "＝ 经营贡献毛利 − 商品采购成本"),
]

# 小计行和计算行的公式，写成数据（core/report_family 的写法）：行名的元组 = 逐行相加，
# "-行名" 表示减；Span(首行, 末行) = 两行之间（含）求和。按**行名**引用 —— 行名
# 在 build_pl_rows() 里唯一，key 为 None 的行也有名字。顺序即写入顺序。
PL_FORMULAS = [
    ("营业收入合计", Span("商品销售收入 (GMV)", "减：卖家承担优惠券/折扣")),
    ("净销售收入", ("营业收入合计", "减：退款与取消（净）", "加：部分退款重复扣佣金调整")),
    ("平台交易费用小计", Span("减：平台销售佣金", "减：订单级未登记科目 ⚠")),
    ("订单毛贡献", ("净销售收入", "平台交易费用小计")),
    # 账单级费用从 FEE_CATALOG 的第一个账单级科目（Product Ads）起，到"账单级未登记科目"止
    ("月度账单费用小计", Span("减：" + FEE_BY_ES["Cargo por campaña de publicidad de Product Ads"][0],
                            "减：账单级未登记科目 ⚠")),
    ("经营贡献毛利（未扣商品成本）", ("订单毛贡献", "月度账单费用小计")),
    ("税前净利润", ("经营贡献毛利（未扣商品成本）", "-减：商品采购成本")),
]


# ════════════════════════════════════════════════════════════════════════
# ② 合并损益表
# ════════════════════════════════════════════════════════════════════════

# 智利版 ② 的说明列：行本身（key、行名、西语科目、行序）与墨西哥完全一样，只换
# 掉说到"代扣税"的说明 —— 智利没有代扣（见 SITE_PROFILES["cl"]）。
PL_NOTES_NO_WITHHOLDING = {
    "cn_adj": "平台对部分退款订单按退款比例退还佣金，但销售报表的“退款与取消”列会把原佣金"
              "再扣一次。本行把多扣的加回来，逐单明细见 ⑩ 页。",
    "com_neg": "账单口径纯佣金（含联盟推广佣金）。智利平台不代扣税，账单佣金与销售报表佣金列"
               "在正常订单上逐单相等",
    "dev_neg": "退货产生的逆向物流/处理费。销售报表把它并进了“Cargo por venta e impuestos”列，"
               "本表按账单单列一行",
}


# 秘鲁：Descuentos y bonificaciones 列的 Flex 配送补贴并进了运费收入（见 load_ventas）
PL_NOTE_ENV_BONIF = ("买家另行支付的运费 ＋ 平台付给卖家的 Mercado Envíos Flex 配送补贴（Ventas 的"
                     "“Descuentos y bonificaciones”列，MercadoPago 里是单独一笔 CASHBACK）。"
                     "两部分各多少见 ① 页注释")


def build_pl_rows():
    """PL_LINES + 按 FEE_CATALOG 顺序生成的账单级费用行 + PL_TAIL。行序恒定。"""
    rows = list(PL_LINES)
    if not site.WITHHOLDING:
        rows = [(k, lab, es, kind, PL_NOTES_NO_WITHHOLDING.get(k, note))
                for k, lab, es, kind, note in rows]
    if site.BONIF_COL:
        rows = [(k, lab, es + " ＋ Descuentos y bonificaciones", kind, PL_NOTE_ENV_BONIF)
                if k == "env_ing" else (k, lab, es, kind, note)
                for k, lab, es, kind, note in rows]
    for es, zh, bucket, order_level in FEE_CATALOG:
        if order_level:
            continue
        rows.append(("fee:" + es, "减：" + zh, es, "d", PL_BILL_NOTES.get(es, "")))
    rows.extend(PL_TAIL)
    return rows


def pl_value(S, key):
    m = S["m"]
    if key is None:
        return None
    if key == "ing":
        return m["ing"]
    if key == "env_ing":
        return m["env_ing"]
    if key == "coupon_neg":
        return -m["coupon"]
    if key == "refund_neg":
        return -m["refund"]
    if key == "cn_adj":
        return S["cn"]["partial_total"]
    if key == "com_neg":
        return -m["com_gross"]
    if key == "shipcost_neg":
        return -m["env_cost"]
    if key == "dev_neg":
        return -m["dev_net"]
    if key == "unc_order_neg":
        return -m["unc_order"]
    if key == "unc_bill_neg":
        return -m["unc_bill"]
    if key == "cogs":
        return 0.0
    if key.startswith("fee:"):
        return -m["bill_fee_net"].get(key[4:], 0.0)
    return 0.0


def pl_memo_no_withholding(stores):
    """智利版 ② 备查项。没有代扣税，也就没有"若代扣税不可抵扣"的敏感性行。

    第一行是 Ventas 佣金列与账单纯佣金的差 —— 它在桥 A / 桥 D 里要用来勾稽，
    但不是税：主要是并进佣金列的退货处理费（已在上表“退货处理费”行按账单扣除）。
    最后一行是 GMV 中所含的 IVA，仅供参考；怎么申报、能抵多少进项是会计判断。
    """
    vat = pct(site.IVA_RATE, 0)
    return [
        ("销售报表佣金列 − 账单纯佣金（非税）", "Cargo por venta e impuestos − Cargo por venta",
         lambda S: S["m"]["tax_ret"], None,
         "智利平台不代扣税（MercadoPago 流水 TAXES_AMOUNT 为 0）。这个差额主要是销售报表并进佣金列的"
         "退货处理费 %s —— 上表已按账单在“退货处理费”行扣过，所以本行不再扣；剩余 %s 是时间/关联差。"
         "详见【%s】桥 E。"
         % (per_store(stores, lambda S: fmt(S["m"].get("vcol_dev", 0.0)), "、"),
            per_store(stores, lambda S: fmt(S["m"].get("vcol_rest", 0.0)), "、"), SHEET_NAMES[2])),
        ("口径尾差（应为 0）", "",
         lambda S: S["m"]["resid"], None,
         "上表的收入与费用逐项加总后与销售报表 Total 列的差额。为 0 表示无任何未解释差额。"),
        ("本期账单实际支付（按计费日归本月）", "Pago con débito automático / Pago voluntario",
         lambda S: S["m"]["pay_auto"], None,
         ("智利账单不按自然月出（本店切期见【①】），本月的费用分属两张账单；这里取两张账单的付款明细里"
          "付给本月计费日费用的部分，与上表的勾稽见【%s】桥 B。" % SHEET_NAMES[2])
         if any(cycle_day(S) != 1 for S in stores) else
         ("本月账单（按自然月出）实际付掉的金额（自动扣款或主动付款，不含贷记单），与上表"
          "“月度账单费用小计”的勾稽见【%s】桥 B。" % SHEET_NAMES[2])),
        ("本期贷记单（平台退还）", "Notas de crédito",
         lambda S: S["m"]["pay_nc"], None,
         "计费日在本月的贷记单行，针对取消/退货订单退还的佣金与运费，已在“退款与取消（净）”中抵减。"),
        ("GMV 中所含 IVA（%s，销项税，参考）" % vat, "IVA débito fiscal (referencia)",
         lambda S: S["m"]["iva_in_gmv"], None,
         "＝ GMV × %s/(1+%s)。智利售价含 IVA，平台不代扣，这笔税由卖家自行向 SII 申报缴纳，"
         "可用进项 IVA 抵扣。⚠ 会计判断：B2B 开票订单、退货冲减与平台费用发票的进项抵扣，"
         "请交智利会计师确认；本行不参与上表任何计算。" % (vat, vat)),
    ]


def sheet_pl(wb, stores, ctx):
    n = len(stores)
    sh = Sheet(wb, SHEET_NAMES[1], first_store_col=3, nstores=n)
    ws = sh.ws
    sh.title("合并损益表（%s）" % ctx["period_cn"],
             "口径：按【订单日】归集收入与订单级费用；月度账单费用按【计费日】归集。单位：%s。"
             "灰底=公式行，浅蓝=小计，黄底=需你录入。行位置每期固定，可直接跨月对比。" % site.CUR)
    incomplete = [S["store"] for S in stores if S.get("ventas_warning")]
    if incomplete:
        ws["A3"] = "⚠ %s 的 Ventas 数据不完整；本店及合计金额可能偏低。" % "、".join(incomplete)
        ws["A3"].font = F(10, True, C_RED)
        ws["A3"].fill = fill(C_WARN)
        ws.row_dimensions[3].height = 26
    pcol = get_column_letter(sh.tcol_idx + 1)      # 占GMV%
    ncol_idx = sh.tcol_idx + 2                      # 说明
    cols = ["项目", "西班牙语原始科目"] + store_cols(stores) + ["合计", "占GMV%", "说明"]
    widths = [30, 42] + [16] * n + [16, 10, 46]
    sh.header(4, cols, widths, 26)

    rows = build_pl_rows()
    r0 = 5
    idx = {}
    for i, (key, label, es, kind, note) in enumerate(rows):
        r = r0 + i
        idx[label] = r
        ws.cell(row=r, column=1, value=label).font = F(10, kind in ("sum", "f"))
        ws.cell(row=r, column=2, value=es or "").font = F(9, False, C_GREY)
        if kind in ("d", "in"):
            for j, S in enumerate(stores):
                ws.cell(row=r, column=sh.c0 + j, value=money(pl_value(S, key)))
        note_cell = ws.cell(row=r, column=ncol_idx, value=note or "")
        note_cell.font = F(9, False, "404040")
        note_cell.alignment = wrap()
        if note and len(note) > 46:
            ws.row_dimensions[r].height = est_height(note, 40)

    def R(label):
        return idx[label]

    # 小计 / 计算行的公式（每店一列；合计列在下面统一按店求和）
    for label, formula in PL_FORMULAS:
        for j in range(n):
            ws.cell(row=R(label), column=sh.c0 + j, value=pl_formula(formula, sh.scol(j), idx))

    gmv_row = R("商品销售收入 (GMV)")
    last = r0 + len(rows) - 1
    for r in range(r0, last + 1):
        ws.cell(row=r, column=sh.tcol_idx,
                value="=SUM(%s%d:%s%d)" % (sh.scol(0), r, sh.scol(n - 1), r))
        ws.cell(row=r, column=sh.tcol_idx + 1,
                value="=IF($%s$%d=0,0,%s%d/$%s$%d)" % (sh.tcol, gmv_row, sh.tcol, r, sh.tcol, gmv_row))
        ws.cell(row=r, column=sh.tcol_idx + 1).number_format = PCT
        for j in range(n + 1):
            ws.cell(row=r, column=sh.c0 + j).number_format = site.MNY
        for c in range(1, ncol_idx + 1):
            ws.cell(row=r, column=c).border = BOX

    for i, (key, label, es, kind, note) in enumerate(rows):
        r = r0 + i
        if kind == "sum":
            for c in range(1, ncol_idx + 1):
                ws.cell(row=r, column=c).fill = fill(C_SUB)
            for c in range(1, ncol_idx):
                ws.cell(row=r, column=c).font = F(10, True)
        elif kind == "f":
            for c in range(1, ncol_idx + 1):
                ws.cell(row=r, column=c).fill = fill(C_CALC)
            for c in range(1, ncol_idx):
                ws.cell(row=r, column=c).font = F(10, True, C_NAVY)
        elif kind == "in":
            for j in range(n):
                cc = ws.cell(row=r, column=sh.c0 + j)
                cc.fill = fill(C_INPUT)
                cc.font = F(10, True, "0000FF")
        if key in ("unc_order_neg", "unc_bill_neg"):
            has = any(abs(pl_value(S, key)) > site.ZERO_TOL for S in stores)
            for c in range(1, ncol_idx + 1):
                ws.cell(row=r, column=c).fill = fill(C_BAD if has else C_WARN)
    rr = R("税前净利润")
    for c in range(1, ncol_idx + 1):
        ws.cell(row=rr, column=c).fill = fill(C_TOTAL)
        ws.cell(row=rr, column=c).font = F(11, True, C_RED)

    # ---------------- 备查项 ----------------
    r = last + 2
    ws.cell(row=r, column=1, value="备查项（不计入上表费用）").font = F(11, True, C_NAVY)
    r += 1
    contrib_row = R("经营贡献毛利（未扣商品成本）")
    memo = [
        ("平台代扣代缴税金（IVA/ISR 预扣）", "Impuestos retenidos",
         lambda S: S["m"]["tax_ret"], None,
         "≈ 含税售价的 %s（＝不含税价的 %s），本期实测。销售报表把它和佣金混在“Cargo por venta e impuestos”里，"
         "是最容易算错的一项。它是预付税款而非成本 —— 详见【%s】桥 E。"
         % (pct(agg_tax_rate(stores), 2), pct(TAX_RATE_NOMINAL_NET, 2), SHEET_NAMES[2])),
        ("口径尾差（应为 0）", "",
         lambda S: S["m"]["resid"], None,
         "上表的收入与费用逐项加总后与销售报表 Total 列的差额。为 0 表示无任何未解释差额。"),
        ("本期账单实际支付", "Pago con débito automático / Pago voluntario",
         lambda S: S["m"]["pay_auto"], None,
         "这张账单实际付掉的金额（自动扣款或主动付款，不含贷记单），与上表“月度账单费用小计”"
         "的勾稽见【%s】桥 B。" % SHEET_NAMES[2]),
        ("本期贷记单（平台退还）", "Notas de crédito",
         lambda S: S["m"]["pay_nc"], None,
         "针对取消/退货订单退还的佣金与运费，已在“退款与取消（净）”中抵减。"),
        ("【敏感性】若代扣税不可抵扣时的贡献毛利", "", None, "sens",
         "＝ 经营贡献毛利 − 代扣代缴税金。请与墨西哥会计师确认该笔预扣能否全额抵扣；"
         "不能抵扣时利润会大幅下降，这是本报表最大的不确定项。"),
    ]
    if not site.WITHHOLDING:
        memo = pl_memo_no_withholding(stores)
    tax_memo_row = r
    for label, es, fn, special, note in memo:
        ws.cell(row=r, column=1, value=label).font = F(10)
        ws.cell(row=r, column=2, value=es).font = F(9, False, C_GREY)
        for j, S in enumerate(stores):
            c = ws.cell(row=r, column=sh.c0 + j)
            if special == "sens":
                c.value = "=%s%d-%s%d" % (sh.scol(j), contrib_row, sh.scol(j), tax_memo_row)
            else:
                c.value = money(fn(S))
            c.number_format = site.MNY
        ws.cell(row=r, column=sh.tcol_idx,
                value="=SUM(%s%d:%s%d)" % (sh.scol(0), r, sh.scol(n - 1), r)).number_format = site.MNY
        ws.cell(row=r, column=sh.tcol_idx + 1,
                value="=IF($%s$%d=0,0,%s%d/$%s$%d)" % (sh.tcol, gmv_row, sh.tcol, r, sh.tcol, gmv_row)
                ).number_format = PCT
        nc = ws.cell(row=r, column=ncol_idx, value=note)
        nc.font = F(9, False, "404040")
        nc.alignment = wrap()
        ws.row_dimensions[r].height = est_height(note, 44)
        for c in range(1, ncol_idx + 1):
            ws.cell(row=r, column=c).border = BOX
        if special == "sens":      # 敏感性行整行加粗标红
            for c in range(1, ncol_idx):
                ws.cell(row=r, column=c).font = F(10, True, C_RED)
        r += 1
    ws.freeze_panes = "%s5" % sh.scol(0)
    return {"contrib_row": contrib_row}
