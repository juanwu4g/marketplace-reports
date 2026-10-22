# -*- coding: utf-8 -*-
"""MercadoLibre report — ④ 平台费用明细: every fee type of the period by settlement
path, reconciled to ②."""

from openpyxl.styles import Alignment

from core.report_family import Col, Family, MONEY, header_row
from core.report_kit import (BOX, C_AMBER, C_BAD, C_BAND, C_GREY, C_NAVY, C_RED, C_SUB, F, INT,
    PCT, est_height, fill, header_cells, title_rows)

from platforms.mercadolibre.report import site
from platforms.mercadolibre.report.catalog import (CHANNEL_MEANING, CHANNEL_ZH, CH_NA, CH_NO,
    CH_SI, FEE_BY_ES, F_DETALLE, SHEET_NAMES, classify, fmt, money)


# ════════════════════════════════════════════════════════════════════════
# ④ 平台费用明细
# ════════════════════════════════════════════════════════════════════════

# 逐科目明细表的列（core/report_family.Col）：表头、列宽、数字格式、字体。
FEE_COLS = [
    Col("店铺", 18, font=F(10)),
    Col("西班牙语科目 (Detalle)", 52, font=F(10)),
    Col("中文科目", 30, font=F(10)),
    Col("结算通道", 16, font=F(10)),
    Col("笔数", 10, INT, font=F(10)),
    Col("金额", 16, MONEY, font=F(10)),          # 本站点的金额格式（智利无小数）
    Col("占该店 GMV%", 14, PCT, font=F(10)),
]


CHAN_ORDER = {CH_SI: 0, CH_NO: 1, CH_NA: 2}


def fee_groups(S):
    """④ 逐科目表的行：账单（含贷记单）按 (Detalle, 结算通道) 汇总，
    列 detalle / channel / count / sum（账单原符号），按页面上的顺序。
    ④ 页和结果文件（report/figures.py）都用它。"""
    FA = S["billing"]
    g = FA.groupby([FA[F_DETALLE].astype(str), FA["channel"]])["amount"].agg(["count", "sum"]).reset_index()
    g.columns = ["detalle", "channel", "count", "sum"]
    g["_c"] = g["channel"].map(lambda x: CHAN_ORDER.get(x, 9))
    return g.sort_values(["_c", "sum"], key=lambda s: s if s.name != "sum" else s.abs(),
                         ascending=[True, False])


def sheet_fees(wb, stores, ctx):
    fam = Family(stores, money=site.MNY)
    ws = wb.create_sheet(SHEET_NAMES[3])
    title_rows(ws, "平台费用明细（%s账单，按计费日归集）" % ctx["period_cn"],
               "数据源：Reporte_Facturacion_MercadoLibre + Reporte_Notas_Credito。"
               "结算通道决定这笔钱是“下单时已扣”还是“月底再扣”，混淆两者是平不了账的主因之一。"
               "未在 FEE_CATALOG 登记的科目以红底标出。单位：%s" % site.CUR)
    header_row(ws, 4, [c.header for c in FEE_COLS], [c.width for c in FEE_COLS])

    r = 5
    for S in stores:
        gmv = S["m"]["ing"] or 1.0
        g = fee_groups(S)
        start = r
        for _, x in g.iterrows():
            zh, bucket, _o, is_rev, known = classify(x["detalle"])
            cells = fam.write_row(ws, r, FEE_COLS, (
                S["store"], x["detalle"], zh, CHANNEL_ZH.get(x["channel"], x["channel"]),
                int(x["count"]), money(x["sum"]), "=F%d/%s" % (r, gmv)))
            if not known:                               # 未登记科目：红底红字
                for c in cells:
                    c.fill = fill(C_BAD)
                    c.font = F(10, True, C_RED)
            elif is_rev:                                # 冲销行：褐色字
                for c in cells:
                    c.font = F(10, False, "8C4A00")
            r += 1
        cells = fam.write_row(ws, r, FEE_COLS, (
            S["store"], None, "账单费用合计", None, "=SUM(E%d:E%d)" % (start, r - 1),
            "=SUM(F%d:F%d)" % (start, r - 1), "=F%d/%s" % (r, gmv)), fill_color=C_SUB)
        for c in cells:
            c.font = F(10, True)
        r += 2

    ws.cell(row=r, column=2, value="按结算通道汇总").font = F(12, True, C_NAVY)
    r += 1
    header_cells(ws, r, ["店铺", "结算通道", "含义", "", "笔数", "金额", "占该店 GMV%"], wrap_text=False)
    r += 1
    for S in stores:
        gmv = S["m"]["ing"] or 1.0
        for k in (CH_SI, CH_NO, CH_NA):
            ws.cell(row=r, column=1, value=S["store"])
            ws.cell(row=r, column=2, value="%s  (Descontado = %s)" % (CHANNEL_ZH[k], k))
            ws.cell(row=r, column=3, value=CHANNEL_MEANING[k]).font = F(9, False, "404040")
            ws.cell(row=r, column=5, value=S["m"]["bill_cnt_by_channel"][k]).number_format = INT
            ws.cell(row=r, column=6, value=money(S["m"]["bill_by_channel"][k])).number_format = site.MNY
            ws.cell(row=r, column=7, value="=F%d/%s" % (r, gmv)).number_format = PCT
            for c in range(1, 8):
                ws.cell(row=r, column=c).border = BOX
            r += 1

    # ---- 与 ② 合并损益表的对账 ----
    # 这两页的费用合计**本来就不相等**，而页面上原先没有任何地方说明为什么。
    # 有人拿两边的合计相减、再减去贷记单文件的金额，得到一个说不清的余数
    # （BOCINA_TA02 2026-08 是 249.22）——那个余数不对应任何东西：冲销行并不
    # 只在贷记单文件里，账单文件自己也带着一堆 Anulación 行（该店 46 行冲销
    # 里有 29 行在账单文件内）。所以这里直接把差异逐项拆开，残差应为 0。
    r += 2
    ws.cell(row=r, column=2, value="与【%s】的对账：为什么两页的费用合计不相等"
            % SHEET_NAMES[1]).font = F(12, True, C_NAVY)
    r += 1
    ws.cell(row=r, column=2, value=(
        "本页按【计费日】列出账单上的每一行；损益表按【订单日】归集订单级费用，"
        "而且佣金取毛额、配送费取自销售报表。差异全部来自这两点，下表逐项拆开。"
        "完整口径对比见【%s】桥 C。" % SHEET_NAMES[2])).font = F(9, False, C_GREY)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    ws.row_dimensions[r].height = 30
    r += 1
    header_cells(ws, r, ["店铺", "项目", "② 损益表", "④ 本页", "差", "原因"], first_col=2, color=C_BAND)
    ws.row_dimensions[r].height = 24
    r += 1

    for S in stores:
        m, FA = S["m"], S["billing"]

        def fa_bucket(k):
            return float(FA[FA["parent_detalle"].map(
                lambda d: FEE_BY_ES.get(str(d), (None, None, None))[1] == k
            ).astype(bool)]["amount"].sum())

        fa_com, fa_env = fa_bucket("commission"), fa_bucket("shipping")
        fa_dev = fa_bucket("returns")
        pl_bill = sum(m["bill_fee_net"].values()) + m["unc_bill"]
        fa_bill = m["bill_total"] - fa_com - fa_env - fa_dev
        rows = [
            ("平台销售佣金", m["com_gross"], fa_com,
             "损益表取**毛额**（不减冲销）；本页是净额。销售报表的“退款与取消”列"
             "已经把平台返还的费用扣进去了，损益表再减一次就重复抵免。"),
            ("平台配送费", m["env_cost"], fa_env,
             "损益表取自**销售报表**的 Costos de envío；本页是**账单**的配送费。"
             "两个数据源，差额见桥 C。"),
            ("退货处理费", m["dev_net"], fa_dev,
             "损益表按订单关联跨账期取数；本页只看本期计费日。"),
            ("广告 / Full / 页面维护等账单级费用", pl_bill, fa_bill,
             "两边同源同口径，应当相等。"),
        ]
        start = r
        for label, a, b, why in rows:
            ws.cell(row=r, column=2, value=S["store"])
            ws.cell(row=r, column=3, value=label)
            ws.cell(row=r, column=4, value=money(a)).number_format = site.MNY
            ws.cell(row=r, column=5, value=money(b)).number_format = site.MNY
            ws.cell(row=r, column=6, value="=D%d-E%d" % (r, r)).number_format = site.MNY
            c = ws.cell(row=r, column=7, value=why)
            c.font = F(9, False, C_GREY)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = est_height(why, 44)
            for c0 in range(2, 8):
                ws.cell(row=r, column=c0).border = BOX
            r += 1
        ws.cell(row=r, column=3, value="合计差异").font = F(10, True)
        for col, cl in ((4, "D"), (5, "E"), (6, "F")):
            ws.cell(row=r, column=col, value="=SUM(%s%d:%s%d)" % (cl, start, cl, r - 1)
                    ).number_format = site.MNY
        ws.cell(row=r, column=7, value=(
            "＝ 上面四项之差相加。这就是两页费用合计的全部差异，没有残数。")
            ).font = F(9, False, C_GREY)
        for c0 in range(2, 8):
            ws.cell(row=r, column=c0).fill = fill(C_SUB)
            ws.cell(row=r, column=c0).font = F(10, True)
            ws.cell(row=r, column=c0).border = BOX
        r += 1
        # 冲销行到底在哪个文件里 —— 这正是上面那个"说不清的余数"的来源
        rev = FA[FA["is_rev"]]
        by_doc = rev.groupby("_doc")["amount"].agg(["count", "sum"]) if len(rev) else None
        nc_n = int(by_doc.loc["NC", "count"]) if by_doc is not None and "NC" in by_doc.index else 0
        nc_v = float(by_doc.loc["NC", "sum"]) if by_doc is not None and "NC" in by_doc.index else 0.0
        c = ws.cell(row=r, column=3, value=(
            "⚠ 本期共 %d 行冲销、合计 %s，其中只有 %d 行（%s）来自贷记单文件 "
            "Reporte_Notas_Credito，其余 %d 行就写在账单文件 Reporte_Facturacion 里。"
            "**不要拿贷记单文件的合计当作本期全部退还** —— 那样算出来的差额对不上任何东西。"
            % (len(rev), fmt(float(rev["amount"].sum()) if len(rev) else 0.0),
               nc_n, fmt(nc_v), len(rev) - nc_n)))
        c.font = F(9, False, C_AMBER)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
        ws.row_dimensions[r].height = est_height(str(c.value), 110)
        r += 2
    ws.freeze_panes = "A5"
