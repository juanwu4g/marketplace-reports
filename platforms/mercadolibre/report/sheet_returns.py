# -*- coding: utf-8 -*-
"""MercadoLibre report — ⑤ 退货退款分析 and ⑥ 退货质检明细."""

import pandas as pd
from openpyxl.utils import get_column_letter

from core.report_family import Col, Family, MONEY, header_row, set_widths
from core.report_kit import (BOX, C_AMBER, C_BAD, C_GREY, C_NAVY, C_OK, C_SUB, F, INT, PCT, fill,
    header_cells, title_rows, wrap)

from platforms.mercadolibre.report import site
from platforms.mercadolibre.report.catalog import (MONEY_KEPT, MONEY_REFUNDED, RETURN_REASON_ZH,
    RETURN_RESULT_ZH, SHEET_NAMES, STATUS_ZH, V_STATUS, V_UNITS, fmt, money, num)
from platforms.mercadolibre.report.loading import date_span
from platforms.mercadolibre.report.sheet_common import no_full_text, store_cols


# ════════════════════════════════════════════════════════════════════════
# ⑤ 退货退款分析
# ════════════════════════════════════════════════════════════════════════

# ⑤ 按订单状态的明细表（core/report_family.Col）
STATUS_COLS = [
    Col("店铺", 16, font=F(10)),
    Col("订单状态 (Estado)", 46, font=F(10)),
    Col("中文含义", 30, font=F(10), align=wrap()),
    Col("单数", 9, INT, font=F(10)),
    Col("件数", 9, INT, font=F(10)),
    Col("原销售额", 15, MONEY, font=F(10)),
    Col("实际退款额", 15, MONEY, font=F(10)),
    Col("退款/原销售额", 13, PCT, font=F(10), formula="=IF(F{r}=0,0,G{r}/F{r})"),
    Col("资金归属", 14, font=F(10)),
    Col("库存处置", 26, font=F(10), align=wrap()),
]


def status_groups(S):
    """⑤ 上表的行：售后/异常订单按 Estado 汇总（estado / 单数 / 件数 / 原销售额 /
    退款，退款为正数），按原销售额从高到低；本月没有售后订单返回 None。
    ⑤ 页和结果文件（report/figures.py）都用它。"""
    A = S["ventas"]
    ab = A[A["abnormal"]]
    if not len(ab):
        return None
    g = ab.groupby(A[V_STATUS].astype(str)).apply(
        lambda x: pd.Series({"单数": len(x), "件数": num(x, V_UNITS).sum(),
                             "原销售额": num(x, site.V_REVENUE).sum(),
                             "退款": -num(x, site.V_REFUND).sum()})).reset_index()
    g.columns = ["estado", "单数", "件数", "原销售额", "退款"]
    return g.sort_values("原销售额", ascending=False)


def sheet_returns_money(wb, stores, ctx):
    fam = Family(stores, money=site.MNY)
    ws = wb.create_sheet(SHEET_NAMES[4])
    title_rows(ws, "退货 · 退款 · 售后分析（%s订单）" % ctx["period_cn"],
               "一笔退货会同时影响：退给买家的货款、平台退还的佣金与运费、额外的退货处理费、以及库存能否回收。"
               "下表按订单状态拆开，并标明责任方。单位：%s" % site.CUR)
    n_par = sum(len(S["cn"]["partial"]) for S in stores)
    if n_par:
        ws["A3"] = ("⚠ 本页的“实际退款额”直接取自销售报表的 Anulaciones y reembolsos 列，未作修正。"
                    "其中 %d 笔**部分退款**订单，平台在这一列里把原佣金重复扣了一次，"
                    "实际退款没有这么多（合计虚高 %s）—— 已用 MercadoPago 现金流水逐单证实。"
                    "利润口径的修正在【%s】的“加：部分退款重复扣佣金调整”行，"
                    "逐单明细见【%s】3.3。本页保持与销售报表原始数据一一对应，故不在此处修正。"
                    % (n_par, fmt(sum(S["cn"]["partial_total"] for S in stores)),
                       SHEET_NAMES[1], SHEET_NAMES[9]))
        ws["A3"].font = F(9, False, C_AMBER)
    header_row(ws, 4, [c.header for c in STATUS_COLS], [c.width for c in STATUS_COLS])

    r = 5
    for S in stores:
        g = status_groups(S)
        if g is None:
            continue
        start = r
        for _, x in g.iterrows():
            zh, who, inv = STATUS_ZH.get(x["estado"], (x["estado"], "待定", "待确认"))
            fam.write_row(ws, r, STATUS_COLS, (
                S["store"], x["estado"], zh, int(x["单数"]), int(x["件数"]), money(x["原销售额"]),
                money(x["退款"]), None, who, inv),
                # 库存不可再售标红，可回收标绿
                fill_color=C_BAD if inv.startswith("✘") else (C_OK if inv.startswith("✔") else None))
            r += 1
        ws.cell(row=r, column=2, value="%s 售后订单合计" % S["store"]).font = F(10, True)
        for col, cl in ((4, "D"), (5, "E"), (6, "F"), (7, "G")):
            ws.cell(row=r, column=col, value="=SUM(%s%d:%s%d)" % (cl, start, cl, r - 1)
                    ).number_format = INT if col in (4, 5) else site.MNY
        ws.cell(row=r, column=8, value="=IF(F%d=0,0,G%d/F%d)" % (r, r, r)).number_format = PCT
        for c in range(1, 11):
            ws.cell(row=r, column=c).fill = fill(C_SUB)
            ws.cell(row=r, column=c).font = F(10, True)
            ws.cell(row=r, column=c).border = BOX
        r += 2

    # ---- KPI ----
    r += 1
    ws.cell(row=r, column=1, value="退货核心指标").font = F(12, True, C_NAVY)
    r += 1
    n = len(stores)
    hdr = ["指标"] + store_cols(stores) + ["合计", "说明"]
    header_cells(ws, r, hdr, wrap_text=False)
    ws.row_dimensions[r].height = 22
    r += 1
    tcol = 1 + n + 1
    ncol = tcol + 1
    kpi_rows = [
        ("本期订单总数", lambda S: S["m"]["rows"], INT, "%s 按下单日，按行计" % site.VENTAS_SHEET),
        ("其中：售后/异常订单数", lambda S: S["m"]["abnormal_orders"], INT, "含取消、退货、仲裁、投诉、换货"),
        ("售后订单占比", "ratio_prev", PCT, "售后订单数 ÷ 订单总数"),
        ("实际退款金额", lambda S: S["m"]["refund"], site.MNY, "Anulaciones y reembolsos 净额"),
        ("退款率（占 GMV）", "ratio_gmv", PCT, "退款金额 ÷ GMV"),
        ("退货处理费（平台另收）", lambda S: S["m"]["dev_net"], site.MNY, "Cargo por devolución 净额，账单中单独收取"),
        ("　— 库存可回收（重新上架）单数", lambda S: S["m"]["n_back_on_sale"], INT, "商品质检合格，已重新上架销售，库存损失为 0"),
        ("　— 库存不可再售（自提/销毁）单数", lambda S: S["m"]["n_stock_lost"], INT,
         "★ 需付费从 Full 仓取回或直接销毁，货款与货物双重损失"),
        ("　— 卖家保留货款单数", lambda S: S["m"]["n_seller_keeps"], INT, "平台仲裁判卖家胜诉或退货被驳回，未退款"),
        ("库存不可再售订单原销售额", lambda S: S["m"]["amt_stock_lost"], site.MNY,
         "★ 这部分是最实在的损失：钱退了、货也拿不回来卖"),
    ]
    for label, fn, code, note in kpi_rows:
        ws.cell(row=r, column=1, value=label).font = F(10)
        if fn == "ratio_prev":
            for j in range(n):
                cl = get_column_letter(2 + j)
                ws.cell(row=r, column=2 + j, value="=IF(%s%d=0,0,%s%d/%s%d)" % (cl, r - 2, cl, r - 1, cl, r - 2))
            ws.cell(row=r, column=tcol, value="")
        elif fn == "ratio_gmv":
            for j, S in enumerate(stores):
                cl = get_column_letter(2 + j)
                gmv = S["m"]["ing"] or 1.0
                ws.cell(row=r, column=2 + j, value="=%s%d/%s" % (cl, r - 1, gmv))
            ws.cell(row=r, column=tcol, value="")
        else:
            for j, S in enumerate(stores):
                ws.cell(row=r, column=2 + j, value=money(fn(S)))
            ws.cell(row=r, column=tcol,
                    value="=SUM(%s%d:%s%d)" % (get_column_letter(2), r, get_column_letter(1 + n), r))
        for j in range(n + 1):
            ws.cell(row=r, column=2 + j).number_format = code
        c = ws.cell(row=r, column=ncol, value=note)
        c.font = F(9, False, "404040")
        c.alignment = wrap()
        for c2 in range(1, ncol + 1):
            ws.cell(row=r, column=c2).border = BOX
        if "不可再售" in label:
            for c2 in range(1, ncol + 1):
                ws.cell(row=r, column=c2).fill = fill(C_BAD)
        elif "可回收" in label:
            for c2 in range(1, ncol + 1):
                ws.cell(row=r, column=c2).fill = fill(C_OK)
        r += 1
    ws.column_dimensions[get_column_letter(ncol)].width = 60


# ════════════════════════════════════════════════════════════════════════
# ⑥ 退货质检明细
# ════════════════════════════════════════════════════════════════════════

# 每家店三张表：按退货原因、按质检处置结果、逐件明细（core/report_family.Col）
REASON_COLS = [Col("退货原因 (Estado del producto)", 42, font=F(10)), Col("中文", 26, font=F(10)),
               Col("件数", 9, INT, font=F(10)), Col("占比", 10, PCT, font=F(10)),
               Col("其中：已退款买家", 15, INT, font=F(10)), Col("其中：货款归卖家", 15, INT, font=F(10)),
               Col("库存后果", 34, font=F(10))]
RESULT_COLS = [Col("处置结果 (Resultado de la revisión)", None, font=F(10)), Col("中文", None, font=F(10)),
               Col("件数", None, INT, font=F(10)), Col("占比", None, PCT, font=F(10)),
               Col("资金结果", None, font=F(10)), Col("库存价值", None, font=F(10)),
               Col("财务影响", None, font=F(10))]
DETAIL_COLS = [Col(h, None, font=F(9)) for h in (
    "订单号", "SKU", "质检日期", "商品状态 (原因)", "质检处置结果", "资金处理", "中文说明")]
# 质检处置结果 → 整行底色：要卖家花钱取回的标红，重新上架的标绿
RESULT_FILL = {"Producto para retirar en centro de distribución": C_BAD,
               "Producto quedó nuevamente para la venta": C_OK}


def sheet_triage(wb, stores, ctx):
    fam = Family(stores, money=site.MNY)
    ws = wb.create_sheet(SHEET_NAMES[5])
    title_rows(ws, "Full 仓退货质检明细与原因分析")
    if not site.HAS_FULL:
        ws["A2"] = no_full_text()
        ws["A2"].font = F(10, False, C_GREY)
        ws["A4"] = ("退货本身（退款金额、订单状态、平台/卖家责任）照常在【%s】分析；"
                    "退回的货直接回到卖家手里，可否再售需卖家自己记录。" % SHEET_NAMES[4])
        ws["A4"].font = F(10, False, C_GREY)
        return
    have = [S for S in stores if len(S["returns"])]
    none = [S["store"] for S in stores if not len(S["returns"])]
    cover = "；".join(
        "%s 覆盖 %s，共 %d 件（其中本期 %d 件）"
        % (S["store"], date_span(S["returns_all"]["fr"]) or "（质检日读不出来）",
           len(S["returns_all"]), len(S["returns"]))
        for S in stores if len(S.get("returns_all", [])))
    ws["A2"] = ("数据源：Returns_*.xlsx（Triages 页），按【质检日】记录，不是按订单日 —— "
                "本期订单的退货可能下月才质检，反之亦然。%s%s"
                % (cover, ("。%s 无退货质检记录。" % "、".join(none)) if none else ""))
    ws["A2"].font = F(9, False, C_GREY)
    r = 4
    if not have:
        ws.cell(row=r, column=1, value="本期没有任何退货质检记录。").font = F(11, False, C_GREY)
        return

    for S in have:
        R = S["returns"]
        tot = len(R) or 1
        ws.cell(row=r, column=1, value="【%s】一、按退货原因汇总（本期质检 %d 件）" % (S["store"], len(R))
                ).font = F(12, True, C_NAVY)
        r += 1
        header_cells(ws, r, [c.header for c in REASON_COLS])
        set_widths(ws, [c.width for c in REASON_COLS])
        r += 1
        start = r
        for k, v in R.groupby("Estado del producto").size().sort_values(ascending=False).items():
            zh, cons = RETURN_REASON_ZH.get(k, (str(k), "待确认"))
            sub = R[R["Estado del producto"] == k]
            fam.write_row(ws, r, REASON_COLS, (
                str(k), zh, int(v), "=C%d/%d" % (r, tot),
                int((sub["Estado del dinero"] == MONEY_REFUNDED).sum()),
                int((sub["Estado del dinero"] == MONEY_KEPT).sum()), cons),
                fill_color=C_BAD if cons.startswith("✘") else None)
            r += 1
        ws.cell(row=r, column=1, value="合计").font = F(10, True)
        for col, cl in ((3, "C"), (5, "E"), (6, "F")):
            ws.cell(row=r, column=col, value="=SUM(%s%d:%s%d)" % (cl, start, cl, r - 1)).number_format = INT
        ws.cell(row=r, column=4, value="=C%d/%d" % (r, tot)).number_format = PCT
        for c in range(1, 8):
            ws.cell(row=r, column=c).fill = fill(C_SUB)
            ws.cell(row=r, column=c).font = F(10, True)
            ws.cell(row=r, column=c).border = BOX
        r += 2

        ws.cell(row=r, column=1, value="【%s】二、按质检处置结果汇总" % S["store"]).font = F(12, True, C_NAVY)
        r += 1
        header_cells(ws, r, [c.header for c in RESULT_COLS])
        r += 1
        for k, v in R.groupby("Resultado de la revisión").size().sort_values(ascending=False).items():
            zh, val, fin = RETURN_RESULT_ZH.get(k, (str(k), "待确认", "待确认"))
            sub = R[R["Resultado de la revisión"] == k]
            ref = int((sub["Estado del dinero"] == MONEY_REFUNDED).sum())
            fam.write_row(ws, r, RESULT_COLS, (
                str(k), zh, int(v), "=C%d/%d" % (r, tot), "退款 %d 件 / 保留 %d 件" % (ref, int(v) - ref),
                val, fin), fill_color=RESULT_FILL.get(k))
            r += 1
        r += 1

        ws.cell(row=r, column=1, value="【%s】三、逐件明细（本期质检记录）" % S["store"]).font = F(12, True, C_NAVY)
        r += 1
        header_cells(ws, r, [c.header for c in DETAIL_COLS])
        r += 1
        for _, x in R.sort_values("fr").iterrows():
            zh = RETURN_REASON_ZH.get(x["Estado del producto"], (str(x["Estado del producto"]), ""))[0]
            dz = RETURN_RESULT_ZH.get(x["Resultado de la revisión"],
                                      (str(x["Resultado de la revisión"]), "", ""))[0]
            mon = "已退款买家" if x["Estado del dinero"] == MONEY_REFUNDED else "货款归卖家"
            fam.write_row(ws, r, DETAIL_COLS, (
                str(x.get("Número de orden", "")).replace(".0", ""), x.get("SKU"),
                x["fr"].strftime("%Y-%m-%d") if pd.notna(x["fr"]) else "", x["Estado del producto"],
                x["Resultado de la revisión"], x["Estado del dinero"], "%s；%s；%s" % (zh, dz, mon)),
                fill_color=RESULT_FILL.get(x["Resultado de la revisión"]))
            r += 1
        r += 2
