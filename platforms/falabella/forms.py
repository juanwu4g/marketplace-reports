# -*- coding: utf-8 -*-
"""The seven Falabella routes for one store and one month (FORMS.md).

Every route takes a Ctx and fills ctx.files[form] with the downloaded file names.
When a form cannot be had it records why — ctx.give_up(form, signal, ...) —
instead of raising: download_diag turns the evidence into a cause code. Only a
store-level problem (wrong language, dead browser) is raised to run_downloads.

Form ids are those of periods.form_ids(); routes are the part before "/".
"""
import datetime
import os
import re
import time
import traceback

from selenium.common.exceptions import TimeoutException

from core import browser_kit
from core import download_diag
from platforms.falabella import login
from platforms.falabella import page_kit as pk
from platforms.falabella import periods

PAYMENTS = "/report/transaction/statement#/ordersandtransactionsoverview"
ORDERS = "/order"
DASHBOARD = "/index/dashboard"
FULFILLMENT = "/fulfillment"
SETTLEMENT = "/invoicing#/settlement-invoice?period=%d"
LIQUIDATION = "/invoicing#/liquidation-invoice?period=%d"

ORDERS_HOST = "div#orders"
MONTH_ABBR = ("ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep",
              "oct", "nov", "dic")


class LanguageWrong(RuntimeError):
    """The seller centre is not in Spanish. Message contains 界面语言 so that
    download_diag classifies it as LANGUAGE_WRONG."""


class Ctx(object):
    """Everything one store's routes share. One per store (never module state)."""

    def __init__(self, driver, out_dir, diag, month, site, today=None,
                 collect_timeout=240, log=print):
        self.d = driver
        self.out_dir = out_dir
        self.diag = diag
        self.month = month
        self.site = site
        self.host = site["host"]
        self.today = today or datetime.date.today()
        self.collect_timeout = collect_timeout
        self.log = log
        self.files = {}          # form -> [file names]
        self.no_data = {}        # form -> reason (platform says: nothing there)
        self.pending = []        # requested, not ready: owed entries
        self.errors = []         # route-level failures, one line each
        self.flags = {}

    def go(self, path, settle=0):
        self.d.get(self.host + path)
        if settle:
            time.sleep(settle)

    def got(self, form, names):
        names = [n for n in names if n]
        if names:
            self.files.setdefault(form, []).extend(names)
            download_diag.delivered(self.diag, form, self.files[form])
        return names

    def give_up(self, form, signal, **kw):
        self.log("  [放弃] %s：%s %s" % (form, signal, kw.get("step") or ""))
        return download_diag.give_up(self.diag, self.d, form, signal, **kw)

    def nothing_there(self, form, reason):
        """The platform has nothing for this form in this period (terminal)."""
        self.log("  [无数据] %s：%s" % (form, reason))
        self.no_data[form] = reason
        self.give_up(form, "no-data", row_state="empty", status_text=reason, shot=False)

    def rename(self, name, new):
        src, dst = os.path.join(self.out_dir, name), os.path.join(self.out_dir, new)
        if os.path.exists(dst):
            os.remove(dst)
        os.replace(src, dst)
        return new


def guarded(ctx, route, forms, fn, *a):
    """Run one route; an exception gives up every form of it not yet delivered."""
    try:
        fn(ctx, *a)
    except LanguageWrong:
        raise
    except Exception as e:
        ctx.errors.append("%s: %s: %s" % (route, e.__class__.__name__, str(e)[:200]))
        ctx.log("  [路线失败] %s：%s" % (route, e))
        if not isinstance(e, pk.PageError):
            traceback.print_exc()
        ev = download_diag.exception_evidence(e, step=route)
        for f in forms:
            if f not in ctx.files and f not in ctx.no_data:
                ctx.give_up(f, "exception", **{k: v for k, v in ev.items() if k != "signal"})


# ================================================================ language gate

FIRST_LOAD_CAP = 60       # seconds the first page of a store may take to load


def check_language(ctx):
    """Open the payments page and read data-localefull. Not Spanish -> raise."""
    # 登录过期时这一下会落到会无限重载的统一登录页（login.py）：Selenium 默认等
    # 300 秒才放弃，然后整家店失败、自动登录根本没试。先把等待压到 60 秒，超时
    # 也照样看看停在哪儿 —— 在登录页就去登录。地址要在压低的时限里读：页面还在
    # 转圈时，读地址也要等它。
    with pk.page_load_cap(ctx.d, FIRST_LOAD_CAP):
        try:
            ctx.go(PAYMENTS)
        except TimeoutException:
            ctx.log("  付款页 %d 秒没加载完，看看是不是停在登录页" % FIRST_LOAD_CAP)
        url = login.current_url(ctx.d)
    if login.on_login_page(url) and login.sign_in(ctx):
        ctx.go(PAYMENTS)
    info = pk.wait_for(lambda: pk.page_locale(ctx.d), timeout=60, poll=2)
    if not info:
        raise pk.PageError("payments page never rendered (no data-localefull found)")
    loc = str(info.get("locale") or "")
    ctx.flags["locale"] = loc
    ctx.flags["seller"] = info.get("seller")
    ctx.flags["seller_name"] = info.get("name")
    want = ctx.site["locale"]
    if not loc.lower().startswith("es"):
        # Someone switched the store to another language (it happened on DAB on
        # 2026-10-19). Every selector and every column name is Spanish, so switch it
        # back — the language is a setting of the account, not of this browser —
        # and say so in the run's record.
        ctx.log("  界面语言是 %s，不是西班牙语：在页脚的语言选项里改回 Español" % loc)
        after = pk.switch_language(ctx.d, "Español")
        new_loc = str((after or {}).get("locale") or "")
        if not new_loc.lower().startswith("es"):
            raise LanguageWrong("界面语言是 %s，不是西班牙语（%s），脚本自动改回 Español 也没成功："
                                "请在这家店的页面最下方把 Language 改成 Español" % (loc, want))
        ctx.flags["language_switched"] = {"from": loc, "to": new_loc}
        ctx.log("  已改回 %s" % new_loc)
        loc = new_loc
        ctx.flags["locale"] = loc
        ctx.go(PAYMENTS)                      # the routes start from a freshly rendered page
        pk.wait_for(lambda: pk.page_locale(ctx.d), timeout=60, poll=2)
    if loc != want:
        ctx.log("  [提醒] 页面区域设置是 %s（期望 %s）" % (loc, want))
    return info


# ================================================================ 1. order

def _orders_root_ready(ctx):
    return pk.shadow(ctx.d, ORDERS_HOST, "#tab-5")


def _picker_cell(ctx, day):
    cells = [c for c in pk.find(ctx.d, "td.ant-picker-cell[title='%s']" % periods.iso(day))
             if "in-view" in (c.get_attribute("class") or "")]
    return cells[0] if cells else None


def _walk_to(ctx, day, prev_sel, next_sel):
    for _ in range(18):
        c = _picker_cell(ctx, day)
        if c is not None:
            return c
        first = pk.find(ctx.d, "td.ant-picker-cell-in-view[title]")
        shown = first[0].get_attribute("title") if first else None
        go_back = shown is not None and periods.iso(day) < shown
        btns = pk.find(ctx.d, prev_sel if go_back else next_sel)
        if not btns:
            raise pk.PageError("calendar navigation button not found (%s)"
                               % (prev_sel if go_back else next_sel))
        pk.js_click(ctx.d, btns[0] if go_back else btns[-1])
        time.sleep(0.6)
    raise pk.PageError("calendar day %s not found" % day)


CHIP = re.compile(r"([a-záéíóú]{3})[a-z.]*\s+(\d{1,2}),\s*(\d{4})\s*-\s*"
                  r"([a-záéíóú]{3})[a-z.]*\s+(\d{1,2}),\s*(\d{4})", re.I)


def chip_matches(text, a, b):
    """Does the search-bar chip ('sep 1, 2026 - sep 14, 2026') show a..b?"""
    m = CHIP.search(text or "")
    if not m:
        return False
    m1, d1, y1, m2, d2, y2 = m.groups()

    def same(mon, day, year, want):
        return (mon.lower()[:3] in (MONTH_ABBR[want.month - 1],
                                    want.strftime("%b").lower()[:3])
                and int(day) == want.day and int(year) == want.year)
    return same(m1, d1, y1, a) and same(m2, d2, y2, b)


def orders_chunk(ctx, form, a, b):
    d = ctx.d
    ctx.go(ORDERS)
    if not pk.wait_for(lambda: _orders_root_ready(ctx), timeout=60, poll=2):
        raise pk.PageError("orders page (div#orders) not found")
    time.sleep(3)
    pk.close_popups(d)
    pk.js_click(d, pk.shadow(d, ORDERS_HOST, "#tab-5"))          # Todas
    time.sleep(4)
    pk.close_popups(d)
    box = pk.need(d, "input[placeholder^='Buscar por n']", what="order search box")
    pk.real_click(d, box)                                        # opens the filter panel
    tab = pk.wait_for(lambda: pk.shadow(d, ORDERS_HOST, "#omnibar-search-tabs-tab-1"),
                      timeout=10)
    if not tab:
        raise pk.PageError("order filter panel tab 'Fecha de orden' not found")
    pk.js_click(d, tab)
    time.sleep(1)
    start = pk.shadow(d, ORDERS_HOST,
                      "#omnibar-search-tabs-tabpane-1 input[date-range=start]")
    if not start:
        raise pk.PageError("order date range input not found")
    pk.real_click(d, start)
    if not pk.wait_for(lambda: pk.find(d, "td.ant-picker-cell-in-view[title]"), timeout=10):
        raise pk.PageError("order calendar did not open")
    time.sleep(0.8)
    prev_sel, next_sel = ".ant-picker-header-prev-btn", ".ant-picker-header-next-btn"
    pk.real_click(d, _walk_to(ctx, a, prev_sel, next_sel))
    time.sleep(0.8)
    end_cell = _walk_to(ctx, b, prev_sel, next_sel)
    if "disabled" in (end_cell.get_attribute("class") or ""):
        raise pk.PageError("end day %s is disabled in the order calendar (range limit "
                           "changed? site orders_max_days=%s)" % (b, ctx.site["orders_max_days"]))
    pk.real_click(d, end_cell)
    omnibar = lambda: d.execute_script(
        "var h=document.querySelector(arguments[0]);var o=h&&h.shadowRoot&&"
        "h.shadowRoot.querySelector('.omnibar');return o?o.innerText:'';", ORDERS_HOST)
    if not pk.wait_for(lambda: chip_matches(omnibar(), a, b), timeout=15):
        raise pk.PageError("order date chip does not show %s..%s (got %r)"
                           % (a, b, (omnibar() or "")[-80:]))
    # wait for the list: rows, or the empty state
    listing = lambda: d.execute_script(
        "var h=document.querySelector(arguments[0]);var r=h&&h.shadowRoot;"
        "if(!r)return '';var t=r.querySelector('.rdt_Table');"
        "return (t?t.innerText:'')+'|'+((r.textContent||'').indexOf('No tienes órdenes')>=0?'EMPTY':'');",
        ORDERS_HOST)
    time.sleep(3)
    pk.wait_for(lambda: re.search(r"\d{9,}|EMPTY", listing() or ""), timeout=40, poll=2)
    export = d.execute_script(
        "var h=document.querySelector(arguments[0]);var r=h.shadowRoot;"
        "return [].filter.call(r.querySelectorAll('button.dropdown-toggle'),"
        "function(x){return x.innerText.indexOf('Exportar')>=0})[0]||null;", ORDERS_HOST)
    if export is None:
        raise pk.PageError("Exportar button not found")
    if export.get_attribute("disabled") is not None:
        if "EMPTY" in (listing() or ""):
            ctx.nothing_there(form, "这段时间没有订单（No tienes órdenes），Exportar 不可用")
            return
        ctx.give_up(form, "ui-disabled", step="Exportar 不可用，但列表里有订单")
        return
    pk.js_click(d, export)
    time.sleep(1.2)
    item = pk.wait_for(lambda: d.execute_script(
        "var r=document.querySelector(arguments[0]).shadowRoot;"
        "return [].filter.call(r.querySelectorAll('.dropdown-menu.show a'),"
        "function(x){return x.innerText.trim()==='Filtrado'})[0]||null;", ORDERS_HOST),
        timeout=8)
    if not item:
        raise pk.PageError("Exportar menu item 'Filtrado' not found")
    pk.js_click(d, item)
    toggle = pk.wait_for(lambda: pk.shadow(d, ORDERS_HOST,
                                           ".modal.show .modal-footer .dropdown-toggle"),
                         timeout=10)
    if not toggle:
        raise pk.PageError("export dialog 'Descargar' not found")
    pk.js_click(d, toggle)
    xlsx = pk.wait_for(lambda: d.execute_script(
        "var r=document.querySelector(arguments[0]).shadowRoot;"
        "return [].filter.call(r.querySelectorAll('.modal.show .dropdown-menu.show a,"
        ".modal.show .dropdown-menu.show button'),function(x){return x.innerText.trim()"
        "==='XLSX'})[0]||null;", ORDERS_HOST), timeout=8)
    if not xlsx:
        raise pk.PageError("export format 'XLSX' not found")
    before = pk.snapshot(ctx.out_dir)
    t0 = time.time()
    pk.js_click(d, xlsx)
    got = pk.wait_new_files(ctx.out_dir, before, timeout=300, settle=3)
    land_order_export(ctx, form, a, b, before, got, t0)


def land_order_export(ctx, form, a, b, before, got, t0):
    """Name the export of a..b Ordenes_<period>.xlsx, or give the chunk up."""
    got, late = own_new_files(before, got)
    if late:
        ctx.log("  [提醒] %s 是上一段订单晚到的文件，不算这一段的" % "、".join(late))
        if not got:
            # 这一段自己的文件可能还没开始下：把晚到的算进"之前"，再等一次
            before = set(before) | set(late)
            got, _ = own_new_files(before, pk.wait_new_files(ctx.out_dir, before,
                                                             timeout=300, settle=3))
    if not got:
        ctx.give_up(form, "download-not-landed", step="订单导出 %s..%s" % (a, b))
        return
    if len(got) != 1:
        # 分不清哪个是这一段的：宁可这一段算没下到，也不能把别的时段的订单
        # 当成 Ordenes_<这一段>
        ctx.give_up(form, "file-rejected", files=got,
                    step="订单导出 %s..%s 得到 %d 个新文件（%s），分不清哪个是这一段的"
                         % (a, b, len(got), "、".join(got)))
        return
    name = ctx.rename(got[0], "Ordenes_%s.xlsx" % periods.period_id(a, b))
    ctx.log("  订单 %s..%s -> %s（%d 秒）" % (a, b, name, time.time() - t0))
    ctx.got(form, [name])


def own_new_files(before, got):
    """Split wait_new_files' answer into (ours, late). A file whose name is a
    partial download (x.crdownload) that was already in the folder before the
    click is the previous export finishing late, not this one."""
    half = {f[:-len(p)] for f in before for p in pk.PARTIAL if f.endswith(p)}
    return ([f for f in got if f not in half and not f.endswith(pk.PARTIAL)],
            [f for f in got if f in half])


def retried(ctx, what, fn, *a, attempts=2):
    """browser_kit.retried for Falabella routes: a PageError is tried again on a
    freshly loaded page (2026-10-19 on production: the order page still showed
    the previous run's range and kept it); a wrong language never is."""
    return browser_kit.retried(ctx, what, fn, *a, attempts=attempts,
                               retry_on=(pk.PageError,), never_retry=(LanguageWrong,))


def _orders_chunk_retried(ctx, form, a, b):
    return retried(ctx, "订单 %s..%s" % (a, b), orders_chunk, form, a, b)


def orders(ctx):
    for a, b in periods.order_chunks(ctx.month, ctx.site["orders_max_days"]):
        form = "order/%s" % periods.period_id(a, b)
        guarded(ctx, form, [form], _orders_chunk_retried, form, a, b)


# ================================================================ 2. billing (Transacciones)

def _norm(s):
    return re.sub(r"\s+", " ", s or "").strip()


def _payments_ready(ctx):
    ctx.go(PAYMENTS)
    if not pk.wait_for(lambda: pk.find(ctx.d, "[data-testid=export-btn-id]"),
                       timeout=60, poll=2):
        raise pk.PageError("payments page 'Descargar reporte' not found")
    time.sleep(2)
    pk.close_popups(ctx.d)


def pick_payments_range(ctx, filter_label, a, b):
    d = ctx.d
    pk.close_popups(d)
    sel = pk.need(d, ".ant-select-selector", what="payments filter select")
    pk.real_click(d, sel)
    opt = pk.need(d, ".ant-select-item-option", filter_label, timeout=10,
                  what="filter option %r" % filter_label)
    pk.real_click(d, opt)
    time.sleep(1.5)
    inp = pk.wait_for(lambda: [e for e in pk.find(d, "input[readonly]")
                               if "per" in (e.get_attribute("placeholder") or "").lower()],
                      timeout=10)
    if not inp:
        raise pk.PageError("period input not found")
    inp = inp[0]
    pk.close_popups(d)
    pk.real_click(d, inp)
    if not pk.wait_for(lambda: pk.find(d, "[data-testid=range-picker-id]"), timeout=10):
        raise pk.PageError("range picker (range-picker-id) not found")

    def cell(day):
        c = [e for e in pk.find(d, "[data-testid=range-picker-id] "
                                   "td.ant-picker-cell[title='%s']" % periods.iso(day))
             if "in-view" in (e.get_attribute("class") or "")]
        return c[0] if c else None

    def walk(day):
        for _ in range(18):
            c = cell(day)
            if c is not None:
                return c
            first = pk.find(d, "[data-testid=range-picker-id] td.ant-picker-cell-in-view[title]")
            shown = first[0].get_attribute("title") if first else ""
            arrow = ".anticon-left" if periods.iso(day) < shown else ".anticon-right"
            btn = pk.find(d, "[data-testid=range-picker-id] " + arrow)
            if not btn:
                raise pk.PageError("range picker arrow %s not found" % arrow)
            pk.js_click(d, btn[0])
            time.sleep(0.6)
        raise pk.PageError("range picker day %s not found" % day)

    apply_ = lambda: pk.find(d, "[data-testid=range-picker-id] button", "Aplicar")
    enabled = lambda: apply_() and apply_()[0].is_enabled()
    pk.close_popups(d)
    pk.real_click(d, walk(a))
    time.sleep(0.6)
    pk.real_click(d, walk(b))
    time.sleep(0.6)
    if not enabled():
        # The 1st of the current month is pre-marked; a click on an already
        # marked day does not start a range. Clicking the start again does.
        pk.real_click(d, walk(a))
        time.sleep(0.6)
    if not enabled():
        raise pk.PageError("range picker 'Aplicar' stays disabled after choosing %s..%s"
                           % (a, b))
    pk.real_click(d, apply_()[0])
    time.sleep(1.5)
    shown = _norm(inp.get_attribute("value"))
    want = "%s - %s" % (periods.es_date(a), periods.es_date(b))
    if shown != want:
        raise pk.PageError("period input shows %r, wanted %r" % (shown, want))
    return shown


def request_payments_report(ctx, menu_key, prefix):
    """Search, then Descargar reporte -> menu item -> XLSX. Returns the newest
    downloads row before the request (to recognise the new one)."""
    d = ctx.d
    pk.close_popups(d)
    btn = pk.need(d, "button.search-btn", what="payments search button")
    pk.js_click(d, btn)
    time.sleep(6)
    pk.close_popups(d)
    # 表格还没画出来时读到的是"没有行"，之后任何一行旧的 Listo 都会被当成新的
    pk.wait_download_rows(d)
    before = ([r for r in pk.download_rows(d) if r.startswith(prefix)] or [None])[0]
    pk.js_click(d, pk.need(d, "[data-testid=export-btn-id]", what="Descargar reporte"))
    item = pk.wait_for(lambda: [e for e in pk.find(d, "li[role=menuitem]")
                                if (e.get_attribute("data-menu-id") or "").endswith(menu_key)
                                and e.get_attribute("aria-disabled") == "false"], timeout=15)
    if not item:
        raise pk.PageError("download menu item %r not found or disabled" % menu_key)
    pk.js_click(d, item[0])
    fmt = pk.need(d, "[data-testid=select-options-id] .ant-select-selector", timeout=10,
                  what="file format select")
    pk.real_click(d, fmt)
    pk.real_click(d, pk.need(d, ".ant-select-item-option", "XLSX", timeout=8,
                             what="format XLSX"))
    time.sleep(0.8)
    go = pk.need(d, "[data-testid=fileType-export-btn-id]", what="dialog Descargar reporte")
    pk.js_click(d, go)
    time.sleep(3)
    return before


def new_tx_rows(rows, before, since):
    """The Transacciones rows that can be ours: listed above the row that was
    newest before the request (by file name — an older request that was still
    Procesando then and is Listo now is not ours), and not made well before the
    request (the file name carries its time; FORMS.md). Newest first."""
    rows = pk.rows_above(rows, pk.file_name_of_row(before) if before else None)
    return [r for r in rows if pk.row_not_older_than(r, since)]


def download_row(ctx, row_text, timeout=90):
    """Click the download button of the downloads row with this file name."""
    name = pk.file_name_of_row(row_text)
    row = pk.row_containing(ctx.d, name)
    btn = pk.inside(ctx.d, row, "[data-testid=settlements-download]") if row else None
    if not btn:
        raise pk.PageError("download button of row %r not found" % name)
    before = pk.snapshot(ctx.out_dir)
    pk.js_click(ctx.d, btn)
    return pk.wait_new_files(ctx.out_dir, before, timeout=timeout, settle=2)


def billing(ctx):
    form = "billing"
    a, b = periods.month_bounds(ctx.month)
    _payments_ready(ctx)
    shown = pick_payments_range(ctx, "Fecha transacción", a, b)
    ctx.log("  Transacciones：%s" % shown)
    asked = time.time()
    before = request_payments_report(ctx, "alltransactionsV2", "NewReportTransaction")
    requested = time.time()
    row, state = pk.wait_new_ready_row(ctx.d, before, timeout=ctx.collect_timeout,
                                       prefix="NewReportTransaction",
                                       select=lambda rows: new_tx_rows(rows, before, asked))
    if state != "ready":
        ctx.log("  Transacciones 还没生成好（%s），记为欠账" % state)
        ctx.pending.append({"form": form, "route": "billing", "row": row,
                            "seen_before": before,
                            "range": ctx.month, "requested_at": requested,
                            "last_state": "generating" if row else "absent",
                            "probes": 1})
        return
    files = download_row(ctx, row)
    if not files:
        ctx.give_up(form, "download-not-landed", step="Transacciones %s" % row[:60])
        return
    ctx.log("  Transacciones -> %s（%d 秒）" % (files, time.time() - requested))
    ctx.got(form, files)


# ================================================================ 3. stock

_EXPORTS = pk.WALK + """var o=[];W(document,function(e){if(e.tagName==='TR'&&/fby-exports/.test(
 (e.querySelector('a[href]')||{}).href||'')){var a=e.querySelector('a[href]');
 o.push([(e.innerText||'').replace(/\\s+/g,' ').trim(),a.href]);}});return o;"""


def stock(ctx):
    form = "stock"
    d = ctx.d
    ctx.go(DASHBOARD, settle=6)
    pk.close_popups(d)
    ctx.go(FULFILLMENT)
    header = pk.wait_for(lambda: pk.find(d, "[data-testid=test_dropdown-header]", "Descargar"),
                         timeout=60, poll=2)
    if not header:
        raise pk.PageError("Fulfillment 'Descargar' dropdown not found")
    exports = lambda: d.execute_script(_EXPORTS) or []
    pk.wait_for(exports, timeout=10)      # read "before" from the rendered list, not an empty one
    before = (exports() or [[None, None]])[0][1]
    item = None
    for _ in range(3):              # the first click is lost while the page is still loading
        pk.real_click(d, header[0])
        got = pk.wait_for(lambda: pk.find(d, "[data-testid=test_list] a", "Todos los productos"),
                          timeout=6)
        if got:
            item = got[0]
            break
        time.sleep(3)
        header = pk.find(d, "[data-testid=test_dropdown-header]", "Descargar") or header
    if item is None:
        raise pk.PageError("'Todos los productos' not found after opening 'Descargar' 3 times")
    requested = time.time()
    pk.real_click(d, item)
    newest = None
    end = time.time() + ctx.collect_timeout
    while time.time() < end:
        time.sleep(5)
        rows = exports()
        if rows and rows[0][1] != before and "Finalizado" in rows[0][0]:
            newest = rows[0]
            break
        if int(time.time() - requested) % 30 < 5:
            # not d.refresh(): /fulfillment/product loaded directly is an error page
            ctx.go(DASHBOARD, settle=5)
            ctx.go(FULFILLMENT, settle=10)
    if not newest:
        ctx.pending.append({"form": form, "route": "stock", "row": None,
                            "seen_before": before,
                            "range": ctx.month, "requested_at": requested,
                            "last_state": "generating", "probes": 1})
        ctx.log("  库存导出还没完成，记为欠账")
        return
    try:
        files = fetch_stock_link(ctx, newest[1])
    except OSError as e:          # network: DNS, timeout, connection (urllib.error.URLError)
        ctx.give_up(form, "download-not-landed",
                    step="Products Export %s：浏览器没下载，直接下载也失败（%s）"
                         % (newest[0][:40], str(e)[:80]))
        return
    if not files:
        ctx.give_up(form, "download-not-landed", step="Products Export %s" % newest[0][:50])
        return
    ctx.flags["stock_snapshot_at"] = datetime.datetime.now().isoformat(timespec="seconds")
    ctx.log("  库存 -> %s" % files)
    ctx.got(form, files)


STOCK_HOSTS = ("https://storage.googleapis.com/fullfillment-by-storage-production/",)


def fetch_stock_link(ctx, href, attempts=2):
    """Download a Fulfillment export.

    The link (target=_blank) points at a public Google Cloud Storage object.
    Clicking it works in a fresh browser profile, but after a few downloads from
    storage.googleapis.com Chrome's "multiple automatic downloads" guard starts to
    swallow the click silently (2026-10-16: no tab, no file, while the object
    answered HTTP 200). Chrome offers no DevTools switch for that guard, so after
    the browser attempts the same URL is fetched directly. Only this bucket is
    accepted, and run_meta records that the fallback was used.
    """
    d = ctx.d
    home = d.current_window_handle
    for i in range(attempts):
        links = [a for a in pk.find(d, "a[href]", hidden=True)
                 if a.get_attribute("href") == href]
        if not links:
            break
        tabs = set(d.window_handles)
        before = pk.snapshot(ctx.out_dir)
        (pk.real_click if i == 0 else pk.js_click)(d, links[0])
        got = pk.wait_new_files(ctx.out_dir, before, timeout=30, settle=2, start_timeout=15)
        for h in d.window_handles:
            if h not in tabs:
                try:
                    d.switch_to.window(h)
                    d.close()
                except Exception:
                    pass
        d.switch_to.window(home)
        if got:
            return got
        time.sleep(3)
    # The browser download sometimes lands late (seen 2026-10-16): give it a moment.
    name = os.path.basename(href.split("?")[0])
    late = pk.wait_for(lambda: os.path.exists(os.path.join(ctx.out_dir, name)) and
                       os.path.getsize(os.path.join(ctx.out_dir, name)) > 0, timeout=20)
    if late:
        ctx.log("  库存文件晚到了，浏览器已下好 %s" % name)
        return [name]
    return fetch_direct(ctx, href)


def fetch_direct(ctx, href, timeout=60, waits=(0, 10, 30)):
    """Fetch the export over HTTPS. Retried: on 2026-10-19 the production machine
    could not resolve storage.googleapis.com for a moment (getaddrinfo failed) and
    the same fetch worked minutes later. Raises OSError when every attempt fails."""
    import urllib.request
    if not href.startswith(STOCK_HOSTS):
        raise pk.PageError("refusing direct download outside the export bucket: %s" % href[:80])
    name = os.path.basename(href.split("?")[0])
    path = os.path.join(ctx.out_dir, name)
    req = urllib.request.Request(href, headers={"User-Agent": "Mozilla/5.0"})
    for i, pause in enumerate(waits):
        time.sleep(pause)
        try:
            with urllib.request.urlopen(req, timeout=timeout) as r, \
                    open(path + ".part", "wb") as fh:
                fh.write(r.read())
            break
        except OSError as e:
            ctx.log("  直接下载库存文件第 %d 次失败：%s" % (i + 1, str(e)[:80]))
            if i == len(waits) - 1:
                raise
    os.replace(path + ".part", path)
    if not os.path.getsize(path):
        os.remove(path)
        return []
    ctx.flags["stock_fetched_direct"] = True
    ctx.log("  浏览器没有下载库存文件，已直接从导出地址取回 %s" % name)
    return [name]


# ================================================================ 4+6. billing1 + invoice1

def _settlement_page(ctx, year):
    ctx.go(SETTLEMENT % year)
    if not pk.wait_for(lambda: pk.find(ctx.d, "[data-testid=settlement-invoice-table]"),
                       timeout=60, poll=2):
        raise pk.PageError("Factura falabella.com table not found")
    time.sleep(3)


def billing1(ctx):
    for a, b in periods.billing_periods(ctx.month, ctx.site["billing_cycle_day"]):
        pid = periods.period_id(a, b)
        f_xlsx, f_pdf = "billing1/%s" % pid, "invoice1/%s" % pid
        if periods.billing_period_state(a, b, ctx.today) == "open":
            why = ("账期 %s 到 %s 还没结束，%s 之后平台才出" % (a, b, b + datetime.timedelta(days=1)))
            for f in (f_xlsx, f_pdf):
                ctx.give_up(f, "not-offered", step=why, shot=False)
            ctx.flags.setdefault("billing_open_periods", []).append(pid)
            continue
        guarded(ctx, f_xlsx, [f_xlsx], billing1_period, f_xlsx, a, b)
        guarded(ctx, f_pdf, [f_pdf], invoice1_period, f_pdf, a, b)


def billing1_period(ctx, form, a, b):
    d = ctx.d
    _settlement_page(ctx, b.year)
    label = periods.billing_row_label(a, b)
    row = pk.wait_for(lambda: pk.row_starting(d, label), timeout=15)
    if not row:
        ctx.give_up(form, "ui-missing", step="账期行 %r 不在列表里" % label)
        return
    before = ([r for r in pk.download_rows(d) if r.startswith("InvoiceReport")] or [None])[0]
    trig = pk.inside(d, row, "[class*=dropdown-trigger]")
    if not trig:
        raise pk.PageError("'Descargar reporte' of period %r not found" % label)
    pk.js_click(d, trig)
    item = pk.need(d, "li[role=menuitem]", "XLSX", timeout=8, what="menu item XLSX")
    pk.js_click(d, item)
    rowtext, state = pk.wait_new_ready_row(d, before, timeout=90, poll=5,
                                           prefix="InvoiceReport")
    if state != "ready":
        ctx.give_up(form, "not-ready", row_state="generating" if rowtext else "absent",
                    status_text=rowtext or "", step=label)
        return
    if "%s-%s" % (periods.iso(a), periods.iso(b)) not in rowtext:
        raise pk.PageError("newest invoice report %r is not for %s..%s" % (rowtext[:80], a, b))
    files = download_row(ctx, rowtext)
    if not files:
        ctx.give_up(form, "download-not-landed", step=label)
        return
    ctx.log("  Factura falabella.com %s -> %s" % (label, files))
    ctx.got(form, files)


def _menu_pdfs(ctx, trigger_text="PDF", wait=15):
    """Open a 'Descargar … (PDF)' menu and download every item. Returns files."""
    d = ctx.d
    out = []
    trig = pk.wait_for(lambda: [t for t in pk.find(d, "[class*=dropdown-trigger]")
                                if trigger_text in (t.text or "")], timeout=wait)
    if not trig:
        return None
    pk.real_click(d, trig[-1])
    time.sleep(1.5)
    names = [i.text.strip() for i in pk.find(d, "li[role=menuitem]") if i.text.strip()]
    for name in names:
        items = [i for i in pk.find(d, "li[role=menuitem]") if i.text.strip() == name]
        if not items:
            pk.real_click(d, trig[-1])
            time.sleep(1.5)
            items = [i for i in pk.find(d, "li[role=menuitem]") if i.text.strip() == name]
        if not items:
            raise pk.PageError("PDF menu item %r disappeared" % name)
        before = pk.snapshot(ctx.out_dir)
        pk.real_click(d, items[0])
        got = pk.wait_new_files(ctx.out_dir, before, timeout=60, settle=2)
        if not got:
            raise pk.PageError("PDF %r did not download" % name)
        out.extend(got)
    return out


INVOICE_PDF_SUFFIX = {"factura": "FACTURA", "nota de credito": "NOTA_DE_CREDITO",
                      "nota de crédito": "NOTA_DE_CREDITO"}


def missing_invoice_pdfs(rows, files):
    """What the detail page should have given but did not.

    rows: [(Numero de documento, Tipo de documento)] of the period's Factura
    falabella.com report. Every document type that has lines (Factura, Nota de
    Crédito) must come with at least one PDF of that type; where the report fills in
    the document number (not every seller's does — DAB's column is empty), that exact
    PDF must be there too. Promotion debit notes (usually 0) have no PDF on the page.
    Returns sorted labels like "FACTURA" or "520878_FACTURA.pdf".
    """
    have = set(files or [])
    kinds, named = set(), set()
    for num, kind in rows:
        suf = INVOICE_PDF_SUFFIX.get(str(kind or "").strip().lower())
        if not suf:
            continue
        kinds.add(suf)
        num = str(num or "").strip().split(".")[0]
        if num and num.lower() != "nan":
            named.add("%s_%s.pdf" % (num, suf))
    out = [n for n in named if n not in have]
    out += [k for k in kinds if not any(f.endswith("_%s.pdf" % k) for f in have)]
    return sorted(out)


def _invoice_rows(ctx, form_xlsx):
    """(document number, type) of every line of the period's invoice report."""
    from openpyxl import load_workbook
    rows = []
    for name in ctx.files.get(form_xlsx) or []:
        try:
            # not read_only: Falabella's file declares a 1x1 sheet size, and read-only
            # mode then sees only the first column
            wb = load_workbook(os.path.join(ctx.out_dir, name), data_only=True)
            it = wb.worksheets[0].iter_rows(values_only=True)
            head = [str(h or "").strip() for h in next(it)]
            i, j = head.index("Numero de documento"), head.index("Tipo de documento")
            rows += [(r[i], r[j]) for r in it if r and len(r) > j]
            wb.close()
        except Exception as e:
            ctx.log("  [提醒] 读不了 %s 来核对发票 PDF：%s" % (name, str(e)[:80]))
    return rows


def invoice1_period(ctx, form, a, b):
    """Every PDF the period's invoice report names must arrive; a missing one is
    tried once more from a fresh detail page, then recorded by name."""
    rows = _invoice_rows(ctx, "billing1/%s" % periods.period_id(a, b))
    files = _invoice1_pdfs(ctx, form, a, b) or []
    missing = missing_invoice_pdfs(rows, files)
    if missing:
        ctx.log("  发票 PDF 缺 %s，重新打开明细页再下一次" % "、".join(missing))
        files += [f for f in (_invoice1_pdfs(ctx, form, a, b) or []) if f not in files]
        missing = missing_invoice_pdfs(rows, files)
    for what in missing:
        ctx.give_up("%s/%s" % (form, what.replace(".pdf", "")), "download-not-landed",
                    step="发票明细里有这类单据，但 PDF（%s）没下到" % what)
    if files:
        ctx.got(form, files)
    elif not rows:
        ctx.give_up(form, "ui-missing", step="明细页上没有 PDF 下载入口")


def _invoice1_pdfs(ctx, form, a, b):
    d = ctx.d
    _settlement_page(ctx, b.year)
    label = periods.billing_row_label(a, b)
    row = pk.wait_for(lambda: pk.row_starting(d, label), timeout=15)
    if not row:
        ctx.give_up(form, "ui-missing", step="账期行 %r 不在列表里" % label)
        return
    btns = [x for x in (pk.inside(d, row, "button", all_=True) or [])
            if (x.text or "").strip() == "Revisar detalle"]
    if not btns:
        raise pk.PageError("'Revisar detalle' of period %r not found" % label)
    pk.js_click(d, btns[0])
    if not pk.wait_for(lambda: pk.find(d, "[role=tab]", "Factura"), timeout=30):
        raise pk.PageError("invoice detail tabs not found")
    time.sleep(2)
    files = []
    for tab in ("Factura", "Notas de crédito"):
        t = pk.find(d, "[role=tab]", tab)
        if not t:
            continue
        pk.js_click(d, t[0])
        time.sleep(2.5)
        got = _menu_pdfs(ctx)
        if got:
            files.extend(got)
    ctx.log("  发票 PDF %s -> %s" % (label, files))
    return files


# ================================================================ 5+7. billing2 + invoice2

def _liquidation_page(ctx, year):
    ctx.go(LIQUIDATION % year)
    time.sleep(4)
    tab = pk.wait_for(lambda: pk.find(ctx.d, "#rc-tabs-0-tab-liquidation-invoice")
                      or pk.find(ctx.d, "[role=tab]", "Liquidación factura"), timeout=60, poll=2)
    if not tab:
        raise pk.PageError("tab 'Liquidación factura' not found")
    pk.js_click(ctx.d, tab[0])
    if not pk.wait_for(lambda: pk.find(ctx.d, "[data-testid=settlement-creditnotes-table]"),
                       timeout=30):
        raise pk.PageError("Liquidación factura table not found")
    time.sleep(3)


def billing2(ctx):
    guarded(ctx, "billing2", ["billing2"], billing2_month)
    guarded(ctx, "invoice2", ["invoice2"], invoice2_month)


def billing2_month(ctx):
    form = "billing2"
    d = ctx.d
    y, _m = periods.parse_month(ctx.month)
    _liquidation_page(ctx, y)
    label = periods.liquidation_row_label(ctx.month)
    row = pk.wait_for(lambda: pk.row_starting(d, label), timeout=15)
    if not row:
        ctx.nothing_there(form, "列表里没有 %s 的 Liquidación（这个月没有平台代开的销售单据）"
                          % label)
        return
    before = ([r for r in pk.download_rows(d) if "FSC_LIQ" in r] or [None])[0]
    btn = pk.inside(d, row, "[data-testid=liquidation-invoice-download-report]")
    if not btn:
        raise pk.PageError("liquidation 'Descargar reporte' not found")
    pk.js_click(d, btn)
    rowtext, state = pk.wait_new_ready_row(d, before, timeout=90, poll=5)
    if state != "ready" or "FSC_LIQ" not in (rowtext or ""):
        ctx.give_up(form, "not-ready", row_state="generating" if rowtext else "absent",
                    status_text=rowtext or "", step=label)
        return
    files = download_row(ctx, rowtext)
    if not files:
        ctx.give_up(form, "download-not-landed", step=label)
        return
    ctx.log("  Liquidación factura %s -> %s" % (label, files))
    ctx.got(form, files)


def invoice2_month(ctx):
    form = "invoice2"
    d = ctx.d
    y, mo = periods.parse_month(ctx.month)
    _liquidation_page(ctx, y)
    label = periods.liquidation_row_label(ctx.month)
    row = pk.wait_for(lambda: pk.row_starting(d, label), timeout=15)
    if not row:
        if "billing2" in ctx.no_data:
            ctx.nothing_there(form, ctx.no_data["billing2"])
        else:
            ctx.give_up(form, "ui-missing", step="月份行 %r 不在列表里" % label)
        return
    btn = pk.inside(d, row, "[data-testid=liquidation-invoice-check-details-btn]")
    if not btn:
        raise pk.PageError("liquidation 'Revisar detalle' not found")
    pk.js_click(d, btn)
    if not pk.wait_for(lambda: pk.find(d, "[class*=dropdown-trigger]"), timeout=30):
        raise pk.PageError("liquidation detail has no Descargar")
    time.sleep(2)
    docs = pk.table_rows(d, r"CLP")
    files = []
    for i, doc in enumerate(docs):
        trig = pk.find(d, "[class*=dropdown-trigger]")
        if i >= len(trig):
            break
        pk.js_click(d, trig[i])
        item = pk.need(d, "li[role=menuitem]", "PDF", timeout=8, what="menu item PDF")
        before = pk.snapshot(ctx.out_dir)
        pk.js_click(d, item)
        got = pk.wait_new_files(ctx.out_dir, before, timeout=60, settle=2)
        if not got:
            raise pk.PageError("liquidation PDF of %r did not download" % doc[:40])
        folio = (doc.split(" ") or ["doc"])[0]
        files.append(ctx.rename(got[0], "Liquidacion_%s_%04d-%02d.pdf" % (folio, y, mo)))
    if not files:
        ctx.give_up(form, "ui-missing", step="明细页上没有单据")
        return
    ctx.log("  Liquidación PDF -> %s" % files)
    ctx.got(form, files)


# ================================================================ owed (collect later)

def collect_owed(ctx, entry):
    """Try once more to fetch an owed async report. Returns files, or [] if still not ready."""
    if entry["route"] == "billing":
        _payments_ready(ctx)
        pk.wait_download_rows(ctx.d)
        rows = [r for r in pk.download_rows(ctx.d) if r.startswith("NewReportTransaction")]
        match = owed_tx_row(entry, rows)
        if match and "Listo" in match:
            return download_row(ctx, match)
        return []
    if entry["route"] == "stock":
        ctx.go(DASHBOARD, settle=6)
        ctx.go(FULFILLMENT, settle=12)
        rows = ctx.d.execute_script(_EXPORTS) or []
        row = owed_stock_row(entry, rows)
        if row and "Finalizado" in row[0]:
            return fetch_stock_link(ctx, row[1])
        return []
    return []


def owed_tx_row(entry, rows):
    """The downloads row of an owed Transacciones request, or None.

    Seen while waiting -> that file name. Never seen -> the oldest row that
    came after the request (new_tx_rows; a later run's request may sit above
    it). Entries written before "seen_before" existed: the newest row, as
    before."""
    name = pk.file_name_of_row(entry.get("row") or "")
    if name:
        return next((r for r in rows if r.startswith(name)), None)
    if "seen_before" not in entry:
        return rows[0] if rows else None
    mine = new_tx_rows(rows, entry["seen_before"], entry.get("requested_at"))
    return mine[-1] if mine else None


def owed_stock_row(entry, rows):
    """The Fulfillment export of an owed stock request, or None: the newest
    export listed above the one that was newest before the request (an export
    older than the request is a stale snapshot). Old entries: the newest."""
    if "seen_before" in entry:
        rows = pk.rows_above(rows, entry["seen_before"], key=lambda r: r[1])
    return rows[0] if rows else None


ROUTES = (("order", orders), ("billing", billing), ("stock", stock),
          ("billing1", billing1), ("billing2", billing2))
