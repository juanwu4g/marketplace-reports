# -*- coding: utf-8 -*-
"""
Amazon Mexico monthly financial workbook — a sister of MercadoLibre's and
Falabella's.

Same ten sheets, same look, same check levels: the layout is written by
core/report_family.py on the frame of core/report_kit.py (docs/REPORT_FAMILY.md);
the content — lines, bridges, columns, rules — is Amazon's. Accounting rules are in
platforms/amazon/README.md ("Accounting rules"); the short version:

* **Basis: Amazon's posting month.** Every money line comes from Amazon's monthly
  transaction report ("Transacciones", option Mensual), which is Amazon's month in
  GMT-6 indexed by posting date. Amazon posts an FBA order when it ships, which is
  close to delivery (IFRS 15). The file IS the month: its rows are never refiltered
  by the printed time (printed in GMT-7).
* **Net of IVA.** Sales, shipping and promotions are printed net of IVA, their IVA in
  separate columns (output IVA, a memo line). Amazon's fee columns are
  IVA-inclusive: the P&L shows fee / 1.16 and the rest is input IVA (memo).
  Advertising comes already split by Amazon (net + IVA) and is not divided again.
* **Withholdings are tax prepayments, not expenses** (as in MercadoLibre Mexico):
  the IVA Amazon withholds (8 %), and the ISR (2.5 %) taken around the 11th for the
  previous month, are memo lines (备查) under the P&L and steps of the cash bridge.
  The ISR of THIS month comes from the withholding certificate (CFDI de
  retenciones).
* **Every transaction kind is registered** (tx_catalog.py). An unregistered one is an
  error check and its amount goes to a visible "未登记科目 ⚠" row.
* **Product cost** is not in Amazon's data: a yellow input row in ② and a unit cost
  column in ⑦.

Inputs: one download folder per store (platforms/amazon/run_downloads.py):
<YYYY><Mon>MonthlyTransaction.csv, <YYYY><Mon>MonthlySummary.pdf,
fee_<uuid>.xml/.pdf + fee_invoices.json, stock_<YYYYMMDD>.csv,
retenciones_<uuid>.xml/.pdf, run_meta.json. Buyer city / state / postcode of the
transaction report are dropped on load and never reach the workbook.
"""
import datetime
import glob
import io
import json
import os
import re
import shutil
import subprocess
import xml.etree.ElementTree as ET

import pandas as pd
from openpyxl import Workbook

from core import report_family as RF
from core import report_kit as K
from core.report_kit import F, fill, BOX
from platforms.amazon import periods
from platforms.amazon import site as platform_site
from platforms.amazon import tx_catalog as TC

SHEETS = K.SHEET_NAMES
COUNTRY = "MX"


def _site_fact(key, default, ok=lambda v: True):
    """A fact of platforms/amazon/site.py, or the default while the site entry is
    still a scaffold placeholder."""
    try:
        v = platform_site.get(COUNTRY).get(key)
    except Exception:
        return default
    return v if v is not None and ok(v) else default


IVA_RATE = float(_site_fact("vat", 0.16, lambda v: isinstance(v, (int, float)) and 0 < v < 1))
CUR = _site_fact("currency", "MXN", lambda v: isinstance(v, str) and v.isalpha() and len(v) == 3)
DECIMALS = int(_site_fact("decimals", 2, lambda v: isinstance(v, int)))
MNY = K.money_format(DECIMALS)
# Amazon's month for Mexico is GMT-6 (Mexico has had no summer time since 2022);
# the transaction report prints times in GMT-7/-8 and says so on every row.
MONTH_UTC_OFFSET = int(_site_fact("utc_offset", -6, lambda v: isinstance(v, int) and v != 0))

TOL_ROW = 0.02                  # MXN; a row's total vs the sum of its columns
TOL_BRIDGE = 0.05               # MXN; residual of the bridges built from one file
TOL_SUMMARY = 1.0               # MXN; summary PDF vs the transaction report
IVA_RATE_BAND = (0.159, 0.161)  # output IVA / product sales on order lines
WITHHOLD_BAND = (0.48, 0.52)    # IVA withheld / output IVA on order lines (8 % of 16 %)

# transaction report columns (Spanish UI)
T_DATE, T_TYPE, T_ORDER, T_SKU, T_DESC, T_QTY = ("fecha/hora", "tipo", "Id. del pedido", "sku",
                                                 "descripción", "cantidad")
T_STATE, T_SETTLE, T_FULFIL = "Estado de la transacción", "Id. de liquidación", "cumplimiento"
HEADER_ES = T_DATE
HEADER_EN = ("date/time",)      # what the same report starts with in an English UI
REQUIRED_COLS = [T_DATE, T_TYPE, T_ORDER, T_SKU, T_DESC, T_QTY] + TC.MONEY_COLS + [TC.TOTAL_COL]
# buyer location columns: never kept
DROP_COLS = ["ciudad del pedido", "estado del pedido", "código postal del pedido"]
POSTED = ("lanzado",)           # Estado de la transacción of a posted (released) line

ORDER_TIPO, REFUND_TIPO = TC.norm("Pedido"), TC.norm("Reembolso")

# ② fixed rows (core/report_family.PLRow): key, 中文, 平台原文, kind, 说明, formula.
# kind: d 取数 / sum 小计 / f 计算 / in 录入. A formula lists the rows it adds
# ("-key" subtracts); Span(a, b) sums every row from a to b.
_R = RF.PLRow
PL_ROWS = [
    _R(TC.SALES, "商品销售收入", "ventas de productos · Pedido", "d",
       "本月 Amazon 入账（FBA 订单≈发货）的商品价款，不含 IVA"),
    _R(TC.REFUND, "减：退货退款", "ventas de productos · Reembolso", "d",
       "本月入账的退款冲回的商品价款，不含 IVA"),
    _R(TC.SHIPPING, "运费收入（净额）", "créditos de envío", "d",
       "买家支付的运费，已扣退款退还的部分；不含 IVA"),
    _R(TC.GIFTWRAP, "礼品包装收入", "créditos por envoltorio de regalo", "d", "不含 IVA"),
    _R(TC.PROMO, "减：促销折扣", "descuentos promocionales", "d",
       "卖家出资的优惠券、促销折扣，冲减收入（净额，含退款冲回）；不含 IVA"),
    _R("net_sales", "净销售收入", "", "f",
       "＝ 商品销售收入 − 退货退款 ＋ 运费 ＋ 礼品包装 − 促销折扣",
       (TC.SALES, TC.REFUND, TC.SHIPPING, TC.GIFTWRAP, TC.PROMO)),
    _R(TC.COMMISSION, "减：销售佣金", "tarifas de venta", "d",
       "按品类费率；退款时退还的佣金已冲减，Amazon 留下的退款管理费在本行；不含 IVA（含税金额 ÷ 1.16）"),
    _R(TC.FBA_FEE, "减：FBA 配送费", "tarifas fba", "d",
       "每件发货的拣货、包装、配送费（退款退还的已冲减）；不含 IVA"),
    _R(TC.STORAGE, "减：FBA 月度仓储费", "Tarifa por almacenamiento de inventario de FBA", "d",
       "上月库存的仓储费，每月 7 日前后扣；不含 IVA"),
    _R(TC.STORAGE_LT, "减：FBA 长期仓储费", "tarifa por almacenamiento prolongado de FBA", "d",
       "库龄超期库存的附加仓储费，每月 15–18 日前后扣；不含 IVA"),
    _R(TC.LOGISTICS, "减：FBA 入仓运输与移除费",
       "Cargos por transporte de envíos a Amazon / Solicitud de retiro … / transportista asociado", "d",
       "入仓运输、合作承运人运费、库存移除与弃置；不含 IVA"),
    _R(TC.ADS, "减：广告费", "Costo de la publicidad", "d",
       "店铺级费用，无法分摊到 SKU；用 Amazon 自己拆出的不含税金额（不再 ÷ 1.16）"),
    _R(TC.SUBSCRIPTION, "减：订阅费", "Suscripción", "d", "专业卖家月费；不含 IVA"),
    _R(TC.OTHER_FEES, "减：其他平台费用", "Tarifa reglamentaria / tarifas de otra transacción", "d",
       "目前为 0；出现时按 Amazon 的列拆分"),
    _R(TC.UNREGISTERED, "减：未登记科目 ⚠", "(tx_catalog 里没有)", "d",
       "字典里没有的交易类型或列。金额照常计入以保证勾稽（不拆 IVA），但需人工确认科目归属，见 ⑩ 页"),
    _R("fee_total", "平台费用小计", "", "sum", "从 Amazon 余额中直接扣除；不含 IVA",
       RF.Span(TC.COMMISSION, TC.UNREGISTERED)),
    _R(TC.OTHER_INCOME, "加：其他收入（FBA 库存赔付）", "Reembolso de inventario de Logística de Amazon", "d",
       "Amazon 对仓内丢失/损坏、退货未退回等库存的赔付（扣除冲回），是存货损失的补偿，不是销售收入，单列在费用之后；"
       "明细见【%s】" % SHEETS[5]),
    _R("contrib", "经营贡献毛利（未扣商品成本）", "", "f",
       "＝ 净销售收入 − 平台费用小计 ＋ 其他收入", ("net_sales", "fee_total", TC.OTHER_INCOME)),
    _R("cogs", "减：商品采购成本", "", "in",
       "⚠ 平台数据不含此项：在【%s】填单位成本，或直接在此填总额" % SHEETS[6]),
    _R("profit", "税前净利润", "", "f", "＝ 经营贡献毛利 − 商品采购成本", ("contrib", "-cogs")),
]
PL_LABELS = dict((x.key, x.label) for x in PL_ROWS)
FEE_ROWS = TC.FEE_LINES
FBA_ROWS = (TC.FBA_FEE, TC.STORAGE, TC.STORAGE_LT, TC.LOGISTICS)
# fee lines that Amazon's seller-fee CFDI covers (③ bridge B). Advertising is billed
# separately (Amazon Ads), so it is shown next to the comparison but not in it.
INVOICED_LINES = (TC.COMMISSION, TC.FBA_FEE, TC.STORAGE, TC.STORAGE_LT, TC.LOGISTICS,
                  TC.SUBSCRIPTION, TC.OTHER_FEES)
MEMO_LABELS = {TC.IVA_OUT: "销项 IVA", TC.IVA_WITHHELD: "Amazon 代扣 IVA",
               TC.ISR: "Amazon 代扣 ISR（上月）", TC.IVA_ADJ: "IVA 调整", TC.PAYOUT: "转账到银行"}

# The accounting policies the report applies (decided, not open questions). Shown on
# ① and explained in README.md.
ACCOUNTING_POLICIES = [
    ("收入确认", "按 Amazon 月度交易报表（Transacciones，月度，GMT-6 的自然月）的入账月份归集。Amazon 在 FBA "
     "订单发货时入账，接近送达、控制权转给买家的时点（IFRS 15）。文件本身就是这个月，不再按打印的时间（GMT-7）重新筛选。"),
    ("不含 IVA", "卖家是 IVA 纳税人：销项 IVA、进项 IVA 都不是收入或成本。商品、运费、促销在 Amazon 的文件里本来就"
     "不含 IVA（IVA 单列一列）；Amazon 的费用列（佣金、FBA、仓储、移除、入仓运输、订阅）是含 IVA 的，"
     "损益表用 ÷ %.2f 后的金额，差额是可抵扣的进项 IVA，列在②备查项。广告费由 Amazon 自己拆好（不含税＋IVA），不再拆。"
     % (1 + IVA_RATE)),
    ("退款", "退款冲减收入（单列一行）；Amazon 退还的佣金、FBA 费冲减同一费用行；退款时 Amazon 留下的退款管理费"
     "包含在佣金行的净额里。"),
    ("促销", "卖家出资的促销折扣（descuentos promocionales）是售价的减项，计入净销售收入。"),
    ("FBA 库存赔付", "Amazon 对丢失、损坏、退货未退回库存的赔付是存货损失的补偿，不是销售：单列为“其他收入”，"
     "放在平台费用之后，按 Amazon 报的金额（不拆 IVA）；冲回此前赔付的“ajuste general”冲减同一行。"),
    ("代扣税", "Amazon 作为平台按月代扣 IVA（8%，即销项 IVA 的一半）和 ISR（2.5%）。两者都是卖家的预缴税款，"
     "可抵减卖家自己的 IVA / ISR 申报，不是费用：列在②备查项，不进损益表；现金桥（③桥 C）把它们作为余额的扣款列出。"
     "本月的 ISR 以扣缴凭证（CFDI de retenciones）为准；交易明细里约 11 日扣的那笔 ISR 是上个月的。"),
    ("平台费用", "按 Amazon 的入账日计入当月；佣金、FBA 配送、月度仓储、长期仓储、入仓与移除、广告、订阅分行列示，"
     "同一科目的退还冲减同一行。"),
    ("商品成本与存货", "平台数据不含采购成本：在②或⑦录入。⑨的库存按 Amazon 的库存快照（下载当天，不是月末）。"),
]
TAX_NOTES = [
    "Amazon 每月出具一张扣缴凭证（Certificado de retención，CFDI de retenciones），列明本月代扣的 IVA 和 ISR；"
    "申报时用它抵减。凭证按付款日（FechaServ）归月，交易明细按入账日，两者的代扣 IVA 有几百比索的时间差属正常（见③桥 E）。",
    "交易明细里每月 11 日前后那笔没有说明的大额 Ajuste 是**上个月**的 ISR 代扣；本月的 ISR 要到下月 11 日前后才从余额扣。",
    "平台费用的进项 IVA 以 Amazon 开具的费用发票（CFDI，类型 I；退还用 E 贷记单）为准，按发票日期申报。"
    "本报表的进项 IVA 是按交易明细推算的（含税金额 × 16/116），与发票可能有尾差（见③桥 B）。",
    "广告费的发票不在卖家费用发票里（由 Amazon Ads 另行开具），申报进项 IVA 时另外收集。",
]


# ======================================================================= helpers

def fmt(x):
    try:
        return "{:,.2f}".format(float(x))
    except (TypeError, ValueError):
        return "-"


def pct(x, dec=1):
    try:
        return "{:.{d}f}%".format(float(x) * 100, d=dec)
    except (TypeError, ValueError):
        return "-"


def money_col(s):
    """'1,234.50' / '-0.5' / '' -> float."""
    return pd.to_numeric(s.astype(str).str.replace(",", "", regex=False)
                         .str.replace("$", "", regex=False).str.strip()
                         .replace({"": None, "nan": None, "None": None}),
                         errors="coerce").fillna(0.0)


def _r(v):
    """Money rounded to the currency's decimals (sums of cents carry float noise)."""
    return round(float(v), DECIMALS) + 0.0


def _s(v):
    """A text cell: None for NaN / empty."""
    if v is None or (isinstance(v, float) and v != v):
        return None
    return str(v) or None


def num(text):
    """'-34,405.11' -> -34405.11; None when it is not a number."""
    try:
        return float(str(text).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


def _same_month(folder):
    """This folder first, then the store's other folders of the same month, newest
    first. A later `run_batch --only stock` (or a rerun of one route) writes a new
    folder with just those files; the report should still find them."""
    root, name = os.path.split(os.path.abspath(folder))
    suffix = name[-8:] if re.search(r"_m\d{6}$", name) else None
    sib = sorted(glob.glob(os.path.join(root, "*" + suffix)), reverse=True) if suffix else []
    return [folder] + [d for d in sib if os.path.abspath(d) != os.path.abspath(folder)]


def _first(folder, pattern):
    for d in _same_month(folder):
        hits = sorted(glob.glob(os.path.join(d, pattern)))
        if hits:
            return hits[-1]
    return None


def _all(folder, pattern):
    """All files of one kind from ONE folder (the first that has any): invoice
    periods must not be mixed across runs."""
    for d in _same_month(folder):
        hits = sorted(glob.glob(os.path.join(d, pattern)))
        if hits:
            return hits
    return []


def rel(path, folder):
    """File name, plus its folder when it comes from another run of the same month."""
    if not path:
        return ""
    d = os.path.dirname(os.path.abspath(path))
    return os.path.basename(path) if d == os.path.abspath(folder) else \
        "%s（取自 %s）" % (os.path.basename(path), os.path.basename(d))


def read_meta(folder):
    try:
        with io.open(os.path.join(folder, "run_meta.json"), encoding="utf-8") as fh:
            return json.load(fh) or {}
    except Exception:
        return {}


def no_data_forms(meta):
    """Forms the downloader found empty (cause NO_DATA): an empty month, not a
    missing file."""
    return set(str(r.get("form")) for r in (meta.get("missing") or [])
               if r.get("cause") == "NO_DATA")


def month_window_utc(month):
    """Amazon's month in UTC: [1st 00:00, next 1st 00:00) at GMT-6."""
    a, b = periods.month_bounds(month)
    lo = datetime.datetime.combine(a, datetime.time()) - datetime.timedelta(hours=MONTH_UTC_OFFSET)
    hi = datetime.datetime.combine(b + datetime.timedelta(days=1), datetime.time()) \
        - datetime.timedelta(hours=MONTH_UTC_OFFSET)
    return lo, hi


# ======================================================================= load

_MONTH_ES = {"ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6, "jul": 7, "ago": 8,
             "sep": 9, "sept": 9, "oct": 10, "nov": 11, "dic": 12}
_WHEN = re.compile(r"^\s*(\d{1,2})\s+([a-z]+)\.?\s+(\d{4})\s+(\d{1,2}):(\d{2}):(\d{2})\s*"
                   r"([ap])\.?\s*m\.?\s*(?:GMT|UTC)\s*([+-]\d{1,2})?", re.I)


def parse_when(text):
    """'31 ago 2026 11:28:19 p.m. GMT-7' -> UTC datetime (naive), None when unreadable."""
    m = _WHEN.match(str(text or ""))
    if not m:
        return None
    d, mon, y, h, mi, s, ap, off = m.groups()
    mo = _MONTH_ES.get(mon.lower())
    if not mo:
        return None
    h = int(h) % 12 + (12 if ap.lower() == "p" else 0)
    local = datetime.datetime(int(y), mo, int(d), h, int(mi), int(s))
    return local - datetime.timedelta(hours=int(off or 0))


class TxFileError(Exception):
    """The transaction file cannot be read as the Spanish report (status, detail)."""

    def __init__(self, status, detail):
        Exception.__init__(self, detail)
        self.status, self.detail = status, detail


def load_transactions(path):
    """The transaction report as a DataFrame (money columns as floats, `_utc`,
    `_day` = the day in Amazon's month time zone). Raises TxFileError when the file is
    not the Spanish report."""
    with io.open(path, encoding="utf-8-sig", errors="replace") as fh:
        raw = fh.read()
    lines = raw.splitlines()
    hdr = None
    for i, ln in enumerate(lines[:60]):
        first = ln.split(",")[0].strip().strip('"').strip().lower()
        if first == HEADER_ES:
            hdr = i
            break
        if first in HEADER_EN:
            raise TxFileError("english", "交易明细是英文的（表头是“%s”而不是“fecha/hora”）：下载时卖家后台不是"
                                         "西班牙语。列名、类型名都会对不上，必须切回西班牙语（Español）后重新下载"
                                         % ln.split(",")[0].strip('"'))
    if hdr is None:
        raise TxFileError("no_header", "交易明细里找不到表头行（应以“fecha/hora”开头）：文件格式变了或不是交易报表")
    df = pd.read_csv(io.StringIO("\n".join(lines[hdr:])), dtype=str, keep_default_na=False)
    df.columns = [str(c).strip() for c in df.columns]
    missing = [c for c in REQUIRED_COLS if c not in df.columns]
    for c in missing:
        df[c] = ""
    df = df.drop(columns=[c for c in DROP_COLS if c in df.columns])
    for c in TC.MONEY_COLS + [TC.TOTAL_COL]:
        df[c] = money_col(df[c])
    df[T_QTY] = pd.to_numeric(df[T_QTY], errors="coerce").fillna(0)
    df["_utc"] = pd.to_datetime(df[T_DATE].map(parse_when), errors="coerce")
    df["_day"] = (df["_utc"] + pd.Timedelta(hours=MONTH_UTC_OFFSET)).dt.normalize()
    df = df.reset_index(drop=True)
    df.attrs["missing_cols"] = missing
    df.attrs["printed_tz"] = sorted(set(re.findall(r"GMT[+-]\d+", " ".join(df[T_DATE].head(50)))))
    return df


def empty_transactions():
    df = pd.DataFrame({c: pd.Series(dtype="object") for c in REQUIRED_COLS})
    for c in TC.MONEY_COLS + [TC.TOTAL_COL, T_QTY]:
        df[c] = pd.Series(dtype="float64")
    df["_utc"] = pd.Series(dtype="datetime64[ns]")
    df["_day"] = pd.Series(dtype="datetime64[ns]")
    df.attrs["missing_cols"] = []
    return df


# ---- summary PDF ("Resumen")

def pdf_text(path):
    """(text, tool) of a PDF: pdftotext -layout when installed, else pypdf / PyPDF2;
    (None, reason) when no tool can read it."""
    exe = shutil.which("pdftotext")
    if exe:
        try:
            out = subprocess.run([exe, "-layout", path, "-"], stdout=subprocess.PIPE,
                                 stderr=subprocess.PIPE, timeout=60)
            if out.returncode == 0 and out.stdout.strip():
                return out.stdout.decode("utf-8", "replace"), "pdftotext"
        except Exception:
            pass
    for mod in ("pypdf", "PyPDF2"):
        try:
            lib = __import__(mod)
        except ImportError:
            continue
        try:
            reader = lib.PdfReader(path)
            text = "\n".join(_lines_by_position(p) for p in reader.pages)
            if text.strip():
                return text, mod
        except Exception:
            continue
    return None, "本机没有能读 PDF 的工具（pdftotext / pypdf）"


def _lines_by_position(page):
    """A page's text as lines, rebuilt from where each piece is drawn. pypdf's
    plain extraction follows the PDF's drawing order, which on Amazon's summary
    separates every label from its amount; its "layout" mode returns almost
    nothing there. Pieces on the same height (2-point bands) form a line, left
    to right, joined with spaces — what pdftotext -layout gives, near enough for
    parse_summary."""
    parts = []

    def visit(text, cm, tm, font_dict, font_size):
        if text and text.strip():
            y = tm[5] * cm[3] + cm[5]
            x = tm[4] * cm[0] + cm[4]
            parts.append((int(round(y / 2.0)), x, text.strip()))
    page.extract_text(visitor_text=visit)
    rows = {}
    for y, x, t in parts:
        rows.setdefault(y, []).append((x, t))
    return "\n".join("   ".join(t for _, t in sorted(rows[y])) for y in sorted(rows, reverse=True))


_NUM = r"(-?[\d,]+(?:\.\d+)?)"
SUMMARY_TOTALS = ("ingresos", "gastos", "transferencias", "impuesto")
# label (normalised, as Amazon prints it) -> key; the number right after the label
SUMMARY_LINES = [
    ("impuestos de producto, envio y envoltura para regalo retenidos", "tax_collected"),
    ("impuestos de producto, envio y envoltura para regalo reembolsados", "tax_refunded"),
    ("impuesto al valor agregado (iva) obligatorio de amazon retenido", "iva_withheld"),
    ("impuesto sobre la renta retenido", "isr"),
    ("ajuste del iva del producto", "iva_adj"),
    ("ventas de producto fba", "sales_fba"),
    ("reembolsos de venta de producto fba", "refund_fba"),
    ("credito de inventario del fba", "inventory_credit"),
    ("reembolsos de tarifa de venta", "fee_refunds"),
    ("tarifas de administracion de reembolso", "refund_admin"),
    ("costo de la publicidad", "ads"),
]


def parse_summary(text):
    """The control totals of the summary PDF's text. {} when it is not the Spanish
    summary; "_lang": "en" when it is the English one."""
    out = {}
    if not text:
        return out
    lines = [TC.norm(x) for x in text.splitlines()]
    joined = "\n".join(lines)
    if "ingresos" not in joined and ("income" in joined or "expenses" in joined):
        out["_lang"] = "en"
        return out
    for key in SUMMARY_TOTALS:
        for ln in lines:
            m = re.match(r"^%s\s+\S.*?\s%s$" % (key, _NUM), ln)
            if m and ln.startswith(key + " "):
                out[key] = num(m.group(1))
                break
    for label, key in SUMMARY_LINES:
        m = re.search(re.escape(label) + r"\s+" + _NUM + r"(?:\s|$)", joined)
        if m:
            out[key] = num(m.group(1))
    m = re.search(r"desde (\w+) (\d{1,2}), (\d{4})", joined)
    if m:
        out["_from"] = "%s %s %s" % m.groups()
    return out


def load_summary(path):
    """{"values": {...}, "status": ok / unreadable / english / unparsed, "tool", "why"}"""
    text, tool = pdf_text(path)
    if text is None:
        return {"status": "unreadable", "why": tool, "values": {}}
    v = parse_summary(text)
    if v.get("_lang") == "en":
        return {"status": "english", "tool": tool, "values": {}}
    if not all(k in v for k in SUMMARY_TOTALS):
        return {"status": "unparsed", "tool": tool, "values": v,
                "why": "用 %s 读出的文字里找不到四个总数（%s）" % (
                    tool, "、".join(k for k in SUMMARY_TOTALS if k not in v))}
    return {"status": "ok", "tool": tool, "values": v}


# ---- CFDI (fee invoices, withholding certificate)

def _local(tag):
    return tag.rsplit("}", 1)[-1]


def _iter(root, name):
    return [e for e in root.iter() if _local(e.tag) == name]


def load_cfdi(path):
    """One seller-fee CFDI: type (I / E / P), totals, UUID, date, payments."""
    root = ET.parse(path).getroot()
    a = root.attrib
    uuid = ""
    for e in _iter(root, "TimbreFiscalDigital"):
        uuid = e.attrib.get("UUID", "")
    tax = 0.0
    for e in root:
        if _local(e.tag) == "Impuestos":
            tax = num(e.attrib.get("TotalImpuestosTrasladados")) or 0.0
    pays = [num(p.attrib.get("Monto")) or 0.0 for p in _iter(root, "Pago")]
    return {"type": a.get("TipoDeComprobante", ""), "total": num(a.get("Total")) or 0.0,
            "subtotal": num(a.get("SubTotal")) or 0.0, "tax": tax, "uuid": uuid.upper(),
            "date": a.get("Fecha", ""), "serie": a.get("Serie", ""), "paid": sum(pays),
            "file": os.path.basename(path)}


def load_invoices(folder):
    """The fee CFDIs of the month (fee_*.xml, one folder) and their index
    (fee_invoices.json, written by the downloader: type, number, start, end, file)."""
    xmls = _all(folder, "fee_*.xml")
    idx_path = os.path.join(os.path.dirname(xmls[0]), "fee_invoices.json") if xmls \
        else _first(folder, "fee_invoices.json")
    index = []
    if idx_path and os.path.exists(idx_path):
        try:
            with io.open(idx_path, encoding="utf-8") as fh:
                index = json.load(fh) or []
        except Exception:
            index = []
    by_uuid = dict((str(x.get("number", "")).upper(), x) for x in index
                   if str(x.get("format", "XML")).upper() == "XML")
    docs, bad = [], []
    for p in xmls:
        try:
            d = load_cfdi(p)
        except Exception as e:
            bad.append("%s（%s）" % (os.path.basename(p), e))
            continue
        x = by_uuid.get(d["uuid"]) or next((i for i in index if i.get("file") == d["file"]), None) or {}
        d["label"] = x.get("type", "")
        d["start"] = _iso(x.get("start"))
        d["end"] = _iso(x.get("end"))
        docs.append(d)
    return {"docs": docs, "index": index, "index_file": idx_path, "files": xmls, "bad": bad}


def _iso(s):
    try:
        return datetime.datetime.strptime(str(s)[:19], "%Y-%m-%dT%H:%M:%S")
    except (TypeError, ValueError):
        return None


def load_taxcert(path):
    """The monthly withholding certificate (CFDI de retenciones 2.0)."""
    root = ET.parse(path).getroot()
    out = {"file": os.path.basename(path), "iva": 0.0, "isr": 0.0}
    for e in _iter(root, "Periodo"):
        out["year"], out["mfrom"], out["mto"] = (e.attrib.get("Ejercicio"), e.attrib.get("MesIni"),
                                                 e.attrib.get("MesFin"))
    for e in _iter(root, "Totales"):
        out["operations"] = num(e.attrib.get("MontoTotOperacion")) or 0.0
        out["withheld"] = num(e.attrib.get("MontoTotRet")) or 0.0
    for e in _iter(root, "ImpRetenidos"):
        kind = e.attrib.get("ImpuestoRet", "")
        amt = num(e.attrib.get("MontoRet")) or 0.0
        if kind in ("002", "02"):
            out["iva"] += amt
        elif kind in ("001", "01"):
            out["isr"] += amt
    for e in _iter(root, "ServiciosPlataformasTecnologicas"):
        out["services"] = int(num(e.attrib.get("NumServ")) or 0)
        out["iva_charged"] = num(e.attrib.get("TotalIVATrasladado")) or 0.0
        out["platform_fees"] = num(e.attrib.get("MonTotalporUsoPlataforma")) or 0.0
    for e in _iter(root, "TimbreFiscalDigital"):
        out["uuid"] = e.attrib.get("UUID", "")
    out["issued"] = root.attrib.get("FechaExp", "")
    return out


def load_stock(path):
    if not path:
        return None
    s = pd.read_csv(path, dtype=str, encoding="utf-8-sig", keep_default_na=False)
    s.columns = [str(c).strip() for c in s.columns]
    for c in s.columns:
        if c in STOCK_TEXT:
            continue
        s[c] = pd.to_numeric(s[c].str.replace(",", "", regex=False), errors="coerce")
    return s


STOCK_TEXT = ("snapshot-date", "sku", "fnsku", "asin", "product-name", "condition", "currency",
              "alert", "recommended-action", "volume-unit-measurement", "storage-type",
              "marketplace", "product-group", "Recommended ship-in date",
              "Inventory age snapshot date")


def load_store(store, folder, month):
    S = {"store": store, "folder": folder, "month": month, "meta": read_meta(folder),
         "files": {}, "no_data": no_data_forms(read_meta(folder))}
    stem = periods.file_stem(month)
    tp = _first(folder, "%s*Transaction.csv" % stem)
    S["files"]["order"] = tp
    S["tx_status"], S["tx_detail"] = "ok", ""
    if tp:
        try:
            S["tx"] = load_transactions(tp)
        except TxFileError as e:
            S["tx_status"], S["tx_detail"] = e.status, e.detail
            S["tx"] = empty_transactions()
    else:
        others = _first(folder, "*Transaction.csv")
        if "order" in S["no_data"]:
            S["tx_status"] = "no_data"
        else:
            S["tx_status"] = "missing"
            S["tx_detail"] = ("目录里的交易明细是别的月份的（%s）" % os.path.basename(others)) if others else ""
        S["tx"] = empty_transactions()
    sp = _first(folder, "%s*Summary.pdf" % stem)
    S["files"]["summary"] = sp
    S["summary"] = load_summary(sp) if sp else None
    S["invoice"] = load_invoices(folder)
    S["files"]["invoice"] = S["invoice"]["files"]
    cp = _first(folder, "retenciones_*.xml")
    S["files"]["taxcert"] = cp
    S["taxcert"], S["taxcert_error"] = None, ""
    if cp:
        try:
            S["taxcert"] = load_taxcert(cp)
        except Exception as e:
            S["taxcert_error"] = str(e)
    kp = _first(folder, "stock_*.csv")
    S["files"]["stock"] = kp
    S["stock"] = load_stock(kp)
    return S


# ======================================================================= derive

def ledger(t):
    """One record per (transaction line, money column with an amount): its report
    line, IVA mode, amount (as printed), net and IVA. Sum of `amt` = sum of `total`
    whenever every row adds up."""
    if not len(t):
        return pd.DataFrame(columns=["_row", "_col", "amt", "line", "mode", "net", "iva", "_key",
                                     "_zh", T_TYPE, T_DESC, T_SKU, T_ORDER, T_QTY, "_utc", "_day"])
    entries = [TC.classify(tp, d, s, tot) for tp, d, s, tot in
               zip(t[T_TYPE], t[T_DESC], t[T_SKU], t[TC.TOTAL_COL])]
    rows = t.assign(_row=range(len(t)), _key=[e.key if e else None for e in entries])
    long = rows.melt(id_vars=["_row"], value_vars=TC.MONEY_COLS, var_name="_col", value_name="amt")
    long = long[long["amt"].abs() > 1e-9].copy()
    tl, md, nets, ivas, zh = [], [], [], [], []
    for r, col, amt in zip(long["_row"], long["_col"], long["amt"]):
        e = entries[r]
        line, mode = TC.column_target(e, col)
        n, v = TC.split(amt, mode, IVA_RATE)
        tl.append(line)
        md.append(mode)
        nets.append(n)
        ivas.append(v)
        zh.append(e.zh if e and line != TC.UNREGISTERED else
                  "⚠ 未登记：%s%s" % (rows.at[r, T_TYPE], "" if e is None else " · " + col))
    long["line"], long["mode"], long["net"], long["iva"], long["_zh"] = tl, md, nets, ivas, zh
    keep = rows[["_row", "_key", T_TYPE, T_DESC, T_SKU, T_ORDER, T_QTY, "_utc", "_day"]]
    return long.merge(keep, on="_row", how="left")


def derive(S):
    t = S["tx"]
    m = {}
    L = ledger(t)
    S["ledger"] = L
    by_line = L.groupby("line")
    net = by_line["net"].sum()
    amt = by_line["amt"].sum()
    # Each P&L line: its net rounded to cents, its IVA = gross − net, so that
    # net + IVA is exactly what Amazon charged and bridge C ties to the cent.
    m["pl"] = dict((k, _r(net.get(k, 0.0))) for k in TC.PL_LINES)
    m["pl_gross"] = dict((k, _r(amt.get(k, 0.0))) for k in TC.PL_LINES)
    m["pl_iva"] = dict((k, _r(m["pl_gross"][k] - m["pl"][k])) for k in TC.PL_LINES)
    for k in TC.TAX_LINES + TC.CASH_LINES:
        m[k] = _r(amt.get(k, 0.0))
    m["fee_iva"] = _r(sum(m["pl_iva"].values()))                      # input IVA (negative)
    m["net_sales"] = _r(sum(m["pl"][k] for k in TC.INCOME_LINES))
    m["fees"] = _r(sum(m["pl"][k] for k in FEE_ROWS))
    m["contrib"] = _r(m["net_sales"] + m["fees"] + m["pl"][TC.OTHER_INCOME])
    m["tx_total"] = _r(t[TC.TOTAL_COL].sum()) if len(t) else 0.0
    m["lines"] = int(len(t))
    # every row: total = sum of its columns
    if len(t):
        diff = (t[TC.MONEY_COLS].sum(axis=1) - t[TC.TOTAL_COL]).abs()
        m["row_bad"] = int((diff > TOL_ROW).sum())
        m["row_bad_amt"] = float((t[TC.MONEY_COLS].sum(axis=1) - t[TC.TOTAL_COL])[diff > TOL_ROW].sum())
    else:
        m["row_bad"], m["row_bad_amt"] = 0, 0.0
    # unregistered: (tipo, descripción or column) -> lines, amount
    U = L[L["line"] == TC.UNREGISTERED]
    m["unregistered"] = []
    for _, x in U.groupby("_row"):
        r0 = x.iloc[0]
        cols = "、".join(x["_col"])
        m["unregistered"].append((str(r0[T_TYPE]), str(r0[T_DESC] or "")[:60], cols, float(x["amt"].sum())))
    # KPIs
    is_order = t[T_TYPE].map(TC.norm) == ORDER_TIPO if len(t) else pd.Series(dtype=bool)
    is_refund = t[T_TYPE].map(TC.norm) == REFUND_TIPO if len(t) else pd.Series(dtype=bool)
    o, rf = t[is_order], t[is_refund]
    m["orders_n"] = int(o[T_ORDER].replace("", pd.NA).dropna().nunique()) if len(o) else 0
    m["order_lines"] = int(len(o))
    m["units"] = int(o[T_QTY].sum()) if len(o) else 0
    m["refund_lines"] = int(len(rf))
    m["refund_units"] = int(rf[T_QTY].sum()) if len(rf) else 0
    m["refund_orders"] = int(rf[T_ORDER].replace("", pd.NA).dropna().nunique()) if len(rf) else 0
    m["gmv"] = _r(o["ventas de productos"].sum() + o["impuesto de ventas de productos"].sum()) \
        if len(o) else 0.0
    o_sales = float(o["ventas de productos"].sum()) if len(o) else 0.0
    o_tax = float(o["impuesto de ventas de productos"].sum()) if len(o) else 0.0
    m["iva_rate"] = o_tax / o_sales if o_sales else None
    o_out = float(L.loc[(L["line"] == TC.IVA_OUT) & (L[T_TYPE].map(TC.norm) == ORDER_TIPO), "amt"].sum()) \
        if len(L) else 0.0
    o_wh = float(o["impuesto de retenciones en la plataforma"].sum()) if len(o) else 0.0
    m["withhold_ratio"] = -o_wh / o_out if o_out else None
    # ISR: the lump(s) of the previous month
    isr_rows = L[L["_key"] == "isr"]
    m["isr_rows"] = [(str(d.date()) if pd.notna(d) else "?", float(a))
                     for d, a in zip(isr_rows["_day"], isr_rows["amt"])]
    # deferred lines (Estado de la transacción other than "Lanzado")
    if len(t) and T_STATE in t.columns:
        st = t[T_STATE].map(TC.norm)
        dfr = t[(st != "") & ~st.isin(POSTED)]
        m["deferred_n"] = int(len(dfr))
        m["deferred_amt"] = float(dfr[TC.TOTAL_COL].sum())
        m["deferred_states"] = dfr[T_STATE].value_counts().to_dict()
    else:
        m["deferred_n"], m["deferred_amt"], m["deferred_states"] = 0, 0.0, {}
    # lines whose posting day is outside the month (the file should be the month)
    if len(t):
        a, b = periods.month_bounds(S["month"])
        days = t["_day"]
        m["outside_n"] = int(((days < pd.Timestamp(a)) | (days > pd.Timestamp(b))).sum())
        m["undated_n"] = int(days.isna().sum())
    else:
        m["outside_n"] = m["undated_n"] = 0
    # ---- bridge A: order and refund lines -> net sales
    def part(tipo, line):
        if not len(L):
            return 0.0
        return float(L.loc[(L[T_TYPE].map(TC.norm) == tipo) & (L["line"] == line), "amt"].sum())
    A = {}
    for side, tipo in (("o", ORDER_TIPO), ("r", REFUND_TIPO)):
        A[side + "_sales"] = part(tipo, TC.SALES) + part(tipo, TC.REFUND)
        A[side + "_ship"] = part(tipo, TC.SHIPPING)
        A[side + "_wrap"] = part(tipo, TC.GIFTWRAP)
        A[side + "_promo"] = part(tipo, TC.PROMO)
        A[side] = A[side + "_sales"] + A[side + "_ship"] + A[side + "_wrap"] + A[side + "_promo"]
    A = dict((k, _r(v)) for k, v in A.items())
    A["resid"] = _r(A["o"] + A["r"] - m["net_sales"])
    m["bridge_a"] = A
    # ---- bridge C: contribution -> change of the Amazon balance
    m["balance_calc"] = _r(m["contrib"] + m[TC.IVA_OUT] + m[TC.IVA_WITHHELD] + m["fee_iva"] + m[TC.ISR]
                           + m[TC.IVA_ADJ] + m[TC.PAYOUT])
    m["bridge_c_resid"] = _r(m["balance_calc"] - m["tx_total"])
    # ---- bridge D: the summary PDF's groups, rebuilt from the transaction report
    g = dict((k, 0.0) for k in TC.SUMMARY_GROUPS)
    for line, a_ in zip(L["line"], L["amt"]):
        grp = TC.SUMMARY_GROUP.get(line, "gastos")
        if grp == TC.SIGNED:
            grp = "ingresos" if a_ > 0 else "gastos"
        g[grp] += a_
    m["groups_csv"] = dict((k, _r(v)) for k, v in g.items())
    sm = S.get("summary") or {}
    m["summary_status"] = sm.get("status") if sm else "missing"
    m["summary"] = sm.get("values") or {}
    # ---- bridge E: the withholding certificate
    c = S.get("taxcert")
    m["cert"] = c
    m["cert_isr"] = -(c or {}).get("isr", 0.0) if c else 0.0           # withheld from the seller: negative
    m["cert_iva"] = -(c or {}).get("iva", 0.0) if c else 0.0
    m["order_net_sales"] = A["o"]
    # ---- bridge B: fee invoices per period
    m["bridge_b"], m["invoice_payments"], m["invoice_unassigned"] = bridge_b(S, L)
    S["m"] = m
    S["sku"] = sku_table(S)
    S["returns"] = returns_table(S)
    S["reimb"] = reimbursements(S)
    return S


def bridge_b(S, L):
    """Per fee period of the month's CFDIs: invoices (I − E, with IVA) vs the
    transaction report's seller fees (with IVA, as charges) posted in the period."""
    lo, hi = month_window_utc(S["month"])
    docs = S["invoice"]["docs"]
    periods_ = {}
    pays, unassigned = [], []
    for d in docs:
        if d["type"] == "P":
            pays.append(d)
            continue
        if d["type"] not in ("I", "E"):
            unassigned.append(d)
            continue
        if not d["start"] or not d["end"]:
            unassigned.append(d)
            continue
        key = (d["start"], d["end"])
        P = periods_.setdefault(key, {"start": d["start"], "end": d["end"], "inv_i": 0.0, "inv_e": 0.0,
                                      "n_i": 0, "n_e": 0, "tax": 0.0})
        if d["type"] == "I":
            P["inv_i"] += d["total"]
            P["n_i"] += 1
            P["tax"] += d["tax"]
        else:
            P["inv_e"] += d["total"]
            P["n_e"] += 1
            P["tax"] -= d["tax"]
    out = []
    fees = L[L["line"].isin(INVOICED_LINES)] if len(L) else L
    ads = L[L["line"] == TC.ADS] if len(L) else L
    for key in sorted(periods_):
        P = periods_[key]
        a, b = P["start"], P["end"]
        w = (fees["_utc"] >= pd.Timestamp(a)) & (fees["_utc"] < pd.Timestamp(b)) if len(fees) else []
        wa = (ads["_utc"] >= pd.Timestamp(a)) & (ads["_utc"] < pd.Timestamp(b)) if len(ads) else []
        P["tx_fees"] = _r(-fees.loc[w, "amt"].sum()) if len(fees) else 0.0
        P["tx_lines"] = int(fees.loc[w, "_row"].nunique()) if len(fees) else 0
        P["tx_ads"] = _r(-ads.loc[wa, "amt"].sum()) if len(ads) else 0.0
        P["inv_i"], P["inv_e"] = _r(P["inv_i"]), _r(P["inv_e"])
        P["inv"] = _r(P["inv_i"] - P["inv_e"])
        P["diff"] = _r(P["inv"] - P["tx_fees"])
        P["full"] = a >= lo and b <= hi
        P["id"] = "%s_%s" % (a.date(), b.date())
        out.append(P)
    return out, pays, unassigned


def sku_table(S):
    L = S["ledger"]
    pl = L[L["line"].isin(TC.PL_LINES)] if len(L) else L
    if not len(pl):
        return pd.DataFrame()
    sk = pl[T_SKU].replace("", pd.NA).fillna("(无 SKU · 店铺级费用)")
    g = pl.assign(_sku=sk).pivot_table(index="_sku", columns="line", values="net", aggfunc="sum",
                                      fill_value=0.0)
    for k in TC.PL_LINES:
        if k not in g.columns:
            g[k] = 0.0
    t = S["tx"]
    is_o = t[T_TYPE].map(TC.norm) == ORDER_TIPO
    is_r = t[T_TYPE].map(TC.norm) == REFUND_TIPO
    g["units"] = t[is_o].groupby(T_SKU)[T_QTY].sum().reindex(g.index).fillna(0).astype(int)
    g["returns"] = t[is_r].groupby(T_SKU)[T_QTY].sum().reindex(g.index).fillna(0).astype(int)
    g["name"] = t[is_o | is_r].groupby(T_SKU)[T_DESC].first().reindex(g.index)
    g["income"] = g[list(TC.INCOME_LINES)].sum(axis=1)
    g["fees"] = g[list(FEE_ROWS)].sum(axis=1)
    g["contrib"] = g["income"] + g["fees"] + g[TC.OTHER_INCOME]
    return g.sort_values(TC.SALES, ascending=False)


def returns_table(S):
    """Refund lines by SKU: lines, units, refunded price / shipping / promotion,
    fees refunded (net of the refund administration fee), the SKU's sales."""
    L = S["ledger"]
    if not len(L):
        return pd.DataFrame()
    R = L[L[T_TYPE].map(TC.norm) == REFUND_TIPO]
    if not len(R):
        return pd.DataFrame()
    g = R.pivot_table(index=T_SKU, columns="line", values="net", aggfunc="sum", fill_value=0.0)
    for k in (TC.REFUND, TC.SHIPPING, TC.PROMO, TC.COMMISSION, TC.FBA_FEE):
        if k not in g.columns:
            g[k] = 0.0
    t = S["tx"]
    rr = t[t[T_TYPE].map(TC.norm) == REFUND_TIPO]
    g["lines"] = rr.groupby(T_SKU).size().reindex(g.index).fillna(0).astype(int)
    g["units"] = rr.groupby(T_SKU)[T_QTY].sum().reindex(g.index).fillna(0).astype(int)
    g["name"] = rr.groupby(T_SKU)[T_DESC].first().reindex(g.index)
    sku = S["sku"]
    g["sales"] = sku[TC.SALES].reindex(g.index).fillna(0.0) if len(sku) else 0.0
    g["sold"] = sku["units"].reindex(g.index).fillna(0).astype(int) if len(sku) else 0
    return g.sort_values(TC.REFUND)


def reimbursements(S):
    """FBA inventory reimbursement lines (⑥): day, SKU, reason, amount."""
    L = S["ledger"]
    if not len(L):
        return []
    R = L[L["_key"] == "fba_reimbursement"]
    out = []
    for _, x in R.sort_values("_day").iterrows():
        reason, zh = TC.reimburse_reason(x[T_DESC])
        out.append({"day": str(x["_day"].date()) if pd.notna(x["_day"]) else "", "sku": x[T_SKU],
                    "reason": reason, "zh": zh, "amt": float(x["amt"])})
    return out


# ======================================================================= checks

def validate(stores):
    C = []
    for S in stores:
        st, m, f = S["store"], S["m"], S["files"]
        status = S["tx_status"]
        C.append(K.check("file_order", st, "交易明细（Transacciones）在",
                         status in ("ok", "no_data") and (status == "no_data" or bool(f.get("order"))),
                         "%s，%d 行" % (os.path.basename(f["order"]), m["lines"]) if status == "ok" else
                         "本月没有交易（下载时 Amazon 显示没有数据）：报表按空月份处理" if status == "no_data" else
                         "%s 在，但读不出（原因见下一条）：本店本月的损益表是空的"
                         % os.path.basename(f.get("order") or "") if status in ("english", "no_header") else
                         "没有 %sMonthlyTransaction.csv：本店本月的损益表是空的，必须重新下载%s"
                         % (periods.file_stem(S["month"]),
                            "（%s）" % S["tx_detail"] if S["tx_detail"] else "")))
        if status in ("english", "no_header"):
            C.append(K.check("language", st, "交易明细是西班牙语", False, S["tx_detail"]))
        miss_cols = S["tx"].attrs.get("missing_cols") or []
        if status == "ok" and miss_cols:
            C.append(K.check("columns", st, "交易明细的列齐全", False,
                             "缺少这些列：%s。Amazon 改了报表格式，或后台不是西班牙语" % "、".join(miss_cols)))
        C.append(K.check("unregistered", st, "交易类型全部已登记", not m["unregistered"],
                         "全部 %d 种已登记" % S["ledger"]["_key"].nunique() if not m["unregistered"] else
                         "未登记 %d 行：%s（金额已计入“未登记科目”行，不拆 IVA）" % (
                             len(m["unregistered"]), "；".join(
                                 "%s %s [%s] %s" % (a, b, c, fmt(d)) for a, b, c, d in m["unregistered"][:8]))))
        C.append(K.check("row_total", st, "每行 total ＝ 各金额列之和", m["row_bad"] == 0,
                         "全部 %d 行相符（容差 %.2f）" % (m["lines"], TOL_ROW) if not m["row_bad"] else
                         "%d 行不相符，差额合计 %s：文件可能被截断或列错位" % (m["row_bad"], fmt(m["row_bad_amt"]))))
        A = m["bridge_a"]
        C.append(K.check("bridge_a", st, "桥 A：订单与退款行 → 净销售收入", abs(A["resid"]) <= TOL_BRIDGE,
                         "残差 %s" % fmt(A["resid"])))
        C.append(K.check("bridge_c", st, "桥 C：贡献毛利 → Amazon 余额变动 ＝ 交易明细 total 合计",
                         abs(m["bridge_c_resid"]) <= TOL_BRIDGE,
                         "推算 %s，明细 %s，残差 %s" % (fmt(m["balance_calc"]), fmt(m["tx_total"]),
                                                  fmt(m["bridge_c_resid"]))))
        C.extend(summary_checks(S))
        if m["iva_rate"] is not None:
            lo, hi = IVA_RATE_BAND
            C.append(K.check("iva_rate", st, "商品销项 IVA ≈ %s" % pct(IVA_RATE, 0),
                             lo <= m["iva_rate"] <= hi, "订单行 IVA ÷ 商品价款 ＝ %s（应在 %s–%s）"
                             % (pct(m["iva_rate"], 2), pct(lo), pct(hi)), severity="warn"))
        if m["withhold_ratio"] is not None:
            lo, hi = WITHHOLD_BAND
            C.append(K.check("withholding_rate", st, "代扣 IVA ≈ 销项 IVA 的一半",
                             lo <= m["withhold_ratio"] <= hi,
                             "订单行代扣 IVA ÷ 销项 IVA ＝ %s（应在 %s–%s）。不在区间说明代扣规则或卖家的税务登记变了"
                             % (pct(m["withhold_ratio"]), pct(lo, 0), pct(hi, 0)), severity="warn"))
        C.extend(taxcert_checks(S))
        C.extend(invoice_checks(S))
        C.append(K.check("file_stock", st, "库存快照在", bool(f.get("stock")),
                         rel(f.get("stock"), S["folder"]) if f.get("stock") else
                         "没有 stock_*.csv：⑨ 库存页为空", severity="warn"))
        C.append(K.check("deferred", st, "交易全部已入账（Lanzado）", not m["deferred_n"],
                         "全部“Lanzado”" if not m["deferred_n"] else
                         "%d 行状态不是“Lanzado”（%s），合计 %s；按入账口径已计入本月，钱要等 Amazon 放款"
                         % (m["deferred_n"], "、".join("%s %d" % kv for kv in m["deferred_states"].items()),
                            fmt(m["deferred_amt"])), severity="warn"))
        if m["lines"]:
            C.append(K.check("tx_in_month", st, "交易明细是本月的",
                             not m["outside_n"] and not m["undated_n"],
                             "全部 %d 行的入账日（GMT%+d）都在本月" % (m["lines"], MONTH_UTC_OFFSET)
                             if not m["outside_n"] and not m["undated_n"] else
                             "%d 行入账日不在本月、%d 行日期读不出：文件可能不是本月的月度报表"
                             % (m["outside_n"], m["undated_n"]), severity="warn"))
    return C


def summary_checks(S):
    st, m = S["store"], S["m"]
    status = m["summary_status"]
    name = "桥 D：摘要 PDF（Resumen）↔ 交易明细"
    if S["tx_status"] not in ("ok",):
        return [K.check("bridge_d", st, name, True, "没有可用的交易明细，无从核对", skipped=True)]
    if status == "missing":
        return [K.check("bridge_d", st, name, True, "没有 %sMonthlySummary.pdf，独立控制总数无法核对"
                        % periods.file_stem(S["month"]), skipped=True)]
    if status == "unreadable":
        return [K.check("bridge_d", st, name, True, "摘要 PDF 在，但%s，无法核对（不影响报表数字；"
                        "装 pdftotext 或 pypdf 后重跑报表即可）" % S["summary"].get("why", ""), skipped=True)]
    if status == "english":
        return [K.check("bridge_d", st, name, False, "摘要 PDF 是英文的：下载时卖家后台不是西班牙语，切回 Español 后重新下载")]
    if status == "unparsed":
        return [K.check("bridge_d", st, name, True, "%s，无法核对（PDF 版式变了？）" % S["summary"].get("why", ""),
                        skipped=True)]
    v, g = m["summary"], m["groups_csv"]
    y, mo = periods.parse_month(S["month"])
    want = "%s 1 %d" % (periods.MONTHS_EN3[mo - 1].lower(), y)
    if v.get("_from") and v["_from"] != want:
        return [K.check("bridge_d", st, name, False, "摘要 PDF 的期间从 %s 开始，不是 %s：下载到的是别的月份的摘要"
                        % (v["_from"], S["month"]))]
    diffs = dict((k, g[k] - v[k]) for k in TC.SUMMARY_GROUPS)
    bad = [k for k in TC.SUMMARY_GROUPS if abs(diffs[k]) > TOL_SUMMARY]
    out = [K.check("bridge_d", st, name, not bad,
                   "四个总数（Ingresos / Gastos / Transferencias / Impuesto）全部相符" if not bad else
                   "不符：%s。交易明细与 Amazon 自己的摘要对不上，先查文件是否完整" % "；".join(
                       "%s 明细 %s ≠ 摘要 %s" % (k, fmt(g[k]), fmt(v[k])) for k in bad))]
    if "isr" in v:
        ok = abs(m[TC.ISR] - v["isr"]) <= TOL_SUMMARY and abs(m[TC.IVA_ADJ] - v.get("iva_adj", 0.0)) <= TOL_SUMMARY
        out.append(K.check("isr_lump", st, "无说明的 Ajuste 行已正确拆成 ISR 与 IVA 调整", ok,
                           "ISR %s（摘要 %s），IVA 调整 %s（摘要 %s）" % (
                               fmt(m[TC.ISR]), fmt(v["isr"]), fmt(m[TC.IVA_ADJ]), fmt(v.get("iva_adj", 0.0)))
                           + ("" if ok else "：拆分规则（|金额| > %d 为 ISR）不再适用，需人工确认" % TC.ISR_MIN)))
    return out


def taxcert_checks(S):
    st, m = S["store"], S["m"]
    c = m["cert"]
    name = "桥 E：扣缴凭证（Certificado de retención）↔ 交易明细的代扣 IVA"
    if S.get("taxcert_error"):
        return [K.check("taxcert", st, name, False, "扣缴凭证读不出：%s" % S["taxcert_error"], severity="warn")]
    if not c:
        due = periods.taxcert_due(S["month"])
        return [K.check("taxcert", st, name, True,
                        "本月的扣缴凭证还没取得（Amazon 在 %s 前后出具%s）：本月 ISR 显示“待取得”，下次下载会取到"
                        % (due, "，现在还没到" if datetime.date.today() < due else "，请重新下载 taxcert"),
                        severity="warn", skipped=True)]
    y, mo = periods.parse_month(S["month"])
    out = []
    right = str(c.get("year")) == str(y) and int(c.get("mfrom") or 0) == mo and int(c.get("mto") or 0) == mo
    if not right:
        out.append(K.check("taxcert_period", st, "扣缴凭证是本月的", False,
                           "凭证期间是 %s 年 %s–%s 月，不是 %s：桥 E 和本月 ISR 不可信，重新下载 taxcert"
                           % (c.get("year"), c.get("mfrom"), c.get("mto"), S["month"]), severity="warn"))
    tol = max(0.01 * c["iva"], 50.0)
    d = -m["cert_iva"] - (-m[TC.IVA_WITHHELD])
    out.append(K.check("taxcert", st, name, abs(d) <= tol,
                       "凭证代扣 IVA %s，交易明细 %s，差 %s（容差 %s；凭证按付款日、明细按入账日，有时间差属正常）"
                       % (fmt(-m["cert_iva"]), fmt(-m[TC.IVA_WITHHELD]), fmt(d), fmt(tol)), severity="warn"))
    return out


def invoice_checks(S):
    st, m = S["store"], S["m"]
    out = []
    inv = S["invoice"]
    if inv["bad"]:
        out.append(K.check("fee_invoice_read", st, "费用发票（CFDI）都能读", False,
                           "读不出：%s" % "；".join(inv["bad"]), severity="warn"))
    if not m["bridge_b"]:
        out.append(K.check("fee_invoice", st, "桥 B：费用发票（CFDI）↔ 交易明细的平台费用", True,
                           "本月没有下载到费用发票（不影响报表数字，费用按交易明细计入）；下月或重新收集时补齐",
                           severity="warn", skipped=True))
    for P in m["bridge_b"]:
        name = "桥 B：费用发票 %s–%s ↔ 交易明细" % (P["start"].date(), P["end"].date())
        if not P["full"]:
            out.append(K.check("fee_invoice", st, name, True,
                               "部分期间：账期跨月，交易明细只含本月部分，不比较（发票 %s，本月部分费用 %s）"
                               % (fmt(P["inv"]), fmt(P["tx_fees"])), severity="warn", skipped=True))
            continue
        tol = max(0.01 * abs(P["inv"]), 5.0)
        out.append(K.check("fee_invoice", st, name, abs(P["diff"]) <= tol,
                           "发票（I − E，含 IVA）%s，交易明细同期费用 %s，差 %s（容差 %s）"
                           % (fmt(P["inv"]), fmt(P["tx_fees"]), fmt(P["diff"]), fmt(tol)), severity="warn"))
    if m["invoice_unassigned"]:
        out.append(K.check("fee_invoice_index", st, "费用发票都能归到账期", False,
                           "%d 张发票在 fee_invoices.json 里找不到账期：%s" % (
                               len(m["invoice_unassigned"]),
                               "、".join(d["file"] for d in m["invoice_unassigned"][:5])), severity="warn"))
    return out


# ======================================================================= workbook
#
# The layout comes from core/report_family.py (the family's sheet writers); what
# follows is Amazon's content: which lines, which bridges, which columns.

def build_workbook(stores, ctx, out_path):
    wb = Workbook()
    wb.remove(wb.active)
    fam = RF.Family(stores, money=MNY)
    sheet_readme(wb, fam, ctx)
    sheet_pl(wb, fam, ctx)
    sheet_bridge(wb, fam, ctx)
    sheet_fees(wb, fam, ctx)
    sheet_returns(wb, fam, ctx)
    sheet_reimbursements(wb, fam, ctx)
    sheet_sku(wb, fam, ctx)
    sheet_trend(wb, fam, ctx)
    sheet_inventory(wb, fam, ctx)
    sheet_gaps(wb, fam, ctx)
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    return out_path


def pl_label(key):
    """② label of a line key; the memo name when it is not a row of ②."""
    return PL_LABELS.get(key) or MEMO_LABELS.get(key) or key


# ----------------------------------------------------------------- ① readme

def sheet_readme(wb, fam, ctx):
    stores = fam.stores
    sources = []
    for S in stores:
        f, meta = S["files"], S["meta"]
        inv = S["invoice"]
        n_p = sum(1 for d in inv["docs"] if d["type"] == "P")
        flags = meta.get("flags") or {}
        sources.append((S["store"], "下载于 %s%s。交易明细：%s；摘要：%s；费用发票：%d 张 XML（%d 张付款补充）；"
                        "扣缴凭证：%s；库存：%s。" % (
                            meta.get("downloaded_at") or "?",
                            "（下载时把后台切回了西班牙语）" if flags.get("language_switched") else "",
                            rel(f.get("order"), S["folder"]) or
                            ("本月没有交易" if S["tx_status"] == "no_data" else "缺"),
                            rel(f.get("summary"), S["folder"]) or "缺",
                            len(inv["docs"]), n_p,
                            rel(f.get("taxcert"), S["folder"]) or "缺",
                            rel(f.get("stock"), S["folder"]) or "缺")))
    tot = lambda k: sum(S["m"].get(k, 0.0) or 0.0 for S in stores)
    plt = lambda k: sum(S["m"]["pl"][k] for S in stores)
    rev = plt(TC.SALES)
    points = [("规模", "本月商品销售收入 %s（不含 IVA），%d 张订单、%d 件；含 IVA 的商品成交额 %s。退款 %d 行、%d 件。"
               % (fmt(rev), int(tot("orders_n")), int(tot("units")), fmt(tot("gmv")),
                  int(tot("refund_lines")), int(tot("refund_units"))))]
    if rev:
        points.append(("费用率", "平台费用合计（不含 IVA）%s，占商品销售收入 %s；其中佣金 %s、FBA 物流与仓储 %s、广告 %s。"
                       % (fmt(-tot("fees")), pct(-tot("fees") / rev), pct(-plt(TC.COMMISSION) / rev),
                          pct(-sum(plt(k) for k in FBA_ROWS) / rev), pct(-plt(TC.ADS) / rev))))
        points.append(("贡献毛利", "经营贡献毛利（未扣商品成本）%s，占商品销售收入 %s。"
                       % (fmt(tot("contrib")), pct(tot("contrib") / rev))))
    points.append(("代扣税", "Amazon 本月从余额代扣 IVA %s、上月 ISR %s；本月 ISR（扣缴凭证）%s。都是预缴税款，不是费用。"
                   % (fmt(-tot(TC.IVA_WITHHELD)), fmt(-tot(TC.ISR)),
                      fmt(-tot("cert_isr")) if all(S["m"]["cert"] for S in stores) else "待取得")))
    sections = [
        ("一、这份报表回答什么", [
            ("本月赚了多少", "【%s】从商品销售收入推到“经营贡献毛利（未扣商品成本）”，"
             "填入商品成本后得到税前净利润。" % SHEETS[1]),
            ("钱对不对得上", "【%s】五座桥：A 订单/退款 → 净销售收入、B 平台费用发票 → 交易明细、"
             "C 贡献毛利 → Amazon 余额变动（现金）、D Amazon 摘要 PDF → 交易明细、E 扣缴凭证 → 代扣税。" % SHEETS[2]),
            ("还缺什么", "【%s】列出未通过的校验、缺的文件，以及报税时要注意的事。" % SHEETS[9])]),
        ("二、口径", [
            ("按 Amazon 入账月", "全部金额来自 Amazon 的月度交易报表（Transacciones，月度）。Amazon 按墨西哥时间（GMT-6）"
             "的自然月、按入账日出报表；FBA 订单在发货时入账。文件就是这个月，不再按打印的时间重新筛选。"),
            ("不含 IVA", "损益表一律不含 IVA。商品、运费、促销在 Amazon 文件里本来就不含 IVA；平台费用列是含 IVA 的，"
             "÷ %.2f 后入表，差额是进项 IVA（②备查项）。" % (1 + IVA_RATE)),
            ("代扣税不是费用", "Amazon 代扣的 IVA 和 ISR 是卖家的预缴税款，列在②备查项、③桥 C 和桥 E，不进损益表"
             "（与美客多墨西哥报表一致）。"),
            ("交易类型", "Amazon 交易报表的每一种类型（tipo＋说明）都登记在字典里（西语原名＋中文＋归入哪一行），"
             "见【%s】。出现字典里没有的类型或列时报 error，金额进“未登记科目 ⚠”行。" % SHEETS[3])]),
        ("三、数据来源", sources),
        ("四、本期要点", points),
        ("五、各页说明", [
            (SHEETS[1], "各店分列 + 合计。行位置每期固定，可直接跨月对比；灰底=公式，浅蓝=小计，黄底=需你录入。"),
            (SHEETS[2], "桥 A 订单/退款 → 净销售、桥 B 费用发票 → 交易明细、桥 C 贡献毛利 → 余额变动、"
                        "桥 D 摘要 PDF → 交易明细、桥 E 扣缴凭证 → 代扣税。"),
            (SHEETS[3], "本月出现的每一种交易类型与金额列：西语原名、中文、归入哪一行、笔数、含 IVA、IVA、不含 IVA。"),
            (SHEETS[4], "退款：本月合计（含 Amazon 摘要里的佣金退还与退款管理费），以及按 SKU 的退款明细与退款率。"),
            (SHEETS[5], "Amazon 不提供逐件退货质检报表；本页列 FBA 库存赔付（丢失、损坏、退货未退回）的逐笔明细。"),
            (SHEETS[6], "按 SKU 的收入、费用与贡献；黄色列填单位成本即得 SKU 毛利。"),
            (SHEETS[7], "本月按周的销售与费用；有往月下载时追加各月对比。"),
            (SHEETS[8], "FBA 库存快照：可售、在途、预留、库龄、近 30/90 天发货、可售天数、预计仓储费，以及本月销量。"),
            (SHEETS[9], "校验结果、缺失文件及原因、报税提示。校验和缺失为空表示没有待办。")]),
        ("六、到账净额 ≠ 利润", [
            ("余额变动是现金指标", "交易明细 total 的合计（本月 %s）是 Amazon 余额的净变动，已含转出到银行的 %s、"
             "销项 IVA、代扣税。利润看②的经营贡献毛利（不含 IVA）。" % (fmt(tot("tx_total")), fmt(-tot(TC.PAYOUT))))]),
        ("七、会计处理", list(ACCOUNTING_POLICIES)),
        ("八、报税时要注意", [("提示 %d" % i, q) for i, q in enumerate(TAX_NOTES, 1)]),
    ]
    bad = RF.failed_errors(ctx["checks"])
    warning = ("本期有 %d 项 error 级校验未通过，数字可能不完整，请先看【%s】顶部的校验结果。"
               % (len(bad), SHEETS[9])) if bad else None
    fam.readme_sheet(wb, SHEETS[0], "亚马逊 月度财务报表 · %s · %s" % (ctx["period_cn"], ctx["who"]),
                     "生成于 %s。所有金额单位：墨西哥比索（%s），保留 %d 位小数。" % (ctx["built"], CUR, DECIMALS),
                     sections, warning=warning)


# ----------------------------------------------------------------- ② P&L

def _fee_share(keys):
    return lambda S: -sum(S["m"]["pl"][k] for k in keys)


def _cert_note(fam):
    return fam.per_store(lambda S: "凭证 %s" % S["m"]["cert"]["file"] if S["m"]["cert"] else "待取得")


def sheet_pl(wb, fam, ctx):
    sh, r, _rows_at, note_col = fam.pl_sheet(
        wb, SHEETS[1], PL_ROWS, value=lambda S, key: S["m"]["pl"][key],
        title="合并损益表（%s）" % ctx["period_cn"],
        subtitle="口径：按 Amazon【入账月】归集（月度交易报表），金额不含 IVA。单位：%s。"
                 "灰底=公式行，浅蓝=小计，黄底=需你录入。行位置每期固定，可直接跨月对比。" % CUR)
    sales = lambda S: S["m"]["pl"][TC.SALES]
    r = fam.ratio_block(sh, r + 1, "费用率（占商品销售收入）", [
        ("销售佣金", _fee_share((TC.COMMISSION,)), sales),
        ("FBA 物流与仓储（配送＋仓储＋入仓移除）", _fee_share(FBA_ROWS), sales),
        ("广告费", _fee_share((TC.ADS,)), sales),
        ("平台费用合计", _fee_share(FEE_ROWS), sales),
        ("退款率（退货退款 ÷ 商品销售收入）", lambda S: -S["m"]["pl"][TC.REFUND], sales),
        ("经营贡献毛利率", lambda S: S["m"]["contrib"], sales)], note_col)
    iva_net = lambda S: _r(S["m"][TC.IVA_OUT] + S["m"]["fee_iva"] + S["m"][TC.IVA_WITHHELD] + S["m"][TC.IVA_ADJ])
    fam.memo_block(sh, r + 1, "备查项（不参与上表计算）", [
        ("含 IVA 的商品成交额（订单行）", "ventas + impuesto de ventas · Pedido", lambda S: S["m"]["gmv"],
         "参考指标：本月入账订单的商品价款含 IVA"),
        ("订单数 / 件数", "Id. del pedido / cantidad", lambda S: S["m"]["orders_n"],
         fam.per_store(lambda S: "%d 张订单，%d 件" % (S["m"]["orders_n"], S["m"]["units"]))),
        ("销项 IVA（商品、运费、礼品包装、促销所含，已扣退款）", "impuesto de ventas / abono de envío / …",
         lambda S: S["m"][TC.IVA_OUT], "卖家应申报的销项 IVA（%s）" % pct(IVA_RATE, 0)),
        ("Amazon 代扣 IVA（预缴，可抵减）", "impuesto de retenciones en la plataforma",
         lambda S: S["m"][TC.IVA_WITHHELD], "≈ 销项 IVA 的一半（8%）；不是费用。凭证金额见③桥 E"),
        ("平台费用所含进项 IVA（可抵扣）", "IVA acreditable", lambda S: S["m"]["fee_iva"],
         "＝ 含税费用 − 含税费用 ÷ %.2f；广告费用 Amazon 拆出的 IVA。以 Amazon 费用发票为准" % (1 + IVA_RATE)),
        ("IVA 调整", "Ajuste del IVA del producto", lambda S: S["m"][TC.IVA_ADJ], "Amazon 对销项 IVA 的尾差调整"),
        ("本月 IVA 净额估算（正数＝还要向 SAT 缴）", "", iva_net,
         "＝ 销项 IVA − 进项 IVA − Amazon 代扣 IVA ＋ IVA 调整；仅供参考，以申报为准"),
        ("本月 ISR 代扣（扣缴凭证）", "ImpRetenidos 001 · Certificado de retención",
         lambda S: S["m"]["cert_isr"], "本月销售对应的 ISR（2.5%%）预缴；%s" % _cert_note(fam)),
        ("本月从余额扣的 ISR（上月的）", "Ajuste（sin descripción）", lambda S: S["m"][TC.ISR],
         "约每月 11 日扣上个月的 ISR；在③桥 C 中作为余额扣款"),
        ("【敏感性】若代扣税不可抵扣时的贡献毛利", "",
         lambda S: _r(S["m"]["contrib"] + S["m"][TC.IVA_WITHHELD] + S["m"]["cert_isr"]),
         "＝ 经营贡献毛利 − 代扣 IVA − 本月 ISR（凭证）。仅当卖家无法抵减时适用（如未按时申报）"),
        ("本月 Amazon 余额净变动（含 IVA、代扣、转账）", "Σ total", lambda S: S["m"]["tx_total"],
         "＝ 交易明细 total 合计；推导见③桥 C"),
    ], note_col)


# ----------------------------------------------------------------- ③ bridges

def sheet_bridge(wb, fam, ctx):
    stores = fam.stores
    Ln = RF.Line
    A = lambda S, k: S["m"]["bridge_a"][k]
    bridge_a = [
        Ln("订单行：商品价款", lambda S: A(S, "o_sales"), "Pedido · ventas de productos"),
        Ln("订单行：运费", lambda S: A(S, "o_ship"), "Pedido · créditos de envío"),
        Ln("订单行：礼品包装", lambda S: A(S, "o_wrap"), "Pedido · créditos por envoltorio de regalo"),
        Ln("订单行：促销折扣", lambda S: A(S, "o_promo"), "Pedido · descuentos promocionales"),
        Ln("＝ 订单行合计（不含 IVA）", lambda S: A(S, "o"), fam.per_store(
            lambda S: "%d 张订单、%d 行" % (S["m"]["orders_n"], S["m"]["order_lines"])), style="subtotal"),
        Ln("退款行：商品价款", lambda S: A(S, "r_sales"), "Reembolso · ventas de productos"),
        Ln("退款行：运费", lambda S: A(S, "r_ship"), "Reembolso · créditos de envío"),
        Ln("退款行：礼品包装", lambda S: A(S, "r_wrap"), ""),
        Ln("退款行：促销折扣冲回", lambda S: A(S, "r_promo"), "Reembolso · descuentos promocionales"),
        Ln("＝ 退款行合计（不含 IVA）", lambda S: A(S, "r"), fam.per_store(
            lambda S: "%d 张订单、%d 行" % (S["m"]["refund_orders"], S["m"]["refund_lines"])), style="subtotal"),
        Ln("②净销售收入", lambda S: S["m"]["net_sales"]),
        Ln("残差（应为 0）", lambda S: A(S, "resid"), style="residual"),
    ]
    bridge_b = []
    pids = sorted(set(P["id"] for S in stores for P in S["m"]["bridge_b"]))
    for pid in pids:
        def B(S, k, pid=pid):
            for x in S["m"]["bridge_b"]:
                if x["id"] == pid:
                    return x[k]
            return 0.0

        def full(pid=pid):
            return all(x["full"] for S in stores for x in S["m"]["bridge_b"] if x["id"] == pid)
        a_, b_ = pid.split("_")
        lines = [
            Ln("账期 %s 至 %s：费用发票 I（含 IVA）" % (a_, b_), lambda S, B=B: B(S, "inv_i"),
               fam.per_store(lambda S, B=B: "%d 张" % B(S, "n_i"))),
            Ln("　　减：贷记单 E（含 IVA）", lambda S, B=B: -B(S, "inv_e"),
               "退款时退还的费用（Ajuste a tarifas de vendedores）"),
            Ln("　　＝ 发票净额", lambda S, B=B: B(S, "inv"), style="subtotal"),
            Ln("　　交易明细同期的平台费用（含 IVA，不含广告）", lambda S, B=B: B(S, "tx_fees"),
               "佣金、FBA、仓储、入仓/移除、订阅，按 fecha/hora 换算成 UTC 落在账期内的行"),
        ]
        if full():
            lines.append(Ln("　　残差（应为 0）", lambda S, B=B: B(S, "diff"), style="residual"))
        else:
            lines.append(Ln("　　差额（部分期间，仅供参考）", lambda S, B=B: B(S, "diff"),
                            "账期跨月：交易明细只含本月这一段，差额不说明问题；完整比较见相邻月份的报表",
                            style="pending"))
        lines.append(Ln("　　参考：同期广告费（含 IVA，另行开票）", lambda S, B=B: B(S, "tx_ads"),
                        "广告费不在卖家费用发票里"))
        bridge_b += lines
    if not pids:
        bridge_b.append(Ln("本月没有下载到费用发票 ⏳", lambda S: 0.0,
                           "不影响报表数字（费用按交易明细计入）；下月或重新收集（--collect-only）时补齐",
                           style="pending"))
    if any(S["m"]["invoice_payments"] for S in stores):
        bridge_b.append(Ln("参考：本月的付款补充（Complemento de pago，类型 P）",
                           lambda S: sum(d["paid"] for d in S["m"]["invoice_payments"]),
                           fam.per_store(lambda S: "%d 张" % len(S["m"]["invoice_payments"]))))
    m = lambda S, k: S["m"][k]
    before = lambda S: _r(m(S, "contrib") + m(S, TC.IVA_OUT) + m(S, TC.IVA_WITHHELD) + m(S, "fee_iva")
                          + m(S, TC.ISR) + m(S, TC.IVA_ADJ))
    bridge_c = [
        Ln("经营贡献毛利（不含 IVA，见②）", lambda S: m(S, "contrib")),
        Ln("加：销项 IVA", lambda S: m(S, TC.IVA_OUT), "Amazon 向买家收、付给卖家，卖家向 SAT 申报"),
        Ln("减：Amazon 代扣 IVA", lambda S: m(S, TC.IVA_WITHHELD), "预缴税款"),
        Ln("减：平台费用所含进项 IVA", lambda S: m(S, "fee_iva")),
        Ln("减：Amazon 代扣 ISR（上月的，本月扣）", lambda S: m(S, TC.ISR),
           fam.per_store(lambda S: "、".join("%s %s" % (d, fmt(a)) for d, a in S["m"]["isr_rows"]) or "本月没有")),
        Ln("加：IVA 调整", lambda S: m(S, TC.IVA_ADJ)),
        Ln("＝ 本月经营带来的余额增加（转账前）", before, style="subtotal"),
        Ln("减：转账到银行（Trasferir）", lambda S: m(S, TC.PAYOUT)),
        Ln("＝ 推算的余额净变动", lambda S: m(S, "balance_calc"), style="subtotal"),
        Ln("交易明细 total 合计", lambda S: m(S, "tx_total")),
        Ln("残差（应为 0）", lambda S: m(S, "bridge_c_resid"), style="residual"),
    ]
    sv = lambda S, k: S["m"]["summary"].get(k, 0.0) or 0.0
    gc = lambda S, k: S["m"]["groups_csv"][k]
    names = {"ingresos": "Ingresos（收入）", "gastos": "Gastos（费用，含 IVA）",
             "transferencias": "Transferencias（转账）", "impuesto": "Impuesto（税）"}
    sum_note = fam.per_store(lambda S: {"ok": "已读取", "missing": "没有摘要 PDF",
                                        "unreadable": "PDF 读不出（缺工具）", "english": "PDF 是英文",
                                        "unparsed": "版式不认识"}.get(S["m"]["summary_status"], ""))
    bridge_d = []
    for k in TC.SUMMARY_GROUPS:
        bridge_d += [
            Ln("%s：交易明细重算" % names[k], lambda S, k=k: gc(S, k),
               "FBA 库存赔付按正负分别归 Ingresos / Gastos，与 Amazon 摘要一致" if k == "ingresos" else ""),
            Ln("　　Amazon 摘要 PDF", lambda S, k=k: sv(S, k), sum_note),
            Ln("　　残差（应为 0）", lambda S, k=k: _r(gc(S, k) - sv(S, k)), style="residual"),
        ]
    bridge_d += [
        Ln("参考：摘要的 ISR 代扣 vs 交易明细无说明的大额 Ajuste", lambda S: _r(sv(S, "isr") - m(S, TC.ISR)),
           "应为 0：证明 ISR 行识别正确"),
        Ln("参考：摘要的佣金退还（Reembolsos de tarifa de venta）", lambda S: sv(S, "fee_refunds"), ""),
        Ln("参考：摘要的退款管理费（Tarifas de administración de reembolso）", lambda S: sv(S, "refund_admin"),
           "交易明细里已并入退款行的佣金净额"),
    ]
    cert = lambda S, k: (S["m"]["cert"] or {}).get(k, 0.0) or 0.0
    bridge_e = [
        Ln("凭证：代扣 IVA（ImpuestoRet 002）", lambda S: -cert(S, "iva"), _cert_note(fam)),
        Ln("交易明细：代扣 IVA（impuesto de retenciones en la plataforma）", lambda S: m(S, TC.IVA_WITHHELD)),
        Ln("差额（时间差）", lambda S: _r(-cert(S, "iva") - m(S, TC.IVA_WITHHELD)) if S["m"]["cert"] else 0.0,
           "凭证按付款日（FechaServ），交易明细按入账日；月初月末的订单落在不同月份。超过 1% 时 ⑩ 提示", style="pending"),
        Ln("凭证：代扣 ISR（ImpuestoRet 001）—— 本月的 ISR", lambda S: -cert(S, "isr"),
           "下月 11 日前后从余额扣（下月交易明细的无说明 Ajuste）。ISR ÷ 交易金额：%s（应为 2.5%%）"
           % fam.per_store(lambda S: pct(cert(S, "isr") / cert(S, "operations"), 2)
                           if cert(S, "operations") else "-")),
        Ln("凭证：交易金额（MontoTotOperacion，不含 IVA）", lambda S: cert(S, "operations"),
           fam.per_store(lambda S: "%d 笔" % ((S["m"]["cert"] or {}).get("services") or 0))),
        Ln("交易明细：订单行合计（不含 IVA，见桥 A）", lambda S: m(S, "order_net_sales"),
           "参考：口径不同（凭证按付款日、不扣退款；明细按入账日），差额不需要为 0"),
    ]
    fam.bridge_sheet(wb, SHEETS[2], "平账勾稽表（%s）" % ctx["period_cn"],
                     "每座桥最后一行是残差，应为 0。单位：%s。桥 A 不含 IVA；桥 B、C、D 用含 IVA 的金额（与 Amazon 的文件一致）。"
                     % CUR, [
                         ("桥 A　订单与退款行 → 净销售收入", bridge_a),
                         ("桥 B　平台费用发票（CFDI，Factura de la tarifa de vendedor）→ 交易明细", bridge_b),
                         ("桥 C　经营贡献毛利 → Amazon 余额净变动（现金）", bridge_c),
                         ("桥 D　Amazon 摘要 PDF（Resumen）→ 交易明细（独立控制总数）", bridge_d),
                         ("桥 E　扣缴凭证（Certificado de retención）→ 交易明细的代扣税", bridge_e)])


# ----------------------------------------------------------------- ④ fees

TEXT10 = F(10)
FEE_COLS = [RF.Col("店铺", 20, font=TEXT10), RF.Col("类型（Amazon 原文 tipo）", 30, font=TEXT10),
            RF.Col("说明 / 金额列（Amazon 原文）", 52, font=F(9, False, K.C_GREY)),
            RF.Col("中文", 30, font=TEXT10), RF.Col("归入", 26, font=TEXT10),
            RF.Col("归类", 16, font=TEXT10), RF.Col("笔数", 8, font=TEXT10),
            RF.Col("金额（Amazon 原数）", 16, RF.MONEY), RF.Col("其中 IVA", 14, RF.MONEY),
            RF.Col("不含 IVA", 16, RF.MONEY)]
_LINE_ORDER = [x.key for x in PL_ROWS] + list(TC.TAX_LINES) + list(TC.CASH_LINES)


def sheet_fees(wb, fam, ctx):
    ws = fam.table_sheet(wb, SHEETS[3], "平台费用明细（%s）" % ctx["period_cn"],
                         "本月交易明细里出现的每一种交易类型与金额列。订单/退款行按金额列分行；其他行按说明分行。"
                         "“金额”是 Amazon 文件里的原数（费用含 IVA），“不含 IVA”是进②的金额。单位 %s。" % CUR, FEE_COLS)
    r = 5
    for S in fam.stores:
        L = S["ledger"]
        if not len(L):
            ws.cell(row=r, column=1, value=S["store"])
            ws.cell(row=r, column=3, value="本月没有交易").font = F(10, False, K.C_GREY)
            r += 2
            continue
        # order/refund lines: one row per column; others: one per description kind
        is_or = L["_key"].isin(("order", "refund"))
        L = L.assign(_what=L["_col"].where(is_or | (L["line"] == TC.UNREGISTERED),
                                           L[T_DESC].map(_short_desc)))
        g = L.groupby([T_TYPE, "_what", "_zh", "line"]).agg(
            n=("_row", "nunique"), amt=("amt", "sum"), iva=("iva", "sum"), net=("net", "sum")).reset_index()
        g["_order"] = g["line"].map(lambda k: _LINE_ORDER.index(k) if k in _LINE_ORDER else 99)
        for _, x in g.sort_values(["_order", T_TYPE, "_what"]).iterrows():
            what = x["_what"]
            zh = x["_zh"]
            if x["line"] != TC.UNREGISTERED and what in TC.COL_ZH:
                zh = "%s · %s" % (zh, TC.COL_ZH[what])
            cells = fam.write_row(ws, r, FEE_COLS, [
                S["store"], x[T_TYPE], what, zh, pl_label(x["line"]),
                TC.LINE_CATEGORY.get(x["line"], ""), int(x["n"]),
                _r(x["amt"]), _r(_r(x["amt"]) - _r(x["net"])), _r(x["net"])])
            if x["line"] == TC.UNREGISTERED:
                for c in cells[:7]:
                    c.font = F(10, False, K.C_RED)
            r += 1
        m = S["m"]
        ws.cell(row=r, column=1, value="%s 合计" % S["store"]).font = F(10, True)
        for j, v in ((8, m["tx_total"]), (9, m["fee_iva"]), (10, m["tx_total"] - m["fee_iva"])):
            c = fam.money_cell(ws.cell(row=r, column=j), v)
            c.font = F(10, True)
            c.fill = fill(K.C_TOTAL)
        ws.cell(row=r, column=3, value="＝ 交易明细 total 合计（Amazon 余额净变动）").font = F(9, False, K.C_GREY)
        r += 2
    ws.freeze_panes = "B5"


def _short_desc(d):
    """A fee description without the varying tail (a reimbursement's reason stays)."""
    d = str(d or "").strip()
    if not d:
        return "（无说明）"
    return d if len(d) <= 120 else d[:117] + "…"


# ----------------------------------------------------------------- ⑤ returns

RET_COLS = [RF.Col("店铺", 18), RF.Col("SKU", 18), RF.Col("商品", 44), RF.Col("退款行数", 9),
            RF.Col("退款件数", 9), RF.Col("本月售出件数", 10)] + \
           [RF.Col(h, w, RF.MONEY) for h, w in (
               ("退款商品价款", 14), ("退还运费", 12), ("冲回促销", 12), ("佣金退还（净）", 14),
               ("FBA 费退还", 12), ("本月商品销售收入", 15))] + \
           [RF.Col("退款率（金额）", 12, K.PCT)]


def sheet_returns(wb, fam, ctx):
    summary_rows = [
        ("退款行 / 退款订单 / 退款件数", lambda S: "%d / %d / %d" % (
            S["m"]["refund_lines"], S["m"]["refund_orders"], S["m"]["refund_units"])),
        ("退款商品价款（不含 IVA）", lambda S: fmt(S["m"]["pl"][TC.REFUND])),
        ("退款率（退款商品价款 ÷ 商品销售收入）", lambda S: pct(-S["m"]["pl"][TC.REFUND] / S["m"]["pl"][TC.SALES])
         if S["m"]["pl"][TC.SALES] else "-"),
        ("摘要：佣金退还（含 IVA）", lambda S: fmt(S["m"]["summary"].get("fee_refunds"))
         if "fee_refunds" in S["m"]["summary"] else "（没有摘要 PDF）"),
        ("摘要：退款管理费（含 IVA）", lambda S: fmt(S["m"]["summary"].get("refund_admin"))
         if "refund_admin" in S["m"]["summary"] else "（没有摘要 PDF）"),
    ]
    head = 4 + len(summary_rows) * len(fam.stores) + 2
    ws = fam.table_sheet(wb, SHEETS[4], "退货退款分析（%s）" % ctx["period_cn"],
                         "上：本月退款合计。下：按 SKU 的退款（交易明细的 Reembolso 行，不含 IVA）。佣金退还是 Amazon 退还的佣金"
                         "扣除退款管理费后的净额。Amazon 的交易报表不含退货原因。", RET_COLS, header_at=head)
    r = 4
    for S in fam.stores:
        for label, fn in summary_rows:
            ws.cell(row=r, column=1, value=S["store"])
            ws.cell(row=r, column=2, value=label).font = F(10, True)
            ws.cell(row=r, column=4, value=fn(S)).border = BOX
            r += 1
    r = head + 1
    for S in fam.stores:
        g = S["returns"]
        if not len(g):
            ws.cell(row=r, column=1, value=S["store"])
            ws.cell(row=r, column=3, value="本月没有退款").font = F(10, False, K.C_GREY)
            r += 2
            continue
        for sku, x in g.iterrows():
            fam.write_row(ws, r, RET_COLS, [
                S["store"], sku, _s(x["name"]), int(x["lines"]), int(x["units"]), int(x["sold"]),
                _r(x[TC.REFUND]), _r(x[TC.SHIPPING]), _r(x[TC.PROMO]), _r(x[TC.COMMISSION]),
                _r(x[TC.FBA_FEE]), _r(x["sales"]),
                (-x[TC.REFUND] / x["sales"]) if x["sales"] else None])
            r += 1
        r += 1
    ws.freeze_panes = ws.cell(row=head + 1, column=3).coordinate


# ----------------------------------------------------------------- ⑥ reimbursements

REIMB_COLS = [RF.Col("店铺", 18), RF.Col("入账日", 12), RF.Col("SKU", 20), RF.Col("原因（Amazon 原文）", 36),
              RF.Col("原因（中文）", 34), RF.Col("金额", 14, RF.MONEY)]


def sheet_reimbursements(wb, fam, ctx):
    ws = fam.table_sheet(wb, SHEETS[5], "退货质检明细 → FBA 库存赔付明细（%s）" % ctx["period_cn"],
                         "Amazon 不提供逐件退货质检报表。本页列 FBA 库存赔付（Reembolso de inventario de Logística de "
                         "Amazon）：仓内丢失/损坏、退货未退回等，合计即②“其他收入”。负数是 Amazon 冲回此前的赔付"
                         "（如商品后来找到或退回了）。", REIMB_COLS)
    r = 5
    for S in fam.stores:
        rows = S["reimb"]
        if not rows:
            ws.cell(row=r, column=1, value=S["store"])
            ws.cell(row=r, column=4, value="本月没有库存赔付").font = F(10, False, K.C_GREY)
            r += 2
            continue
        for x in rows:
            fam.write_row(ws, r, REIMB_COLS, [S["store"], x["day"], x["sku"], x["reason"], x["zh"], x["amt"]])
            r += 1
        by = {}
        for x in rows:
            by.setdefault(x["zh"], [0, 0.0])
            by[x["zh"]][0] += 1
            by[x["zh"]][1] += x["amt"]
        for zh, (n, a) in sorted(by.items(), key=lambda kv: -kv[1][1]):
            ws.cell(row=r, column=4, value="小计（%d 笔）" % n).font = F(9, False, K.C_GREY)
            ws.cell(row=r, column=5, value=zh).font = F(9, False, K.C_GREY)
            fam.money_cell(ws.cell(row=r, column=6), a)
            r += 1
        ws.cell(row=r, column=1, value="%s 合计" % S["store"]).font = F(10, True)
        c = fam.money_cell(ws.cell(row=r, column=6), S["m"]["pl"][TC.OTHER_INCOME])
        c.font = F(10, True)
        c.fill = fill(K.C_TOTAL)
        ws.cell(row=r, column=4, value="＝ ②其他收入（FBA 库存赔付）").font = F(9, False, K.C_GREY)
        r += 2
    ws.freeze_panes = "A5"


# ----------------------------------------------------------------- ⑦ SKU

_SKU_FORMULA = dict(fmt=RF.MONEY, border=False, fill=K.C_CALC)
SKU_COLS = ([RF.Col("店铺", 18), RF.Col("SKU", 18), RF.Col("商品", 44), RF.Col("售出件数", 9),
             RF.Col("退款件数", 9)]
            + [RF.Col(h, w, RF.MONEY) for h, w in (
                ("商品销售收入", 14), ("退货退款", 12), ("运费与礼品包装", 12), ("促销折扣", 12),
                ("佣金", 12), ("FBA 配送费", 12), ("其他费用", 12), ("库存赔付", 12),
                ("贡献（未扣成本）", 14))]
            + [RF.Col("单位成本（录入）", 12, RF.MONEY, fill=K.C_INPUT),
               RF.Col("商品成本", 14, formula="=-O{r}*D{r}", **_SKU_FORMULA),
               RF.Col("SKU 毛利", 14, formula="=N{r}+P{r}", **_SKU_FORMULA)])


def sheet_sku(wb, fam, ctx):
    ws = fam.table_sheet(wb, SHEETS[6], "SKU 盈利分析（%s）" % ctx["period_cn"],
                         "按交易明细的 SKU 汇总（不含 IVA）。仓储、广告、订阅等没有 SKU 的店铺级费用单列一行"
                         "“(无 SKU · 店铺级费用)”。黄色列填单位采购成本（不含 IVA），右边两列自动算出 SKU 毛利。", SKU_COLS)
    r = 5
    other = [TC.STORAGE, TC.STORAGE_LT, TC.LOGISTICS, TC.ADS, TC.SUBSCRIPTION, TC.OTHER_FEES, TC.UNREGISTERED]
    for S in fam.stores:
        g = S["sku"]
        for sku, x in g.iterrows():
            money = [x[TC.SALES], x[TC.REFUND], x[TC.SHIPPING] + x[TC.GIFTWRAP], x[TC.PROMO],
                     x[TC.COMMISSION], x[TC.FBA_FEE], sum(x[k] for k in other), x[TC.OTHER_INCOME],
                     x["contrib"]]
            fam.write_row(ws, r, SKU_COLS, [S["store"], sku, _s(x["name"]), int(x["units"]), int(x["returns"])]
                          + [_r(v) for v in money] + [0, None, None])
            r += 1
        r += 1
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = "A4:Q%d" % max(4, r - 1)


# ----------------------------------------------------------------- ⑧ trend

TREND_COLS = [RF.Col("店铺", 22), RF.Col("周（起始日）", 16), RF.Col("售出件数", 10),
              RF.Col("商品销售收入", 16, RF.MONEY), RF.Col("退货退款", 14, RF.MONEY),
              RF.Col("平台费用", 16, RF.MONEY), RF.Col("贡献（未扣成本）", 16, RF.MONEY),
              RF.Col("费用率", 10, K.PCT, border=False)]
HIST_COLS = [RF.Col("店铺", 22, border=False), RF.Col("月份", 16, border=False),
             RF.Col("售出件数", 10, border=False)] + TREND_COLS[3:]


def _period_values(L, t):
    pl = L[L["line"].isin(TC.PL_LINES)]
    o = t[t[T_TYPE].map(TC.norm) == ORDER_TIPO]
    sales = float(pl.loc[pl["line"] == TC.SALES, "net"].sum())
    fees = float(pl.loc[pl["line"].isin(FEE_ROWS), "net"].sum())
    return {"items": int(o[T_QTY].sum()), "sales": _r(sales),
            "refund": _r(pl.loc[pl["line"] == TC.REFUND, "net"].sum()), "fees": _r(fees),
            "contrib": _r(pl["net"].sum())}


def sheet_trend(wb, fam, ctx):
    ws = fam.table_sheet(wb, SHEETS[7], "月度趋势（%s）" % ctx["period_cn"],
                         "本月按周（周一开始，按入账日 GMT%+d，不含 IVA）。往月趋势需要往月的下载：同一台机器上有更早月份的"
                         "Amazon 下载目录时，下方自动追加各月合计。" % MONTH_UTC_OFFSET, TREND_COLS)
    r = 5
    for S in fam.stores:
        L, t = S["ledger"], S["tx"]
        if not len(L):
            continue
        wk = lambda d: (d - pd.to_timedelta(d.dt.weekday, unit="D")).dt.date
        L = L.assign(_wk=wk(L["_day"]))
        t = t.assign(_wk=wk(t["_day"]))
        for w, x in L.groupby("_wk"):
            v = _period_values(x, t[t["_wk"] == w])
            fam.write_row(ws, r, TREND_COLS, [
                S["store"], str(w), v["items"], v["sales"], v["refund"], v["fees"], v["contrib"],
                (-v["fees"] / v["sales"]) if v["sales"] else 0.0])
            r += 1
        r += 1
    hist = ctx.get("history") or {}
    if hist:
        r = RF.section(ws, r + 1, "往月（来自本机更早的 Amazon 下载）", 8)
        for store, months in hist.items():
            for mon, v in sorted(months.items()):
                fam.write_row(ws, r, HIST_COLS, [
                    store, mon, v["items"], v["sales"], v["refund"], v["fees"], v["contrib"],
                    (-v["fees"] / v["sales"]) if v["sales"] else 0.0])
                r += 1


def history(stores):
    """Monthly totals of earlier months, from the store's other Amazon folders."""
    out = {}
    for S in stores:
        root = os.path.dirname(os.path.abspath(S["folder"]))
        seen = {}
        for d in sorted(glob.glob(os.path.join(root, "*_m??????"))):
            mon = d[-6:-2] + "-" + d[-2:]
            if mon >= S["month"] or mon in seen:
                continue
            p = _first(d, "%s*Transaction.csv" % periods.file_stem(mon))
            if not p:
                continue
            try:
                t = load_transactions(p)
                seen[mon] = _period_values(ledger(t), t)
            except Exception:
                continue
        if seen:
            out[S["store"]] = dict(sorted(seen.items())[-6:])
    return out


# ----------------------------------------------------------------- ⑨ inventory

INV_COLS = [RF.Col("店铺", 18), RF.Col("SKU", 18), RF.Col("ASIN", 12), RF.Col("商品", 40),
            RF.Col("售价", 12, RF.MONEY), RF.Col("可售", 8), RF.Col("在途入仓", 9), RF.Col("预留", 8),
            RF.Col("库龄 0–90 天", 10), RF.Col("91–180", 8), RF.Col("181–270", 8), RF.Col("271–365", 8),
            RF.Col("365 天以上", 9), RF.Col("近 30 天发货", 10), RF.Col("近 90 天发货", 10),
            RF.Col("本月售出（交易明细）", 12), RF.Col("可售天数", 9),
            RF.Col("预计下月仓储费", 13, RF.MONEY), RF.Col("预计长期仓储费", 13, RF.MONEY),
            RF.Col("库存货值（售价）", 15, RF.MONEY), RF.Col("判断", 14)]


def _n(x, k):
    try:
        v = float(x.get(k))
        return 0.0 if v != v else v
    except (TypeError, ValueError):
        return 0.0


def sheet_inventory(wb, fam, ctx):
    ws = fam.table_sheet(wb, SHEETS[8], "库存与动销分析（FBA，亚马逊物流）",
                         "库存快照取自下载当时（不是月末），近 30/90 天发货、可售天数、预计仓储费是 Amazon 给的。"
                         "滞销：有库存、近 90 天没发货；断货风险：近 30 天有发货、可售不到 14 天；"
                         "超龄：有 181 天以上的库存（长期仓储费风险）。", INV_COLS)
    r = 5
    for S in fam.stores:
        s = S["stock"]
        if s is None:
            ws.cell(row=r, column=1, value=S["store"])
            ws.cell(row=r, column=4, value="本月没有下载到库存快照").font = F(10, False, K.C_RED)
            r += 2
            continue
        sold = S["sku"]["units"] if len(S["sku"]) else pd.Series(dtype=int)
        for _, x in s.iterrows():
            avail = _n(x, "available")
            t30, t90 = _n(x, "units-shipped-t30"), _n(x, "units-shipped-t90")
            dos = x.get("days-of-supply")
            dos = None if dos is None or dos != dos else float(dos)
            old = _n(x, "inv-age-181-to-270-days") + _n(x, "inv-age-271-to-365-days") + \
                _n(x, "inv-age-365-plus-days")
            verdict = ("滞销" if avail > 0 and not t90 else
                       "断货风险" if t30 and dos is not None and dos < 14 else
                       "超龄" if old > 0 else "无库存" if not avail else "")
            reserved = _n(x, "Reserved FC Processing") + _n(x, "Reserved Customer Order")
            price = _n(x, "your-price")
            cells = fam.write_row(ws, r, INV_COLS, [
                S["store"], _s(x.get("sku")), _s(x.get("asin")), _s(x.get("product-name")), price, int(avail),
                int(_n(x, "inbound-quantity")), int(reserved),
                int(_n(x, "inv-age-0-to-90-days")), int(_n(x, "inv-age-91-to-180-days")),
                int(_n(x, "inv-age-181-to-270-days")), int(_n(x, "inv-age-271-to-365-days")),
                int(_n(x, "inv-age-365-plus-days")), int(t30), int(t90),
                int(sold.get(x.get("sku"), 0)) if len(sold) else 0, dos,
                _n(x, "estimated-storage-cost-next-month"), _n(x, "estimated-ltsf-next-charge"),
                avail * price, verdict])
            if verdict in ("滞销", "断货风险", "超龄"):
                cells[-1].font = F(10, True, K.C_AMBER)
            r += 1
        r += 1
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = "A4:U%d" % max(4, r - 1)


# ----------------------------------------------------------------- ⑩ gaps

def store_meta_missing(S):
    return RF.meta_missing(S["meta"])


def sheet_gaps(wb, fam, ctx):
    fam.gaps_sheet(wb, SHEETS[9], "数据缺口与待办（%s）" % ctx["period_cn"], ctx["checks"],
                   notes=[("报税时要注意（不影响本报表的数字）",
                           [("提示 %d" % i, q) for i, q in enumerate(TAX_NOTES, 1)])])


# ======================================================================= entry

def build_context(stores, month, build_date=None):
    y, mo = periods.parse_month(month)
    return {"month": month, "period_cn": "%d 年 %d 月" % (y, mo),
            "who": stores[0]["store"] if len(stores) == 1 else "%d 家店铺汇总" % len(stores),
            "built": build_date or datetime.datetime.now().strftime("%Y-%m-%d %H:%M")}


def build(pairs, month, out_path, strict=False, build_date=None):
    """pairs: [(store, folder)]. Returns a result dict (outcome, checks, out)."""
    stores = [derive(load_store(s, f, month)) for s, f in pairs]
    ctx = build_context(stores, month, build_date)
    ctx["checks"] = validate(stores)
    ctx["history"] = history(stores)
    bad = RF.failed_errors(ctx["checks"])
    if strict and bad:
        return {"outcome": "strict_aborted", "checks": ctx["checks"], "out": out_path}
    build_workbook(stores, ctx, out_path)
    return {"outcome": "checks_failed" if bad else "ok", "checks": ctx["checks"], "out": out_path,
            "metrics": {S["store"]: S["m"] for S in stores}}
