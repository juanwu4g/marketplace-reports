# -*- coding: utf-8 -*-
"""
Liverpool monthly financial workbook — a sister of the other platforms'.

Same ten sheets, same look, same check levels: the layout is written by
core/report_family.py on the frame of core/report_kit.py (docs/REPORT_FAMILY.md);
the content — lines, bridges, columns, rules — is Liverpool's. Accounting rules are
in platforms/liverpool/README.md ("Accounting rules"); the short version:

* **Basis: Liverpool's ledger.** Every money line is a line of the supplier ledger
  ("Saldos Proveedor") posted in the month (Fecha de contabilización). Liverpool
  posts a sale on the day it is ordered and paid (measured: 122 of 122 sales of
  2026-09 posted on their order date, about 6 days before delivery), its commission
  and withholding with it; a cancellation or a return is a reversal document on
  its own day; fees post on their charge day. Like MercadoLibre's bill basis, the
  platform's own posting date decides the month. Sales posted but not delivered at
  month end are a memo line; orders are reconciled in ③ A.
* **With IVA**, as Liverpool charges and pays, like MercadoLibre Mexico's official
  (bill) basis for the same company. IVA in the figures is a memo line.
* **Withheld taxes** (IVA 8 %, ISR 2.5 % of the price net of IVA) are a prepayment
  of the seller's taxes: a row below the profit, not a cost — as on MercadoLibre.
* **Every ledger line kind is registered** (catalog.py); an unknown one is an error
  check and its amount goes to "未登记科目 ⚠".
* **Product cost** is not in Liverpool's data: the shared SKU cost table
  (core/sku_costs.py) or yellow input cells.

Inputs: one download folder per store (platforms/liverpool/run_downloads.py), the
files named in readers.py. Customer columns of the order list are dropped on load.
"""
import collections
import datetime
import glob
import io
import json
import os
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from core import report_family as RF
from core import report_kit as K
from core import sku_costs as SC
from core.report_kit import F, fill, wrap, BOX
from platforms.liverpool import catalog as C
from platforms.liverpool import periods
from platforms.liverpool import readers as R

SHEETS = K.SHEET_NAMES
PLATFORM, COUNTRY = "LV", "MX"  # core/markets.py codes (and the SKU cost table's 平台 / 国家)
CUR = "MXN"
DECIMALS = 2
MNY = K.money_format(DECIMALS)
IVA_RATE = 0.16
# The cost table holds costs net of IVA; this report is with IVA (core/sku_costs.apply factor).
COST_FACTOR = 1 + IVA_RATE
COST_IVA_NOTE = "（成本表的成本不含 IVA，本报表 × 1.16 换算为含 IVA）"
COST_UNITS_BASIS = "本月入账的销售件数 − 本月入账的退货件数（退回的商品回到库存）"
TOL = 0.05                      # MXN; rounding tolerance per bridge
NO_SKU = "(无 SKU · 店铺级费用)"
UNLINKED = "(未关联到订单的销售单据)"
IMPLEMENTED = True

# ② fixed rows (core/report_family.PLRow): key, 中文, 平台原文, kind, 说明, formula.
_R = RF.PLRow
PL_ROWS = [
    _R(C.SALES, "商品销售收入（买家实付）", "F/ · Venta (YV)", "d",
       "Liverpool 在下单付款当天入账的销售价款（含 IVA），按入账日（Fecha de contabilización）归月；"
       "取消或退货另有冲回单据"),
    _R(C.PROMO, "加：Liverpool 出资的促销补偿", "F/ · Descuento (YV, C2)", "d",
       "Liverpool 自己的促销让买家少付的部分，由 Liverpool 补给卖家；佣金按原价计"),
    _R("revenue", "商品销售收入合计（按原价）", "", "sum", "＝ 买家实付 ＋ 促销补偿", (C.SALES, C.PROMO)),
    _R(C.REFUND, "减：取消与退货退款", "DEV · Devolución (YD)", "d",
       "本月入账的冲回单据：已入账销售的取消和退货（逐单见⑤）"),
    _R("net_sales", "净销售收入", "", "f", "＝ 商品销售收入合计 − 退货退款", ("revenue", C.REFUND)),
    _R(C.COMMISSION, "减：销售佣金（净）", "Comisión % DILISA − comisión devuelta", "d",
       "每笔销售按原价计的佣金；退货时 Liverpool 退还的佣金已抵减"),
    _R(C.CARD_FEE, "减：外部银行卡附加佣金", "Sobrecomis % T.Exter", "d",
       "买家用非 Liverpool 银行卡付款时的附加佣金"),
    _R(C.LOGISTICS, "减：发货物流费", "Rec Logística", "d", "按单收取的配送费（物流费明细逐单可查，见③桥 D）"),
    _R(C.LOGISTICS_EXTRA, "减：补充物流费", "Balance logístico · Cargo logístico complementario", "d",
       "尺寸/重量复核后补收的运费"),
    _R(C.LOGISTICS_ADJ, "加/减：物流费调整（往月）", "Ajuste log", "d", "往月物流费的补收或退还"),
    _R(C.MSI, "减：无息分期（MSI）成本", "Ajuste MSI", "d", "买家分期付款（Meses sin intereses）由卖家承担的成本"),
    _R(C.SUBSCRIPTION, "减：订阅费（Cuota digital）", "Suscripción MKP", "d", "每月固定"),
    _R(C.OTHER, "减：其他调整", "", "d", ""),
    _R(C.UNREGISTERED, "减：未登记科目 ⚠", "(catalog.py 里没有)", "d",
       "账本里出现了字典里没有的行。金额照常计入以保证勾稽，但需人工确认科目归属，见 ⑩ 页"),
    _R("fee_total", "平台费用小计", "", "sum", "", RF.Span(C.COMMISSION, C.UNREGISTERED)),
    _R("contrib", "经营贡献毛利（未扣商品成本）", "", "f", "＝ 净销售收入 − 平台费用小计",
       ("net_sales", "fee_total")),
    _R("cogs", "减：商品采购成本", "", "in",
       "⚠ 平台数据不含此项：在【⑦ SKU盈利分析】填单位成本，或直接在此填总额"),
    _R("profit", "扣除商品成本后的贡献", "", "f", "＝ 经营贡献毛利 − 商品采购成本", ("contrib", "-cogs")),
    _R("withheld", "平台代扣代缴税金（预缴税款，不计入贡献）", "Q6 IVA retenido · Q8 ISR retenido", "d",
       "Liverpool 按不含 IVA 价款代扣 8% IVA、2.5% ISR，退货时退还；是预缴税款，申报时抵减，所以不从贡献中扣（与美客多相同）"),
    _R("profit_after_wh", "【敏感性】若代扣税不能抵减：扣除后的贡献", "", "f",
       "＝ 扣除商品成本后的贡献 − 代扣税金；仅供比较", ("profit", "withheld")),
]
PL_LABELS = dict((x.key, x.label) for x in PL_ROWS)
FEE_ROWS = C.FEE_LINES
LOGISTICS_ROWS = (C.LOGISTICS, C.LOGISTICS_EXTRA, C.LOGISTICS_ADJ)

ACCOUNTING_POLICIES = [
    ("收入确认", "以 Liverpool 账本的入账日为准：Liverpool 在买家下单付款当天入账销售（比送达早约一周），"
     "取消、退货在发生当天以冲回单据入账。这与美客多账单基准“以平台的账为准”相同；"
     "月末已入账、尚未送达的销售金额见②备查项（按送达确认收入时，这部分属于下个月）。"),
    ("含 IVA", "全部金额含 IVA，与 Liverpool 账本、付款和美客多墨西哥账单基准报表的口径相同。"
     "损益里所含的 IVA（销项 − 进项）要向 SAT 缴纳，不是利润，见②备查项。"),
    ("促销", "Liverpool 出资的促销（买家少付、Liverpool 补给卖家）是这笔销售对价的一部分：买家实付与补偿"
     "合起来就是原价，计入收入。"),
    ("退货", "退货冲减收入；Liverpool 退还的佣金冲减佣金；退还的代扣税冲减代扣税金。"),
    ("平台费用", "佣金、银行卡附加佣金随每笔销售入账；物流费、补充物流费、订阅费、调整按 Liverpool 的入账日计入当月。"
     "这些费用记在 Liverpool 的另一个账户上，在付款时从应付卖家的款项里抵扣。"),
    ("代扣税金", "Liverpool 代扣的 IVA（8%）和 ISR（2.5%）是卖家的预缴税款，可在申报时抵减：单列在利润之后，"
     "不计入贡献；最后一行给出“若不能抵减”时的数字作比较。"),
    ("商品成本与存货", "平台数据不含采购成本：在②或⑦录入，或配置 SKU 成本表。未使用 Liverpool Fulfillment（FBL），"
     "没有平台库存。"),
]
TAX_NOTES = [
    "代扣税：Liverpool 每笔销售代扣 IVA 8% 与 ISR 2.5%（按不含 IVA 价款），本月合计见②“平台代扣代缴税金”和备查项；"
    "申报时作为已预缴税款抵减，以 Liverpool 开具的代扣凭证（CFDI de retenciones）为准。",
    "进项 IVA：Liverpool 的费用发票（Notas de Cargo：MCOM 佣金、MLOG 物流、MCUO 订阅）按周开具，"
    "月末最后几天的费用在下个月的发票里，进项 IVA 按发票日期申报。只影响申报月份，不影响本报表的损益。",
    "销项 IVA：门店与网站的销售票据由 Liverpool 开给买家；卖家向 Liverpool 开具的销售发票以 Liverpool 的要求为准。",
]


# ======================================================================= helpers

def fmt(x):
    try:
        return "{:,.2f}".format(float(x))
    except (TypeError, ValueError):
        return "-"


def pct(x, dec=1):
    try:
        return "{:.{d}f}%".format(float(x) * 100, d=dec)
    except (TypeError, ValueError):
        return "-"


def _same_month(folder):
    """This folder first, then the store's other folders of the same month, newest
    first (a rerun of one route writes a new folder with just those files)."""
    root, name = os.path.split(os.path.abspath(folder))
    suffix = name[-8:] if re.search(r"_m\d{6}$", name) else None
    sib = sorted(glob.glob(os.path.join(root, "*" + suffix)), reverse=True) if suffix else []
    return [folder] + [d for d in sib if os.path.abspath(d) != os.path.abspath(folder)]


def _first(folder, pattern):
    for d in _same_month(folder):
        hits = sorted(glob.glob(os.path.join(d, pattern)))
        if hits:
            return hits[-1]
    return None


def rel(path, folder):
    if not path:
        return ""
    d = os.path.dirname(os.path.abspath(path))
    return os.path.basename(path) if d == os.path.abspath(folder) else \
        "%s（取自 %s）" % (os.path.basename(path), os.path.basename(d))


def read_meta(folder):
    try:
        with io.open(os.path.join(folder, "run_meta.json"), encoding="utf-8") as fh:
            return json.load(fh) or {}
    except Exception:
        return {}


def _empty(cols):
    return pd.DataFrame({c: pd.Series(dtype="object") for c in cols})


# ======================================================================= load

def load_store(store, folder, month):
    st = periods.stamp(month)
    S = {"store": store, "folder": folder, "month": month, "meta": read_meta(folder),
         "files": {}, "problems": []}

    def one(key, pattern, reader):
        p = _first(folder, pattern)
        S["files"][key] = p
        if not p:
            return None
        try:
            return reader(p)
        except Exception as e:                 # noqa: BLE001 — reported as a check
            S["problems"].append((key, os.path.basename(p), "%s: %s" % (e.__class__.__name__, e)))
            return None

    S["ledger"] = one("ledger", R.LEDGER % st, R.read_ledger)
    S["sales"] = one("sales", R.SALES % st, R.read_sales)
    S["orders"] = one("orders", R.ORDERS, R.read_orders)
    S["logistics"] = one("logistics", R.LOGISTICS % st, R.read_logistics)
    S["balance"] = one("balance", R.BALANCE % st, R.read_balance)
    S["payments"] = one("payments", R.PAYMENTS % st, R.read_payments)
    invs = []
    S["files"]["invoices"] = []
    for mon in (month, periods.next_month(month)):
        d = _first(folder, R.INVOICE_DIR % periods.stamp(mon))
        if d and os.path.isdir(d):
            S["files"]["invoices"].append(d)
            invs.append(R.read_invoices(d))
    S["invoices"] = pd.concat(invs, ignore_index=True) if invs else None
    no_data = set(r.get("form") for r in (S["meta"].get("missing") or []) if r.get("cause") == "NO_DATA")
    S["no_data"] = no_data
    return S


# ======================================================================= derive

def classify(L):
    """Each ledger line's report line, 中文 and original term; amount = −ledger amount."""
    looked = [C.lookup(a, c, t, x) for a, c, t, x in zip(L["account"], L["cls"], L["text"], L["tax"])]
    L = L.assign(line=[x[0] if x else C.UNREGISTERED for x in looked],
                 zh=[x[1] if x else "⚠ 未登记：%s %s %s" % (a, c, C.clean_text(t))
                     for x, a, c, t in zip(looked, L["account"], L["cls"], L["text"])],
                 orig=[x[2] if x else C.clean_text(t) for x, t in zip(looked, L["text"])])
    L["amount"] = -L["v"]
    return L


def _sku_maps(S):
    """order -> (sku, title, qty); article -> (sku, title) from the order list."""
    o = S["orders"]
    by_order, by_article = {}, {}
    if o is not None and len(o):
        for r in o.itertuples(index=False):
            by_order.setdefault(r.order, (r.sku or "", r.title, r.qty or 1.0, r.article))
            if r.sku:
                by_article.setdefault(r.article, (r.sku, r.title))
    return by_order, by_article


def _doc_parts(S):
    """Ledger sale document -> [(order, sku, title, units, weight)] from Ventas
    Marketplace (one row per order of the document) and the order list."""
    v = S["sales"]
    by_order, by_article = _sku_maps(S)
    parts = collections.defaultdict(list)
    S["unknown_orders"] = []
    if v is None:
        return parts
    for r in v.itertuples(index=False):
        o = by_order.get(r.order)
        if o and o[0]:
            sku, title, qty = o[0], o[1], o[2]
        elif r.article in by_article:
            sku, title = by_article[r.article]
            qty = 1.0
            S["unknown_orders"].append(r.order)
        else:
            sku, title, qty = "LV-%s" % r.article, "", 1.0
            S["unknown_orders"].append(r.order)
        parts[r.doc].append((r.order, sku, title, qty, abs(r.gross) or 1.0))
    return parts


def _split(amount, parts):
    tot = sum(p[4] for p in parts)
    out = [(p, amount * p[4] / tot) for p in parts[:-1]]
    out.append((parts[-1], amount - sum(a for _p, a in out)))
    return out


def derive(S):
    a, b = periods.month_bounds(S["month"])
    m = {"pl": {}}
    L = S["ledger"]
    if L is None or not len(L):
        L = _empty(["account", "doc", "cls", "text", "tax", "ref", "assign", "clearing", "post",
                    "doc_date", "cleared"]).assign(v=pd.Series(dtype=float))
    L = classify(L)
    inside = (L["post"] >= pd.Timestamp(a)) & (L["post"] <= pd.Timestamp(b))
    S["outside"] = int((~inside).sum())
    L = L[inside].copy()
    S["L"] = L
    m["unregistered"] = sorted(set(L.loc[L["line"] == C.UNREGISTERED, "zh"]))
    by = L.groupby("line")["amount"].sum()
    for k in C.INCOME_LINES + C.FEE_LINES + C.WITHHOLDING_LINES:
        m["pl"][k] = float(by.get(k, 0.0))
    pl = m["pl"]
    pl["withheld"] = pl[C.WH_IVA] + pl[C.WH_ISR]
    m["revenue"] = pl[C.SALES] + pl[C.PROMO]
    m["net_sales"] = m["revenue"] + pl[C.REFUND]
    m["fees"] = sum(pl[k] for k in FEE_ROWS)
    m["contrib"] = m["net_sales"] + m["fees"]
    m["book_total"] = float(L.loc[L["line"] != C.PAYMENT, "amount"].sum())
    pay = L[L["line"] == C.PAYMENT]
    m["cash"] = float(pay["v"].sum())
    m["lines"] = int(len(L))
    m["iva_out"] = m["net_sales"] * IVA_RATE / (1 + IVA_RATE)
    m["iva_in"] = -m["fees"] * IVA_RATE / (1 + IVA_RATE)
    # payments: attempts the bank sent back (a clearing document whose payable lines net to 0)
    pv = pay[pay["account"] == C.PAYABLE]
    nets = pv.groupby("clearing")["v"].sum() if len(pv) else pd.Series(dtype=float)
    gross = pv[pv["v"] > 0].groupby("clearing")["v"].sum() if len(pv) else pd.Series(dtype=float)
    rej = [c for c in nets.index if abs(nets[c]) <= TOL and gross.get(c, 0) > 0]
    m["rejected"] = [(c, float(gross.get(c, 0.0))) for c in rej]
    m["paid_clearings"] = [(c, float(nets[c])) for c in nets.index if abs(nets[c]) > TOL]
    P = S["payments"]
    m["paid_detail"] = float(P.loc[(P["paid"] >= pd.Timestamp(a)) & (P["paid"] <= pd.Timestamp(b)),
                                   "amount"].sum()) if P is not None and len(P) else None
    # ---- sales documents: the link to orders and SKUs
    parts = _doc_parts(S)
    sale_docs = L[L["line"] == C.SALES]
    ref_doc = dict(zip(sale_docs["ref"], sale_docs["doc"]))
    # a promotion credit is its own document carrying the sale's reference
    L["sale_doc"] = [ref_doc.get(r, d) if ln == C.PROMO else d
                     for ln, r, d in zip(L["line"], L["ref"], L["doc"])]
    linked = L[L["line"].isin(C.SALE_LINES)]
    m["docs_sale"] = int(L.loc[L["line"].isin((C.SALES, C.REFUND)), "doc"].nunique())
    unl = linked[~linked["sale_doc"].isin(parts.keys())]
    m["docs_unlinked"] = int(unl.loc[unl["line"].isin((C.SALES, C.REFUND)), "doc"].nunique())
    m["unlinked_sales"] = float(unl.loc[unl["line"].isin((C.SALES, C.PROMO, C.REFUND)), "amount"].sum())
    m["orders_unknown"] = len(set(S.get("unknown_orders") or []))
    S["parts"] = parts
    sku_rows(S)
    bridge_a(S, m)
    bridge_b(S, m)
    bridge_d(S, m)
    orders_kpi(S, m)
    S["m"] = m
    S["sku"] = sku_table(S)
    S["returns"] = returns_table(S)
    return S


def sku_rows(S):
    """S["by_sku"]: one row per (SKU, line) with the amount; every ledger line of the
    month lands exactly once: sale lines by their document's orders, shipping by
    the logistic-cost list's orders, the rest allocated by net sales."""
    L, parts = S["L"], S["parts"]
    out = []          # (sku, title, line, amount, units, returned)
    for r in L[L["line"].isin(C.SALE_LINES)].itertuples(index=False):
        ps = parts.get(r.sale_doc)
        if not ps:
            out.append((UNLINKED, "", r.line, r.amount, 1.0 if r.line == C.SALES else 0.0,
                        1.0 if r.line == C.REFUND else 0.0))
            continue
        for p, amt in _split(r.amount, ps):
            out.append((p[1], p[2], r.line, amt, p[3] if r.line == C.SALES else 0.0,
                        p[3] if r.line == C.REFUND else 0.0))
    by_order, by_article = _sku_maps(S)
    v = S["sales"]
    order_sku = dict(by_order)
    if v is not None:
        for r in v.itertuples(index=False):
            if r.order not in order_sku:
                sku = by_article.get(r.article, ("LV-%s" % r.article, ""))
                order_sku[r.order] = (sku[0], sku[1], 1.0, r.article)
    # shipping and extra shipping: per order from the marketplace lists (ledger doc = Doc contable)
    alloc = []
    for line, lst in ((C.LOGISTICS, S["logistics"]), (C.LOGISTICS_EXTRA, S["balance"])):
        rows = L[L["line"] == line]
        per_doc = lst.groupby("doc") if lst is not None and len(lst) else None
        for doc, g in rows.groupby("doc"):
            amt = float(g["amount"].sum())
            if per_doc is None or doc not in per_doc.groups:
                alloc.append((line, amt))
                continue
            items = per_doc.get_group(doc)
            w = items.groupby("order")["amount"].sum()
            tot = float(w.sum())
            if tot <= 0:
                alloc.append((line, amt))
                continue
            done = 0.0
            keys = list(w.index)
            for i, o in enumerate(keys):
                part = amt - done if i == len(keys) - 1 else amt * float(w[o]) / tot
                done += part
                sku = order_sku.get(o, (UNLINKED, ""))
                out.append((sku[0] or UNLINKED, sku[1], line, part, 0.0, 0.0))
    for line in (C.LOGISTICS_ADJ, C.MSI, C.SUBSCRIPTION, C.OTHER, C.UNREGISTERED):
        amt = float(L.loc[L["line"] == line, "amount"].sum())
        if amt:
            alloc.append((line, amt))
    df = pd.DataFrame(out, columns=["sku", "title", "line", "amount", "units", "returned"])
    # store-level charges: allocated by each SKU's net sales (negative counted as 0)
    base = df[df["line"].isin(C.INCOME_LINES) & (df["sku"] != UNLINKED)].groupby("sku")["amount"].sum()
    base = base[base > 0]
    extra = []
    for line, amt in alloc:
        if base.sum() <= 0:
            extra.append((NO_SKU, "", line, amt, 0.0, 0.0))
            continue
        keys = list(base.index)
        done = 0.0
        for i, k in enumerate(keys):
            part = amt - done if i == len(keys) - 1 else amt * float(base[k]) / float(base.sum())
            done += part
            extra.append((k, "", line, part, 0.0, 0.0))
    S["by_sku"] = pd.concat([df, pd.DataFrame(extra, columns=df.columns)], ignore_index=True) \
        if extra else df
    S["allocated"] = {line: amt for line, amt in alloc}


def sku_table(S):
    d = S["by_sku"]
    if not len(d):
        return pd.DataFrame()
    g = d.pivot_table(index="sku", columns="line", values="amount", aggfunc="sum", fill_value=0.0)
    for k in C.INCOME_LINES + C.FEE_LINES + C.WITHHOLDING_LINES:
        if k not in g.columns:
            g[k] = 0.0
    g["units"] = d.groupby("sku")["units"].sum().reindex(g.index).fillna(0.0)
    g["returns"] = d.groupby("sku")["returned"].sum().reindex(g.index).fillna(0.0)
    titles = d[d["title"] != ""].groupby("sku")["title"].first()
    g["name"] = titles.reindex(g.index).fillna("")
    g["income"] = g[list(C.INCOME_LINES)].sum(axis=1)
    g["fees"] = g[list(C.FEE_LINES)].sum(axis=1)
    g["contrib"] = g["income"] + g["fees"]
    g["withheld"] = g[C.WH_IVA] + g[C.WH_ISR]
    return g.sort_values(C.SALES, ascending=False)


def bridge_a(S, m):
    """Orders created in the month (order list) -> sales posted in the month
    (ledger, full price = buyer paid + Liverpool promotion), order by order through
    Ventas Marketplace."""
    o, v = S["orders"], S["sales"]
    if o is None or v is None:
        m["bridge_a"] = None
        return
    a, b = periods.month_bounds(S["month"])
    # every order of the month: Liverpool posts a sale on its order day, also when it
    # is cancelled later (the cancellation is a reversal document, ⑤)
    month_orders = o[(o["created"] >= pd.Timestamp(a)) & (o["created"] <= pd.Timestamp(b))]
    posted_docs = set(S["L"].loc[S["L"]["line"] == C.SALES, "doc"])
    sold = v[(v["gross"] > 0) & v["doc"].isin(posted_docs)]
    mo = month_orders.groupby("order")["subtotal"].sum()
    sv = sold.groupby("order")["gross"].sum()
    created = o.set_index("order")["created"].groupby(level=0).first()
    both = mo.index.intersection(sv.index)
    earlier = [x for x in sv.index if x not in mo.index and x in created.index
               and created[x] < pd.Timestamp(a)]
    unknown = [x for x in sv.index if x not in created.index]
    not_posted = [x for x in mo.index if x not in sv.index]
    A = {"orders": float(mo.sum()), "orders_n": int(len(mo)),
         "not_posted": float(mo.reindex(not_posted).sum()), "not_posted_n": len(not_posted),
         "earlier": float(sv.reindex(earlier).sum()), "earlier_n": len(earlier),
         "unknown": float(sv.reindex(unknown).sum()), "unknown_n": len(unknown),
         "price_diff": float(sv.reindex(both).sum() - mo.reindex(both).sum()),
         "vm": float(sold["gross"].sum()),
         "ledger": m["revenue"] - float(S["L"].loc[S["L"]["line"].isin((C.SALES, C.PROMO)) &
                                                   ~S["L"]["sale_doc"].isin(set(sold["doc"])),
                                                   "amount"].sum())}
    A["computed"] = A["orders"] - A["not_posted"] + A["earlier"] + A["unknown"] + A["price_diff"]
    A["resid"] = A["ledger"] - A["computed"]
    A["not_posted_states"] = o[o["order"].isin(not_posted)]["status"].value_counts().to_dict()
    A["unlinked"] = m["revenue"] - A["ledger"]
    m["bridge_a"] = A


def bridge_b(S, m):
    """Commission per sale (Ventas Marketplace) <-> the ledger's commission lines."""
    v = S["sales"]
    L = S["L"]
    docs = set(L.loc[L["line"].isin((C.SALES, C.REFUND)), "doc"])
    m["bridge_b"] = None if v is None else {
        "vm": -float(v.loc[v["doc"].isin(docs), "commission"].sum()),
        "ledger": m["pl"][C.COMMISSION]}


def bridge_d(S, m):
    """The logistic-cost list (one row per order) <-> the ledger's shipping charges,
    by accounting document."""
    lst = S["logistics"]
    rows = S["L"][S["L"]["line"] == C.LOGISTICS]
    led = rows.groupby("doc")["amount"].sum()
    if lst is None:
        m["bridge_d"] = None
        return
    fl = -lst.groupby("doc")["amount"].sum()
    both = led.index.intersection(fl.index)
    m["bridge_d"] = {"list": float(fl.sum()), "list_n": int(len(lst)),
                     "matched_list": float(fl.reindex(both).sum()),
                     "matched_ledger": float(led.reindex(both).sum()),
                     "ledger_only": float(led.drop(both).sum()), "ledger_only_n": int(len(led) - len(both)),
                     "list_only": float(fl.drop(both).sum()), "ledger": float(led.sum())}


def orders_kpi(S, m):
    o = S["orders"]
    if o is None:
        return
    a, b = periods.month_bounds(S["month"])
    mo = o[(o["created"] >= pd.Timestamp(a)) & (o["created"] <= pd.Timestamp(b))]
    live = mo[mo["status"] != "Cancelado"]
    m["gmv_orders"] = float(live["subtotal"].sum())
    m["orders_n"] = int(live["order"].nunique())
    m["cancel_n"] = int(mo.loc[mo["status"] == "Cancelado", "order"].nunique())
    m["cancel_amt"] = float(mo.loc[mo["status"] == "Cancelado", "subtotal"].sum())
    # sales posted in the month whose order was not delivered by month end (and not
    # cancelled or refunded): the timing difference to delivery-based revenue
    v = S["sales"]
    if v is not None:
        posted = v[(v["gross"] > 0) & v["doc"].isin(set(S["L"].loc[S["L"]["line"] == C.SALES, "doc"]))]
        od = o.drop_duplicates("order").set_index("order")
        st = posted["order"].map(od["status"])
        dl = posted["order"].map(od["delivered"])
        # no date: "Recibido" means received (date not given), else still on its way
        open_ = posted[~st.isin(["Cancelado", "Reembolsado"]) &
                       ((dl > pd.Timestamp(b)) | (dl.isna() & (st != "Recibido")))]
        m["undelivered"] = float(open_["gross"].sum())
        m["undelivered_n"] = int(open_["order"].nunique())
    m["orders_from"] = o["created"].min()


def returns_table(S):
    L = S["L"]
    ret_docs = set(L.loc[L["line"] == C.REFUND, "doc"])
    r = L[L["doc"].isin(ret_docs) & L["line"].isin(C.SALE_LINES)]
    out = {"docs": len(ret_docs)}
    out["money"] = r.groupby(["line", "zh"])["amount"].agg(["count", "sum"]) if len(r) else pd.DataFrame()
    rows = []
    parts = S["parts"]
    for doc, g in r.groupby("doc"):
        ps = parts.get(doc) or [("", UNLINKED, "", 0, 1)]
        amt = lambda k: float(g.loc[g["line"] == k, "amount"].sum())
        rows.append({"date": g["post"].min(), "doc": doc, "orders": "、".join(p[0] for p in ps if p[0]),
                     "sku": "、".join(sorted(set(p[1] for p in ps))), "refund": amt(C.REFUND),
                     "commission": amt(C.COMMISSION), "withheld": amt(C.WH_IVA) + amt(C.WH_ISR)})
    out["list"] = pd.DataFrame(rows)
    o = S["orders"]
    if o is not None:
        a, b = periods.month_bounds(S["month"])
        mo = o[(o["created"] >= pd.Timestamp(a)) & (o["created"] <= pd.Timestamp(b))]
        flag = mo["status"].where(mo["status"] == "Cancelado", mo["refund_flag"].replace("", "（无）"))
        out["by_status"] = mo.assign(_f=flag).groupby("_f")["subtotal"].agg(["count", "sum"])
    return out


# ======================================================================= checks

def validate(stores):
    Cs = []
    for S in stores:
        st, m, f = S["store"], S["m"], S["files"]
        for key, name, why in S["problems"]:
            Cs.append(K.check("load_" + key, st, "读取 %s" % name, False, "读不出来：%s" % why))
        Cs.append(K.check("file_ledger", st, "Liverpool 账本（Cobros y Pagos）在", bool(f.get("ledger")),
                          "损益表的全部金额来自它" if f.get("ledger") else
                          "没有账本文件（LV_mayor_*.xlsx）：本店本月的损益表是空的，必须重新下载"))
        Cs.append(K.check("file_sales", st, "Ventas Marketplace 在", bool(f.get("sales")),
                          "销售单据 → 订单与 SKU 的对照来自它" if f.get("sales") else
                          "没有 Ventas Marketplace：⑦ 无法按 SKU 归集，③桥 A、B 无法核对"))
        for key, name, why in (("orders", "订单列表", "SKU 编码与订单口径 GMV 来自它；缺了 SKU 只显示 Liverpool 商品号"),
                               ("logistics", "物流费明细（Costos logísticos）", "运费逐单归 SKU 与③桥 D 用它；缺了运费按销售额分摊"),
                               ("balance", "补充物流费（Balance de cuenta）", "补充物流费逐单归 SKU；缺了按销售额分摊")):
            Cs.append(K.check("file_" + key, st, "%s在" % name, bool(f.get(key)),
                              rel(f.get(key), S["folder"]) if f.get(key) else "没有：%s" % why,
                              severity="warn"))
        Cs.append(K.check("ledger_in_month", st, "账本只取本月入账的行", m["lines"] > 0,
                          "本月 %d 行；文件里另有 %d 行不在本月，已排除" % (m["lines"], S.get("outside") or 0)
                          if m["lines"] else "账本里没有本月入账的行"))
        Cs.append(K.check("unregistered", st, "账本行全部已登记", not m["unregistered"],
                          "全部已登记（catalog.py）" if not m["unregistered"] else
                          "未登记：%s（金额已计入“未登记科目”行）" % "；".join(m["unregistered"])))
        Cs.append(K.check("pl_tie", st, "损益表各行 ＝ 账本本月非付款行合计",
                          abs(m["contrib"] + m["pl"]["withheld"] - m["book_total"]) <= TOL,
                          "损益（贡献 − 代扣税）%s，账本 %s" % (fmt(m["contrib"] + m["pl"]["withheld"]),
                                                       fmt(m["book_total"]))))
        if S["sales"] is not None:
            Cs.append(K.check("sale_link", st, "销售单据都能对到订单", m["docs_unlinked"] == 0,
                              "全部 %d 张销售/退货单据都在 Ventas Marketplace 里" % m["docs_sale"]
                              if not m["docs_unlinked"] else
                              "%d 张单据（金额 %s）不在 Ventas Marketplace 里，⑦ 列为“未关联到订单”" % (
                                  m["docs_unlinked"], fmt(m["unlinked_sales"])), severity="warn"))
            Cs.append(K.check("order_link", st, "订单都在订单列表里（SKU 编码）", not m["orders_unknown"],
                              "全部能找到卖家 SKU" if not m["orders_unknown"] else
                              "%d 张订单不在订单列表里（订单列表只有最近四个月），SKU 按 Liverpool 商品号对照"
                              % m["orders_unknown"], severity="warn"))
        A = m.get("bridge_a")
        Cs.append(K.check("bridge_a", st, "桥 A：订单 → 本月入账的销售",
                          True if not A else abs(A["resid"]) <= TOL,
                          "没有订单列表或 Ventas Marketplace" if not A else
                          "残差 %s；价格差 %s" % (fmt(A["resid"]), fmt(A["price_diff"])),
                          severity="warn", skipped=not A))
        B = m.get("bridge_b")
        Cs.append(K.check("bridge_b", st, "桥 B：每笔销售的佣金 ＝ 账本佣金",
                          True if not B else abs(B["vm"] - B["ledger"]) <= TOL,
                          "没有 Ventas Marketplace" if not B else
                          "Ventas Marketplace %s，账本 %s，差 %s" % (fmt(B["vm"]), fmt(B["ledger"]),
                                                                fmt(B["ledger"] - B["vm"])),
                          severity="warn", skipped=not B))
        D = m.get("bridge_d")
        Cs.append(K.check("bridge_d", st, "桥 D：物流费明细 ＝ 账本发货物流费（按凭证号）",
                          True if not D else abs(D["matched_list"] - D["matched_ledger"]) <= TOL,
                          "没有物流费明细" if not D else
                          "对上的凭证：明细 %s，账本 %s；账本另有 %d 张凭证 %s 不在明细里（按销售额分摊）" % (
                              fmt(D["matched_list"]), fmt(D["matched_ledger"]), D["ledger_only_n"],
                              fmt(D["ledger_only"])), severity="warn", skipped=not D))
        rej = m["rejected"]
        Cs.append(K.check("payments_rejected", st, "本月付款没有被银行退回", not rej,
                          "没有" if not rej else
                          "%d 次付款被银行退回（%s），之后的付款已补付；反复出现请检查在 Liverpool 登记的收款账户"
                          % (len(rej), "、".join("%s %s" % (c, fmt(a)) for c, a in rej)), severity="warn"))
    return Cs


def store_meta_missing(S):
    return RF.meta_missing(S["meta"])


# ======================================================================= workbook

def build_workbook(stores, ctx, out_path):
    wb = Workbook()
    wb.remove(wb.active)
    fam = RF.Family(stores, money=MNY)
    sheet_readme(wb, fam, ctx)
    sheet_pl(wb, fam, ctx)
    sheet_bridge(wb, fam, ctx)
    sheet_fees(wb, fam, ctx)
    sheet_returns(wb, fam, ctx)
    RF.not_applicable_sheet(wb, SHEETS[5], "退货质检明细",
                            "不适用：Liverpool 不提供逐件退货质检报表。退货的金额与佣金、代扣税的退还见【%s】。"
                            % SHEETS[4])
    sheet_sku(wb, fam, ctx)
    sheet_trend(wb, fam, ctx)
    RF.not_applicable_sheet(wb, SHEETS[8], "库存与动销分析",
                            "不适用：本店没有使用 Liverpool Fulfillment（FBL），平台上没有库存（订单表 "
                            "Fulfillment = No）。本月各 SKU 的件数见【%s】。" % SHEETS[6])
    sheet_gaps(wb, fam, ctx)
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    return out_path


def pl_label(key):
    return PL_LABELS.get(key, key)


def sheet_readme(wb, fam, ctx):
    stores = fam.stores
    sources = []
    for S in stores:
        f, meta = S["files"], S["meta"]
        flags = meta.get("flags") or {}
        sources.append((S["store"], "下载于 %s%s。账本：%s；Ventas Marketplace：%s；订单列表：%s；物流费明细：%s；"
                        "补充物流费：%s；付款：%s；费用发票：%s。"
                        % (meta.get("downloaded_at") or "?",
                           "（订单等取自 %s）" % flags["marketplace_store"] if flags.get("marketplace_store") else "",
                           rel(f.get("ledger"), S["folder"]) or "缺", rel(f.get("sales"), S["folder"]) or "缺",
                           rel(f.get("orders"), S["folder"]) or "缺", rel(f.get("logistics"), S["folder"]) or "缺",
                           rel(f.get("balance"), S["folder"]) or "缺", rel(f.get("payments"), S["folder"]) or "缺",
                           "、".join(os.path.basename(p) for p in f.get("invoices") or []) or "缺")))
    tot = lambda k: sum(S["m"].get(k, 0.0) or 0.0 for S in stores)
    rev = tot("revenue")
    pl = lambda k: sum(S["m"]["pl"][k] for S in stores)
    points = [("规模", "本月入账商品销售（按原价，含 IVA）%s，退货 %s，净销售收入 %s；按下单日的订单 GMV %s，%d 张订单。"
               % (fmt(rev), fmt(-pl(C.REFUND)), fmt(tot("net_sales")), fmt(tot("gmv_orders")),
                  int(tot("orders_n"))))]
    if rev:
        points.append(("费用率", "平台费用合计 %s，占商品销售收入 %s；其中佣金（含银行卡附加）%s、物流 %s。"
                       % (fmt(-tot("fees")), pct(-tot("fees") / rev),
                          pct(-(pl(C.COMMISSION) + pl(C.CARD_FEE)) / rev),
                          pct(-sum(pl(k) for k in LOGISTICS_ROWS) / rev))))
        points.append(("贡献毛利", "经营贡献毛利（未扣商品成本）%s，占商品销售收入 %s。"
                       % (fmt(tot("contrib")), pct(tot("contrib") / rev))))
    points.append(("代扣税", "Liverpool 本月净代扣 %s（IVA %s、ISR %s），是预缴税款，不在贡献里扣。"
                   % (fmt(-pl("withheld")), fmt(-pl(C.WH_IVA)), fmt(-pl(C.WH_ISR)))))
    points.append(("结算", "本月 Liverpool 实际付款（扣除费用后）%s。" % fmt(tot("cash"))))
    costs = ctx.get("sku_costs")
    policies = list(ACCOUNTING_POLICIES)
    if costs is not None:
        sources.append(SC.source_line(costs))
        results = [S.get("cost") for S in stores]
        points.append(SC.point_line(results, tot("contrib") - sum(r.cogs for r in results if r)))
        policies.append(SC.policy_line(COST_UNITS_BASIS, COST_IVA_NOTE.strip("（）") + "。" + SC.STOCK_NOT_VALUED))
    sections = [
        ("一、这份报表回答什么", [
            ("本月赚了多少", "【%s】从商品销售收入推到“经营贡献毛利（未扣商品成本）”，填入商品成本后得到利润；"
             "Liverpool 代扣的税单列在最后。" % SHEETS[1]),
            ("钱对不对得上", "【%s】四座桥：A 订单 → 入账的销售、B 每笔销售的佣金 → 账本、C 账本 → 付款、"
             "D 物流费明细 → 账本。" % SHEETS[2]),
            ("还缺什么", "【%s】列出未通过的校验、缺的文件，以及报税时要注意的事。" % SHEETS[9])]),
        ("二、口径（与其他平台报表的不同之处）", [
            ("以 Liverpool 账本为准", "Liverpool 在供应商门户里给每家卖家记两个账户：应付卖家的（销售、退货、代扣税、付款）"
             "和卖家应付的（佣金、物流费、订阅费）。两边本月入账的每一行都计入本报表，按入账日归月；"
             "和美客多的账单基准一样，以平台的账为唯一金额来源。销售在下单付款当天入账，取消也先入账再冲回。"),
            ("含 IVA", "全部金额含 IVA，与 Liverpool 的账和付款一致，也与美客多墨西哥账单基准报表的口径相同。"),
            ("代扣税", "Liverpool 每笔销售代扣 IVA 8%、ISR 2.5%（按不含 IVA 价款），退货时退还；"
             "和美客多墨西哥一样单列，不计入贡献。"),
            ("账本行的科目", "账本里每一种行都登记在字典里（西语原文＋中文＋归入哪一行），见【%s】。"
             "出现字典里没有的行时报 error，金额进“未登记科目 ⚠”行。" % SHEETS[3])]),
        ("三、数据来源", sources),
        ("四、本期要点", points),
        ("五、各页说明", [
            (SHEETS[1], "各店分列（＋合计）。行位置每期固定，可直接跨月对比；灰底=公式，浅蓝=小计，黄底=需你录入。"),
            (SHEETS[2], "桥 A 订单 → 入账的销售、桥 B 佣金逐笔 → 账本、桥 C 账本 → 付款、桥 D 物流费明细 → 账本。"),
            (SHEETS[3], "本月账本里出现的每一种行：西语原文、中文、归入哪一行、笔数、金额（含 IVA）。"),
            (SHEETS[4], "本月入账的退货逐单：订单、SKU、退款、退还的佣金和代扣税；以及本月订单的取消/退款状态。"),
            (SHEETS[5], "不适用：Liverpool 不提供逐件退货质检报表。"),
            (SHEETS[6], "按 SKU 的收入、费用、贡献与代扣税；运费按物流费明细逐单归属，订阅费等店铺级费用按销售额分摊。"),
            (SHEETS[7], "本月按周（按入账日）的销售、退货与费用。"),
            (SHEETS[8], "不适用：未使用 Liverpool Fulfillment，没有平台库存。"),
            (SHEETS[9], "校验结果、缺失文件及原因、报税提示。校验和缺失为空表示没有待办。")]),
        ("六、会计处理", policies),
        ("七、报税时要注意", [("提示 %d" % i, q) for i, q in enumerate(TAX_NOTES, 1)]),
    ]
    bad = RF.failed_errors(ctx["checks"])
    warning = ("本期有 %d 项 error 级校验未通过，数字可能不完整，请先看【%s】顶部的校验结果。"
               % (len(bad), SHEETS[9])) if bad else None
    fam.readme_sheet(wb, SHEETS[0], "利物浦 Liverpool 月度财务报表 · %s · %s" % (ctx["period_cn"], ctx["who"]),
                     "生成于 %s。所有金额单位：墨西哥比索（MXN），含 IVA。" % ctx["built"],
                     sections, warning=warning)


def _fee_share(keys):
    return lambda S: -sum(S["m"]["pl"][k] for k in keys)


PL_MEMO = [
    ("gmv_orders", ("订单口径 GMV（含 IVA，不含已取消）", "Pedidos · Subtotal",
                    lambda S: S["m"].get("gmv_orders", 0.0),
                    "本月创建的订单金额（原价）；与入账口径的差见③桥 A")),
    (None, ("订单数（不含已取消）", "", lambda S: S["m"].get("orders_n", 0), "")),
    ("cancel_amt", ("已取消订单金额", "Cancelado", lambda S: S["m"].get("cancel_amt", 0.0),
                    "已入账的取消在“取消与退货退款”里冲回")),
    ("undelivered", ("本月入账、月末尚未送达的销售（参考）", "Fecha de entrega",
                     lambda S: S["m"].get("undelivered", 0.0),
                     "按送达确认收入时属于下个月的部分（订单列表的送达/提货日在月末之后，或下载时还没送达；不含取消/退款）")),
    ("wh_iva", ("其中：代扣 IVA", "Q6 IVA retenido", lambda S: S["m"]["pl"][C.WH_IVA], "按不含 IVA 价款的 8%，已扣除退货退还")),
    ("wh_isr", ("其中：代扣 ISR", "Q8 ISR retenido", lambda S: S["m"]["pl"][C.WH_ISR], "按不含 IVA 价款的 2.5%，已扣除退货退还")),
    ("iva_out", ("净销售收入所含 IVA（销项，参考）", "IVA trasladado", lambda S: S["m"]["iva_out"],
                 "＝ 净销售收入 × 16/116；仅供参考，申报以 SAT 单据为准")),
    ("iva_in", ("平台费用所含 IVA（进项，参考）", "IVA acreditable", lambda S: S["m"]["iva_in"],
                "＝ 平台费用 × 16/116；能否抵扣取决于 Liverpool 的费用发票（Notas de Cargo）")),
    ("cash", ("本月 Liverpool 实际付款（净额）", "ZP · Pago", lambda S: S["m"]["cash"],
              "付给卖家的款项已扣除卖家应付的费用；被银行退回又补付的付款已抵消，见③桥 C")),
]


def fill_cogs(fam, sh, rows_at, note_col, ctx):
    if ctx.get("sku_costs") is None:
        return
    r = rows_at["cogs"]
    for i, S in enumerate(fam.stores):
        sh.ws.cell(row=r, column=sh.c0 + i).value = S["cost"].cogs
    sh.ws.cell(row=r, column=note_col).value = fam.per_store(lambda S: SC.pl_note(S["cost"])) + COST_IVA_NOTE


def sheet_pl(wb, fam, ctx):
    sh, r, rows_at, note_col = fam.pl_sheet(
        wb, SHEETS[1], PL_ROWS, value=lambda S, key: S["m"]["pl"][key],
        title="合并损益表（%s）" % ctx["period_cn"],
        subtitle="口径：Liverpool 账本本月入账的全部行（按入账日），金额含 IVA。单位：MXN。"
                 "灰底=公式行，浅蓝=小计，黄底=需你录入。行位置每期固定，可直接跨月对比。")
    fill_cogs(fam, sh, rows_at, note_col, ctx)
    rev = lambda S: S["m"]["revenue"]
    r = fam.ratio_block(sh, r + 1, "费用率（占商品销售收入）", [
        ("佣金（含银行卡附加）", _fee_share((C.COMMISSION, C.CARD_FEE)), rev),
        ("物流（发货＋补充＋调整）", _fee_share(LOGISTICS_ROWS), rev),
        ("平台费用合计", _fee_share(FEE_ROWS), rev),
        ("经营贡献毛利率", lambda S: S["m"]["contrib"], rev)], note_col)
    fam.memo_block(sh, r + 1, "备查项（不参与上表计算）", [row for _k, row in PL_MEMO], note_col)


def sheet_bridge(wb, fam, ctx):
    Ln = RF.Line
    A = lambda S, k: (S["m"].get("bridge_a") or {}).get(k, 0.0)
    bridge_a = [
        Ln("本月创建的订单（原价，含后来取消的）", lambda S: A(S, "orders"),
           "订单列表 · Subtotal。Liverpool 在下单当天入账销售，取消的订单也先入账、再以冲回单据冲销"),
        Ln("减：其中本月没有入账的", lambda S: -A(S, "not_posted"),
           "付款前就取消等（状态：%s）" % fam.per_store(lambda S: "、".join(
               "%s %d" % (k, v) for k, v in ((S["m"].get("bridge_a") or {}).get("not_posted_states")
                                             or {}).items()))),
        Ln("加：以前月份的订单，本月入账", lambda S: A(S, "earlier"), "正常为 0（销售在下单当天入账）"),
        Ln("加：订单列表里没有的订单（早于四个月），本月入账", lambda S: A(S, "unknown")),
        Ln("加：入账价与订单价的差", lambda S: A(S, "price_diff"), "逐单比较，正常为 0"),
        Ln("＝ 推算的本月入账销售（原价）", lambda S: A(S, "computed"), style="subtotal"),
        Ln("账本：本月入账的销售 ＋ 促销补偿（能对到订单的）", lambda S: A(S, "ledger")),
        Ln("残差（应为 0）", lambda S: A(S, "resid"), style="residual"),
        Ln("参考：对不到订单的销售单据", lambda S: A(S, "unlinked"), "见⑩“销售单据都能对到订单”"),
    ]
    B = lambda S, k: (S["m"].get("bridge_b") or {}).get(k, 0.0)
    bridge_b = [
        Ln("Ventas Marketplace：本月单据的佣金（销售 − 退货）", lambda S: B(S, "vm"), "Comisión Dilisa，逐笔"),
        Ln("账本：销售佣金（净）", lambda S: B(S, "ledger"), "Comisión % DILISA − 退货退还的佣金"),
        Ln("残差（应为 0）", lambda S: B(S, "ledger") - B(S, "vm"), style="residual"),
    ]
    m = lambda S, k: S["m"][k]
    bridge_c = [
        Ln("经营贡献毛利（见②）", lambda S: m(S, "contrib")),
        Ln("加：代扣税金（见②）", lambda S: S["m"]["pl"]["withheld"]),
        Ln("＝ 本月账面净额", lambda S: m(S, "contrib") + S["m"]["pl"]["withheld"], style="subtotal"),
        Ln("账本：本月非付款行合计", lambda S: m(S, "book_total")),
        Ln("残差（应为 0）", lambda S: m(S, "contrib") + S["m"]["pl"]["withheld"] - m(S, "book_total"),
           style="residual"),
        Ln("本月实际付款（净额，ZP）", lambda S: m(S, "cash"),
           "Liverpool 付给卖家的款项扣除卖家应付的费用后的净额；本月付款 %s" % fam.per_store(
               lambda S: "、".join("%s %s" % (c, fmt(a)) for c, a in S["m"]["paid_clearings"]) or "无")),
        Ln("账面净额 − 实际付款（应收应付余额的变动）", lambda S: m(S, "contrib") + S["m"]["pl"]["withheld"] - m(S, "cash"),
           "正数：本月入账但还没付的（下次付款结算）；负数：本月付清了以前月份的"),
        Ln("参考：Detalle de Pagos 的本月付款（应付卖家一侧）",
           lambda S: S["m"]["paid_detail"] if S["m"]["paid_detail"] is not None else 0.0,
           "付款明细只列应付卖家一侧（销售、退货、代扣税），费用在账本另一侧抵扣"),
        Ln("参考：被银行退回的付款", lambda S: sum(a for _c, a in S["m"]["rejected"]),
           "退回后已抵消，不影响上面的数字；反复出现请检查收款账户"),
    ]
    D = lambda S, k: (S["m"].get("bridge_d") or {}).get(k, 0.0)
    bridge_d = [
        Ln("物流费明细：本月扣费的订单（Costos logísticos）", lambda S: D(S, "list"),
           "%s 单" % fam.per_store(lambda S: str((S["m"].get("bridge_d") or {}).get("list_n", 0)))),
        Ln("其中凭证号在账本里的", lambda S: D(S, "matched_list")),
        Ln("账本：同一批凭证的发货物流费", lambda S: D(S, "matched_ledger")),
        Ln("残差（应为 0）", lambda S: D(S, "matched_list") - D(S, "matched_ledger"), style="residual"),
        Ln("参考：账本里另有、明细里没有的物流费凭证", lambda S: D(S, "ledger_only"),
           "多为退货取回等不在明细里的物流费；⑦ 里按销售额分摊"),
    ]
    fam.bridge_sheet(wb, SHEETS[2], "平账勾稽表（%s）" % ctx["period_cn"],
                     "每座桥最后一行是残差，应为 0。单位：MXN，金额含 IVA。", [
                         ("桥 A　订单（下单日）→ 本月入账的销售（入账日）", bridge_a),
                         ("桥 B　每笔销售的佣金（Ventas Marketplace）→ 账本", bridge_b),
                         ("桥 C　损益 → 账本 → Liverpool 的付款", bridge_c),
                         ("桥 D　物流费明细（逐单）→ 账本的发货物流费", bridge_d)])


TEXT10 = F(10)
FEE_COLS = [RF.Col("店铺", 22, font=TEXT10), RF.Col("账本原文（科目）", 46, font=TEXT10),
            RF.Col("中文", 30, font=TEXT10), RF.Col("损益表行", 30, font=TEXT10),
            RF.Col("归类", 12, font=TEXT10), RF.Col("笔数", 8, font=TEXT10),
            RF.Col("金额（含 IVA）", 16, RF.MONEY)]


def pl_row_of(line):
    return "withheld" if line in C.WITHHOLDING_LINES else line


def fee_groups(S):
    """④'s rows of one store: ledger term × 中文 × ② line, n and amount, in ②'s row
    order then by amount; payments left out. ④ and results_figures read this."""
    L = S["L"]
    L = L[L["line"] != C.PAYMENT]
    if not len(L):
        return None
    keys = [x.key for x in PL_ROWS]
    g = L.groupby(["orig", "zh", "line"]).agg(n=("amount", "size"), amt=("amount", "sum")).reset_index()
    g["_order"] = g["line"].map(lambda k: keys.index(pl_row_of(k)) if pl_row_of(k) in keys else 99)
    return g.sort_values(["_order", "amt"])


def sheet_fees(wb, fam, ctx):
    ws = fam.table_sheet(wb, SHEETS[3], "平台费用明细（%s）" % ctx["period_cn"],
                         "本月 Liverpool 账本里出现的每一种行（付款除外）。归入②的哪一行见“损益表行”。金额 MXN，含 IVA。",
                         FEE_COLS)
    r = 5
    for S in fam.stores:
        g = fee_groups(S)
        if g is None:
            continue
        for _, x in g.iterrows():
            cells = fam.write_row(ws, r, FEE_COLS, [
                S["store"], x["orig"], x["zh"], pl_label(pl_row_of(x["line"])),
                C.LINE_CATEGORY.get(x["line"], "⚠ 未登记"), int(x["n"]), float(x["amt"])])
            if x["line"] == C.UNREGISTERED:
                for c in cells[:6]:
                    c.font = F(10, False, K.C_RED)
            r += 1
        ws.cell(row=r, column=1, value="%s 合计" % S["store"]).font = F(10, True)
        c = fam.money_cell(ws.cell(row=r, column=7), S["m"]["book_total"])
        c.font = F(10, True)
        c.fill = fill(K.C_TOTAL)
        ws.cell(row=r, column=3, value="＝ ② 收入 ＋ 费用 ＋ 代扣税金").font = F(9, False, K.C_GREY)
        r += 2
    ws.freeze_panes = "A5"


RET_COLS = [RF.Col("店铺", 20), RF.Col("入账日", 12), RF.Col("单据号", 14), RF.Col("订单号", 26),
            RF.Col("SKU", 18), RF.Col("退货退款", 14, RF.MONEY), RF.Col("退还的佣金", 14, RF.MONEY),
            RF.Col("退还的代扣税", 14, RF.MONEY), RF.Col("净影响", 14, RF.MONEY)]


def sheet_returns(wb, fam, ctx):
    ws = fam.table_sheet(wb, SHEETS[4], "退货退款分析（%s）" % ctx["period_cn"],
                         "上半：本月入账的退货逐单（账本），退款与 Liverpool 退还的佣金、代扣税。下半：本月创建的订单里"
                         "被取消或退款的（订单列表）。金额含 IVA。", RET_COLS)
    r = 5
    for S in fam.stores:
        R_ = S["returns"]
        lst = R_.get("list")
        if lst is not None and len(lst):
            for _, x in lst.sort_values("date").iterrows():
                fam.write_row(ws, r, RET_COLS, [S["store"], str(x["date"].date()) if pd.notna(x["date"]) else "",
                                                x["doc"], x["orders"], x["sku"], x["refund"], x["commission"],
                                                x["withheld"], x["refund"] + x["commission"] + x["withheld"]])
                r += 1
        ws.cell(row=r, column=1, value="%s：本月入账的退货 %d 张单据" % (S["store"], R_["docs"])).font = F(10, True)
        r += 2
        bs = R_.get("by_status")
        if bs is not None and len(bs):
            ws.cell(row=r, column=1, value="本月创建的订单").font = F(10, True)
            r += 1
            for k, row in bs.iterrows():
                ws.cell(row=r, column=2, value="状态")
                ws.cell(row=r, column=4, value=k)
                ws.cell(row=r, column=5, value=int(row["count"])).border = BOX
                fam.money_cell(ws.cell(row=r, column=6), float(row["sum"]))
                r += 1
            r += 1


_SKU_FORMULA = dict(fmt=RF.MONEY, border=False, fill=K.C_CALC)
SKU_COLS = ([RF.Col("店铺", 18), RF.Col("SKU", 18), RF.Col("商品", 44), RF.Col("入账件数", 9),
             RF.Col("退货件数", 9)]
            + [RF.Col(h, w, RF.MONEY) for h, w in (
                ("商品销售收入（原价）", 14), ("退货退款", 12), ("佣金（含银行卡附加）", 14),
                ("物流", 12), ("其他费用", 12), ("贡献（未扣成本）", 14))]
            + [RF.Col("单位成本（录入，含 IVA）", 12, RF.MONEY, fill=K.C_INPUT),
               RF.Col("商品成本", 14, formula="=-L{r}*(D{r}-E{r})", **_SKU_FORMULA),
               RF.Col("SKU 毛利", 14, formula="=K{r}+M{r}", **_SKU_FORMULA),
               RF.Col("代扣税金（不计入贡献）", 14, RF.MONEY)])
SKU_COLS_COST = SKU_COLS[:11] + [
    RF.Col("单位成本（含 IVA）", 12, RF.MONEY),
    RF.Col("商品成本", 14, formula="=-L{r}*(D{r}-E{r})", **_SKU_FORMULA),
    RF.Col("SKU 毛利", 14, formula="=K{r}+M{r}", **_SKU_FORMULA),
    RF.Col("代扣税金（不计入贡献）", 14, RF.MONEY),
    RF.Col("成本来源", 40, font=F(9, False, K.C_GREY))]


def _is_real_sku(sku):
    return sku not in (NO_SKU, UNLINKED)


def sku_items(S):
    g = S["sku"]
    return [SC.Item(sku, x["name"] if isinstance(x["name"], str) else "", int(round(x["units"])),
                    int(round(x["returns"])), float(x[C.SALES] + x[C.PROMO]))
            for sku, x in g.iterrows() if _is_real_sku(sku)] if len(g) else []


def sku_list(store, folder, month):
    """(platform, country, currency, [Item]) of one store's month — for tools/sku_cost_template."""
    return PLATFORM, COUNTRY, CUR, sku_items(derive(load_store(store, folder, month)))


def _cost_cells(cells, S, sku):
    if not _is_real_sku(sku):
        cells[15].value = "店铺级费用，不计商品成本" if sku == NO_SKU else "未关联到订单，无法计成本"
        return
    ln = S["cost"].line(sku)
    if ln is not None and ln.status == SC.OK:
        cells[11].value = ln.unit_cost
        cells[15].value = SC.status_text(ln)
    else:
        cells[11].fill = fill(K.C_INPUT)
        cells[15].value = SC.status_text(ln)
        cells[15].font = F(9, True, K.C_AMBER)


def sheet_sku(wb, fam, ctx):
    costs = ctx.get("sku_costs") is not None
    cols = SKU_COLS_COST if costs else SKU_COLS
    ws = fam.table_sheet(wb, SHEETS[6], "SKU 盈利分析（%s）" % ctx["period_cn"],
                         "按销售单据对到的订单与 SKU 汇总（含 IVA）。运费按物流费明细逐单归属；订阅费、调整等店铺级费用"
                         "按各 SKU 的净销售额分摊。" + (
                             "单位成本 ＝ 商品成本表的成本（不含 IVA）× 1.16，商品成本 ＝ 单位成本 ×（入账件数 − 退货件数）；"
                             "“未录入成本”的 SKU 黄色格可临时填写（含 IVA），长期请补录到成本表。" if costs else
                             "黄色列填单位采购成本（含 IVA），右边两列自动算出 SKU 毛利。"), cols)
    r = 5
    other = [C.LOGISTICS_ADJ, C.MSI, C.SUBSCRIPTION, C.OTHER, C.UNREGISTERED]
    for S in fam.stores:
        g = S["sku"]
        for sku, x in (g.iterrows() if len(g) else ()):
            money = [x[C.SALES] + x[C.PROMO], x[C.REFUND], x[C.COMMISSION] + x[C.CARD_FEE],
                     x[C.LOGISTICS] + x[C.LOGISTICS_EXTRA], sum(x[k] for k in other), x["contrib"]]
            vals = [S["store"], sku, x["name"], round(float(x["units"]), 2), round(float(x["returns"]), 2)] + \
                [float(v) for v in money]
            vals += ([None, None, None, float(x["withheld"]), None] if costs else
                     [0, None, None, float(x["withheld"])])
            cells = fam.write_row(ws, r, cols, vals)
            if costs:
                _cost_cells(cells, S, sku)
            r += 1
        r += 1
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = "A4:%s%d" % (get_column_letter(len(cols)), max(4, r - 1))


TREND_COLS = [RF.Col("店铺", 22), RF.Col("周（起始日）", 16), RF.Col("销售单据", 10),
              RF.Col("商品销售收入（原价）", 16, RF.MONEY), RF.Col("退货退款", 14, RF.MONEY),
              RF.Col("平台费用", 16, RF.MONEY), RF.Col("贡献（未扣成本）", 16, RF.MONEY),
              RF.Col("费用率", 10, K.PCT, border=False)]


def sheet_trend(wb, fam, ctx):
    ws = fam.table_sheet(wb, SHEETS[7], "月度趋势（%s）" % ctx["period_cn"],
                         "本月按周（周一开始，按入账日，含 IVA）。", TREND_COLS)
    r = 5
    for S in fam.stores:
        L = S["L"]
        L = L[L["line"] != C.PAYMENT]
        if not len(L):
            continue
        wk = (L["post"] - pd.to_timedelta(L["post"].dt.weekday, unit="D")).dt.date
        for w, x in L.assign(_wk=wk).groupby("_wk"):
            sales = float(x.loc[x["line"].isin((C.SALES, C.PROMO)), "amount"].sum())
            fees = float(x.loc[x["line"].isin(FEE_ROWS), "amount"].sum())
            inc = float(x.loc[x["line"].isin(C.INCOME_LINES), "amount"].sum())
            fam.write_row(ws, r, TREND_COLS, [
                S["store"], str(w), int(x.loc[x["line"] == C.SALES, "doc"].nunique()), sales,
                float(x.loc[x["line"] == C.REFUND, "amount"].sum()), fees, inc + fees,
                (-fees / sales) if sales else 0.0])
            r += 1
        r += 1


def sheet_gaps(wb, fam, ctx):
    fam.gaps_sheet(wb, SHEETS[9], "数据缺口与待办（%s）" % ctx["period_cn"], ctx["checks"],
                   notes=[("报税时要注意（不影响本报表的数字）",
                           [("提示 %d" % i, q) for i, q in enumerate(TAX_NOTES, 1)])])


# ======================================================================= entry

def build_context(stores, month, build_date=None):
    y, mo = periods.parse_month(month)
    return {"month": month, "period_cn": "%d 年 %d 月" % (y, mo),
            "who": stores[0]["store"] if len(stores) == 1 else "%d 家店铺汇总" % len(stores),
            "built": build_date or datetime.datetime.now().strftime("%Y-%m-%d %H:%M")}


def apply_costs(stores, month, checks, table=None):
    table = table or SC.load_configured()
    if table is None:
        return None
    for S in stores:
        S["cost"] = SC.apply(table, PLATFORM, COUNTRY, CUR, month, sku_items(S), COST_FACTOR)
        S["m"]["cogs"] = S["cost"].cogs
    checks.extend(SC.checks(table, [(S["store"], S["cost"]) for S in stores], PLATFORM, COUNTRY))
    return table


def results_figures(S):
    """This store's numbers for the monthly results file (core/results), from the
    same S and row specs that fill the workbook: ② (PL_ROWS, S["m"]["pl"], the cost
    table's cogs), PL_MEMO, ④ fee_groups (fee lines), ⑤ the returned documents'
    money by line, ⑦ S["sku"]; no inventory (Liverpool holds none for the seller)."""
    from core import results
    cost = S.get("cost")
    spec = [{"key": x.key, "label": x.label, "original": x.original, "kind": x.kind, "formula": x.formula,
             "value": (S["m"]["pl"][x.key] if x.kind == "d" else
                       (cost.cogs if cost is not None else 0.0) if x.kind == "in" else None)}
            for x in PL_ROWS]
    memo = [{"key": k, "label": label, "value": float(fn(S))} for k, (label, _o, fn, _n) in PL_MEMO if k]
    fees = []
    g = fee_groups(S)
    if g is not None:
        for _, x in g.iterrows():
            if x["line"] in C.FEE_LINES:
                fees.append({"name": x["orig"], "label": x["zh"], "pl_key": x["line"], "amount": float(x["amt"])})
    rets = []
    money = S["returns"].get("money")
    if money is not None and len(money):
        for (line_, zh), row in money.iterrows():
            rets.append({"key": line_, "label": zh, "count": int(row["count"]), "amount": float(row["sum"])})
    skus = []
    for sku, x in (S["sku"].iterrows() if len(S["sku"]) else ()):
        skus.append({"sku": sku, "title": x["name"] if isinstance(x["name"], str) else "",
                     "units": float(x["units"]), "units_returned": float(x["returns"]),
                     "net_sales": float(x["income"]), "fees": float(x["fees"]),
                     "contribution": float(x["contrib"]),
                     "cost": results.sku_cost(cost, sku) if _is_real_sku(sku) else None})
    return {"pl": results.pl_rows(spec), "memo": memo, "fees": fees, "returns": rets, "sku": skus,
            "inventory": None}


def build(pairs, month, out_path, strict=False, build_date=None):
    """pairs: [(store, folder)]. Returns a result dict (outcome, checks, out, and
    "results": {store: results inputs} when the workbook was written)."""
    stores = [derive(load_store(s, f, month)) for s, f in pairs]
    ctx = build_context(stores, month, build_date)
    ctx["checks"] = validate(stores)
    ctx["sku_costs"] = apply_costs(stores, month, ctx["checks"])
    bad = RF.failed_errors(ctx["checks"])
    if strict and bad:
        return {"outcome": "strict_aborted", "checks": ctx["checks"], "out": out_path}
    build_workbook(stores, ctx, out_path)
    return {"outcome": "checks_failed" if bad else "ok", "checks": ctx["checks"], "out": out_path,
            "metrics": {S["store"]: S["m"] for S in stores},
            "results": RF.results_inputs(stores, results_figures)}
