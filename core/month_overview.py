# -*- coding: utf-8 -*-
"""One month, one page: 本月总览_<YYYY-MM>.xlsx and the single ！需要处理_<YYYYMM>.txt.

Rebuilt after every run (core/run_monthly) from what is already on disk — the
monthly results files (core/results), the run log's pending human actions
(core/runlog), the download batch summaries (logs/batch_*.json), config.json and
the period-report results the runner just produced. It computes no money: every
figure is copied from a results file, and a figure the file does not carry is
left blank, never 0.

    build(month, alert_folder, run_notes=…, period_results=…) -> summary dict

Sheets: 总览 (one row per store and basis, every configured store), 需要你处理
(only what a person must act on), 版本历史 (every version, deltas), 说明.
The alert file replaces the former per-region ！报表异常_*.txt and the period
reports' ！期间报表异常_*.txt; both are removed here as superseded. It exists only
while there is at least one item, so its presence alone means "look".

build() never raises: a failure here must never change a run's exit code.
"""
import datetime
import io
import os
import re
import tempfile
import traceback

from core import layout
from core import markets
from core import report_kit as RK
from core import results as R

FIGURES = ("revenue", "fees", "contrib", "profit")
FIGURE_CN = {"revenue": "营业收入", "fees": "平台费用", "contrib": "经营贡献毛利", "profit": "税前利润"}

# (platform, basis) -> which results "pl" keys add up to each headline figure.
# tests/test_month_overview.py checks every key against the platform's P&L spec.
HEADLINE = {
    ("ML", "order"): {"revenue": ["net_sales"], "fees": ["order_fees_total", "bill_fees_total"],
                      "contrib": ["contrib"], "profit": ["profit"]},
    ("ML", "bill"): {"revenue": ["revenue"], "fees": ["fees_total"],
                     "contrib": ["contrib"], "profit": ["profit"]},
    ("FA", "main"): {"revenue": ["net_sales"], "fees": ["fee_total"],
                     "contrib": ["contrib"], "profit": ["profit"]},
    ("AZ", "main"): {"revenue": ["net_sales"], "fees": ["fee_total"],
                     "contrib": ["contrib"], "profit": ["profit"]},
    ("LV", "main"): {"revenue": ["net_sales"], "fees": ["fee_total"],
                     "contrib": ["contrib"], "profit": ["profit"]},
}

BASIS_LABEL = {"bill": "账单基准 正式", "order": "订单基准 参考", "main": "—"}
NO_REPORT = "没有报表"
CONFIG_BAD = "配置有问题"
GRADE_FILL = {R.COMPLETE: RK.C_OK, R.COMPLETE_WARN: RK.C_WARN, R.PROVISIONAL: RK.C_WARN,
              R.GAPS: RK.C_BAD, R.FAILED: RK.C_BAD, NO_REPORT: RK.C_BAD, CONFIG_BAD: RK.C_BAD}
RED_GRADES = (R.GAPS, R.FAILED, NO_REPORT, CONFIG_BAD)

GRADE_PLAIN = {
    R.COMPLETE: "数据齐全、校验都过：数字可以直接用。",
    R.COMPLETE_WARN: "数字可以用；缺的只是不影响金额的文件，或有不影响金额的提醒。",
    R.PROVISIONAL: "数字会变：账单还没关账，或这次下载完整不完整说不准。关账后点「只出报表」重出。",
    R.GAPS: "数字不完整：报表要读的文件没下到，或有校验没过。要处理（见「需要你处理」）。",
    R.FAILED: "这个版本的报表没生成出来。要处理。",
}

SHEET_OVERVIEW, SHEET_ACTIONS, SHEET_HISTORY, SHEET_NOTES = "总览", "需要你处理", "版本历史", "说明"
CHANGED_AFTER_FINAL = "定稿后有变动"

LEGACY_ALERT_RE = re.compile(r"^！报表异常_.*\.txt$")
LEGACY_PERIOD_RE = re.compile(r"^！期间报表异常_(?P<rest>.+)\.txt$")
PERIOD_LABEL_RE = re.compile(r"(\d{4}-Q[1-4]|\d{4}-\d{2}_\d{4}-\d{2}|\d{6}|\d{4})")


# ------------------------------------------------------------------ names

def overview_file_name(month):
    return "本月总览_%s.xlsx" % month


def overview_fallback_name(month):
    return "本月总览_%s(新).xlsx" % month


def alert_file_name(month):
    """！需要处理_202609.txt — the one alert file of a month."""
    return "！需要处理_%s.txt" % str(month).replace("-", "")


# ------------------------------------------------------------------ figures

def figures(doc):
    """{'revenue', 'fees', 'contrib', 'profit'} copied from a results document's
    pl rows; None where the document lacks a key (shown blank, never 0)."""
    spec = HEADLINE.get((doc.get("platform"), doc.get("basis")))
    out = dict((k, None) for k in FIGURES)
    if not spec:
        return out
    by = {}
    for r in doc.get("pl") or []:
        if isinstance(r, dict) and r.get("key") not in by:
            by[r.get("key")] = r.get("value")
    for name in FIGURES:
        keys = spec[name]
        if keys and all(by.get(k) is not None for k in keys):
            try:
                out[name] = round(sum(float(by[k]) for k in keys), 6)
            except (TypeError, ValueError):
                out[name] = None
    return out


def _fmt_run(run):
    """'20261021_233002_m202609' -> '2026-10-21 23:30'."""
    m = re.match(r"^(\d{4})(\d{2})(\d{2})_(\d{2})(\d{2})", str(run or ""))
    return "%s-%s-%s %s:%s" % m.groups() if m else (str(run or "") or "")


def _fmt_built(iso):
    s = str(iso or "")
    return s[:16].replace("T", " ") if len(s) >= 16 else s


def _same(a, b):
    return all(_eq(a.get(k), b.get(k)) for k in FIGURES)


def _eq(x, y):
    if x is None or y is None:
        return x is None and y is None
    return abs(float(x) - float(y)) < 0.005


# ------------------------------------------------------------------ gathering

class _Store(object):
    """One store of the month: where it belongs and every version per basis."""

    def __init__(self, platform, country, name, configured=True, problem=None):
        self.platform, self.country, self.name = platform, country, name
        self.configured = configured
        self.problem = problem                  # config_problems() text, or None
        self.versions = {}                      # basis -> [doc newest first]

    @property
    def bases(self):
        return ("bill", "order") if self.platform == "ML" else ("main",)

    @property
    def label(self):
        return self.name if self.configured else "%s（已不在配置里）" % self.name

    def final(self, basis):
        for d in self.versions.get(basis) or []:
            if d["grade"]["level"] in R.COUNTABLE:
                return d
        return None

    def newest(self, basis):
        vs = self.versions.get(basis) or []
        return vs[0] if vs else None

    def shown(self, basis):
        """The version the overview row shows: the final, else the newest."""
        return self.final(basis) or self.newest(basis)

    def has_any(self):
        return any(self.versions.get(b) for b in self.bases)


def _order_key(platform, country, name):
    pf = list(markets.PLATFORMS)
    cc = list(markets.COUNTRIES)
    return (pf.index(platform) if platform in pf else len(pf),
            cc.index(country) if country in cc else len(cc), str(name))


def configured_stores(notes):
    """[(platform, country, store)] for every unit of core/markets, from
    config.json. A unit whose list cannot be read becomes a note (an action)."""
    from core import store_config
    out = []
    for unit in markets.all_units():
        market = markets.unit_market(unit)
        pf, cc = markets.split(market)
        try:
            names = store_config.stores_for_market(market)
        except Exception as e:                                    # noqa: BLE001
            notes.append("config.json 有问题，定不了%s（%s）有哪些店铺：%s"
                         % (markets.label(market), unit, str(e)[:200]))
            continue
        for n in names:
            if (pf, cc, n) not in out:
                out.append((pf, cc, n))
    return out


def config_problem_stores():
    from core import store_config
    try:
        return list(store_config.config_problems())
    except Exception:                                             # noqa: BLE001
        return []


def stores_with_results(month, config_path=None):
    """[(platform, country, store)] that have a results folder for the month."""
    root = layout.results_root(config_path)
    out = []
    if not os.path.isdir(root):
        return out
    for pf in sorted(os.listdir(root)):
        p1 = os.path.join(root, pf)
        if not os.path.isdir(p1):
            continue
        for cc in sorted(os.listdir(p1)):
            p2 = os.path.join(p1, cc)
            if not os.path.isdir(p2):
                continue
            for name in sorted(os.listdir(p2)):
                if os.path.isdir(os.path.join(p2, name, month)):
                    out.append((pf, cc, name))
    return out


def gather(month, notes, config_path=None):
    """Every store the overview lists, in order, with their versions."""
    stores = {}
    for pf, cc, name in configured_stores(notes):
        stores[(pf, cc, name)] = _Store(pf, cc, name)
    for pf, cc, name in stores_with_results(month, config_path):
        if (pf, cc, name) not in stores:
            stores[(pf, cc, name)] = _Store(pf, cc, name, configured=False)
    for st in stores.values():
        for basis in st.bases:
            try:
                st.versions[basis] = R.versions_at(st.platform, st.country, st.name, month, basis,
                                                   config_path)
            except Exception:                                     # noqa: BLE001
                st.versions[basis] = []
    ordered = sorted(stores.values(), key=lambda s: _order_key(s.platform, s.country, s.name))
    for name, why in config_problem_stores():
        if not any(s.name == name for s in ordered):
            ordered.append(_Store("?", "?", name, problem=why))
    return ordered


# ------------------------------------------------------------------ why / who / next

def _cause_info(code):
    try:
        from core import download_diag
        return download_diag.CAUSES.get(code)
    except Exception:                                             # noqa: BLE001
        return None


def _join(items):
    out = []
    for x in items:
        if x and x not in out:
            out.append(x)
    return "；".join(out)


def explain_version(doc):
    """(缺什么/原因, 谁来处理, 下一步) for the version a row shows."""
    level = doc["grade"]["level"]
    reasons = list(doc["grade"].get("reasons") or [])
    why = _join(reasons)
    if level == R.COMPLETE:
        return "", "", ""
    if level == R.COMPLETE_WARN:
        return why, "—", "不用处理（不影响金额）"
    if level == R.PROVISIONAL:
        if any("关账" in r for r in reasons):
            return why, "财务", "账单关账后点「只出报表」重出"
        return why, "运营", "重跑这家店的下载（在「立即运行」里只勾它），再出报表"
    if level == R.FAILED:
        return why or "报表没有生成", "维护人", "看运行日志和 logs/report_diag/ 里这次生成的输出，修好后「只出报表」"
    # gaps
    owners, steps = [], []
    for m in doc.get("missing") or []:
        info = _cause_info((m or {}).get("cause"))
        if info:
            owners.append(info[1])
            steps.append(info[3])
    if any(r.startswith("校验未过") for r in reasons):
        owners.append("维护人 / 财务")
        steps.append("看报表【⑩ 数据缺口与待办】里这条校验的明细")
    if any("导出不完整" in r for r in reasons):
        owners.append("运营")
        steps.append("重新下载这家店（销售报表导出不完整），再出报表")
    if any("公式错误" in r for r in reasons):
        owners.append("维护人")
        steps.append("把这份报表发给维护人")
    return why, _join(owners) or "运营", _join(steps) or "处理后重跑这家店并重出报表"


def explain_missing(store, basis, batch_rows, human, other_has=False):
    """(缺什么/原因, 谁来处理, 下一步) for a store-basis without any version.
    other_has: the store's other basis does have a version this month."""
    acts = [a for a in human if a.get("store") == store]
    if acts:
        a = acts[0]
        info = _cause_info(a.get("cause"))
        return (_join("%s：%s" % (x.get("title") or "", x.get("how") or "") for x in acts),
                info[1] if info else "运营", a.get("how") or "")
    row = batch_rows.get(store)
    if row:
        status = row.get("status")
        if status == "not-allowed":
            return ("店铺不在白名单", "维护人", "把它加进 config.json 的 stores.allowed 后重跑")
        errs = [str(e) for e in row.get("errors") or []]
        causes = [(m or {}).get("cause") for m in row.get("missing") or [] if isinstance(m, dict)]
        if status not in ("ok", "partial"):
            if any("找不到" in e or "not found" in e.lower() for e in errs):
                return ("紫鸟里找不到这家店：%s" % errs[0][:160], "运营",
                        "确认店名和紫鸟里的一致，改 config.json 后重跑")
            info = _cause_info(causes[0]) if causes else None
            if info:
                return ("这家店没下到数据：%s" % info[0], info[1], info[3])
            if errs:
                return ("这家店没下到数据：%s" % errs[0][:160], "运营", "看运行日志，处理后重跑这家店")
            return "这家店没下到数据", "运营", "看运行日志，处理后重跑这家店"
        return ("下载了但没出报表", "维护人",
                "在「立即运行」里只勾这家店点「只出报表」；还不行看运行日志")
    if other_has and basis in ("bill", "order"):
        other = "订单基准" if basis == "bill" else "账单基准"
        return ("这种口径的报表没有生成（%s有）" % other, "维护人",
                "在「立即运行」里只勾这家店点「只出报表」；还不行看运行日志里出报表那一段")
    return ("本月没有下载目录或没出过报表", "运营",
            "在「立即运行」里只勾这家店跑一次「下载并出报表」")


# ------------------------------------------------------------------ run log reads

def pending_actions(log_dir=None):
    try:
        from core import runlog
        return list(runlog.pending_human_actions(log_dir).get("actions") or [])
    except Exception:                                             # noqa: BLE001
        return []


def batch_rows_for(month, log_dir=None, extra_paths=()):
    """{store: its newest batch-summary row this month} from the run log's batch
    summaries (newest first) plus the paths the runner hands over."""
    paths = [p for p in extra_paths or () if p]
    try:
        from core import runlog
        paths += [p for _r, p in runlog.batch_summaries_for_month(month, log_dir) if p]
    except Exception:                                             # noqa: BLE001
        pass
    out = {}
    seen = set()
    for p in paths:
        if p in seen:
            continue
        seen.add(p)
        try:
            from core import runlog
            doc = runlog.read_json(p)
        except Exception:                                         # noqa: BLE001
            doc = None
        if not doc:
            continue
        for row in doc.get("stores") or []:
            if isinstance(row, dict) and row.get("store") and row["store"] not in out:
                out[row["store"]] = row
    return out


# ------------------------------------------------------------------ rows

def overview_rows(stores, batch_rows, human):
    """The 总览 table as plain dicts (one per store and basis, subtotals added
    by the writer)."""
    rows = []
    for st in stores:
        if st.problem:
            rows.append({"platform": "？", "country": "？", "store": st.name, "basis": "—",
                         "grade": CONFIG_BAD, "final": "—", "version": "", "run": "", "built": "",
                         "currency": "", "decimals": 2, "figures": dict((k, None) for k in FIGURES),
                         "why": st.problem, "who": "维护人",
                         "next": "改 config.json 的 stores.catalog（写明平台和国家）后重跑",
                         "key": ("?", "?", st.name, "—"), "store_obj": st, "basis_code": None})
            continue
        pf_name = markets.platform(st.platform).name_zh if st.platform in markets.PLATFORMS else st.platform
        for basis in st.bases:
            doc = st.shown(basis)
            base = {"platform": pf_name, "country": markets.country_name(st.country), "store": st.label,
                    "basis": BASIS_LABEL.get(basis, basis), "basis_code": basis, "store_obj": st,
                    "key": (st.platform, st.country, st.name, basis)}
            if doc is None:
                why, who, nxt = explain_missing(st.name, basis, batch_rows, human,
                                                other_has=any(st.versions.get(b) for b in st.bases))
                cur, dec = _currency_of(st)
                base.update(grade=NO_REPORT, final="—", version="", run="", built="", currency=cur,
                            decimals=dec, figures=dict((k, None) for k in FIGURES), why=why, who=who, next=nxt)
            else:
                why, who, nxt = explain_version(doc)
                base.update(grade=doc["grade"]["level"],
                            final="是" if doc["grade"]["level"] in R.COUNTABLE else "否",
                            version=doc.get("version") or "", run=_fmt_run(doc.get("run")),
                            built=_fmt_built(doc.get("built_at")), currency=doc.get("currency") or "",
                            decimals=int(doc.get("decimals") or 0), figures=figures(doc),
                            why=why, who=who, next=nxt)
            rows.append(base)
    _fill_currency(rows)
    return rows


def _fill_currency(rows):
    """A store without any version this month borrows the currency the other
    stores of its market report in, so it is counted in that market's subtotal
    line (as "没有报表，未计入"). Nothing is borrowed when the market has none."""
    by_group = {}
    for r in rows:
        if r["currency"] and r["grade"] != CONFIG_BAD:
            g = (r["platform"], r["country"])
            by_group.setdefault(g, {})
            by_group[g][r["currency"]] = by_group[g].get(r["currency"], 0) + 1
    for r in rows:
        if not r["currency"] and r["grade"] != CONFIG_BAD:
            counts = by_group.get((r["platform"], r["country"]))
            if counts and len(counts) == 1:
                r["currency"] = list(counts)[0]
                r["decimals"] = max([x["decimals"] for x in rows if x["currency"] == r["currency"]
                                     and x["grade"] not in (NO_REPORT, CONFIG_BAD)] or [2])


def _currency_of(st):
    """Currency of a store without a version this month: from any other
    version of the store (other basis), else blank."""
    for b in st.bases:
        for d in st.versions.get(b) or []:
            return d.get("currency") or "", int(d.get("decimals") or 0)
    return "", 2


def action_rows(rows, human, run_notes, period_results, config_notes):
    """The 需要你处理 table: [{unit, store, problem, who, next, source}]."""
    out = []
    # (i) pending human actions from the run log
    for a in human:
        info = _cause_info(a.get("cause"))
        unit = a.get("region") or ""
        try:
            unit_cn = markets.unit_label(unit) if unit else ""
        except Exception:                                         # noqa: BLE001
            unit_cn = unit
        out.append({"unit": unit_cn, "store": a.get("store") or "", "problem": a.get("title") or "",
                    "who": info[1] if info else "运营", "next": a.get("how") or "",
                    "source": "下载记录（%s）" % (a.get("run_id") or "")})
    # (iv) what the runner says about this run (config problems, crashes, failures)
    for text in list(config_notes) + list(run_notes or ()):
        out.append(_note_row(text))
    # (ii) stores in config without a version; (iii) bad versions
    seen_missing = set()
    for r in rows:
        st = r.get("store_obj")
        if r["grade"] == CONFIG_BAD:
            out.append({"unit": "？", "store": r["store"], "problem": r["why"], "who": r["who"],
                        "next": r["next"], "source": "config.json"})
            continue
        unit = "%s %s" % (r["platform"], r["country"])
        if r["grade"] == NO_REPORT:
            if not st.configured:
                continue
            if st.has_any():
                problem = "没有%s的报表" % r["basis"].split(" ")[0]
            else:
                if st.name in seen_missing:
                    continue
                seen_missing.add(st.name)
                problem = "本月没有报表"
            out.append({"unit": unit, "store": r["store"], "problem": "%s：%s" % (problem, r["why"]),
                        "who": r["who"], "next": r["next"], "source": "月度结果"})
        elif r["grade"] == R.FAILED:
            out.append({"unit": unit, "store": r["store"], "problem": "报表没生成（%s）：%s"
                        % (r["basis"].split(" ")[0] if r["basis"] != "—" else "", r["why"]),
                        "who": r["who"], "next": r["next"], "source": "月度结果"})
        elif r["grade"] == R.GAPS:
            doc = st.shown(r["basis_code"])
            blocking = [x for x in (doc["grade"].get("reasons") or [])
                        if not x.startswith("缺下载（不影响金额）") and not x.startswith("提醒")]
            if blocking:
                out.append({"unit": unit, "store": r["store"], "problem": "有缺口（%s）：%s"
                            % (r["basis"].split(" ")[0] if r["basis"] != "—" else "", _join(blocking)),
                            "who": r["who"], "next": r["next"], "source": "月度结果"})
    # (v) period reports that are incomplete or failed
    for res in period_results or ():
        try:
            from core import period_report as PR
            if res.get("status") not in (PR.INCOMPLETE, PR.ERROR):
                continue
            lines = PR.summary_lines(res)
        except Exception:                                         # noqa: BLE001
            continue
        try:
            unit_cn = markets.label(res.get("market"))
        except Exception:                                         # noqa: BLE001
            unit_cn = str(res.get("market") or "")
        out.append({"unit": unit_cn, "store": "", "problem": " ".join(x.strip() for x in lines),
                    "who": "财务", "source": "期间报表",
                    "next": "把缺的那个月的报表重出（「只出报表」），有了完整版本后在窗口「期间报表」页再生成一次"})
    return out


def _note_row(text):
    try:
        from core import runlog
        region, body = runlog.split_region(text)
    except Exception:                                             # noqa: BLE001
        region, body = None, text
    try:
        unit_cn = markets.unit_label(region) if region else ""
    except Exception:                                             # noqa: BLE001
        unit_cn = region or ""
    body = str(body or "").strip()
    if "config.json" in body or "白名单" in body:
        who, nxt = "维护人", "改好 config.json 后重跑"
    elif "运行器内部出错" in body or "期间报表" in body:
        who, nxt = "维护人", "看运行日志最后一段"
    elif "需要人工处理" in body:
        who, nxt = "运营", "处理完在「需要人工处理」页点「重跑这些店铺并重出全部报表」"
    elif "报表" in body and ("没有写出来" in body or "失败" in body or "未生成" in body):
        who, nxt = "维护人", "看运行日志里这一步的输出，修好后「只出报表」"
    elif "没下到数据" in body or "下载" in body:
        who, nxt = "运营", "看运行日志，处理后重跑这些店铺"
    else:
        who, nxt = "运营 / 维护人", "看运行日志"
    return {"unit": unit_cn, "store": "", "problem": body, "who": who, "next": nxt, "source": "本轮运行"}


def history_rows(stores):
    """The 版本历史 table: every version of every store-basis, newest first."""
    out = []
    for st in stores:
        if st.problem:
            continue
        for basis in st.bases:
            vs = list(st.versions.get(basis) or [])
            if not vs:
                continue
            final = st.final(basis)
            figs = [figures(d) for d in vs]
            for i, d in enumerate(vs):
                prev = figs[i + 1] if i + 1 < len(vs) else None
                delta = {}
                for k in FIGURES:
                    a, b = figs[i].get(k), prev.get(k) if prev else None
                    delta[k] = round(a - b, 6) if (a is not None and b is not None) else None
                note = ""
                if final is not None and d is final:
                    older = [j for j in range(i + 1, len(vs)) if vs[j]["grade"]["level"] in R.COUNTABLE]
                    if any(not _same(figs[i], figs[j]) for j in older):
                        note = CHANGED_AFTER_FINAL
                out.append({"store": st.label, "basis": BASIS_LABEL.get(basis, basis),
                            "version": d.get("version") or "", "run": _fmt_run(d.get("run")),
                            "built": _fmt_built(d.get("built_at")), "grade": d["grade"]["level"],
                            "final": "是" if d is final else "否", "figures": figs[i], "delta": delta,
                            "decimals": int(d.get("decimals") or 0), "note": note})
    return out


# ------------------------------------------------------------------ workbook

OVERVIEW_COLS = [("平台", 10), ("国家", 8), ("店铺", 26), ("口径", 13), ("完整度", 11), ("是否最终", 9),
                 ("报表版本", 30), ("数据下载于", 16), ("报表生成于", 16), ("币种", 7),
                 ("营业收入", 15), ("平台费用", 15), ("经营贡献毛利", 15), ("税前利润", 15),
                 ("缺什么 / 原因", 55), ("谁来处理", 12), ("下一步", 45)]
ACTION_COLS = [("平台 / 国家", 16), ("店铺", 26), ("问题", 70), ("谁来处理", 12), ("下一步", 60), ("来源", 22)]
HISTORY_COLS = [("店铺", 26), ("口径", 13), ("报表版本", 30), ("下载于", 16), ("生成于", 16), ("完整度", 11),
                ("是否最终", 9), ("营业收入", 15), ("平台费用", 15), ("经营贡献毛利", 15), ("税前利润", 15),
                ("较上一版本：营业收入", 15), ("较上一版本：平台费用", 15), ("较上一版本：经营贡献毛利", 15),
                ("较上一版本：税前利润", 15), ("备注", 16)]
HEADER_ROW = 4
MONEY_FIRST = 11                                # 总览: K..N are the money columns


def _grade_text(level):
    return R.GRADE_CN.get(level, level)


def _setup(ws, cols, n_rows):
    from openpyxl.utils import get_column_letter
    RK.header_cells(ws, HEADER_ROW, [c for c, _ in cols])
    ws.row_dimensions[HEADER_ROW].height = 30
    for i, (_c, w) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A%d" % (HEADER_ROW + 1)
    last = max(HEADER_ROW + n_rows, HEADER_ROW)
    ws.auto_filter.ref = "A%d:%s%d" % (HEADER_ROW, get_column_letter(len(cols)), last)
    ws.print_title_rows = "%d:%d" % (HEADER_ROW, HEADER_ROW)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _cell(ws, row, col, value, fmt=None, bold=False, color="000000", fill=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = RK.F(10, bold, color)
    c.border = RK.BOX
    c.alignment = RK.wrap("top") if isinstance(value, str) and len(value) > 20 else RK.wrap("center")
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = RK.fill(fill)
    return c


def _grade_cell(ws, row, col, level):
    red = level in RED_GRADES
    return _cell(ws, row, col, _grade_text(level) if level not in (NO_REPORT, CONFIG_BAD) else level,
                 bold=red, color=RK.C_RED if red else "000000", fill=GRADE_FILL.get(level))


def write_overview_sheet(ws, month, rows, built_at):
    RK.title_rows(ws, "%s 本月总览" % month,
                  "每家店每种口径一行；金额抄自月度结果文件，不在这里算。美客多以账单基准为正式口径，订单基准供参考。",
                  "生成于 %s。小计只在同一币种内相加，不同币种各一行。" % built_at)
    r = HEADER_ROW + 1
    groups = []
    for row in rows:
        g = (row["platform"], row["country"])
        if not groups or groups[-1][0] != g:
            groups.append((g, []))
        groups[-1][1].append(row)
    for (_g, grows) in groups:
        for row in grows:
            _write_overview_row(ws, r, row)
            r += 1
        for basis in _bases_in(grows):
            for cur in _currencies_in(grows, basis):
                members = [x for x in grows if x["basis"] == basis and x["currency"] == cur]
                _write_subtotal(ws, r, members, basis, cur)
                r += 1
    _setup(ws, OVERVIEW_COLS, r - HEADER_ROW - 1)


def _bases_in(rows):
    out = []
    for x in rows:
        if x["basis"] not in out and x["grade"] != CONFIG_BAD:
            out.append(x["basis"])
    return out


def _currencies_in(rows, basis):
    out = []
    for x in rows:
        if x["basis"] == basis and x["currency"] and x["currency"] not in out:
            out.append(x["currency"])
    return out


def _write_overview_row(ws, r, row):
    fmt = RK.money_format(row["decimals"])
    vals = [row["platform"], row["country"], row["store"], row["basis"]]
    for i, v in enumerate(vals, 1):
        _cell(ws, r, i, v)
    _grade_cell(ws, r, 5, row["grade"])
    _cell(ws, r, 6, row["final"])
    _cell(ws, r, 7, row["version"])
    _cell(ws, r, 8, row["run"])
    _cell(ws, r, 9, row["built"])
    _cell(ws, r, 10, row["currency"])
    for i, k in enumerate(FIGURES):
        _cell(ws, r, MONEY_FIRST + i, row["figures"].get(k), fmt=fmt)
    red = row["grade"] in RED_GRADES
    _cell(ws, r, 15, row["why"], color=RK.C_RED if red else "000000")
    _cell(ws, r, 16, row["who"])
    _cell(ws, r, 17, row["next"])


def _write_subtotal(ws, r, members, basis, cur):
    decimals = max([m["decimals"] for m in members] or [2])
    fmt = RK.money_format(decimals)
    all_final = all(m["final"] == "是" for m in members)
    _cell(ws, r, 1, members[0]["platform"], bold=True, fill=RK.C_TOTAL)
    _cell(ws, r, 2, members[0]["country"], bold=True, fill=RK.C_TOTAL)
    _cell(ws, r, 3, "小计（%s）" % cur, bold=True, fill=RK.C_TOTAL)
    _cell(ws, r, 4, basis, bold=True, fill=RK.C_TOTAL)
    _cell(ws, r, 5, "", fill=RK.C_TOTAL)
    _cell(ws, r, 6, "全部定稿" if all_final else "含未定稿", bold=True,
          color="000000" if all_final else RK.C_AMBER, fill=RK.C_TOTAL)
    for col in range(7, 11):
        _cell(ws, r, col, "", fill=RK.C_TOTAL)
    for i, k in enumerate(FIGURES):
        vals = [m["figures"].get(k) for m in members if m["figures"].get(k) is not None]
        _cell(ws, r, MONEY_FIRST + i, round(sum(vals), 6) if vals else None, fmt=fmt, bold=True,
              fill=RK.C_TOTAL)
    n_blank = len([m for m in members if all(m["figures"].get(k) is None for k in FIGURES)])
    _cell(ws, r, 15, "%d 家店；%d 家没有数字（没有报表 / 失败），未计入" % (len(members), n_blank) if n_blank
          else "%d 家店" % len(members), fill=RK.C_TOTAL)
    _cell(ws, r, 16, "", fill=RK.C_TOTAL)
    _cell(ws, r, 17, "", fill=RK.C_TOTAL)


def write_actions_sheet(ws, month, actions):
    RK.title_rows(ws, "%s 需要你处理" % month,
                  "只列要人动手的事；平台还在生成、报表不读的文件、不影响金额的提醒都不在这里。",
                  "没有一行 = 这个月没有要处理的事。" if not actions else "共 %d 项。" % len(actions))
    r = HEADER_ROW + 1
    for a in actions:
        for i, k in enumerate(("unit", "store", "problem", "who", "next", "source"), 1):
            _cell(ws, r, i, a.get(k) or "", color=RK.C_RED if k == "problem" else "000000",
                  bold=(k == "problem"))
        r += 1
    if not actions:
        _cell(ws, r, 1, "这个月没有需要处理的事。", color=RK.C_GREEN, bold=True)
        r += 1
    _setup(ws, ACTION_COLS, r - HEADER_ROW - 1)


def write_history_sheet(ws, month, hist):
    RK.title_rows(ws, "%s 版本历史" % month,
                  "每家店每种口径的每一个版本，新的在上；「较上一版本」是和下面那一行（更早的版本）相比。",
                  "「%s」：定稿版本和更早的一个定稿版本数字不同 —— 已发出去的数字变了，要告诉收件人。" % CHANGED_AFTER_FINAL)
    r = HEADER_ROW + 1
    for h in hist:
        fmt = RK.money_format(h["decimals"])
        _cell(ws, r, 1, h["store"])
        _cell(ws, r, 2, h["basis"])
        _cell(ws, r, 3, h["version"])
        _cell(ws, r, 4, h["run"])
        _cell(ws, r, 5, h["built"])
        _grade_cell(ws, r, 6, h["grade"])
        _cell(ws, r, 7, h["final"])
        for i, k in enumerate(FIGURES):
            _cell(ws, r, 8 + i, h["figures"].get(k), fmt=fmt)
        for i, k in enumerate(FIGURES):
            _cell(ws, r, 12 + i, h["delta"].get(k), fmt=fmt)
        _cell(ws, r, 16, h["note"], bold=bool(h["note"]), color=RK.C_RED if h["note"] else "000000")
        r += 1
    _setup(ws, HISTORY_COLS, r - HEADER_ROW - 1)


def notes_lines(month, config_path=None):
    """The 说明 sheet's text: [(heading, text)]."""
    out = [("完整度是什么意思", "")]
    for level in R.GRADES:
        out.append(("  " + R.GRADE_CN[level], GRADE_PLAIN[level]))
    out.append(("  " + NO_REPORT, "这家店这个月一个版本都没有：没下载过，或下载了没出报表。"))
    out.append(("是否最终", "「是」= 这是这个月这家店这种口径的最终版本（最新的一个「完整」或「完整·有提醒」版本），期间报表用的就是它；"
                           "「否」= 显示的是最新的一个版本，但它还不算数。"))
    out.append(("哪种口径是正式的",
                "美客多每家店两份报表：账单基准（_账单基准.xlsx，以平台账单为准，会计入账用的正式口径）和订单基准"
                "（以订单日为准，经营分析参考）。其他平台只有一种口径（口径一栏是 —）。"))
    out.append(("营业收入 / 平台费用是什么", ""))
    out.append(("  美客多 账单基准", "营业收入 = 净商品销售额 ＋ 买家运费 − 卖家优惠 ＋ 平台补贴（本期计费）；平台费用 = 本期账单全部费用。"))
    out.append(("  美客多 订单基准", "营业收入 = 净销售收入（按订单日，扣退款）；平台费用 = 订单级交易费用 ＋ 月度账单费用。"))
    out.append(("  Falabella 智利", "净销售收入与费用都含 IVA（会计师 2026-10 的决定）。"))
    out.append(("  亚马逊 墨西哥", "净销售收入与费用都不含 IVA。"))
    out.append(("  利物浦 墨西哥", "按平台账本日期归集，含 IVA。"))
    out.append(("经营贡献毛利 / 税前利润", "经营贡献毛利 = 营业收入 − 平台费用（未扣商品成本）；税前利润 = 经营贡献毛利 − 商品采购成本"
                                        "（没配 SKU 成本表时成本为 0，两者相等）。"))
    out.append(("文件在哪", ""))
    try:
        for k, v in layout.describe(config_path):
            out.append(("  " + k, str(v)))
    except Exception as e:                                        # noqa: BLE001
        out.append(("  （读不到目录设置）", str(e)[:200]))
    out.append(("怎么重跑一家店",
                '.venv\\Scripts\\python run_monthly.py --region <区域> --stores "<店名>" --month %s'
                "　（区域：MX / CL / PE 是美客多的国家，其他平台写 FA-CL、AZ-MX、LV-MX 这样的市场代码；"
                "或者双击 月度报表.cmd，在「立即运行」里只勾这家店）" % month))
    out.append(("这份总览什么时候更新", "每跑完一轮（手动、定时、命令行都一样）自动重建；「！需要处理_%s.txt」只在有事要处理时存在，"
                                      "没事就自动删掉。" % month.replace("-", "")))
    return out


def write_notes_sheet(ws, month, config_path=None):
    RK.title_rows(ws, "%s 总览说明" % month)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 120
    r = 3
    for head, text in notes_lines(month, config_path):
        c = ws.cell(row=r, column=1, value=head)
        c.font = RK.F(10, not head.startswith("  "), RK.C_NAVY if not head.startswith("  ") else "000000")
        c.alignment = RK.wrap("top")
        d = ws.cell(row=r, column=2, value=text)
        d.font = RK.F(10)
        d.alignment = RK.wrap("top")
        ws.row_dimensions[r].height = RK.est_height(text, per_line=90)
        r += 1
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def write_workbook(path, month, rows, actions, hist, built_at, config_path=None):
    """Write the workbook atomically. Returns the path written (the fallback
    name when the target is locked by Excel)."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_OVERVIEW
    write_overview_sheet(ws, month, rows, built_at)
    write_actions_sheet(wb.create_sheet(SHEET_ACTIONS), month, actions)
    write_history_sheet(wb.create_sheet(SHEET_HISTORY), month, hist)
    write_notes_sheet(wb.create_sheet(SHEET_NOTES), month, config_path)
    folder = os.path.dirname(path)
    os.makedirs(folder, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=".本月总览_", suffix=".xlsx.part", dir=folder)
    os.close(fd)
    try:
        wb.save(tmp)
        try:
            os.replace(tmp, path)
            return path
        except PermissionError:
            alt = os.path.join(folder, overview_fallback_name(month))
            os.replace(tmp, alt)
            return alt
    finally:
        try:
            if os.path.exists(tmp):
                os.remove(tmp)
        except OSError:
            pass


# ------------------------------------------------------------------ alert file

def alert_text(month, actions, entrance=None, now=None, log_path=None):
    now = now or datetime.datetime.now()
    lines = ["%s 需要处理的事（%d 项）" % (month, len(actions)),
             "时间：%s" % now.strftime("%Y-%m-%d %H:%M:%S")]
    if entrance:
        try:
            from core import runlog
            lines.append("这一轮是：%s（%s）" % (runlog.ENTRANCE_ZH.get(entrance, entrance),
                                             runlog.ENTRANCE_WHAT.get(entrance, "")))
        except Exception:                                         # noqa: BLE001
            lines.append("这一轮是：%s" % entrance)
    if log_path:
        lines.append("完整日志：%s（也可以双击 月度报表.cmd，在右边的「运行日志」里挑这一轮来看）" % log_path)
    lines.append("")
    for i, a in enumerate(actions, 1):
        head = " ".join(x for x in (a.get("unit"), a.get("store")) if x)
        lines.append("%d. %s%s" % (i, (head + " —— ") if head else "", a.get("problem") or ""))
        if a.get("who"):
            lines.append("   谁来处理：%s" % a["who"])
        if a.get("next"):
            lines.append("   下一步：%s" % a["next"])
        if a.get("source"):
            lines.append("   来源：%s" % a["source"])
        lines.append("")
    lines.append("完整清单见 %s 的「%s」页；处理办法见 README.md 第 12 节。" % (overview_file_name(month), SHEET_ACTIONS))
    return "\n".join(lines)


def write_alert(folder, month, actions, entrance=None, log_path=None):
    """Write the alert when there is something to do, else remove it.
    Returns (path written or None, path removed or None)."""
    p = os.path.join(folder, alert_file_name(month))
    if actions:
        os.makedirs(folder, exist_ok=True)
        tmp = p + ".part"
        with io.open(tmp, "w", encoding="utf-8") as f:
            f.write(alert_text(month, actions, entrance, log_path=log_path))
        os.replace(tmp, p)
        return p, None
    if os.path.exists(p):
        os.remove(p)
        return None, p
    return None, None


def _period_is_for(rest, month):
    """Does a ！期间报表异常_<rest>.txt concern this month? (period label in the name)"""
    ym = month.replace("-", "")
    for lab in PERIOD_LABEL_RE.findall(rest):
        if lab == ym:
            return True
        if re.match(r"^\d{4}-Q[1-4]$", lab) and month in R.quarter_months(lab):
            return True
        if re.match(r"^\d{4}-\d{2}_\d{4}-\d{2}$", lab):
            a, b = lab.split("_")
            try:
                if month in R.months_between(a, b):
                    return True
            except ValueError:
                pass
        if re.match(r"^\d{4}$", lab) and lab == month[:4]:
            return True
    return False


def remove_legacy_alerts(folder, month):
    """Delete the per-region ！报表异常_*<YYYYMM>*.txt and the ！期间报表异常_*
    files of this month: the single alert and the overview supersede them."""
    gone = []
    if not os.path.isdir(folder):
        return gone
    ym = month.replace("-", "")
    for name in sorted(os.listdir(folder)):
        p = os.path.join(folder, name)
        if LEGACY_ALERT_RE.match(name) and ym in name:
            pass
        else:
            m = LEGACY_PERIOD_RE.match(name)
            if not (m and _period_is_for(m.group("rest"), month)):
                continue
        try:
            os.remove(p)
            gone.append(p)
        except OSError:
            continue
    return gone


# ------------------------------------------------------------------ entry point

def build(month, alert_folder, run_notes=None, period_results=None, config_path=None, say=print,
          entrance=None, log_dir=None, batch_summaries=None, now=None, log_path=None):
    """Rebuild the month's overview workbook and its single alert file.

    Returns {"path", "actions", "alert_path", "rows", "removed", "error"}; never
    raises. run_notes: plain strings from the runner (its urgent lines, config
    problems, crashes). period_results: core.period_report.generate results of
    this run. log_dir: where the run log lives (default: the layout's).
    """
    out = {"path": None, "actions": [], "alert_path": None, "rows": 0, "removed": [], "error": None}
    now = now or datetime.datetime.now()
    try:
        notes = []
        stores = gather(month, notes, config_path)
        human = pending_actions(log_dir)
        batch_rows = batch_rows_for(month, log_dir, batch_summaries)
        rows = overview_rows(stores, batch_rows, human)
        actions = action_rows(rows, human, run_notes, period_results, notes)
        hist = history_rows(stores)
        out["rows"] = len(rows)
        out["actions"] = actions
        target = os.path.join(alert_folder, overview_file_name(month))
        written = write_workbook(target, month, rows, actions, hist, now.strftime("%Y-%m-%d %H:%M"),
                                 config_path)
        out["path"] = written
        if written != target:
            say("本月总览写不进 %s（文件多半正被 Excel 打开着），写成了 %s" % (target, written))
        else:
            say("已更新本月总览：%s" % written)
    except Exception as e:                                        # noqa: BLE001
        out["error"] = "%s: %s" % (e.__class__.__name__, e)
        say("[警告] 本月总览没生成：%s\n%s" % (out["error"], traceback.format_exc()))
        actions = out["actions"] or [{"unit": "", "store": "", "problem": "本月总览没生成：%s" % out["error"],
                                      "who": "维护人", "next": "看运行日志", "source": "本月总览"}]
        for x in run_notes or ():
            actions.append(_note_row(x))
        out["actions"] = actions
    try:
        wrote, removed = write_alert(alert_folder, month, out["actions"], entrance, log_path)
        out["alert_path"] = wrote
        if wrote:
            say("已写出待处理清单：%s" % wrote)
        elif removed:
            say("没有需要处理的事，删掉了 %s" % removed)
    except Exception as e:                                        # noqa: BLE001
        out["error"] = (out["error"] or "") + " 待处理清单写不进去：%s" % str(e)[:120]
        say("[警告] 待处理清单写不进去：%s" % str(e)[:160])
    try:
        out["removed"] = remove_legacy_alerts(alert_folder, month)
        for p in out["removed"]:
            say("删掉了旧的告警文件（已由本月总览取代）：%s" % p)
    except Exception as e:                                        # noqa: BLE001
        say("[警告] 旧告警文件删不掉：%s" % str(e)[:160])
    return out
