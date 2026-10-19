# -*- coding: utf-8 -*-
"""
The SKU cost table — one optional file, shared by every platform's report, in which
finance enters the unit cost of each SKU once. With it, ② of every workbook shows
the cost of goods and the profit after it; ⑦ shows each SKU's cost and margin.

Where it is: config.json

    "reports": {"sku_costs": "sku_costs.xlsx"}

A relative path is relative to the repository root. **No key, no table**: the
workbooks are exactly what they were before this module existed. The file is
business data and lives outside git (.gitignore).

The file (first sheet of an .xlsx, or a UTF-8 .csv), one header row:

    平台      platform code (core/markets.py: ML, FA, AZ …); blank = any platform
    国家      ISO country code (MX, CL …);                   blank = any country
    SKU       the seller's SKU, as the platform's files show it (required)
    单位成本  landed cost per unit — purchase + freight + duties (required, ≥ 0)
    币种      currency of the cost (required; must be the report's currency —
              there is no FX conversion)
    生效日期  optional: the row applies from that date; for a month, the latest
              row effective on or before the month's last day is used
    备注      ignored (any other column is ignored too)

English headers are accepted as well (platform, country, sku, unit_cost,
currency, effective_date, note).

A report's SKU is matched (trimmed, case-insensitive) on (platform, country) first,
then (platform, any country), then (any platform, country), then (any, any).

How a platform's report_build uses it (the whole contract):

    from core import sku_costs as SC

    table = SC.load_configured()                     # None: not configured
    items = [SC.Item(sku, name, units, returned, sales), …]   # one per ⑦ SKU
    res = SC.apply(table, "FA", "CL", "CLP", "2026-09", items)   # None if table is None
    checks += SC.checks(table, [(store, res)], "FA", "CL")
    res.cogs                      # ② 商品采购成本 (covered SKUs only, positive)
    res.line(sku)                 # per SKU: .unit_cost (None = uncovered), .cost, .status
    SC.pl_note(res), SC.status_text(line), SC.source_line(table), SC.policy_line(basis)

Platform names never appear here; each report passes its own codes and words.
"""
import calendar
import collections
import csv
import datetime
import io
import json
import os
import re

from core import markets
from core import report_kit as K

CONFIG_KEY = ("reports", "sku_costs")
SUGGESTED_NAME = "sku_costs.xlsx"

# canonical column -> accepted headers (compared after _norm_header)
HEADERS = collections.OrderedDict([
    ("platform", ("平台", "平台代码", "platform", "platformcode")),
    ("country", ("国家", "国家代码", "country", "countrycode")),
    ("sku", ("sku", "卖家sku", "sellersku")),
    ("cost", ("单位成本", "成本", "unitcost", "cost", "landedcost")),
    ("currency", ("币种", "货币", "currency")),
    ("effective", ("生效日期", "生效", "effectivedate", "effective", "effectivefrom", "validfrom")),
    ("note", ("备注", "note", "notes", "remark", "comment")),
])
REQUIRED = ("sku", "cost", "currency")
# the columns finance sees, in this order (tools/sku_cost_template writes them)
COLUMN_LABELS = collections.OrderedDict([
    ("platform", "平台"), ("country", "国家"), ("sku", "SKU"), ("cost", "单位成本"),
    ("currency", "币种"), ("effective", "生效日期"), ("note", "备注")])

# status of one SKU of a report
OK, NO_COST, CURRENCY, CONFLICT = "ok", "no_cost", "currency", "conflict"


def sku_key(sku):
    """The matching key of a SKU: trimmed, case-insensitive. A number read from a
    spreadsheet (12345.0) is the text 12345."""
    if sku is None:
        return ""
    if isinstance(sku, float):
        if sku != sku:
            return ""
        if sku.is_integer():
            sku = int(sku)
    return str(sku).strip().upper()


def _norm_header(h):
    s = str(h or "").strip().lower()
    s = re.sub(r"[（(].*?[)）]", "", s)            # "单位成本（录入）" -> "单位成本"
    return re.sub(r"[\s_\-/·.]+", "", s)


_ALIASES = dict((_norm_header(a), k) for k, names in HEADERS.items() for a in names)


def _text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v != v:
        return ""
    return str(v).strip()


# ------------------------------------------------------------------ where

def configured_path(config_path=None):
    """The table's absolute path from config.json → reports.sku_costs, or None
    when the key is absent (or config.json does not exist). A config.json that
    cannot be parsed raises store_config.StoreConfigError, as everywhere."""
    from core import store_config
    from core.paths import ROOT
    path = config_path or store_config.CONFIG_PATH
    if not os.path.exists(path):
        return None
    try:
        with io.open(path, encoding="utf-8") as fh:
            cfg = json.load(fh)
    except (ValueError, IOError, OSError) as e:
        raise store_config.StoreConfigError("config.json 读不出来（%s）" % e)
    node = cfg
    for k in CONFIG_KEY:
        node = node.get(k) if isinstance(node, dict) else None
    if not node or not isinstance(node, str) or not node.strip():
        return None
    p = os.path.expanduser(node.strip())
    return p if os.path.isabs(p) else os.path.normpath(os.path.join(ROOT, p))


# ------------------------------------------------------------------ the table

Row = collections.namedtuple("Row", "line platform country sku key cost currency effective")
Row.__doc__ = """One usable row of the file. line = row number in the file (header = 1
for csv, the sheet row for xlsx); effective = a date or None (= always)."""

Problem = collections.namedtuple("Problem", "line platform country sku text")


class Table(object):
    """The cost file as read. `error` set = the file could not be used at all."""

    def __init__(self, path):
        self.path = path
        self.name = os.path.basename(path or "")
        self.rows = []
        self.problems = []          # rows ignored, with the reason
        self.duplicates = []        # [(platform, country, sku, effective, [lines])]
        self.error = None
        self.modified = None
        if path and os.path.exists(path):
            self.modified = datetime.datetime.fromtimestamp(os.path.getmtime(path))

    @property
    def usable(self):
        return self.error is None

    def in_scope(self, platform, country):
        """Rows that could apply to a report of this platform and country."""
        p, c = platform.upper(), country.upper()
        return [r for r in self.rows if r.platform in ("", p) and r.country in ("", c)]


def _scoped(platform, country, p, c):
    return (p or "") in ("", platform.upper()) and (c or "") in ("", country.upper())


def _parse_date(v):
    if v is None or _text(v) == "":
        return None, None
    if isinstance(v, datetime.datetime):
        return v.date(), None
    if isinstance(v, datetime.date):
        return v, None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if 20000 < v < 80000:                       # an Excel serial date in a csv
            return (datetime.date(1899, 12, 30) + datetime.timedelta(days=int(v))), None
        return None, "生效日期 %r 看不懂" % v
    s = _text(v)
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.datetime.strptime(s, fmt).date(), None
        except ValueError:
            pass
    m = re.match(r"^(\d{4})[-/.](\d{1,2})$", s)
    if m and 1 <= int(m.group(2)) <= 12:
        return datetime.date(int(m.group(1)), int(m.group(2)), 1), None
    return None, "生效日期 %r 看不懂（写成 2026-09-01）" % s


def _parse_cost(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if v == v else None
    s = _text(v).replace(" ", "").replace("$", "").replace(",", "")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _read_cells(path):
    """[(line number, [cell values])] of the first sheet / the csv."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            return [(i, list(r)) for i, r in enumerate(ws.iter_rows(values_only=True), 1)]
        finally:
            wb.close()
    if ext in (".csv", ".txt"):
        with io.open(path, encoding="utf-8-sig", newline="") as fh:
            text = fh.read()
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [(i, r) for i, r in enumerate(csv.reader(io.StringIO(text), dialect), 1)]
    raise ValueError("只认 .xlsx 或 .csv（这是 %s）" % (ext or "没有扩展名"))


def load(path):
    """Read the cost file. Never raises: a file that cannot be used comes back
    with `error` set; rows that cannot be used are in `problems`."""
    T = Table(path)
    if not path or not os.path.exists(path):
        T.error = "找不到成本表文件 %s" % (path or "(未写路径)")
        return T
    try:
        rows = _read_cells(path)
    except Exception as e:
        T.error = "成本表 %s 读不出来：%s" % (T.name, e)
        return T
    header_at, cols = None, {}
    for i, (line, cells) in enumerate(rows[:20]):
        found = {}
        for j, h in enumerate(cells):
            k = _ALIASES.get(_norm_header(h))
            if k and k not in found:
                found[k] = j
        if all(k in found for k in REQUIRED):
            header_at, cols = i, found
            break
    if header_at is None:
        T.error = ("成本表 %s 没有表头行：需要列 %s（或英文 sku / unit_cost / currency）"
                   % (T.name, "、".join(COLUMN_LABELS[k] for k in REQUIRED)))
        return T
    get = lambda cells, k: cells[cols[k]] if k in cols and cols[k] < len(cells) else None
    known = set(markets.PLATFORMS)
    seen = collections.defaultdict(list)
    for line, cells in rows[header_at + 1:]:
        if not any(_text(c) for c in cells):
            continue
        p = _text(get(cells, "platform")).upper()
        c = _text(get(cells, "country")).upper()
        raw_sku = get(cells, "sku")
        key = sku_key(raw_sku)
        sku = _text(raw_sku) if not isinstance(raw_sku, float) else key
        bad = lambda text: T.problems.append(Problem(line, p, c, sku, text))
        # a bad code is shown in every report (it could have been meant for any)
        if p and p not in known:
            T.problems.append(Problem(line, "", "", sku, "平台代码 %r 不认识（可用：%s）"
                                      % (p, "、".join(sorted(known)))))
            continue
        if c and not re.match(r"^[A-Z]{2}$", c):
            T.problems.append(Problem(line, p, "", sku, "国家代码 %r 应是两个字母（MX、CL …）" % c))
            continue
        if not key:
            if _text(get(cells, "cost")):
                bad("没有 SKU")
            continue
        raw_cost = get(cells, "cost")
        if _text(raw_cost) == "":
            continue        # not entered yet (a template row): the SKU is simply uncovered
        cost = _parse_cost(raw_cost)
        if cost is None:
            bad("单位成本 %r 不是数字" % (_text(raw_cost) or "（空）"))
            continue
        if cost < 0:
            bad("单位成本 %s 是负数" % raw_cost)
            continue
        cur = _text(get(cells, "currency")).upper()
        if not re.match(r"^[A-Z]{3}$", cur):
            bad("币种 %r 应是三个字母（MXN、CLP …）" % (cur or "（空）"))
            continue
        eff, why = _parse_date(get(cells, "effective"))
        if why:
            bad(why)
            continue
        r = Row(line, p, c, sku, key, cost, cur, eff)
        T.rows.append(r)
        seen[(p, c, key, eff)].append(r)
    for (p, c, key, eff), rs in seen.items():
        if len(rs) > 1:
            T.duplicates.append((p, c, rs[0].sku, eff, [r.line for r in rs]))
    return T


def load_configured(config_path=None):
    """The configured table (read), or None when config.json names none."""
    path = configured_path(config_path)
    return load(path) if path else None


# ------------------------------------------------------------------ one report

Item = collections.namedtuple("Item", "sku name units returned sales")
Item.__doc__ = """One SKU of a report's ⑦: units delivered in the month, units
returned in the month (0 when the platform's data does not identify them), sales
(the ② revenue line's amount for the SKU — used to weigh coverage)."""


class Line(object):
    """One SKU's cost in one report."""

    def __init__(self, item, status, row=None, rows=()):
        self.item = item
        self.status = status
        self.row = row              # the Row used (OK) or rejected (CURRENCY)
        self.rows = list(rows)      # CONFLICT: the rows in conflict

    @property
    def cost_units(self):
        return int(self.item.units or 0) - int(self.item.returned or 0)

    @property
    def unit_cost(self):
        return self.row.cost if self.status == OK else None

    @property
    def cost(self):
        """Cost of goods of the SKU in the month (positive); 0 when uncovered."""
        return self.row.cost * self.cost_units if self.status == OK else 0.0

    @property
    def counts(self):
        """Does this SKU count for coverage (something sold or returned)?"""
        return bool(self.item.units or self.item.returned or self.item.sales)


def month_end(month):
    y, m = int(str(month)[:4]), int(str(month)[5:7])
    return datetime.date(y, m, calendar.monthrange(y, m)[1])


def lookup(table, platform, country, sku, until):
    """(status, row, rows): the row that applies to this SKU on date `until`.
    status OK (row), NO_COST, or CONFLICT (two rows with the same key and date)."""
    key = sku_key(sku)
    p, c = platform.upper(), country.upper()
    for lp, lc in ((p, c), (p, ""), ("", c), ("", "")):
        cand = [r for r in table.rows if r.key == key and r.platform == lp and r.country == lc
                and (r.effective is None or r.effective <= until)]
        if not cand:
            continue
        best = max(r.effective or datetime.date.min for r in cand)
        top = [r for r in cand if (r.effective or datetime.date.min) == best]
        if len(top) > 1:
            return CONFLICT, None, top
        return OK, top[0], top
    return NO_COST, None, []


class Result(object):
    """The cost table applied to one store's month."""

    def __init__(self, table, platform, country, currency, month, items):
        self.table = table
        self.platform, self.country = platform.upper(), country.upper()
        self.currency = currency.upper()
        self.month = month
        self.until = month_end(month)
        self.lines = []
        self._by = {}
        for it in items:
            if table.usable:
                st, row, rows = lookup(table, platform, country, it.sku, self.until)
                if st == OK and row.currency != self.currency:
                    st = CURRENCY
            else:
                st, row, rows = NO_COST, None, []
            ln = Line(it, st, row, rows)
            self.lines.append(ln)
            self._by[sku_key(it.sku)] = ln

    def line(self, sku):
        return self._by.get(sku_key(sku))

    @property
    def counted(self):
        return [ln for ln in self.lines if ln.counts]

    @property
    def covered(self):
        return [ln for ln in self.counted if ln.status == OK]

    @property
    def uncovered(self):
        """SKUs without a usable cost, most sales first."""
        return sorted((ln for ln in self.counted if ln.status != OK),
                      key=lambda ln: (-(ln.item.sales or 0.0), -(ln.item.units or 0), str(ln.item.sku)))

    @property
    def cogs(self):
        return sum(ln.cost for ln in self.lines if ln.status == OK)

    def share(self, what):
        """Covered share of "sales", "units" or "skus" (1.0 when there is nothing)."""
        if what == "skus":
            tot, cov = len(self.counted), len(self.covered)
        elif what == "units":
            tot = sum(ln.item.units or 0 for ln in self.counted)
            cov = sum(ln.item.units or 0 for ln in self.covered)
        else:
            tot = sum(ln.item.sales or 0.0 for ln in self.counted)
            cov = sum(ln.item.sales or 0.0 for ln in self.covered)
        return (cov / float(tot)) if tot else 1.0

    def coverage_text(self):
        c = self.counted
        return ("SKU %d/%d、件数 %s/%s、销售额 %s 已有成本" % (
            len(self.covered), len(c), "{:,}".format(sum(ln.item.units or 0 for ln in self.covered)),
            "{:,}".format(sum(ln.item.units or 0 for ln in c)), _pct(self.share("sales"))))


def sales_share(results):
    """Covered share of sales over several stores' Results (1.0 when nothing sold)."""
    lines = [ln for res in results if res is not None for ln in res.counted]
    tot = sum(ln.item.sales or 0.0 for ln in lines)
    cov = sum(ln.item.sales or 0.0 for ln in lines if ln.status == OK)
    return (cov / float(tot)) if tot else 1.0


def point_line(results, profit):
    """① 本期要点: (label, text) — cost of goods, coverage and the profit after it."""
    cogs = sum(res.cogs for res in results if res is not None)
    return ("商品成本与利润", "商品采购成本 %s（按成本表，%s 的销售额已有成本），扣除后税前净利润 %s。" % (
        _fmt(cogs), _pct(sales_share(results)), _fmt(profit)))


def apply(table, platform, country, currency, month, items):
    """The table applied to one store's ⑦ SKUs; None when no table is configured."""
    if table is None:
        return None
    return Result(table, platform, country, currency, month, items)


def _pct(x):
    return "{:.1f}%".format(100.0 * x)


def _fmt(x):
    try:
        v = float(x)
    except (TypeError, ValueError):
        return "-"
    return "{:,.0f}".format(v) if v == int(v) else "{:,.2f}".format(v)


# ------------------------------------------------------------------ checks

def checks(table, results, platform, country, file_store=None):
    """Check records (report_kit.check) for one workbook.

    results     [(store name, Result)] — one per store of the workbook
    file_store  the "store" shown on the file-level checks (default: the only
                store, or "全部店铺")
    Nothing when no table is configured.
    """
    if table is None:
        return []
    who = file_store or (results[0][0] if len(results) == 1 else "全部店铺")
    C = []
    if not table.usable:
        C.append(K.check("sku_costs_file", who, "商品成本表可读", False,
                         "%s。config.json 的 reports.sku_costs 指向 %s：放好文件或改正路径；"
                         "在那之前②的商品采购成本为 0" % (table.error, table.path)))
        return C
    C.append(K.check("sku_costs_file", who, "商品成本表可读", True,
                     "%s：%d 行可用%s" % (table.name, len(table.rows),
                                       "（修改于 %s）" % table.modified.strftime("%Y-%m-%d %H:%M")
                                       if table.modified else "")))
    probs = [x for x in table.problems if _scoped(platform, country, x.platform, x.country)]
    C.append(K.check("sku_costs_rows", who, "商品成本表各行有效", not probs,
                     "本平台、本国家适用的行都有效" if not probs else
                     "以下行没有采用：%s" % "；".join(
                         "第 %d 行（%s）%s" % (x.line, x.sku or "?", x.text) for x in probs[:15])
                     + ("；另有 %d 行" % (len(probs) - 15) if len(probs) > 15 else "")))
    dups = [d for d in table.duplicates if _scoped(platform, country, d[0], d[1])]
    C.append(K.check("sku_costs_dup", who, "商品成本表没有重复的 SKU", not dups,
                     "没有重复" if not dups else
                     "同一 SKU、同一生效日期写了不止一行（这些 SKU 本月不计成本，直到只留一行）：%s" % "；".join(
                         "%s%s（%s 第 %s 行）" % (sku, " 生效 %s" % eff if eff else "",
                                              "/".join(x for x in (p, c) if x) or "通用",
                                              "、".join(str(n) for n in lines))
                         for p, c, sku, eff, lines in dups[:15])))
    for store, res in results:
        if res is None:
            continue
        wrong = [ln for ln in res.counted if ln.status == CURRENCY]
        C.append(K.check("sku_costs_currency", store, "商品成本的币种 ＝ 报表币种（%s）" % res.currency,
                         not wrong,
                         "全部相符" if not wrong else
                         "%d 个 SKU 的成本币种不是 %s，未采用（报表不做汇率换算，请按 %s 录入）：%s" % (
                             len(wrong), res.currency, res.currency, "、".join(
                                 "%s（%s，第 %d 行）" % (ln.item.sku, ln.row.currency, ln.row.line)
                                 for ln in wrong[:15]))))
        un = res.uncovered
        C.append(K.check("sku_costs_coverage", store, "商品成本覆盖全部售出 SKU", not un,
                         ("%s；商品采购成本 %s" % (res.coverage_text(), _fmt(res.cogs))) + (
                             "" if not un else
                             "。没有成本的 SKU（按销售额）：%s%s。在成本表补录后重新生成报表" % (
                                 "、".join("%s（%d 件，销售额 %s）" % (ln.item.sku, ln.item.units or 0,
                                                                  _fmt(ln.item.sales or 0))
                                          for ln in un[:20]),
                                 "；另有 %d 个" % (len(un) - 20) if len(un) > 20 else "")),
                         severity="warn"))
    return C


# ------------------------------------------------------------------ words for the sheets

def source_line(table):
    """① sources: (label, text)."""
    if not table.usable:
        return ("商品成本表", "⚠ %s（config.json → reports.sku_costs）。本期没有采用任何成本。" % table.error)
    return ("商品成本表", "%s（%s；修改于 %s），%d 行。成本按录入的币种使用，不做汇率换算。" % (
        table.name, table.path, table.modified.strftime("%Y-%m-%d %H:%M") if table.modified else "?",
        len(table.rows)))


def policy_line(units_basis):
    """① accounting policies: (label, text). units_basis: the platform's words for
    which units count (e.g. "本月结算件数 − 本月退货件数")."""
    return ("商品成本", "商品采购成本 ＝ 成本表的单位成本 × 成本件数。单位成本是财务录入的到岸成本（采购价＋运费＋关税），"
            "取本月最后一天有效的那一行。成本件数：%s，与收入同口径、同月份。成本表里没有的 SKU 不计成本，"
            "②注明覆盖率。⑨的库存只有数量，暂不按成本计价（以后的一步）。" % units_basis)


def pl_note(res):
    """② note of the 商品采购成本 row when a table is configured."""
    if not res.table.usable:
        return "⚠ 成本表读不出来（见⑩），本行为 0：可直接在此填总额"
    c = res.counted
    return ("来自成本表 %s（SKU %d/%d、件数 %s/%s 已有成本）。其中 %s 销售额已有成本；其余请补录或在此直接改总额" % (
        res.table.name, len(res.covered), len(c),
        "{:,}".format(sum(ln.item.units or 0 for ln in res.covered)),
        "{:,}".format(sum(ln.item.units or 0 for ln in c)), _pct(res.share("sales"))))


def status_text(line):
    """⑦ 成本来源 column for one SKU."""
    if line is None or line.status == NO_COST:
        return "未录入成本"
    if line.status == CURRENCY:
        return "未录入成本（成本表币种 %s ≠ 报表币种，第 %d 行未采用）" % (line.row.currency, line.row.line)
    if line.status == CONFLICT:
        return "未录入成本（成本表第 %s 行重复）" % "、".join(str(r.line) for r in line.rows)
    r = line.row
    scope = "/".join(x for x in (r.platform, r.country) if x) or "通用"
    return "成本表第 %d 行（%s%s）" % (r.line, scope, "，%s 起" % r.effective if r.effective else "")
