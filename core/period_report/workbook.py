# -*- coding: utf-8 -*-
"""The period workbook: the report family's ten sheets over several months.

Built only from a model.PeriodData (final monthly results); every figure is a
value, not a formula, so a reader (or a test) sees the same number Excel shows.

  ①  what the report is; the version and grade of every store-month; what is
     missing (red banner); what changed since the previous file
  ②  P&L: one block per store and, with several stores, the market total first;
     one column per month + 期间合计; 备查项 under each block
  ③  the period's own cross-checks (period total = sum of months, market = sum
     of stores, only final versions used); the monthly bridges stay monthly
  ④  fees by type and month    ⑤ returns by class and month
  ⑥  not applicable (the monthly reports hold the per-order detail)
  ⑦  SKU totals over the period, by store and SKU
  ⑧  monthly trend of the market    ⑨ the last month's inventory snapshot
  ⑩  missing store-months, checks carried from the monthly versions, P&L rows
     missing in a month, versions changed since the previous file

A hidden sheet (VERSIONS_SHEET) records the version used for every
store-month, so the next generation of the same file can say what changed.
"""
import math

from openpyxl import Workbook, load_workbook

from core import markets
from core import report_family as RF
from core import report_kit as K
from core import results as R
from core.period_report import model as MD
from core.report_kit import BOX, F, fill, wrap

SHEETS = K.SHEET_NAMES
VERSIONS_SHEET = "_版本记录"
VERSIONS_HEAD = ("店铺", "月份", "基准", "版本", "等级", "生成于")
MARKET_TOTAL = "全部店铺合计"
NO_FINAL = "无完整版本"


# ------------------------------------------------------------------ helpers

def title_of(D):
    b = MD.basis_cn(D.basis)
    return "期间财务报表 · %s · %s%s" % (markets.label(D.market), MD.period_cn(D.period),
                                     " · %s" % b if b else "")


def grade_text(doc):
    g = doc["grade"]
    return R.GRADE_CN.get(g["level"], g["level"])


def money_cell(ws, r, c, v, money, bold=False, color=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.number_format = money
    cell.border = BOX
    if bold:
        cell.font = F(10, True)
    if color:
        cell.fill = fill(color)
    return cell


def text_cell(ws, r, c, v, font=None, color=None, align=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.border = BOX
    cell.font = font or F(10)
    if color:
        cell.fill = fill(color)
    cell.alignment = align or wrap()
    return cell


def month_head(D, month, store=None):
    """Header text of a month column: '2026-08', with 无完整版本 when the store
    (or, for the market total, some store) has no final version."""
    if store is not None:
        return month if D.final(store, month) else "%s\n%s" % (month, NO_FINAL)
    n, total = len(D.complete_stores(month)), len(D.stores)
    return month if n == total else "%s\n%d/%d 家完整" % (month, n, total)


def build(D, out_path, previous=None):
    """Write the workbook. previous: versions_of() of the file this one
    replaces (None = first generation, "unreadable" = could not be read)."""
    wb = Workbook()
    wb.remove(wb.active)
    money = K.money_format(D.decimals)
    changes = compare(previous, D)
    sheet_readme(wb, D, changes)
    sheet_pl(wb, D, money)
    sheet_checks(wb, D)
    sheet_fees(wb, D, money)
    sheet_returns(wb, D, money)
    RF.not_applicable_sheet(wb, SHEETS[5], "退货质检明细",
                            "不适用：期间报表不重复逐单明细。每笔退货的质检结果在各月的月度报表⑥里。")
    sheet_sku(wb, D, money)
    sheet_trend(wb, D, money)
    sheet_inventory(wb, D)
    sheet_gaps(wb, D, changes)
    sheet_versions(wb, D)
    wb.save(out_path)
    return changes


# ------------------------------------------------------------------ versions

def sheet_versions(wb, D):
    ws = wb.create_sheet(VERSIONS_SHEET)
    ws.append(list(VERSIONS_HEAD))
    for s in D.stores:
        for m in D.months:
            doc = D.final(s, m)
            ws.append([s, m, D.basis, doc["version"] if doc else "", grade_text(doc) if doc else NO_FINAL,
                       doc["built_at"] if doc else ""])
    ws.sheet_state = "hidden"


def versions_of(path):
    """{(store, month): version or ''} recorded in a period workbook; None when
    the file does not exist; 'unreadable' when it cannot be read."""
    import os
    if not os.path.exists(path):
        return None
    try:
        wb = load_workbook(path, read_only=True)
        try:
            ws = wb[VERSIONS_SHEET]
            out = {}
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and row[0]:
                    out[(str(row[0]), str(row[1]))] = str(row[3] or "")
            return out
        finally:
            wb.close()
    except Exception:                                  # noqa: BLE001 - any bad file
        return "unreadable"


def compare(previous, D):
    """What changed since the previous file of the same name:
    {"first": bool, "unreadable": bool, "rows": [(store, month, old, new, what)]}."""
    out = {"first": previous is None, "unreadable": previous == "unreadable", "rows": []}
    if not isinstance(previous, dict):
        return out
    now = D.version_map()
    for (s, m) in sorted(now, key=lambda k: (D.stores.index(k[0]), k[1])):
        new = now[(s, m)]
        if (s, m) not in previous:
            if new:
                out["rows"].append((s, m, "", new, "新加进来的店铺月份"))
            continue
        old = previous[(s, m)]
        if old == new:
            continue
        what = ("有了完整版本" if not old else "完整版本没有了（被清理或文件读不出来）" if not new
                else "换成了更新的完整版本")
        out["rows"].append((s, m, old, new, what))
    for (s, m), old in sorted(previous.items()):
        if (s, m) not in now and old:
            out["rows"].append((s, m, old, "", "这一份不再包含这个店铺月份"))
    return out


# ------------------------------------------------------------------ ①

def sheet_readme(wb, D, changes):
    ws = wb.create_sheet(SHEETS[0])
    RF.set_widths(ws, (4, 30, 118))
    ws["B1"] = title_of(D)
    ws["B1"].font = F(16, True, K.C_NAVY)
    ws["B2"] = "生成于 %s　%d 家店铺　%d 个月　币种 %s" % (
        D.built.strftime("%Y-%m-%d %H:%M"), len(D.stores), len(D.months), D.currency)
    ws["B2"].font = F(9, False, K.C_GREY)
    missing = D.missing()
    if missing:
        ws.merge_cells("B3:C3")
        c = ws["B3"]
        c.value = ("⚠ 有 %d 个店铺月份没有完整版本，没有计入任何合计：%s。原因见下面第二节和⑩。"
                   % (len(missing), "、".join("%s %s" % (x.store, x.month) for x in missing[:12])
                      + ("…" if len(missing) > 12 else "")))
        c.font = F(11, True, "FFFFFF")
        c.fill = fill(K.C_RED)
        c.alignment = wrap()
        ws.row_dimensions[3].height = K.est_height(c.value, 70, 18, 22)
    b = MD.basis_cn(D.basis)
    rows = [("S", "一、这份报表是什么", ""),
            ("", "期间", "%s：%s" % (MD.KIND_CN[D.period.kind], MD.period_cn(D.period))),
            ("", "市场 / 币种", "%s ／ %s" % (markets.label(D.market), D.currency)),
            ("", "基准", (b + "。" if b else "") + (
                "会计入账以账单基准为准；订单基准作经营分析参考。" if D.platform == "ML" else "本平台只有一种基准。")),
            ("", "数据从哪来", "只用每个店铺每个月的定稿月度结果：同一个月有几个完整版本时，用最新生成的那个"
                           "（等级「完整」或「完整·有提醒」）。不重新读平台数据，所以和各月月度报表的数字一致。"),
            ("", "店铺", "、".join(D.stores) or "（无）"),
            ("", "", ""),
            ("S", "二、各店各月用的版本", "")]
    flagged = []
    for s in D.stores:
        for m in D.months:
            cell = D.cell(s, m)
            if cell.final:
                f = cell.final
                text = "%s · 版本 %s · 生成于 %s" % (grade_text(f), f["version"], f["built_at"].replace("T", " "))
                if f["grade"].get("reasons"):
                    text += " · " + "；".join(f["grade"]["reasons"])
            else:
                text = "%s：%s" % (NO_FINAL, cell.reason())
                flagged.append(len(rows))
            rows.append(("", "%s · %s" % (s, m), text))
    rows += [("", "", ""), ("S", "三、缺什么", "")]
    if missing:
        rows.append(("", "没有完整版本", "%d 个店铺月份（红字）。它们的数字不在任何合计里 —— 不是 0，是还没有。"
                                      "等那个月重出报表、有了完整版本，再生成一次这份报表。" % len(missing)))
    else:
        rows.append(("", "没有完整版本", "没有。每个店铺每个月都有完整版本。"))
    gaps = D.pl_gaps()
    if gaps:
        rows.append(("", "损益行缺月", "%d 处：某个月的结果里没有这一行，按 0 计（见⑩）。" % len(gaps)))
    rows += [("", "", ""), ("S", "四、和上一份的差别", "")]
    if changes["first"]:
        rows.append(("", "上一份", "这是第一次生成这份报表。"))
    elif changes["unreadable"]:
        rows.append(("", "上一份", "上一份同名文件读不出版本记录，没法比较。"))
    elif changes["rows"]:
        rows.append(("", "换了版本", "%d 个店铺月份和上一份用的版本不同（见⑩）。" % len(changes["rows"])))
    else:
        rows.append(("", "换了版本", "没有：每个店铺月份用的版本和上一份相同。"))
    rows += [("", "", ""), ("S", "五、各页说明", ""),
             ("", SHEETS[1], "损益：%s每家店一块；每块每月一列，最后一列期间合计。没有完整版本的月份留空（浅红）。"
                             % ("最上面一块是全部店铺合计，然后" if len(D.stores) > 1 else "")),
             ("", SHEETS[2], "期间报表自己的核对：期间合计 = 各月之和、全部店铺 = 各店之和、只用了完整版本。"
                             "每个月的平账在各月月度报表③。"),
             ("", SHEETS[3], "各费用类型每月的金额与期间合计（全部店铺，然后逐店）。"),
             ("", SHEETS[4], "各退货类别每月的件数和金额。"),
             ("", SHEETS[5], "不适用：逐单质检在各月月度报表⑥。"),
             ("", SHEETS[6], "期间内每家店每个 SKU 的件数、净销售额、费用、贡献和成本之和。"),
             ("", SHEETS[7], "全部店铺按月：完整的店铺数和主要损益行。"),
             ("", SHEETS[8], "每家店期间内最后一个有库存数据的月份的库存快照。库存是时点数，不跨月相加。"),
             ("", SHEETS[9], "没有完整版本的店铺月份、月度版本带过来的提醒、损益行缺月、和上一份相比换了的版本。")]
    first = 5 if missing else 4
    readme_end = RF.readme_rows(ws, first, rows)
    for i in flagged:
        c = ws.cell(row=first + i, column=3)
        c.font = F(10, True, K.C_RED)
    return readme_end


# ------------------------------------------------------------------ ②

def sheet_pl(wb, D, money):
    ws = wb.create_sheet(SHEETS[1])
    nm = len(D.months)
    total_col = 3 + nm
    note_col = total_col + 1
    K.title_rows(ws, "合并损益表（期间） · %s" % MD.period_cn(D.period),
                 "每月一列＋期间合计；金额取各月定稿月度结果的②（同一行、同一符号）。没有完整版本的月份留空，不计入合计。",
                 "币种 %s　%s" % (D.currency, MD.basis_cn(D.basis)))
    RF.set_widths(ws, [34, 26] + [15] * nm + [17, 52])
    ws.freeze_panes = "C5"
    blocks = ([(MARKET_TOTAL + "（%d 家）" % len(D.stores), None)] if len(D.stores) > 1 else []) + \
             [(s, s) for s in D.stores]
    gaps = {}
    for s, m, key, _ in D.pl_gaps():
        gaps.setdefault((s, key), []).append(m)
    r = 5
    for name, store in blocks:
        r = RF.section(ws, r, name, note_col)
        heads = ["项目", "平台科目（原文）"] + [month_head(D, m, store) for m in D.months] + ["期间合计", "说明"]
        K.header_cells(ws, r, heads)
        ws.row_dimensions[r].height = 30
        r += 1
        for row in D.pl_spec():
            r = _pl_row(ws, r, D, row, store, money, gaps, note_col)
        memo = D.memo_spec()
        if memo:
            ws.cell(row=r, column=1, value="备查项（不进损益）").font = F(10, True, K.C_GREY)
            r += 1
            for key, label in memo.items():
                ws.cell(row=r, column=1, value=label).font = F(10)
                vals = [_memo_value(D, store, m, key) for m in D.months]
                _value_cells(ws, r, D, store, vals, money)
                tot = (D.field_store_total(store, "memo", "key", key, "value") if store
                       else D.sum(v for v in vals if v is not None))
                money_cell(ws, r, total_col, tot, money, bold=True, color=K.C_TOTAL)
                r += 1
        r += 1
    return ws


def _memo_value(D, store, month, key):
    if store:
        return D.field_value(store, month, "memo", "key", key, "value")
    if not D.complete_stores(month):
        return None
    return D.field_month_total(month, "memo", "key", key, "value")


def _value_cells(ws, r, D, store, vals, money, bold=False, color=None):
    for j, v in enumerate(vals):
        if v is None:
            text_cell(ws, r, 3 + j, None, color=K.C_BAD)
        else:
            money_cell(ws, r, 3 + j, round(v, D.decimals), money, bold=bold, color=color)


def _pl_row(ws, r, D, row, store, money, gaps, note_col):
    kind = row["kind"]
    strong = kind in ("sum", "f")
    color = K.C_SUB if kind == "sum" else K.C_CALC if kind == "f" else None
    ws.cell(row=r, column=1, value=row["label"]).font = F(10, strong)
    if color:
        ws.cell(row=r, column=1).fill = fill(color)
    ws.cell(row=r, column=2, value=row["original"]).font = F(9, False, K.C_GREY)
    key = row["key"]
    if store:
        vals = [D.pl_value(store, m, key) for m in D.months]
        tot = D.store_total(store, key)
        miss = gaps.get((store, key), [])
    else:
        vals = [D.month_total(m, key) if D.complete_stores(m) else None for m in D.months]
        tot = D.total(key)
        miss = sorted(set(m for (s, k), ms in gaps.items() if k == key for m in ms))
    _value_cells(ws, r, D, store, vals, money, bold=strong, color=color)
    money_cell(ws, r, 3 + len(D.months), tot, money, bold=True, color=K.C_TOTAL)
    notes = []
    if kind == "in":
        notes.append("月度报表的录入行（商品成本表填的值；没填为 0）")
    if miss:
        notes.append("%s 的结果里没有这一行，按 0 计" % "、".join(miss))
    c = ws.cell(row=r, column=note_col, value="；".join(notes))
    c.font = F(9, False, K.C_RED if miss else K.C_GREY)
    c.alignment = wrap()
    return r + 1


# ------------------------------------------------------------------ ③

def checks(D):
    """The period's own checks: [(scope, name, ok, detail)]. Every total is
    compared with an independent binary sum (math.fsum) of the figures it adds:
    it must be within half a cent (half a peso for CLP) — i.e. the sum to the cent."""
    tol = 0.5 * 10 ** -D.decimals + 1e-9
    out = []
    keys = [row["key"] for row in D.pl_spec()]

    def diff(total, values):
        return abs(total - math.fsum(v for v in values if v is not None))

    for s in D.stores:
        worst = max([diff(D.store_total(s, k), (D.pl_value(s, m, k) for m in D.months)) for k in keys] or [0])
        out.append((s, "② 期间合计 = 各月之和", worst <= tol, "%d 行，最大差额 %.6f" % (len(keys), worst)))
    if len(D.stores) > 1:
        worst = max([diff(D.month_total(m, k), (D.pl_value(s, m, k) for s in D.stores))
                     for m in D.months for k in keys]
                    + [diff(D.total(k), (D.pl_value(s, m, k) for s in D.stores for m in D.months))
                       for k in keys] or [0])
        out.append((MARKET_TOTAL, "全部店铺 = 各店之和（每月和期间）", worst <= tol, "最大差额 %.6f" % worst))
    used = D.finals()
    bad = [x for x in used if x[2]["grade"]["level"] not in R.COUNTABLE]
    out.append((MARKET_TOTAL, "只用了完整版本", not bad,
                "用了 %d 个店铺月份的定稿版本；%d 个没有完整版本，没有计入" % (len(used), len(D.missing()))))
    return out


def sheet_checks(wb, D):
    ws = wb.create_sheet(SHEETS[2])
    K.title_rows(ws, "平账勾稽表（期间）",
                 "期间报表不重做平账：每个月的订单 ↔ 回款、费用发票 ↔ 费用等桥在各月月度报表③里，期间报表只用已经定稿的结果。",
                 "这里核对的是期间报表自己：合计用另一种算法（二进制精确相加）再算一遍，差额应在半分以内（即到分相等）。")
    cols = [RF.Col("范围", 26), RF.Col("核对", 40), RF.Col("结果", 10), RF.Col("说明", 90, align=wrap())]
    RF.header_row(ws, 5, [c.header for c in cols], [c.width for c in cols])
    fam = RF.Family([{"store": ""}])
    r = 6
    for scope, name, ok, detail in checks(D):
        fam.write_row(ws, r, cols, (scope, name, "通过" if ok else "未通过", detail),
                      fill_color=K.C_OK if ok else K.C_BAD)
        r += 1
    return ws


# ------------------------------------------------------------------ ④ ⑤ (month tables)

def _month_table(ws, D, head_cols, rows, money, header_at=4):
    """rows: [(fixed values, [month values or None], total, fmt, color)]."""
    nm = len(D.months)
    heads = list(head_cols) + list(D.months) + ["期间合计"]
    r = header_at
    K.header_cells(ws, r, [h[0] for h in heads[:len(head_cols)]] + heads[len(head_cols):])
    ws.row_dimensions[r].height = 30
    RF.set_widths(ws, [h[1] for h in head_cols] + [15] * nm + [17])
    r += 1
    nf = len(head_cols)
    for fixed, vals, total, fmt, color in rows:
        for j, v in enumerate(fixed, 1):
            text_cell(ws, r, j, v, color=color)
        for j, v in enumerate(vals):
            if v is None:
                text_cell(ws, r, nf + 1 + j, None, color=K.C_BAD)
            else:
                money_cell(ws, r, nf + 1 + j, v, fmt, color=color)
        money_cell(ws, r, nf + nm + 1, total, fmt, bold=True, color=K.C_TOTAL)
        r += 1
    return r


def _scopes(D):
    """(label, store or None): the market total first when there are several stores."""
    return ([(MARKET_TOTAL, None)] if len(D.stores) > 1 else []) + [(s, s) for s in D.stores]


def _field_vals(D, store, field, key_field, key, value_field, money=True):
    if store:
        vals = [D.field_value(store, m, field, key_field, key, value_field) for m in D.months]
        tot = D.field_store_total(store, field, key_field, key, value_field, money)
    else:
        vals = [D.field_month_total(m, field, key_field, key, value_field, money)
                if D.complete_stores(m) else None for m in D.months]
        tot = D.sum(v for v in vals if v is not None) if money else int(sum(v for v in vals if v is not None))
    if not money:
        vals = [None if v is None else int(round(v)) for v in vals]
    return vals, tot


def sheet_fees(wb, D, money):
    ws = wb.create_sheet(SHEETS[3])
    K.title_rows(ws, "平台费用明细（期间）",
                 "各月定稿结果④的费用，按费用类型（原文）每月一列；金额与②同号（费用为负）。没有完整版本的月份留空。")
    spec = D.fee_spec()
    rows = []
    for label, store in _scopes(D):
        for name, (cn, pl_key) in spec.items():
            vals, tot = _field_vals(D, store, "fees", "name", name, "amount")
            if not any(vals) and not tot:
                continue
            rows.append(((label, cn, name, pl_key), vals, tot, money, None))
    r = _month_table(ws, D, [("店铺", 22), ("费用类型", 30), ("原文", 34), ("计入②行", 16)], rows, money)
    if not rows:
        ws.cell(row=r, column=1, value="期间内的月度结果里没有费用明细。").font = F(10, False, K.C_GREY)
    return ws


def sheet_returns(wb, D, money):
    ws = wb.create_sheet(SHEETS[4])
    K.title_rows(ws, "退货退款分析（期间）",
                 "各月定稿结果⑤的退货类别，每月的件数和金额（金额与②同号）。没有完整版本的月份留空。")
    spec = D.return_spec()
    rows = []
    for label, store in _scopes(D):
        for key, cn in spec.items():
            n, ntot = _field_vals(D, store, "returns", "key", key, "count", money=False)
            a, atot = _field_vals(D, store, "returns", "key", key, "amount")
            rows.append(((label, cn, "件数"), n, ntot, K.INT, None))
            rows.append(((label, cn, "金额"), a, atot, money, None))
    r = _month_table(ws, D, [("店铺", 22), ("类别", 30), ("指标", 8)], rows, money)
    if not rows:
        ws.cell(row=r, column=1, value="期间内的月度结果里没有退货分类。").font = F(10, False, K.C_GREY)
    return ws


# ------------------------------------------------------------------ ⑦

SKU_COLS = [RF.Col("店铺", 22), RF.Col("SKU", 22), RF.Col("名称", 44, align=wrap()),
            RF.Col("销售件数", 11, K.INT), RF.Col("退回件数", 11, K.INT),
            RF.Col("净销售额", 15, RF.MONEY), RF.Col("平台费用", 15, RF.MONEY),
            RF.Col("经营贡献", 15, RF.MONEY), RF.Col("商品成本", 15, RF.MONEY),
            RF.Col("成本说明", 30, align=wrap()), RF.Col("有数据的月份", 12, K.INT)]


def sheet_sku(wb, D, money):
    fam = RF.Family([{"store": s} for s in D.stores] or [{"store": ""}], money=money)
    ws = fam.table_sheet(wb, SHEETS[6], "SKU 盈利分析（期间）",
                         "期间内各月定稿结果⑦逐 SKU 相加（同一 SKU 在不同店铺各占一行）；金额与②同号。"
                         "商品成本只加有成本的月份。", SKU_COLS)
    r = 5
    for s in D.stores:
        rows = D.sku_rows(s)
        first = r
        for x in rows:
            if x["cost"] is None:
                note = "未录入成本"
            elif x["cost_missing"]:
                note = "%s 没有成本，未计入" % "、".join(sorted(x["cost_missing"]))
            else:
                note = ""
            cells = fam.write_row(ws, r, SKU_COLS, (
                s, x["sku"], x["title"], int(round(x["units"])), int(round(x["units_returned"])),
                x["net_sales"], x["fees"], x["contribution"], x["cost"], note, len(x["months"])))
            if note:
                cells[9].fill = fill(K.C_INPUT if x["cost"] is None else K.C_WARN)
            r += 1
        if rows:
            fam.write_row(ws, r, SKU_COLS, (
                "%s 合计" % s, "", "", int(round(sum(x["units"] for x in rows))),
                int(round(sum(x["units_returned"] for x in rows))),
                D.sum(x["net_sales"] for x in rows), D.sum(x["fees"] for x in rows),
                D.sum(x["contribution"] for x in rows),
                D.sum(x["cost"] for x in rows if x["cost"] is not None)
                if any(x["cost"] is not None for x in rows) else None, "", None),
                fill_color=K.C_TOTAL)
            for c in range(1, len(SKU_COLS) + 1):
                ws.cell(row=r, column=c).font = F(10, True)
            r += 2
        elif first == r:
            ws.cell(row=r, column=1, value=s).font = F(10)
            ws.cell(row=r, column=2, value="期间内的月度结果里没有 SKU 明细").font = F(10, False, K.C_GREY)
            r += 2
    return ws


# ------------------------------------------------------------------ ⑧

def trend_rows(D):
    """The P&L rows shown on ⑧: the first row, then every subtotal and computed row."""
    spec = D.pl_spec()
    return [row for i, row in enumerate(spec) if i == 0 or row["kind"] in ("sum", "f")]


def sheet_trend(wb, D, money):
    ws = wb.create_sheet(SHEETS[7])
    lines = trend_rows(D)
    K.title_rows(ws, "月度趋势（期间）",
                 "全部店铺按月（只加有完整版本的店铺）；完整店铺数少于店铺数的月份标浅红，和别的月份不可比。")
    heads = ["月份", "完整店铺", "缺的店铺"] + [row["label"] for row in lines] + ["销售件数"]
    K.header_cells(ws, 4, heads)
    ws.row_dimensions[4].height = 42
    RF.set_widths(ws, [10, 10, 30] + [16] * len(lines) + [12])
    r = 5
    for m in D.months:
        done = D.complete_stores(m)
        lacking = [s for s in D.stores if s not in done]
        color = K.C_BAD if lacking else None
        text_cell(ws, r, 1, m, color=color)
        text_cell(ws, r, 2, "%d/%d" % (len(done), len(D.stores)), color=color)
        text_cell(ws, r, 3, "、".join(lacking), color=color)
        for j, row in enumerate(lines):
            money_cell(ws, r, 4 + j, D.month_total(m, row["key"]) if done else None, money, color=color)
        units = sum(float(x.get("units") or 0) for s in done for x in (D.final(s, m).get("sku") or []))
        money_cell(ws, r, 4 + len(lines), int(round(units)), K.INT, color=color)
        r += 1
    text_cell(ws, r, 1, "期间合计", font=F(10, True), color=K.C_TOTAL)
    text_cell(ws, r, 2, "", color=K.C_TOTAL)
    text_cell(ws, r, 3, "", color=K.C_TOTAL)
    for j, row in enumerate(lines):
        money_cell(ws, r, 4 + j, D.total(row["key"]), money, bold=True, color=K.C_TOTAL)
    units = sum(float(x.get("units") or 0) for s, m, doc in D.finals() for x in (doc.get("sku") or []))
    money_cell(ws, r, 4 + len(lines), int(round(units)), K.INT, bold=True, color=K.C_TOTAL)
    return ws


# ------------------------------------------------------------------ ⑨

INV_COLS = [RF.Col("店铺", 22), RF.Col("截至", 14), RF.Col("来自月份", 10), RF.Col("SKU", 22),
            RF.Col("名称", 44, align=wrap()), RF.Col("库存", 11, K.INT), RF.Col("近 30 天销量", 12, K.INT)]


def sheet_inventory(wb, D):
    fam = RF.Family([{"store": s} for s in D.stores] or [{"store": ""}])
    ws = fam.table_sheet(wb, SHEETS[8], "库存与动销分析（期间）",
                         "每家店取期间内最后一个有完整版本、带库存数据的月份的快照。库存是时点数，不跨月相加。", INV_COLS)
    r = 5
    last = D.months[-1]
    for s in D.stores:
        m, snap = D.inventory(s)
        if not snap:
            fam.write_row(ws, r, INV_COLS, (s, "", "", "", "期间内没有库存快照（平台没有仓库数据，或月度结果里没有）",
                                            None, None), fill_color=K.C_CALC)
            r += 1
            continue
        color = K.C_WARN if m != last else None
        for x in snap.get("rows") or []:
            fam.write_row(ws, r, INV_COLS, (s, snap.get("as_of") or "", m, x.get("sku"), x.get("title"),
                                            x.get("stock"), x.get("sold_30d")), fill_color=color)
            r += 1
    if any(D.inventory(s)[0] not in (None, last) for s in D.stores):
        r += 1
        ws.cell(row=r, column=1, value="黄底：快照不是期间最后一个月（%s）的 —— 那个月没有完整版本或没有库存数据。"
                % last).font = F(9, False, K.C_AMBER)
    return ws


# ------------------------------------------------------------------ ⑩

GAP_COLS = [RF.Col("店铺", 22), RF.Col("月份", 10), RF.Col("项目", 40, align=wrap()),
            RF.Col("级别", 12), RF.Col("说明", 100, align=wrap())]


def sheet_gaps(wb, D, changes):
    fam = RF.Family([{"store": s} for s in D.stores] or [{"store": ""}])
    ws = fam.table_sheet(wb, SHEETS[9], "数据缺口与待办（期间）",
                         "空表示没有待办。没有完整版本的店铺月份不在任何合计里；补上之后重新生成这份报表。", GAP_COLS)
    r = 5
    r = RF.section(ws, r, "一、没有完整版本的店铺月份（没有计入合计）", 5)
    miss = D.missing()
    for c in miss:
        lvl = R.GRADE_CN.get(c.newest["grade"]["level"], "") if c.newest else "没有版本"
        fam.write_row(ws, r, GAP_COLS, (c.store, c.month, NO_FINAL, lvl, c.reason()), fill_color=K.C_BAD)
        r += 1
    if not miss:
        ws.cell(row=r, column=3, value="没有。每个店铺每个月都有完整版本。")
        r += 1
    r += 1
    r = RF.section(ws, r, "二、月度版本里没通过的校验（数字已按月度报表定稿；看一眼是否影响期间判断）", 5)
    carried = D.carried_checks()
    for s, m, c in carried:
        sev = c.get("severity") or "warn"
        fam.write_row(ws, r, GAP_COLS, (s, m, c.get("name") or c.get("key"), sev, c.get("detail") or ""),
                      fill_color=K.C_BAD if sev == "error" else K.C_WARN)
        r += 1
    if not carried:
        ws.cell(row=r, column=3, value="没有。")
        r += 1
    r += 1
    r = RF.section(ws, r, "三、某个月的结果里没有的损益行（按 0 计）", 5)
    gaps = D.pl_gaps()
    for s, m, key, label in gaps:
        fam.write_row(ws, r, GAP_COLS, (s, m, label, "按 0 计",
                                        "这个月的定稿结果里没有「%s」（%s）这一行：多半是这一行是后来的月度报表"
                                        "才加的。期间合计里这个月按 0 计。" % (label, key)), fill_color=K.C_WARN)
        r += 1
    if not gaps:
        ws.cell(row=r, column=3, value="没有。")
        r += 1
    r += 1
    r = RF.section(ws, r, "四、和上一份同名报表相比换了的版本", 5)
    if changes["first"]:
        ws.cell(row=r, column=3, value="这是第一次生成，没有上一份可比。")
        r += 1
    elif changes["unreadable"]:
        ws.cell(row=r, column=3, value="上一份同名文件读不出版本记录，没法比较。")
        r += 1
    else:
        for s, m, old, new, what in changes["rows"]:
            fam.write_row(ws, r, GAP_COLS, (s, m, what, "",
                                            "上一份：%s；这一份：%s" % (old or NO_FINAL, new or NO_FINAL)),
                          fill_color=K.C_WARN)
            r += 1
        if not changes["rows"]:
            ws.cell(row=r, column=3, value="没有：每个店铺月份用的版本和上一份相同。")
            r += 1
    return ws
