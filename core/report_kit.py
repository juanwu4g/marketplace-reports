# -*- coding: utf-8 -*-
"""
The report family's frame, shared by every platform's workbook
(docs/REPORT_FAMILY.md): look, sheet scaffolding, check records, names.

What lives here is what makes two platforms' workbooks recognisable as sisters;
what they *say* stays in each platform's report_build. MercadoLibre's
report_build imports these instead of defining its own — the values were moved,
not changed, and the golden tests prove it.

Only openpyxl; no pandas, no project imports: a report builder that runs as a
plain script can import it after putting the repository root on sys.path.
"""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ------------------------------------------------------------------ sheets

# Ten sheets, always, in this order. A sheet that does not apply to a platform
# says so ("不适用：…") instead of disappearing: people learn where things are.
SHEET_NAMES = [
    "① 阅读说明", "② 合并损益表", "③ 平账勾稽表", "④ 平台费用明细",
    "⑤ 退货退款分析", "⑥ 退货质检明细", "⑦ SKU盈利分析", "⑧ 月度趋势",
    "⑨ 库存与动销分析", "⑩ 数据缺口与待办",
]

# ------------------------------------------------------------------ look

FONT_NAME = "Arial"
MNY = '#,##0.00;(#,##0.00);-'
MNY0 = '#,##0;(#,##0);-'
INT = '#,##0;(#,##0);-'
PCT = '0.0%;(0.0%);-'
PCT2 = '0.00%'
PCT3 = '0.000%'

C_NAVY, C_BAND, C_SUB, C_CALC = "1F3864", "4472C4", "D9E2F3", "EDEDED"
C_TOTAL, C_INPUT, C_OK, C_BAD, C_WARN = "FCE4D6", "FFFF00", "E2EFDA", "FBE5E5", "FFF2CC"
C_RED, C_GREEN, C_AMBER, C_GREY = "C00000", "375623", "BF8F00", "595959"

_THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def money_format(decimals):
    """Number format for money in a currency with this many decimals (CLP: 0)."""
    return MNY0 if not decimals else MNY


def F(size=10, bold=False, color="000000", italic=False):
    return Font(name=FONT_NAME, size=size, bold=bold, color=color, italic=italic)


def fill(hexcolor):
    return PatternFill("solid", fgColor=hexcolor)


def wrap(vertical="center"):
    return Alignment(wrap_text=True, vertical=vertical)


def est_height(text, per_line=60, line_h=15, minimum=15):
    """按字符数粗估行高。中文宽字符，60 字/行是保守值。"""
    if not text:
        return minimum
    return max(minimum, line_h * (len(str(text)) // per_line + 1))


class Sheet(object):
    """薄封装：多店铺时列数随店铺数变化，这里统一算列位。"""

    def __init__(self, wb, name, first_store_col, nstores, has_total=True):
        self.ws = wb.create_sheet(name)
        self.c0 = first_store_col
        self.n = nstores
        self.has_total = has_total

    def scol(self, i):
        return get_column_letter(self.c0 + i)

    @property
    def tcol_idx(self):
        return self.c0 + self.n

    @property
    def tcol(self):
        return get_column_letter(self.tcol_idx)

    def title(self, text, sub="", sub2=""):
        title_rows(self.ws, text, sub, sub2)

    def header(self, row, cols, widths, height=30, color=C_NAVY):
        ws = self.ws
        header_cells(ws, row, cols, color=color)
        ws.row_dimensions[row].height = height
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def title_rows(ws, text, sub="", sub2=""):
    """A sheet's head: A1 the title (navy, bold, row 24 high), A2 / A3 grey notes
    when given (Sheet.title on any worksheet)."""
    ws["A1"] = text
    ws["A1"].font = F(14, True, C_NAVY)
    ws.row_dimensions[1].height = 24
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F(9, False, C_GREY)
    if sub2:
        ws["A3"] = sub2
        ws["A3"].font = F(9, False, C_GREY)


def header_cells(ws, row, cols, first_col=1, color=C_NAVY, wrap_text=True):
    """The cells of a table header: white bold text on `color`, centred, boxed —
    one per entry of cols, from column first_col. Row height and column widths
    are left to the caller (Sheet.header sets both). wrap_text=False leaves the
    text on one line (the alignment then carries no wrap flag at all)."""
    if wrap_text:
        align = dict(horizontal="center", vertical="center", wrap_text=True)
    else:
        align = dict(horizontal="center", vertical="center")
    for i, t in enumerate(cols, first_col):
        c = ws.cell(row=row, column=i, value=t)
        c.font = F(10, True, "FFFFFF")
        c.fill = fill(color)
        c.alignment = Alignment(**align)
        c.border = BOX


# ------------------------------------------------------------------ checks

def check(key, store, name, ok, detail, severity="error", skipped=False):
    """One check record (REPORT_FAMILY §4): error = numbers not trustworthy until
    fixed; warn = look at this; skipped = not applicable here, with the reason."""
    return {"key": key, "store": store, "name": name, "ok": bool(ok),
            "detail": detail, "severity": severity, "skipped": skipped}


# ------------------------------------------------------------------ names

def store_workbook_name(store, month):
    return "%s_%s.xlsx" % (store, month)


def combined_workbook_name(country, platform_code, month):
    """CL_FA_全店汇总_2026-09.xlsx (MercadoLibre keeps its historical CL_ML_ / MX_ML_)."""
    return "%s_%s_全店汇总_%s.xlsx" % (str(country).upper(), str(platform_code).upper(), month)


def alert_file_name(country, platform_code, month):
    """！报表异常_CL_FA_202609.txt"""
    return "！报表异常_%s_%s_%s.txt" % (str(country).upper(), str(platform_code).upper(),
                                     str(month).replace("-", ""))


# ------------------------------------------------------------------ recalculation

WIN_SOFFICE = (r"C:\Program Files\LibreOffice\program\soffice.exe",
               r"C:\Program Files (x86)\LibreOffice\program\soffice.exe")


def soffice_path():
    """LibreOffice 的可执行文件，找不到返回 None。

    顺序：环境变量 SOFFICE > Program Files 里的标准安装位置 > PATH。has_libreoffice()
    和 recalc() 都用它 —— 以前 recalc 只认 PATH 上的 soffice，装在 Program Files
    但没进 PATH 的机器上，前面说"有 LibreOffice"，这里却悄悄跳过了重算。"""
    import os
    env = (os.environ.get("SOFFICE") or "").strip().strip('"')
    if env:
        return env
    if os.name == "nt":
        for p in WIN_SOFFICE:
            if os.path.isfile(p):
                return p
    for d in os.environ.get("PATH", "").split(os.pathsep):
        for exe in ("soffice.exe", "soffice"):
            if d and os.path.isfile(os.path.join(d, exe)):
                return os.path.join(d, exe)
    return None


def _kill_tree(p):
    """结束 soffice 连同它起的 soffice.bin。永不抛异常。"""
    import os
    import signal
    import subprocess
    try:
        if os.name == "nt":
            subprocess.run(["taskkill", "/T", "/F", "/PID", str(p.pid)],
                           capture_output=True, timeout=30)
        else:
            os.killpg(p.pid, signal.SIGKILL)
    except Exception:                                             # noqa: BLE001
        pass
    try:
        p.kill()
    except Exception:                                             # noqa: BLE001
        pass
    try:
        p.wait(timeout=10)
    except Exception:                                             # noqa: BLE001
        pass


def recalc(path, timeout=180):
    """用 LibreOffice 重算并回填公式缓存值。openpyxl 只写公式不写结果，
    不重算的话 pandas / data_only=True 读到的全是 None。

    timeout 秒还没完就结束 soffice（连同子进程），当成"没重算"返回：卡死的
    soffice 以前会让出报表那一步永远挂着，锁也一直不放，后面每一轮都退出码 5。"""
    import os
    import subprocess
    import time
    out_dir = os.path.join(os.path.dirname(os.path.abspath(path)), "_recalc")
    exe = soffice_path() or "soffice"
    cmd = [exe, "--headless", "--norestore", "--convert-to",
           "xlsx:Calc MS Excel 2007 XML", "--outdir", out_dir, path]
    kw = {"start_new_session": True} if os.name != "nt" else {}
    produced = os.path.join(out_dir, os.path.basename(path))
    try:
        p = subprocess.Popen(cmd, stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                             stderr=subprocess.DEVNULL, **kw)
    except OSError:
        return False, "找不到 soffice，跳过重算（公式仍然正确，Excel 打开时会自动算）"
    try:
        p.wait(timeout=timeout)
    except subprocess.TimeoutExpired:
        _kill_tree(p)
        time.sleep(0.2)
        for leftover in (produced, out_dir):
            try:
                (os.remove if leftover == produced else os.rmdir)(leftover)
            except OSError:
                pass
        return False, ("LibreOffice 重算超过 %g 秒没完成，已结束它，跳过重算"
                       "（公式仍然正确，Excel 打开时会自动算）" % timeout)
    if not os.path.exists(produced):
        return False, "LibreOffice 未产出文件，跳过重算"
    try:
        os.remove(path)
    except OSError:
        pass
    os.rename(produced, path)
    try:
        os.rmdir(out_dir)
    except OSError:
        pass
    return True, "已重算"


def scan_formula_errors(path):
    """重算后扫一遍，确保没有 #REF! / #NAME? 之类烂在文件里。"""
    ERR = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A", "#NULL!", "#NUM!", "Err:")
    from openpyxl import load_workbook
    wbv = load_workbook(path, data_only=True)
    wbf = load_workbook(path)
    n_formula = n_err = n_none = 0
    bad = []
    for sn in wbf.sheetnames:
        sf, sv = wbf[sn], wbv[sn]
        for row in sf.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n_formula += 1
                    v = sv[c.coordinate].value
                    if isinstance(v, str) and any(e in v for e in ERR):
                        n_err += 1
                        bad.append("%s!%s = %s" % (sn, c.coordinate, v))
                    elif v is None:
                        n_none += 1
    return n_formula, n_err, n_none, bad[:20]


# ------------------------------------------------------------------ grade banner
#
# Every monthly store workbook carries its grade (core/results.GRADES) on sheet ①
# so finance sees it without opening the results file. The runner stamps it after
# the builder saved and before LibreOffice recalculation (openpyxl drops cached
# formula values when it reopens and saves a workbook); builders never write it,
# so the golden tests see no banner. docs/REPORT_FAMILY.md "Grade banner (sheet ①)".
# The five grades are spelled here as in core/results.GRADE_CN — this module must
# stay free of project imports.

GRADE_LABEL = "完整度"
GRADE_CN = {"complete": "完整", "complete_warn": "完整·有提醒", "provisional": "暂定",
            "gaps": "有缺口", "failed": "失败"}
GRADE_TAB = {"complete": C_GREEN, "complete_warn": C_GREEN, "provisional": C_AMBER,
             "gaps": C_RED, "failed": C_RED}
GRADE_NOTE = "等级含义见 本月总览 的「说明」页；重出报表会产生新版本，以月份文件夹和本月总览为准。"
GRADE_ROWS_MAX = 8          # label rows + the note; a shorter block clears the rows below it


def stamp_grade(path, info):
    """Write the grade banner on the first sheet of the saved workbook at `path`.

    info: {"level": one of GRADE_CN, "reasons": [str], "version": str,
           "downloaded_at": "YYYY-MM-DD HH:MM" or "", "built_at": "YYYY-MM-DD HH:MM",
           "basis_cn": "账单基准" / "订单基准" / "", "final": bool}.

    The block starts at row 1 of the first empty column after the sheet's content
    (one spacer column; content in A–C -> E1:F8): a label column and a value
    column. No row or column is inserted and no existing cell is touched.
    Stamping again overwrites the same block (found by the GRADE_LABEL cell in row
    1). Returns the range written ("E1:F8"), or None when the workbook has no
    sheet."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    if not wb.worksheets:
        return None
    ws = wb.worksheets[0]
    level = str(info.get("level") or "")
    grade_cn = GRADE_CN.get(level, level or "—")
    reasons_text = "；".join(str(r) for r in (info.get("reasons") or ()) if r) or "—"

    col = None
    for c in range(1, ws.max_column + 1):
        if ws.cell(row=1, column=c).value == GRADE_LABEL:
            col = c
            break
    if col is None:
        col = ws.max_column + 2
    L, V = get_column_letter(col), get_column_letter(col + 1)

    rows = [(GRADE_LABEL, grade_cn), ("原因", reasons_text), ("报表版本", str(info.get("version") or "—")),
            ("数据下载于", str(info.get("downloaded_at") or "—")),
            ("报表生成于", str(info.get("built_at") or "—"))]
    if info.get("basis_cn"):
        rows.append(("口径", str(info["basis_cn"])))
    rows.append(("是否最终版本", "是（可以定稿的等级，且是最新的版本）" if info.get("final")
                 else "否（等级不够，只是未定稿）"))

    for i, (label, value) in enumerate(rows, 1):
        lc = ws.cell(row=i, column=col, value=label)
        lc.font = F(10, True)
        lc.fill = fill(C_SUB)
        lc.border = BOX
        lc.alignment = Alignment(vertical="center")
        vc = ws.cell(row=i, column=col + 1, value=value)
        vc.border = BOX
        vc.alignment = Alignment(vertical="center", wrap_text=True)
        if i == 1:
            if level in ("gaps", "failed"):
                vc.font = F(12, True, C_RED)
                vc.fill = fill(C_BAD)
            else:
                vc.font = F(12, True)
                vc.fill = fill(C_WARN if level == "provisional" else C_OK)
        else:
            vc.font = F(10)
    last = len(rows) + 1
    note = ws.cell(row=last, column=col, value=GRADE_NOTE)
    note.font = F(8, False, C_GREY)
    note.fill = PatternFill(fill_type=None)
    note.border = Border()
    note.alignment = Alignment(vertical="top")
    # the value cell beside the note, and the rows below a shorter block (an
    # earlier stamp of the same workbook had a 口径 row), are cleared
    stale = [(last, col + 1)] + [(r, c) for r in range(last + 1, GRADE_ROWS_MAX + 1) for c in (col, col + 1)]
    for r, c in stale:
        cell = ws.cell(row=r, column=c)
        cell.value = None
        cell.fill = PatternFill(fill_type=None)
        cell.border = Border()
    ws.row_dimensions[2].height = max(ws.row_dimensions[2].height or 15,
                                      est_height(reasons_text, per_line=40))
    ws.column_dimensions[L].width = max(ws.column_dimensions[L].width or 0, 14)
    ws.column_dimensions[V].width = max(ws.column_dimensions[V].width or 0, 60)

    ws.sheet_properties.tabColor = GRADE_TAB.get(level, C_GREY)
    wb.properties.subject = "%s：%s" % (GRADE_LABEL, grade_cn)
    wb.properties.keywords = str(info.get("version") or "")
    wb.properties.description = reasons_text[:255]
    wb.save(path)
    return "%s1:%s%d" % (L, V, last)


def stamp_grade_safely(path, info, log=print):
    """stamp_grade that never raises — a banner must not cost a report. Prints
    "[提示] 完整度标记没写上：…" and returns None on any failure."""
    try:
        return stamp_grade(path, info)
    except Exception as e:                                        # noqa: BLE001
        log("[提示] 完整度标记没写上：%s: %s" % (e.__class__.__name__, e))
        return None
