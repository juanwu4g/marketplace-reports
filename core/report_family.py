# -*- coding: utf-8 -*-
"""
The report family's sheet writers (docs/REPORT_FAMILY.md), shared by every
platform's workbook.

core/report_kit.py is the frame — styles, colours, number formats, the ten sheet
names, check records, file names. This module builds the sheets out of that
frame, so that a new platform supplies **data and rules** and gets the same
layout as the others:

  ① readme_sheet         sections of (label, text) rows, error banner
  ② pl_sheet             P&L from a row spec: data rows, input rows, subtotals
                         and computed rows whose formulas are declared as data
     ratio_block         "费用率" block under the P&L
     memo_block          "备查项" block under the P&L
  ③ bridge_sheet         bridges from sections of Line(label, value, note, style)
                         where style "residual" is the row that should be 0
  ④⑦⑧⑨ table_sheet + write_row   long tables from column specs (Col)
  ⑥ (any) not_applicable_sheet    "不适用：…" — the numbering stays stable
  ⑩ gaps_sheet           failed checks, missing files from run_meta.json, notes

Stores: a workbook is about one store or several. Every per-store value is a
function of one store object (a dict with "store" = its name, by convention; pass
`name=` to Family for another shape). With more than one store a 合计 column
follows the store columns; with one store it is left out — the same in every
platform's workbook.

Platform names do not appear here: labels in these functions are the family's
own Chinese words (合计, 说明, 校验 …); everything that names a platform's
files, terms or rules is passed in by the platform's report_build.

Example (a P&L of two data rows and a computed row):

    fam = Family(stores, money=K.money_format(0))
    spec = [PLRow("sales", "商品销售收入", "Sales", "d", ""),
            PLRow("fees", "减：平台费用", "Fees", "d", ""),
            PLRow("contrib", "经营贡献毛利", "", "f", "", formula=("sales", "fees"))]
    sh, r, at, note_col = fam.pl_sheet(wb, K.SHEET_NAMES[1], spec,
                                       value=lambda S, key: S["m"][key],
                                       title="合并损益表（2026 年 9 月）", subtitle="…")
"""
import collections

from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from core import report_kit as K
from core.report_kit import BOX, F, fill, wrap

# ------------------------------------------------------------------ specs

PLRow = collections.namedtuple("PLRow", "key label original kind note formula")
PLRow.__new__.__defaults__ = (None,)
PLRow.__doc__ = """One fixed row of the P&L (②).

key       stable id of the row; formulas and other sheets refer to it
label     Chinese line label
original  the platform's own term (grey, next to the label)
kind      "d"   data: value(S, key) per store, 合计 = sum
          "in"  input: yellow 0 per store for a person to fill, 合计 = SUM formula
          "sum" subtotal (light blue), formula
          "f"   computed row (grey), formula
note      explanation in the last column
formula   for "sum" / "f": a tuple of row keys added together, a key written
          "-key" is subtracted — ("revenue", "refund", "-cogs") -> =C8+C9-C27 —
          or Span(first, last) -> =SUM(C10:C19) over the rows between them
"""

Span = collections.namedtuple("Span", "first last")

Line = collections.namedtuple("Line", "label value note style")
Line.__new__.__defaults__ = ("", None)
Line.__doc__ = """One row of a bridge (③): label, value(S) per store, note, style.

style  None        plain
       "subtotal"  bold, light blue — the computed side of the bridge
       "residual"  bold, total colour — should be 0
       "pending"   amber — waiting for a document that is not out yet
"""

LINE_STYLES = {None: (False, None), "subtotal": (True, K.C_SUB),
               "residual": (True, K.C_TOTAL), "pending": (False, K.C_WARN)}

MONEY = "money"         # Col.fmt sentinel: the workbook's money format

Col = collections.namedtuple("Col", "header width fmt border fill font formula align")
Col.__new__.__defaults__ = (None, True, None, None, None, None)
Col.__doc__ = """One column of a long table (④ ⑦ ⑧ ⑨).

fmt      None, MONEY (the currency's money format) or a number format string
border   thin box around the cell (default True)
fill     a colour for every cell of the column (K.C_INPUT for input columns)
font     a Font for every cell of the column
formula  a template such as "=-O{r}*D{r}"; {r} becomes the row number and the
         value passed for this column is ignored
align    an Alignment for every cell of the column (e.g. wrap() for long text)
"""


# ------------------------------------------------------------------ small writers

def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def section(ws, r, text, ncols):
    """A blue band across ncols columns with white bold text. Returns the next row."""
    c = ws.cell(row=r, column=1, value=text)
    c.font = F(11, True, "FFFFFF")
    for j in range(1, ncols + 1):
        ws.cell(row=r, column=j).fill = fill(K.C_BAND)
    return r + 1


def sheet_title(ws, title, subtitle=None):
    """A1 title, A2 grey subtitle — the head of every long-table sheet."""
    ws["A1"] = title
    ws["A1"].font = F(14, True, K.C_NAVY)
    if subtitle is not None:
        ws["A2"] = subtitle
        ws["A2"].font = F(9, False, K.C_GREY)


def header_row(ws, row, cols, widths):
    """report_kit.Sheet.header on a sheet that has no store columns."""
    hdr = K.Sheet.__new__(K.Sheet)
    hdr.ws = ws
    K.Sheet.header(hdr, row, cols, widths)


def not_applicable_sheet(wb, name, title, text):
    """A sheet that does not apply to this platform: it stays, and says why."""
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = F(14, True, K.C_NAVY)
    ws["A3"] = text
    ws["A3"].font = F(11, False, K.C_GREY)
    ws.column_dimensions["A"].width = 120
    return ws


def readme_rows(ws, r, rows):
    """The body of ①: rows of (tag, label, text) from row r, label in column B and
    text in column C. tag "S" makes the row a section head (navy on light blue,
    B:C merged); any other tag is a plain row — ("", "", "") an empty one. Long
    texts get a taller row. Returns the next free row."""
    for tag, b, c in rows:
        if tag == "S":
            cell = ws.cell(row=r, column=2, value=b)
            cell.font = F(12, True, K.C_NAVY)
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
            cell.fill = fill(K.C_SUB)
        else:
            ws.cell(row=r, column=2, value=b).font = F(10, True)
            ws.cell(row=r, column=2).alignment = Alignment(vertical="top")
            c2 = ws.cell(row=r, column=3, value=c)
            c2.font = F(10)
            c2.alignment = wrap("top")
            if c and len(c) > 90:
                ws.row_dimensions[r].height = K.est_height(c, 78)
        r += 1
    return r


def meta_missing(meta):
    """The missing-form records of one download folder's run_meta.json (those
    written by core/download_diag through the platform's run_downloads)."""
    return [r for r in ((meta or {}).get("missing") or []) if r.get("form")]


def failed_errors(checks):
    return [c for c in checks if not c["ok"] and not c["skipped"] and c["severity"] == "error"]


def pl_formula(formula, col, rows_at):
    """'=C7+C8-C9' or '=SUM(C10:C19)' for one column of a sum / computed row."""
    if isinstance(formula, Span):
        return "=SUM(%s%d:%s%d)" % (col, rows_at[formula.first], col, rows_at[formula.last])
    parts = []
    for i, key in enumerate(formula):
        neg = key.startswith("-")
        ref = "%s%d" % (col, rows_at[key[1:] if neg else key])
        parts.append(("-" if neg else ("" if i == 0 else "+")) + ref)
    return "=" + "".join(parts)


# ------------------------------------------------------------------ the family

class Family(object):
    """The sheet writers for one workbook: its stores and its money format."""

    def __init__(self, stores, money=K.MNY, name=None):
        self.stores = list(stores)
        self.money = money
        self.name = name or (lambda S: S["store"])
        self.n = len(self.stores)
        self.has_total = self.n > 1

    # ---- columns and cells

    @property
    def names(self):
        return [self.name(S) for S in self.stores]

    def value_columns(self):
        """Store columns plus 合计 (when more than one store)."""
        return self.n + (1 if self.has_total else 0)

    def money_cell(self, c, v):
        c.value = v
        c.number_format = self.money
        c.border = BOX
        return c

    def values_row(self, ws, r, c0, vals, total=True, bold=False, color=None):
        """Per-store values from column c0, then their sum in the 合计 column.
        Returns the column after the last one written."""
        for i, v in enumerate(vals):
            c = self.money_cell(ws.cell(row=r, column=c0 + i), v)
            if bold:
                c.font = F(10, True)
            if color:
                c.fill = fill(color)
        if total and len(vals) > 1:
            c = self.money_cell(ws.cell(row=r, column=c0 + len(vals)), sum(vals))
            c.font = F(10, True)
            if color:
                c.fill = fill(color)
        return c0 + len(vals) + (1 if total and len(vals) > 1 else 0)

    def store_header(self, sh, r, first, extra=("说明",), widths=None,
                     first_widths=(34, 30), store_width=16, note_width=60):
        """Header row: first columns, one per store, 合计, then extra columns."""
        cols = list(first) + self.names + (["合计"] if self.has_total else []) + list(extra)
        sh.header(r, cols, widths or (list(first_widths) + [store_width] * self.value_columns()
                                      + [note_width]))
        return r + 1

    def per_store(self, fn, sep="；"):
        """fn(S) for a one-store workbook; "store fn(S)；store fn(S)" otherwise."""
        if self.n == 1:
            return fn(self.stores[0])
        return sep.join("%s %s" % (self.name(S), fn(S)) for S in self.stores)

    # ---- ① readme

    def readme_sheet(self, wb, name, title, subtitle, sections, warning=None,
                     widths=(4, 30, 118), first_row=4):
        """sections: [(section title, [(label, text), …]), …], an empty row
        between sections. warning: a red line under everything (e.g. "本期有 2 项
        error 级校验未通过 …") or None."""
        ws = wb.create_sheet(name)
        set_widths(ws, widths)
        ws["B1"] = title
        ws["B1"].font = F(16, True, K.C_NAVY)
        ws["B2"] = subtitle
        ws["B2"].font = F(9, False, K.C_GREY)
        rows = []
        for i, (head, lines) in enumerate(sections):
            if i:
                rows.append(("", "", ""))
            rows.append(("S", head, ""))
            rows.extend(("", b, c) for b, c in lines)
        r = readme_rows(ws, first_row, rows)
        if warning:
            ws.cell(row=r + 1, column=2, value="⚠ 注意").font = F(10, True, K.C_RED)
            c2 = ws.cell(row=r + 1, column=3, value=warning)
            c2.font = F(10, True, K.C_RED)
        return ws

    # ---- ② P&L

    def pl_sheet(self, wb, name, spec, value, title, subtitle,
                 first=("项目", "平台科目（原文）"), header_row=5, freeze="C6"):
        """The P&L: one row per PLRow, store columns from C, 合计, note column.

        value(S, key) gives a data row's amount for one store. Returns
        (sheet, next free row, {key: row number}, note column index).
        """
        n = self.n
        sh = K.Sheet(wb, name, first_store_col=len(first) + 1, nstores=n)
        ws = sh.ws
        sh.title(title, subtitle)
        r = self.store_header(sh, header_row, list(first))
        if freeze:
            ws.freeze_panes = freeze
        note_col = sh.c0 + self.value_columns()
        rows_at = {}
        for row in spec:
            kind = row.kind
            ws.cell(row=r, column=1, value=row.label).font = F(10, kind in ("sum", "f"))
            ws.cell(row=r, column=2, value=row.original).font = F(9, False, K.C_GREY)
            if kind == "d":
                self.values_row(ws, r, sh.c0, [value(S, row.key) for S in self.stores])
            elif kind == "in":
                for i in range(n):
                    c = self.money_cell(ws.cell(row=r, column=sh.c0 + i), 0)
                    c.fill = fill(K.C_INPUT)
                if self.has_total:
                    c = ws.cell(row=r, column=sh.c0 + n,
                                value="=SUM(%s%d:%s%d)" % (sh.scol(0), r, sh.scol(n - 1), r))
                    c.number_format = self.money
                    c.border = BOX
            elif kind in ("sum", "f"):
                if not row.formula:
                    raise ValueError("P&L row %r (%s) needs a formula" % (row.key, kind))
                for i in range(self.value_columns()):
                    col = get_column_letter(sh.c0 + i)
                    c = ws.cell(row=r, column=sh.c0 + i, value=pl_formula(row.formula, col, rows_at))
                    c.number_format = self.money
                    c.border = BOX
                    c.font = F(10, True)
                    c.fill = fill(K.C_SUB if kind == "sum" else K.C_CALC)
                ws.cell(row=r, column=1).fill = fill(K.C_SUB if kind == "sum" else K.C_CALC)
            else:
                raise ValueError("P&L row %r: unknown kind %r" % (row.key, kind))
            ws.cell(row=r, column=note_col, value=row.note).font = F(9, False, K.C_GREY)
            ws.cell(row=r, column=note_col).alignment = wrap()
            if row.key in rows_at:
                raise ValueError("P&L row key %r used twice" % row.key)
            rows_at[row.key] = r
            r += 1
        return sh, r, rows_at, note_col

    def ratio_block(self, sh, r, title, rows, note_col):
        """rows: [(label, numerator(S), base(S))]. Each store shows numerator /
        base (0 when base is 0); 合计 = Σ numerator / Σ base. Returns next row."""
        ws = sh.ws
        r = section(ws, r, title, note_col)
        for label, num, base in rows:
            ws.cell(row=r, column=1, value=label).font = F(10)
            for i, S in enumerate(self.stores):
                b = base(S)
                c = ws.cell(row=r, column=sh.c0 + i, value=num(S) / b if b else 0.0)
                c.number_format = K.PCT
                c.border = BOX
            if self.has_total:
                b = sum(base(S) for S in self.stores)
                v = sum(num(S) for S in self.stores)
                c = ws.cell(row=r, column=sh.c0 + self.n, value=v / b if b else 0.0)
                c.number_format = K.PCT
                c.border = BOX
            r += 1
        return r

    def memo_block(self, sh, r, title, rows, note_col):
        """rows: [(label, original term, value(S), note)] — figures that do not
        enter the P&L. Returns next row."""
        ws = sh.ws
        r = section(ws, r, title, note_col)
        for label, original, fn, note in rows:
            ws.cell(row=r, column=1, value=label).font = F(10)
            ws.cell(row=r, column=2, value=original).font = F(9, False, K.C_GREY)
            self.values_row(ws, r, sh.c0, [fn(S) for S in self.stores])
            ws.cell(row=r, column=note_col, value=note).font = F(9, False, K.C_GREY)
            r += 1
        return r

    # ---- ③ bridges

    def bridge_sheet(self, wb, name, title, subtitle, sections,
                     label_width=52, store_width=16, note_width=70, header_row=5):
        """sections: [(section title, [Line, …]), …], an empty row between them.
        Returns the sheet."""
        sh = K.Sheet(wb, name, first_store_col=2, nstores=self.n)
        ws = sh.ws
        sh.title(title, subtitle)
        note_col = sh.c0 + self.value_columns()
        sh.header(header_row, ["项目"] + self.names + (["合计"] if self.has_total else []) + ["说明"],
                  [label_width] + [store_width] * self.value_columns() + [note_width])
        r = header_row + 1
        for i, (head, lines) in enumerate(sections):
            if i:
                r += 1
            r = section(ws, r, head, note_col)
            for ln in lines:
                if ln.style not in LINE_STYLES:
                    raise ValueError("bridge line %r: unknown style %r" % (ln.label, ln.style))
                bold, color = LINE_STYLES[ln.style]
                ws.cell(row=r, column=1, value=ln.label).font = F(10, bold)
                self.values_row(ws, r, sh.c0, [ln.value(S) for S in self.stores],
                                bold=bold, color=color)
                ws.cell(row=r, column=note_col, value=ln.note).font = F(9, False, K.C_GREY)
                ws.cell(row=r, column=note_col).alignment = wrap()
                r += 1
        return ws

    # ---- ④ ⑦ ⑧ ⑨ long tables

    def table_sheet(self, wb, name, title, subtitle, cols, header_at=4):
        """Title, subtitle and a header row from Col specs. Returns the sheet;
        rows start at header_at + 1 (write them with write_row)."""
        ws = wb.create_sheet(name)
        sheet_title(ws, title, subtitle)
        header_row(ws, header_at, [c.header for c in cols], [c.width for c in cols])
        return ws

    def write_row(self, ws, r, cols, values, fill_color=None, alignment=None):
        """One table row by its Col specs; fill_color / alignment apply to every
        cell of the row. Returns the cells (for per-cell touches such as a red
        font on a flagged row)."""
        cells = []
        for j, (col, v) in enumerate(zip(cols, values), 1):
            if col.formula:
                v = col.formula.format(r=r)
            c = ws.cell(row=r, column=j, value=v)
            if col.fmt:
                c.number_format = self.money if col.fmt == MONEY else col.fmt
            if col.border:
                c.border = BOX
            if col.fill:
                c.fill = fill(col.fill)
            if col.font is not None:
                c.font = col.font
            if col.align is not None:
                c.alignment = col.align
            if fill_color:
                c.fill = fill(fill_color)
            if alignment is not None:
                c.alignment = alignment
            cells.append(c)
        return cells

    # ---- ⑩ gaps

    GAP_COLS = [Col("店铺", 22), Col("校验", 46), Col("级别", 9), Col("结果", 9), Col("说明", 100)]

    def gaps_sheet(self, wb, name, title, checks, missing=None, notes=()):
        """⑩: every check (failed errors first, then warnings, then passed and
        skipped), the missing files with cause and next step, then note sections.

        missing  [(store name, record)] — default: meta_missing(S["meta"]) of
                 every store
        notes    [(section title, [(label, text), …])] — e.g. tax notes
        """
        ws = self.table_sheet(wb, name, title,
                              "error = 数字在修好之前不能信；warn = 请看一眼，数字仍成立；"
                              "skipped = 不适用。空表示没有待办。", self.GAP_COLS)
        r = 5
        order = sorted(checks, key=lambda c: (c["ok"] or c["skipped"], c["severity"] != "error"))
        for c in order:
            res = "跳过" if c["skipped"] else ("通过" if c["ok"] else "未通过")
            color = K.C_OK if (c["ok"] or c["skipped"]) else (
                K.C_BAD if c["severity"] == "error" else K.C_WARN)
            self.write_row(ws, r, self.GAP_COLS,
                           (c["store"], c["name"], c["severity"], res, c["detail"]),
                           fill_color=color, alignment=wrap())
            r += 1
        r += 1
        r = section(ws, r, "没有下载到的文件及原因（来自下载目录的 run_meta.json）", 5)
        if missing is None:
            missing = [(self.name(S), rec) for S in self.stores
                       for rec in meta_missing(S.get("meta"))]
        for store, rec in missing:
            vals = (store, rec.get("form"), rec.get("cause"),
                    "影响报表" if rec.get("in_report", True) else "报表不读",
                    "%s　下一步：%s" % (rec.get("explain") or "", rec.get("action") or ""))
            for j, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=j, value=v)
                cell.border = BOX
                cell.alignment = wrap()
            r += 1
        if not missing:
            ws.cell(row=r, column=2, value="没有。应下的文件全部到齐。")
            r += 1
        for head, lines in notes:
            r += 1
            r = section(ws, r, head, 5)
            for label, text in lines:
                ws.cell(row=r, column=2, value=label).font = F(10, True)
                c = ws.cell(row=r, column=5, value=text)
                c.alignment = wrap()
                ws.row_dimensions[r].height = K.est_height(text, 95)
                r += 1
        return ws
